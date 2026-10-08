"""
scripts/import_lab_price_notes_20261007.py — lab fiyat notlarının yüklenmesi.

Sentetik veri dosyası + küçük kart/tedarikçi kurulumu ile: malzeme eşleştirme
(TR katlama, name_tr, material_key, alım belgesinin "lab tablosu karşılığı"
ipucu, belirsiz aday, ayrı form) · tedarikçi takma adları (PHARMATEM,
DOLAIN/Doalin, ULUDĞ / ULUDAĞ AGRO, VESER KİMYA, GÜLER KİMYA (YEŞİL KİL),
TATLI DİLİMLER, HAMMADDE SEPETİ) · öncelik (en yeni belge > lab tablosu,
ezilen değer notta) · birim uyumsuzluğu · elle (manual) ve Stok Son Durum
satırlarına dokunmama · idempotent ikinci çalıştırma · --create-suppliers ·
KDV alanları · kontrol Excel'inin sayfaları · onaylı Excel'den tercih
(grup/kart kapsamı, rank sırası, mevcut tercihi ezmeme) ve bitirilecek
(yalnız E) · Excel satır kayması koruması · CLI kuru çalıştırma · gerçek
veri dosyasıyla duman testi.  İnceleme düzeltmeleri (08.10.2026): üstü
çizili / fiyat listesi kaydı yazılmaz ve yarışmaz · ad ↔ ipucu çelişkisi
belirsiz · alternatif tercih lab satırının kartına · yalnız E/H (X, ✓
hata; Excel listesi girişi reddeder) · işlenmemiş "Doğru değer / Doğru
IMS kartı" hücresi yazımı durdurur · yeni firma yalnız E ile açılır ·
tercih ↔ bitirilecek çelişkisi hata · Tercih önerisinde KDV etiketi ve
"Uludağ ucuz" uyarısı · 2 sarıda ucuz olan üstte.
    MINERVA_TEST_DB=minerva_test2 .venv/bin/pytest tests/test_import_lab_price_notes.py -q
"""
import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from database import (AdminAuditLog, Item, MaterialGroup, MaterialSupplierPref, Supplier,
                      SupplierPrice)

_SPEC = importlib.util.spec_from_file_location(
    "_import_lab_notes",
    Path(__file__).resolve().parent.parent / "scripts" / "import_lab_price_notes_20261007.py")
lab = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(lab)


def rec(rid, malzeme, tedarikci, fiyat, *, tur="lab_tablo", tarih="2026-06-23", cur="EUR",
        unit="kg", pkg=25.0, kdv=None, oran=None, secili=False, guven="kesin", note="", sayfa=1,
        birim="g"):
    return {"id": rid, "kaynak_sayfa": sayfa, "tur": tur, "tarih": tarih, "malzeme": malzeme,
            "birim_kart": birim if tur == "lab_tablo" else None, "tedarikci": tedarikci,
            "fiyat": fiyat, "para_birimi": cur, "fiyat_birimi": unit, "ambalaj_kg": pkg,
            "kdv_dahil": kdv, "kdv_orani": oran, "secili": secili, "guven": guven, "not": note,
            "ihtiyac": None}


def data(*recs, **extra):
    return {"kaynak": "test", "kapsam": "test", "kurallar": [], "kayitlar": list(recs), **extra}


def _items(db, *specs):
    out = {}
    for spec in specs:
        name, unit = spec[0], spec[1]
        kw = spec[2] if len(spec) > 2 else {}
        it = Item(name=name, unit=unit, category=kw.pop("category", "Hammadde"), **kw)
        db.add(it)
        db.flush()
        out[name] = it
    return out


def _sups(db, *names):
    out = {}
    for n in names:
        s = Supplier(name=n)
        db.add(s)
        db.flush()
        out[n] = s
    return out


def _prices(db):
    db.expire_all()
    return db.query(SupplierPrice).order_by(SupplierPrice.id).all()


def _lab_prices(db):
    return [p for p in _prices(db) if (p.source_label or "").startswith(lab.LABEL_PREFIX)]


def _by_id(p):
    return {r["id"]: r for r in p["records"]}


# ─── Eşleştirme ─────────────────────────────────────────────────────────────

def test_material_match_tr_fold_name_tr_material_key_and_ambiguous(db_session: Session):
    its = _items(db_session, ("COCO BETAİNE%30", "g"), ("SHEA YAĞI", "g", {"name_tr": "KARİTE"}),
                 ("JOJOBA YAĞI", "g"), ("BADEM YAĞI", "g"), ("BADEM YAGI (NUMUNE)", "g"),
                 ("ÜRÜN X", "adet", {"category": "Bitmiş Ürün"}))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    p = lab.plan(db_session, data(
        rec("A1-01-S1", "coco betaine%30", "UMAYCHEM", 1.0),            # TR katlama + küçük harf
        rec("A1-02-S1", "KARİTE", "UMAYCHEM", 2.0),                      # name_tr
        rec("A1-03-S1", "JOJOBA OIL", "UMAYCHEM", 3.0),                  # material_key (OIL = YAĞI)
        rec("A1-04-S1", "ALMOND OIL", "UMAYCHEM", 4.0),                  # 2 kart aynı anahtar → belirsiz
        rec("A1-05-S1", "ÜRÜN X", "UMAYCHEM", 5.0),                      # bitmiş ürün sayılmaz
        rec("A1-06-S1", "COCO BETAİNE %31", "UMAYCHEM", 6.0),            # bulunamaz → benzer öneri
    ))
    r = _by_id(p)
    assert r["A1-01-S1"]["item"]["id"] == its["COCO BETAİNE%30"].id and r["A1-01-S1"]["mat"]["how"] == "ad"
    assert r["A1-02-S1"]["item"]["id"] == its["SHEA YAĞI"].id
    assert r["A1-03-S1"]["item"]["id"] == its["JOJOBA YAĞI"].id
    assert r["A1-03-S1"]["mat"]["how"] == "material_key"
    assert r["A1-04-S1"]["status"] == "belirsiz"
    assert set(r["A1-04-S1"]["mat"]["candidates"]) == {its["BADEM YAĞI"].id, its["BADEM YAGI (NUMUNE)"].id}
    assert r["A1-05-S1"]["status"] == "eslesmedi"
    assert r["A1-06-S1"]["status"] == "eslesmedi"
    assert its["COCO BETAİNE%30"].id in r["A1-06-S1"]["mat"]["suggestions"]
    assert p["counts"]["malzeme_belirsiz"] == 1 and p["counts"]["malzeme_eslesmeyen"] == 2


def test_document_uses_lab_table_hint_and_strips_package_suffix(db_session: Session):
    its = _items(db_session, ("ASPİR YAĞI", "g"), ("SUSAM YAĞI", "g"), ("PAPATYA HİDRASOLÜ", "g"),
                 ("TATLI BADEM YAGI", "g"), ("BADEM YAĞI", "g"))
    _sups(db_session, "NATURALYA", "KRK GIDA", "ULUDAĞ HERBAL")
    db_session.commit()
    p = lab.plan(db_session, data(
        rec("A1-07-S1", "ASPİR YAĞI", "NATURALYA", 12.9),
        rec("A1-14-S1", "PAPATYA HİDRASOLÜ", "ULUDAĞ HERBAL", 4.36),
        rec("A1-26-S1", "TATLI BADEM YAGI", "NATURALYA", 12.9),
        rec("A1-41-S1", "BADEM YAĞI", "NATURALYA", 15.17),
        rec("C-s19-2", "SAFFLOWER OIL", "NATURALYA KİMYA", 15.0, tur="proforma", tarih="2026-07-01",
            cur="USD", note="aspir yağı; lab tablosu karşılığı: ASPİR YAĞI (A1-07)"),
        rec("C-s13-2", "SUSAM YAĞI (KG)", "KRK GIDA", 340.0, tur="proforma", tarih="2026-06-25",
            cur="TRY", pkg=5.0),
        rec("C-s9-3", "PAPATYA HİDROSOLÜ 1LT", "ULUDAĞ AGRO", 208.33, tur="fatura", tarih="2026-06-26",
            cur="TRY", unit="l", pkg=1.0,
            note="lab tablosu karşılığı: PAPATYA HİDRASOLÜ (A1-14, g birimli kart)"),
        # İki ipucu farklı karta → material_key ayırır (BADEM YAĞI (TATLI) = TATLI BADEM YAGI)
        rec("C-s13-1", "BADEM YAĞI (TATLI) 1 LT.", "KRK GIDA", 750.0, tur="proforma",
            tarih="2026-06-25", cur="TRY", unit="l", pkg=1.0,
            note="lab tablosu karşılığı: TATLI BADEM YAGI (A1-26) / BADEM YAĞI (A1-41)"),
    ))
    r = _by_id(p)
    assert r["C-s19-2"]["item"]["id"] == its["ASPİR YAĞI"].id and r["C-s19-2"]["mat"]["how"] == "lab tablosu ipucu"
    assert r["C-s13-2"]["item"]["id"] == its["SUSAM YAĞI"].id and r["C-s13-2"]["mat"]["how"] == "ad"
    assert r["C-s9-3"]["item"]["id"] == its["PAPATYA HİDRASOLÜ"].id
    assert r["C-s13-1"]["item"]["id"] == its["TATLI BADEM YAGI"].id
    assert r["C-s13-1"]["mat"]["how"] == "ipucu + material_key"
    assert lab.clean_doc_material("XANTHAN GUM (25 KG) (MEIUHA)") == "XANTHAN GUM (MEIUHA)"
    assert lab.clean_doc_material("MİSK ADAÇAYI YAĞI LT") == "MİSK ADAÇAYI YAĞI"


def test_separate_form_records_never_auto_matched(db_session: Session):
    _items(db_session, ("ALEOVERA EKSTRAKTI", "ml"))
    _sups(db_session, "NATURALYA")
    db_session.commit()
    p = lab.plan(db_session, data(
        rec("A1-08-S2", "ALEOVERA EKSTRAKTI", "NATURALYA", 10.32, secili=True),
        rec("A1-08-S3", "ALEOVERA EKSTRAKTI", "Naturalya", 39.6, note="toz hali"),
        rec("C-s8-1", "ALOEVERA EXTRACT PE", "NATURALYA KİMYA", 43.0, tur="proforma", tarih="2026-07-16",
            cur="USD", pkg=1.0, note="lab tablosu karşılığı: ALEOVERA EKSTRAKTI (A1-08, sıvı kart)"),
    ))
    r = _by_id(p)
    assert r["A1-08-S3"]["status"] == "ayri_form" and r["C-s8-1"]["status"] == "ayri_form"
    # Sıvı kartın fiyatı tozun proformasıyla EZİLMEZ
    assert r["A1-08-S2"]["status"] == "yazilacak" and not r["A1-08-S2"]["overridden"]


def test_supplier_aliases_and_new_suppliers(db_session: Session):
    _items(db_session, ("X", "g"))
    s = _sups(db_session, "PHARMATERM", "DOALİNN", "ULUDAĞ HERBAL", "VESER KİMYEVİ", "GÜLER KİMYA",
              "TATLIDİLİMLER", "HAMMADDESEPETİ", "DOGASA", "YİĞİTOĞLU KİMYA")
    db_session.commit()
    names = ["PHARMATEM", "PHARMATERM İLAÇ", "Pharmaterm", "DOLAIN", "DOALIN", "Doalin", "ULUDĞ HERBAL",
             "ULUDAĞ AGRO", "VESER KİMYA", "GÜLER KİMYA (YEŞİL KİL)", "TATLI DİLİMLER", "HAMMADDE SEPETİ",
             "DOĞASA", "YİĞİTOĞLU", "BEFCHEM", "ATAMAN KİMYA", "MARKET"]
    p = lab.plan(db_session, data(*[rec(f"A1-{i:02d}-S1", "X", n, 1.0 + i) for i, n in enumerate(names, 1)]))
    got = {r["supplier"]: r["sup"] for r in p["records"]}
    want = {"PHARMATEM": "PHARMATERM", "PHARMATERM İLAÇ": "PHARMATERM", "Pharmaterm": "PHARMATERM",
            "DOLAIN": "DOALİNN", "DOALIN": "DOALİNN", "Doalin": "DOALİNN", "ULUDĞ HERBAL": "ULUDAĞ HERBAL",
            "ULUDAĞ AGRO": "ULUDAĞ HERBAL", "VESER KİMYA": "VESER KİMYEVİ",
            "GÜLER KİMYA (YEŞİL KİL)": "GÜLER KİMYA", "TATLI DİLİMLER": "TATLIDİLİMLER",
            "HAMMADDE SEPETİ": "HAMMADDESEPETİ", "DOĞASA": "DOGASA", "YİĞİTOĞLU": "YİĞİTOĞLU KİMYA"}
    for doc, card in want.items():
        assert got[doc]["kind"] == "id" and got[doc]["id"] == s[card].id, doc
    assert {n["name"] for n in p["new_suppliers"]} == {"BEFCHEM", "ATAMAN KİMYA", "MARKET"}
    for n in ("BEFCHEM", "ATAMAN KİMYA", "MARKET"):
        assert got[n]["kind"] == "new"


def test_inactive_duplicate_and_merged_card(db_session: Session):
    _items(db_session, ("X", "g"))
    s = _sups(db_session, "TATLİDİLİMLER", "SURYA KİMYA", "SURYA ESKİ")
    s["TATLİDİLİMLER"].is_active = False                       # pasif mükerrer — eşlenmez
    s["SURYA ESKİ"].is_active = False
    s["SURYA ESKİ"].merged_into_id = s["SURYA KİMYA"].id
    db_session.commit()
    p = lab.plan(db_session, data(rec("A1-01-S1", "X", "TATLI DİLİMLER", 1.0),
                                  rec("A1-02-S1", "X", "SURYA ESKİ", 2.0)))
    r = _by_id(p)
    # Yalnız pasif kartı olan firma: yeni kart AÇILMAZ, satır yazılmaz
    assert r["A1-01-S1"]["sup"]["kind"] == "inactive" and r["A1-01-S1"]["status"] == "pasif_tedarikci"
    assert p["new_suppliers"] == []
    assert r["A1-02-S1"]["sup"]["id"] == s["SURYA KİMYA"].id and r["A1-02-S1"]["status"] == "yazilacak"
    # Aktif ikiz varsa pasif mükerrer önemsiz: aktif karta bağlanır
    twin = Supplier(name="TATLIDİLİMLER")
    db_session.add(twin)
    db_session.commit()
    r = _by_id(lab.plan(db_session, data(rec("A1-01-S1", "X", "TATLI DİLİMLER", 1.0))))
    assert r["A1-01-S1"]["sup"]["id"] == twin.id
    p = lab.apply(db_session, data(rec("A1-01-S1", "X", "TATLI DİLİMLER", 1.0)), create_suppliers=True)
    assert p["created_suppliers"] == [] and p["written"] == 1


# ─── Öncelik, birim, elle fiyat ─────────────────────────────────────────────

def test_newest_document_wins_and_overridden_values_in_note(db_session: Session):
    its = _items(db_session, ("LAURYL GLUCOSİDE", "g"))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    p = lab.apply(db_session, data(
        rec("A1-05-S1", "LAURYL GLUCOSİDE", "UMAYCHEM", 1.93, secili=True, pkg=220.0),
        rec("C-s12-1", "Lauryl Glucoside", "UMAYCHEM", 1.90, tur="proforma", tarih="2026-06-26",
            cur="USD", pkg=220.0, kdv=False, oran=20.0, sayfa=12),
        rec("C-s99-1", "Lauryl Glucoside", "UMAYCHEM", 2.10, tur="siparis", tarih="2026-07-10",
            cur="USD", pkg=25.0, kdv=False, oran=20.0, sayfa=15),
    ))
    r = _by_id(p)
    assert r["C-s99-1"]["status"] == "yazilacak"
    assert r["C-s12-1"]["status"] == "ezildi" and r["A1-05-S1"]["status"] == "ezildi"
    rows = _lab_prices(db_session)
    assert len(rows) == 1
    sp = rows[0]
    assert sp.item_id == its["LAURYL GLUCOSİDE"].id and sp.unit_price == 2.10 and sp.currency == "USD"
    assert sp.source == "siparis" and sp.source_label == f"{lab.LABEL_PREFIX} — s.15"
    assert str(sp.quoted_at) == "2026-07-10" and sp.package_size == 25.0
    assert sp.vat_included is False and sp.vat_rate == 20.0 and sp.created_by == lab.ACTOR
    assert "C-s12-1" in sp.note and "A1-05-S1" in sp.note and "Ezilen" in sp.note


def test_lab_rows_vat_and_selected_note(db_session: Session):
    _items(db_session, ("VANİLYA UÇUCU YAĞI", "ml"))
    _sups(db_session, "PHARMATERM", "ULUDAĞ HERBAL")
    db_session.commit()
    lab.apply(db_session, data(
        rec("A2-40-S1", "VANİLYA UÇUCU YAĞI", "ULUDAĞ HERBAL", 47.58, birim="ml"),
        rec("A2-40-S2", "VANİLYA UÇUCU YAĞI", "PHARMATERM İLAÇ", 82.56, kdv=True, oran=20.0,
            secili=True, birim="ml"),
    ))
    rows = {p.supplier_name: p for p in _lab_prices(db_session)}
    ph = rows["PHARMATERM"]
    assert ph.vat_included is True and ph.vat_rate == 20.0 and ph.source == "lab_notu"
    assert ph.price_unit == "kg" and ph.currency == "EUR" and "SARI" in ph.note
    assert rows["ULUDAĞ HERBAL"].vat_included is None and rows["ULUDAĞ HERBAL"].vat_rate is None


def test_unit_mismatch_listed_and_not_competing(db_session: Session):
    its = _items(db_session, ("POMPA", "adet"), ("ARGAN YAĞI", "g"))
    _sups(db_session, "NATURALYA")
    db_session.commit()
    p = lab.apply(db_session, data(
        rec("A1-01-S1", "POMPA", "NATURALYA", 1.0),                                   # kg ↔ adet
        rec("A1-12-S1", "ARGAN YAĞI", "NATURALYA", 25.8),
        rec("C-s1-1", "ARGAN YAĞI", "NATURALYA", 30.0, tur="proforma", tarih="2026-07-01",
            unit="adet", cur="USD"),                                                   # uyumsuz → yarışmaz
    ))
    r = _by_id(p)
    assert r["A1-01-S1"]["status"] == "uyumsuz" and r["C-s1-1"]["status"] == "uyumsuz"
    assert r["A1-12-S1"]["status"] == "yazilacak"
    rows = _lab_prices(db_session)
    assert [(x.item_id, x.unit_price) for x in rows] == [(its["ARGAN YAĞI"].id, 25.8)]


def test_manual_and_stok_son_durum_rows_untouched(db_session: Session):
    its = _items(db_session, ("ARGAN YAĞI", "g"), ("JOJOBA YAĞI", "g"))
    s = _sups(db_session, "NATURALYA", "SURYA KİMYA")
    man = SupplierPrice(item_id=its["ARGAN YAĞI"].id, supplier_id=s["NATURALYA"].id, supplier_name="NATURALYA",
                        unit_price=24.0, currency="EUR", price_unit="kg", source="manual",
                        created_by="Songül", domain="cosmetics")
    ssd = SupplierPrice(item_id=its["JOJOBA YAĞI"].id, supplier_id=None, supplier_name="NATURALYA",
                        unit_price=14.0, currency="USD", price_unit="kg", source="stok_son_durum",
                        source_label="Stok Son Durum — Eylül", domain="cosmetics")
    db_session.add_all([man, ssd])
    db_session.commit()
    snap = [(x.id, x.unit_price, x.source, x.source_label, x.note) for x in _prices(db_session)]
    d = data(rec("A1-12-S3", "ARGAN YAĞI", "NATURALYA", 25.8, secili=True),
             rec("A1-12-S2", "ARGAN YAĞI", "SURYA KİMYA", 26.0),
             rec("A1-17-S1", "JOJOBA YAĞI", "NATURALYA", 12.9))
    p = lab.apply(db_session, d)
    r = _by_id(p)
    assert r["A1-12-S3"]["status"] == "elle_var" and "24" in "; ".join(r["A1-12-S3"]["flags"])
    assert r["A1-17-S1"]["status"] == "yazilacak"
    assert any("Stok Son Durum" in f for f in r["A1-17-S1"]["flags"])
    after = {x.id: (x.id, x.unit_price, x.source, x.source_label, x.note) for x in _prices(db_session)}
    for row in snap:
        assert after[row[0]] == row                                    # elle + Stok Son Durum aynen
    labs = _lab_prices(db_session)
    assert {(x.item_id, x.supplier_name) for x in labs} == {(its["ARGAN YAĞI"].id, "SURYA KİMYA"),
                                                           (its["JOJOBA YAĞI"].id, "NATURALYA")}
    # İkinci çalıştırma: elle satır hâlâ kazanır, Stok Son Durum hâlâ yerinde
    lab.apply(db_session, d)
    after2 = {x.id: (x.id, x.unit_price, x.source, x.source_label, x.note) for x in _prices(db_session)}
    for row in snap:
        assert after2[row[0]] == row


def test_second_run_is_idempotent(db_session: Session):
    _items(db_session, ("GLİSERİN", "g"), ("SHEA BUTTER", "g"))
    _sups(db_session, "UMAYCHEM", "BEFCHEM")
    db_session.commit()
    d = data(rec("A1-09-S1", "GLİSERİN", "UMAYCHEM", 4.52, secili=True),
             rec("A1-18-S1", "SHEA BUTTER", "BEFCHEM", 8.0, secili=True))
    first = lab.apply(db_session, d)
    rows1 = [(x.item_id, x.supplier_id, x.unit_price, x.source, x.source_label, x.note, x.package_size)
             for x in _lab_prices(db_session)]
    second = lab.apply(db_session, d)
    rows2 = [(x.item_id, x.supplier_id, x.unit_price, x.source, x.source_label, x.note, x.package_size)
             for x in _lab_prices(db_session)]
    assert first["written"] == 2 and first["deleted"] == 0
    assert second["written"] == 2 and second["deleted"] == 2
    assert sorted(rows1) == sorted(rows2) and len(_prices(db_session)) == 2
    db_session.expire_all()
    audits = db_session.query(AdminAuditLog).filter(AdminAuditLog.action == lab.AUDIT_IMPORT).all()
    assert len(audits) == 2 and audits[0].actor_id is None
    det = json.loads(audits[1].details)
    assert det["actor"] == lab.ACTOR and det["written"] == 2 and det["deleted_previous"] == 2


def test_new_supplier_rows_skipped_without_flag_created_with_flag(db_session: Session, tmp_path):
    its = _items(db_session, ("D-PANTHENOL", "g"), ("GLİSERİN", "g"))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    d = data(rec("A1-20-S1", "D-PANTHENOL", "ATAMAN KİMYA", 11.63, secili=True),
             rec("A1-09-S1", "GLİSERİN", "UMAYCHEM", 4.52),
             rec("A1-21-S1", "GLİSERİN", "MARKET", 0.29),
             rec("A1-22-S1", "YOK BÖYLE KART", "KİMYACINIZ", 5.46))
    p = lab.apply(db_session, d)
    assert _by_id(p)["A1-20-S1"]["status"] == "yeni_tedarikci"
    assert [x.item_id for x in _lab_prices(db_session)] == [its["GLİSERİN"].id]
    db_session.expire_all()
    assert db_session.query(Supplier).filter(Supplier.name == "ATAMAN KİMYA").count() == 0
    # Onaylı Excel'siz kart açılmaz
    p = lab.apply(db_session, d, create_suppliers=True)
    assert p["applied"] is False and any("onaylı Excel" in e for e in p["errors"])
    # 'Yeni tedarikçiler': yalnız E açılır; H (MARKET) açılmaz; yazılabilir
    # satırı olmayan (KİMYACINIZ — kartı yok) "—", açılmaz
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    ws = load_workbook(path)["Yeni tedarikçiler"]
    head = [c.value for c in ws[1]]
    new_rows = {r[head.index("Açılacak ad")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert new_rows["KİMYACINIZ"][head.index(lab.NEW_COL)] == "—"
    assert ws.data_validations.dataValidation[0].showErrorMessage
    ids = {r[head.index("Açılacak ad")]: r[head.index("Kayıt")] for r in new_rows.values()}
    _fill(path, "Yeni tedarikçiler", lab.NEW_COL, {ids["ATAMAN KİMYA"]: "E", ids["MARKET"]: "H"})
    appr = lab.read_approvals(path)
    assert not appr["errors"] and appr["new"][ids["ATAMAN KİMYA"]]["ok"] is True
    p = lab.apply(db_session, d, create_suppliers=True, approvals=appr)
    assert p["applied"]
    db_session.expire_all()
    assert db_session.query(Supplier).filter(Supplier.name.in_(("MARKET", "KİMYACINIZ"))).count() == 0
    at = db_session.query(Supplier).filter(Supplier.name == "ATAMAN KİMYA").one()
    assert at.domain == "cosmetics" and at.is_active and lab.SUPPLIER_NOTE in (at.notes or "")
    assert p["created_suppliers"] == [{"id": at.id, "name": "ATAMAN KİMYA"}]
    rows = {x.item_id: x for x in _lab_prices(db_session)}
    assert rows[its["D-PANTHENOL"].id].supplier_id == at.id
    # Kart artık var: üçüncü çalıştırma yeni kart açmaz, aynı sonucu yazar
    p = lab.apply(db_session, d, create_suppliers=True, approvals=appr)
    assert p["created_suppliers"] == [] and p["written"] == 2
    assert db_session.query(Supplier).filter(Supplier.name == "ATAMAN KİMYA").count() == 1


def test_bad_manual_override_raises(db_session: Session):
    _items(db_session, ("X", "g"))
    db_session.commit()
    with pytest.raises(ValueError):
        lab.plan(db_session, data(rec("A1-01-S1", "Y", "Z", 1.0), elle_eslesme={"kart": {"A1-01": 999999}}))


def test_manual_override_maps_row(db_session: Session):
    its = _items(db_session, ("NAR ÇEKİRDEĞİ YAĞI", "g"))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    p = lab.plan(db_session, data(rec("A2-01-S2", "NAR ÇEKİDEĞİ YAGI", "UMAYCHEM", 18.57),
                                  elle_eslesme={"kart": {"A2-01": its["NAR ÇEKİRDEĞİ YAĞI"].id}}))
    r = p["records"][0]
    assert r["item"]["id"] == its["NAR ÇEKİRDEĞİ YAĞI"].id and r["mat"]["how"] == "elle eşleme"


# ─── Kontrol Excel'i + onay uygulaması ──────────────────────────────────────

def _pref_setup(db):
    its = _items(db, ("ALEOVERA EKSTRAKTI", "ml"), ("AKTİF KÖMÜR", "g"), ("GLİSERİN", "g"),
                 ("GLYCERIN", "g"), ("SQUALANE", "g"))
    grp = MaterialGroup(name="GLİSERİN", domain="cosmetics")
    db.add(grp)
    db.flush()
    its["GLİSERİN"].material_group_id = grp.id
    its["GLYCERIN"].material_group_id = grp.id
    s = _sups(db, "SURYA KİMYA", "NATURALYA", "TATLIDİLİMLER", "BEFCHEM", "UMAYCHEM", "HAMMADDESEPETİ",
              "ROSECE")
    # Mevcut tercih (elle): squalane'de UMAYCHEM 1. sıra — ezilmemeli
    db.add(MaterialSupplierPref(domain="cosmetics", item_id=its["SQUALANE"].id, supplier_id=s["UMAYCHEM"].id,
                                preference="preferred", rank=1, created_by="Songül"))
    db.commit()
    d = data(
        rec("A1-08-S1", "ALEOVERA EKSTRAKTI", "SURYA KİMYA", 16.34, secili=True, birim="ml"),
        rec("A1-08-S2", "ALEOVERA EKSTRAKTI", "NATURALYA", 10.32, secili=True, birim="ml"),
        rec("A1-09-S1", "GLİSERİN", "UMAYCHEM", 4.52, secili=True),
        rec("A2-13-S1", "AKTİF KÖMÜR", "TATLI DİLİMLER", 23.55, pkg=1.0),
        rec("A2-13-S2", "AKTİF KÖMÜR", "BEFCHEM", 38.0, pkg=20.0, secili=True),
        rec("A1-34-S1", "SQUALANE", "HAMMADDE SEPETİ", 120.37, pkg=1.0),
        rec("A1-34-S2", "SQUALANE", "BEFCHEM", 45.0, pkg=5.0, secili=True),
        rec("A1-31-S1", "SQUALANE", "ROSECE", 28.67, pkg=1.0),
        rec("A1-35-S1", "SQUALANE", "SURYA KİMYA", 124.0, pkg=1.0),
    )
    return its, s, grp, d


def _fill(path, sheet, col, values):
    wb = load_workbook(path)
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    ik, ic = head.index("Kayıt") + 1, head.index(col) + 1
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, ik).value
        if rid in values:
            ws.cell(r, ic, values[rid])
    wb.save(path)


def test_xlsx_sheets_and_columns(db_session: Session, tmp_path):
    _its, _s, _grp, d = _pref_setup(db_session)
    p = lab.plan(db_session, d)
    path = lab.build_xlsx(p, tmp_path / "kontrol.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == list(lab.SHEETS)
    fiy = wb["Fiyatlar"]
    assert fiy.max_row == 1 + len(d["kayitlar"]) and fiy.freeze_panes == "B2" and fiy.auto_filter.ref
    pref = wb["Tercih önerisi"]
    head = [c.value for c in pref[1]]
    assert lab.APPROVE_COL in head
    rows = [[c.value for c in r] for r in pref.iter_rows(min_row=2)]
    ids = [r[0] for r in rows]
    # sarılar (2 sarıda ucuz olan — NATURALYA 10,32 — üstte) + nottan
    # alternatif (A2-13-S1, aynı malzemenin sarısından sonra)
    assert ids == ["A1-08-S2", "A1-08-S1", "A1-09-S1", "A2-13-S2", "A2-13-S1", "A1-34-S2"]
    warn = {r[0]: r[head.index("Uyarı")] for r in rows}
    assert "2 sarı seçim" in warn["A1-08-S1"] and "tatlı dilimlerden" in warn["A2-13-S2"]
    assert "yalnız NATURALYA" in warn["A1-08-S1"]
    assert pref.data_validations.dataValidation[0].showErrorMessage
    assert all(r[head.index(lab.APPROVE_COL)] is None for r in rows)
    cur = {r[0]: r[head.index("Mevcut tercih")] for r in rows}
    assert "UMAYCHEM (tercih 1)" in cur["A1-34-S2"]
    small = wb["Küçük satıcı önerisi"]
    sh = [c.value for c in small[1]]
    srows = {r[sh.index("Öneri türü")]: r for r in small.iter_rows(min_row=2, values_only=True)}
    kinds = [r[sh.index("Öneri türü")] for r in small.iter_rows(min_row=2, values_only=True)]
    assert kinds.count("küçük satıcı") == 3 and "bilgi" in kinds           # Surya: karışık → bilgi
    assert srows["bilgi"][sh.index(lab.PHASE_COL)] == "—"
    for name in ("Eşleşmeyenler", "Belirsiz okumalar", "Yeni tedarikçiler", "Nasıl doldurulur"):
        assert wb[name].max_row >= 1


def test_apply_prefs_e_only_group_scope_rank_and_existing_kept(db_session: Session, tmp_path):
    its, s, grp, d = _pref_setup(db_session)
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    _fill(path, "Tercih önerisi", lab.APPROVE_COL,
          {"A1-08-S1": "E", "A1-08-S2": "e", "A1-09-S1": "Evet", "A2-13-S2": "H", "A2-13-S1": "E",
           "A1-34-S2": "E"})
    small_ids = {x["sup"]["name"]: x["id"] for x in lab.plan(db_session, d)["small"]}
    _fill(path, "Küçük satıcı önerisi", lab.PHASE_COL,
          {small_ids["HAMMADDESEPETİ"]: "E", small_ids["ROSECE"]: "H"})
    appr = lab.read_approvals(path)
    assert not appr["errors"]
    # Kuru: hiçbir şey yazılmaz
    pp = lab.plan_prefs(db_session, appr, d)
    assert pp["counts"] == {"tercih_yazilacak": 5, "tercih_atlanan": 0, "bitirilecek": 1}
    db_session.expire_all()
    assert db_session.query(MaterialSupplierPref).count() == 1
    pp = lab.apply_prefs(db_session, appr, d)
    assert pp["applied"]
    db_session.expire_all()
    prefs = {(p.item_id, p.material_group_id, p.supplier_id): p
             for p in db_session.query(MaterialSupplierPref).all()}
    alo = its["ALEOVERA EKSTRAKTI"].id
    assert prefs[(alo, None, s["NATURALYA"].id)].rank == 1             # Excel sırası = rank (ucuz üstte)
    assert prefs[(alo, None, s["SURYA KİMYA"].id)].rank == 2
    assert prefs[(None, grp.id, s["UMAYCHEM"].id)].rank == 1           # grubu olan kart → gruba
    assert prefs[(its["AKTİF KÖMÜR"].id, None, s["TATLIDİLİMLER"].id)].rank == 1
    assert (its["AKTİF KÖMÜR"].id, None, s["BEFCHEM"].id) not in prefs     # H
    sq = its["SQUALANE"].id
    assert prefs[(sq, None, s["UMAYCHEM"].id)].rank == 1 and prefs[(sq, None, s["UMAYCHEM"].id)].created_by == "Songül"
    assert prefs[(sq, None, s["BEFCHEM"].id)].rank == 2                # mevcut korunur, arkasına
    assert all(p.preference == "preferred" for p in prefs.values())
    sups = {x.name: x for x in db_session.query(Supplier).all()}
    assert sups["HAMMADDESEPETİ"].purchase_status == "phase_out"
    assert sups["HAMMADDESEPETİ"].status_reason == lab.PHASE_OUT_REASON
    assert sups["ROSECE"].purchase_status == "normal"                   # H → dokunulmaz
    assert sups["TATLIDİLİMLER"].purchase_status == "normal"            # boş → dokunulmaz
    acts = [a.action for a in db_session.query(AdminAuditLog).all()]
    assert acts.count("material_pref.create") == 5 and acts.count("supplier.status") == 1
    assert lab.AUDIT_PREFS in acts
    # İkinci uygulama: zaten var → atlanır, yeni satır yok
    pp = lab.apply_prefs(db_session, lab.read_approvals(path), d)
    assert pp["counts"]["tercih_yazilacak"] == 0 and pp["counts"]["bitirilecek"] == 0
    db_session.expire_all()
    assert db_session.query(MaterialSupplierPref).count() == 6


def test_apply_prefs_shifted_row_or_bad_value_writes_nothing(db_session: Session, tmp_path):
    _its, _s, _grp, d = _pref_setup(db_session)
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    _fill(path, "Tercih önerisi", lab.APPROVE_COL, {"A1-08-S1": "E", "A1-09-S1": "belki"})
    appr = lab.read_approvals(path)
    assert any("belki" in e for e in appr["errors"])
    pp = lab.apply_prefs(db_session, appr, d)
    assert pp["applied"] is False
    db_session.expire_all()
    assert db_session.query(MaterialSupplierPref).count() == 1
    # Satır kayması: malzeme metni plana uymuyor → hata
    wb = load_workbook(path)
    ws = wb["Tercih önerisi"]
    head = [c.value for c in ws[1]]
    ws.cell(2, head.index("Malzeme (not)") + 1, "BAŞKA MALZEME")
    ws.cell(3, head.index(lab.APPROVE_COL) + 1, "H")
    wb.save(path)
    pp = lab.apply_prefs(db_session, lab.read_approvals(path), d)
    assert pp["applied"] is False and any("kaymış" in e for e in pp["errors"])


def test_cli_dry_run_writes_nothing_and_xlsx(db_session: Session, tmp_path, capsys):
    _items(db_session, ("GLİSERİN", "g"))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    f = tmp_path / "veri.json"
    f.write_text(json.dumps(data(rec("A1-09-S1", "GLİSERİN", "UMAYCHEM", 4.52, secili=True, guven="belirsiz"),
                                 rec("A1-20-S1", "GLİSERİN", "BEFCHEM", 5.0)), ensure_ascii=False),
                 encoding="utf-8")
    x = tmp_path / "k.xlsx"
    assert lab.main(["--data", str(f), "--xlsx", str(x)]) == 0
    out = capsys.readouterr().out
    assert "KURU ÇALIŞTIRMA" in out and "BEFCHEM" in out and x.exists()
    assert _prices(db_session) == []
    # Fiyat yazımı onaylı Excel ister
    assert lab.main(["--data", str(f), "--commit"]) == 2
    assert lab.main(["--data", str(f), "--create-suppliers"]) == 2
    assert _prices(db_session) == []
    capsys.readouterr()
    assert lab.main(["--data", str(f), "--onayli", str(x)]) == 0           # kuru, onaylı denetim
    assert "KURU" in capsys.readouterr().out and _prices(db_session) == []
    assert lab.main(["--data", str(f), "--onayli", str(x), "--commit"]) == 0
    assert "YAZILDI" in capsys.readouterr().out
    assert len(_lab_prices(db_session)) == 1
    # İşlenmemiş düzeltme → çıkış 2, yazılan satırlar değişmez
    _fill(x, "Belirsiz okumalar", lab.FIX_VALUE_COL, {"A1-09-S1": "4,25"})
    assert lab.main(["--data", str(f), "--onayli", str(x), "--commit"]) == 2
    assert "DURDURULDU" in capsys.readouterr().out
    assert [p.unit_price for p in _lab_prices(db_session)] == [4.52]
    assert lab.main(["--data", str(f), "--apply-prefs", str(tmp_path / "yok.xlsx")]) == 2


# ─── İnceleme düzeltmeleri (08.10.2026) ─────────────────────────────────────

def test_struck_and_price_list_records_not_written_nor_competing(db_session: Session, tmp_path):
    """Üstü çizili teklif (CLOVE OIL proforması, Fındık Naturalya) ve s.11
    Doalin fiyat listesi yazılmaz; çizili belge 'belge > lab' önceliğiyle
    sarı lab satırını EZMEZ.  Doalin yine tercih önerisinde alternatiftir."""
    its = _items(db_session, ("KARANFİL TOMURCUĞU UÇUCU YAGI", "ml"), ("BERGAMOT UÇUCU YAĞI", "ml"),
                 ("FINDIK YAGI", "g"))
    _sups(db_session, "NATURALYA", "DOALİNN", "SURYA KİMYA")
    db_session.commit()
    d = data(
        rec("A2-36-S1", "KARANFİL TOMURCUĞU UÇUCU YAGI", "NATURALYA", 37.84, secili=True, birim="ml"),
        rec("C-s16-4", "CLOVE OIL", "NATURALYA KİMYA", 42.0, tur="proforma", tarih="2026-06-25", cur="USD",
            pkg=1.0, sayfa=16, note="satır üstü çizili; lab tablosu karşılığı: KARANFİL TOMURCUĞU UÇUCU YAGI (A2-36)"),
        rec("A2-16-S2", "FINDIK YAGI", "NATURALYA", 17.2),
        rec("A2-16-S3", "FINDIK YAGI", "SURYA KİMYA", 12.9, secili=True),
        rec("A2-28-S1", "BERGAMOT UÇUCU YAĞI", "NATURALYA", 36.12, secili=True, birim="ml"),
        rec("C-s11-1", "Bergamot Yağı", "Doalin", 175.0, tur="fiyat_listesi", tarih=None, cur="USD", unit="l",
            pkg=1.0, guven="belirsiz", sayfa=11, note="lab tablosu karşılığı: BERGAMOT UÇUCU YAĞI (A2-28)"),
    )
    d["kayitlar"][1]["cizili"] = True
    d["kayitlar"][2]["cizili"] = True
    p = lab.apply(db_session, d)
    r = _by_id(p)
    assert r["C-s16-4"]["status"] == "cizili" and r["A2-16-S2"]["status"] == "cizili"
    assert r["A2-36-S1"]["status"] == "yazilacak" and not r["A2-36-S1"]["overridden"]
    assert r["C-s11-1"]["status"] == "kapsam_disi" and p["counts"]["kapsam_disi"] == 1
    kar, ber, fin = (its["KARANFİL TOMURCUĞU UÇUCU YAGI"].id, its["BERGAMOT UÇUCU YAĞI"].id,
                     its["FINDIK YAGI"].id)
    rows = {(x.item_id, x.supplier_name): x for x in _lab_prices(db_session)}
    assert set(rows) == {(kar, "NATURALYA"), (fin, "SURYA KİMYA"), (ber, "NATURALYA")}
    assert rows[(kar, "NATURALYA")].unit_price == 37.84 and rows[(kar, "NATURALYA")].source == "lab_notu"
    alt = next(x for x in p["prefs"] if x["id"] == "C-s11-1")
    assert alt["kind"] == "nottan alternatif" and alt["item"]["id"] == ber and alt["sup"]["name"] == "DOALİNN"
    wb = load_workbook(lab.build_xlsx(p, tmp_path / "k.xlsx"))
    head = [c.value for c in wb["Fiyatlar"][1]]
    fiy = {row[0]: row for row in wb["Fiyatlar"].iter_rows(min_row=2, values_only=True)}
    assert "üstü çizili" in fiy["C-s16-4"][head.index("Yazılacak mı / neden")]
    assert "fiyat listesi" in fiy["C-s11-1"][head.index("Yazılacak mı / neden")]
    bel = wb["Belirsiz okumalar"]
    bh = [c.value for c in bel[1]]
    brow = next(x for x in bel.iter_rows(min_row=2, values_only=True) if x[0] == "C-s11-1")
    assert brow[bh.index(lab.FIX_VALUE_COL)] == "—"


def test_document_name_vs_hint_conflict_is_ambiguous_and_alt_scope_from_row(db_session: Session):
    """TR/EN mükerrer kart: belgenin adı BERGAMOT YAĞI kartına birebir uyuyor
    ama ipucu A2-28 (BERGAMOT UÇUCU YAĞI) → belirsiz; Uludağ faturası lab
    satırından ayrı karta yazılıp (kart, firma) kuralını delmez.  Nottan
    alternatif tercih lab satırının kartına bağlanır."""
    its = _items(db_session, ("BERGAMOT UÇUCU YAĞI", "ml"), ("BERGAMOT YAĞI", "ml"))
    _sups(db_session, "NATURALYA", "ULUDAĞ HERBAL", "DOALİNN")
    db_session.commit()
    p = lab.plan(db_session, data(
        rec("A2-28-S1", "BERGAMOT UÇUCU YAĞI", "NATURALYA", 36.12, secili=True, birim="ml"),
        rec("A2-28-S2", "BERGAMOT UÇUCU YAĞI", "ULUDAĞ HERBAL", 44.99, birim="ml"),
        rec("C-s9-2", "BERGAMOT YAĞI 1LT", "ULUDAĞ AGRO", 11054.17, tur="fatura", tarih="2026-06-26", cur="TRY",
            unit="l", pkg=1.0, note="lab tablosu karşılığı: BERGAMOT UÇUCU YAĞI (A2-28)"),
        rec("C-s11-1", "Bergamot Yağı", "Doalin", 175.0, tur="fiyat_listesi", tarih=None, cur="USD", unit="l",
            pkg=1.0, note="lab tablosu karşılığı: BERGAMOT UÇUCU YAĞI (A2-28)"),
    ))
    r = _by_id(p)
    uc, dup = its["BERGAMOT UÇUCU YAĞI"].id, its["BERGAMOT YAĞI"].id
    assert r["C-s9-2"]["status"] == "belirsiz" and "ipucu çelişiyor" in "; ".join(r["C-s9-2"]["flags"])
    assert set(r["C-s9-2"]["mat"]["candidates"]) == {uc, dup}
    assert r["A2-28-S2"]["status"] == "yazilacak"
    alt = next(x for x in p["prefs"] if x["id"] == "C-s11-1")
    assert alt["item"]["id"] == uc and alt["scope"] == ("item", uc)
    # İpucu adın kartını gösteriyorsa (normal durum) birebir ad eşleşmesi aynen
    p = lab.plan(db_session, data(
        rec("A2-28-S1", "BERGAMOT YAĞI", "NATURALYA", 36.12, secili=True, birim="ml"),
        rec("C-s9-2", "BERGAMOT YAĞI 1LT", "ULUDAĞ AGRO", 11054.17, tur="fatura", tarih="2026-06-26", cur="TRY",
            unit="l", pkg=1.0, note="lab tablosu karşılığı: BERGAMOT YAĞI (A2-28)")))
    assert _by_id(p)["C-s9-2"]["item"]["id"] == dup and _by_id(p)["C-s9-2"]["mat"]["how"] == "ad"


def test_approval_cells_accept_only_e_or_h(db_session: Session, tmp_path):
    """'X' (Türkçede çoğu zaman hayır) ve '✓' gibi işaretler hata; Excel
    listesi de liste dışı girişi reddeder (showErrorMessage + stop)."""
    for v in ("X", "x", "✓", "+", "1", "0", "Y", "yes", "N", "-"):
        with pytest.raises(ValueError):
            lab._yes_no(v)
    assert lab._yes_no("e") is True and lab._yes_no(" Evet ") is True and lab._yes_no("E.") is True
    assert lab._yes_no("hayır") is False and lab._yes_no(None) is None and lab._yes_no("  ") is None
    _its, _s, _grp, d = _pref_setup(db_session)
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    _fill(path, "Tercih önerisi", lab.APPROVE_COL, {"A1-08-S1": "✓", "A1-09-S1": "X", "A1-34-S2": "E"})
    appr = lab.read_approvals(path)
    assert len(appr["errors"]) == 2 and "A1-08-S1" not in appr["prefs"] and "A1-09-S1" not in appr["prefs"]
    pp = lab.apply_prefs(db_session, appr, d)
    assert pp["applied"] is False
    db_session.expire_all()
    assert db_session.query(MaterialSupplierPref).count() == 1
    dv = load_workbook(path)["Tercih önerisi"].data_validations.dataValidation[0]
    assert dv.showErrorMessage and dv.errorStyle == "stop"


def test_unprocessed_corrections_block_price_and_pref_writes(db_session: Session, tmp_path):
    """'Belirsiz okumalar → Doğru değer' ve 'Eşleşmeyenler → Doğru IMS kartı'
    otomatik işlenmez; veri dosyasına aktarılmadan HİÇBİR ŞEY yazılmaz
    (2,15 yazılıp 7,15 düzeltmesi sessizce kaybolmasın)."""
    its = _items(db_session, ("COCO GLUCOSİDE", "g"), ("NAR ÇEKİRDEĞİ YAĞI", "g"))
    _sups(db_session, "UMAYCHEM")
    db_session.commit()
    d = data(rec("A1-06-S1", "COCO GLUCOSİDE", "UMAYCHEM", 2.15, secili=True, guven="belirsiz", pkg=220.0),
             rec("A2-01-S2", "NAR TOHUMU ÖZÜ", "UMAYCHEM", 18.57))
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    _fill(path, "Belirsiz okumalar", lab.FIX_VALUE_COL, {"A1-06-S1": "7,15"})
    _fill(path, "Eşleşmeyenler", lab.FIX_CARD_COL, {"A2-01-S2": "NAR ÇEKİRDEĞİ YAĞI"})
    appr = lab.read_approvals(path)
    assert not appr["errors"] and len(appr["fixes"]) == 2
    p = lab.apply(db_session, d, approvals=appr)
    assert p["applied"] is False and len(p["errors"]) == 2 and _lab_prices(db_session) == []
    assert any("7,15" in e for e in p["errors"]) and any("elle_eslesme" in e for e in p["errors"])
    assert lab.apply_prefs(db_session, appr, d)["applied"] is False
    # Düzeltmeler veri dosyasına işlendi → düzeltilmiş değerle yazılır
    d["kayitlar"][0].update(fiyat=7.15, guven="kesin")
    d["elle_eslesme"] = {"kart": {"A2-01": its["NAR ÇEKİRDEĞİ YAĞI"].id}}
    p = lab.apply(db_session, d, approvals=appr)
    assert p["applied"] and p["errors"] == []
    assert sorted(x.unit_price for x in _lab_prices(db_session)) == [7.15, 18.57]


def test_pref_and_phase_out_conflict_is_error(db_session: Session, tmp_path):
    """Aynı firma hem tercih (E) hem bitirilecek (E) ya da tercih edilen firma
    zaten bitirilecek → hata, hiçbir şey yazılmaz (bitirilecek tercihi
    geçersiz kılar)."""
    _its, _s, _grp, d = _pref_setup(db_session)
    path = lab.build_xlsx(lab.plan(db_session, d), tmp_path / "onay.xlsx")
    _fill(path, "Tercih önerisi", lab.APPROVE_COL, {"A2-13-S1": "E", "A1-09-S1": "E"})
    small_ids = {x["sup"]["name"]: x["id"] for x in lab.plan(db_session, d)["small"]}
    _fill(path, "Küçük satıcı önerisi", lab.PHASE_COL, {small_ids["TATLIDİLİMLER"]: "E"})
    pp = lab.apply_prefs(db_session, lab.read_approvals(path), d)
    assert pp["applied"] is False
    assert any("TATLIDİLİMLER" in e and "bitirilecek" in e for e in pp["errors"])
    db_session.expire_all()
    assert db_session.query(MaterialSupplierPref).count() == 1
    t = db_session.query(Supplier).filter(Supplier.name == "TATLIDİLİMLER").one()
    assert t.purchase_status == "normal"
    _fill(path, "Küçük satıcı önerisi", lab.PHASE_COL, {small_ids["TATLIDİLİMLER"]: ""})
    t.purchase_status = "phase_out"
    db_session.commit()
    pp = lab.plan_prefs(db_session, lab.read_approvals(path), d)
    assert any("zaten 'bitirilecek'" in e for e in pp["errors"])


def test_pref_rows_vat_label_and_uludag_cheaper_warning(db_session: Session, tmp_path):
    _items(db_session, ("VANİLYA UÇUCU YAĞI", "ml"), ("ITIR HİDROSOLÜ", "g"))
    _sups(db_session, "PHARMATERM", "ULUDAĞ HERBAL")
    db_session.commit()
    p = lab.plan(db_session, data(
        rec("A2-40-S1", "VANİLYA UÇUCU YAĞI", "ULUDAĞ HERBAL", 47.58, birim="ml"),
        rec("A2-40-S2", "VANİLYA UÇUCU YAĞI", "PHARMATERM İLAÇ", 82.56, kdv=True, oran=20.0, secili=True,
            birim="ml"),
        rec("A1-27-S1", "ITIR HİDROSOLÜ", "ULUDAĞ HERBAL", 5.65),
        rec("A1-27-S2", "ITIR HİDROSOLÜ", "PHARMATERM İLAÇ", 7.95, secili=True),
    ))
    ws = load_workbook(lab.build_xlsx(p, tmp_path / "k.xlsx"))["Tercih önerisi"]
    head = [c.value for c in ws[1]]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["A2-40-S2"][head.index("Fiyat")] == "82,56 EUR/kg (KDV %20 dahil; net 68,8)"
    assert rows["A1-27-S2"][head.index("Fiyat")] == "7,95 EUR/kg"
    w = rows["A1-27-S2"][head.index("Uyarı")]
    assert "Uludağ ucuz olanları alacağız" in w and "5,65 EUR ↔ sarı 7,95 EUR" in w
    assert "Uludağ" not in (rows["A2-40-S2"][head.index("Uyarı")] or "")      # listede olmayan satır
    assert rows["A1-27-S1"][head.index("Öneri türü")] == "nottan alternatif"   # s.17 yeşil Uludağ


# ─── Gerçek veri dosyası ────────────────────────────────────────────────────

def test_real_data_file_plans_and_builds_excel(db_session: Session, tmp_path):
    d = lab.load_data()
    assert len(d["kayitlar"]) == 160
    names = sorted({r["malzeme"] for r in d["kayitlar"] if r["tur"] == "lab_tablo"})
    _items(db_session, *[(n, "g") for n in names])
    _sups(db_session, "PHARMATERM", "DOALİNN", "ULUDAĞ HERBAL", "VESER KİMYEVİ", "NATURALYA",
          "SURYA KİMYA", "UMAYCHEM", "KRK GIDA", "TATLIDİLİMLER", "HAMMADDESEPETİ", "YİĞİTOĞLU KİMYA")
    db_session.commit()
    p = lab.plan(db_session, d)
    c = p["counts"]
    assert c["kayit"] == 160 and sum(c[k] for k in lab.STATUS_TEXT) == 160
    r = _by_id(p)
    # Pharmaterm KDV dahil uçucu yağlar (7) — yazılacak satırda KDV taşınır
    assert sum(1 for x in p["records"] if x["vat_included"] is True) == 7
    assert r["A2-40-S2"]["status"] == "yazilacak" and r["A2-40-S2"]["vat_rate"] == 20.0
    assert r["A1-08-S3"]["status"] == "ayri_form"
    # Tek (kart, firma) kuralı: yazılacaklarda çift yok
    keys = [(x["item"]["id"], x["sup"]["id"]) for x in p["records"] if x["status"] == "yazilacak"]
    assert len(keys) == len(set(keys))
    assert {"BEFCHEM", "ATAMAN KİMYA", "MARKET"} <= {n["name"] for n in p["new_suppliers"]}
    small = {x["sup"]["name"]: x for x in p["small"]}
    assert small["DOALİNN"]["kind"] == "Doalin notu" and small["SURYA KİMYA"]["kind"] == "bilgi"
    for nm in ("HAMMADDESEPETİ", "TATLIDİLİMLER", "ROSECE", "MARKET", "KİMYASAL EVİ"):
        assert small[nm]["kind"] == "küçük satıcı", nm
    pref_ids = [x["id"] for x in p["prefs"]]
    assert sum(1 for x in p["prefs"] if x["kind"] == "sarı seçim") == 80
    assert {"A2-13-S1", "C-s11-1", "C-s11-2", "A1-27-S1", "A1-28-S1"} <= set(pref_ids)
    # İnceleme düzeltmeleri: çizili teklifler + Doalin fiyat listesi yazılmaz;
    # çizili CLOVE OIL sarı karanfili ezmez; 2 sarılı aloe'de Naturalya üstte
    assert c["cizili"] == 2 and c["kapsam_disi"] == 2
    assert r["C-s16-4"]["status"] == "cizili" and r["A2-36-S1"]["status"] == "yazilacak"
    assert r["C-s11-1"]["status"] == "kapsam_disi" and r["C-s11-2"]["status"] == "kapsam_disi"
    assert pref_ids.index("A1-08-S2") < pref_ids.index("A1-08-S1")
    assert r["C-s10-1"]["material"] == "Kakao Yağı (Katı)"
    warn = {x["id"]: x["warn"] for x in p["prefs"]}
    for rid in ("A1-04-S3", "A1-10-S2", "A1-11-S2", "A1-14-S3", "A1-27-S2", "A1-28-S2", "A2-01-S3", "A2-08-S2"):
        assert "Uludağ ucuz olanları alacağız" in warn[rid], rid
    path = lab.build_xlsx(p, tmp_path / "gercek.xlsx")
    assert load_workbook(path).sheetnames == list(lab.SHEETS)
