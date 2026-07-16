"""
İçindekiler Raporu testleri.

  • Seçim paneli: bitmiş ürünler + marka + reçete durumu; soyut ebeveynler hariç
  • Excel export: 4 sabit sayfa, düz başlıklar, bilinen satır değerleri (fire dahil),
    ürün başına Bileşim % toplamı ≈ 100, Hammadde/Ambalaj ayrımı, reçetesiz sayfası
  • A4 baskı ayarı: her sayfada paperSize=A4 + fitToWidth + başlık tekrarı (NET-YENİ)
  • PDF export: antetli, tüm sayfalar A4, içerik metni doğru
  • Güvenlik: domain izolasyonu, boş seçim 400, anonim 401, Distributor 403
"""
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy.orm import Session

from database import Item, Recipe, RecipeIngredient

_HDR = {"Origin": "http://testserver"}
_A4_W, _A4_H = 595.276, 841.89   # pt

_EXPORT = "/api/recipes/ingredients-report/export"
_PRODUCTS = "/api/recipes/ingredients-report/products"


def _all_pages_a4(pdf_bytes: bytes) -> bool:
    pages = PdfReader(BytesIO(pdf_bytes)).pages
    assert len(pages) >= 1
    for pg in pages:
        b = pg.mediabox
        if not (abs(float(b.width) - _A4_W) < 2 and abs(float(b.height) - _A4_H) < 2):
            return False
    return True


def _product(db, *, brand="Serenida", sku="P1", with_recipe=True, waste=10.0,
             domain="cosmetics"):
    """Bitmiş ürün (+ opsiyonel reçete: 2 hammadde + şişe + TR etiket) kurar."""
    tgt = Item(name=f"{brand} Krem {sku}", name_tr=f"{brand} KREM TR {sku}",
               sku=f"tgt-{sku}", category="Bitmiş Ürün", unit="adet",
               barcode=f"869{sku}", variation_name="200ml",
               current_stock=0, domain=domain)
    db.add(tgt); db.flush()
    if not with_recipe:
        db.commit()
        return tgt.id, None
    su = Item(name=f"Su {sku}", name_tr=f"SU TR {sku}", sku=f"su-{sku}",
              category="Hammadde", unit="g", current_stock=1000, domain=domain)
    gli = Item(name=f"Gliserin {sku}", sku=f"gli-{sku}", category="Hammadde",
               unit="g", current_stock=1000, domain=domain)
    sise = Item(name=f"Şişe {sku}", sku=f"sise-{sku}", category="Ambalaj",
                unit="adet", pkg_type="şişe", current_stock=500, domain=domain)
    eti = Item(name=f"Etiket TR {sku}", sku=f"eti-{sku}", category="Ambalaj",
               unit="adet", pkg_type="etiket", language="TR",
               label_group=f"lbl-{sku}", current_stock=500, domain=domain)
    db.add_all([su, gli, sise, eti]); db.flush()
    rec = Recipe(name=tgt.name, output_quantity=1, output_unit="adet",
                 target_item_id=tgt.id, waste_percentage=waste, domain=domain)
    db.add(rec); db.flush()
    db.add_all([
        RecipeIngredient(recipe_id=rec.id, item_id=su.id, quantity=70.0, unit="g", phase="A"),
        RecipeIngredient(recipe_id=rec.id, item_id=gli.id, quantity=30.0, unit="g", phase="B"),
        RecipeIngredient(recipe_id=rec.id, item_id=sise.id, quantity=1.0, unit="adet"),
        RecipeIngredient(recipe_id=rec.id, item_id=eti.id, quantity=1.0, unit="adet"),
    ])
    db.commit()
    return tgt.id, rec.id


def _export_wb(client, item_ids):
    r = client.post(_EXPORT, json={"item_ids": item_ids}, headers=_HDR)
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(BytesIO(r.content))


def _rows(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


# ─── Seçim paneli ────────────────────────────────────────────────────────────

def test_products_endpoint_lists_finished_goods(authed_client: TestClient, db_session: Session):
    iid, rid = _product(db_session)
    r = authed_client.get(_PRODUCTS)
    assert r.status_code == 200
    prods = {p["item_id"]: p for p in r.json()["products"]}
    p = prods[iid]
    assert p["brand"] == "Serenida" and p["has_recipe"] is True and p["recipe_id"] == rid
    assert p["variation_name"] == "200ml" and p["name_tr"].startswith("Serenida KREM TR")


def test_products_endpoint_flags_recipeless(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session, with_recipe=False)
    prods = {p["item_id"]: p for p in authed_client.get(_PRODUCTS).json()["products"]}
    assert prods[iid]["has_recipe"] is False and prods[iid]["recipe_id"] is None


def test_products_endpoint_excludes_abstract_parents(authed_client: TestClient, db_session: Session):
    parent = Item(name="Minerva Serum", sku="par-1", category="Bitmiş Ürün",
                  unit="adet", current_stock=0, domain="cosmetics")
    db_session.add(parent); db_session.flush()
    child = Item(name="Minerva Serum 30ML", sku="var-1", category="Bitmiş Ürün",
                 unit="adet", parent_id=parent.id, variation_name="30ml",
                 current_stock=0, domain="cosmetics")
    db_session.add(child); db_session.commit()
    prods = {p["item_id"]: p for p in authed_client.get(_PRODUCTS).json()["products"]}
    assert parent.id not in prods                      # soyut ebeveyn asla listede değil
    assert prods[child.id]["variation_name"] == "30ml"
    assert prods[child.id]["parent_name"] == "Minerva Serum"


def test_brand_variants_coalesce(authed_client: TestClient, db_session: Session):
    """'Minerva' / 'MİNERVA' yazım varyantları tek markada toplanır (prod'da gerçek durum)."""
    i1, _ = _product(db_session, brand="Minerva", sku="V1", with_recipe=False)
    i2, _ = _product(db_session, brand="Minerva", sku="V2", with_recipe=False)
    i3, _ = _product(db_session, brand="MİNERVA", sku="V3", with_recipe=False)
    prods = {p["item_id"]: p for p in authed_client.get(_PRODUCTS).json()["products"]}
    brands = {prods[i]["brand"] for i in (i1, i2, i3)}
    assert brands == {"Minerva"}          # en yaygın yazım kazanır, varyant ayrı chip olmaz


# ─── Excel export ────────────────────────────────────────────────────────────

def test_export_returns_xlsx(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    r = authed_client.post(_EXPORT, json={"item_ids": [iid]}, headers=_HDR)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"                      # xlsx = zip
    assert "icindekiler_raporu" in r.headers["content-disposition"]
    r.headers["content-disposition"].encode("latin-1")  # header-safe (ASCII)


def test_export_sheet_names(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    wb = _export_wb(authed_client, [iid])
    assert wb.sheetnames == ["Özet", "Hammaddeler", "Ambalaj-Paketleme", "Reçetesiz Ürünler"]


def test_export_headers(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    wb = _export_wb(authed_client, [iid])
    assert [c.value for c in wb["Hammaddeler"][1]] == [
        "Marka", "Ürün (EN)", "Ürün (TR)", "Ürün SKU", "Varyasyon", "Reçete ID",
        "Faz", "Hammadde (EN)", "Hammadde (TR)", "Hammadde SKU", "Net Miktar",
        "Fire %", "Brüt Miktar", "Birim", "Bileşim %"]
    assert [c.value for c in wb["Ambalaj-Paketleme"][1]] == [
        "Marka", "Ürün (EN)", "Ürün (TR)", "Ürün SKU", "Varyasyon", "Reçete ID",
        "Bileşen (EN)", "Bileşen (TR)", "Bileşen SKU", "Ambalaj Tipi",
        "Etiket Dili", "Etiket Grubu", "Miktar", "Birim"]


def test_export_known_raw_row(authed_client: TestClient, db_session: Session):
    iid, rid = _product(db_session, waste=10.0)        # Su 70g faz A, fire %10
    wb = _export_wb(authed_client, [iid])
    su = next(r for r in _rows(wb["Hammaddeler"]) if str(r[7]).startswith("Su "))
    assert su[6] == "A"                                # Faz
    assert su[10] == 70.0 and not isinstance(su[10], str)  # Net (numeric — openpyxl int'e düşürebilir)
    assert su[11] == 10.0                              # Fire %
    assert abs(su[12] - 77.0) < 1e-6                   # Brüt = 70 × 1.10
    assert abs(su[14] - 70.0) < 1e-6                   # Bileşim % = 70/(70+30)
    # Özet satırı da tutarlı
    ozet = _rows(wb["Özet"])[0]
    assert ozet[6] == rid and ozet[7] == "Var"
    assert ozet[11] == 2 and ozet[12] == 2             # 2 hammadde + 2 ambalaj kalemi
    assert abs(ozet[13] - 100.0) < 1e-6                # toplam hammadde net


def test_export_pct_sums_100(authed_client: TestClient, db_session: Session):
    i1, _ = _product(db_session, sku="A1")
    i2, _ = _product(db_session, brand="Minerva", sku="B2", waste=25.0)
    wb = _export_wb(authed_client, [i1, i2])
    by_prod = {}
    for r in _rows(wb["Hammaddeler"]):
        by_prod.setdefault(r[1], 0.0)
        by_prod[r[1]] += r[14]
    assert len(by_prod) == 2
    for total in by_prod.values():
        assert abs(total - 100.0) < 0.01


def test_ambalaj_split(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    wb = _export_wb(authed_client, [iid])
    raw_names = {r[7] for r in _rows(wb["Hammaddeler"])}
    pkg = {r[6]: r for r in _rows(wb["Ambalaj-Paketleme"])}
    assert not any(n.startswith(("Şişe", "Etiket")) for n in raw_names)
    assert not any(n.startswith(("Su", "Gliserin")) for n in pkg)
    eti = next(v for k, v in pkg.items() if k.startswith("Etiket"))
    assert eti[9] == "etiket" and eti[10] == "TR" and str(eti[11]).startswith("lbl-")
    sise = next(v for k, v in pkg.items() if k.startswith("Şişe"))
    assert sise[9] == "şişe" and sise[12] == 1.0


def test_recipeless_product_on_recipeless_sheet(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session, sku="R1", with_recipe=False)
    wb = _export_wb(authed_client, [iid])
    rl = _rows(wb["Reçetesiz Ürünler"])
    assert len(rl) == 1 and rl[0][1] == "Serenida Krem R1"
    # 1-3. sayfalarda İZİ YOK
    assert not _rows(wb["Özet"]) and not _rows(wb["Hammaddeler"]) \
        and not _rows(wb["Ambalaj-Paketleme"])


def test_a4_page_setup_every_sheet(authed_client: TestClient, db_session: Session):
    """A4 baskı garantisi — repo'daki ilk Excel print-setup'ının regresyon kilidi."""
    iid, _ = _product(db_session)
    wb = _export_wb(authed_client, [iid])
    for name in wb.sheetnames:
        ws = wb[name]
        assert str(ws.page_setup.paperSize) == "9", name          # A4
        assert int(ws.page_setup.fitToWidth) == 1, name
        assert ws.sheet_properties.pageSetUpPr.fitToPage, name
        assert ws.print_title_rows == "$1:$1", name               # başlık tekrarı
    assert wb["Hammaddeler"].page_setup.orientation == "landscape"
    assert wb["Reçetesiz Ürünler"].page_setup.orientation == "portrait"


# ─── PDF export ──────────────────────────────────────────────────────────────

def test_export_pdf(authed_client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    r = authed_client.post(_EXPORT, json={"item_ids": [iid], "format": "pdf"}, headers=_HDR)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    assert "icindekiler_raporu" in r.headers["content-disposition"]
    assert _all_pages_a4(r.content)
    text = "".join(pg.extract_text() for pg in PdfReader(BytesIO(r.content)).pages)
    assert "Serenida" in text and "Gliserin" in text


# ─── Güvenlik ────────────────────────────────────────────────────────────────

def test_export_domain_isolation(authed_client: TestClient, db_session: Session):
    cos_id, _ = _product(db_session, sku="C1", domain="cosmetics")
    sup_id, _ = _product(db_session, brand="Takviye", sku="S1", domain="supplement")
    # cosmetics panelinde (varsayılan) supplement ürünü listede YOK
    prods = {p["item_id"] for p in authed_client.get(_PRODUCTS).json()["products"]}
    assert cos_id in prods and sup_id not in prods
    # supplement id'siyle export → assemble'da düşer → 400
    r = authed_client.post(_EXPORT, json={"item_ids": [sup_id]}, headers=_HDR)
    assert r.status_code == 400


def test_export_empty_selection_400(authed_client: TestClient, db_session: Session):
    r = authed_client.post(_EXPORT, json={"item_ids": [999999]}, headers=_HDR)
    assert r.status_code == 400
    assert r.json()["detail"] == "Seçilen ürün bulunamadı."


def test_permissions(authed_client: TestClient, client: TestClient, db_session: Session):
    iid, _ = _product(db_session)
    # Distribütör oluştur (SuperAdmin) ve onunla login ol — recipes.view yok → 403
    r = authed_client.post("/api/distributors", headers=_HDR, json={
        "username": "bayi-icr", "password": "Distpass123", "company_name": "Bayi ICR",
        "currency": "TRY"})
    assert r.status_code == 201, r.text
    r = authed_client.post("/api/login", headers=_HDR,
                           json={"username": "bayi-icr", "password": "Distpass123"})
    assert r.status_code == 200
    assert authed_client.get(_PRODUCTS).status_code == 403
    assert authed_client.post(_EXPORT, json={"item_ids": [iid]}, headers=_HDR).status_code == 403
    # Anonim → 401/403 (repo davranışı: kimliksiz istek 403 dönebilir)
    r = client.post(_EXPORT, json={"item_ids": [iid]}, headers=_HDR)
    assert r.status_code in (401, 403)
    assert client.get(_PRODUCTS).status_code in (401, 403)
