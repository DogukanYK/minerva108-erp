"""
"Aynı malzeme" grup stoğu — raporlar, Stoklar, dashboard, bildirim (13.10.2026).

Kilitlenen sözleşmeler (çekirdek: core/group_stock.py, tests/test_group_stock.py):
  • Stoklar: `/api/inventory/summary` kartlara `group_id`/`group_key` +
    birim ailesi çubuk alanları ekler; `/api/inventory/groups-summary` grup ×
    kova başına SUNUCUDA biçimli toplam + döküm döner (ayardan bağımsız,
    panel izole, oturum şart).
  • Aylık Stok Raporu her zaman gruplar (bugünkü grup): grup girdisi, üyeler
    ardışık, bölüm toplamı DEĞİŞMEZ; dondurulmuş ayda dipnot; Excel'de kalın
    grup satırı (toplam METİN — SUM çift saymaz), girintili üyeler, "Grup
    dökümü" sütunu.
  • Eksik Hammaddeler / dashboard / push — grup kuralı YALNIZ ayar açıkken:
    grubu yeterli kart düşer, eksik grup tek satır (üyeler altında), sipariş
    bayrağı kart başına, grup "sipariş verildi" görünür; ayar kapalı → eski
    kart başı davranış birebir.
  • Asistan: item_detail grup toplamı + döküm; stock_gaps grup satırı.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.production_plan import CFG_ENABLED
from database import (AppSetting, Inventory, Item, MaterialGroup, Recipe, RecipeIngredient,
                      StockOrderFlag, StockSnapshot, Supplier)

_HDR = {"Origin": "http://testserver"}
_GAPS = "/api/reports/stock-gaps"


# ─── Kurulum ────────────────────────────────────────────────────────────────

def _grp(db, name, domain="cosmetics", active=True):
    g = MaterialGroup(name=name, domain=domain, is_active=active)
    db.add(g)
    db.flush()
    return g


def _sup(db, name, domain="cosmetics"):
    s = Supplier(name=name, domain=domain)
    db.add(s)
    db.flush()
    return s


def _card(db, name, group=None, *, stock, unit="g", mn=0.0, sup=None, category="Hammadde",
          pkg_type=None, domain="cosmetics", active=True, barcode=None):
    it = Item(name=name, sku=f"gsr-{name}", category=category, unit=unit, current_stock=stock,
              min_stock_level=mn, supplier_id=sup.id if sup else None, pkg_type=pkg_type,
              material_group_id=group.id if group else None, domain=domain, is_active=active,
              barcode=barcode)
    db.add(it)
    db.flush()
    return it


def _set(db, value):
    row = db.query(AppSetting).filter_by(key=CFG_ENABLED).first()
    if row is None:
        db.add(AppSetting(key=CFG_ENABLED, value=value))
    else:
        row.value = value
    db.commit()


def _xhantan(db):
    """Canlıdaki grup #5: kart #105 HAMMADDESEPETİ 180,31 g + #636 YİĞİTOGLU
    KİMYA 24.687,81 g."""
    g = _grp(db, "XHANTAN GUM")
    a = _card(db, "XHANTAN GUM", g, stock=180.31, sup=_sup(db, "HAMMADDESEPETİ"))
    b = _card(db, "XANTHAN GUM", g, stock=24687.81, sup=_sup(db, "YİĞİTOGLU KİMYA"))
    return g, a, b


# ─── Stoklar: summary + groups-summary ─────────────────────────────────────

def test_summary_marks_group_cards_and_bar_family(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _xhantan(db)
    solo = _card(db, "SOLO", stock=2, unit="kg")
    bottle = _card(db, "ŞİŞE", stock=40, unit="adet", category="Ambalaj", pkg_type="şişe")
    db.commit()
    rows = {r["item_id"]: r for r in authed_client.get("/api/inventory/summary").json()}
    assert rows[a.id]["group_id"] == g.id and rows[a.id]["group_key"] == f"{g.id}:hammadde"
    assert rows[b.id]["group_key"] == f"{g.id}:hammadde"
    assert rows[solo.id]["group_id"] is None and rows[solo.id]["group_key"] is None
    # Çubuk: aile + küçük birim (kg → g); g ile adet aynı ölçekte kıyaslanmaz
    assert rows[solo.id]["unit_family"] == "mass" and rows[solo.id]["bar_value"] == pytest.approx(2000)
    assert rows[bottle.id]["unit_family"] == "count" and rows[bottle.id]["bar_value"] == 40
    assert rows[a.id]["total_stock"] == 180.31                  # kart stoğu değişmez


def test_groups_summary_server_formatted(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _xhantan(db)
    db.commit()
    d = authed_client.get("/api/inventory/groups-summary").json()
    assert d["rule_active"] is True                              # ayar satırı yok = açık
    (grp,) = d["groups"]
    assert grp["key"] == f"{g.id}:hammadde" and grp["bucket"] == "hammadde"
    assert grp["name"] == "XHANTAN GUM" and grp["base_unit"] == "g"
    assert grp["total"] == pytest.approx(24868.12)
    assert grp["total_text"] == "24,9 kg"
    assert grp["breakdown_text"] == "HAMMADDESEPETİ 180 g · YİĞİTOGLU KİMYA 24,7 kg"
    assert grp["members_label"] == "2 tedarikçi" and grp["supplier_count"] == 2
    assert sorted(grp["member_ids"]) == sorted([a.id, b.id])
    assert grp["unit_family"] == "mass" and grp["bar_value"] == pytest.approx(24868.12)
    assert {m["item_id"]: m["stock_text"] for m in grp["members"]} == {a.id: "180 g", b.id: "24,7 kg"}
    # Yalnız görünüm — ayar kapalıyken de döner
    _set(db, "0")
    d = authed_client.get("/api/inventory/groups-summary").json()
    assert d["rule_active"] is False and len(d["groups"]) == 1


def test_groups_summary_aside_label_domain_and_auth(client: TestClient, authed_client: TestClient,
                                                    db_session):
    db = db_session
    g = _grp(db, "GÜL HİDROSOLÜ")
    _card(db, "Gül ml", g, stock=5000, unit="ml")
    _card(db, "Gül g", g, stock=10)
    _card(db, "Gül ml 2", g, stock=1000, unit="ml")
    sg = _grp(db, "SUPP GRUP", domain="supplement")
    _card(db, "S1", sg, stock=1, domain="supplement")
    _card(db, "S2", sg, stock=1, domain="supplement")
    db.commit()
    (grp,) = authed_client.get("/api/inventory/groups-summary").json()["groups"]
    assert grp["base_unit"] == "ml" and grp["total_text"] == "6,0 l"
    assert grp["aside_text"] == "Gül g 10 g"
    assert grp["members_label"] == "3 kart"                      # tedarikçisiz kartlar
    # Panel izolasyonu
    authed_client.post("/api/domain/switch", json={"domain": "supplement"}, headers=_HDR)
    names = [x["name"] for x in authed_client.get("/api/inventory/groups-summary").json()["groups"]]
    assert names == ["SUPP GRUP"]
    # Oturumsuz → 401
    authed_client.post("/api/logout", headers=_HDR)
    assert client.get("/api/inventory/groups-summary").status_code == 401


def test_groups_summary_labtech_can_read(labtech_client: TestClient, db_session):
    _xhantan(db_session)
    db_session.commit()
    r = labtech_client.get("/api/inventory/groups-summary")
    assert r.status_code == 200 and len(r.json()["groups"]) == 1


def test_stocks_page_has_group_ui(authed_client: TestClient):
    r = authed_client.get("/stocks")
    assert r.status_code == 200
    for needle in ("/api/inventory/groups-summary", "toggleStockGroup", "grp-row", "kpiTotalSub"):
        assert needle in r.text


# ─── Aylık Stok Raporu ─────────────────────────────────────────────────────

def _now_ym():
    import datetime as _dt
    n = _dt.datetime.utcnow()
    return n.year, n.month


def test_monthly_live_groups_members_contiguous_total_unchanged(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _xhantan(db)
    solo = _card(db, "XB SOLO", stock=5)
    db.commit()
    y, m = _now_ym()
    d = authed_client.get(f"/api/reports/monthly-stock?year={y}&month={m}&category=hammadde").json()
    (sec,) = d["sections"]
    assert [r["name"] for r in sec["rows"]] == ["XANTHAN GUM", "XHANTAN GUM", "XB SOLO"]
    assert [r["group_id"] for r in sec["rows"]] == [g.id, g.id, None]
    (grp,) = sec["groups"]
    assert grp["total_text"] == "24,9 kg" and grp["member_count"] == 2
    assert grp["breakdown_text"] == "HAMMADDESEPETİ 180 g · YİĞİTOGLU KİMYA 24,7 kg"
    assert sec["row_count"] == 3 and sec["total"] == pytest.approx(180.31 + 24687.81 + 5)
    assert "bugünkü" in d["group_note"] and "Dondurulmuş" not in d["group_note"]


def test_monthly_snapshot_uses_today_groups_with_note(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _xhantan(db)
    for it, st in ((a, 100.0), (b, 2000.0)):
        db.add(StockSnapshot(year=2026, month=1, item_id=it.id, item_name=it.name,
                             category="Hammadde", unit="g", stock=st, domain="cosmetics"))
    db.commit()
    d = authed_client.get("/api/reports/monthly-stock?year=2026&month=1&category=hammadde").json()
    assert d["source"] == "snapshot"
    (sec,) = d["sections"]
    (grp,) = sec["groups"]
    assert grp["total_text"] == "2,1 kg"                          # o ayın stokları
    assert sec["total"] == pytest.approx(2100.0)
    assert "Dondurulmuş" in d["group_note"]


def test_monthly_single_member_or_inactive_group_not_grouped(authed_client: TestClient, db_session):
    db = db_session
    g = _grp(db, "TEK")
    _card(db, "TEK A", g, stock=5)
    _card(db, "TEK B", g, stock=0)                                 # o ay stoksuz → satır yok
    off = _grp(db, "PASİF", active=False)
    _card(db, "PASİF A", off, stock=1)
    _card(db, "PASİF B", off, stock=1)
    db.commit()
    y, m = _now_ym()
    d = authed_client.get(f"/api/reports/monthly-stock?year={y}&month={m}&category=hammadde").json()
    (sec,) = d["sections"]
    assert sec["groups"] == [] and all(r["group_id"] is None for r in sec["rows"])
    assert d["group_note"] == ""


def test_monthly_excel_group_row_text_total_and_column(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _xhantan(db)
    _card(db, "XB SOLO", stock=5)
    db.commit()
    y, m = _now_ym()
    r = authed_client.get(f"/api/reports/monthly-stock/export?year={y}&month={m}&category=hammadde")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content))["Hammadde"]
    assert ws.cell(row=4, column=7).value == "Grup dökümü"
    assert "bugünkü" in (ws["A3"].value or "")
    grow = next(rw for rw in range(5, ws.max_row + 1)
                if "(aynı malzeme, 2 kart)" in str(ws.cell(row=rw, column=2).value or ""))
    assert ws.cell(row=grow, column=2).font.bold
    assert ws.cell(row=grow, column=6).value == "24,9 kg"          # METİN — SUM çift saymaz
    assert ws.cell(row=grow, column=7).value == "HAMMADDESEPETİ 180 g · YİĞİTOGLU KİMYA 24,7 kg"
    assert ws.cell(row=grow, column=1).value in (None, "")
    assert [ws.cell(row=grow + k, column=2).value for k in (1, 2)] == ["XANTHAN GUM", "XHANTAN GUM"]
    assert ws.cell(row=grow + 1, column=2).alignment.indent == 2
    assert [ws.cell(row=grow + k, column=1).value for k in (1, 2, 3)] == [1, 2, 3]
    assert ws.cell(row=grow + 3, column=2).value == "XB SOLO"
    total_row = grow + 4
    assert ws.cell(row=total_row, column=2).value == "TOPLAM"
    assert ws.cell(row=total_row, column=6).value == pytest.approx(180.31 + 24687.81 + 5)
    nums = [ws.cell(row=rw, column=6).value for rw in range(5, total_row)]
    assert sum(v for v in nums if isinstance(v, (int, float))) == pytest.approx(180.31 + 24687.81 + 5)


# ─── Eksik Hammaddeler ─────────────────────────────────────────────────────

def _short_group(db):
    """Toplam 50 g ≤ grup min (en büyük) 200 g → kritik grup."""
    g = _grp(db, "STEARİL ALKOL")
    a = _card(db, "STEARİL ALKOL", g, stock=0, mn=100, sup=_sup(db, "Ana Kimya"))
    b = _card(db, "CETEARYL ALCOHOL", g, stock=50, mn=200, sup=_sup(db, "Tercih Kimya"))
    return g, a, b


def test_gaps_healthy_group_card_drops_out(authed_client: TestClient, db_session):
    db = db_session
    g = _grp(db, "XHANTAN")
    zero = _card(db, "XHANTAN SIFIR", g, stock=0, mn=100)
    _card(db, "XHANTAN DOLU", g, stock=5000, mn=200)
    solo = _card(db, "GRUPSUZ SIFIR", stock=0)
    db.commit()
    d = authed_client.get(_GAPS).json()
    assert [r["item_id"] for r in d["rows"]] == [solo.id]
    assert d["group_rule"] is True and d["summary"]["sifir"] == 1 and d["summary"]["groups"] == 0
    # Ayar kapalı → kart başı eski kural: sıfır kart yeniden görünür
    _set(db, "0")
    d = authed_client.get(_GAPS).json()
    assert {r["item_id"] for r in d["rows"]} == {zero.id, solo.id}
    assert d["group_rule"] is False and all(r["kind"] == "item" for r in d["rows"])


def test_gaps_short_group_single_row_with_members_and_order_flag(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _short_group(db)
    db.commit()
    r = authed_client.post(f"{_GAPS}/{b.id}/order", headers=_HDR, json={"quantity": 1000})
    assert r.status_code == 201
    d = authed_client.get(_GAPS).json()
    (row,) = d["rows"]
    assert row["kind"] == "group" and row["group_id"] == g.id and row["item_id"] is None
    assert row["status"] == "kritik" and row["unit"] == "g"
    assert row["current_stock"] == 50 and row["min_stock_level"] == 200
    assert row["current_text"] == "50 g" and row["min_text"] == "200 g" and row["deficit_text"] == "150 g"
    assert row["breakdown_text"] == "Ana Kimya 0 g · Tercih Kimya 50 g"
    assert row["member_count"] == 2 and sorted(row["member_ids"]) == sorted([a.id, b.id])
    members = {m["item_id"]: m for m in row["members"]}
    assert members[b.id]["order"]["quantity"] == 1000 and members[a.id]["order"] is None
    assert members[a.id]["status"] == "sifir" and members[a.id]["kind"] == "member"
    assert row["order"]["quantity"] == 1000 and row["ordered_count"] == 1
    s = d["summary"]
    assert (s["sifir"], s["kritik"], s["numune"], s["groups"], s["open_orders"]) == (0, 1, 0, 1, 1)
    # Durum süzgeci grubun durumuna bakar
    assert authed_client.get(f"{_GAPS}?status=sifir").json()["rows"] == []
    # Arama üye tedarikçisini de tarar
    assert len(authed_client.get(f"{_GAPS}?q=tercih").json()["rows"]) == 1
    assert authed_client.get(f"{_GAPS}?q=yokboyle").json()["rows"] == []
    # Ayar kapalı → iki kart satırı
    _set(db, "0")
    d = authed_client.get(_GAPS).json()
    assert {r["item_id"]: r["status"] for r in d["rows"]} == {a.id: "sifir", b.id: "kritik"}


def test_gaps_group_numune_and_min_undefined(authed_client: TestClient, db_session):
    db = db_session
    g = _grp(db, "NUMUNE GRUBU")
    a = _card(db, "N A", g, stock=0, unit="g")
    b = _card(db, "N B", g, stock=0, unit="kg")
    db.add(Inventory(item_id=b.id, lot_number="NUM-1", quantity=0.5, status="APPROVED",
                     is_sample=True, domain="cosmetics"))
    g2 = _grp(db, "MİNSİZ GRUP")
    _card(db, "M A", g2, stock=3)
    _card(db, "M B", g2, stock=0)
    db.commit()
    d = authed_client.get(_GAPS).json()
    (row,) = d["rows"]
    assert row["status"] == "numune" and row["unit"] == "g"
    assert row["sample_qty"] == 500 and row["sample_text"] == "500 g"     # kg → g çevrildi
    assert row["sample_lots"][0]["item_name"] == "N B" and row["sample_lots"][0]["unit"] == "kg"
    assert d["summary"]["min_undefined"] == 1                              # minsiz grup bir kez
    d = authed_client.get(f"{_GAPS}?status=sifir,kritik,numune,min_tanimsiz&include_undefined_min=1").json()
    assert {r["name"]: r["status"] for r in d["rows"]} == {"NUMUNE GRUBU": "numune",
                                                           "MİNSİZ GRUP": "min_tanimsiz"}
    assert a.id in d["rows"][0]["member_ids"]


def test_gaps_aside_unit_card_keeps_card_rule(authed_client: TestClient, db_session):
    db = db_session
    g = _grp(db, "GÜL HİDROSOLÜ")
    _card(db, "Gül ml A", g, stock=5000, unit="ml", mn=100)
    _card(db, "Gül ml B", g, stock=100, unit="ml", mn=50)
    gram = _card(db, "Gül g", g, stock=0, mn=50)                        # başka aile → kart kuralı
    db.commit()
    d = authed_client.get(_GAPS).json()
    assert [(r["kind"], r["item_id"]) for r in d["rows"]] == [("item", gram.id)]


def test_gaps_selection_mode_per_card_with_group_note(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _short_group(db)
    db.commit()
    from core.stock_gaps import assemble
    rep = assemble(db, "cosmetics", item_ids=[a.id, b.id])
    assert [r["kind"] for r in rep["rows"]] == ["item", "item"]
    note = {r["item_id"]: r["group"] for r in rep["rows"]}
    assert note[a.id]["name"] == "STEARİL ALKOL" and note[a.id]["total_text"] == "50 g"
    assert note[a.id]["status"] == "kritik" and note[a.id]["counted"] is True
    _set(db, "0")
    rep = assemble(db, "cosmetics", item_ids=[a.id])
    assert rep["rows"][0]["group"] is None


def test_gaps_export_group_then_members(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _short_group(db)
    solo = _card(db, "GRUPSUZ", stock=0)
    db.commit()
    authed_client.post(f"{_GAPS}/{b.id}/order", headers=_HDR, json={"quantity": 1000})
    r = authed_client.get(f"{_GAPS}/export?format=xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Eksikler"]
    hdr = next(rw for rw in range(1, ws.max_row + 1) if ws.cell(rw, 1).value == "Hammadde")
    assert ws.cell(hdr, 14).value == "Grup dökümü"
    names = [ws.cell(rw, 1).value for rw in range(hdr + 1, ws.max_row + 1)]
    gi = names.index("STEARİL ALKOL (aynı malzeme, 2 kart)")
    assert set(names[gi + 1:gi + 3]) == {"STEARİL ALKOL", "CETEARYL ALCOHOL"}
    assert "GRUPSUZ" in names
    assert ws.cell(hdr + 1 + gi, 1).font.bold
    assert ws.cell(hdr + 1 + gi, 14).value == "Ana Kimya 0 g · Tercih Kimya 50 g"
    assert ws.cell(hdr + 1 + gi, 13).value == "1 kartta sipariş"
    assert ws.cell(hdr + 2 + gi, 1).alignment.indent == 2
    # Sipariş listesi kart başına: grubun siparişli üyesi görünür
    vals = [c.value for row in wb["Sipariş Listesi"].iter_rows() for c in row if c.value]
    assert "CETEARYL ALCOHOL" in vals
    # Reçete etkisi sayfası kırılmaz; PDF üretilir
    pdf = authed_client.get(f"{_GAPS}/export?format=pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    # Seçim modu (ids) kart başı
    r = authed_client.get(f"{_GAPS}/export?format=xlsx&ids={a.id},{solo.id}")
    ws = load_workbook(BytesIO(r.content))["Eksikler"]
    hdr = next(rw for rw in range(1, ws.max_row + 1) if ws.cell(rw, 1).value == "Hammadde")
    names = [ws.cell(rw, 1).value for rw in range(hdr + 1, ws.max_row + 1)]
    assert set(names) == {"STEARİL ALKOL", "GRUPSUZ"}
    # ...ama gruptaki kart grup bağlamını taşır (ekrandaki grup satırıyla çelişmesin)
    notes = {ws.cell(rw, 1).value: ws.cell(rw, 14).value for rw in range(hdr + 1, ws.max_row + 1)}
    assert notes["STEARİL ALKOL"] == ("«STEARİL ALKOL» grubu: 50 g / min 200 g · Kritik altı"
                                      " — Ana Kimya 0 g · Tercih Kimya 50 g")
    assert notes["GRUPSUZ"] in (None, "")
    pdf = authed_client.get(f"{_GAPS}/export?format=pdf&ids={a.id},{solo.id}")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    from core.stock_gaps import _story, assemble
    from reportlab.platypus import Table
    rep = assemble(db, "cosmetics", item_ids=[a.id])
    tbl = next(f for f in _story(rep) if isinstance(f, Table))
    cells = [str(getattr(c, "text", c)) for row in tbl._cellvalues for c in row]
    assert any("«STEARİL ALKOL» grubu: 50 g · Kritik altı)" in c for c in cells)


def _near_min_group(db):
    """Toplam 999,6 g, grup min 1.000 g → kritik; eksik yalnız 0,4 g."""
    g = _grp(db, "SETİL ALKOL")
    a = _card(db, "SETİL ALKOL", g, stock=600, mn=1000, sup=_sup(db, "Bir Kimya"))
    b = _card(db, "CETYL ALCOHOL", g, stock=399.6, mn=500, sup=_sup(db, "İki Kimya"))
    return g, a, b


def test_gaps_group_texts_exact_near_min(authed_client: TestClient, db_session):
    """amount_short 999,6 g / 1.000 g'ı aynı, 0,4 g eksiği "0 g" gösteriyordu —
    grup satırı, seçim notu, PDF ve Excel notu kesin metin kullanır."""
    db = db_session
    g, a, b = _near_min_group(db)
    db.commit()
    (row,) = authed_client.get(_GAPS).json()["rows"]
    assert row["kind"] == "group" and row["status"] == "kritik"
    assert row["current_text"] == "999,6 g" and row["min_text"] == "1.000 g"
    assert row["deficit_text"] == "0,4 g" and row["current_text"] != row["min_text"]
    # Kart dökümü kısa metinde kalır
    assert row["breakdown_text"] == "Bir Kimya 600 g · İki Kimya 400 g"
    from reportlab.platypus import Table
    from core.stock_gaps import _story, assemble, group_note_text
    rep = assemble(db, "cosmetics")
    tbl = next(f for f in _story(rep) if isinstance(f, Table))
    cells = [str(getattr(c, "text", c)) for r in tbl._cellvalues for c in r]
    assert "999,6 g / 1.000 g" in cells and "0,4 g" in cells
    # Seçim modu: kart başı satır + kesin grup notu (Excel "Grup dökümü", PDF ad hücresi)
    rep = assemble(db, "cosmetics", item_ids=[a.id])
    note = rep["rows"][0]["group"]
    assert (note["total_text"], note["min_text"]) == ("999,6 g", "1.000 g")
    assert group_note_text(note).startswith("«SETİL ALKOL» grubu: 999,6 g / min 1.000 g · Kritik altı")
    tbl = next(f for f in _story(rep) if isinstance(f, Table))
    cells = [str(getattr(c, "text", c)) for r in tbl._cellvalues for c in r]
    assert any("«SETİL ALKOL» grubu: 999,6 g · Kritik altı)" in c for c in cells)


def test_dashboard_group_texts_exact_near_min(authed_client: TestClient, db_session):
    db = db_session
    _near_min_group(db)
    db.commit()
    d = authed_client.get("/api/dashboard/stats").json()
    (crit,) = [r for r in d["critical_stock_list"] if r["name"] == "SETİL ALKOL"]
    sg = crit["group"]
    assert (sg["total_text"], sg["min_text"], sg["deficit_text"]) == ("999,6 g", "1.000 g", "0,4 g")
    assert crit["deficit"] == pytest.approx(0.4)


def test_gaps_export_old_13_columns_when_rule_off(authed_client: TestClient, db_session):
    db = db_session
    _short_group(db)
    _card(db, "GRUPSUZ", stock=0, mn=10)
    db.commit()
    _set(db, "0")
    r = authed_client.get(f"{_GAPS}/export?format=xlsx")
    ws = load_workbook(BytesIO(r.content))["Eksikler"]
    hdr = next(rw for rw in range(1, ws.max_row + 1) if ws.cell(rw, 1).value == "Hammadde")
    header = [c.value for c in ws[hdr] if c.value is not None]
    assert len(header) == 13 and header[-1] == "Sipariş"
    assert all(ws.cell(rw, 14).value is None for rw in range(hdr, ws.max_row + 1))


def test_gaps_domain_isolation(authed_client: TestClient, db_session):
    db = db_session
    _short_group(db)
    sg = _grp(db, "SUPP", domain="supplement")
    _card(db, "SA", sg, stock=0, mn=10, domain="supplement")
    _card(db, "SB", sg, stock=0, mn=10, domain="supplement")
    db.commit()
    assert [r["name"] for r in authed_client.get(_GAPS).json()["rows"]] == ["STEARİL ALKOL"]
    authed_client.post("/api/domain/switch", json={"domain": "supplement"}, headers=_HDR)
    assert [r["name"] for r in authed_client.get(_GAPS).json()["rows"]] == ["SUPP"]


# ─── Dashboard ─────────────────────────────────────────────────────────────

def test_dashboard_group_rule_on_and_off(authed_client: TestClient, db_session):
    db = db_session
    g, a, b = _short_group(db)                                  # kritik grup, a sıfır
    h = _grp(db, "YETERLİ GRUP")
    hz = _card(db, "Y SIFIR", h, stock=0, mn=100)
    _card(db, "Y DOLU", h, stock=5000, mn=200)
    z = _grp(db, "BOŞ GRUP")
    _card(db, "Z A", z, stock=0, mn=10)
    _card(db, "Z B", z, stock=0)
    _card(db, "GRUPSUZ SIFIR", stock=0)
    db.commit()
    d = authed_client.get("/api/dashboard/stats").json()
    # sıfır: BOŞ GRUP (bir kez) + grupsuz; STEARİL a ve Y SIFIR grubu dolu → düşer
    assert d["zero_stock_raw_count"] == 2
    crit = {r["name"]: r for r in d["critical_stock_list"]}
    assert set(crit) == {"STEARİL ALKOL", "BOŞ GRUP"}
    sg = crit["STEARİL ALKOL"]["group"]
    assert crit["STEARİL ALKOL"]["id"] is None and sg["member_count"] == 2
    assert (sg["total_text"], sg["min_text"], sg["deficit_text"]) == ("50 g", "200 g", "150 g")
    assert sg["breakdown_text"] == "Ana Kimya 0 g · Tercih Kimya 50 g"
    assert d["critical_stock_count"] == 2
    # Ayar kapalı → kart başı (a, b, hz, Z A)
    _set(db, "0")
    d = authed_client.get("/api/dashboard/stats").json()
    assert d["zero_stock_raw_count"] == 5
    assert {r["id"] for r in d["critical_stock_list"]} >= {a.id, b.id, hz.id}
    assert all(r["group"] is None for r in d["critical_stock_list"])
    assert d["critical_stock_count"] == 4


def test_dashboard_page_renders_group_badge_code(authed_client: TestClient):
    r = authed_client.get("/")
    assert r.status_code == 200 and "grp-badge" in r.text


# ─── Push bildirimi ────────────────────────────────────────────────────────

def test_item_edit_push_suppressed_when_group_healthy(authed_client: TestClient, db_session,
                                                       monkeypatch):
    import routers.inventory as inv
    sent = []
    monkeypatch.setattr(inv, "notify_low_stock", lambda *a: sent.append(a))
    db = db_session
    g = _grp(db, "XHANTAN")
    a = _card(db, "XHANTAN A", g, stock=100)
    _card(db, "XHANTAN B", g, stock=5000)
    db.commit()
    body = {"name": "XHANTAN A", "category": "Hammadde", "unit": "g", "min_stock": 500}
    assert authed_client.put(f"/api/items/{a.id}", headers=_HDR, json=body).status_code == 200
    assert sent == []                                            # grup 5.100 g > 500 g
    _set(db, "0")
    body["min_stock"] = 600
    assert authed_client.put(f"/api/items/{a.id}", headers=_HDR, json=body).status_code == 200
    assert sent == [("XHANTAN A", 100.0, 600.0, "g")]


def test_production_push_uses_group_rule(authed_client: TestClient, db_session, monkeypatch):
    import routers.production as prod
    sent = []
    monkeypatch.setattr(prod, "notify_low_stock", lambda *a: sent.append(a))
    db = db_session
    g = _grp(db, "ALKOL")
    rc = _card(db, "ALKOL A", g, stock=100, mn=200)
    alt = _card(db, "ALKOL B", g, stock=1000)
    tgt = _card(db, "KREM", stock=0, unit="adet", category="Bitmiş Ürün")
    rec = Recipe(name="Krem", output_quantity=1, output_unit="adet", target_item_id=tgt.id,
                 waste_percentage=0)
    db.add(rec)
    db.flush()
    db.add(RecipeIngredient(recipe_id=rec.id, item_id=rc.id, quantity=2.0, unit="g"))
    db.commit()
    r = authed_client.post("/api/production", headers=_HDR, json={
        "recipe_id": rec.id, "produced_quantity": 10,
        "ingredient_sources": {str(rc.id): [{"item_id": rc.id, "quantity": 20}]}})
    assert r.status_code == 201, r.text
    assert sent == []                                            # grup 1.080 g > 200 g
    _set(db, "0")
    r = authed_client.post("/api/production", headers=_HDR,
                           json={"recipe_id": rec.id, "produced_quantity": 10})
    assert r.status_code == 201, r.text
    assert sent == [("ALKOL A", 60.0, 200.0, "g")]
    assert alt.id                                                # kart dokunulmadı


def test_quotation_confirm_push_uses_group_rule(authed_client: TestClient, db_session, monkeypatch):
    import routers.b2b as b2b
    from database import Quotation, QuotationItem
    sent = []
    monkeypatch.setattr(b2b, "notify_low_stock", lambda *a: sent.append(a))
    db = db_session
    g = _grp(db, "AMBALAJ GRUBU")
    a = _card(db, "Kavanoz A", g, stock=10, mn=8, unit="adet", category="Ambalaj", pkg_type="kavanoz")
    _card(db, "Kavanoz B", g, stock=100, unit="adet", category="Ambalaj", pkg_type="kavanoz")

    def _quote(no):
        q = Quotation(quote_number=no, customer_name="Müşteri", total_amount=10, status="DRAFT",
                      domain="cosmetics")
        db.add(q)
        db.flush()
        db.add(QuotationItem(quotation_id=q.id, item_id=a.id, quantity=1, unit_price_foreign=1,
                             line_total=1))
        db.commit()
        return q.id

    q1 = _quote("TQ-1")
    assert authed_client.post(f"/api/quotations/{q1}/confirm", headers=_HDR).status_code == 200
    assert sent == []                                            # grup 109 adet > 8
    _set(db, "0")
    q2 = _quote("TQ-2")
    assert authed_client.post(f"/api/quotations/{q2}/confirm", headers=_HDR).status_code == 200
    assert sent == [("Kavanoz A", 8.0, 8.0, "adet")]


def test_shopify_low_stock_after_group_rule(db_session):
    from core.shopify import _low_stock_after
    db = db_session
    g = _grp(db, "BİTMİŞ GRUP")
    _card(db, "Ürün A", g, stock=1, mn=5, unit="adet", category="Bitmiş Ürün", barcode="868A")
    _card(db, "Ürün B", g, stock=50, unit="adet", category="Bitmiş Ürün", barcode="868B")
    _card(db, "Ürün C", stock=2, mn=5, unit="adet", category="Bitmiş Ürün", barcode="868C")
    db.commit()
    order = {"lines": [{"barcode": "868A"}, {"barcode": "868C"}]}
    assert _low_stock_after(db, order) == [("Ürün C", 2.0, 5.0, "adet")]
    _set(db, "0")
    assert sorted(_low_stock_after(db, order)) == [("Ürün A", 1.0, 5.0, "adet"),
                                                   ("Ürün C", 2.0, 5.0, "adet")]


# ─── Asistan ───────────────────────────────────────────────────────────────

def test_assistant_item_detail_group_total_and_gaps_group_row(db_session):
    import json
    from core.assistant.tools import ToolContext, run_tool
    from database import User
    db = db_session
    g, a, b = _xhantan(db)
    sg, sa, sb = _short_group(db)
    db.commit()
    ctx = ToolContext.from_user(db.query(User).filter_by(username="dogukan").one(), "cosmetics", db)
    p = json.loads(run_tool(ctx, "item_detail", {"item_id": a.id}).content)
    grp = p["group"]
    assert grp["group_total"] == pytest.approx(24868.12) and grp["group_unit"] == "g"
    assert grp["group_total_text"] == "24,9 kg" and grp["counted"] is True
    assert grp["breakdown"] == "HAMMADDESEPETİ 180 g · YİĞİTOGLU KİMYA 24,7 kg"
    assert grp["rule_active"] is True and [m["id"] for m in grp["members"]] == [b.id]
    p = json.loads(run_tool(ctx, "stock_gaps", {"status": None, "q": None}).content)
    (row,) = p["rows"]
    assert row["kind"] == "group" and row["name"] == "STEARİL ALKOL"
    assert row["total_text"] == "50 g" and {m["item_id"] for m in row["members"]} == {sa.id, sb.id}
    assert "group_rule" in p
