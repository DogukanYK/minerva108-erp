"""
P2 — Üretimde "Hangi tedarikçiden?" (core/production_plan.py +
POST /api/production/preview + start_production `ingredient_sources`).

Kritik sözleşmeler:
  • Grup/alternatif yokken davranış BİREBİR eski: Output'lar, notlar, lotlar,
    döküm; önizleme satırları = expand_recipe.
  • Reçete kartının "aynı malzeme" grubundaki başka aktif kartta stok varsa
    satır seçim ister — kaynak gönderilmezse 400 `source_choice_required` ve
    HİÇBİR ŞEY yazılmaz.  Öneri: bitirilecek/alma → reçete kartı → tercih →
    diğer.
  • Bölme: her kart ayrı Output; toplam = brüt; not "Kaynak kart" taşır ve
    "Üretim Lot: X" ile biter; döküm recipe_item_id ≠ item_id.
  • Kart başına TOPLAM stok kapısı (çift satır taşması), aynı lot iki satıra
    verilmez, seçili lot kısaysa 400.
  • Kill switch kapalı → eski davranış.
  • Bölünmüş üretimin iptali her kartı/lotu iade eder; föy reçete sonradan
    düzenlense de değişmez.
"""
import io
import queue
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta

import bcrypt
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.orm import Session

from core.consumption import expand_recipe, load_recipe_recs
from core.production_plan import CFG_ENABLED
from core.production_plan import plan as plan_production
from core.snapshots import compute_stock_at
from database import (AppSetting, Inventory, Item, MaterialGroup, MaterialSupplierPref,
                      ProductionConsumption, ProductionHistory, Recipe, RecipeIngredient,
                      SessionLocal, Supplier, Transaction, User)

_HDR = {"Origin": "http://testserver"}


# ─── Kurulum ────────────────────────────────────────────────────────────────

def _sup(db, name, status="normal", domain="cosmetics"):
    s = Supplier(name=name, purchase_status=status, domain=domain,
                 status_reason="küçük satıcı" if status == "phase_out" else None)
    db.add(s); db.flush()
    return s


def _card(db, name, *, stock, supplier=None, unit="g", category="Hammadde",
          domain="cosmetics", group=None, active=True, **kw):
    it = Item(name=name, sku=f"sku-{name}", category=category, unit=unit,
              current_stock=stock, supplier_id=supplier.id if supplier else None,
              domain=domain, material_group_id=group.id if group else None,
              is_active=active, **kw)
    db.add(it); db.flush()
    return it


def _lot(db, item, lot, qty, *, supplier=None, age_days=0, sample=False):
    inv = Inventory(item_id=item.id, supplier_id=supplier.id if supplier else None,
                    lot_number=lot, quantity=qty, status="APPROVED", is_sample=sample,
                    domain=item.domain,
                    created_at=datetime.utcnow() - timedelta(days=age_days))
    db.add(inv); db.flush()
    return inv


def _world(db, *, group=True, rc_stock=100.0, po_stock=30.0, pref_stock=50.0):
    """Stearil alkol üç tedarikçide (reçete kartı + bitirilecek + tercih) +
    grupsuz su + şişe + TR/EN etiket.  Reçete: 1 adet = 2 g alkol + 5 ml su +
    1 şişe + 1 etiket, fire %10."""
    s_main = _sup(db, "Ana Kimya")
    s_po = _sup(db, "Küçük Satıcı", "phase_out")
    s_pref = _sup(db, "Tercih Kimya", "preferred")
    g = None
    if group:
        g = MaterialGroup(name="Stearil alkol", domain="cosmetics")
        db.add(g); db.flush()
    rc = _card(db, "STEARİL ALKOL", stock=rc_stock, supplier=s_main, group=g)
    po = _card(db, "SETİL STEARİL ALKOL", stock=po_stock, supplier=s_po, group=g)
    pref = _card(db, "CETEARYL ALCOHOL", stock=pref_stock, supplier=s_pref, group=g)
    water = _card(db, "Saf Su", stock=1000, unit="ml")
    bottle = _card(db, "Şişe 100", stock=1000, unit="adet", category="Ambalaj", pkg_type="şişe")
    ltr = _card(db, "Etiket TR", stock=1000, unit="adet", category="Ambalaj", pkg_type="etiket",
                language="TR", label_group="KRM")
    len_ = _card(db, "Etiket EN", stock=1000, unit="adet", category="Ambalaj", pkg_type="etiket",
                 language="EN", label_group="KRM")
    lots = {}
    if rc_stock:
        lots["rc1"] = _lot(db, rc, "RC-1", rc_stock * 0.6, supplier=s_main, age_days=20)
        lots["rc2"] = _lot(db, rc, "RC-2", rc_stock * 0.4, supplier=s_main, age_days=5)
    if po_stock:
        lots["po1"] = _lot(db, po, "PO-1", po_stock / 3, supplier=s_po, age_days=60)
        lots["po2"] = _lot(db, po, "PO-2", po_stock * 2 / 3, supplier=s_po, age_days=40)
    if pref_stock:
        lots["pf1"] = _lot(db, pref, "PF-1", pref_stock, supplier=s_pref, age_days=10)
    tgt = _card(db, "Minerva Kaynak Krem 50 ml", stock=0, unit="adet", category="Bitmiş Ürün")
    rec = Recipe(name="Kaynak Krem", output_quantity=1, output_unit="adet",
                 target_item_id=tgt.id, waste_percentage=10)
    db.add(rec); db.flush()
    for it, q, ph in ((rc, 2.0, "A"), (water, 5.0, "A"), (bottle, 1, None), (ltr, 1, None)):
        db.add(RecipeIngredient(recipe_id=rec.id, item_id=it.id, quantity=q, unit=it.unit, phase=ph))
    db.commit()
    return {"rec": rec.id, "rc": rc.id, "po": po.id, "pref": pref.id, "water": water.id,
            "bottle": bottle.id, "ltr": ltr.id, "len": len_.id, "tgt": tgt.id,
            "group": g.id if g else None, "s_main": s_main.id, "s_po": s_po.id,
            "s_pref": s_pref.id, "lots": {k: v.id for k, v in lots.items()}}


def _preview(client, rid, qty=10, **body):
    r = client.post("/api/production/preview",
                    json={"recipe_id": rid, "produced_quantity": qty, **body}, headers=_HDR)
    assert r.status_code == 200, r.text
    return r.json()


def _start(client, rid, qty=10, **body):
    return client.post("/api/production",
                       json={"recipe_id": rid, "produced_quantity": qty, **body}, headers=_HDR)


def _line(pv, key):
    return next(ln for ln in pv["lines"] if ln["key"] == str(key))


def _state(db):
    db.expire_all()
    return {
        "items": {i.id: round(i.current_stock or 0.0, 6) for i in db.query(Item).all()},
        "inv": {i.id: round(i.quantity or 0.0, 6) for i in db.query(Inventory).all()},
        "tx": db.query(Transaction).count(),
        "ph": db.query(ProductionHistory).count(),
        "pc": db.query(ProductionConsumption).count(),
    }


# ─── 1) Grup yok → bugünkü davranış ─────────────────────────────────────────

def test_no_group_preview_equals_expand_and_start_is_legacy(authed_client, db_session):
    w = _world(db_session, group=False)
    pv = _preview(authed_client, w["rec"], qty=10, label_language="EN")
    recs, items, sibs = load_recipe_recs(db_session, [w["rec"]], "cosmetics")
    exp = expand_recipe(recs[0], 10, items, sibs, label_mode="EN")
    assert [(ln["item_id"], ln["gross"]) for ln in pv["lines"]] == \
        [(ln.item_id, pytest.approx(ln.gross)) for ln in exp.lines]
    assert all(ln["needs_choice"] is False for ln in pv["lines"])
    assert pv["can_start"] is True and pv["blockers"] == [] and pv["needs_choice"] == []
    assert _line(pv, w["rc"])["options"][0]["is_recipe_card"] is True
    assert len(_line(pv, w["rc"])["options"]) == 1
    # Etiket kardeşe çözüldü, anahtar reçetedeki kart
    lbl = _line(pv, w["ltr"])
    assert lbl["item_id"] == w["len"] and lbl["kind"] == "label" and lbl["options"] == []

    r = _start(authed_client, w["rec"], qty=10, label_language="EN")
    assert r.status_code == 201, r.text
    lot = r.json()["lot_number"]
    db_session.expire_all()
    outs = (db_session.query(Transaction).filter(Transaction.transaction_type == "Output")
            .order_by(Transaction.id).all())
    base = "Üretim tüketimi — Reçete: Kaynak Krem"
    # Alkol 22 g FIFO: RC-1 (60) yeter → tek Output; su lotsuz; şişe + EN etiket
    assert [(t.item_id, t.lot_number, round(t.quantity, 6), t.notes) for t in outs] == [
        (w["rc"], "RC-1", 22.0,
         f"{base} | %10.0 fire dahil, brüt girdi | Dil: İngilizce | Tedarikçi: Ana Kimya | "
         f"Kaynak Lot: RC-1 | Üretim Lot: {lot}"),
        (w["water"], None, 55.0,
         f"{base} | %10.0 fire dahil, brüt girdi | Dil: İngilizce | Üretim Lot: {lot}"),
        (w["bottle"], None, 10.0, f"{base} | Dil: İngilizce | Üretim Lot: {lot}"),
        (w["len"], None, 10.0, f"{base} | Dil: İngilizce | Üretim Lot: {lot}"),
    ]
    snap = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.kind != "output")
            .order_by(ProductionConsumption.id).all())
    assert [(p.kind, p.recipe_item_id, p.item_id, p.lot_number, p.phase) for p in snap] == [
        ("raw", w["rc"], w["rc"], "RC-1", "A"), ("raw", w["water"], w["water"], None, "A"),
        ("packaging", w["bottle"], w["bottle"], None, None),
        ("label", w["ltr"], w["len"], None, None)]
    assert all("Kaynak kart" not in t.notes for t in outs)
    assert r.json()["substituted"] == []


def test_group_without_alternative_stock_does_not_ask(authed_client, db_session):
    """Grup var ama diğer kartlarda stok yok → soru YOK, eski akış."""
    w = _world(db_session, po_stock=0, pref_stock=0)
    pv = _preview(authed_client, w["rec"])
    ln = _line(pv, w["rc"])
    assert ln["needs_choice"] is False and ln["group"]["id"] == w["group"]
    assert {o["item_id"] for o in ln["options"]} == {w["rc"], w["po"], w["pref"]}
    r = _start(authed_client, w["rec"])
    assert r.status_code == 201, r.text
    db_session.expire_all()
    assert db_session.get(Item, w["rc"]).current_stock == pytest.approx(78)


# ─── 2) Öneri sırası ────────────────────────────────────────────────────────

def test_suggestion_phase_out_first_then_recipe_card(authed_client, db_session):
    w = _world(db_session)
    pv = _preview(authed_client, w["rec"], qty=20)          # brüt 44 g
    ln = _line(pv, w["rc"])
    assert ln["needs_choice"] is True and pv["needs_choice"] == [str(w["rc"])]
    assert [o["item_id"] for o in ln["options"]] == [w["po"], w["rc"], w["pref"]]
    po = ln["options"][0]
    assert po["supplier_status"] == "phase_out" and po["supplier_name"] == "Küçük Satıcı"
    assert [l["lot_number"] for l in po["lots"]] == ["PO-1", "PO-2"]          # FIFO
    assert [(s["item_id"], s["quantity"]) for s in ln["suggested"]] == [(w["po"], 30), (w["rc"], 14)]
    assert ln["shortfall"] == 0
    assert pv["can_start"] is False and pv["code"] == "source_choice_required"
    assert ln["answered"] is False and ln["chosen"] == []


def test_suggestion_preferred_before_other_when_recipe_card_empty(authed_client, db_session):
    w = _world(db_session, rc_stock=0)
    s_norm = _sup(db_session, "Normal Kimya")
    norm = _card(db_session, "CETYL STEARYL ALCOHOL", stock=100, supplier=s_norm,
                 group=db_session.get(MaterialGroup, w["group"]))
    _lot(db_session, norm, "NR-1", 100, supplier=s_norm, age_days=300)   # en eski — yine sonda
    db_session.commit()
    pv = _preview(authed_client, w["rec"], qty=50)          # brüt 110 g
    ln = _line(pv, w["rc"])
    assert [o["item_id"] for o in ln["options"]] == [w["po"], w["rc"], w["pref"], norm.id]
    assert [(s["item_id"], s["quantity"]) for s in ln["suggested"]] == \
        [(w["po"], 30), (w["pref"], 50), (norm.id, 30)]


def test_material_pref_avoid_and_rank(authed_client, db_session):
    """"Bu malzemede alma" kartı bitirilecek gibi önce; malzeme tercihi rank'i
    firma durumundan önce gelir."""
    w = _world(db_session, rc_stock=0, po_stock=0)
    g = db_session.get(MaterialGroup, w["group"])
    s_a = _sup(db_session, "A Firma"); s_b = _sup(db_session, "B Firma")
    a = _card(db_session, "STEARYL ALCOHOL A", stock=10, supplier=s_a, group=g)
    b = _card(db_session, "STEARYL ALCOHOL B", stock=10, supplier=s_b, group=g)
    db_session.add_all([
        MaterialSupplierPref(material_group_id=g.id, supplier_id=s_a.id, preference="avoid"),
        MaterialSupplierPref(material_group_id=g.id, supplier_id=s_b.id, preference="preferred",
                             rank=1)])
    db_session.commit()
    ln = _line(_preview(authed_client, w["rec"]), w["rc"])
    order = [o["item_id"] for o in ln["options"]]
    assert order.index(a.id) < order.index(w["rc"])         # alma = bitirilecek gibi önce
    assert order.index(w["po"]) < order.index(w["rc"])
    assert order.index(w["rc"]) < order.index(b.id)
    assert order.index(b.id) < order.index(w["pref"])       # malzeme tercihi > firma tercihi
    assert order[-1] == w["pref"]
    opt_a = next(o for o in ln["options"] if o["item_id"] == a.id)
    assert opt_a["material_pref"] == "avoid"


# ─── 3) Seçim yoksa 400, hiçbir şey yazılmaz ────────────────────────────────

def test_needs_choice_without_sources_400_nothing_written(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    r = _start(authed_client, w["rec"])
    assert r.status_code == 400
    body = r.json()
    assert body["code"] == "source_choice_required"
    assert "STEARİL ALKOL" in body["detail"]
    ln = next(x for x in body["lines"] if x["key"] == str(w["rc"]))
    assert ln["needs_choice"] is True and ln["suggested"]
    assert _state(db_session) == before


def test_legacy_lot_choice_does_not_answer_choice(authed_client, db_session):
    w = _world(db_session)
    r = _start(authed_client, w["rec"],
               ingredient_lot_choices={str(w["rc"]): w["lots"]["rc2"]})
    assert r.status_code == 400 and r.json()["code"] == "source_choice_required"


# ─── 4) Bölme ───────────────────────────────────────────────────────────────

def test_split_two_cards_outputs_notes_snapshot(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    src = {str(w["rc"]): [{"item_id": w["po"], "quantity": 20, "inventory_id": w["lots"]["po2"]},
                          {"item_id": w["rc"], "quantity": 24.000001}]}   # tolerans → 24'e oturur
    pv = _preview(authed_client, w["rec"], qty=20, ingredient_sources=src)
    assert pv["can_start"] is True, pv["blockers"]
    ln = _line(pv, w["rc"])
    assert ln["answered"] and ln["ok"]
    assert [(c["item_id"], c["quantity"]) for c in ln["chosen"]] == [(w["po"], 20), (w["rc"], 24)]
    assert ln["chosen"][0]["lots"] == [{"inventory_id": w["lots"]["po2"], "lot_number": "PO-2",
                                        "quantity": 20.0, "supplier_name": "Küçük Satıcı"}]
    assert [lt["lot_number"] for lt in ln["chosen"][1]["lots"]] == ["RC-1"]
    r = _start(authed_client, w["rec"], qty=20, ingredient_sources=src)
    assert r.status_code == 201, r.text
    lot = r.json()["lot_number"]
    assert r.json()["substituted"][0]["recipe_item_id"] == w["rc"]
    db_session.expire_all()
    outs = (db_session.query(Transaction)
            .filter(Transaction.transaction_type == "Output",
                    Transaction.item_id.in_([w["rc"], w["po"]]))
            .order_by(Transaction.id).all())
    assert sum(t.quantity for t in outs) == pytest.approx(44)
    po_tx = [t for t in outs if t.item_id == w["po"]]
    rc_tx = [t for t in outs if t.item_id == w["rc"]]
    assert [(t.lot_number, t.quantity) for t in po_tx] == [("PO-2", 20)]       # seçilen lot
    assert [(t.lot_number, t.quantity) for t in rc_tx] == [("RC-1", 24)]       # FIFO
    note = po_tx[0].notes
    assert "| Kaynak kart: SETİL STEARİL ALKOL (reçetede: STEARİL ALKOL) | Tedarikçi: Küçük Satıcı" in note
    assert note.startswith("Üretim tüketimi — Reçete: Kaynak Krem | ")
    assert note.endswith(f"Üretim Lot: {lot}")
    assert "Kaynak kart" not in rc_tx[0].notes and rc_tx[0].notes.endswith(f"Üretim Lot: {lot}")
    snap = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.recipe_item_id == w["rc"]).all())
    assert sorted((p.item_id, p.quantity, p.supplier_name) for p in snap) == sorted(
        [(w["po"], 20, "Küçük Satıcı"), (w["rc"], 24, "Ana Kimya")])
    after = _state(db_session)
    assert after["items"][w["po"]] == pytest.approx(before["items"][w["po"]] - 20)
    assert after["items"][w["rc"]] == pytest.approx(before["items"][w["rc"]] - 24)
    assert after["inv"][w["lots"]["po2"]] == pytest.approx(0)
    assert after["inv"][w["lots"]["po1"]] == pytest.approx(10)                 # dokunulmadı


def test_full_switch_to_alternative_card(authed_client, db_session):
    w = _world(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["pref"], "quantity": 22}]})
    assert r.status_code == 201, r.text
    db_session.expire_all()
    assert db_session.get(Item, w["rc"]).current_stock == pytest.approx(100)
    assert db_session.get(Item, w["pref"]).current_stock == pytest.approx(28)
    pc = (db_session.query(ProductionConsumption)
          .filter(ProductionConsumption.recipe_item_id == w["rc"]).one())
    assert pc.item_id == w["pref"] and pc.lot_number == "PF-1"


# ─── 5) Geçersiz kaynak ─────────────────────────────────────────────────────

def _extra_card(db, w, kind):
    g = db.get(MaterialGroup, w["group"])
    if kind == "outside":
        return _card(db, "GRUP DIŞI ALKOL", stock=100).id
    if kind == "unit":
        return _card(db, "STEARİL ALKOL KG", stock=100, unit="kg", group=g).id
    if kind == "domain":
        return _card(db, "SUPPLEMENT ALKOL", stock=100, domain="supplement", group=g).id
    if kind == "inactive":
        return _card(db, "PASİF ALKOL", stock=100, group=g, active=False).id
    if kind == "packaging":
        return _card(db, "AMBALAJ ALKOL", stock=100, category="Ambalaj", group=g).id
    raise AssertionError(kind)


@pytest.mark.parametrize("kind", ["outside", "unit", "domain", "inactive", "packaging"])
def test_invalid_source_card_400(authed_client, db_session, kind):
    w = _world(db_session)
    bad = _extra_card(db_session, w, kind)
    db_session.commit()
    pv = _preview(authed_client, w["rec"])
    assert bad not in [o["item_id"] for o in _line(pv, w["rc"])["options"]]
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": bad, "quantity": 22}]})
    assert r.status_code == 400 and r.json()["code"] == "invalid_source", r.text
    assert _state(db_session) == before


@pytest.mark.parametrize("entries", [
    [{"q": 10}, {"q": 10}],                 # toplam 20 ≠ 22
    [{"q": 30}],                            # fazla
    [{"q": -1}, {"q": 23}],                 # negatif
    [{"q": 0}, {"q": 22}],                  # sıfır
    [],                                     # boş
])
def test_invalid_source_quantities_400(authed_client, db_session, entries):
    w = _world(db_session)
    ids = [w["po"], w["rc"]]
    src = [{"item_id": ids[i % 2], "quantity": e["q"]} for i, e in enumerate(entries)]
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={str(w["rc"]): src})
    assert r.status_code == 400, r.text
    assert r.json()["code"] == "invalid_source"
    assert _state(db_session) == before


def test_invalid_source_nan_and_foreign_lot_and_unknown_key(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    # NaN — JSON gövdesi taşıyabilir (Python json); planlayıcı reddeder
    r = authed_client.post("/api/production", content=(
        '{"recipe_id": %d, "produced_quantity": 10, "ingredient_sources": '
        '{"%d": [{"item_id": %d, "quantity": NaN}]}}' % (w["rec"], w["rc"], w["po"])),
        headers={**_HDR, "Content-Type": "application/json"})
    assert r.status_code in (400, 422), r.text
    # Başka kartın lotu → bulunamadı
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 22, "inventory_id": w["lots"]["pf1"]}]})
    assert r.status_code == 400 and r.json()["code"] == "lot_not_found"
    assert "bulunamadı" in r.json()["detail"]
    # Reçetede olmayan anahtar + etiket satırına kaynak
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["rc"], "quantity": 22}], "999999": [
            {"item_id": w["rc"], "quantity": 1}]})
    assert r.status_code == 400 and r.json()["code"] == "invalid_source"
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["rc"], "quantity": 22}],
        str(w["ltr"]): [{"item_id": w["ltr"], "quantity": 10}]})
    assert r.status_code == 400 and r.json()["code"] == "invalid_source"
    # Numune lotu seçilemez
    smp = _lot(db_session, db_session.get(Item, w["po"]), "PO-NUM", 100, sample=True)
    db_session.commit()
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 22, "inventory_id": smp.id}]})
    assert r.status_code == 400 and r.json()["code"] == "lot_not_found"
    assert _state(db_session) == before


# ─── 6) Kart başına toplam kapı + lot rezervasyonu ──────────────────────────

def _two_row_recipe(db, w, *, card, q1, q2):
    """Aynı kartın İKİ satırda geçtiği reçete (A ve C fazı)."""
    rec = Recipe(name="Çift Satır", output_quantity=1, output_unit="adet",
                 target_item_id=w["tgt"], waste_percentage=0)
    db.add(rec); db.flush()
    db.add(RecipeIngredient(recipe_id=rec.id, item_id=card, quantity=q1, unit="g", phase="A"))
    db.add(RecipeIngredient(recipe_id=rec.id, item_id=card, quantity=q2, unit="g", phase="C"))
    db.commit()
    return rec.id


def test_double_row_overflow_regression(authed_client, db_session):
    """Eskiden kapı satır başınaydı: 60 + 60 > 100 stoğu geçiyordu."""
    w = _world(db_session, group=False)
    rid = _two_row_recipe(db_session, w, card=w["rc"], q1=60, q2=60)
    before = _state(db_session)
    r = _start(authed_client, rid, qty=1)
    assert r.status_code == 400 and r.json()["code"] == "insufficient_stock", r.text
    assert "yeterli stok yok" in r.json()["detail"]
    assert _state(db_session) == before
    # Yeterliyse tek satır (birleşik), faz 'A+C', lotlar iki kez verilmez
    rid2 = _two_row_recipe(db_session, w, card=w["rc"], q1=40, q2=40)
    pv = _preview(authed_client, rid2, qty=1)
    assert len(pv["lines"]) == 1 and pv["lines"][0]["gross"] == 80
    assert pv["lines"][0]["phase"] == "A+C"
    r = _start(authed_client, rid2, qty=1)
    assert r.status_code == 201, r.text
    after = _state(db_session)
    assert after["inv"][w["lots"]["rc1"]] == pytest.approx(0)
    assert after["inv"][w["lots"]["rc2"]] == pytest.approx(20)
    assert min(after["inv"].values()) >= 0
    # Her reçete satırı eskisi gibi kendi Output'u + fazı; lotlar sırayla dilimlenir
    pid = r.json()["production_id"]
    rows = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.production_id == pid,
                    ProductionConsumption.kind != "output")
            .order_by(ProductionConsumption.id).all())
    assert [(p.phase, p.lot_number, p.quantity) for p in rows] == [
        ("A", "RC-1", 40), ("C", "RC-1", 20), ("C", "RC-2", 20)]
    assert r.json()["consumed_hammadde"] == 2                  # reçete satırı sayısı
    sheet = authed_client.get(f"/api/production/{pid}").json()
    assert [(i["phase"], i["gross"]) for i in sheet["ingredients"]] == [("A", 40), ("C", 40)]


def test_shared_alternative_card_combined_gate_and_lot_reservation(authed_client, db_session):
    """İki satır (iki ayrı reçete kartı, aynı grup) aynı alternatif kartı
    seçerse kapı TOPLAMA bakar; FIFO aynı lotu iki satıra vermez."""
    w = _world(db_session)
    g = db_session.get(MaterialGroup, w["group"])
    other = _card(db_session, "CETOSTEARYL ALCOHOL", stock=0, group=g)
    rec = Recipe(name="İki Kart", output_quantity=1, output_unit="adet",
                 target_item_id=w["tgt"], waste_percentage=0)
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=w["rc"], quantity=8, unit="g"))
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=other.id, quantity=8, unit="g"))
    db_session.commit()
    pv = _preview(authed_client, rec.id, qty=1)
    # Ortak havuz: önce reçete kartının satırı PO'dan 8 alır, ikinci satıra 22 kalır
    assert [(s["item_id"], s["quantity"]) for s in _line(pv, w["rc"])["suggested"]] == [(w["po"], 8)]
    assert [(s["item_id"], s["quantity"]) for s in _line(pv, other.id)["suggested"]] == [(w["po"], 8)]
    # Birleşik kapı: 20 + 20 > 30
    before = _state(db_session)
    rec2 = Recipe(name="İki Kart Büyük", output_quantity=1, output_unit="adet",
                  target_item_id=w["tgt"], waste_percentage=0)
    db_session.add(rec2); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec2.id, item_id=w["rc"], quantity=20, unit="g"))
    db_session.add(RecipeIngredient(recipe_id=rec2.id, item_id=other.id, quantity=20, unit="g"))
    db_session.commit()
    before = _state(db_session)
    r = _start(authed_client, rec2.id, qty=1, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 20}],
        str(other.id): [{"item_id": w["po"], "quantity": 20}]})
    assert r.status_code == 400 and r.json()["code"] == "insufficient_stock", r.text
    assert "2 satırın toplamı" in r.json()["detail"]
    pc = next(c for c in r.json()["per_card"] if c["item_id"] == w["po"])
    assert pc["requested"] == 40 and pc["ok"] is False
    assert _state(db_session) == before
    # Sığan durumda: PO-1 (10) + PO-2 (20); satır 1 8 → PO-1, satır 2 8 → PO-1 2 + PO-2 6
    r = _start(authed_client, rec.id, qty=1, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 8}],
        str(other.id): [{"item_id": w["po"], "quantity": 8}]})
    assert r.status_code == 201, r.text
    after = _state(db_session)
    assert after["inv"][w["lots"]["po1"]] == pytest.approx(0)
    assert after["inv"][w["lots"]["po2"]] == pytest.approx(14)
    assert after["items"][w["po"]] == pytest.approx(14)
    rows = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.item_id == w["po"])
            .order_by(ProductionConsumption.id).all())
    assert [(r.recipe_item_id, r.lot_number, r.quantity) for r in rows] == [
        (w["rc"], "PO-1", 8), (other.id, "PO-1", 2), (other.id, "PO-2", 6)]


def test_same_lot_chosen_for_two_lines_cannot_exceed_lot(authed_client, db_session):
    w = _world(db_session)
    g = db_session.get(MaterialGroup, w["group"])
    other = _card(db_session, "CETOSTEARYL ALCOHOL", stock=0, group=g)
    rec = Recipe(name="Aynı Lot", output_quantity=1, output_unit="adet",
                 target_item_id=w["tgt"], waste_percentage=0)
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=w["rc"], quantity=6, unit="g"))
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=other.id, quantity=6, unit="g"))
    db_session.commit()
    before = _state(db_session)
    r = _start(authed_client, rec.id, qty=1, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 6, "inventory_id": w["lots"]["po1"]}],
        str(other.id): [{"item_id": w["po"], "quantity": 6, "inventory_id": w["lots"]["po1"]}]})
    assert r.status_code == 400 and r.json()["code"] == "lot_insufficient", r.text
    assert "yetersiz" in r.json()["detail"]
    assert _state(db_session) == before


def test_chosen_lot_short_on_alternative_card_400(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 22, "inventory_id": w["lots"]["po1"]}]})
    assert r.status_code == 400 and r.json()["code"] == "lot_insufficient"
    assert "PO-1" in r.json()["detail"] and "yetersiz" in r.json()["detail"]
    assert _state(db_session) == before


# ─── 7) Kill switch ─────────────────────────────────────────────────────────

def test_kill_switch_off_restores_old_behaviour(authed_client, db_session):
    w = _world(db_session)
    db_session.add(AppSetting(key=CFG_ENABLED, value="0")); db_session.commit()
    pv = _preview(authed_client, w["rec"])
    ln = _line(pv, w["rc"])
    assert pv["source_choice_enabled"] is False
    assert ln["needs_choice"] is False and [o["item_id"] for o in ln["options"]] == [w["rc"]]
    assert ln["group"] is None and pv["can_start"] is True
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 22}]})
    assert r.status_code == 400 and r.json()["code"] == "invalid_source"
    assert _state(db_session) == before
    r = _start(authed_client, w["rec"])
    assert r.status_code == 201, r.text
    after = _state(db_session)
    assert after["items"][w["rc"]] == pytest.approx(78)
    assert after["items"][w["po"]] == before["items"][w["po"]]
    # Açınca soru geri gelir
    db_session.get(AppSetting, CFG_ENABLED).value = "1"; db_session.commit()
    assert _line(_preview(authed_client, w["rec"]), w["rc"])["needs_choice"] is True


# ─── 8) Bölünmüş üretimin iptali ────────────────────────────────────────────

def test_cancel_split_production_restores_every_card_and_lot(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    t0 = datetime.utcnow()
    touched = [w["rc"], w["po"], w["pref"], w["water"], w["bottle"], w["ltr"], w["tgt"]]
    s0, _ = compute_stock_at(db_session, t0, Item.id.in_(touched))
    r = _start(authed_client, w["rec"], qty=20, witness_quantity=2, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 25},          # PO-1 10 + PO-2 15
                       {"item_id": w["pref"], "quantity": 10},
                       {"item_id": w["rc"], "quantity": 9}]})
    assert r.status_code == 201, r.text
    ph_id = r.json()["production_id"]
    mid = _state(db_session)
    assert mid["items"][w["po"]] == pytest.approx(5) and mid["items"][w["pref"]] == pytest.approx(40)
    pv = authed_client.get(f"/api/production/{ph_id}/cancel-preview").json()
    assert pv["can_cancel"] is True, pv["blockers"]
    rc = authed_client.post(f"/api/production/{ph_id}/cancel",
                            json={"reason": "Yanlış tedarikçi", "fingerprint": pv["fingerprint"]},
                            headers=_HDR)
    assert rc.status_code == 200, rc.text
    after = _state(db_session)
    for iid in touched:
        assert after["items"][iid] == pytest.approx(before["items"][iid]), iid
    for inv_id, q in before["inv"].items():
        assert after["inv"][inv_id] == pytest.approx(q), inv_id
    s_after, _ = compute_stock_at(db_session, t0, Item.id.in_(touched))
    for iid in touched:
        assert s_after[iid] == pytest.approx(s0[iid]), iid
    # Telafiler kaynak kartlara ve lotlara yazıldı
    adj = (db_session.query(Transaction)
           .filter(Transaction.transaction_type == "Adjustment",
                   Transaction.item_id.in_([w["po"], w["pref"], w["rc"]])).all())
    assert sorted((t.item_id, t.lot_number, t.quantity) for t in adj) == sorted([
        (w["po"], "PO-1", 10), (w["po"], "PO-2", 15), (w["pref"], "PF-1", 10),
        (w["rc"], "RC-1", 9)])


# ─── 9) Föy reçete sonradan değişse de aynı ─────────────────────────────────

def test_sheet_from_snapshot_survives_recipe_edit_and_delete(authed_client, db_session):
    w = _world(db_session)
    r = _start(authed_client, w["rec"], qty=10, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 12}, {"item_id": w["rc"], "quantity": 10}]})
    assert r.status_code == 201, r.text
    pid = r.json()["production_id"]
    sheet = authed_client.get(f"/api/production/{pid}").json()
    assert sheet["source"] == "snapshot"
    alc = sheet["ingredients"][0]
    assert alc["item_name"] == "STEARİL ALKOL" and alc["phase"] == "A"
    assert alc["gross"] == pytest.approx(22) and alc["net"] == pytest.approx(20)
    assert alc["substituted"] is True
    assert [(s["item_name"], s["lot_number"], s["quantity"], s["is_recipe_card"])
            for s in alc["sources"]] == [
        ("SETİL STEARİL ALKOL", "PO-1", 10, False), ("SETİL STEARİL ALKOL", "PO-2", 2, False),
        ("STEARİL ALKOL", "RC-1", 10, True)]
    # % bileşim hammadde netlerinden: alkol 20 / (20 + 50)
    assert alc["percent"] == pytest.approx(20 / 70 * 100, abs=1e-3)
    assert sheet["ingredients"][3]["item_name"] == "Etiket TR"
    assert sheet["totals"]["gross"] == pytest.approx(22 + 55 + 10 + 10)
    assert sheet["waste_percentage"] == 10

    # Reçeteyi düzenle (satırları sil + yeniden yarat, fire değiştir) → föy aynı
    rec = db_session.get(Recipe, w["rec"])
    for ing in list(rec.ingredients):
        db_session.delete(ing)
    rec.waste_percentage = 25
    db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=w["water"], quantity=99, unit="ml"))
    db_session.commit()
    assert authed_client.get(f"/api/production/{pid}").json() == sheet

    # Excel: ana sayfa lab formatı + "Tüketim Kaynakları"
    x = authed_client.get(f"/api/production/{pid}/export")
    assert x.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Üretim Föyü", "Tüketim Kaynakları"]
    src = wb["Tüketim Kaynakları"]
    rows = [[c.value for c in row] for row in src.iter_rows(min_row=2)]
    assert ["A", "STEARİL ALKOL", "SETİL STEARİL ALKOL (kaynak kart)", "Küçük Satıcı",
            "PO-1", 10, "g"] in rows

    # Reçete silinse de föy döner (eskiden 409)
    rec_id = rec.id
    ph = db_session.get(ProductionHistory, pid); ph.recipe_id = None; db_session.commit()
    db_session.query(RecipeIngredient).filter(RecipeIngredient.recipe_id == rec_id).delete()
    db_session.delete(db_session.get(Recipe, rec_id)); db_session.commit()
    again = authed_client.get(f"/api/production/{pid}")
    assert again.status_code == 200
    assert again.json()["ingredients"] == sheet["ingredients"]


# ─── 10) Önizleme sözleşmesi, yetki, panel, izlenebilirlik ──────────────────

def test_preview_shape_and_writes_nothing(authed_client, db_session):
    w = _world(db_session)
    before = _state(db_session)
    pv = _preview(authed_client, w["rec"], qty=10, label_language="EN")
    assert set(pv) >= {"recipe", "quantity", "label_language", "lines", "per_card", "totals",
                       "label_warnings", "can_start", "blockers", "warnings", "errors",
                       "source_choice_enabled", "needs_choice", "code"}
    ln = _line(pv, w["rc"])
    assert set(ln) >= {"key", "recipe_item_id", "recipe_item_name", "kind", "unit", "phase",
                       "net", "gross", "factor", "group", "needs_choice", "options",
                       "suggested", "chosen", "shortfall", "ok", "answered", "stock"}
    opt = ln["options"][0]
    assert set(opt) >= {"item_id", "name", "supplier_id", "supplier_name", "supplier_status",
                        "material_pref", "is_recipe_card", "stock", "unit", "lots"}
    assert set(opt["lots"][0]) >= {"inventory_id", "lot_number", "qty", "created_at",
                                   "supplier_name"}
    assert ln["net"] == pytest.approx(20) and ln["gross"] == pytest.approx(22)
    water = _line(pv, w["water"])
    assert water["chosen"][0]["uncovered"] == pytest.approx(55)       # lotsuz → toplam stoktan
    assert pv["totals"]["raw_gross"] == pytest.approx(22 + 55)
    assert _state(db_session) == before


def test_preview_rbac_and_domain(client, authed_client, db_session):
    w = _world(db_session)
    sup_rec = Recipe(name="Supplement Reçete", output_quantity=1, domain="supplement")
    db_session.add(sup_rec); db_session.commit()
    r = authed_client.post("/api/production/preview",
                           json={"recipe_id": sup_rec.id, "produced_quantity": 1}, headers=_HDR)
    assert r.status_code == 404
    u = User(username="uretim_staff",
             password_hash=bcrypt.hashpw(b"minerva123", bcrypt.gensalt()).decode(),
             full_name="Staff", role="Staff", is_active=True)
    db_session.add(u); db_session.commit()
    authed_client.post("/api/logout", headers=_HDR)
    assert client.post("/api/login", json={"username": "uretim_staff", "password": "minerva123"},
                       headers=_HDR).status_code == 200
    r = client.post("/api/production/preview",
                    json={"recipe_id": w["rec"], "produced_quantity": 1}, headers=_HDR)
    assert r.status_code == 403
    client.post("/api/logout", headers=_HDR)
    r = client.post("/api/production/preview",
                    json={"recipe_id": w["rec"], "produced_quantity": 1}, headers=_HDR)
    assert r.status_code == 401


def test_trace_lot_shows_recipe_card_for_substituted_source(authed_client, db_session):
    w = _world(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 12},
                       {"item_id": w["rc"], "quantity": 10}]})
    assert r.status_code == 201, r.text
    lot = r.json()["lot_number"]
    tr = authed_client.get(f"/api/traceability/lot/{lot}").json()
    alc = [x for x in tr["production"]["ingredients_consumed"] if x["item_id"] == w["po"]]
    assert alc and all(x["substituted"] and x["recipe_item_id"] == w["rc"]
                       and x["recipe_item_name"] == "STEARİL ALKOL" for x in alc)
    assert {x["supplier_name"] for x in alc} == {"Küçük Satıcı"}
    assert sum(x["quantity"] for x in alc) == pytest.approx(12)
    direct = [x for x in tr["production"]["ingredients_consumed"] if x["item_id"] == w["rc"]]
    assert sum(x["quantity"] for x in direct) == pytest.approx(10)
    assert all(not x["substituted"] for x in direct)


def test_label_sibling_is_resolved_inside_active_domain(authed_client, db_session):
    foreign = _card(db_session, "Supplement Etiket EN", stock=100, unit="adet",
                    category="Ambalaj", pkg_type="etiket", language="EN",
                    label_group="KRM", domain="supplement")
    w = _world(db_session, group=False)
    pv = _preview(authed_client, w["rec"], label_language="EN")
    assert _line(pv, w["ltr"])["item_id"] == w["len"]
    r = _start(authed_client, w["rec"], label_language="EN")
    assert r.status_code == 201, r.text
    db_session.expire_all()
    assert db_session.get(Item, foreign.id).current_stock == 100
    assert db_session.get(Item, w["len"]).current_stock == 990


def test_packaging_cannot_select_label_card_in_same_group(authed_client, db_session):
    w = _world(db_session, group=False)
    g = MaterialGroup(name="Ambalaj kontrolü", domain="cosmetics")
    db_session.add(g); db_session.flush()
    for iid in (w["bottle"], w["ltr"]):
        db_session.get(Item, iid).material_group_id = g.id
    db_session.commit()
    pv = _preview(authed_client, w["rec"])
    assert [o["item_id"] for o in _line(pv, w["bottle"])["options"]] == [w["bottle"]]
    before = _state(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["bottle"]): [{"item_id": w["ltr"], "quantity": 10}]})
    assert r.status_code == 400 and r.json()["code"] == "invalid_source"
    assert _state(db_session) == before


def test_snapshot_keeps_supplier_for_cards_without_inventory_lot(authed_client, db_session):
    w = _world(db_session, group=False)
    for iid in (w["water"], w["bottle"]):
        db_session.get(Item, iid).supplier_id = w["s_main"]
    db_session.commit()
    r = _start(authed_client, w["rec"])
    assert r.status_code == 201, r.text
    pid = r.json()["production_id"]
    rows = db_session.query(ProductionConsumption).filter(
        ProductionConsumption.production_id == pid,
        ProductionConsumption.item_id.in_([w["water"], w["bottle"]])).all()
    assert len(rows) == 2
    assert all(p.supplier_id == w["s_main"] and p.supplier_name == "Ana Kimya" for p in rows)
    sheet = authed_client.get(f"/api/production/{pid}").json()
    sources = [s for ing in sheet["ingredients"] for s in ing["sources"]
               if s["item_id"] in (w["water"], w["bottle"])]
    assert all(s["supplier_name"] == "Ana Kimya" for s in sources)
    tr = authed_client.get(f"/api/traceability/lot/{r.json()['lot_number']}").json()
    traced = [s for s in tr["production"]["ingredients_consumed"]
              if s["item_id"] in (w["water"], w["bottle"])]
    assert len(traced) == 2 and all(s["supplier_name"] == "Ana Kimya" for s in traced)


def test_explicit_lot_is_reserved_before_other_line_fifo(authed_client, db_session):
    w = _world(db_session)
    other = _card(db_session, "İkinci reçete kartı", stock=0,
                  group=db_session.get(MaterialGroup, w["group"]))
    rec = Recipe(name="Seçili lot rezervasyonu", output_quantity=1, target_item_id=w["tgt"])
    db_session.add(rec); db_session.flush()
    for iid in (w["rc"], other.id):
        db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=iid, quantity=8, unit="g"))
    db_session.commit()
    r = _start(authed_client, rec.id, qty=1, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 8}],
        str(other.id): [{"item_id": w["po"], "quantity": 8,
                        "inventory_id": w["lots"]["po1"]}]})
    assert r.status_code == 201, r.text
    db_session.expire_all()
    assert db_session.get(Inventory, w["lots"]["po1"]).quantity == 0
    assert db_session.get(Inventory, w["lots"]["po2"]).quantity == 14
    rows = db_session.query(ProductionConsumption).filter(
        ProductionConsumption.production_id == r.json()["production_id"],
        ProductionConsumption.kind == "raw").order_by(ProductionConsumption.id).all()
    assert [(p.recipe_item_id, p.lot_number, p.quantity) for p in rows] == [
        (w["rc"], "PO-1", 2), (w["rc"], "PO-2", 6), (other.id, "PO-1", 8)]


def test_split_rounding_preserves_exact_gross_total(authed_client, db_session):
    w = _world(db_session)
    r = _start(authed_client, w["rec"], ingredient_sources={
        str(w["rc"]): [{"item_id": w["rc"], "quantity": 22 / 3},
                       {"item_id": w["po"], "quantity": 22 / 3},
                       {"item_id": w["pref"], "quantity": 22 / 3}]})
    assert r.status_code == 201, r.text
    quantities = [pc.quantity for pc in db_session.query(ProductionConsumption).filter(
        ProductionConsumption.production_id == r.json()["production_id"],
        ProductionConsumption.recipe_item_id == w["rc"]).all()]
    assert round(sum(quantities), 6) == 22


def test_cross_domain_target_cannot_be_produced(authed_client, db_session):
    w = _world(db_session, group=False)
    foreign = _card(db_session, "Supplement hedef", stock=10, unit="adet",
                    category="Bitmiş Ürün", domain="supplement")
    db_session.get(Recipe, w["rec"]).target_item_id = foreign.id
    db_session.commit()
    before = _state(db_session)
    for endpoint in ("/api/production/preview", "/api/production"):
        r = authed_client.post(endpoint,
                               json={"recipe_id": w["rec"], "produced_quantity": 10}, headers=_HDR)
        assert r.status_code == 404 and r.json()["code"] == "item_not_found"
    assert _state(db_session) == before


def test_production_waiting_for_target_leaves_source_free_for_cancellation(db_session):
    """İptal hedefi tutarken başlatma kaynağı tutup ters yönde beklememeli.
    İki gerçek PostgreSQL transaction ile hedef önce yaratılır (küçük id)."""
    target = _card(db_session, "Kilit hedefi", stock=0, unit="adet", category="Bitmiş Ürün")
    raw = _card(db_session, "Kilit hammaddesi", stock=20)
    _lot(db_session, raw, "LOCK-RAW", 20)
    rec = Recipe(name="Kilit reçetesi", output_quantity=1, target_item_id=target.id)
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=raw.id, quantity=2, unit="g"))
    db_session.commit()
    target_id, raw_id, recipe_id = target.id, raw.id, rec.id
    pid_queue = queue.Queue()

    def starting_production():
        with SessionLocal() as starting:
            starting.execute(text("SET LOCAL lock_timeout = '5s'"))
            pid_queue.put(starting.execute(text("SELECT pg_backend_pid()")).scalar())
            recipe = starting.get(Recipe, recipe_id)
            result = plan_production(starting, recipe, 5, "TR", domain="cosmetics", lock=True)
            starting.rollback()
            return result.can_start

    with SessionLocal() as cancelling, ThreadPoolExecutor(max_workers=1) as executor:
        # İptalin id sıralı kart kilitlerinde hedefi aldığı anı temsil eder.
        cancelling.query(Item).filter(Item.id == target_id).with_for_update().one()
        future = executor.submit(starting_production)
        try:
            pid = pid_queue.get(timeout=3)
            deadline, blocked = time.monotonic() + 3, False
            while time.monotonic() < deadline:
                blocked = db_session.execute(text(
                    "SELECT EXISTS (SELECT 1 FROM pg_locks WHERE pid = :pid AND NOT granted)"
                ), {"pid": pid}).scalar()
                if blocked or future.done():
                    break
                time.sleep(0.01)
            assert blocked, "Üretim hedef kartın kilidini beklemeliydi."
            # Başlatma kaynağı erken kilitleseydi NOWAIT burada hata verirdi;
            # gerçek iptal de aynı kartı bekleyerek deadlock yaratırdı.
            cancelling.query(Item).filter(Item.id == raw_id).with_for_update(nowait=True).one()
        finally:
            cancelling.rollback()
        assert future.result(timeout=6) is True


def test_same_card_as_source_and_target_has_net_stock_change(authed_client, db_session):
    item = _card(db_session, "Yarı mamul yeniden üretim", stock=20)
    source_lot = _lot(db_session, item, "REWORK-RAW", 20)
    rec = Recipe(name="Yeniden üretim", output_quantity=1, target_item_id=item.id)
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=item.id, quantity=2, unit="g"))
    db_session.commit()
    r = _start(authed_client, rec.id, qty=5)
    assert r.status_code == 201, r.text
    db_session.expire_all()
    assert db_session.get(Item, item.id).current_stock == 15
    assert db_session.get(Inventory, source_lot.id).quantity == 10
    assert db_session.query(Inventory).filter(
        Inventory.item_id == item.id, Inventory.lot_number == r.json()["lot_number"]).one().quantity == 5


def test_multiphase_split_rounding_matches_stock_lots_ledger_and_cancel(authed_client, db_session):
    w = _world(db_session)
    rec = Recipe(name="Üç fazda bölünmüş tüketim", output_quantity=3,
                 target_item_id=w["tgt"], waste_percentage=0)
    db_session.add(rec); db_session.flush()
    for phase in ("A", "B", "C"):
        db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=w["rc"],
                                       quantity=2, unit="g", phase=phase))
    db_session.commit()
    before = _state(db_session)
    r = _start(authed_client, rec.id, qty=10, witness_quantity=2, ingredient_sources={
        str(w["rc"]): [{"item_id": w["po"], "quantity": 12},
                       {"item_id": w["rc"], "quantity": 8}]})
    assert r.status_code == 201, r.text
    pid = r.json()["production_id"]
    middle = _state(db_session)
    rows = db_session.query(ProductionConsumption).filter(
        ProductionConsumption.production_id == pid,
        ProductionConsumption.kind == "raw").order_by(ProductionConsumption.id).all()
    assert round(sum(p.quantity for p in rows), 6) == 20
    assert [(p.phase, p.lot_number, p.quantity) for p in rows] == [
        ("A", "PO-1", 6.666667), ("B", "PO-1", 3.333333),
        ("B", "PO-2", 2), ("B", "RC-1", 1.333334), ("C", "RC-1", 6.666666)]
    for iid, quantity in ((w["po"], 12), (w["rc"], 8)):
        assert before["items"][iid] - middle["items"][iid] == quantity
        assert round(sum(p.quantity for p in rows if p.item_id == iid), 6) == quantity
        assert round(sum(p.transaction.quantity for p in rows if p.item_id == iid), 6) == quantity
    for inv_id in (w["lots"]["po1"], w["lots"]["po2"], w["lots"]["rc1"]):
        assert round(sum(p.quantity for p in rows if p.inventory_id == inv_id), 6) == (
            before["inv"][inv_id] - middle["inv"][inv_id])
    pv = authed_client.get(f"/api/production/{pid}/cancel-preview").json()
    assert pv["can_cancel"] is True, pv["blockers"]
    cancelled = authed_client.post(f"/api/production/{pid}/cancel", headers=_HDR,
                                  json={"reason": "Yuvarlama kontrolü", "fingerprint": pv["fingerprint"]})
    assert cancelled.status_code == 200, cancelled.text
    after = _state(db_session)
    assert after["items"] == before["items"]
    assert all(after["inv"][iid] == quantity for iid, quantity in before["inv"].items())


def test_fifo_fractional_lots_and_uncovered_keep_full_consumption(authed_client, db_session):
    raw = _card(db_session, "Hassas lot hammaddesi", stock=1)
    first = _lot(db_session, raw, "MICRO-1", 0.3333334, age_days=2)
    second = _lot(db_session, raw, "MICRO-2", 0.3333334, age_days=1)
    target = _card(db_session, "Hassas lot hedefi", stock=0, unit="adet", category="Bitmiş Ürün")
    rec = Recipe(name="Hassas lot reçetesi", output_quantity=1, target_item_id=target.id)
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=raw.id, quantity=1, unit="g"))
    db_session.commit()
    r = _start(authed_client, rec.id, qty=1)
    assert r.status_code == 201, r.text
    db_session.expire_all()
    rows = db_session.query(ProductionConsumption).filter(
        ProductionConsumption.production_id == r.json()["production_id"],
        ProductionConsumption.kind == "raw").order_by(ProductionConsumption.id).all()
    assert [p.quantity for p in rows] == [0.333333, 0.333333, 0.333334]
    assert round(sum(p.quantity for p in rows), 6) == 1
    assert round(sum(p.transaction.quantity for p in rows), 6) == 1
    assert db_session.get(Item, raw.id).current_stock == 0
    assert db_session.get(Inventory, first.id).quantity >= 0
    assert db_session.get(Inventory, second.id).quantity >= 0
