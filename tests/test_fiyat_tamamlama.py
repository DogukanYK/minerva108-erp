"""
scripts/fiyat_tamamlama_20261016.py — fiyat tamamlama kademesi (Paraşüt →
Songül notları → fiyat formu).  AĞ YOK: Paraşüt satırları sentetik JSON:API
fikstüründen (`parse_bill` + `upsert_bill`) ya da düz dökümden; kur çekilmez.

Kilitlenenler:
  • Kuru çalıştırma üç aday sınıfı üretir (Paraşüt kesin / belirsiz, Songül
    yazılabilir / ezildi / belirsiz, forma kalan) ve HİÇBİR ŞEY yazmaz; hedef
    = aktif reçetedeki fiyatsız hammadde (fiyatlı, pasif reçete, ambalaj hariç).
  • Çalışma kitabı: 5 sayfa, giriş sütunları, E/H listesi `showErrorMessage`,
    ambalaj hatırlama.
  • E/H doğrulaması içe aktarma betiğiyle aynı ("X" hata); ambalaj / birim /
    kazanan-değil hataları → hiçbir şey yazılmaz (tek işlem).
  • Yazım: Paraşüt `approve_line` + Songül `apply(record_selection)`; idempotent;
    önce tüm kartlar → satırlar kilitlenir (arayüzle kilitlenme döngüsü yok),
    kilit çakışması `busy` → hiçbir şey yazılmaz.
  • Elle (manual) fiyat hiç ezilmez — araya giren elle satır yazımı durdurur.
  • --parasut-json: JSON:API dökümü (kimlikli, --commit ile tablolara yazılır)
    ve parasut_read.py düz dökümü (beyaz liste, supplier_contact atılır, kimliksiz
    → yazılamaz).
  • CLI kapıları; gerçek veri dosyasıyla duman testi.
    MINERVA_TEST_DB=minerva_test_n .venv/bin/pytest tests/test_fiyat_tamamlama.py -q
"""
import importlib.util
import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from core import parasut_purchases as PP
from database import (AdminAuditLog, Inventory, Item, MaterialGroup, ParasutProductPack, ParasutPurchaseLine,
                      Recipe, RecipeIngredient, Supplier, SupplierPrice)

_SPEC = importlib.util.spec_from_file_location(
    "_fiyat_tamamlama", Path(__file__).resolve().parent.parent / "scripts" / "fiyat_tamamlama_20261016.py")
ft = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(ft)
lab = ft.lab


# ─── Fikstürler (sentetik; iletişim alanları UYDURMA, saklanmamalı) ─────────

def _contact(cid, name):
    return {"id": str(cid), "type": "contacts", "attributes": {
        "name": name, "email": "gizli@example.com", "phone": "+90 555 000 00 00",
        "tax_number": "1234567890", "address": "Gizli Sk. No:1"}}


def _bill(pid, no, issue, supplier, lines, *, currency="TRL", description=None):
    """lines: [(detail_id, product_id, product_name, unit, qty, unit_price, vat, net, extra_attrs)]"""
    sid, sname = supplier
    details, included = [], [_contact(sid, sname)]
    for (did, prid, pname, unit, qty, up, vat, net, extra) in lines:
        details.append({"id": str(did), "type": "purchase_bill_details"})
        att = {"quantity": str(qty), "unit_price": str(up), "vat_rate": str(vat), "net_total": str(net),
               "description": pname}
        att.update(extra or {})
        det = {"id": str(did), "type": "purchase_bill_details", "attributes": att}
        if prid is not None:
            det["relationships"] = {"product": {"data": {"id": str(prid), "type": "products"}}}
            included.append({"id": str(prid), "type": "products", "attributes": {"name": pname, "unit": unit}})
        included.append(det)
    data = {"id": str(pid), "type": "purchase_bills", "attributes": {
        "item_type": "purchase_bill", "invoice_no": no, "issue_date": issue, "due_date": issue,
        "currency": currency, "exchange_rate": "1.0", "net_total": str(sum(l[7] for l in lines)),
        "gross_total": None, "description": description},
        "relationships": {"supplier": {"data": {"id": str(sid), "type": "contacts"}},
                          "details": {"data": details}}}
    return data, included


def _payload(bills):
    return {"data": [b[0] for b in bills], "included": [x for b in bills for x in b[1]],
            "meta": {"current_page": 1, "total_pages": 1, "total_count": len(bills)}}


NILKIM = (41, "NİLKİM TEKNİK KİMYA SAN.TİC.LTD.ŞTİ")
NATURALYA = (44, "NATURALYA KİMYA GIDA LTD. ŞTİ.")
ARAS = (43, "ARAS KARGO YURT İÇİ YURT DIŞI TAŞIMACILIK A.Ş.")
B_NILKIM = _bill(1001, "NIL2024000000554", "2024-07-29", NILKIM, [
    (5001, 777, "CAPRYLİC / CAPRİC TRİGLYCERİDE", "Adet", 1, 625.0, 20.0, 750.0, None)])
B_NAT = _bill(1004, "M012026000002134", "2026-08-14", NATURALYA, [
    (5005, 781, "LICORICE EXTRACT LIQ", "Adet", 5, 620.798, 20.0, 3724.79, None),
    (5006, 782, "GINSENG EXTRACT LIQ", "Adet", 5, 716.306, 20.0, 4297.84, None)])
B_ARAS = _bill(1003, "AAA2024000988833", "2024-05-23", ARAS, [
    (5004, 780, "ARAS STANDART KARGO", "Adet", 1, 50.73, 20.0, 60.88, None)], description="Kargo")


def rec(rid, malzeme, tedarikci, fiyat, *, tur="lab_tablo", tarih="2026-06-23", cur="EUR", unit="kg",
        pkg=25.0, kdv=None, oran=None, secili=False, guven="kesin", note="", sayfa=1, birim="g"):
    return {"id": rid, "kaynak_sayfa": sayfa, "tur": tur, "tarih": tarih, "malzeme": malzeme,
            "birim_kart": birim if tur == "lab_tablo" else None, "tedarikci": tedarikci, "fiyat": fiyat,
            "para_birimi": cur, "fiyat_birimi": unit, "ambalaj_kg": pkg, "kdv_dahil": kdv, "kdv_orani": oran,
            "secili": secili, "guven": guven, "not": note, "ihtiyac": None}


def _data(*recs):
    return {"kaynak": "test", "kapsam": "test", "kurallar": [], "kayitlar": list(recs)}


SONGUL = _data(
    rec("A1-02-S1", "SHEA YAĞI", "UMAYCHEM", 9.5),                                   # ezilir (belge daha yeni)
    rec("C-s12-1", "SHEA BUTTER", "UMAYCHEM", 10.0, tur="proforma", tarih="2026-07-01", pkg=20.0,
        kdv=False, oran=20.0, note="shea; lab tablosu karşılığı: SHEA YAĞI (A1-02)"),
    rec("A1-04-S1", "ALMOND OIL", "UMAYCHEM", 4.0),                                  # 2 kart → belirsiz
    rec("A1-09-S1", "FİYATLI KART", "UMAYCHEM", 3.0),                                # hedef değil (fiyatlı)
    rec("A1-10-S1", "BİLİNMEYEN MALZEME", "UMAYCHEM", 2.0),                          # eşleşmez, hedefte benzer yok
)


def _item(db, name, unit="g", **kw):
    kw.setdefault("category", "Hammadde")
    kw.setdefault("domain", "cosmetics")
    it = Item(name=name, unit=unit, **kw)
    db.add(it)
    db.flush()
    return it


def _seed(db):
    sups = {}
    for n in ("NİLKİM TEKNİK KİMYA", "UMAYCHEM", "NATURALYA"):
        s = Supplier(name=n, domain="cosmetics")
        db.add(s)
        db.flush()
        sups[n] = s
    its = {
        "cct": _item(db, "CAPRYLIC CAPRIC TRIGLYCERIDE", name_tr="KAPRİLİK KAPRİK TRİGLİSERİT", sku="CCT-1"),
        "lic": _item(db, "MEYAN KÖKÜ EKSTRAKTI", name_tr="LICORICE EXTRACT LIQ", sku="LIC-1"),
        "gin1": _item(db, "GINSENG EKSTRAKTI", name_tr="GINSENG EXTRACT LIQ", sku="GIN-1"),
        "gin2": _item(db, "GİNSENG EKSTRAKTI", name_tr="GINSENG EXTRACT LIQ", sku="GIN-2"),
        "shea": _item(db, "SHEA YAĞI", sku="SHEA-1"),
        "badem": _item(db, "BADEM YAĞI", sku="BAD-1"),
        "badem2": _item(db, "BADEM YAGI (NUMUNE)", sku="BAD-2"),
        "alla": _item(db, "ALLANTOIN", sku="ALL-1"),
        "alla_b": _item(db, "ALLANTOIN B", sku="ALL-2"),
        "priced": _item(db, "FİYATLI KART", sku="PRC-1"),
        "inactive_only": _item(db, "PASİF REÇETE KARTI", sku="PAS-1"),
        "jar": _item(db, "KAVANOZ 50 ML", unit="adet", category="Ambalaj", sku="JAR-1"),
        "krem": _item(db, "LICORICE KREM", unit="adet", category="Bitmiş Ürün", sku="KREM-1"),
    }
    grp = MaterialGroup(name="ALLANTOIN GRUBU", domain="cosmetics", is_active=True, source="manual")
    db.add(grp)
    db.flush()
    its["alla"].material_group_id = its["alla_b"].material_group_id = grp.id
    db.add(SupplierPrice(item_id=its["priced"].id, supplier_id=sups["UMAYCHEM"].id, supplier_name="UMAYCHEM",
                         unit_price=5.0, currency="USD", price_unit="kg", source="manual", domain="cosmetics"))
    db.add(SupplierPrice(item_id=its["alla_b"].id, supplier_id=sups["UMAYCHEM"].id, supplier_name="UMAYCHEM",
                         unit_price=7.0, currency="USD", price_unit="kg", source="manual", domain="cosmetics"))
    r1 = Recipe(name="KREM A", domain="cosmetics", is_active=True, target_item_id=its["krem"].id)
    r2 = Recipe(name="ESKİ KREM", domain="cosmetics", is_active=False)
    r3 = Recipe(name="KREM B", domain="cosmetics", is_active=True)
    db.add_all([r1, r2, r3])
    db.flush()
    for k in ("cct", "lic", "gin1", "gin2", "shea", "badem", "badem2", "alla", "priced", "jar"):
        db.add(RecipeIngredient(recipe_id=r1.id, item_id=its[k].id, quantity=1.0, unit="g"))
    db.add(RecipeIngredient(recipe_id=r2.id, item_id=its["inactive_only"].id, quantity=1.0, unit="g"))
    db.add(RecipeIngredient(recipe_id=r3.id, item_id=its["lic"].id, quantity=2.0, unit="g"))
    db.add(Inventory(item_id=its["alla"].id, supplier_id=sups["NATURALYA"].id, lot_number="ALL-LOT-1",
                     quantity=0.0, status="APPROVED", domain="cosmetics", created_at=datetime(2026, 9, 1)))
    db.commit()
    return its, sups


def _bills_to_db(db, *bills):
    out = {}
    for b in bills:
        parsed = PP.parse_bill(b[0], PP._incmap({"included": b[1]}))
        row, lines, _ = PP.upsert_bill(db, parsed)
        for ln in lines:
            out[ln.parasut_detail_id] = ln
    db.commit()
    return out


def _rows(plan, part):
    return {r["key"]: r for r in plan[part]["rows"]}


def _fill(path, sheet, values):
    """values: {kayıt: {sütun: değer}}"""
    wb = load_workbook(path)
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    ik = head.index(ft.KEY_COL) + 1
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, ik).value
        if rid in values:
            for col, v in values[rid].items():
                ws.cell(r, head.index(col) + 1, v)
    wb.save(path)


def _prices(db):
    db.expire_all()
    return db.query(SupplierPrice).order_by(SupplierPrice.id).all()


def _audits(db):
    return db.query(AdminAuditLog).filter(AdminAuditLog.action == ft.AUDIT_ACTION).all()


# ─── Kuru çalıştırma ────────────────────────────────────────────────────────

def test_dry_run_targets_and_three_candidate_classes(db_session: Session):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT, B_ARAS)
    plan = ft.build_plan(db_session, data=SONGUL)
    # Hedef: aktif reçetedeki fiyatsız hammadde — fiyatlı, pasif reçete, ambalaj, bitmiş ürün HARİÇ
    tg = plan["targets"]
    assert set(tg) == {its[k].id for k in ("cct", "lic", "gin1", "gin2", "shea", "badem", "badem2", "alla")}
    assert tg[its["lic"].id]["recipes"] == 2 and tg[its["cct"].id]["recipes"] == 1
    assert tg[its["alla"].id]["last_supplier"] == "NATURALYA" and tg[its["alla"].id]["last_date"] == date(2026, 9, 1)
    assert tg[its["alla"].id]["group_priced_via"] == f"{its['alla_b'].id} ALLANTOIN B"
    # Paraşüt: LIC kesin (taze), CCT kesin ama eski fatura, GINSENG iki kart → belirsiz, kargo listelenmez
    pr = _rows(plan, "parasut")
    lic_key, cct_key, gin_key = (f"P-{lines[d].id}" for d in (5005, 5001, 5006))
    assert pr[lic_key]["match"] == ft.EXACT and pr[lic_key]["card"]["id"] == its["lic"].id
    assert pr[lic_key]["stale"] is False and pr[lic_key]["sup"]["id"] == sups["NATURALYA"].id
    assert pr[lic_key]["unit"] == "kg" and pr[lic_key]["pack_qty"] is None and pr[lic_key]["input"]
    assert pr[cct_key]["match"] == ft.EXACT and pr[cct_key]["stale"] is True and "eski fatura" in pr[cct_key]["note"]
    assert pr[gin_key]["match"] == ft.AMBIG and {pr[gin_key]["card"]["id"]} | {
        int(x.split()[0]) for x in pr[gin_key]["others"].split("; ")} == {its["gin1"].id, its["gin2"].id}
    assert f"P-{lines[5004].id}" not in pr
    s = plan["summary"]
    assert s["parasut_off_target"] == 1 and s["parasut_expense"] == 1 and s["parasut_pending_total"] == 4
    assert s["parasut_exact_cards"] == 2 and s["parasut_ambig_cards"] == 1 and s["parasut_stale_rows"] == 1
    # Songül: C-s12-1 yazılabilir, A1-02-S1 ezildi, ALMOND belirsiz; fiyatlı kart / bilinmeyen listelenmez
    sr = _rows(plan, "songul")
    assert sr["C-s12-1"]["writable"] and sr["C-s12-1"]["status"] == "yazilacak"
    assert sr["C-s12-1"]["card"]["id"] == its["shea"].id and sr["C-s12-1"]["sup"]["id"] == sups["UMAYCHEM"].id
    assert sr["A1-02-S1"]["status"] == "ezildi" and not sr["A1-02-S1"]["writable"] and "C-s12-1" in sr["A1-02-S1"]["status_text"]
    assert sr["A1-04-S1"]["status"] == "belirsiz" and sr["A1-04-S1"]["match"] == ft.AMBIG
    assert {sr["A1-04-S1"]["card"]["id"]} | {int(x.split()[0]) for x in sr["A1-04-S1"]["others"].split("; ")} == {
        its["badem"].id, its["badem2"].id}
    assert "A1-09-S1" not in sr and "A1-10-S1" not in sr
    assert s["songul_writable_rows"] == 1 and s["songul_overridden"] == 1 and s["songul_ambig"] == 1
    # Form: karşılanmayanlar (eski faturalı CCT, belirsiz GINSENG/BADEM, ALLANTOIN) — LIC ve SHEA yok
    form = {t["item_id"]: t for t in plan["form"]}
    assert set(form) == {its[k].id for k in ("cct", "gin1", "gin2", "badem", "badem2", "alla")}
    assert "eski fatura" in form[its["cct"].id]["note"] and "grup üzerinden fiyatlı" in form[its["alla"].id]["note"]
    assert ft.AMBIG in form[its["gin1"].id]["note"] and "Songül notu" in form[its["badem"].id]["note"]
    assert s["covered"] == 2 and s["form"] == 6 and s["targets"] == 8 and s["group_priced"] == 1
    # Hiçbir şey yazılmadı
    assert len(_prices(db_session)) == 2 and not _audits(db_session)
    assert all(ln.status == "pending" for ln in db_session.query(ParasutPurchaseLine).all())


ORLIFE = (42, "ORLIFE GLOBAL İÇ VE DIŞ TİCARET A.Ş.")
B_NILKIM_KAR = _bill(1005, "NIL2026000000777", "2026-09-02", NILKIM, [
    (5007, 783, "KARNAUBA MUMU (CARNAUBA WAX)", "Adet", 1, 4250.0, 20.0, 4250.0, None)])
B_ORLIFE = _bill(1002, "ORA2026000000058", "2026-05-27", ORLIFE, [
    (5003, 779, "DUY.FRK-SE.07 AKGÜNLÜK UÇUÇU YAĞI (Boswellia serrata) 10 ml.", "Adet", 4, 249.96, 20.0, 799.87,
     {"discount_type": "percentage", "discount_value": "20.0"})], description="Hammadde")


def test_dry_run_real_parasut_names_reach_candidates_with_purchase_unit(db_session: Session, tmp_path):
    """Gerçek Paraşüt adları (kart adını İÇEREN uzun metin) TR adı boş kartlara
    'kesin' düşer — eskiden Nilkim KARNAUBA 40 puanda 'belirsiz', Orlife AKGÜNLÜK
    hiç aday değildi ('hedef dışı').  Hatırlanan ambalaj yoksa birim kartın ALIM
    birimi: ml kartı → 'l' (default_price_unit 'kg' verirdi), g kartı → 'kg'."""
    its, sups = _seed(db_session)
    orl = Supplier(name="ORLIFE GLOBAL", domain="cosmetics")
    db_session.add(orl)
    db_session.flush()
    kar = _item(db_session, "CARNAUBA WAX", unit="g")
    akg = _item(db_session, "AKGÜNLÜK UÇUCU YAĞI", unit="ml", supplier_id=orl.id)
    rec_ = db_session.query(Recipe).filter(Recipe.name == "KREM A").one()
    for it in (kar, akg):
        db_session.add(RecipeIngredient(recipe_id=rec_.id, item_id=it.id, quantity=1.0, unit=it.unit))
    db_session.commit()
    lines = _bills_to_db(db_session, B_NILKIM_KAR, B_ORLIFE)
    plan = ft.build_plan(db_session, data=SONGUL)
    assert {kar.id, akg.id} <= set(plan["targets"])
    pr = _rows(plan, "parasut")
    k, a = pr[f"P-{lines[5007].id}"], pr[f"P-{lines[5003].id}"]
    assert k["match"] == ft.EXACT and k["card"]["id"] == kar.id and k["score"] >= 80 and k["unit"] == "kg"
    assert a["match"] == ft.EXACT and a["card"]["id"] == akg.id and a["unit"] == "l"
    assert a["sup"]["id"] == orl.id and a["score"] >= 85                     # faturanın firmasının kartı +5
    assert a["discount_rate"] == pytest.approx(0.2) and a["pack_qty"] is None and a["net_preview"] is None
    # Fiyat formu aynı birimi önerir; iki kart Paraşüt'le karşılandı → formda yok
    assert plan["targets"][akg.id]["purchase_unit"] == "l"
    assert not {kar.id, akg.id} & {t["item_id"] for t in plan["form"]}
    # Çalışma kitabı: birim hücresi + iletişim verisi (fikstürdeki uydurma e-posta/telefon/vergi/adres) yok
    path = ft.build_xlsx(plan, tmp_path / "ft.xlsx")
    wb = load_workbook(path)
    ws = wb[ft.SHEETS[1]]
    head = [c.value for c in ws[1]]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows[f"P-{lines[5003].id}"][head.index(ft.UNIT_COL)] == "l"
    assert rows[f"P-{lines[5007].id}"][head.index(ft.UNIT_COL)] == "kg"
    text = " ".join(str(c) for w in wb.worksheets for r in w.iter_rows(values_only=True) for c in r if c is not None)
    for secret in ("gizli@example.com", "+90 555", "1234567890", "Gizli Sk."):
        assert secret not in text
    # Hatırlanan ambalaj varsa onun birimi kazanır (ml kartında 'kg' hatırlandıysa 'kg')
    db_session.add(ParasutProductPack(product_parasut_id=779, package_qty=0.01, price_unit="kg", item_id=akg.id))
    db_session.commit()
    a2 = _rows(ft.build_plan(db_session, data=SONGUL), "parasut")[f"P-{lines[5003].id}"]
    assert a2["unit"] == "kg" and a2["pack_qty"] == 0.01 and a2["net_preview"] == pytest.approx(249.96 * 0.8 / 0.01)


def test_workbook_structure_prefill_and_validation(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT, B_ARAS)
    db_session.add(ParasutProductPack(product_parasut_id=781, product_name="LICORICE", package_qty=5.0,
                                      price_unit="kg", item_id=its["lic"].id))
    db_session.commit()
    plan = ft.build_plan(db_session, data=SONGUL)
    path = ft.build_xlsx(plan, tmp_path / "ft.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == list(ft.SHEETS)
    ws = wb[ft.SHEETS[1]]
    head = [c.value for c in ws[1]]
    for col in (ft.KEY_COL, ft.APPROVE_COL, ft.CARD_COL, ft.PACK_COL, ft.UNIT_COL, ft.SUP_COL, "Net önizleme"):
        assert col in head
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    lic = rows[f"P-{lines[5005].id}"]
    assert lic[head.index(ft.CARD_COL)] == its["lic"].id and lic[head.index(ft.PACK_COL)] == 5.0
    assert lic[head.index(ft.UNIT_COL)] == "kg" and lic[head.index(ft.SUP_COL)] == sups["NATURALYA"].id
    assert lic[head.index("Net önizleme")] == pytest.approx(620.798 / 5) and lic[head.index(ft.APPROVE_COL)] is None
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref
    assert any(dv.showErrorMessage and dv.formula1 == '"E,H"' for dv in ws.data_validations.dataValidation)
    ws2 = wb[ft.SHEETS[2]]
    h2 = [c.value for c in ws2[1]]
    r2 = {r[0]: r for r in ws2.iter_rows(min_row=2, values_only=True)}
    assert r2["C-s12-1"][h2.index(ft.APPROVE_COL)] is None and r2["C-s12-1"][h2.index(ft.CARD_COL)] == its["shea"].id
    assert r2["C-s12-1"][h2.index(ft.SUP_COL)] == sups["UMAYCHEM"].id
    assert r2["A1-02-S1"][h2.index(ft.APPROVE_COL)] == ft.NOT_APPLICABLE
    assert r2["A1-04-S1"][h2.index(ft.APPROVE_COL)] == ft.NOT_APPLICABLE
    ws3 = wb[ft.SHEETS[3]]
    h3 = [c.value for c in ws3[1]]
    for col in ("Kart no", "Kart", "Birim", "Reçete sayısı", "Son alım tedarikçisi", ft.FORM_PRICE_COL,
                "Para birimi", "Fiyat birimi (kg/l/adet)", "Ambalaj", "Tedarikçi", "Tarih", "Not"):
        assert col in h3
    assert ws3.max_row == 1 + len(plan["form"])
    alla = {r[0]: r for r in ws3.iter_rows(min_row=2, values_only=True)}[its["alla"].id]
    assert alla[h3.index("Son alım tedarikçisi")] == "NATURALYA" and alla[h3.index("Tedarikçi")] == "NATURALYA"
    assert alla[h3.index("Fiyat birimi (kg/l/adet)")] == "kg" and alla[h3.index(ft.FORM_PRICE_COL)] is None
    assert wb[ft.SHEETS[4]].max_row > 10 and wb[ft.SHEETS[0]].max_row >= 5
    # Boş kitap: onay yok, hata yok
    appr = ft.read_workbook(path)
    assert appr == {"parasut": {}, "songul": {}, "form_filled": [], "errors": []}


# ─── Doğrulama ──────────────────────────────────────────────────────────────

def test_eh_validation_and_errors_block_everything(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT)
    lic_key, gin_key = f"P-{lines[5005].id}", f"P-{lines[5006].id}"
    plan = ft.build_plan(db_session, data=SONGUL)
    path = ft.build_xlsx(plan, tmp_path / "onay.xlsx")
    # "X" / "✓" E/H değil
    _fill(path, ft.SHEETS[1], {lic_key: {ft.APPROVE_COL: "X"}})
    _fill(path, ft.SHEETS[2], {"C-s12-1": {ft.APPROVE_COL: "✓"}})
    appr = ft.read_workbook(path)
    assert len(appr["errors"]) == 2 and all("yalnız E ya da H" in e for e in appr["errors"])
    # E ama ambalaj yok · birim uyumsuz · ezilen Songül satırına E · form dolu uyarısı
    _fill(path, ft.SHEETS[1], {lic_key: {ft.APPROVE_COL: "E", ft.PACK_COL: None},
                               gin_key: {ft.APPROVE_COL: "e", ft.PACK_COL: 1, ft.UNIT_COL: "adet"}})
    _fill(path, ft.SHEETS[2], {"C-s12-1": {ft.APPROVE_COL: "H"},
                               "A1-02-S1": {ft.APPROVE_COL: "E", ft.CARD_COL: its["shea"].id,
                                            ft.SUP_COL: sups["UMAYCHEM"].id}})      # ezilen satıra zorla E
    wb = load_workbook(path)
    ws = wb[ft.SHEETS[3]]
    ws.cell(2, [c.value for c in ws[1]].index(ft.FORM_PRICE_COL) + 1, 12.5)
    wb.save(path)
    appr = ft.read_workbook(path)
    assert not appr["errors"] and appr["parasut"][lic_key]["ok"] is True and appr["songul"]["C-s12-1"]["ok"] is False
    assert len(appr["form_filled"]) == 1
    res = ft.resolve(db_session, appr, data=SONGUL)
    errs = "\n".join(res["errors"])
    assert ft.PACK_COL in errs and ft.UNIT_COL in errs and "yazılacak kazanan değil" in errs
    assert res["warnings"] and "yazılmaz" in res["warnings"][0]
    res = ft.apply(db_session, appr, data=SONGUL)
    assert res["applied"] is False and len(res["errors"]) == 3
    assert len(_prices(db_session)) == 2 and not _audits(db_session)
    assert all(ln.status == "pending" for ln in db_session.query(ParasutPurchaseLine).all())
    # Olmayan kayıt / bilinmeyen anahtar / geçersiz kart
    bad = {"parasut": {"P-999999": {"ok": True, "row": 2, "item_id": its["lic"].id, "package": 5, "unit": "kg",
                                   "supplier_id": None},
                       lic_key: {"ok": True, "row": 3, "item_id": its["krem"].id, "package": 5, "unit": "kg",
                                 "supplier_id": None}},
           "songul": {"YOK-1": {"ok": True, "row": 2, "item_id": 1, "supplier_id": 1}},
           "form_filled": [], "errors": []}
    res = ft.resolve(db_session, bad, data=SONGUL)
    assert len(res["errors"]) == 3 and "veritabanında yok" in res["errors"][0]
    assert "aktif hammadde değil" in res["errors"][1] and "tam planda yok" in res["errors"][2]


# ─── Yazım ──────────────────────────────────────────────────────────────────

def test_commit_writes_expected_rows_and_is_idempotent(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT)
    lic_key, cct_key = f"P-{lines[5005].id}", f"P-{lines[5001].id}"
    plan = ft.build_plan(db_session, data=SONGUL)
    path = ft.build_xlsx(plan, tmp_path / "onay.xlsx")
    _fill(path, ft.SHEETS[1], {lic_key: {ft.APPROVE_COL: "E", ft.PACK_COL: 5},
                               cct_key: {ft.APPROVE_COL: "Evet", ft.PACK_COL: 25, ft.SUP_COL: sups["NİLKİM TEKNİK KİMYA"].id}})
    _fill(path, ft.SHEETS[2], {"C-s12-1": {ft.APPROVE_COL: "E"}})
    appr = ft.read_workbook(path)
    res = ft.apply(db_session, appr, data=SONGUL)
    assert res["applied"] is True and not res["errors"] and res["songul_written"] == 1
    assert [w["key"] for w in res["parasut_written"]] == [lic_key, cct_key]
    ps = _prices(db_session)
    assert len(ps) == 5                                   # 2 tohum (elle) + LIC + CCT + SHEA
    lic_sp = next(p for p in ps if p.item_id == its["lic"].id)
    assert lic_sp.source == "parasut" and lic_sp.unit_price == pytest.approx(620.798 / 5) and lic_sp.currency == "TRY"
    assert lic_sp.price_unit == "kg" and lic_sp.package_size == 5.0 and lic_sp.quoted_at == date(2026, 8, 14)
    assert lic_sp.supplier_id == sups["NATURALYA"].id and lic_sp.vat_included is False and lic_sp.vat_rate == 20.0
    assert lic_sp.created_by == ft.ACTOR and lic_sp.source_label.startswith("Paraşüt alış faturası M012026000002134")
    cct_sp = next(p for p in ps if p.item_id == its["cct"].id)
    assert cct_sp.unit_price == 25.0 and cct_sp.supplier_id == sups["NİLKİM TEKNİK KİMYA"].id
    shea_sp = next(p for p in ps if p.item_id == its["shea"].id)
    assert shea_sp.source == "proforma" and shea_sp.source_label.startswith(lab.LABEL_PREFIX)
    assert shea_sp.unit_price == 10.0 and shea_sp.package_size == 20.0 and shea_sp.supplier_id == sups["UMAYCHEM"].id
    assert shea_sp.created_by == lab.ACTOR and "[C-s12-1]" in shea_sp.note
    for d in (5005, 5001):
        db_session.refresh(lines[d])
        assert lines[d].status == "approved" and lines[d].decided_by == ft.ACTOR
    assert lines[5006].status == "pending"
    assert db_session.query(ParasutProductPack).filter(ParasutProductPack.product_parasut_id == 781).one().package_qty == 5.0
    aud = _audits(db_session)
    det = json.loads(aud[0].details)
    assert len(aud) == 1 and det["songul_written"] == 1 and len(det["parasut_written"]) == 2
    assert det["parasut_written"][0]["key"] == lic_key and det["actor"] == ft.ACTOR
    # Artık hedef değiller
    plan2 = ft.build_plan(db_session, data=SONGUL)
    assert its["lic"].id not in plan2["targets"] and its["shea"].id not in plan2["targets"]
    assert its["cct"].id not in plan2["targets"]
    before = {(p.item_id, p.unit_price, p.source) for p in _prices(db_session)}
    # İkinci çalıştırma aynı kitapla: Paraşüt "zaten yazılmış", lab satırı yenilenir, sayı değişmez
    res2 = ft.apply(db_session, appr, data=SONGUL)
    assert res2["applied"] is True and not res2["errors"]
    assert res2["parasut_written"] == [] and sorted(res2["parasut_done"]) == sorted([lic_key, cct_key])
    assert res2["songul_written"] == 1 and res2["songul_deleted"] == 1
    ps2 = _prices(db_session)
    assert len(ps2) == 5 and {(p.item_id, p.unit_price, p.source) for p in ps2} == before
    assert len(_audits(db_session)) == 2


def _lock_waiter(timeout=10.0) -> bool:
    """Test veritabanında satır kilidinde bekleyen oturum var mı (yoklama başına ayrı bağlantı)."""
    import time
    from sqlalchemy import text
    from database import engine
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        with engine.connect() as c:
            n = c.execute(text("SELECT count(*) FROM pg_stat_activity WHERE datname = current_database() "
                               "AND wait_event_type = 'Lock'")).scalar()
        if n:
            return True
        time.sleep(0.02)
    return False


def test_commit_prelocks_cards_so_concurrent_ui_approve_cannot_deadlock(db_session: Session, monkeypatch):
    """Doğrulayıcı bulgusu: --commit satırları tek tek onaylarken (L1→CCT, sonra
    L2→LIC) /fiyatlar'dan aynı anda gelen L1→LIC onayı LIC'i tutup L1'i
    bekliyordu → PostgreSQL kilitlenmesi (arayüz 500).  Betik artık yazımdan
    önce tüm kartları, sonra tüm satırları kilitler: arayüz LIC'te bekler, betik
    tamamlanır, arayüz 409 not_pending alır (L1 betiğin kartında kalır)."""
    import threading
    from database import SessionLocal
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT)
    l_cct, l_lic = lines[5001].id, lines[5005].id
    cct, lic = its["cct"].id, its["lic"].id
    appr = {"parasut": {
        f"P-{l_cct}": {"ok": True, "row": 2, "item_id": cct, "package": 25, "unit": "kg",
                       "supplier_id": sups["NİLKİM TEKNİK KİMYA"].id},
        f"P-{l_lic}": {"ok": True, "row": 3, "item_id": lic, "package": 5, "unit": "kg", "supplier_id": None}},
        "songul": {}, "form_filled": [], "errors": []}
    real = PP.approve_line
    first_done, go = threading.Event(), threading.Event()

    def paused(*a, **kw):                                # betik ilk satırı onayladıktan sonra bekler
        out = real(*a, **kw)
        if not first_done.is_set():
            first_done.set()
            go.wait(10)
        return out
    monkeypatch.setattr(PP, "approve_line", paused)
    res = {}

    def cascade():
        db = SessionLocal()
        try:
            res["A"] = ft.apply(db, appr, data=SONGUL)
        except Exception as e:                           # noqa: BLE001
            res["A"] = type(e).__name__
        finally:
            db.close()

    def ui():
        db = SessionLocal()
        try:
            real(db, l_cct, item_id=lic, supplier_id=None, supplier_name="UI", package_qty=5, price_unit="kg",
                 domain="cosmetics", actor="ui")
            db.commit()
            res["B"] = "ok"
        except PP.ApproveError as e:
            db.rollback()
            res["B"] = e.code
        except Exception as e:                           # noqa: BLE001
            db.rollback()
            res["B"] = type(e).__name__
        finally:
            db.close()
    ta = threading.Thread(target=cascade, daemon=True)
    ta.start()
    assert first_done.wait(10)
    tb = threading.Thread(target=ui, daemon=True)
    tb.start()
    try:
        assert _lock_waiter()                            # arayüz LIC kartında bekliyor
    finally:
        go.set()
        ta.join(20)
        tb.join(20)
    assert not ta.is_alive() and not tb.is_alive()
    assert isinstance(res["A"], dict) and res["A"]["applied"] is True and not res["A"]["errors"], res
    assert res["B"] == "not_pending", res
    db_session.expire_all()
    lc, ll = (db_session.query(ParasutPurchaseLine).get(i) for i in (l_cct, l_lic))
    assert (lc.chosen_item_id, ll.chosen_item_id) == (cct, lic)
    ps = [p for p in _prices(db_session) if p.source == "parasut"]
    assert sorted(p.item_id for p in ps) == sorted([cct, lic])


def test_commit_lock_busy_writes_nothing(db_session: Session, monkeypatch):
    """Kalan nadir kilit çakışması (kilitlenme kurbanı betik): hata `busy`, rollback,
    hiçbir fiyat / onay / audit yazılmaz; diğer veritabanı hataları yükselir."""
    from sqlalchemy.exc import OperationalError

    class Orig(Exception):
        def __init__(self, code):
            super().__init__(code)
            self.pgcode = code
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NILKIM, B_NAT)
    l_cct = lines[5001].id
    appr = {"parasut": {f"P-{l_cct}": {"ok": True, "row": 2, "item_id": its["cct"].id, "package": 25, "unit": "kg",
                                       "supplier_id": None}},
            "songul": {}, "form_filled": [], "errors": []}
    before = len(_prices(db_session))
    real = PP.approve_line

    def deadlock_after_write(*a, **kw):
        real(*a, **kw)
        raise OperationalError("SELECT 1", {}, Orig("40P01"))
    monkeypatch.setattr(PP, "approve_line", deadlock_after_write)
    res = ft.apply(db_session, appr, data=SONGUL)
    assert res["applied"] is False and res["errors"][0].startswith("busy:")
    assert len(_prices(db_session)) == before and not _audits(db_session)
    db_session.expire_all()
    assert db_session.query(ParasutPurchaseLine).get(l_cct).status == "pending"
    monkeypatch.setattr(PP, "approve_line",
                        lambda *a, **kw: (_ for _ in ()).throw(OperationalError("SELECT 1", {}, Orig("08006"))))
    with pytest.raises(OperationalError):
        ft.apply(db_session, appr, data=SONGUL)
    assert len(_prices(db_session)) == before


def test_manual_rows_never_overwritten(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NAT)
    lic_key = f"P-{lines[5005].id}"
    plan = ft.build_plan(db_session, data=SONGUL)
    path = ft.build_xlsx(plan, tmp_path / "onay.xlsx")
    _fill(path, ft.SHEETS[1], {lic_key: {ft.APPROVE_COL: "E", ft.PACK_COL: 5}})
    _fill(path, ft.SHEETS[2], {"C-s12-1": {ft.APPROVE_COL: "E"}})
    appr = ft.read_workbook(path)
    # Kitap ile yazım arasında lab elle fiyat girdi (aynı firma) → ikisi de durur, elle satır aynen kalır
    man1 = SupplierPrice(item_id=its["lic"].id, supplier_id=sups["NATURALYA"].id, supplier_name="NATURALYA",
                         unit_price=99.0, currency="TRY", price_unit="kg", source="manual", domain="cosmetics",
                         created_by="Songül")
    man2 = SupplierPrice(item_id=its["shea"].id, supplier_id=sups["UMAYCHEM"].id, supplier_name="UMAYCHEM",
                         unit_price=11.0, currency="EUR", price_unit="kg", source="manual", domain="cosmetics",
                         created_by="Songül")
    db_session.add_all([man1, man2])
    db_session.commit()
    res = ft.apply(db_session, appr, data=SONGUL)
    assert res["applied"] is False and any("elle_var" in e for e in res["errors"])
    ps = _prices(db_session)
    assert len(ps) == 4 and {p.unit_price for p in ps if p.source == "manual"} == {5.0, 7.0, 99.0, 11.0}
    assert db_session.query(ParasutPurchaseLine).get(lines[5005].id).status == "pending"
    # Yalnız Paraşüt E (Songül H): approve_line manual_exists → yine hiçbir şey yazılmaz
    appr["songul"]["C-s12-1"]["ok"] = False
    res = ft.apply(db_session, appr, data=SONGUL)
    assert res["applied"] is False and res["errors"] and "manual_exists" in res["errors"][0]
    assert len(_prices(db_session)) == 4 and not _audits(db_session)
    assert db_session.query(ParasutPurchaseLine).get(lines[5005].id).status == "pending"


# ─── --parasut-json ─────────────────────────────────────────────────────────

def test_parasut_json_jsonapi_dump_rows_and_commit_upserts(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    f = tmp_path / "parasut_api.json"
    f.write_text(json.dumps(_payload([B_NAT, B_NILKIM, B_ARAS])), encoding="utf-8")
    bills, kind = ft.load_parasut_json(f)
    assert kind == "jsonapi" and len(bills) == 3
    plan = ft.build_plan(db_session, data=SONGUL, parasut_bills=bills, parasut_kind=kind)
    pr = _rows(plan, "parasut")
    assert set(pr) == {"PJ-5005", "PJ-5006", "PJ-5001"} and pr["PJ-5005"]["input"] and pr["PJ-5005"]["line_id"] is None
    assert plan["summary"]["parasut_source"] == "db+jsonapi" and plan["summary"]["parasut_json_lines"] == 4
    assert plan["summary"]["parasut_pending_total"] == 0
    # Veritabanında olan döküm satırı atlanır (DB esastır) — çift satır yok
    db_lines = _bills_to_db(db_session, B_NILKIM)
    plan = ft.build_plan(db_session, data=SONGUL, parasut_bills=bills, parasut_kind=kind)
    pr = _rows(plan, "parasut")
    assert set(pr) == {"PJ-5005", "PJ-5006", f"P-{db_lines[5001].id}"} and plan["summary"]["parasut_json_in_db"] == 1
    # --commit JSON:API dökümüyle: fatura aynı işlemde tablolara yazılır, satır onaylanır
    path = ft.build_xlsx(plan, tmp_path / "onay.xlsx")
    _fill(path, ft.SHEETS[1], {"PJ-5005": {ft.APPROVE_COL: "E", ft.PACK_COL: 5}})
    appr = ft.read_workbook(path)
    res = ft.apply(db_session, appr, data=SONGUL, parasut_bills=bills)
    assert res["applied"] is True and [w["key"] for w in res["parasut_written"]] == ["PJ-5005"]
    ln = db_session.query(ParasutPurchaseLine).filter(ParasutPurchaseLine.parasut_detail_id == 5005).one()
    assert ln.status == "approved" and ln.bill.invoice_no == "M012026000002134"
    assert db_session.query(ParasutPurchaseLine).filter(ParasutPurchaseLine.parasut_detail_id == 5006).one().status == "pending"
    lic_sp = next(p for p in _prices(db_session) if p.item_id == its["lic"].id)
    assert lic_sp.source == "parasut" and lic_sp.unit_price == pytest.approx(620.798 / 5)
    # Döküm olmadan aynı E satırı: artık veritabanında → "zaten yazılmış"
    res2 = ft.apply(db_session, appr, data=SONGUL)
    assert res2["applied"] is True and res2["parasut_done"] == ["PJ-5005"] and res2["parasut_written"] == []


def test_parasut_json_flat_dump_whitelist_discount_and_not_writable(db_session: Session, tmp_path):
    its, sups = _seed(db_session)
    dump = {"bills": [
        {"date": "2024-05-27", "no": "ORA2024000000058", "currency": "TRL", "exchange_rate": "1.0",
         "net_total": "959.84", "description": "Hammadde", "item_type": "purchase_bill",
         "supplier": "ORLIFE GLOBAL İÇ VE DIŞ TİCARET A.Ş.",
         "supplier_contact": {"email": "gizli@example.com", "phone": "+90 555", "tax_number": "123", "city": "X",
                              "address": "Gizli"},
         "lines": [{"desc": "AKGÜNLÜK UÇUCU YAĞI 10 ml", "qty": "4", "unit_price": "249.96", "vat_rate": "20.0",
                    "net_total": "959.84", "product": "AKGÜNLÜK UÇUCU YAĞI (Boswellia serrata) 10 ml.",
                    "unit": "Adet"}]},
        {"date": "2026-08-14", "no": "M012026000002134", "currency": "TRL", "exchange_rate": "1.0",
         "net_total": "3724.79", "description": None, "item_type": "purchase_bill",
         "supplier": "NATURALYA KİMYA GIDA LTD. ŞTİ.", "supplier_contact": {"email": "gizli@example.com"},
         "lines": [{"desc": "LICORICE EXTRACT LIQ", "qty": "5", "unit_price": "620.798", "vat_rate": "20.0",
                    "net_total": "3724.79", "product": "LICORICE EXTRACT LIQ", "unit": "Adet"}]},
        {"date": "2024-01-01", "no": "X", "item_type": "sales_invoice", "supplier": "Y", "lines": []},
    ], "products": [{"name": "gizli", "buying_price": "1"}], "product_total": 65}
    f = tmp_path / "parasut.json"
    f.write_text(json.dumps(dump, ensure_ascii=False), encoding="utf-8")
    bills, kind = ft.load_parasut_json(f)
    assert kind == "dump" and len(bills) == 2                         # sales_invoice atıldı
    for b in bills:
        assert set(b) == PP.BILL_KEYS and all(set(ln) == PP.LINE_KEYS for ln in b["lines"])
        assert "supplier_contact" not in json.dumps(b, default=str) and "gizli" not in json.dumps(b, default=str)
    orl = bills[0]
    assert orl["currency"] == "TRY" and orl["parasut_id"] is None and orl["issue_date"] == date(2024, 5, 27)
    assert orl["lines"][0]["discount_rate"] == pytest.approx(0.2, abs=1e-3)     # net toplamdan türetilen %20
    assert bills[1]["lines"][0]["discount_rate"] is None
    plan = ft.build_plan(db_session, data=SONGUL, parasut_bills=bills, parasut_kind=kind)
    pr = _rows(plan, "parasut")
    assert set(pr) == {"PJx-2"} and pr["PJx-2"]["input"] is False and "YAZILAMAZ" in pr["PJx-2"]["note"]
    path = ft.build_xlsx(plan, tmp_path / "onay.xlsx")
    ws = load_workbook(path)[ft.SHEETS[1]]
    head = [c.value for c in ws[1]]
    row = next(ws.iter_rows(min_row=2, values_only=True))
    assert row[head.index(ft.APPROVE_COL)] == ft.NOT_APPLICABLE and row[head.index(ft.PACK_COL)] == ft.NOT_APPLICABLE
    # Kimliksiz satıra zorla E → yazılamaz, hiçbir şey yazılmaz
    appr = {"parasut": {"PJx-2": {"ok": True, "row": 2, "item_id": its["lic"].id, "package": 5, "unit": "kg",
                                  "supplier_id": None}}, "songul": {}, "form_filled": [], "errors": []}
    res = ft.apply(db_session, appr, data=SONGUL, parasut_bills=bills)
    assert res["applied"] is False and "kimliksiz" in res["errors"][0] and len(_prices(db_session)) == 2
    with pytest.raises(ValueError):
        g = tmp_path / "bad.json"
        g.write_text('{"x": 1}', encoding="utf-8")
        ft.load_parasut_json(g)


# ─── CLI ────────────────────────────────────────────────────────────────────

def test_cli_dry_run_workbook_and_commit_gates(db_session: Session, tmp_path, capsys):
    its, sups = _seed(db_session)
    lines = _bills_to_db(db_session, B_NAT)
    lic_key = f"P-{lines[5005].id}"
    f = tmp_path / "songul.json"
    f.write_text(json.dumps(SONGUL, ensure_ascii=False), encoding="utf-8")
    x = tmp_path / "ft.xlsx"
    assert ft.main(["--data", str(f), "--xlsx", str(x)]) == 0
    out = capsys.readouterr().out
    assert "kuru çalıştırma" in out and "Çalışma kitabı yazıldı" in out and x.exists()
    assert len(_prices(db_session)) == 2
    assert ft.main(["--data", str(f), "--commit"]) == 2                              # onaylı Excel şart
    assert ft.main(["--data", str(f), "--parasut-json", str(tmp_path / "yok.json")]) == 2
    assert ft.main(["--data", str(f), "--onayli", str(x)]) == 0                      # boş onay, kuru denetim
    _fill(x, ft.SHEETS[1], {lic_key: {ft.APPROVE_COL: "E"}})                         # ambalajsız E → hata
    assert ft.main(["--data", str(f), "--onayli", str(x), "--commit"]) == 2
    assert len(_prices(db_session)) == 2 and not _audits(db_session)
    _fill(x, ft.SHEETS[1], {lic_key: {ft.PACK_COL: 5}})
    _fill(x, ft.SHEETS[2], {"C-s12-1": {ft.APPROVE_COL: "E"}})
    assert ft.main(["--data", str(f), "--onayli", str(x)]) == 0                      # kuru: yazmaz
    assert len(_prices(db_session)) == 2
    assert ft.main(["--data", str(f), "--onayli", str(x), "--commit"]) == 0
    out = capsys.readouterr().out
    assert "✓ yazıldı: Paraşüt 1 fiyat · Songül 1 fiyat" in out
    assert len(_prices(db_session)) == 4 and len(_audits(db_session)) == 1


def test_real_data_file_dry_run_smoke(db_session: Session):
    its, sups = _seed(db_session)
    db_session.add(Supplier(name="SURYA KİMYA", domain="cosmetics"))      # A2-37-S1 MEYAN KÖKÜ / SURYA → yazılabilir
    db_session.commit()
    _bills_to_db(db_session, B_NAT)
    d = lab.load_data()
    assert isinstance(d["kaynak"], list) and len(d["kaynak"]) == 2 and "SongulHanimFiyat2.pdf" in d["kaynak"][1]
    assert len(d["kayitlar"]) == 178
    plan = ft.build_plan(db_session, data=d)
    s = plan["summary"]
    assert s["songul_records"] == 178 and s["targets"] == 8 and s["parasut_rows"] >= 2
    # Gerçek dosyadaki SHEA / BADEM kayıtları hedef kartlara düşer, eşleşmeyenler listelenmez
    sr = _rows(plan, "songul")
    assert sr and all(r["card"]["id"] in plan["targets"] for r in sr.values())
    assert set(r["status"] for r in sr.values()) <= set(lab.STATUS_TEXT)
    assert sr["A2-37-S1"]["writable"] and sr["A2-37-S1"]["card"]["id"] == its["lic"].id
    assert sr["A2-29-S1"]["status"] == "belirsiz" and not sr["A2-29-S1"]["writable"]      # iki GİNSENG kartı
    assert its["lic"].id not in {t["item_id"] for t in plan["form"]} and s["covered"] >= 1
