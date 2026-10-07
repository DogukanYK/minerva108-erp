"""
Üretim iptali (core/production_cancel.py + /api/production/{id}/cancel*) testleri.

Kritik sözleşmeler:
  • Başlatma her Output/Input'u `production_consumptions`'a döker; toplamlar
    defterle birebir, şahit kaydı üretime bağlı.
  • İptal hiçbir Transaction'ı silmez/değiştirmez — yalnız telafi
    Adjustment'ları ("Üretim iptali" önekli); hammadde stok + lot miktarları ve
    bitmiş ürün stoğu üretim öncesine döner, `compute_stock_at(t0)` değişmez.
  • Bitmiş ürün sonradan hareket gördüyse (teslimat/QC reddi/şahit çıkışı/boy
    düzeltmesi/elle düzeltme) 409 — HİÇBİR ŞEY değişmez.
  • Dökümü olmayan eski üretim defterden kurulur — AYNI lot no'lu başka
    ürünün (ve aynı reçete adlı SR0011'in) satırları karışmaz.
  • İptaller raporlardan dışlanır; lot no yalnız `release_lot` ile serbest.
"""
import json
from datetime import date, datetime, timedelta

import bcrypt
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from core import lots, stock_lots
from core.production_cancel import NOTE_MARK
from core.snapshots import compute_stock_at
from database import (AdminAuditLog, AppSetting, Inventory, Item, ProductionConsumption,
                      ProductionHistory, Recipe, RecipeIngredient, RetentionSample,
                      RetentionSampleMovement, Supplier, Transaction, User)

_HDR = {"Origin": "http://testserver"}


# ─── Kurulum ────────────────────────────────────────────────────────────────

def _materials(db: Session, *, domain="cosmetics"):
    """Ortak malzemeler: lotlu hammadde (iki lot, iki tedarikçi), lotsuz
    hammadde, şişe, TR/EN etiket çifti."""
    s1 = Supplier(name="Tedarikçi Bir", domain=domain)
    s2 = Supplier(name="Tedarikçi İki", domain=domain)
    db.add_all([s1, s2]); db.flush()
    raw = Item(name="İptal Yağı", sku="c-raw", category="Hammadde", unit="g",
               current_stock=1000, domain=domain)
    raw2 = Item(name="İptal Yağı (2. firma)", sku="c-raw2", category="Hammadde", unit="g",
                current_stock=0, domain=domain)
    water = Item(name="İptal Suyu", sku="c-water", category="Hammadde", unit="ml",
                 current_stock=1000, domain=domain)
    bottle = Item(name="İptal Şişe", sku="c-bottle", category="Ambalaj", pkg_type="şişe",
                  unit="adet", current_stock=1000, domain=domain)
    ltr = Item(name="İptal Etiket TR", sku="c-ltr", category="Ambalaj", pkg_type="etiket",
               language="TR", label_group="CNL", unit="adet", current_stock=1000, domain=domain)
    len_ = Item(name="İptal Etiket EN", sku="c-len", category="Ambalaj", pkg_type="etiket",
                language="EN", label_group="CNL", unit="adet", current_stock=1000, domain=domain)
    db.add_all([raw, raw2, water, bottle, ltr, len_]); db.flush()
    # FIFO: L-A (3 g, eski) biter, kalan L-B'den — tek kalemde iki Output.
    db.add(Inventory(item_id=raw.id, supplier_id=s1.id, lot_number="L-A", quantity=3.0,
                     status="APPROVED", domain=domain,
                     created_at=datetime.utcnow() - timedelta(days=20)))
    db.add(Inventory(item_id=raw.id, supplier_id=s2.id, lot_number="L-B", quantity=997.0,
                     status="APPROVED", domain=domain,
                     created_at=datetime.utcnow() - timedelta(days=10)))
    db.commit()
    return {"raw": raw.id, "raw2": raw2.id, "water": water.id, "bottle": bottle.id,
            "ltr": ltr.id, "len": len_.id}


def _product(db: Session, m: dict, *, suffix="A", recipe_name=None, domain="cosmetics"):
    """Varyasyonlu bitmiş ürün (ana + bu boy + kardeş boy) + reçete."""
    parent = Item(name=f"Serenida İptal Şampuan {suffix}", sku=f"c-p-{suffix}",
                  category="Bitmiş Ürün", unit="adet", current_stock=0, domain=domain)
    db.add(parent); db.flush()
    tgt = Item(name=f"Serenida İptal Şampuan {suffix} 200 ml", sku=f"c-t-{suffix}",
               category="Bitmiş Ürün", unit="adet", current_stock=0, parent_id=parent.id,
               variation_name="200 ml", domain=domain)
    sib = Item(name=f"Serenida İptal Şampuan {suffix} 500 ml", sku=f"c-s-{suffix}",
               category="Bitmiş Ürün", unit="adet", current_stock=0, parent_id=parent.id,
               variation_name="500 ml", domain=domain)
    db.add_all([tgt, sib]); db.flush()
    rec = Recipe(name=recipe_name or tgt.name, output_quantity=1, output_unit="adet",
                 target_item_id=tgt.id, waste_percentage=10, domain=domain)
    db.add(rec); db.flush()
    for key, q in (("raw", 1.0), ("water", 2.0), ("bottle", 1), ("ltr", 1)):
        db.add(RecipeIngredient(recipe_id=rec.id, item_id=m[key], quantity=q, phase="A"))
    db.commit()
    return rec.id, tgt.id, sib.id


def _produce(client: TestClient, recipe_id: int, qty=20, *, witness=2, lang="EN", lot=None):
    body = {"recipe_id": recipe_id, "produced_quantity": qty, "label_language": lang,
            "witness_quantity": witness}
    if lot:
        body["lot_number"] = lot
    r = client.post("/api/production", json=body, headers=_HDR)
    assert r.status_code == 201, r.text
    return r.json()


def _state(db: Session) -> dict:
    """Karşılaştırma için tüm kart stokları + lot miktarları/durumları + şahit."""
    db.expire_all()
    return {
        "items": {i.id: round(i.current_stock or 0.0, 6) for i in db.query(Item).all()},
        "inv": {i.id: (round(i.quantity or 0.0, 6), i.status, i.item_id)
                for i in db.query(Inventory).all()},
        "rs": {r.id: (r.quantity, r.is_active, r.status, r.item_id)
               for r in db.query(RetentionSample).all()},
        "tx": db.query(Transaction).count(),
        "ph": {p.id: p.cancelled_at for p in db.query(ProductionHistory).all()},
    }


def _preview(client: TestClient, pid: int):
    r = client.get(f"/api/production/{pid}/cancel-preview")
    assert r.status_code == 200, r.text
    return r.json()


def _cancel(client: TestClient, pid: int, fp: str, *, reason="Yanlış ürün seçildi",
            release_lot=False):
    return client.post(f"/api/production/{pid}/cancel",
                       json={"reason": reason, "fingerprint": fp, "release_lot": release_lot},
                       headers=_HDR)


def _login(client: TestClient, username: str):
    client.post("/api/logout", headers=_HDR)
    r = client.post("/api/login", json={"username": username, "password": "minerva123"},
                    headers=_HDR)
    assert r.status_code == 200, r.text
    return client


# ─── 1) Başlatma döküm yazar ────────────────────────────────────────────────

def test_start_writes_consumption_snapshot(authed_client, db_session):
    m = _materials(db_session)
    rid, tgt, _ = _product(db_session, m)
    out = _produce(authed_client, rid)
    db_session.expire_all()
    ph = db_session.query(ProductionHistory).one()
    assert out["production_id"] == ph.id and ph.lot_number == "SR001"
    rows = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.production_id == ph.id).all())
    outputs = db_session.query(Transaction).filter(Transaction.transaction_type == "Output").all()
    cons = [r for r in rows if r.kind != "output"]
    # Output başına tam bir satır, aynı tx, aynı kart, aynı miktar
    assert sorted(r.transaction_id for r in cons) == sorted(t.id for t in outputs)
    by_tx = {t.id: t for t in outputs}
    for r in cons:
        t = by_tx[r.transaction_id]
        assert r.item_id == t.item_id and r.quantity == pytest.approx(t.quantity)
        assert r.lot_number == t.lot_number and r.source == "live" and r.domain == "cosmetics"
    kinds = {r.item_id: r.kind for r in cons}
    assert kinds[m["raw"]] == "raw" and kinds[m["water"]] == "raw"
    assert kinds[m["bottle"]] == "packaging" and kinds[m["len"]] == "label"
    # Etiket EN kardeşine çözüldü — reçetedeki kart TR
    lbl = next(r for r in cons if r.item_id == m["len"])
    assert lbl.recipe_item_id == m["ltr"] and lbl.factor == 1.0
    # Lotlu hammadde: iki lot, tedarikçi + inventory bağı, fire çarpanı
    raw_rows = sorted((r for r in cons if r.item_id == m["raw"]), key=lambda r: r.lot_number)
    assert [r.lot_number for r in raw_rows] == ["L-A", "L-B"]
    assert [r.supplier_name for r in raw_rows] == ["Tedarikçi Bir", "Tedarikçi İki"]
    assert all(r.inventory_id and r.factor == pytest.approx(1.1) for r in raw_rows)
    assert sum(r.quantity for r in raw_rows) == pytest.approx(22.0)
    water = [r for r in cons if r.item_id == m["water"]]
    assert len(water) == 1 and water[0].inventory_id is None and water[0].lot_number is None
    # Bitmiş ürün: showroom + -S, ikisi de tek Input
    inp = (db_session.query(Transaction)
           .filter(Transaction.transaction_type == "Input", Transaction.item_id == tgt).one())
    outs = sorted((r for r in rows if r.kind == "output"), key=lambda r: r.lot_number)
    assert [(r.lot_number, r.quantity) for r in outs] == [("SR001", 18), ("SR001-S", 2)]
    assert all(r.transaction_id == inp.id and r.item_id == tgt and r.inventory_id for r in outs)
    rs = db_session.query(RetentionSample).one()
    assert rs.production_history_id == ph.id


def test_start_production_scoped_to_active_domain(authed_client, db_session):
    m = _materials(db_session, domain="supplement")
    rid, _, _ = _product(db_session, m, domain="supplement")
    r = authed_client.post("/api/production", json={"recipe_id": rid, "produced_quantity": 1},
                           headers=_HDR)
    assert r.status_code == 404
    assert db_session.query(Transaction).count() == 0
    authed_client.cookies.set("active_domain", "supplement")
    assert authed_client.post("/api/production", json={"recipe_id": rid, "produced_quantity": 1},
                              headers=_HDR).status_code == 201


# ─── 2) Mutlu yol + 3) defter tutarlılığı + 4) idempotent ──────────────────

def test_cancel_happy_path_restores_everything(authed_client, db_session):
    m = _materials(db_session)
    rid, tgt, _ = _product(db_session, m)
    before = _state(db_session)
    t0 = datetime.utcnow()
    touched = [m["raw"], m["water"], m["bottle"], m["len"], tgt]
    s0, _ = compute_stock_at(db_session, t0, Item.id.in_(touched))

    _produce(authed_client, rid)
    db_session.expire_all()
    ph = db_session.query(ProductionHistory).one()
    orig_tx = {t.id: (t.item_id, t.transaction_type, t.quantity, t.notes)
               for t in db_session.query(Transaction).all()}

    pv = _preview(authed_client, ph.id)
    assert pv["can_cancel"] is True and pv["blockers"] == [] and pv["source"] == "snapshot"
    assert pv["production"]["lot"] == "SR001" and pv["production"]["qty"] == 20
    assert {r["tx_id"] for r in pv["restore"]} == {t for t, v in orig_tx.items()
                                                    if v[1] == "Output"}
    assert sorted((r["lot"], r["qty"]) for r in pv["remove"]) == [("SR001", 18), ("SR001-S", 2)]
    assert pv["retention"]["qty"] == 2

    r = _cancel(authed_client, ph.id, pv["fingerprint"])
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["already_cancelled"] is False and body["source"] == "snapshot"
    assert body["cancelled"]["reason"] == "Yanlış ürün seçildi"

    after = _state(db_session)
    # Hammadde / ambalaj / etiket / hedef stok üretim öncesine döndü
    for iid in touched:
        assert after["items"][iid] == pytest.approx(before["items"][iid]), iid
    # Kaynak lotlar üretim öncesine döndü
    for inv_id, (q, st, iid) in before["inv"].items():
        assert after["inv"][inv_id][0] == pytest.approx(q), inv_id
    # Bitmiş satırlar kapandı
    fin = db_session.query(Inventory).filter(Inventory.item_id == tgt).all()
    assert {(i.lot_number, i.quantity, i.status, i.qc_required) for i in fin} == {
        ("SR001", 0.0, "CANCELLED", False), ("SR001-S", 0.0, "CANCELLED", False)}
    assert all(NOTE_MARK in (i.qc_notes or "") for i in fin)
    # Şahit dolaptan kalktı + hareket
    rs = db_session.query(RetentionSample).one()
    assert rs.is_active is False and rs.quantity == 0
    mv = (db_session.query(RetentionSampleMovement)
          .filter(RetentionSampleMovement.sample_id == rs.id,
                  RetentionSampleMovement.movement_type == "duzeltme").one())
    assert mv.quantity == 2 and NOTE_MARK in mv.note
    # Orijinal defter satırlarının HİÇBİRİ silinmedi/değişmedi
    now_tx = {t.id: (t.item_id, t.transaction_type, t.quantity, t.notes)
              for t in db_session.query(Transaction).all()}
    for tid, v in orig_tx.items():
        assert now_tx[tid] == v
    new = [t for tid, t in ((tid, db_session.get(Transaction, tid)) for tid in now_tx)
           if tid not in orig_tx]
    assert new and all(t.transaction_type == "Adjustment" and t.notes.startswith(NOTE_MARK)
                       for t in new)
    # Hedefte TEK −Adjustment (Input aynası), lot no'lu
    tgt_adj = [t for t in new if t.item_id == tgt]
    assert len(tgt_adj) == 1 and tgt_adj[0].quantity == -20 and tgt_adj[0].lot_number == "SR001"
    # Kaynak lot no'su telafiye yazıldı
    assert sorted(t.lot_number for t in new if t.item_id == m["raw"]) == ["L-A", "L-B"]
    # Döküm satırları telafi kaydına bağlandı
    rows = db_session.query(ProductionConsumption).all()
    assert all(r.cancel_transaction_id for r in rows)
    # PH damgası + audit
    ph = db_session.get(ProductionHistory, ph.id)
    assert ph.cancelled_at and ph.cancelled_by and ph.lot_released is False
    audit = db_session.query(AdminAuditLog).filter(AdminAuditLog.action == "production.cancel").one()
    det = json.loads(audit.details)
    assert det["lot"] == "SR001" and det["source"] == "snapshot"
    assert set(det["tx_ids"]) == {r.transaction_id for r in rows}
    assert set(det["cancel_tx_ids"]) == {t.id for t in new}

    # 3) Defter tutarlılığı — t0 anındaki stok değişmedi, bugünkü stok = t0
    s_after, _ = compute_stock_at(db_session, t0, Item.id.in_(touched))
    for iid in touched:
        assert s_after[iid] == pytest.approx(s0[iid]), iid
        assert db_session.get(Item, iid).current_stock == pytest.approx(s0[iid]), iid

    # 4) İdempotent — ikinci POST hiçbir şey yazmaz
    n_tx = db_session.query(Transaction).count()
    r2 = _cancel(authed_client, ph.id, pv["fingerprint"])
    assert r2.status_code == 200 and r2.json()["already_cancelled"] is True
    db_session.expire_all()
    assert db_session.query(Transaction).count() == n_tx
    assert db_session.query(AdminAuditLog).filter(
        AdminAuditLog.action == "production.cancel").count() == 1


def test_reason_required(authed_client, db_session):
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    fp = _preview(authed_client, pid)["fingerprint"]
    assert _cancel(authed_client, pid, fp, reason="    a    ").status_code == 422
    assert _cancel(authed_client, pid, fp, reason="abc").status_code == 422
    db_session.expire_all()
    assert db_session.get(ProductionHistory, pid).cancelled_at is None


# ─── 5) Engeller — 409, hiçbir şey değişmez ────────────────────────────────

def _block_delivery(client, db, ids):
    tgt = db.get(Item, ids["tgt"])
    stock_lots.consume(db, tgt, 1, note="Teslimat — test", actor="test")
    db.commit()


def _block_qc_reject(client, db, ids):
    inv = (db.query(Inventory).filter(Inventory.item_id == ids["tgt"],
                                      Inventory.lot_number == "SR001").one())
    r = client.post(f"/api/qc/process/{inv.id}", json={"status": "REJECTED", "notes": "bozuk"},
                    headers=_HDR)
    assert r.status_code == 200, r.text


def _block_witness_checkout(client, db, ids):
    rs = db.query(RetentionSample).one()
    r = client.post(f"/api/retention/samples/{rs.id}/checkout",
                    json={"quantity": 1, "reason": "test"}, headers=_HDR)
    assert r.status_code == 200, r.text


def _block_witness_variation(client, db, ids):
    # Uç üretim kaydının boyunu değiştirmeye izin vermiyor (inventory_id
    # bağlı); toplu düzeltme script'leri (scripts/assign_retention_variation)
    # aynı yardımcıyı doğrudan çağırıyor — o yol.
    from routers.retention import _apply_variation
    rs = db.query(RetentionSample).one()
    assert _apply_variation(db, rs, db.get(Item, ids["sib"]), "script")
    db.commit()


def _block_manual_adjust(client, db, ids):
    r = client.post("/api/inventory/adjust",
                    json={"item_id": ids["tgt"], "new_quantity": 15, "reason": "sayım farkı"},
                    headers=_HDR)
    assert r.status_code == 201, r.text


@pytest.mark.parametrize("mutate,expect", [
    (_block_delivery, "sonradan stok hareketi"),
    (_block_qc_reject, "QC'de reddedilmiş"),
    (_block_witness_checkout, "Şahit numune"),
    (_block_witness_variation, "boyu düzeltilmiş"),
    (_block_manual_adjust, "miktarı değişmiş"),
])
def test_blockers_return_409_and_change_nothing(authed_client, db_session, mutate, expect):
    m = _materials(db_session)
    rid, tgt, sib = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    mutate(authed_client, db_session, {"tgt": tgt, "sib": sib})

    pv = _preview(authed_client, pid)
    assert pv["can_cancel"] is False
    assert any(expect in b for b in pv["blockers"]), pv["blockers"]

    before = _state(db_session)
    r = _cancel(authed_client, pid, pv["fingerprint"])
    assert r.status_code == 409, r.text
    assert r.json()["code"] == "blocked" and r.json()["blockers"]
    assert _state(db_session) == before
    assert db_session.query(AdminAuditLog).filter(
        AdminAuditLog.action == "production.cancel").count() == 0


# ─── 6) Eski üretim — defterden yeniden kurma ──────────────────────────────

def test_legacy_production_reconstructed_from_ledger_without_mixing(authed_client, db_session):
    """Dökümü olmayan üretim: aynı kişi, aynı dakikalar, AYNI lot no (SR001)
    başka ürün + aynı reçete ADLI başka ürünün SR0011'i — yalnız kendi
    Output'ları eşlenir."""
    m = _materials(db_session)
    rid_a, tgt_a, _ = _product(db_session, m, suffix="A")
    rid_b, tgt_b, _ = _product(db_session, m, suffix="B")
    a_name = db_session.get(Recipe, rid_a).name
    rid_c, tgt_c, _ = _product(db_session, m, suffix="C", recipe_name=a_name)

    pid_a = _produce(authed_client, rid_a)["production_id"]
    pid_b = _produce(authed_client, rid_b, qty=5, witness=1)["production_id"]
    pid_c = _produce(authed_client, rid_c, qty=3, witness=0, lot="SR0011")["production_id"]
    db_session.expire_all()
    assert db_session.get(ProductionHistory, pid_b).lot_number == "SR001"   # aynı lot no, başka ürün
    a_tx = {r.transaction_id for r in db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.production_id == pid_a,
                    ProductionConsumption.kind != "output")}
    a_input = {r.transaction_id for r in db_session.query(ProductionConsumption)
               .filter(ProductionConsumption.production_id == pid_a,
                       ProductionConsumption.kind == "output")}
    # Özellik öncesi gibi: hiçbir üretimin dökümü yok, şahit bağı yok
    db_session.query(ProductionConsumption).delete()
    for rs in db_session.query(RetentionSample).all():
        rs.production_history_id = None
    db_session.commit()

    state_b_c = lambda: {i.id: (i.quantity, i.status) for i in db_session.query(Inventory)  # noqa: E731
                         .filter(Inventory.item_id.in_([tgt_b, tgt_c])).all()}
    before_bc = state_b_c()
    stock_bc = {i: db_session.get(Item, i).current_stock for i in (tgt_b, tgt_c)}

    pv = _preview(authed_client, pid_a)
    assert pv["source"] == "ledger", pv
    assert pv["blockers"] == [] and pv["can_cancel"] is True
    assert {r["tx_id"] for r in pv["restore"]} == a_tx
    assert any("defterden" in w for w in pv["warnings"])
    assert sorted((r["lot"], r["qty"]) for r in pv["remove"]) == [("SR001", 18), ("SR001-S", 2)]
    # Kaynak lotlar ve tedarikçi defterden bulundu
    raw_rows = sorted((r for r in pv["restore"] if r["item_id"] == m["raw"]),
                      key=lambda r: r["lot"])
    assert [r["lot"] for r in raw_rows] == ["L-A", "L-B"] and all(r["inventory_id"] for r in raw_rows)
    assert [r["supplier"] for r in raw_rows] == ["Tedarikçi Bir", "Tedarikçi İki"]

    r = _cancel(authed_client, pid_a, pv["fingerprint"])
    assert r.status_code == 200, r.text
    assert r.json()["source"] == "ledger" and set(r.json()["tx_ids"]) == a_tx | a_input
    db_session.expire_all()
    # Döküm kalıcılaştı (source='ledger')
    rows = (db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.production_id == pid_a).all())
    assert rows and all(r.source == "ledger" and r.cancel_transaction_id for r in rows)
    assert {r.transaction_id for r in rows if r.kind != "output"} == a_tx
    # B ve C'ye dokunulmadı
    assert state_b_c() == before_bc
    assert {i: db_session.get(Item, i).current_stock for i in (tgt_b, tgt_c)} == stock_bc
    assert db_session.get(ProductionHistory, pid_b).cancelled_at is None
    assert db_session.get(ProductionHistory, pid_c).cancelled_at is None
    # B'nin şahidi hâlâ dolapta
    rs_b = (db_session.query(RetentionSample).filter(RetentionSample.item_id == tgt_b).one())
    assert rs_b.is_active and rs_b.quantity == 1
    # A tamamen geri döndü
    assert db_session.get(Item, tgt_a).current_stock == 0
    # B de (artık A'nın txleri sahiplenildiği hâlde) defterden kurulabilir
    pv_b = _preview(authed_client, pid_b)
    assert pv_b["source"] == "ledger" and pv_b["can_cancel"] is True, pv_b
    assert not ({r["tx_id"] for r in pv_b["restore"]} & a_tx)


def test_ledger_without_outputs_is_blocked(authed_client, db_session):
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    db_session.query(ProductionConsumption).delete()
    ph = db_session.get(ProductionHistory, pid)
    ph.produced_by = "Başka Biri"           # imza tutmuyor → eşleşme yok
    db_session.commit()
    pv = _preview(authed_client, pid)
    assert pv["source"] == "ledger" and pv["can_cancel"] is False
    assert any("defterde bulunamadı" in b for b in pv["blockers"])
    r = _cancel(authed_client, pid, pv["fingerprint"])
    assert r.status_code == 409


# ─── 7) Bayat önizleme ──────────────────────────────────────────────────────

def test_stale_fingerprint_rejected(authed_client, db_session):
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    before = _state(db_session)
    r = _cancel(authed_client, pid, "deadbeef")
    assert r.status_code == 409 and r.json()["code"] == "preview_stale"
    assert _state(db_session) == before


def test_source_lot_moved_after_preview(authed_client, db_session):
    """Önizlemeden sonra kaynak lot başka karta taşındı → 409 preview_stale;
    yeni önizleme iadeyi lotun BUGÜNKÜ kartına yazar (uyarılı)."""
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    pv = _preview(authed_client, pid)

    lot_b = (db_session.query(Inventory).filter(Inventory.item_id == m["raw"],
                                                Inventory.lot_number == "L-B").one())
    stock_lots.move_lot(db_session, lot_b, db_session.get(Item, m["raw"]),
                        db_session.get(Item, m["raw2"]), None, actor="test", reason="yanlış kart")
    db_session.commit()

    r = _cancel(authed_client, pid, pv["fingerprint"])
    assert r.status_code == 409 and r.json()["code"] == "preview_stale"

    pv2 = _preview(authed_client, pid)
    moved = next(r for r in pv2["restore"] if r["lot"] == "L-B")
    assert moved["item_id"] == m["raw2"] and moved["consumed_item_id"] == m["raw"]
    assert "target_card_note" in moved and pv2["can_cancel"] is True
    assert any("taşınmış" in w for w in pv2["warnings"])
    lot_b_qty = db_session.get(Inventory, lot_b.id).quantity
    raw2_stock = db_session.get(Item, m["raw2"]).current_stock

    assert _cancel(authed_client, pid, pv2["fingerprint"]).status_code == 200
    db_session.expire_all()
    assert db_session.get(Inventory, lot_b.id).quantity == pytest.approx(lot_b_qty + moved["qty"])
    assert db_session.get(Item, m["raw2"]).current_stock == pytest.approx(raw2_stock + moved["qty"])


# ─── 8) RBAC + panel ────────────────────────────────────────────────────────

def test_rbac_and_domain(client, db_session):
    _login(client, "dogukan")
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(client, rid)["production_id"]

    _login(client, "meltem")                                       # LabTech
    assert client.get(f"/api/production/{pid}/cancel-preview").status_code == 403
    assert _cancel(client, pid, "x").status_code == 403

    _login(client, "songul")                                       # LabLead
    pv = client.get(f"/api/production/{pid}/cancel-preview")
    assert pv.status_code == 200
    client.cookies.set("active_domain", "supplement")
    assert client.get(f"/api/production/{pid}/cancel-preview").status_code == 404
    assert _cancel(client, pid, pv.json()["fingerprint"]).status_code == 404
    client.cookies.set("active_domain", "cosmetics")
    assert _cancel(client, pid, pv.json()["fingerprint"]).status_code == 200

    client.post("/api/logout", headers=_HDR)
    client.cookies.clear()
    assert client.get(f"/api/production/{pid}/cancel-preview").status_code == 401
    assert _cancel(client, pid, "x").status_code == 401


def test_manager_default_has_cancel():
    from core.permissions import _DEFAULT_PERMISSIONS, PERMISSION_CATEGORIES
    assert "cancel" in PERMISSION_CATEGORIES["production"]
    assert _DEFAULT_PERMISSIONS["Manager"]["production"]["cancel"] is True
    assert _DEFAULT_PERMISSIONS["LabLead"]["production"]["cancel"] is True
    for role in ("LabTech", "Staff", "Distributor"):
        assert _DEFAULT_PERMISSIONS[role]["production"]["cancel"] is False


def test_permission_backfill_for_override_users(db_session):
    from database import _backfill_perm_production_cancel
    pw = bcrypt.hashpw(b"minerva123", bcrypt.gensalt()).decode()

    def _u(name, role, perms):
        u = User(username=name, password_hash=pw, full_name=name, role=role, is_active=True,
                 permissions=json.dumps(perms))
        db_session.add(u)
        return u

    mgr = _u("isik2", "Manager", {"production": {"view": True, "create": True}})
    lt = _u("tech2", "LabTech", {"production": {"view": True, "create": True}})
    off = _u("mgr_off", "Manager", {"production": {"view": True, "create": False}})
    keep = _u("mgr_keep", "Manager", {"production": {"view": True, "create": True,
                                                      "cancel": False}})
    nocat = _u("lead_nocat", "LabLead", {"items": {"view": True}})
    db_session.query(AppSetting).filter(
        AppSetting.key == "backfill.perm.production_cancel.v1").delete()
    db_session.commit()

    _backfill_perm_production_cancel()
    db_session.expire_all()
    perms = lambda u: json.loads(db_session.get(User, u.id).permissions)  # noqa: E731
    assert perms(mgr)["production"]["cancel"] is True
    assert perms(lt)["production"]["cancel"] is False
    assert perms(off)["production"]["cancel"] is False         # üretimi kapalı → iptal de kapalı
    assert perms(keep)["production"]["cancel"] is False        # bilerek kapatılmış, dokunulmadı
    assert perms(nocat)["production"] == {"cancel": False}      # create yok → kapalı
    assert perms(nocat)["items"] == {"view": True}
    audits = (db_session.query(AdminAuditLog)
              .filter(AdminAuditLog.action == "permissions.backfill").all())
    assert {a.target_name for a in audits} == {"isik2", "tech2", "mgr_off", "lead_nocat"}

    # Sentinel — ikinci koşu no-op (yönetici sonradan değiştirse de ezilmez)
    u = db_session.get(User, mgr.id)
    p = json.loads(u.permissions); p["production"]["cancel"] = False
    u.permissions = json.dumps(p); db_session.commit()
    _backfill_perm_production_cancel()
    db_session.expire_all()
    assert perms(mgr)["production"]["cancel"] is False


# ─── 9) Raporlar ────────────────────────────────────────────────────────────

def test_reports_exclude_cancelled(authed_client, db_session):
    from core.monthly_report import gather_report_data
    from core.production_history_report import history_quantities
    m = _materials(db_session)
    rid, tgt, _ = _product(db_session, m)
    p1 = _produce(authed_client, rid, qty=20)["production_id"]
    p2 = _produce(authed_client, rid, qty=10, witness=0)["production_id"]
    fp = _preview(authed_client, p1)["fingerprint"]
    assert _cancel(authed_client, p1, fp).status_code == 200

    trends = authed_client.get("/api/reports/production-trends").json()
    assert sum(d["total_quantity"] for d in trends) == 10
    assert sum(d["count"] for d in trends) == 1

    dash = authed_client.get("/api/dashboard/stats").json()
    assert [r["produced_quantity"] for r in dash["recent_productions"]] == [10]

    today = datetime.utcnow().date()
    hq = history_quantities(db_session, "cosmetics", today - timedelta(days=2),
                            today + timedelta(days=2))
    assert hq == {tgt: 10}

    top = {r["item_id"]: r["total_used"] for r in
           authed_client.get("/api/reports/top-usage").json()}
    assert top[m["raw"]] == pytest.approx(11.0)            # yalnız p2: 10 × 1 × 1.1
    assert top[m["bottle"]] == pytest.approx(10)

    now = datetime.utcnow()
    rep = gather_report_data(db_session, now.year, now.month)
    assert [p["quantity"] for p in rep["production"]] == [10]
    assert [c["quantity"] for c in rep["cancelled"]] == [20]
    mats = {x["name"]: x["qty"] for x in rep["materials"]}
    assert mats["İptal Yağı"] == pytest.approx(11.0)
    assert all(NOTE_MARK not in (a["notes"] or "") for a in rep["adjustments"])
    assert [x["qty"] for x in rep["receiving"]] == [10]
    from io import BytesIO
    from openpyxl import load_workbook
    from core.monthly_report import render_excel, render_pdf
    assert render_pdf(rep)[:4] == b"%PDF"
    assert "İptal Edilen Üretim" in load_workbook(BytesIO(render_excel(rep))).sheetnames

    lst = {r["id"]: r for r in authed_client.get("/api/production").json()}
    assert lst[p1]["cancelled"]["reason"] == "Yanlış ürün seçildi"
    assert lst[p2]["cancelled"] is None
    assert authed_client.get(f"/api/production/{p1}").json()["cancelled"] is not None

    # İzlenebilirlik: iptal edilmemiş üretim tercih edilir, iptal bilgisi döner
    tr = authed_client.get("/api/traceability/lot/SR002").json()
    assert tr["production"]["id"] == p2 and tr["production"]["cancelled"] is None
    tr1 = authed_client.get("/api/traceability/lot/SR001").json()
    assert tr1["production"]["cancelled"] is not None
    assert tr1["lot_info"]["status"] == "CANCELLED"
    assert {i["source_lot"] for i in tr1["production"]["ingredients_consumed"]
            if i["item_id"] == m["raw"]} == {"L-A", "L-B"}


# ─── 10) Lot no serbest bırakma ─────────────────────────────────────────────

@pytest.mark.parametrize("release", [True, False])
def test_release_lot(authed_client, db_session, release):
    m = _materials(db_session)
    rid, tgt, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    fp = _preview(authed_client, pid)["fingerprint"]
    r = _cancel(authed_client, pid, fp, release_lot=release)
    assert r.status_code == 200 and r.json()["release_lot"] is release
    db_session.expire_all()
    item = db_session.get(Item, tgt)
    sug = lots.suggest(db_session, item)
    if release:
        assert lots.is_taken(db_session, tgt, "SR001") is False
        assert sug["lot_number"] == "SR001" and r.json()["lot_seq_rolled_back"] is True
        # Aynı lot no ile tekrar üretilebilir; yeni üretim de iptal edilebilir
        pid2 = _produce(authed_client, rid, qty=4, witness=0, lot="SR001")["production_id"]
        pv2 = _preview(authed_client, pid2)
        assert pv2["can_cancel"] is True, pv2["blockers"]
    else:
        assert lots.is_taken(db_session, tgt, "SR001") is True
        assert sug["lot_number"] == "SR002"
        bad = authed_client.post("/api/production", json={
            "recipe_id": rid, "produced_quantity": 1, "lot_number": "SR001"}, headers=_HDR)
        assert bad.status_code == 400


# ─── Üretimsiz hedef / boş tüketim kenarları ────────────────────────────────

def test_preview_404_for_unknown(authed_client):
    assert authed_client.get("/api/production/99999/cancel-preview").status_code == 404
    assert _cancel(authed_client, 99999, "x").status_code == 404


# ─── Eski üretim kenarları ──────────────────────────────────────────────────

def test_legacy_witness_split_read_from_input_note(authed_client, db_session):
    """Eski satırda witness_quantity 0 ise şahit adedi Input notundaki
    "Showroom: X, Şahit: Y"den okunur."""
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    db_session.query(ProductionConsumption).delete()
    db_session.get(ProductionHistory, pid).witness_quantity = 0
    db_session.commit()
    pv = _preview(authed_client, pid)
    assert pv["source"] == "ledger" and pv["can_cancel"] is True, pv["blockers"]
    assert sorted((r["lot"], r["qty"]) for r in pv["remove"]) == [("SR001", 18), ("SR001-S", 2)]
    assert pv["retention"]["qty"] == 2
    assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 200


def test_legacy_prd_lot_without_target(authed_client, db_session):
    """Hedefsiz reçete (lot_number NULL, notlarda PRD- lotu) — pencerede tek
    PRD- lotu varsa o üretimin Output'ları iade edilir."""
    m = _materials(db_session)
    rec = Recipe(name="Yarı mamul deneme", output_quantity=1, output_unit="adet",
                 waste_percentage=0, domain="cosmetics")
    db_session.add(rec); db_session.flush()
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=m["raw"], quantity=1))
    db_session.add(RecipeIngredient(recipe_id=rec.id, item_id=m["water"], quantity=2))
    db_session.commit()
    before = _state(db_session)["items"]
    pid = _produce(authed_client, rec.id, qty=5, witness=0)["production_id"]
    db_session.expire_all()
    assert db_session.get(ProductionHistory, pid).lot_number is None
    db_session.query(ProductionConsumption).delete()
    db_session.commit()

    pv = _preview(authed_client, pid)
    assert pv["source"] == "ledger" and pv["can_cancel"] is True, pv["blockers"]
    assert pv["production"]["lot"].startswith("PRD-") and pv["remove"] == []
    assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 200
    after = _state(db_session)["items"]
    for key in ("raw", "water"):
        assert after[m[key]] == pytest.approx(before[m[key]])


@pytest.mark.parametrize("card_active", [True, False])
def test_missing_source_lot_row(authed_client, db_session, card_active):
    """Kaynak lot satırı artık yok (birleştirmede silinmiş): kart aktifse
    yalnız kart stoğuna iade (uyarı), pasifse engel."""
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    lot_a = (db_session.query(Inventory).filter(Inventory.item_id == m["raw"],
                                                Inventory.lot_number == "L-A").one())
    db_session.delete(lot_a)                         # FK ondelete=SET NULL
    db_session.get(Item, m["raw"]).is_active = card_active
    db_session.commit()
    pv = _preview(authed_client, pid)
    if card_active:
        assert pv["can_cancel"] is True, pv["blockers"]
        row = next(r for r in pv["restore"] if r["lot"] == "L-A")
        assert row["inventory_id"] is None and "target_card_note" in row
        stock = db_session.get(Item, m["raw"]).current_stock
        assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 200
        db_session.expire_all()
        assert db_session.get(Item, m["raw"]).current_stock == pytest.approx(stock + 22.0)
    else:
        assert pv["can_cancel"] is False
        assert any("pasif ve kaynak lot" in b for b in pv["blockers"]), pv["blockers"]
        assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 409


def test_preview_of_cancelled_production(authed_client, db_session):
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    fp = _preview(authed_client, pid)["fingerprint"]
    assert _cancel(authed_client, pid, fp).status_code == 200
    pv = _preview(authed_client, pid)
    assert pv["can_cancel"] is False and pv["blockers"] == ["Bu üretim zaten iptal edilmiş."]
    assert pv["cancelled"]["by"] and pv["cancelled"]["reason"] == "Yanlış ürün seçildi"


# ─── İnceleme bulguları (07.10) ─────────────────────────────────────────────

def _strip_snapshots(db: Session, *pids):
    """Özellik öncesi gibi: verilen üretimlerin dökümü yok, şahit bağı yok."""
    q = db.query(ProductionConsumption)
    if pids:
        q = q.filter(ProductionConsumption.production_id.in_(pids))
    q.delete(synchronize_session=False)
    rq = db.query(RetentionSample)
    if pids:
        rq = rq.filter(RetentionSample.production_history_id.in_(pids))
    for rs in rq.all():
        rs.production_history_id = None
    db.commit()


def test_legacy_same_recipe_name_same_lot_is_ambiguous(authed_client, db_session):
    """İki FARKLI ürün, AYNI reçete adı, ikisi de kendi sayacından SR001, aynı
    kişi aynı dakikalarda — defter imzası ayrılamaz → engel (çift iade yok).
    Öbürünün dökümü duruyorsa (sahiplenilmiş) belirsizlik yoktur."""
    m = _materials(db_session)
    rid_a, tgt_a, _ = _product(db_session, m, suffix="A")
    a_name = db_session.get(Recipe, rid_a).name
    rid_c, tgt_c, _ = _product(db_session, m, suffix="C", recipe_name=a_name)
    pid_a = _produce(authed_client, rid_a)["production_id"]
    pid_c = _produce(authed_client, rid_c, qty=10, witness=0)["production_id"]
    db_session.expire_all()
    assert db_session.get(ProductionHistory, pid_c).lot_number == "SR001"
    a_tx = {r.transaction_id for r in db_session.query(ProductionConsumption)
            .filter(ProductionConsumption.production_id == pid_a,
                    ProductionConsumption.kind != "output")}

    # (1) Yalnız A'nın dökümü yok — C'ninkiler sahiplenilmiş, karışmaz.
    _strip_snapshots(db_session, pid_a)
    pv = _preview(authed_client, pid_a)
    assert pv["source"] == "ledger" and pv["can_cancel"] is True, pv["blockers"]
    assert {r["tx_id"] for r in pv["restore"]} == a_tx

    # (2) İkisinin de dökümü yok — belirsiz, iki yönde de engel; hiçbir şey değişmez.
    _strip_snapshots(db_session)
    for pid in (pid_a, pid_c):
        pv = _preview(authed_client, pid)
        assert pv["can_cancel"] is False
        assert any("ayrılamıyor" in b for b in pv["blockers"]), pv["blockers"]
    before = _state(db_session)
    r = _cancel(authed_client, pid_a, _preview(authed_client, pid_a)["fingerprint"])
    assert r.status_code == 409 and r.json()["code"] == "blocked"
    assert _state(db_session) == before

    # (3) C'nin PH kaydı olmasa bile defterdeki ikinci Input belirsizliği gösterir.
    db_session.query(ProductionHistory).filter(ProductionHistory.id == pid_c).delete()
    db_session.commit()
    pv = _preview(authed_client, pid_a)
    assert pv["can_cancel"] is False
    assert any("ayrılamıyor" in b for b in pv["blockers"]), pv["blockers"]


def test_absorbed_source_lot_restores_to_twin_on_new_card(authed_client, db_session):
    """Kaynak lot üretimden sonra aynı lot no'lu ikizi olan karta taşındı
    (satır ikize katılıp silindi): döküm ikize yönlenir, iade lotun BUGÜNKÜ
    kartına gider — eski kartta hayalî stok oluşmaz."""
    m = _materials(db_session)
    s2 = db_session.query(Supplier).filter(Supplier.name == "Tedarikçi İki").one()
    twin = Inventory(item_id=m["raw2"], supplier_id=s2.id, lot_number="L-B", quantity=5.0,
                     status="APPROVED", domain="cosmetics",
                     created_at=datetime.utcnow() - timedelta(days=5))
    db_session.add(twin)
    db_session.get(Item, m["raw2"]).current_stock = 5.0
    db_session.commit()
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]

    lot_b = (db_session.query(Inventory).filter(Inventory.item_id == m["raw"],
                                                Inventory.lot_number == "L-B").one())
    lot_b_id = lot_b.id
    res = stock_lots.move_lot(db_session, lot_b, db_session.get(Item, m["raw"]),
                              db_session.get(Item, m["raw2"]), None, actor="test",
                              reason="yanlış kart")
    db_session.commit()
    assert res["merged"] and db_session.get(Inventory, lot_b_id) is None
    pc = (db_session.query(ProductionConsumption)
          .filter(ProductionConsumption.production_id == pid,
                  ProductionConsumption.lot_number == "L-B").one())
    assert pc.inventory_id == twin.id                    # SET NULL'a düşmedi

    pv = _preview(authed_client, pid)
    assert pv["can_cancel"] is True, pv["blockers"]
    row = next(r for r in pv["restore"] if r["lot"] == "L-B")
    assert row["item_id"] == m["raw2"] and row["inventory_id"] == twin.id
    raw_before = db_session.get(Item, m["raw"]).current_stock
    raw2_before = db_session.get(Item, m["raw2"]).current_stock
    twin_before = db_session.get(Inventory, twin.id).quantity
    assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 200
    db_session.expire_all()
    assert db_session.get(Inventory, twin.id).quantity == pytest.approx(twin_before + row["qty"])
    assert db_session.get(Item, m["raw2"]).current_stock == pytest.approx(raw2_before + row["qty"])
    # Eski kart yalnız kendi lotu L-A'nın iadesini alır (3 g) — lot toplamıyla tutarlı.
    raw = db_session.get(Item, m["raw"])
    lots_on_raw = sum(i.quantity for i in db_session.query(Inventory)
                      .filter(Inventory.item_id == m["raw"]).all())
    assert raw.current_stock == pytest.approx(raw_before + 3.0)
    assert raw.current_stock == pytest.approx(lots_on_raw)


def test_absorbed_source_lot_blocks_ledger_reconstruction(authed_client, db_session):
    """Eski üretim (döküm yok) + kaynak lot taşınıp ikize katılmış: lotun
    bugünkü yeri defterden bilinemez → engel (eski karta lotsuz iade yok)."""
    m = _materials(db_session)
    s2 = db_session.query(Supplier).filter(Supplier.name == "Tedarikçi İki").one()
    db_session.add(Inventory(item_id=m["raw2"], supplier_id=s2.id, lot_number="L-B",
                             quantity=5.0, status="APPROVED", domain="cosmetics",
                             created_at=datetime.utcnow() - timedelta(days=5)))
    db_session.get(Item, m["raw2"]).current_stock = 5.0
    db_session.commit()
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    _strip_snapshots(db_session)
    lot_b = (db_session.query(Inventory).filter(Inventory.item_id == m["raw"],
                                                Inventory.lot_number == "L-B").one())
    stock_lots.move_lot(db_session, lot_b, db_session.get(Item, m["raw"]),
                        db_session.get(Item, m["raw2"]), None, actor="test", reason="yanlış kart")
    db_session.commit()
    pv = _preview(authed_client, pid)
    assert pv["source"] == "ledger" and pv["can_cancel"] is False
    assert any("taşınıp oradaki aynı lotla birleşmiş" in b for b in pv["blockers"]), pv["blockers"]
    before = _state(db_session)
    assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 409
    assert _state(db_session) == before


def test_lotless_row_on_merged_card_is_blocked(authed_client, db_session):
    """Lotsuz tüketim (etiket) kartı sonradan kopya karta birleştirildi
    (pasif) — iade pasif karta yazılmaz, engel."""
    from core.item_merge import merge_items
    m = _materials(db_session)
    dup = Item(name="İptal Etiket EN (kopya)", sku="c-len2", category="Ambalaj",
               pkg_type="etiket", language="EN", unit="adet", current_stock=1010,
               domain="cosmetics")
    db_session.add(dup); db_session.commit()
    rid, _, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    merge_items(db_session, m["len"], dup.id, "test")
    db_session.commit()
    pv = _preview(authed_client, pid)
    assert pv["can_cancel"] is False
    assert any("İptal Etiket EN" in b and "pasif" in b for b in pv["blockers"]), pv["blockers"]
    before = _state(db_session)
    assert _cancel(authed_client, pid, pv["fingerprint"]).status_code == 409
    assert _state(db_session) == before


def test_recipe_label_card_referenced_only_by_snapshot_soft_deletes(authed_client, db_session):
    """Reçetedeki TR etiket yalnız `recipe_item_id`'de geçiyor (üretim EN
    kardeşinden düştü): reçeteden çıkarılıp silinince 500 değil soft-delete."""
    m = _materials(db_session)
    rid, _, _ = _product(db_session, m)
    _produce(authed_client, rid, lang="EN")
    db_session.query(RecipeIngredient).filter(RecipeIngredient.item_id == m["ltr"]).delete()
    db_session.commit()
    assert db_session.query(Transaction).filter(Transaction.item_id == m["ltr"]).count() == 0
    r = authed_client.delete(f"/api/items/{m['ltr']}", headers=_HDR)
    assert r.status_code == 200, r.text
    assert r.json()["soft_deleted"] is True
    db_session.expire_all()
    assert db_session.get(Item, m["ltr"]).is_active is False


def test_trace_lot_pairs_lot_row_with_same_product(authed_client, db_session):
    """Aynı lot no (SR001) iki üründe: lot satırı ile üretim kaydı AYNI ürüne
    bağlanır; öbür ürün `other_items`'ta, `item_id` ile seçilir.  Dökümü
    olmayan eski üretimde de yalnız kendi Output'ları listelenir."""
    m = _materials(db_session)
    rid_a, tgt_a, _ = _product(db_session, m, suffix="A")
    rid_b, tgt_b, _ = _product(db_session, m, suffix="B")
    pid_a = _produce(authed_client, rid_a)["production_id"]
    pid_b = _produce(authed_client, rid_b, qty=5, witness=1)["production_id"]
    tx_of = lambda pid: {r.transaction_id for r in db_session.query(ProductionConsumption)  # noqa: E731
                         .filter(ProductionConsumption.production_id == pid,
                                 ProductionConsumption.kind != "output")}
    a_tx, b_tx = tx_of(pid_a), tx_of(pid_b)

    d = authed_client.get("/api/traceability/lot/SR001").json()
    assert d["lot_info"]["item_id"] == d["production"]["target_item_id"] == tgt_b
    assert [o["item_id"] for o in d["other_items"]] == [tgt_a]

    d = authed_client.get(f"/api/traceability/lot/SR001?item_id={tgt_a}").json()
    assert d["lot_info"]["item_id"] == d["production"]["target_item_id"] == tgt_a
    assert d["production"]["id"] == pid_a
    assert len(d["production"]["ingredients_consumed"]) == len(a_tx)

    # Eski kayıt gibi (döküm yok) — defter imzasıyla, B karışmadan
    _strip_snapshots(db_session)
    d = authed_client.get(f"/api/traceability/lot/SR001?item_id={tgt_a}").json()
    assert len(d["production"]["ingredients_consumed"]) == len(a_tx)
    assert not (a_tx & b_tx)


@pytest.mark.parametrize("action", ["qc_reject", "witness_checkout"])
def test_concurrent_qc_or_checkout_waits_for_cancel(authed_client, db_session, action):
    """İptal transaction'ı lotları kilitliyken gelen QC reddi / şahit çıkışı
    beklemeli ve kilit bırakılınca güncel durumu (CANCELLED / pasif) görüp
    reddedilmeli — bayat kopya üzerinden ikinci düşüm YOK."""
    import threading
    from core.production_cancel import apply_cancel
    m = _materials(db_session)
    rid, tgt, _ = _product(db_session, m)
    pid = _produce(authed_client, rid)["production_id"]
    showroom = (db_session.query(Inventory).filter(Inventory.item_id == tgt,
                                                   Inventory.lot_number == "SR001").one())
    rs = db_session.query(RetentionSample).one()
    if action == "qc_reject":
        call = lambda: authed_client.post(f"/api/qc/process/{showroom.id}",  # noqa: E731
                                          json={"status": "REJECTED", "notes": "bozuk"},
                                          headers=_HDR)
        expect = 400
    else:
        call = lambda: authed_client.post(f"/api/retention/samples/{rs.id}/checkout",  # noqa: E731
                                          json={"quantity": 1, "reason": "test"}, headers=_HDR)
        expect = 404

    prod = (db_session.query(ProductionHistory).filter(ProductionHistory.id == pid)
            .with_for_update().one())
    apply_cancel(db_session, prod, actor="test", reason="eşzamanlı iptal testi")   # kilitli, commit yok
    out = {}
    th = threading.Thread(target=lambda: out.setdefault("r", call()))
    th.start()
    th.join(timeout=1.5)
    assert th.is_alive(), "istek iptalin kilidini beklemeliydi"
    db_session.commit()
    th.join(timeout=20)
    assert not th.is_alive()
    assert out["r"].status_code == expect, out["r"].text
    db_session.expire_all()
    assert db_session.get(Item, tgt).current_stock == pytest.approx(0.0)
    assert db_session.get(Inventory, showroom.id).status == "CANCELLED"
    extra = (db_session.query(Transaction)
             .filter(Transaction.item_id == tgt,
                     ~Transaction.notes.like(f"{NOTE_MARK}%"),
                     Transaction.transaction_type.in_(("Output", "Adjustment"))).count())
    assert extra == 0
