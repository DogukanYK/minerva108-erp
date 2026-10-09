"""SKT kontrol raporu + lot SKT düzeltme (İzlenebilirlik, 09.10.2026).

Kilitlenen sözleşmeler:
  • Tek ayrıştırıcı (core/lots.parse_expiry): YYYY-MM-DD · GG.AA.YYYY · GG/AA/YYYY;
    "bugün" TR takvim günü.  Rapor ile B2B sıkı havuzu (strict_lot_ok) AYNI kararı verir.
  • Rapor: aktif panel, numune olmayan, miktarı > 0, APPROVED/QC bekleyen hammadde
    lotları; eksik / okunamayan / geçmiş; scope=recipes yalnız aktif reçetedekiler.
  • PATCH: inventory.adjust, panel, ISO saklar, eski → yeni audit, numune 400; stok değişmez.
Yalnız izole test DB'si; sentetik adlar.
"""
import json
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core import expiry_report
from core.b2b_orders import strict_lot_ok
from core.lots import expiry_today, parse_expiry
from database import AdminAuditLog, Inventory, Item, Recipe, RecipeIngredient, Transaction

HDR = {"Origin": "http://testserver"}


def _item(db, name, category="Hammadde", unit="g", domain="cosmetics", active=True):
    it = Item(name=name, sku=f"SKT-{abs(hash(name)) % 10**7}", unit=unit, category=category,
              current_stock=100, domain=domain, is_active=active)
    db.add(it)
    db.flush()
    return it


def _lot(db, item, lot_number, expiry, qty=10, status="APPROVED", sample=False, domain=None):
    inv = Inventory(item_id=item.id, lot_number=lot_number, expiry_date=expiry, quantity=qty, status=status,
                    qc_required=(status == "QUARANTINE"), is_sample=sample, domain=domain or item.domain)
    db.add(inv)
    db.flush()
    return inv


def _recipe(db, name, target, ingredients, active=True, domain="cosmetics"):
    rec = Recipe(name=name, target_item_id=target.id, output_quantity=1, output_unit="adet",
                 waste_percentage=0, domain=domain, is_active=active)
    db.add(rec)
    db.flush()
    for it in ingredients:
        db.add(RecipeIngredient(recipe_id=rec.id, item_id=it.id, quantity=1, unit=it.unit))
    return rec


def test_parse_expiry_formats_and_tr_today():
    assert parse_expiry("2027-03-01") == date(2027, 3, 1)
    assert parse_expiry(" 01.03.2027 ") == date(2027, 3, 1)
    assert parse_expiry("01/03/2027") == date(2027, 3, 1)
    for bad in (None, "", "Ekim 2027", "2027-13-01", "03/2027", "31.02.2027"):
        assert parse_expiry(bad) is None, bad
    from database import tr_now
    assert expiry_today() == tr_now().date()
    lot = SimpleNamespace(is_sample=False, outsourcing_receipt_id=None, status="APPROVED")
    assert expiry_report.editable(lot)
    assert not expiry_report.editable(SimpleNamespace(**{**vars(lot), "outsourcing_receipt_id": 5}))
    assert not expiry_report.editable(SimpleNamespace(**{**vars(lot), "is_sample": True}))
    assert not expiry_report.editable(SimpleNamespace(**{**vars(lot), "status": "REJECTED"}))


def _world(db):
    today = expiry_today()
    future = (today + timedelta(days=400)).isoformat()
    soon_tr = (today + timedelta(days=10)).strftime("%d.%m.%Y")
    a = _item(db, "SKT Baz Yağ")                  # aktif reçetede
    b = _item(db, "SKT Reçetesiz Yağ")            # yalnız pasif reçetede
    c = _item(db, "SKT Şişe", category="Ambalaj", unit="adet")
    cream = _item(db, "SKT Krem", category="Bitmiş Ürün", unit="adet")
    other = _item(db, "SKT Takviye Tozu", domain="supplement")
    _recipe(db, "SKT Krem Reçetesi", cream, [a, c])
    _recipe(db, "SKT Eski Reçete", cream, [b], active=False)
    lots = {
        "missing": _lot(db, a, "A-BOS", None),
        "expired": _lot(db, a, "A-ESKI", "31/12/2020"),
        "unreadable": _lot(db, a, "A-YAZI", "Ekim 2027"),
        "valid": _lot(db, a, "A-GECERLI", future),
        "soon_tr": _lot(db, a, "A-YAKIN", soon_tr),
        "quarantine": _lot(db, a, "A-QC", "", status="QUARANTINE"),
        "sample": _lot(db, a, "A-NUMUNE", None, sample=True),
        "empty": _lot(db, a, "A-SIFIR", None, qty=0),
        "rejected": _lot(db, a, "A-RED", None, status="REJECTED"),
        "b_missing": _lot(db, b, "B-BOS", None),
        "c_missing": _lot(db, c, "C-BOS", None),
        "finished": _lot(db, cream, "K-BOS", None),
        "other_domain": _lot(db, other, "S-BOS", None),
    }
    db.commit()
    return {"a": a, "b": b, "c": c, "lots": lots, "today": today}


def test_report_scopes_classification_and_b2b_agreement(db_session):
    db = db_session
    w = _world(db)
    lots = w["lots"]
    rep = expiry_report.expiry_issues(db, "cosmetics", "recipes")
    by_lot = {r["lot_number"]: r for r in rep["rows"]}
    assert set(by_lot) == {"A-BOS", "A-ESKI", "A-YAZI", "A-QC"}
    assert rep["summary"] == {"expired": 1, "missing": 2, "unreadable": 1, "total": 4}
    assert [r["issue"] for r in rep["rows"]][0] == "expired"            # en acil önce
    assert by_lot["A-ESKI"]["days_expired"] == (w["today"] - date(2020, 12, 31)).days
    assert by_lot["A-ESKI"]["expiry_iso"] == "2020-12-31"
    assert by_lot["A-QC"]["qc_pending"] and by_lot["A-QC"]["issue"] == "missing"
    assert by_lot["A-BOS"]["recipes"] == ["SKT Krem Reçetesi"] and by_lot["A-BOS"]["editable"]
    # scope=all: reçetesiz hammadde de; ambalaj, bitmiş ürün, numune, boş, reddedilen, başka panel ASLA
    all_rows = {r["lot_number"] for r in expiry_report.expiry_issues(db, "cosmetics", "all")["rows"]}
    assert all_rows == {"A-BOS", "A-ESKI", "A-YAZI", "A-QC", "B-BOS"}
    sup = {r["lot_number"] for r in expiry_report.expiry_issues(db, "supplement", "all")["rows"]}
    assert sup == {"S-BOS"}
    # Rapordaki lot sıkı havuzda yok; raporda olmayan APPROVED lot havuzda (GG.AA.YYYY dahil)
    for key in ("missing", "expired", "unreadable"):
        assert strict_lot_ok(lots[key]) is False
    assert strict_lot_ok(lots["valid"]) and strict_lot_ok(lots["soon_tr"])


def test_patch_fixes_expiry_with_audit_and_no_stock_change(client: TestClient, db_session):
    db = db_session
    w = _world(db)
    lots = w["lots"]
    client.post("/api/login", json={"username": "dogukan", "password": "minerva123"}, headers=HDR)
    r = client.get("/api/traceability/expiry-issues")
    assert r.status_code == 200 and r.json()["summary"]["total"] == 4
    stock_before = db.get(Item, w["a"].id).current_stock
    tx_before = db.query(Transaction).count()

    r = client.patch(f"/api/inventory/lots/{lots['missing'].id}/expiry", headers=HDR,
                     json={"expiry_date": "31.12.2027", "note": "CoA'dan"})
    assert r.status_code == 200, r.text
    assert r.json() == {"ok": True, "changed": True, "id": lots["missing"].id, "expiry_date": "2027-12-31",
                        "previous": "", "issue": None}
    db.expire_all()
    assert db.get(Inventory, lots["missing"].id).expiry_date == "2027-12-31"
    assert db.get(Item, w["a"].id).current_stock == stock_before and db.query(Transaction).count() == tx_before
    log = db.query(AdminAuditLog).filter_by(action="inventory.lot_expiry").one()
    assert json.loads(log.details) == {"item_id": w["a"].id, "lot_number": "A-BOS", "old": "",
                                       "new": "2027-12-31", "note": "CoA'dan"}
    assert strict_lot_ok(db.get(Inventory, lots["missing"].id))
    # Aynı tarih tekrar → değişiklik yok, ikinci audit yok
    r = client.patch(f"/api/inventory/lots/{lots['missing'].id}/expiry", headers=HDR,
                     json={"expiry_date": "2027-12-31"})
    assert r.json()["changed"] is False
    assert db.query(AdminAuditLog).filter_by(action="inventory.lot_expiry").count() == 1
    # Geçmiş tarih kaydedilebilir (gerçekten geçmişse) ama sorun "expired" kalır
    r = client.patch(f"/api/inventory/lots/{lots['unreadable'].id}/expiry", headers=HDR,
                     json={"expiry_date": "2021-01-01"})
    assert r.status_code == 200 and r.json()["issue"] == "expired"
    rows = {x["lot_number"]: x for x in client.get("/api/traceability/expiry-issues").json()["rows"]}
    assert "A-BOS" not in rows and rows["A-YAZI"]["issue"] == "expired"

    for bad in ("2027-02-30", "Ekim 2027", "1899-01-01"):
        r = client.patch(f"/api/inventory/lots/{lots['expired'].id}/expiry", headers=HDR,
                         json={"expiry_date": bad})
        assert r.status_code in (400, 422) and (r.status_code == 422 or r.json()["code"] == "invalid_expiry")
    r = client.patch(f"/api/inventory/lots/{lots['sample'].id}/expiry", headers=HDR,
                     json={"expiry_date": "2027-01-01"})
    assert r.status_code == 400 and r.json()["code"] == "not_editable"
    r = client.patch(f"/api/inventory/lots/{lots['other_domain'].id}/expiry", headers=HDR,
                     json={"expiry_date": "2027-01-01"})
    assert r.status_code == 404                                         # başka panelin lotu

    # LabTech (inventory.adjust yok) düzeltemez, raporu görür
    client.post("/api/login", json={"username": "meltem", "password": "minerva123"}, headers=HDR)
    assert client.get("/api/traceability/expiry-issues").status_code == 200
    r = client.patch(f"/api/inventory/lots/{lots['expired'].id}/expiry", headers=HDR,
                     json={"expiry_date": "2027-01-01"})
    assert r.status_code == 403


def test_excel_export_and_expiring_reads_tr_dates(authed_client: TestClient, db_session):
    w = _world(db_session)
    r = authed_client.get("/api/traceability/expiry-issues/export?scope=all")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(r.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    head = next(i for i, row in enumerate(values) if row[0] == "Sorun")
    lots = {row[3] for row in values[head + 1:]}
    assert lots == {"A-BOS", "A-ESKI", "A-YAZI", "A-QC", "B-BOS"}
    # "Yaklaşan SKT" artık GG.AA.YYYY biçimini de okur
    exp = authed_client.get("/api/traceability/expiring").json()
    soon = next(x for x in exp if x["lot_number"] == "A-YAKIN")
    assert soon["days_left"] == 10 and w["lots"]["soon_tr"].id == soon["id"]
