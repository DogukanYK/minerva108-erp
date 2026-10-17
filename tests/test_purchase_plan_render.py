"""
Satın Alma Planı çıktıları — PDF (core/purchase_plan_pdf.py), Excel
(core/purchase_plan_xlsx.py) ve formül önbelleği (core/xlsx_cache.py).

    MINERVA_TEST_DB=minerva_test_u .venv/bin/pytest tests/test_purchase_plan_render.py -q

Rapor sözlüğü `compute()` + `attach()` ile sentetik girdilerden kurulur (DB'ye
yazılmaz).  Kapsam: PDF (16.10.2026 yeniden tasarım: Detaylı Rapor = dikey
kapak + yönetici bölümü + içindekiler, sonra yatay operasyon bölümleri; Yönetici
Özeti = dikey ≤ 4 sayfa, 3 grafik; "Sayfa X / Y"; yer imleri; Ek A teklif /
ilişki detayı, Ek B fiyatı eksikler) + "Alındı" + bölüm başlıkları/numaraları + boş
bölüm basılmaz + kaçışlı ad + brüt/yolda sütun başlıkları; Excel sayfaları,
Tutar hücresinde formül VE gömülü `<v>`, `data_only` okuma, ara/genel toplam,
brüt modda G=D, serbest metin formül kalkanı ("=HYPERLINK…" adı / notu /
serbest satırı formül olarak yazılmaz); birim uyuşmazlığındaki teklifin
nedeni ("fiyat yazılmamış" değil) PDF + Excel'de, ambalaj katı Excel "Not"
sütununda, ondalıklı serbest miktar biçimi, kontrol karakterli ad Excel'i
düşürmez; xlsx_cache hata ve tür durumları; "aynı malzeme" + sipariş
satırları (PDF'te en çok 3 + "+N", Excel'de "Not"tan sonra Y/Z sütunları,
Tedarikçiler K "İlişki türleri"); grup stoğu dökümü PDF stok hücresinde ve
Excel AA (27.) sütununda (24–26 aynı); stoktan karşılanan değer —
"Stoktan karşılananlar" sayfası (formül + önbellek + TOPLAM, açık siparişte
L/M), Alım listesi GENEL TOPLAM'ı değişmez, Özet plan toplamı, PDF "Değer" /
"stoktan" / "Planın toplam malzeme maliyeti" (brüt modda hiçbiri); ambalaj
uyarısı (16.10.2026) Alım listesi K turuncu + hücre yorumu, PDF ⚠ + "Not:";
"Ürün maliyeti" sayfası (formül + önbellek, ara toplam, TOPLAM, tür bazında
maliyet, eksik kalemler) ve PDF "Maliyet özeti" bölümü.
Excel (16.10.2026): ilk sayfa "Yönetici Özeti" (`summary_data`), lacivert
başlık, süzgeç, yazdırma ayarı.  `summary_data` toplamları `pricing.totals` /
`cost_summary` / `cost_overview` ile aynı; grafik serileri tablo hücreleriyle.
Sondaki test `RUSYA_FIXTURE_DIR` varsa Rusya kabul senaryosunu iki çıktıya da
basar (toplamlar PDF ↔ Excel ↔ rapor aynı; özet ≤ 3, yönetici bölümü ≤ 3 dikey
sayfa, detay ≤ 40 sayfa; bölüm 7 bölüm 6'nın sayfasında).  17.10.2026 görsel QA
düzeltmeleri: iç içe KeepTogether yok, dipnot numarası kesintisiz + tek kaçış,
lejant sırası = çubuk sırası, etiket boşluğu veriden, eksen "0", özet PDF'te yer
imi paneli yok, Excel süzgeci ara toplamı kapsamaz, kapakta tek "fiyatı yok"
tanımı, yeterli satır bilgi notu / ürün eksik adları / bitmiş stok erişilebilir.
İkinci tur: tür "Alınacak" = kart + yeni kalem ve anahtarla sayım, maliyet uyarısı
KPI terimleriyle, `_amt` yarım yukarı, eksen hanesi adımdan, lejant kenarsız, grafik
başlıkları, gram ölçeği, sayı–birim bölünmez, Ek B satır sırası + "tam liste" yalnız
kırpılmadıysa, ürün bilgi uyarısı, Excel tür sütunları moda göre.
"""
import importlib.util
import io
import os
import re
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfReader

from core.consumption import IngredientRec, ItemRec, RecipeRec
from core.purchase_plan import LotRec, OpenOrderRec, PlanInputs, compute
from core.purchase_plan_models import PlanRequest
from core.purchase_plan_pdf import render_pdf
from core.purchase_plan_xlsx import SHEET_DETAIL, SHEET_EXEC, SHEET_LIST, SHEET_STOCK, SHEET_SUP, build_workbook
from core.purchase_pricing import (OfferRec, OrderRec, PriceInputs, ReceiptRec, SupplierRec, alt_texts, attach,
                                   money, order_texts, sections, summary_data)
from core.report_theme import Theme
from core.xlsx_cache import inject_cached_values

NOW = datetime(2026, 10, 5, 9, 0)
EVIL_NAME = '=HYPERLINK("http://evil.example","tıkla")'


def _it(id, name, *, cat="Hammadde", unit="g", stock=0.0, pkg=None):
    return ItemRec(id=id, name=name, category=cat, unit=unit, pkg_type=pkg, current_stock=stock)


def _report(*, mode="net", open_orders=False, with_packaging=False, surya_preferred=False):
    items = [
        _it(1, "GLİSERİN", stock=500),
        _it(2, "SHEA & BUTTER <özel>"),
        _it(3, "GÜL ABSOLÜ", unit="ml"),
        _it(4, EVIL_NAME),
        _it(5, "KSANTAN GAM", stock=1_000_000),
        _it(6, "ALOE VERA TOZU"),
    ]
    ings = [(1, 10), (2, 5), (3, 2), (4, 1), (5, 3), (6, 1)]
    if with_packaging:
        items.append(_it(7, "50 ML KAVANOZ", cat="Ambalaj", unit="adet", pkg="kavanoz"))
        ings.append((7, 1))
    recipe = RecipeRec(id=1, name="Gece Kremi 50 ml",
                       ingredients=tuple(IngredientRec(item_id=i, quantity=q) for i, q in ings))
    oo = {1: [OpenOrderRec(item_id=1, quantity=2000, unit="g")]} if open_orders else {}
    inp = PlanInputs(items={i.id: i for i in items}, recipes={1: recipe}, open_orders=oo, stock_as_of=NOW)
    req = PlanRequest(
        title="Deneme Siparişi",
        lines=[{"recipe_id": 1, "qty": 1000, "display_name": "Gece Kremi 50 ml", "label_faces": 2}],
        manual_lines=[{"section": "shipping", "name": "=1+1 koli", "qty": 40, "unit": "adet",
                       "note": "=2 adet kırık"},
                      {"section": "shipping", "name": "Streç film", "qty_text": "2 palet için"}],
        options={"stock_mode": mode, "subtract_open_orders": open_orders, "label_mode": "new",
                 "new_label_title": "Rusça etiket", "excluded_item_ids": [],
                 "held_items": [{"item_id": 6, "reason": "Lab teyidi bekleniyor"}],
                 "checklist_owner": "Satın alma ekibi",
                 "notes": ["=cmd|' /C calc'!A0", "Normal bir not."]})
    res = compute(inp, req)
    pin = PriceInputs(
        offers={1: [OfferRec(item_id=1, supplier_name="UMAYCHEM", supplier_id=1, unit_price=4.52,
                             package_size=25, price_unit="kg", source="stok_son_durum",
                             source_label="Stok Son Durum — Eylül"),
                    OfferRec(item_id=1, supplier_name="SURYA KİMYA", unit_price=5.0, package_size=25,
                             price_unit="kg", source="stok_son_durum")],
                2: [OfferRec(item_id=2, supplier_name="BEFCHEM", unit_price=8.0, package_size=25,
                             price_unit="kg", source="stok_son_durum")]},
        suppliers={1: SupplierRec(id=1, name="UMAYCHEM", contact_person="Semih Bey", phone="05537953163",
                                  email="info@umaychem.example"),
                   2: SupplierRec(id=2, name="DOALİNN")},
        card_supplier={3: 2})
    if surya_preferred:
        pin.suppliers[3] = SupplierRec(id=3, name="SURYA KİMYA", status="preferred")
    return attach(res, pin, None, currency="USD")


def _pdf_text(pdf: bytes):
    """(okuyucu, boşlukları tekleştirilmiş metin).  PDF'te kesme işareti
    tipografik (’, `report_theme.e`); içerik denetimleri düz yazımla yazılsın
    diye burada geri çevrilir — tipografinin kendisi
    `test_pdf_typographic_apostrophe_…` testinde ham metinle denetlenir."""
    rd = PdfReader(io.BytesIO(pdf))
    return rd, " ".join(" ".join((p.extract_text() or "") for p in rd.pages).split()).replace("\u2019", "'")


def _mat(res, name):
    return next(m for m in res["materials"] if m["name"] == name)


def _row_of(ws, name):
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == name)


# ─── PDF ────────────────────────────────────────────────────────────────────

def _portrait(p) -> bool:
    return float(p.mediabox.width) < float(p.mediabox.height)


def test_pdf_landscape_sections_and_alindi():
    """Detaylı Rapor: kapak + yönetici bölümü dikey, operasyon bölümleri yatay
    (16.10.2026 — eskiden her sayfa yataydı)."""
    res = _report()
    pdf = render_pdf(res)
    assert pdf.startswith(b"%PDF")
    rd, text = _pdf_text(pdf)
    assert len(rd.pages) >= 2
    assert _portrait(rd.pages[0]) and not _portrait(rd.pages[-1])
    assert "Alındı" in text
    assert "Deneme Siparişi" in text and re.search(rf"Sayfa 1 / {len(rd.pages)}\b", text)
    # Bölüm numaraları sections() ile: ambalaj/etiket bölümleri boş → atlanır, numara kaymaz
    for t in ("1. Hammadde — fiyat listesinde olanlar", "2. Hammadde — fiyatı olmayanlar, teklif alınacak",
              "3. Yeni etiket / koli / palet — teklif alınacak", "4. Kimlerle görüşeceğiz — tedarikçi bazında",
              "Fiyatı olmayanlar için teklif istenebilecek firmalar",
              "Tedarikçi bilgi eksikleri (Satın alma ekibi için)",
              "Ürünler, hedef adetler ve kapasite", "Yeterli olanlar", "Notlar ve kaynaklar"):
        assert t in text, t
    assert "Ambalaj — fiyatlı" not in text and "Ambalaj — fiyatı olmayanlar" not in text
    # Kutular + tedarikçi hücresi
    assert "teklif alınacak; gerçek toplam, fiyatı olmayan kalemler eklenince bundan yüksek olacak." in text
    assert "bekletiliyor" in text and "ALOE VERA TOZU" in text           # turuncu uyarı kutusu
    assert "UMAYCHEM — 4,52 $/kg · 25 kg'lık ambalaj" in text
    assert "SURYA KİMYA — 5,00 $/kg · 25 kg'lık ambalaj (daha pahalı)" in text
    assert "Fiyat yok — teklif alınacak" in text
    assert "Stok kartında yazan: DOALİNN" in text
    assert "İletişim bilgisi sistemde yok — Satın alma ekibi girecek" in text
    # XML'e özel karakterli ad kaçışlanır (PDF düşmez, ad aynen görünür)
    assert "SHEA & BUTTER <özel>" in text


def test_pdf_subtitle_amount_basis_follows_selection_rule():
    """Başlık altındaki "Tutar = …" cümlesi notlar / bölüm alt başlıklarıyla
    çelişmesin: seçimi tercih değiştirdiyse "seçilen teklif" der."""
    _, text = _pdf_text(render_pdf(_report()))
    assert "Tutar = alınacak × en ucuz birim fiyat." in text
    res = _report(surya_preferred=True)
    assert _mat(res, "GLİSERİN")["supplier"] == "SURYA KİMYA" and res["pricing"]["status_used"] is True
    _, text = _pdf_text(render_pdf(res))
    assert "en ucuz birim fiyat" not in text
    assert "Tutar = alınacak × seçilen teklifin birim fiyatı (önce malzeme tercihi" in text


def test_pdf_empty_section_not_printed_and_packaging_when_present():
    _, text = _pdf_text(render_pdf(_report(with_packaging=True)))
    assert "Ambalaj — fiyatı olmayanlar" in text
    assert "Ambalaj — fiyatlı" not in text                    # fiyatlı ambalaj yok → bölüm yok
    assert "3. Ambalaj — fiyatı olmayanlar" in text


def test_pdf_gross_and_open_order_columns():
    _, text = _pdf_text(render_pdf(_report(mode="gross")))
    assert "Elimizde (bilgi)" in text
    assert "Yeterli olanlar" not in text                      # brüt modda bu bölüm yok
    _, text = _pdf_text(render_pdf(_report(open_orders=True)))
    assert "Yolda" in text and "Elimizde (bilgi)" not in text


# ─── Excel ──────────────────────────────────────────────────────────────────

def test_xlsx_sheets_formulas_and_cached_values():
    res = _report()
    data = build_workbook(res)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # "Stoktan karşılananlar" (14.10.2026) net modda Alım listesi'nden hemen sonra;
    # "Ürün maliyeti" (16.10.2026) "Ürünler"den sonra
    # "Yönetici Özeti" (16.10.2026) İLK sayfa; ötekilerin sırası değişmez
    assert wb.sheetnames == ["Yönetici Özeti", "Alım listesi", "Stoktan karşılananlar", "Tedarikçiler", "Ürünler",
                             "Ürün maliyeti", "Ürün detayı", "Özet"]
    ws = wb[SHEET_LIST]
    assert ws.freeze_panes == "D2"
    r = _row_of(ws, "GLİSERİN")
    assert ws[f"G{r}"].value == f"=MAX(0,ROUND(D{r}-E{r}-F{r},3))"
    assert ws[f"K{r}"].value == f'=IF(J{r}="","",ROUND(G{r}*J{r},0))'

    # Tutar hücresinde formül VE gömülü <v> (XML'den) — sayfa dosyası ada göre bulunur
    xml = zipfile.ZipFile(io.BytesIO(data)).read(f"xl/worksheets/sheet{wb.sheetnames.index(SHEET_LIST) + 1}.xml").decode()
    gl = _mat(res, "GLİSERİN")
    cell = re.search(rf'<c r="K{r}"[^>]*>(.*?)</c>', xml).group(1)
    assert "<f>" in cell and f"<v>{gl['amount']}</v>" in cell
    assert not re.search(r"<f>[^<]*</f><v ?/>", xml)                  # önbelleksiz formül kalmadı

    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_LIST]
    assert wv[f"K{r}"].value == gl["amount"]
    assert wv[f"G{r}"].value == pytest.approx(gl["display"]["buy_num"])
    ru = _row_of(ws, "GÜL ABSOLÜ")                                  # fiyatsız → boş metin sonucu
    assert wv[f"K{ru}"].value in (None, "")
    tot = next(x for x in range(2, ws.max_row + 1) if ws.cell(x, 3).value == "GENEL TOPLAM (fiyatı belli olanlar)")
    assert wv[f"K{tot}"].value == res["pricing"]["totals"]["all"] > 0
    sub = next(x for x in range(2, ws.max_row + 1)
               if str(ws.cell(x, 3).value or "").startswith("1. Hammadde — fiyat listesinde olanlar — toplam"))
    assert ws[f"K{sub}"].value.startswith("=SUM(K")
    assert wv[f"K{sub}"].value == res["pricing"]["totals"]["raw"]
    assert ws[f"J{r}"].number_format == "#,##0.00" and ws[f"K{r}"].number_format == "#,##0"

    # Ürün detayı: ürün × malzeme, SUM önbelleği
    wd = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_DETAIL]
    rd = next(x for x in range(2, wd.max_row + 1) if wd.cell(x, 1).value == "GLİSERİN")
    assert wd.cell(rd, 5).value == pytest.approx(gl["need_production"])


def test_xlsx_gross_mode_g_equals_need():
    res = _report(mode="gross")
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))[SHEET_LIST]
    r = _row_of(ws, "GLİSERİN")
    assert ws[f"G{r}"].value == f"=D{r}"
    assert ws["E1"].value == "Elimizde (bilgi)"


def test_xlsx_free_text_never_becomes_formula():
    data = build_workbook(_report())
    allowed = re.compile(r'^(MAX\(0,ROUND\(D\d+-E\d+-F\d+,3\)\)|D\d+|IF\(J\d+="","",ROUND\([GL]\d+\*J\d+,0\)\)'
                         r'|SUM\([A-Z]+\d+:[A-Z]+\d+\)|K\d+(\+K\d+)*'
                         r'|MIN\(D\d+,E\d+\)|MIN\(F\d+,ROUND\(D\d+-G\d+,6\)\)'    # + Stoktan karşılananlar
                         r'|H\d+|IF\(G\d+="","",[FJ]\d+\*G\d+\)|F\d+\*I\d+'       # + Ürün maliyeti
                         r'|ROUND\(SUM\(K\d+:K\d+\),0\))$')
    z = zipfile.ZipFile(io.BytesIO(data))
    for n in z.namelist():
        if n.startswith("xl/worksheets/sheet"):
            for f in re.findall(r"<f>([^<]*)</f>", z.read(n).decode()):
                assert allowed.match(f), (n, f)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[SHEET_LIST]
    r = _row_of(ws, EVIL_NAME)
    assert ws.cell(r, 3).data_type == "s" and ws.cell(r, 3).value == EVIL_NAME
    r = _row_of(ws, "=1+1 koli")
    assert ws.cell(r, 3).data_type == "s"
    assert ws.cell(r, 24).value == "Sistemde stok kaydı yok · =2 adet kırık"
    texts = [c.value for row in wb["Özet"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "=cmd|' /C calc'!A0" in texts


def _edge_report():
    """Adet kartına kg fiyatı (birim uyuşmazlığı), ambalaj katı açık, ondalıklı
    serbest satır, adında/başlıkta kontrol karakteri."""
    items = [_it(1, "GLİSERİN", stock=500), _it(2, "KAPSÜL\x01 KABUĞU", unit="adet")]
    recipe = RecipeRec(id=1, name="Krem", ingredients=(IngredientRec(item_id=1, quantity=10),
                                                       IngredientRec(item_id=2, quantity=2)))
    inp = PlanInputs(items={i.id: i for i in items}, recipes={1: recipe}, open_orders={}, stock_as_of=NOW)
    req = PlanRequest(title="Deneme\x02 Siparişi", lines=[{"recipe_id": 1, "qty": 1000}],
                      manual_lines=[{"section": "shipping", "name": "Koli bandı", "qty": 2.5, "unit": "rulo"},
                                    {"section": "shipping", "name": "Koli", "qty": 40, "unit": "adet"}],
                      options={"round_to_package": True})
    pin = PriceInputs(offers={1: [OfferRec(item_id=1, supplier_name="UMAYCHEM", unit_price=4.37, package_size=25,
                                           price_unit="kg", source="stok_son_durum")],
                              2: [OfferRec(item_id=2, supplier_name="FİRMA D", unit_price=2.0, price_unit="kg",
                                           source="stok_son_durum")]})
    return attach(compute(inp, req), pin, None, currency="USD", round_to_package=True)


def test_unit_mismatch_offer_shows_reason_not_missing_price():
    res = _edge_report()
    _, text = _pdf_text(render_pdf(res))
    why = "FİRMA D — fiyat listesinde, fiyat birimi (kg) malzemenin alım birimiyle (adet) uyuşmuyor"
    assert why in text and "FİRMA D — fiyat listesinde, fiyat yazılmamış" not in text
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))[SHEET_LIST]
    assert why in ws.cell(_row_of(ws, "KAPSÜL KABUĞU"), 24).value


def test_xlsx_package_rounding_written_and_decimal_manual_qty():
    res = _edge_report()
    gl = _mat(res, "GLİSERİN")
    assert gl["pkg_buy"] == 25 and gl["pkg_amount"] == 109
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))
    ws = wb[SHEET_LIST]
    # Not: ambalaj katı satırı + ambalaj uyarısı (9,5 kg < 25 kg'ın yarısı; seçenek açık →
    # "tam ambalaj alınırsa" eki yok, yuvarlanmış satır zaten yazıyor)
    assert ws.cell(_row_of(ws, "GLİSERİN"), 24).value == (
        "Ambalaj katına yuvarlanırsa: 25,0 kg · 109 $. Bu fiyat 25 kg'lık ambalaj için; alınacak 9,5 kg. "
        "Bu miktarda birim fiyat daha yüksek olabilir — bu miktara özel fiyat alınmalı.")
    r = _row_of(ws, "Koli bandı")
    assert ws.cell(r, 7).value == 2.5 and ws.cell(r, 7).number_format == "#,##0.0#"
    assert ws.cell(_row_of(ws, "Koli"), 7).number_format == "#,##0"
    # Özet'teki "ayrıca gösterildi" cümlesi artık doğru
    assert any("Ambalaj katına yuvarlanmış" in str(c.value) for row in wb["Özet"].iter_rows() for c in row)


def test_xlsx_control_characters_stripped_not_500():
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(_edge_report())))
    ws = wb[SHEET_LIST]
    assert _row_of(ws, "KAPSÜL KABUĞU")
    assert wb.properties.title == "Deneme Siparişi"


# ─── xlsx_cache ─────────────────────────────────────────────────────────────

def _small_book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sayfa & 1"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = 2, "=A1*2", '=IF(A1=2,"","x")', '=IF(A1=2,"a<b","")'
    wb.create_sheet("İkinci")["A1"] = "=1+1"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_cache_types_and_lookup_by_sheet_name():
    out = inject_cached_values(_small_book(), {"Sayfa & 1": {"B1": 4, "C1": "", "D1": "a<b"},
                                               "İkinci": {"A1": 2.5}})
    z = zipfile.ZipFile(io.BytesIO(out))
    x1 = z.read("xl/worksheets/sheet1.xml").decode()
    assert re.search(r'<c r="B1"[^>]*><f>A1\*2</f><v>4</v></c>', x1)
    assert re.search(r'<c r="C1"[^>]*t="str"[^>]*>.*?<v></v></c>', x1)
    assert re.search(r'<c r="D1"[^>]*t="str"[^>]*>.*?<v>a&lt;b</v></c>', x1)
    wv = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    assert wv["Sayfa & 1"]["B1"].value == 4 and wv["Sayfa & 1"]["D1"].value == "a<b"
    assert wv["İkinci"]["A1"].value == 2.5
    wf = openpyxl.load_workbook(io.BytesIO(out))
    assert wf["Sayfa & 1"]["B1"].value == "=A1*2"                     # formül korunur


def test_xlsx_cache_rejects_unknown_sheet_or_missing_formula_cell():
    raw = _small_book()
    with pytest.raises(ValueError):
        inject_cached_values(raw, {"Yok": {"A1": 1}})
    with pytest.raises(ValueError):
        inject_cached_values(raw, {"Sayfa & 1": {"A1": 1}})          # A1 formül değil
    with pytest.raises(ValueError):
        inject_cached_values(raw, {"Sayfa & 1": {"Z9": 1}})
    assert inject_cached_values(raw, {}) == raw


# ─── "Aynı malzeme" + sipariş geçmişi ───────────────────────────────────────

def _stearyl_report():
    """Prod'daki stearil alkol: 126 Tatlıdilimler (g, stok) reçetede; 599
    Yiğitoğlu ve 593 Veser (adet) aynı malzeme grubunda; 4 sipariş işareti."""
    items = [_it(126, "SETİL STEARİL ALKOL", stock=3944), _it(599, "CETYL STEARYL ALCOHOL", unit="adet"),
             _it(593, "CETEARYL ALCOHOL", unit="adet")]
    recipe = RecipeRec(id=1, name="Krem", ingredients=(IngredientRec(item_id=126, quantity=40),))
    inp = PlanInputs(items={i.id: i for i in items}, recipes={1: recipe}, stock_as_of=NOW,
                     mgroup_of={126: 7, 599: 7, 593: 7}, mgroup_names={7: "Setil stearil alkol"})
    res = compute(inp, PlanRequest(title="Stearil", lines=[{"recipe_id": 1, "qty": 1000}]))
    sup = {1: SupplierRec(id=1, name="TATLIDİLİMLER"), 2: SupplierRec(id=2, name="YİĞİTOGLU KİMYA"),
           3: SupplierRec(id=3, name="VESER KİMYEVİ")}
    orders = {126: [OrderRec(126, supplier_id=1, supplier_name="TATLIDİLİMLER", quantity=25000, unit="g",
                             ordered_at=datetime(2026, 10, 1, 7), expected_date=datetime(2026, 10, 10).date())]
              + [OrderRec(126, supplier_id=1, supplier_name="TATLIDİLİMLER", quantity=10000, unit="g",
                          ordered_at=datetime(2026, m, 1, 7), closed_at=datetime(2026, m, 3),
                          closed_reason="received") for m in (8, 6, 4)]}
    pin = PriceInputs(suppliers=sup, card_supplier={126: 1, 599: 2, 593: 3},
                      lots={126: [LotRec(item_id=126, lot_number="TD-1", supplier_id=1, quantity=3944,
                                         created_at=datetime(2026, 8, 12, 7), inventory_id=1)]},
                      receipts={126: [ReceiptRec(126, "TD-1", 25000, datetime(2026, 8, 12, 7), "purchase")]},
                      orders=orders)
    return attach(res, pin, None, currency="USD")


def test_pdf_same_material_and_order_lines_capped():
    res = _stearyl_report()
    _, text = _pdf_text(render_pdf(res))
    assert "Stok kartında yazan: TATLIDİLİMLER (son alım 12.08.2026)" in text
    assert "Sipariş: TATLIDİLİMLER 01.10.2026 · 25,0 kg · açık · beklenen 10.10.2026" in text
    assert "Sipariş: TATLIDİLİMLER 01.06.2026 · 10,0 kg · teslim alındı" in text
    assert "01.04.2026" not in text and "+1 sipariş daha" in text            # en çok 3 + "+N"
    assert "Aynı malzeme: «CETEARYL ALCOHOL» (VESER KİMYEVİ, stok yok, birimi adet)" in text
    assert "Aynı malzeme: «CETYL STEARYL ALCOHOL» (YİĞİTOGLU KİMYA, stok yok, birimi adet)" in text
    # eşdeğer kartların tedarikçileri fiyatsız kalem için aday firma
    assert "eşdeğer kart «CETYL STEARYL ALCOHOL»" in text


def test_xlsx_same_material_and_order_history_after_note():
    res = _stearyl_report()
    m = res["materials"][0]
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))
    ws = wb[SHEET_LIST]
    assert (ws.cell(1, 24).value, ws.cell(1, 25).value, ws.cell(1, 26).value) == (
        "Not", "Aynı malzeme (diğer kartlar)", "Sipariş geçmişi")
    r = _row_of(ws, "SETİL STEARİL ALKOL")
    assert ws.cell(r, 25).value == "; ".join(alt_texts(m)) and "«CETYL STEARYL ALCOHOL»" in ws.cell(r, 25).value
    assert ws.cell(r, 26).value == "; ".join(order_texts(m))                  # Excel'de sınır yok (4 sipariş)
    assert ws.cell(r, 26).value.count("TATLIDİLİMLER") == 4
    assert ws[f"K{r}"].value == f'=IF(J{r}="","",ROUND(G{r}*J{r},0))'         # formüller kaymadı
    sp = wb[SHEET_SUP]
    assert sp.cell(1, 11).value == "İlişki türleri"
    types = {sp.cell(x, 1).value: sp.cell(x, 11).value for x in range(2, sp.max_row + 1)}
    assert types["TATLIDİLİMLER"] == "stok kartı, alım, sipariş"
    assert types["YİĞİTOGLU KİMYA"] == types["VESER KİMYEVİ"] == "eşdeğer kart"


# ─── Stoktan karşılanan değer (14.10.2026) ──────────────────────────────────

def _sheet_rows(ws):
    return {ws.cell(r, 3).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value}


def test_xlsx_stock_sheet_formulas_cache_and_total():
    """GLİSERİN: 10 kg gerek, 0,5 kg stok (kısmi, 4,52 $) · KSANTAN: yeterli,
    fiyatsız.  Yeni sayfa: önce yeterli, sonra kısmi; G = MIN(D,E), K =
    IF(J="","",ROUND(G*J,0)) + önbellek; TOPLAM = stoktan değer.  Alım
    listesi GENEL TOPLAM'ı değişmez; Özet plan toplamını gösterir."""
    res = _report()
    cs = res["pricing"]["cost_summary"]
    assert (cs["buy_total"], cs["stock_total"], cs["plan_total"]) == (83, 2, 85)
    data = build_workbook(res)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws, sv = wb[SHEET_STOCK], wv[SHEET_STOCK]
    assert [ws.cell(1, c).value for c in range(1, 14)] == [
        "Durum", "Tür", "Malzeme", "Gereken", "Elimizde", "Yolda", "Stoktan kullanılan", "Birim",
        "Seçilen tedarikçi", "Birim fiyat ($)", "Değer ($)", "Fiyat kaynağı", "Not"]
    rows = _sheet_rows(ws)
    assert list(rows)[:2] == ["KSANTAN GAM", "GLİSERİN"]                      # önce yeterli, sonra kısmi
    k, g = rows["KSANTAN GAM"], rows["GLİSERİN"]
    assert ws.cell(k, 1).value == "Yeterli" and ws.cell(g, 1).value == "Kısmen stoktan"
    assert ws[f"G{g}"].value == f"=MIN(D{g},E{g})" and ws[f"K{g}"].value == f'=IF(J{g}="","",ROUND(G{g}*J{g},0))'
    assert sv[f"G{g}"].value == pytest.approx(0.5) and sv[f"K{g}"].value == 2 == _mat(res, "GLİSERİN")["stock_value"]
    assert sv[f"G{k}"].value == pytest.approx(3.0) and sv[f"K{k}"].value in (None, "")    # fiyatsız
    assert "Fiyat yok" in ws.cell(k, 13).value
    tot = rows["TOPLAM (fiyatı belli olanlar)"]
    assert ws[f"K{tot}"].value == f"=SUM(K2:K{tot - 1})" and sv[f"K{tot}"].value == cs["stock_total"] == 2
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet2.xml").decode()
    assert not re.search(r"<f>[^<]*</f><v ?/>", xml)                          # önbelleksiz formül yok
    # Alım listesi GENEL TOPLAM'ı birebir aynı (stok değeri girmez)
    lv = wv[SHEET_LIST]
    gt = next(x for x in range(2, lv.max_row + 1) if lv.cell(x, 3).value == "GENEL TOPLAM (fiyatı belli olanlar)")
    assert lv[f"K{gt}"].value == res["pricing"]["totals"]["all"] == 83
    # Özet: GENEL TOPLAM'ın altında stoktan değer + PLANIN TOPLAM MALZEME MALİYETİ + sayılar
    oz = {r[0].value: r[1].value for r in wv["Özet"].iter_rows(min_row=2) if r[0].value}
    labels = [r[0].value for r in wv["Özet"].iter_rows(min_row=2)]
    i = labels.index("GENEL TOPLAM — fiyatı belli ($)")
    assert labels[i + 1:i + 4] == ["Stoktan karşılananların değeri — fiyatı belli ($)",
                                   "PLANIN TOPLAM MALZEME MALİYETİ ($)", "Fiyatlı / fiyatsız stoktan karşılanan"]
    assert oz["PLANIN TOPLAM MALZEME MALİYETİ ($)"] == 85 and oz["Fiyatlı / fiyatsız stoktan karşılanan"] == "1 / 1"
    assert any("yenileme değeri" in str(r[1].value) for r in wv["Özet"].iter_rows(min_row=2))


def test_xlsx_stock_sheet_open_order_columns():
    res = _report(open_orders=True)
    cs = res["pricing"]["cost_summary"]
    assert (cs["buy_total"], cs["stock_total"], cs["open_total"], cs["plan_total"]) == (74, 2, 9, 85)
    data = build_workbook(res)
    ws = openpyxl.load_workbook(io.BytesIO(data))[SHEET_STOCK]
    sv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_STOCK]
    assert (ws["L1"].value, ws["M1"].value, ws["N1"].value) == ("Yoldan kullanılan", "Yoldan değer ($)",
                                                                "Fiyat kaynağı")
    rows = _sheet_rows(ws)
    g = rows["GLİSERİN"]
    assert ws.cell(g, 1).value == "Kısmen stoktan + yoldan"
    assert ws[f"L{g}"].value == f"=MIN(F{g},ROUND(D{g}-G{g},6))"
    assert ws[f"M{g}"].value == f'=IF(J{g}="","",ROUND(L{g}*J{g},0))'
    assert sv[f"L{g}"].value == pytest.approx(2.0) and sv[f"M{g}"].value == 9
    tot = rows["TOPLAM (fiyatı belli olanlar)"]
    assert sv[f"K{tot}"].value == 2 and sv[f"M{tot}"].value == cs["open_total"] == 9
    oz = {r[0].value: r[1].value for r in openpyxl.load_workbook(io.BytesIO(data), data_only=True)["Özet"]
          .iter_rows(min_row=2) if r[0].value}
    assert oz["Yoldan karşılananların değeri (açık sipariş) — fiyatı belli ($)"] == 9
    assert oz["PLANIN TOPLAM MALZEME MALİYETİ ($)"] == 85


def test_pdf_stock_value_lines_and_gross_mode_has_none():
    _, text = _pdf_text(render_pdf(_report()))
    assert "Değer" in text and "stoktan 2 $" in text                          # Tutar altı + Yeterli sütunu
    assert "Planın toplam malzeme maliyeti: 85 $" in text
    assert "Stoktan karşılananların değeri: 2 $" in text
    assert ("Stoktan kısmen ya da tamamen karşılanan 2 kalemden 1 tanesinin fiyatı yok; bu stoğun değeri "
            "toplama girmedi.") in text
    # "Alınacak kalem" = kart + yeni kalem (KPI ve özet kutusuyla aynı tanım)
    assert ("Fiyatı olmayan 5 alınacak kalem (2 kart, 3 yeni etiket / koli / palet kalemi) planın toplam "
            "maliyetine de girmedi.") in text
    assert "Yeterli olanlar (1 kalem — fiyatı yok)" in text
    assert "yenileme değeri" in text
    _, text = _pdf_text(render_pdf(_report(open_orders=True)))
    assert "yoldan 9 $" in text and "Yoldan karşılananların değeri (açık sipariş): 9 $" in text
    # Brüt mod: hiçbiri
    g = _report(mode="gross")
    _, text = _pdf_text(render_pdf(g))
    assert "Değer" not in text and "Planın toplam malzeme maliyeti" not in text
    assert not re.search(r"stoktan [\d.]+ \$", text) and "Stoktan karşılanan" not in text
    assert "yenileme değeri" not in text
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(g)))
    assert SHEET_STOCK not in wb.sheetnames
    assert not any("PLANIN" in str(c.value) for row in wb["Özet"].iter_rows() for c in row)


def test_xlsx_stock_sheet_unsafe_rounding_writes_number_not_min():
    """Düz yuvarlamada adet satırının stoktan kullanılanı gereken − alınacak
    (10,6 − 8 = 2,6); MIN(D,E) = 3 rapordan ayrışırdı → G sayı yazılır (Alım
    listesi G'sinin koruması gibi), K formülü kalır, TOPLAM = cost_summary."""
    items = [_it(1, "100 ML ŞİŞE", cat="Ambalaj", unit="adet", stock=3, pkg="şişe")]
    recipe = RecipeRec(id=1, name="Tonik", ingredients=(IngredientRec(item_id=1, quantity=1),))
    inp = PlanInputs(items={1: items[0]}, recipes={1: recipe}, stock_as_of=NOW)
    res = compute(inp, PlanRequest(lines=[{"recipe_id": 1, "qty": 10}], options={
        "safe_rounding": False, "extra_waste_pct": {"raw": 0, "packaging": 6, "label": 0}}))
    attach(res, PriceInputs(offers={1: [OfferRec(item_id=1, supplier_name="CAM", unit_price=10.0,
                                                 price_unit="adet", source="stok_son_durum")]}), None)
    m = _mat(res, "100 ML ŞİŞE")
    assert (m["stock_qty_used"], m["stock_value"]) == (pytest.approx(2.6), 26)
    data = build_workbook(res)
    ws = openpyxl.load_workbook(io.BytesIO(data))[SHEET_STOCK]
    sv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_STOCK]
    r = _sheet_rows(ws)["100 ML ŞİŞE"]
    assert ws[f"G{r}"].value == pytest.approx(2.6)                              # sayı, MIN(D,E) değil
    assert ws[f"K{r}"].value == f'=IF(J{r}="","",ROUND(G{r}*J{r},0))' and sv[f"K{r}"].value == 26
    tot = _sheet_rows(ws)["TOPLAM (fiyatı belli olanlar)"]
    assert sv[f"K{tot}"].value == res["pricing"]["cost_summary"]["stock_total"] == 26


def test_cover_text_same_in_preview_and_pdf():
    """"stoktan X · yoldan Y" tek yardımcıdan (`cover_text`): önizlemenin
    `ui.cover_text`'i PDF'teki metinle birebir aynı."""
    from core.purchase_pricing import cover_text
    from routers.purchase_plan import _preview_view
    res = _report(open_orders=True)
    assert cover_text(_mat(res, "GLİSERİN"), "USD") == "stoktan 2 $ · yoldan 9 $"
    assert cover_text(_mat(res, "SHEA & BUTTER <özel>"), "USD") is None
    _, text = _pdf_text(render_pdf(res))
    assert "stoktan 2 $ · yoldan 9 $" in text and "stoktan: " not in text
    view = _preview_view(_report(open_orders=True))
    ui = _mat(view, "GLİSERİN")["ui"]
    assert ui["cover_text"] == "stoktan 2 $ · yoldan 9 $"
    assert (ui["stock_value"], ui["open_value"]) == ("2 $", "9 $")


# ─── Rusya kabul senaryosu (ortam değişkeniyle kapılı) ──────────────────────

@pytest.mark.skipif(not os.environ.get("RUSYA_FIXTURE_DIR"), reason="RUSYA_FIXTURE_DIR tanımlı değil (kabul testi)")
def test_rusya_scenario_renders_with_matching_totals():
    spec = importlib.util.spec_from_file_location(
        "_rusya_acceptance", Path(__file__).with_name("test_purchase_plan_acceptance.py"))
    acc = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(acc)
    res, _, _, has_xls = acc.run_plan(label_faces=acc._label_layout())
    if not has_xls:
        pytest.skip("STOK SON DURUM.xlsx yok")
    total = res["pricing"]["totals"]["all"]
    rd, text = _pdf_text(render_pdf(res))
    # Yeniden tasarım hedefi: detay ≤ 40 sayfa (eski 67), kapak dikey, operasyon yatay
    assert len(rd.pages) <= 40 and _portrait(rd.pages[0]) and not _portrait(rd.pages[-1])
    assert f"Alınacakların tutarı (fiyatı belli): {money(total)}" in text
    assert "Fiyatı belli kalemlerin tutarı" not in text           # aynı sayı iki adla yazılmaz
    assert "Rusça etiket, toplam 119.520 adet" in text
    # 17.10.2026 görsel QA: yönetici bölümü 3 dikey sayfa (özet 3; ≥ %40 boş ara sayfa yok),
    # "7. Tedarikçi bilgi eksikleri" bölüm 6'nın tablosuyla AYNI sayfada (zorunlu sayfa sonu yok)
    assert sum(_portrait(p) for p in rd.pages) <= 3
    ops0 = sum(_portrait(p) for p in rd.pages)                 # içindekiler dikey sayfada: atla
    pages = [" ".join((p.extract_text() or "").split()) for p in rd.pages]
    p7 = next(i for i, t in enumerate(pages) if i >= ops0 and "7. Tedarikçi bilgi eksikleri" in t)
    assert "Teklif istenecek kalemler" in pages[p7]
    rs, stext = _pdf_text(render_pdf(res, variant="summary"))
    assert len(rs.pages) <= 3 and all(_portrait(p) for p in rs.pages)
    assert f"Alınacakların tutarı (fiyatı belli): {money(total)}" in stext and "Hammadde ihtiyacı" in stext
    # Kapakta tek "alınacak" tanımı: tür tablosunun "Alınacak" sütunu = kutudaki sayı (kart + yeni kalem)
    sd = summary_data(res)
    n_buy = sd["unpriced_split"]["to_buy"] + sd["unpriced_split"]["new"]
    assert sum(g["to_buy"] for g in sd["unpriced_groups"]) == n_buy
    assert f"{n_buy} alınacak kalemin fiyatı yok" in stext and "stok değerine girmedi" not in stext
    # Detaylı Rapor'un kapağı da AYNI tanım (eski madde listesi 36 / 17 / "Rusça etiket …" sayısız yazıyordu)
    assert f"{n_buy} alınacak kalemin fiyatı yok" in text and "Ambalajın (" not in text
    assert "bu ikisinin toplamı" not in text and "bu ikisinin toplamı" not in stext
    # Ek B tam liste (≤ COST_UNPRICED_MAX): "tam liste Ek B'de" ile "+N kalem daha" çelişmez
    assert "kalem daha (Excel" not in text
    assert f"Fiyatı eksik kalemler ({sd['unpriced_count']}): tam liste Ek B'de." in text
    wx = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))
    assert wx.sheetnames[0] == SHEET_EXEC
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(res)), data_only=True)[SHEET_LIST]
    tot = next(x for x in range(2, wb.max_row + 1) if wb.cell(x, 3).value == "GENEL TOPLAM (fiyatı belli olanlar)")
    assert wb[f"K{tot}"].value == total


# ─── "Aynı malzeme" grup stoğu (13.10.2026) ─────────────────────────────────

def _xhantan_report(qty=100):
    """Grup #5: 105 HAMMADDESEPETİ 180,31 g (reçetede) + 636 YİĞİTOGLU KİMYA
    24.687,81 g — grup stoğu sayılır, döküm stok hücresinin altında."""
    items = [_it(105, "XHANTAN GUM", stock=180.31), _it(636, "XANTHAN GUM", stock=24687.81)]
    recipe = RecipeRec(id=1, name="Jel", ingredients=(IngredientRec(item_id=105, quantity=2),))
    inp = PlanInputs(items={i.id: i for i in items}, recipes={1: recipe}, stock_as_of=NOW,
                     mgroup_of={105: 5, 636: 5}, mgroup_names={5: "XHANTAN GUM"}, card_supplier={105: 1, 636: 2},
                     supplier_names={1: "HAMMADDESEPETİ", 2: "YİĞİTOGLU KİMYA"})
    res = compute(inp, PlanRequest(title="Grup", lines=[{"recipe_id": 1, "qty": qty}]))
    sup = {1: SupplierRec(id=1, name="HAMMADDESEPETİ"), 2: SupplierRec(id=2, name="YİĞİTOGLU KİMYA")}
    return attach(res, PriceInputs(suppliers=sup, card_supplier={105: 1, 636: 2}), None, currency="USD")


def test_pdf_group_stock_breakdown_under_stock_cell():
    res = _xhantan_report()
    assert res["materials"][0]["status"] == "yeterli"
    _, text = _pdf_text(render_pdf(res))
    assert "24.868 gram" in text and "YİĞİTOGLU KİMYA 24,7 kg" in text          # "Yeterli olanlar" tablosu
    assert "Aynı malzeme: «XANTHAN GUM» (YİĞİTOGLU KİMYA, stok 24,7 kg) — stoğu sayıldı" not in text  # bölümde tedarikçi hücresi yok
    big = _xhantan_report(qty=20000)
    _, text = _pdf_text(render_pdf(big))
    assert "15,2 kg" in text and "YİĞİTOGLU KİMYA 24,7 kg" in text              # alım tablosu
    assert "Aynı malzeme: «XANTHAN GUM» (YİĞİTOGLU KİMYA, stok 24,7 kg) — stoğu sayıldı" in text


def test_xlsx_group_stock_column_aa_after_existing_columns():
    res = _xhantan_report(qty=20000)
    m = res["materials"][0]
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))
    ws = wb[SHEET_LIST]
    assert [ws.cell(1, c).value for c in (24, 25, 26, 27)] == [
        "Not", "Aynı malzeme (diğer kartlar)", "Sipariş geçmişi", "Grup stoğu (kart dökümü)"]
    r = _row_of(ws, "XHANTAN GUM")
    assert ws.cell(r, 27).value == m["group_stock"]["note"] == "HAMMADDESEPETİ 180 g · YİĞİTOGLU KİMYA 24,7 kg"
    assert ws.cell(r, 25).value == "«XANTHAN GUM» (YİĞİTOGLU KİMYA, stok 24,7 kg) — stoğu sayıldı"
    assert ws.cell(r, 5).value == pytest.approx(24.8)                          # E = elimizde (grup dahil)
    assert ws[f"G{r}"].value == f"=MAX(0,ROUND(D{r}-E{r}-F{r},3))"            # formüller kaymadı
    assert ws.column_dimensions["AA"].width == 50


# ─── Ambalaj uyarısı + "Ürün maliyeti" / "Maliyet özeti" (16.10.2026) ───────

def test_xlsx_amount_warn_k_cell_orange_with_comment_total_unchanged():
    """GLİSERİN 9,5 kg ve SHEA 5 kg alınacak, fiyatlar 25 kg ambalaj için →
    K turuncu + uyarı hücre yorumu; formül ve GENEL TOPLAM aynı."""
    from core.purchase_plan_xlsx import WARN_COLOR
    from core.purchase_pricing import amount_warn
    res = _report()
    gl = _mat(res, "GLİSERİN")
    assert amount_warn(gl).startswith("Bu fiyat 25 kg'lık ambalaj için; alınacak 9,5 kg.")
    data = build_workbook(res)
    ws = openpyxl.load_workbook(io.BytesIO(data))[SHEET_LIST]
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_LIST]
    r = _row_of(ws, "GLİSERİN")
    k = ws[f"K{r}"]
    assert k.value == f'=IF(J{r}="","",ROUND(G{r}*J{r},0))' and wv[f"K{r}"].value == gl["amount"]
    assert k.font.color.rgb.endswith(WARN_COLOR) and k.comment.text == amount_warn(gl)
    assert amount_warn(gl) in ws.cell(r, 24).value                         # X "Not" da yazar
    ru = _row_of(ws, "GÜL ABSOLÜ")                                          # fiyatsız → işaret yok
    assert ws[f"K{ru}"].comment is None
    tot = next(x for x in range(2, ws.max_row + 1) if ws.cell(x, 3).value == "GENEL TOPLAM (fiyatı belli olanlar)")
    assert wv[f"K{tot}"].value == res["pricing"]["totals"]["all"] == 83


def test_xlsx_cost_sheet_formulas_cached_values_and_totals():
    from core.purchase_plan_xlsx import SHEET_COST
    from core.purchase_pricing import product_cost_lines
    res = _report()
    co = res["pricing"]["cost_overview"]
    p = co["products"][0]
    # Eksik: fiyatsız malzemeler + ürünün yeni etiketi (label_mode="new")
    assert p["total"] == 85 and p["unit_cost"] == pytest.approx(0.0852) and p["missing"] == 4
    assert p["missing_names"][-1].endswith("— Gece Kremi 50 ml")
    assert co["plan_total"] == res["pricing"]["cost_summary"]["plan_total"] == 85
    data = build_workbook(res)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    assert wb.sheetnames.index(SHEET_COST) == wb.sheetnames.index("Ürünler") + 1
    ws, sv = wb[SHEET_COST], wv[SHEET_COST]
    assert ws.freeze_panes == "C1"          # yalnız No / Ürün: iki blokta da aynı anlam
    assert [ws.cell(1, c).value for c in range(1, 8)] == ["No", "Ürün", "Adet", "1 adet maliyeti ($)", "Toplam ($)",
                                                          "Fiyatı eksik kalem", "Eksik kalemler"]
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == (1, "Gece Kremi 50 ml", 1000)
    sub = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == "1. Gece Kremi 50 ml — ara toplam")
    assert ws["D2"].value == f"=H{sub}" and ws["E2"].value == f"=K{sub}"
    assert sv["D2"].value == pytest.approx(p["unit_cost"]) and sv["E2"].value == 85
    assert ws["B3"].value == "TOPLAM" and ws["E3"].value == "=SUM(E2:E2)" and sv["E3"].value == co["products_total"]
    assert ws["G2"].value == ", ".join(p["missing_names"])
    # Döküm: H = IF(G="","",F*G), J = F*I, K = IF(G="","",J*G) — önbellek product_cost_lines ile aynı
    lines = product_cost_lines(res)
    rows = [r for r in range(2, sub) if ws.cell(r, 3).value in {x["name"] for x in lines} and ws.cell(r, 9).value]
    assert len(rows) == len(lines)
    for r, ln in zip(rows, lines):
        assert ws[f"H{r}"].value == f'=IF(G{r}="","",F{r}*G{r})' and ws[f"J{r}"].value == f"=F{r}*I{r}"
        assert ws[f"K{r}"].value == f'=IF(G{r}="","",J{r}*G{r})'
        assert ws[f"F{r}"].value == pytest.approx(ln["per_unit"]) and ws[f"I{r}"].value == 1000
        assert sv[f"J{r}"].value == pytest.approx(ln["total_qty"])
        if ln["total"] is None:                                                  # fiyatsız → boş sonuç
            assert sv[f"K{r}"].value in (None, "") and sv[f"H{r}"].value in (None, "")
        else:
            assert sv[f"K{r}"].value == pytest.approx(ln["total"])
    gl = next(r for r, ln in zip(rows, lines) if ln["name"] == "GLİSERİN")
    assert sv[f"H{gl}"].value == pytest.approx(0.0452) and sv[f"K{gl}"].value == pytest.approx(45.2)
    assert ws[f"H{sub}"].value == f"=SUM(H{rows[0]}:H{rows[-1]})"
    assert ws[f"K{sub}"].value == f"=ROUND(SUM(K{rows[0]}:K{rows[-1]}),0)" and sv[f"K{sub}"].value == 85
    # Tür bazında maliyet: TOPLAM = planın toplam malzeme maliyeti; eksik kalemler tam liste
    tr = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == "TOPLAM (planın toplam malzeme maliyeti)")
    assert ws.cell(tr, 7).value == 85 and ws.cell(tr, 9).value == co["total"]["unpriced"]
    texts = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
    assert f"Fiyatı eksik kalemler ({len(co['unpriced'])})" in texts and "=1+1 koli" in texts
    assert ws.cell(texts.index("=1+1 koli") + 1, 2).data_type == "s"           # serbest metin formül olmaz
    for n in zipfile.ZipFile(io.BytesIO(data)).namelist():                       # önbelleksiz formül yok
        if n.startswith("xl/worksheets/sheet"):
            assert not re.search(r"<f>[^<]*</f><v ?/>", zipfile.ZipFile(io.BytesIO(data)).read(n).decode()), n
    # Özet: yalnız yeni satırlar (eski sıra aynı), Alım listesi GENEL TOPLAM aynı
    oz = {r[0].value: r[1].value for r in wv["Özet"].iter_rows(min_row=2) if r[0].value}
    assert oz["Ürün başına maliyet toplamı ($)"] == 85 and oz["Ambalaj uyarılı satır"] == 2
    assert oz["Fiyatı eksik kalem (maliyet özeti)"] == co["total"]["unpriced"]
    labels = [r[0].value for r in wv["Özet"].iter_rows(min_row=2)]
    assert labels.index("Hariç kalem") < labels.index("Ürün başına maliyet toplamı ($)")


def test_pdf_cost_section_and_amount_warn_mark():
    res = _report()
    _, text = _pdf_text(render_pdf(res))
    assert "Maliyet özeti (1 ürün, 1.000 adet — malzeme maliyeti 85 $)" in text
    assert "Ürün başına maliyet" in text and "Gereken miktarlar" in text and "Fiyatı eksik kalemler (6)" in text
    # Ürünler + ürün başına maliyet TEK tabloda (16.10.2026): … 1 adet maliyeti, Toplam, Fiyatı eksik
    assert "Gece Kremi 50 ml" in text and "0,09 $ 85 $ eksik: 4 kalem" in text     # 1 adet 0,0852 → 0,09
    assert "Bu fiyat 25 kg'lık ambalaj için; alınacak 9,5 kg." in text      # ⚠ dipnotu (bölüm sonunda)
    secs = [s["key"] for s in sections(res)]
    no = secs.index("cost") + 1
    assert f"{no}. Maliyet özeti" in text and f"{no - 1}. Ürünler, hedef adetler ve kapasite" in text
    # Brüt modda da bölüm var, stok değeri dili yok (bölüm başlığından "Notlar" alt başlığına kadar)
    gres = _report(mode="gross")
    gno = [s["key"] for s in sections(gres)].index("cost") + 1
    _, gtext = _pdf_text(render_pdf(gres))
    assert f"{gno}. Maliyet özeti (" in gtext
    assert "Stoktan" not in gtext.split(f"{gno}. Maliyet özeti (", 1)[1].split("Notlar")[0]


# ─── Yeniden tasarım (16.10.2026): summary_data, Yönetici Özeti, Detaylı Rapor ─

def _flat_outline(rd):
    out = []

    def walk(items):
        for it in items:
            if isinstance(it, list):
                walk(it)
            else:
                out.append(it.title)
    walk(rd.outline)
    return out


def test_summary_data_matches_pricing_totals_and_cost_summary():
    """Özet verisi yeniden hesaplamaz: KPI'lar `pricing.totals` / `cost_summary`,
    kategoriler `cost_overview`, tedarikçiler `suppliers.chosen`, fiyatsızlar
    `cost_overview.unpriced`; `raw_to_buy` = Σ display.buy_num (alım birimi)."""
    res = _report()
    pr = res["pricing"]
    sd = summary_data(res)
    k = {x["key"]: x for x in sd["kpis"]}
    assert k["buy_total"]["value"] == pr["totals"]["all"] == 83
    assert k["plan_total"]["value"] == pr["cost_summary"]["plan_total"] == 85
    assert k["stock_total"]["value"] == pr["cost_summary"]["stock_total"]
    assert k["to_buy"]["value"] == 7 and k["to_buy"]["note"] == "4 kart + 3 yeni kalem"
    assert k["products"]["text"] == "1.000 adet" and len(sd["kpis"]) == 6
    co = pr["cost_overview"]
    assert [c["kind"] for c in sd["categories"]] == ["raw", "packaging", "label"]
    assert sum(c["total"] for c in sd["categories"]) == co["total"]["total"] == co["plan_total"]
    assert sd["categories_total"]["buy"] == pr["totals"]["all"]
    assert sd["unpriced_count"] == len(co["unpriced"]) == k["unpriced"]["value"] == 6
    assert sum(g["count"] for g in sd["unpriced_groups"]) == 6
    # "Alınacak" = alınacak kart + yeni kalem (iki serbest koli/streç satırı, yeni etiket) — toplamı
    # özet kutusundaki "N alınacak kalemin fiyatı yok" (= split.to_buy + split.new)
    assert [(g["kind"], g["count"], g["to_buy"]) for g in sd["unpriced_groups"]] == [("raw", 3, 2), ("packaging", 2, 2),
                                                                                   ("label", 1, 1)]
    sp = sd["unpriced_split"]
    assert sum(g["to_buy"] for g in sd["unpriced_groups"]) == sp["to_buy"] + sp["new"] == 5
    assert sd["unpriced_groups"][0]["names"][0] == "GÜL ABSOLÜ"           # önce alınacaklar
    assert [b["name"] for b in sd["top_suppliers"]] == [b["name"] for b in res["suppliers"]["chosen"]]
    assert sum(b["total"] for b in sd["top_suppliers"]) == pr["totals"]["all"]
    exp = {"kg": 0.0, "l": 0.0, "adet": 0.0}
    for m in res["materials"]:
        d = m["display"]
        if m["kind"] == "raw" and m["status"] == "to_buy":
            exp[d["buy_unit"] if d["buy_unit"] in ("kg", "l") else "adet"] += d["buy_num"]
    assert sd["raw_to_buy"] == pytest.approx(exp) and exp["kg"] == pytest.approx(15.5) and exp["l"] == 2.0
    kg = next(b for b in sd["raw_need"] if b["unit"] == "kg")
    assert [r["name"] for r in kg["rows"]][:2] == ["GLİSERİN", "SHEA & BUTTER <özel>"]      # ihtiyaca göre
    assert kg["need_total"] == pytest.approx(19.0) and kg["buy_total"] == pytest.approx(15.5)
    assert sd["held"][0]["name"] == "ALOE VERA TOZU" and any("2 satırda" in w for w in sd["warnings"])
    # Brüt mod: plan/stok KPI'sı yok, bekletilen kutusu var
    g = summary_data(_report(mode="gross"))
    assert [x["key"] for x in g["kpis"]] == ["buy_total", "to_buy", "unpriced", "products", "held"]
    assert g["show_stock"] is False


def test_theme_hbar_chart_series_equal_table_values_and_single_label():
    """Grafik verisi verilen seriyle birebir (ters sırada — aşağıdaki sıra
    testi); yalnız en büyük çubuk etiketli; ≥2 seride lejant, tek seride yok;
    'Diğer' çubuğu gri; değer ekseni sıfırı '0'."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.charts.legends import Legend

    th = Theme()
    d = th.hbar_chart(["Hammadde", "Ambalaj", "Diğer"], [("Alınacak", [83, 0, 5]), ("Stoktan", [2, 0, 0])],
                      width_mm=170, unit="$", muted=[2])
    bc = next(x for x in d.contents if isinstance(x, HorizontalBarChart))
    # reportlab grupta 0. seriyi en alta çizer → seriler ters verilir
    assert bc.data == [[2.0, 0.0, 0.0], [83.0, 0.0, 5.0]]
    assert list(bc.categoryAxis.categoryNames) == ["Hammadde", "Ambalaj", "Diğer"]
    # Çok seride her serinin en büyüğü etiketli (grafik serisi 1 = Alınacak, 0 = Stoktan — ters sıra)
    fmt_buy, fmt_stock = bc.barLabelFormat[1], bc.barLabelFormat[0]
    assert fmt_buy(83.0) == "83 $" and fmt_buy(5.0) == "" and fmt_stock(2.0) == "2 $" and fmt_stock(0.0) == ""
    assert any(isinstance(x, Legend) for x in d.contents)
    assert bc.bars[(1, 2)].fillColor.hexval().upper().endswith("C3C7D1")       # "Diğer" gri
    assert bc.bars[(1, 0)].fillColor.hexval().upper().endswith("3B5BC4")       # 1. seri (Alınacak)
    assert bc.bars[(0, 0)].fillColor.hexval().upper().endswith("BD8A10")       # 2. seri (Stoktan)
    assert bc.valueAxis.labelTextFormat(0) == "0"
    # Lejant örneği çubuk gibi kenarsız (strokeWidth 0 PDF'te kıl çizgi demek)
    lg = next(x for x in d.contents if isinstance(x, Legend))
    assert lg.strokeColor is None
    d1 = th.hbar_chart(["a", "b"], [("Tutar", [1, 2])], width_mm=100)
    assert not any(isinstance(x, Legend) for x in d1.contents)
    b1 = next(x for x in d1.contents if isinstance(x, HorizontalBarChart))
    assert b1.barLabelFormat(2.0) == "2" and b1.barLabelFormat(1.0) == ""      # tek seride yalnız en büyük
    assert th.hbar_chart(["a"], [("Tutar", [0])], width_mm=100) is None
    # Ayrı çubuk etiketi biçimi (tabloyla aynı yazım); `fmt` verilirse eksen onu kullanır
    d2 = th.hbar_chart(["a", "b"], [("Gereken", [10, 2.5])], width_mm=100, unit="kg", fmt=lambda v: f"<{v:g}>",
                       label_fmt=lambda v: f"{v:.1f}".replace(".", ","))
    b2 = next(x for x in d2.contents if isinstance(x, HorizontalBarChart))
    assert b2.barLabelFormat(10.0) == "10,0 kg" and b2.valueAxis.labelTextFormat(2.5) == "<2.5>"


def _axis_labels(d):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart

    bc = next(x for x in d.contents if isinstance(x, HorizontalBarChart))
    bc.draw()                                         # eksen işaretleri (_tickValues) hesaplanır
    va = bc.valueAxis
    return [va.labelTextFormat(t) for t in va._tickValues]


def test_theme_hbar_chart_axis_places_from_tick_step():
    """Eksen ondalığı ADIMDAN: 0,05 adımlı eksen "0 0,1 0,1 0,2 0,2" diye
    tekrar etmez, 0,25 "0,2" diye yanlış basılmaz; aynı eksende aynı hane,
    sıfır "0"; büyük tutar binlik ayraçlı tam sayı."""
    from core.report_theme import axis_places

    th = Theme()
    assert axis_places([0, 50, 100]) == 0 and axis_places([0, 0.5, 1.0]) == 1
    assert axis_places([0, 0.05, 0.1, 0.15000000000000002]) == 2
    for vals in ([0.2], [0.35, 0.1], [1.2, 0.4], [2, 1], [38455, 2077]):
        labels = _axis_labels(th.hbar_chart([f"k{i}" for i in range(len(vals))], [("t", vals)], width_mm=80))
        assert labels[0] == "0" and len(set(labels)) == len(labels), (vals, labels)
        assert len({len(x.split(",")[1]) if "," in x else 0 for x in labels[1:]}) == 1, (vals, labels)
    assert _axis_labels(th.hbar_chart(["a"], [("t", [38455])], width_mm=80))[1].endswith(".000")


def test_theme_hbar_chart_bar_order_matches_legend_and_labels_fit_gutter():
    """Çok seride lejantın İLK girdisi her grubun EN ÜST çubuğu (aynı renk);
    kategori etiketleri sol boşluğa (`bc.x`) sığar — uzun adlar komşu sütuna /
    kenar boşluğuna taşmaz (gerekirse kısaltılır)."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.pdfbase.pdfmetrics import stringWidth

    th = Theme()
    for k in (2, 3):
        names = ["Alınacak", "Stoktan", "Yoldan"][:k]
        series = [(n, [10 + i, 4 + i, 1 + i]) for i, n in enumerate(names)]
        d = th.hbar_chart(["Hammadde", "Ambalaj", "Etiket"], series, width_mm=170, unit="$")
        bc = next(x for x in d.contents if isinstance(x, HorizontalBarChart))
        lg = next(x for x in d.contents if isinstance(x, Legend))
        bc.draw()                                   # _barPositions hesaplanır
        assert [n for _, n in lg.colorNamePairs] == names
        for g in range(3):
            ys = [bc._barPositions[i][g][1] for i in range(k)]          # grafik serisi i'nin y'si
            top_first = sorted(range(k), key=lambda i: -ys[i])          # yukarıdan aşağı
            colors_top_down = [bc.bars[i].fillColor.hexval() for i in top_first]
            assert colors_top_down == [col.hexval() for col, _ in lg.colorNamePairs]
    long = ["COCO CAPRYLATE CAPRATE", "MELEZ LAVANTA UÇUCU YAĞI", "KAYISI ÇEKİRDEĞİ YAĞI",
            "ÇOK ÇOK UZUN BİR HAMMADDE ADI — ÜÇÜNCÜ TEDARİKÇİ KARTI (BİLİNMEYEN)"]
    for width, lw, mc in ((84, 30, 22), (174, 52, 34), (60, 20, 40)):
        d = th.hbar_chart(long, [("Toplam gereken", [5, 4, 3, 2])], width_mm=width, unit="kg",
                          label_w_mm=lw, max_chars=mc)
        bc = next(x for x in d.contents if isinstance(x, HorizontalBarChart))
        ca = bc.categoryAxis
        room = bc.x + ca.labels.dx                       # dx < 0: etiket bc.x'in solunda biter
        for t in ca.categoryNames:
            assert stringWidth(t, ca.labels.fontName, ca.labels.fontSize) <= room, (width, t)
        assert bc.x <= 0.45 * width * 72 / 25.4 + 0.01


def test_summary_pdf_portrait_le_4_pages_charts_and_no_ops_sections():
    res = _report()
    rd, text = _pdf_text(render_pdf(res, variant="summary"))
    assert 1 <= len(rd.pages) <= 4 and all(_portrait(p) for p in rd.pages)
    assert re.search(rf"Sayfa 1 / {len(rd.pages)}\b", text)
    assert "Yönetici Özeti" in text and "Planın toplam malzeme maliyeti 85 $" in text
    assert "Maliyet — tür bazında" in text and "Hammadde 83 $ 2 $ 85 $ 2 3" in text
    assert "Alınacaklar — tedarikçi bazında (2 firma, toplam 83 $)" in text
    assert "Hammadde ihtiyacı" in text and "Alınacak hammadde: 15,5 kg · 2,0 l" in text
    assert "Kilogram — 4 kalem gereken 19,0 kg · alınacak 15,5 kg" in text and "Litre — 1 kalem" in text
    assert "Fiyatı eksik kalemler (6)" in text and "Nasıl tamamlanır: Malzeme Fiyatları sayfası" in text
    assert "Uyarılar ve bekletilenler" in text and "ALOE VERA TOZU (1,0 kg) bekletiliyor: Lab teyidi bekleniyor" in text
    # Grafik etiketleri (en büyük çubuk) tablo hücreleriyle aynı yazım: tür 83 $, tedarikçi 43 $,
    # hammadde "10,0 kg" (tabloda da "10,0 kg"; eksen "0  2  4 …")
    assert text.count("83 $") >= 3 and text.count("43 $") >= 2 and text.count("10,0 kg") >= 2
    # Operasyon bölümleri, ekler ve içindekiler yok
    for t in ("1. Hammadde — fiyat listesinde olanlar", "Alındı", "Ek A", "Ek B", "İçindekiler", "Yeterli olanlar"):
        assert t not in text, t
    with pytest.raises(ValueError):
        render_pdf(res, variant="kısa")


def test_full_pdf_cover_portrait_ops_landscape_toc_bookmarks_and_appendices():
    res = _report()
    rd, text = _pdf_text(render_pdf(res, variant="full"))
    por = [_portrait(p) for p in rd.pages]
    assert por[0] and not por[-1] and por == sorted(por, reverse=True)          # dikey → yatay, geri dönmez
    n = len(rd.pages)
    assert f"Sayfa {n} / {n}" in text and "…" not in text.split("Sayfa 1 /", 1)[1][:6]
    assert "İçindekiler" in text and "Yönetici özeti 1" in text
    titles = _flat_outline(rd)
    assert titles[0] == "Yönetici özeti" and "1. Hammadde — fiyat listesinde olanlar" in titles
    assert "Ek A — Teklif ve ilişki detayı" in titles and "Ek B — Fiyatı eksik kalemler (6)" in titles
    # Kısaltma: satırda yalnız seçilen teklif; diğeri "+1 teklif" + Ek A'da
    assert "+1 teklif · ayrıntı Ek A" in text
    assert "SURYA KİMYA — 5,00 $/kg · 25 kg'lık ambalaj (daha pahalı)" in text.split("Ek A — Teklif ve ilişki detayı Alım", 1)[1]
    # Fiyatsız liste yalnız Ek B'de (yeni kalem işaretleri bir kez), ürün + maliyet tek tabloda
    assert text.count("(yeni kalem, kartı yok)") == 3
    assert "1 adet maliyeti Toplam Fiyatı eksik" in text
    # Yeterli satırı tek satır: ad, tür, gereken, elimizde, değer, dayanağı
    assert "KSANTAN GAM Hammadde 3.000 gram 1.000.000 gram — Fiyat yok" in text
    # Notlar gruplu
    assert "Kaynaklar ve kur" in text and "Hesap kuralları" in text and "Kapsam dışı kalemler" in text


def _story_of(monkeypatch, res, variant):
    """render_pdf'in derlediği akış listesi (multiBuild yakalanır)."""
    from core import report_theme

    seen = {}
    orig = report_theme.ReportDoc.multiBuild

    def spy(self, story, *a, **kw):
        seen["story"] = list(story)
        return orig(self, story, *a, **kw)
    monkeypatch.setattr(report_theme.ReportDoc, "multiBuild", spy)
    pdf = render_pdf(res, variant=variant)
    monkeypatch.undo()
    return seen["story"], pdf


def test_pdf_no_nested_keeptogether_and_outline_only_with_bookmarks(monkeypatch):
    """İç içe KeepTogether YOK (reportlab iç bloğu 0xffffff yükseklik sayar,
    dış blok sayfa başında değilse hep yeni sayfaya atlardı — "Tedarikçi
    bilgi eksikleri" %70 boş sayfa bırakıyordu).  Yer imi paneli yalnız yer
    imi olan PDF'te: özet /UseOutlines DEĞİL, detay öyle ve ana hat dolu."""
    from reportlab.platypus import KeepTogether, Table

    def walk(flows, inside):
        for f in flows:
            if isinstance(f, KeepTogether):
                assert not inside, "iç içe KeepTogether"
                walk(f._content, True)
            elif isinstance(f, Table):
                for row in f._cellvalues:
                    for cell in row:
                        if isinstance(cell, (list, tuple)):
                            walk(cell, inside)
    res = _report()
    for variant in ("summary", "full"):
        story, pdf = _story_of(monkeypatch, res, variant)
        walk(story, False)
        rd = PdfReader(io.BytesIO(pdf))
        mode = rd.trailer["/Root"].get("/PageMode")
        if variant == "summary":
            assert mode != "/UseOutlines" and not rd.outline
        else:
            assert mode == "/UseOutlines" and len(rd.outline) >= 5
    # Kontrol listesi (bölüm 6) aday firmalarla (bölüm 5) aynı sayfada başlar
    rd, _ = _pdf_text(render_pdf(res))
    ops0 = sum(_portrait(p) for p in rd.pages)                 # içindekiler dikey sayfada: atla
    pages = [" ".join((p.extract_text() or "").split()) for p in rd.pages]
    p5 = next(i for i, t in enumerate(pages) if i >= ops0 and "5. Fiyatı olmayanlar için teklif" in t)
    p6 = next(i for i, t in enumerate(pages) if i >= ops0 and "6. Tedarikçi bilgi eksikleri" in t)
    assert p6 == p5


def test_pdf_footnotes_continuous_and_escaped_once():
    """Dipnot numarası belge boyunca kesintisiz (bölümler sayfa paylaşır; aynı
    ⚠ numarası iki farklı nota gitmez); ürün notu / uyarısı tek kez kaçışlanır
    ("&" → "&", "&amp;" değil)."""
    res = _report()
    res["products"][0]["note"] = "Müşteri A & B <acil>"
    rd, text = _pdf_text(render_pdf(res))
    assert "Not: Müşteri A & B <acil>" in text and "&amp;" not in text and "&lt;" not in text
    nums = [int(n) for n in re.findall(r"⚠ (\d+)\. ", text)]
    assert nums and nums == list(range(1, len(nums) + 1))


def test_pdf_reachability_sufficient_info_notes_missing_names_finished_stock():
    """Yeterli satırın bilgi notları "Dayanağı" hücresinde; ürün başına fiyatı
    eksik malzeme adları dipnotta; bitmiş stok (düşülmese de) Adet altında;
    malzeme bölümü alt başlığı PDF'e özel ("Gri satırlar…" yok, Ek A'ya yönlendirir)."""
    res = _report()
    ks = _mat(res, "KSANTAN GAM")
    ks["cautions"] = list(ks.get("cautions") or []) + [
        {"code": "item_note", "severity": "info", "text": "Lab: raf ömrü kontrol & yenile"},
        {"code": "samples_available", "severity": "info", "text": "Numune lotu var (50 gram)."}]
    res["products"][0]["finished_stock"] = 12.0
    rd, text = _pdf_text(render_pdf(res))
    suff = text.split("Yeterli olanlar", 1)[1]
    assert "Not: Lab: raf ömrü kontrol & yenile" in suff and "Not: Numune lotu var (50 gram)." in suff
    assert re.search(r"Fiyatı eksik: GÜL ABSOLÜ, .*KSANTAN GAM", text)
    assert "1.000 stok 12" in text and "hedeften düşülmedi" in text
    assert "Gri satırlar" not in text and "Satırda, sistemde bu malzeme için kayıtlı firma" not in text
    assert "Diğer teklifler ve sistemdeki kayıtlar Ek A'da" in text


def test_summary_cover_unpriced_counts_and_quantity_spelling():
    """Kapakta tek "fiyatı yok" tanımı: KPI = alınacak kart + yeni kalem +
    yeterli kalem, kutu "alınacak" sayısını aynı kaynaktan yazar ve yeterli
    olanları ayrıca söyler.  Miktar her yerde aynı yazım (1 ondalık,
    "litre"/"gram")."""
    from core.purchase_plan_pdf import _amt

    res = _report()
    sd = summary_data(res)
    sp = sd["unpriced_split"]
    assert sp == {"to_buy": 2, "new": 3, "stock": 1} and sum(sp.values()) == sd["unpriced_count"] == 6
    k = next(x for x in sd["kpis"] if x["key"] == "unpriced")
    assert k["note"] == "2 alınacak kart + 3 yeni kalem + 1 yeterli kalem; toplama girmedi"
    for variant in ("summary", "full"):
        _, text = _pdf_text(render_pdf(res, variant=variant))
        assert "5 alınacak kalemin fiyatı yok" in text, variant
        # Stoğu yeten fiyatsızlar ayrı cümle; KPI toplamı (6) kutuda çıplak yazılmaz — KPI notu mutabık kılar
        assert "Stoğu yeten (alınmayacak) 1 kalemin de fiyatı yok." in text and "bu ikisinin" not in text
        assert "stok değerine girmedi" not in text and "Fiyatı belli kalemlerin tutarı" not in text
        assert "Fiyatı olmayan 5 alınacak kalem (2 kart, 3 yeni" in text
        assert "hammadde fiyat listesinden" not in text           # tür maddeleri yok (tablo + KPI söyler)
    assert _amt(344.56, "kg") == "344,6 kg" and _amt(1250, "kg") == "1.250,0 kg"
    assert _amt(199.629, "l") == "199,6 litre" and _amt(0.026, "kg") == "26 gram" and _amt(3, "adet") == "3 adet"
    # Yarım yukarı (Excel "#,##0.0" gösterimi): Python biçimi 1,2 / 11,2 / 12 / 2 yazardı
    assert _amt(1.25, "kg") == "1,3 kg" and _amt(11.25, "kg") == "11,3 kg" and _amt(10.25, "l") == "10,3 litre"
    assert _amt(0.0125, "kg") == "13 gram" and _amt(2.5, "adet") == "3 adet" and _amt(0.9996, "kg") == "1,0 kg"


def test_summary_unpriced_groups_count_each_item_by_key():
    """Fiyatı eksik tür grupları kalemi TEKİL ANAHTARLA sayar: aynı adlı iki
    serbest "Koli" satırı iki kalemdir (grup toplamı = KPI = Ek B); örnek ad
    tekil kalır."""
    recipe = RecipeRec(id=1, name="Krem", ingredients=(IngredientRec(item_id=1, quantity=10),))
    inp = PlanInputs(items={1: _it(1, "GLİSERİN")}, recipes={1: recipe}, stock_as_of=NOW)
    req = PlanRequest(title="Koli", lines=[{"recipe_id": 1, "qty": 100}],
                      manual_lines=[{"section": "shipping", "name": "Koli", "qty_text": "40×30×20 cm, 20 adet"},
                                    {"section": "shipping", "name": "Koli", "qty_text": "44×30×13 cm, 15 adet"}])
    res = attach(compute(inp, req), PriceInputs(), None, currency="USD")
    sd = summary_data(res)
    assert sd["unpriced_count"] == 3
    assert [(g["kind"], g["count"], g["to_buy"], g["names"], g["more"]) for g in sd["unpriced_groups"]] == [
        ("raw", 1, 1, ["GLİSERİN"], 0), ("packaging", 2, 2, ["Koli"], 0)]
    k = next(x for x in sd["kpis"] if x["key"] == "unpriced")
    assert k["value"] == sum(g["count"] for g in sd["unpriced_groups"]) == 3
    _, text = _pdf_text(render_pdf(res, variant="summary"))
    assert "Fiyatı eksik kalemler (3)" in text and "Ambalaj 2 2 Koli" in text


def test_pdf_product_info_cautions_reachable_priced_and_unpriced():
    """Bilgi düzeyi ürün uyarısı ("Tahmini hesap: …") düşmez: fiyatlı raporda
    adın altında gri (⚠ yok, reçete adıyla tek satır grubu), fiyatsız raporda
    "Not" sütununda — eski rapor ikisinde de basıyordu."""
    import copy

    res = _report()
    p = res["products"][0]
    p["estimated"], p["scale"] = True, 2.0
    p["cautions"] = list(p.get("cautions") or []) + [
        {"code": "estimated", "severity": "info", "text": "Tahmini hesap: reçete ×2,00 ölçekle kullanıldı."}]
    _, text = _pdf_text(render_pdf(res))
    prod = text.split("Ürünler, hedef adetler")[-1].split("Maliyet özeti", 1)[0]      # [-1]: içindekiler değil
    assert "Tahmini hesap: reçete ×2,00 ölçekle kullanıldı — Gece Kremi 50 ml" in prod
    assert "Reçete: Gece Kremi" not in prod                    # ölçek tek yerde (uyarı metninde)
    un = copy.deepcopy(res)
    un["pricing"]["cost_overview"] = None                       # fiyatlanmamış rapor: "Not" sütunu
    _, text = _pdf_text(render_pdf(un))
    prod = text.split("Ürünler, hedef adetler")[-1]
    assert "Reçete: Gece Kremi 50 ml ×2,00" in prod and "Tahmini hesap: reçete ×2,00 ölçekle kullanıldı." in prod


def test_pdf_appendix_b_row_major_and_full_list_wording(monkeypatch):
    """Ek B iki sütunu SATIR sırasıyla doldurur (sayfalar arası okuma sırası
    bozulmaz); "tam liste Ek B'de" yalnız liste kırpılmadıysa — kırpılınca
    tamamının Excel'de olduğu yazılır (Ek B'nin "+N kalem daha" satırıyla aynı)."""
    from core import purchase_plan_pdf as P
    from core.purchase_pricing import cost_overview_view

    res = _report()
    order = [u["name"] for u in cost_overview_view(res)["unpriced"]]
    assert len(order) == 6
    _, text = _pdf_text(render_pdf(res))
    ekb = text.split("Ek B — Fiyatı eksik kalemler (6)")[-1]
    pos = [ekb.find(n) for n in order]
    assert all(x >= 0 for x in pos) and pos == sorted(pos), (order, pos)
    assert "Tam liste Ek B'de." in text and "Fiyatı eksik kalemler (6): tam liste Ek B'de." in text
    assert "kalem daha (Excel" not in text
    monkeypatch.setattr(P, "COST_UNPRICED_MAX", 4)
    _, text = _pdf_text(render_pdf(res))
    assert "tam liste Ek B'de" not in text.lower()
    assert "İlk 4 kalem Ek B'de, tam liste Excel “Ürün maliyeti” sayfasında." in text
    assert "Fiyatı eksik kalemler (6): ilk 4 kalem Ek B'de, tam liste Excel “Ürün maliyeti” sayfasında." in text
    assert "+2 kalem daha (Excel “Ürün maliyeti” sayfasında tam liste)." in text


def test_pdf_charts_captioned_small_blocks_in_gram_and_numbers_bound_to_units():
    """Tutar grafikleri kendi gri başlığıyla (sayfa başına taşınsa da birimi
    görünür); 1 kg altındaki hammadde bloğu gram ölçeğinde, çubuk etiketi tablo
    hücresiyle aynı ("200 gram"); sayı ile birim bölünmez boşlukla bağlı; yan
    yana hammadde tablosunda ad tek satıra kısaltılır (iki taraf satır
    yüksekliğini paylaşır)."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    from core.purchase_plan_pdf import _Ctx, _nb, _raw_rows

    res = _report()
    _, text = _pdf_text(render_pdf(res, variant="summary"))
    # Yalnız Hammadde'nin tutarı var → tür grafiği yok (tablonun tekrarı olurdu); tedarikçi grafiği var
    assert "Tutar ($) — tür bazında" not in text and "Tutar ($) — tedarikçi bazında" in text
    res["pricing"]["cost_overview"]["categories"][1].update(buy=7, total=7)
    _, text = _pdf_text(render_pdf(res, variant="summary"))
    assert "Tutar ($) — tür bazında" in text and "Tutar ($) — tedarikçi bazında" in text
    _, text = _pdf_text(render_pdf(_xhantan_report(), variant="summary"))
    assert "Toplam gereken (gram)" in text and text.count("200 gram") >= 3       # başlık, tablo, çubuk
    assert "0,0 kg" not in text and "alınacak yok" in text
    assert _nb("KERATİN (170 ml) · 6.912 adet · alınacak 6.407 adet") == \
        "KERATİN (170\u00a0ml) · 6.912\u00a0adet · alınacak 6.407\u00a0adet"
    assert _nb("95,4 litre · 3 gram · 2 kalem") == "95,4\u00a0litre · 3\u00a0gram · 2 kalem"
    c = _Ctx(_report(), "summary")
    blk = {"unit": "kg", "rows": [{"name": "COCO CAPRYLATE / CAPRATE ÇOK UZUN BİR AD", "need": 155.4, "buy": 139.5,
                                   "other": False}]}
    rows = _raw_rows(c, blk, compact=True, name_pt=80)
    nm = rows[1][0].text
    assert nm.endswith("…") and stringWidth(nm, c.S["cells"].fontName, c.S["cells"].fontSize) <= 80


def test_xlsx_exec_cost_columns_follow_mode_like_pdf():
    """Excel "Yönetici Özeti" tür tablosunun sütunları PDF'tekiyle aynı:
    Stoktan yalnız net modda, Yoldan yalnız açık siparişte (brüt modda stok
    değeri dili yok)."""
    def header(res):
        ws = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))[SHEET_EXEC]
        r = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Maliyet — tür bazında")
        return [ws.cell(r, c).value for c in range(2, 8) if ws.cell(r, c).value]
    assert header(_report(mode="gross")) == ["Alınacak ($)", "Toplam ($)", "Fiyatlı", "Fiyatsız"]
    assert header(_report()) == ["Alınacak ($)", "Stoktan ($)", "Toplam ($)", "Fiyatlı", "Fiyatsız"]
    assert header(_report(open_orders=True)) == ["Alınacak ($)", "Stoktan ($)", "Yoldan ($)", "Toplam ($)",
                                                 "Fiyatlı", "Fiyatsız"]


def test_xlsx_exec_sheet_first_header_style_filter_and_print_setup():
    res = _report()
    data = build_workbook(res)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == SHEET_EXEC
    ws = wb[SHEET_EXEC]
    pr = res["pricing"]
    val = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert ws["A1"].value == "Deneme Siparişi — Yönetici Özeti"
    assert val["Alınacakların tutarı (fiyatı belli)"] == pr["totals"]["all"] == 83
    assert val["Planın toplam malzeme maliyeti"] == pr["cost_summary"]["plan_total"] == 85
    assert val["Alınacak kalem"] == 7 and val["Fiyatı eksik kalem"] == 6
    assert val["Toplam"] == pr["cost_overview"]["total"]["buy"]               # tür bazında toplam satırı
    assert val["UMAYCHEM"] == 1 and val["Hammadde"] == 3                      # tedarikçi kalemi; fiyatsız hammadde
    r_gl = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "GLİSERİN")
    assert (ws.cell(r_gl, 2).value, ws.cell(r_gl, 3).value, ws.cell(r_gl, 4).value) == ("kg", 10.0, 9.5)
    # Malzeme satırı Alım listesinin biçimiyle (PDF satırı da Alım listesi metnini basar);
    # toplam satırı `_amt` yazımı: 1 ondalık (1'in altı 3 ondalığa kadar); litre sözcüğü
    assert ws.cell(r_gl, 3).number_format == "#,##0.0##"
    r_tot = next(r for r in range(1, ws.max_row + 1)
                 if str(ws.cell(r, 1).value or "").startswith("Alınacak hammadde — kg"))
    assert ws.cell(r_tot, 3).number_format == "[<1]0.0##;#,##0.0"
    r_gul = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "GÜL ABSOLÜ")
    assert ws.cell(r_gul, 2).value == "litre"
    # Formül kalkanı özet sayfasında da
    evil = next(ws.cell(r, 1) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == EVIL_NAME)
    assert evil.data_type == "s" and evil.quotePrefix
    # Başlık biçimi lacivert / beyaz kalın; süzgeç; sabit başlık; yazdırma ayarı
    wl = wb[SHEET_LIST]
    assert wl["A1"].fill.fgColor.rgb.endswith("232E6E") and wl["A1"].font.bold and wl["A1"].font.color.rgb.endswith("FFFFFF")
    # Süzgeç yalnız tek bloklu sayfalarda: Alım listesi ara toplam / GENEL TOPLAM taşır → süzgeç yok
    assert not wl.auto_filter.ref and wl.freeze_panes == "D2"
    if SHEET_STOCK in wb.sheetnames:
        assert not wb[SHEET_STOCK].auto_filter.ref
    assert wb[SHEET_SUP].auto_filter.ref.startswith("A1:")
    for wsx in wb.worksheets:
        ref = wsx.auto_filter.ref
        if not ref:
            continue
        last = int(re.search(r"(\d+)$", ref).group(1))
        for r in range(2, last + 1):
            vals = [str(wsx.cell(r, cc).value or "") for cc in range(1, wsx.max_column + 1)]
            assert not any("TOPLAM" in v.upper() for v in vals[:4]), (wsx.title, r)
            # sütun aralıklı SUM (=SUM(K2:K42)) ara toplamdır; satır içi SUM (D5:Y5) değil
            assert not any(a != b for v in vals for a, b in re.findall(r"SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", v)), \
                (wsx.title, r)
    assert wl.print_title_rows.replace("$", "") == "1:1" and wl.page_setup.orientation == "landscape"
    assert wl.oddHeader.center.text == "Minerva 108 — Satın Alma Planı · Deneme Siparişi"
    assert wl.sheet_properties.pageSetUpPr.fitToPage and wl.page_setup.fitToWidth == 1
    assert ws.page_setup.orientation == "portrait" and not ws.auto_filter.ref
    # Zebra: veri satırlarında dönüşümlü dolgu, ara toplam satırında yok
    r1 = _row_of(wl, "GLİSERİN")
    fills = {wl.cell(r, 3).fill.fgColor.rgb for r in (r1, r1 + 1)}
    assert fills == {"00000000", "00F5F6F9"}
    sub = next(x for x in range(2, wl.max_row + 1)
               if str(wl.cell(x, 3).value or "").startswith("1. Hammadde — fiyat listesinde olanlar — toplam"))
    assert wl.cell(sub, 3).fill.fgColor.rgb == "00000000"
    # Önbellek değerleri ad üzerinden çözüldü: Alım listesi K formülü gömülü
    wv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET_LIST]
    assert wv[f"K{r1}"].value == _mat(res, "GLİSERİN")["amount"]


# ─── Son kapı: kapak, tipografi, yetim satır (17.10.2026 doğrulayıcı bulguları) ──

def test_pdf_typographic_apostrophe_and_cover_dates():
    """İki harf arasındaki kesme işareti PDF'te tipografik (’): sabit metin
    ("Ek A’da") ve ortak veri ("25 kg’lık") — önizleme / Excel düz kalır.
    Kapak alt başlığı tek satırlık bilgi (oluşturma saati yalnız sayfa altında,
    kapakla aynı tarih yazımı)."""
    from core.purchase_pricing import PKG_SUFFIX

    res = _report()
    for variant in ("summary", "full"):
        rd = PdfReader(io.BytesIO(render_pdf(res, variant=variant)))
        raw = " ".join(" ".join((p.extract_text() or "") for p in rd.pages).split())
        assert not re.search(r"\w'\w", raw), re.findall(r".{20}\w'\w.{20}", raw)[:3]
        if variant == "full":
            assert "Ek A’da" in raw and "kg’lık" in raw and "Ek B’de" in raw
        p1 = rd.pages[0].extract_text() or ""
        assert "oluşturma 05.10" not in p1                       # sayfa altında uzun tarih
        assert re.search(r"oluşturma \d{1,2} Ekim 2026 \d{2}:\d{2}", " ".join(p1.split())) or \
            "oluşturma" not in p1
    assert "'" in PKG_SUFFIX["kg"]                                 # ortak metin (UI / Excel) düz kalır


def test_summary_kpi_open_orders_reconcile_and_ekb_marks_stock_covered():
    """Açık sipariş düşülünce plan kutusu = alınacak + stoktan + yoldan; yoldan
    kutusu yok, plan kutusunun notu söyler.  Ek B'de stoğu yeten fiyatsız kalem
    "stok yetiyor" diye ayrılır (KPI'daki "yeterli kalem" ile aynı sayı)."""
    from core.purchase_pricing import cost_overview_view

    res = _report(open_orders=True)
    sd = summary_data(res)
    k = {x["key"]: x for x in sd["kpis"]}
    cs = res["pricing"]["cost_summary"]
    assert k["plan_total"]["value"] == k["buy_total"]["value"] + k["stock_total"]["value"] + cs["open_total"]
    assert k["plan_total"]["note"] == f"yoldan karşılanan {money(cs['open_total'], 'USD')} dahil"
    assert summary_data(_report())["kpis"][0].get("note") is None          # yoldan yoksa not yok
    _, text = _pdf_text(render_pdf(res, variant="summary"))
    assert "yoldan karşılanan 9 $ dahil" in text
    res = _report()
    view = cost_overview_view(res)
    marked = [u for u in view["unpriced"] if u["need"].endswith(" · stok yetiyor")]
    assert len(marked) == summary_data(res)["unpriced_split"]["stock"] == 1
    assert not any(u["new"] for u in marked)


def test_pdf_exec_raw_rows_use_row_text_and_limiting_stock_spelling():
    """Hammadde ihtiyacı tablosu satırın Alım listesi metnini basar ("1.250
    gram" — `_amt` 1,3 kg yazıyordu); "Diğer" toplamı `_amt`.  Excel Yönetici
    Özeti satır hücresi Alım listesinin biçimiyle (1,25).  Kısıtlayan malzeme:
    stok yoksa "stok yok", sıfıra yuvarlanan "< 1 ml", parantez önü boşluk yok."""
    from core.purchase_plan_pdf import _stock_amt, _tight

    items = [_it(1, "GLİSERİN", stock=500), _it(2, "STEARİK ASİT", stock=500)]
    r = RecipeRec(id=1, name="Krem", ingredients=(IngredientRec(item_id=1, quantity=10),
                                                  IngredientRec(item_id=2, quantity=1.25)))
    res = compute(PlanInputs(items={i.id: i for i in items}, recipes={1: r}, stock_as_of=NOW),
                  PlanRequest(title="Küçük", lines=[{"recipe_id": 1, "qty": 1000, "display_name": "Krem"}],
                              options={"stock_mode": "net", "label_mode": "exclude"}))
    attach(res, PriceInputs(offers={2: [OfferRec(item_id=2, supplier_name="X KİMYA", unit_price=3.0, package_size=25,
                                                  price_unit="kg", source="stok_son_durum")]}, suppliers={}),
           None, currency="USD")
    st = _mat(res, "STEARİK ASİT")["display"]
    assert (st["need_text"], st["buy_text"]) == ("1.250 gram", "750 gram")
    _, text = _pdf_text(render_pdf(res, variant="summary"))
    assert "STEARİK ASİT 1.250 gram 750 gram" in text and "STEARİK ASİT 1,3 kg" not in text
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(res)))[SHEET_EXEC]
    row = next(x for x in range(1, ws.max_row + 1) if ws.cell(x, 1).value == "STEARİK ASİT")
    assert ws.cell(row, 3).value == 1.25 and ws.cell(row, 3).number_format == "#,##0.0##"
    assert _stock_amt(0, "adet") == _stock_amt(-2, "kg") == "stok yok"
    assert _stock_amt(0.3, "ml") == "< 1 ml" and _stock_amt(0.02, "kg") == "20 gram"
    assert _stock_amt(4400, "g") == "4,4 kg" and _stock_amt(94, "adet") == "94 adet"
    assert _tight("PAPATYA UÇUCU YAĞI(SARI )") == "PAPATYA UÇUCU YAĞI(SARI)"


def test_pdf_min_tail_table_never_leaves_short_continuation():
    """Yan yana hammadde tablosu bölünürken devam sayfasına 3 veri satırından
    azı kalmaz; ilk parçaya da kalmıyorsa tablo bütün olarak taşınır."""
    from core.purchase_plan_pdf import _MinTailTable

    data = [["Başlık", ""], ["Malzeme", "Miktar"]] + [[f"m{i}", str(i)] for i in range(14)]
    t = _MinTailTable(data, colWidths=[100, 50], rowHeights=[10] * len(data), repeatRows=2)
    t.wrap(150, 1000)
    for h in range(20, 165, 5):
        parts = t.split(150, h)
        if len(parts) == 2:
            assert len(parts[1]._cellvalues) - 2 >= 3, h
    # 15 satırlık yükseklik: düz bölme devam sayfasına 1 satır bırakırdı → kesim yukarı kayar
    parts = t.split(150, 155.5)
    assert [len(p._cellvalues) for p in parts] == [13, 5]
    # 5 veri satırı, 3'ü sığıyor: devamda 2 kalırdı, kesim kayınca ilk parçada 2 → bütün taşınır
    small = _MinTailTable(data[:7], colWidths=[100, 50], rowHeights=[10] * 7, repeatRows=2)
    small.wrap(150, 1000)
    assert small.split(150, 55.5) == []
