"""
Üretim geçmişi raporu testleri (core/production_history_report.py +
GET /api/reports/production-history).

  • Tarih filtresi TR-yerel gün: 12.05 21:30 UTC = 13.05 00:30 TR → 13'üne aittir
  • Marka filtresi (MİNERVA-108 / Minerva108 → 'Minerva 108')
  • Aynı ürünün farklı ad yazımları target_item_id ile tek satır, ad = kartın bugünkü adı
  • Ürün grubu anahtar kelimeleri — SPF'li losyon vücut losyonu DEĞİL
  • Domain izolasyonu, Excel ('Özet' + 'Aylık'), history_quantities
"""
import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from core import production_history_report as phr
from database import Item, ProductionHistory

URL = "/api/reports/production-history"


def _item(db, name, domain="cosmetics"):
    it = Item(name=name, sku=f"ph-{abs(hash((name, domain))) % 10**9}", category="Bitmiş Ürün",
              unit="adet", current_stock=0, domain=domain)
    db.add(it); db.flush()
    return it


def _ph(db, item, qty, at, *, name=None, domain="cosmetics", recipe_name=None):
    db.add(ProductionHistory(
        recipe_id=None, recipe_name=recipe_name or (name or (item.name if item else "")),
        target_item_id=item.id if item else None,
        target_item_name=name or (item.name if item else None),
        produced_quantity=qty, produced_at=at, domain=domain))


def _seed(db):
    sh = _item(db, "MİNERVA-108 ALTIN SERİ KURU SAÇ ŞAMPUANI")
    tonik = _item(db, "Serenida Tonik 100 ml")
    _ph(db, sh, 100, dt.datetime(2026, 5, 20, 9, 0))
    _ph(db, sh, 24, dt.datetime(2026, 6, 3, 9, 0))
    # Aynı ürün, eski ad yazımı — target_item_id ile birleşmeli
    _ph(db, tonik, 30, dt.datetime(2026, 5, 21, 9, 0), name="Serenida Tonik (100ml)")
    _ph(db, tonik, 20, dt.datetime(2026, 6, 10, 9, 0))
    db.commit()
    return sh, tonik


# ─── Saf: ürün grubu anahtar kelimeleri ─────────────────────────────────────

@pytest.mark.parametrize("name,cat", [
    ("MİNERVA 108 ALTIN SERİ KURU SAÇ SAMPUANI", "Şampuan"),
    ("Serenida Yağlı Saç Şampuan", "Şampuan"),
    ("Minerva 108 Hair Shampoo", "Şampuan"),
    ("MİNERVA 108 ALTIN SERİ SAC KREMİ", "Saç Kremi"),
    ("Serenida Conditioner", "Saç Kremi"),
    ("SERENİDE DUŞ JELİ", "Duş Jeli"),
    ("Minerva 108 Shower Gel", "Duş Jeli"),
    ("MİNERVA 108 ALTIN SERİ SU BAZLI SAÇ MASKESİ", "Saç Maskesi"),
    ("Minerva 108 Intense Hair Mask", "Saç Maskesi"),
    ("Serenida Besleyici Vücut Losyonu 200ml", "Vücut Losyonu"),
    ("Evanira Niacinamide Radiance Body Lotion", "Vücut Losyonu"),
    ("Minerva 108 Anti-Blemish Sun Protection Lotion SPF 50+", "Güneş Ürünleri"),
    ("MİNERVA-108 GÜNEŞ KREMİ 50 SPF-200 ML", "Güneş Ürünleri"),
    ("SPF 30 Vücut Losyonu", "Güneş Ürünleri"),           # SPF'li losyon güneş ürünüdür
    ("SERENİDA YÜZ KREMİ", "Yüz Kremi"),
    ("Minerva 108 Red Clover Night Cream", "Yüz Kremi"),
    ("MİNERVA 108 EL KREMİ", "El Kremi"),
    ("MİNERVA-108 ANTI-AGING HAND CREAM", "El Kremi"),
    ("MİNERVA ALTIN SERİ AYAKBAKIM KREMİ", "Ayak Kremi"),
    ("Minerva 108 Intensive Foot Care Cream", "Ayak Kremi"),
    ("minerva 108 altın seri göz çevresi kremi", "Göz Kremi"),
    ("Minerva 108 Red Clover Eye Contour Cream", "Göz Kremi"),
    ("SERENİDA SİVİLCE SERUM", "Serum"),
    ("Minerva 108 Face Scrub", "Peeling"),
    ("MİNERVA ALTIN SERİ TONİK", "Tonik"),
    ("Serenida Facial Toner", "Tonik"),
    ("Aloe Jel Kremi", "Krem"),                          # 'jel krem' içindeki 'el krem' sayılmaz
    # Genel krem grubu — özel krem gruplarına uymayan kremler "Diğer"e düşmez
    ("Evanira Kırmızı Yonca Leke Karşıtı Krem", "Krem"),
    ("Serenida Bikini Bölgesi Krem", "Krem"),
    ("MİNERVA 108 KIRMIZI YONCA KREMİ", "Krem"),
    ("Minerva 108 Vücut Kremi", "Krem"),
    ("Evanira Body Cream", "Krem"),
    # Araya giren tek "bakım"/"care" özel grubu bozmaz
    ("Serenida Saç Bakım Kremi", "Saç Kremi"),
    ("Saç Bakım Maskesi", "Saç Maskesi"),
    ("Serenida Hair Care Mask", "Saç Maskesi"),
    ("Minerva 108 Yüz Bakım Kremi", "Yüz Kremi"),
    ("Minerva 108 Gece Bakım Kremi", "Yüz Kremi"),
    ("Minerva 108 Hand Care Cream", "El Kremi"),
    ("Vücut Bakım Losyonu", "Vücut Losyonu"),
    ("Göz Altı Kremi", "Göz Kremi"),
    ("El Yapımı Krem", "Krem"),                          # ara kelime serbest değil
    ("Saç Bakım Yağı", "Diğer"),
    ("Krem Şampuan", "Şampuan"),
    ("Serenida Charcoal Whitening Toothpaste", "Diğer"),
    ("", "Diğer"),
])
def test_product_category_keywords(name, cat):
    assert phr.product_category(name) == cat


def test_body_lotion_rule_itself_excludes_spf(monkeypatch):
    """Sıra değişse bile SPF'li losyon vücut losyonuna düşmemeli (kuralın kendi hariç listesi)."""
    body = next(r for r in phr.CATEGORY_RULES if r[0] == "Vücut Losyonu")
    monkeypatch.setattr(phr, "CATEGORY_RULES", (body,))
    assert phr.product_category("Body Lotion SPF 30") == "Diğer"
    assert phr.product_category("Güneş Koruyucu Vücut Losyonu") == "Diğer"
    assert phr.product_category("Besleyici Vücut Losyonu") == "Vücut Losyonu"


def test_utc_bounds_are_tr_local_days():
    lo, hi = phr.utc_bounds(dt.date(2026, 5, 13), dt.date(2026, 5, 13))
    assert lo == dt.datetime(2026, 5, 12, 21, 0) and hi == dt.datetime(2026, 5, 13, 21, 0)


def test_default_range_is_last_12_months():
    s, e = phr.default_range(dt.date(2026, 10, 5))
    assert (s, e) == (dt.date(2025, 10, 6), dt.date(2026, 10, 5))
    s, e = phr.default_range(dt.date(2028, 2, 29))
    assert e == dt.date(2028, 2, 29) and s == dt.date(2027, 3, 1)


# ─── DB ─────────────────────────────────────────────────────────────────────

def test_date_filter_uses_tr_local_day(db_session: Session):
    it = _item(db_session, "Serenida Tonik")
    _ph(db_session, it, 5, dt.datetime(2026, 5, 12, 20, 0))    # 12.05 23:00 TR → dışarıda
    _ph(db_session, it, 7, dt.datetime(2026, 5, 12, 21, 30))   # 13.05 00:30 TR → içeride
    _ph(db_session, it, 11, dt.datetime(2026, 6, 30, 20, 59))  # 30.06 23:59 TR → içeride
    _ph(db_session, it, 13, dt.datetime(2026, 6, 30, 21, 0))   # 01.07 00:00 TR → dışarıda
    db_session.commit()
    rep = phr.build_report(db_session, "cosmetics", dt.date(2026, 5, 13), dt.date(2026, 6, 30))
    assert len(rep["rows"]) == 1
    row = rep["rows"][0]
    assert row["total"] == 18 and row["count"] == 2
    assert row["first"] == "2026-05-13" and row["last"] == "2026-06-30"
    assert row["monthly"] == {"2026-05": 7, "2026-06": 11}
    assert rep["months"] == ["2026-05", "2026-06"]
    # İlk kayıt filtreden bağımsız, TR günüyle
    assert rep["first_record"] == "2026-05-12"


def test_name_variants_merge_by_target_item(db_session: Session):
    sh, tonik = _seed(db_session)
    rep = phr.build_report(db_session, "cosmetics", dt.date(2026, 5, 1), dt.date(2026, 6, 30))
    rows = {r["name"]: r for r in rep["rows"]}
    assert set(rows) == {sh.name, "Serenida Tonik 100 ml"}
    t = rows["Serenida Tonik 100 ml"]
    assert t["total"] == 50 and t["count"] == 2 and t["target_item_id"] == tonik.id
    assert t["names"] == ["Serenida Tonik (100ml)"]           # eski yazım görünür kalır
    assert rows[sh.name]["brand"] == "Minerva 108" and rows[sh.name]["category"] == "Şampuan"
    assert rep["rows"][0]["name"] == sh.name                  # en çok üretilen başta
    assert rep["totals"]["quantity"] == 174 and rep["totals"]["products"] == 2


def test_current_item_name_wins_over_stored(db_session: Session):
    it = _item(db_session, "Serenida Tonik")
    _ph(db_session, it, 5, dt.datetime(2026, 5, 20, 9, 0), name="SERENİDA TONİK ESKİ")
    db_session.commit()
    it.name = "Serenida Facial Toner"; db_session.commit()
    rep = phr.build_report(db_session, "cosmetics", dt.date(2026, 5, 1), dt.date(2026, 5, 31))
    assert rep["rows"][0]["name"] == "Serenida Facial Toner"
    assert rep["rows"][0]["names"] == ["SERENİDA TONİK ESKİ"]


def test_brand_and_search_filters(db_session: Session):
    _seed(db_session)
    other = _item(db_session, "Minerva108 Duş Jeli")
    _ph(db_session, other, 10, dt.datetime(2026, 5, 22, 9, 0)); db_session.commit()
    s, e = dt.date(2026, 5, 1), dt.date(2026, 6, 30)
    rep = phr.build_report(db_session, "cosmetics", s, e, brand="Minerva 108")
    assert {r["brand"] for r in rep["rows"]} == {"Minerva 108"} and len(rep["rows"]) == 2
    assert rep["brands"] == ["Minerva 108", "Serenida"]      # çipler filtreden bağımsız
    rep = phr.build_report(db_session, "cosmetics", s, e, brand="minerva 108,SERENİDA")
    assert len(rep["rows"]) == 3
    # Türkçe katlamalı arama + eski ad yazımında da arar
    assert [r["name"] for r in phr.build_report(db_session, "cosmetics", s, e, q="sampuan")["rows"]] \
        == ["MİNERVA-108 ALTIN SERİ KURU SAÇ ŞAMPUANI"]
    assert len(phr.build_report(db_session, "cosmetics", s, e, q="(100ml)")["rows"]) == 1


def test_group_by_brand_and_category(db_session: Session):
    _seed(db_session)
    s, e = dt.date(2026, 5, 1), dt.date(2026, 6, 30)
    rep = phr.build_report(db_session, "cosmetics", s, e, group="brand")
    rows = {r["name"]: r for r in rep["rows"]}
    assert rows["Minerva 108"]["total"] == 124 and rows["Serenida"]["total"] == 50
    assert rows["Minerva 108"]["products"] == 1 and rows["Minerva 108"]["count"] == 2
    assert rows["Serenida"]["first"] == "2026-05-21" and rows["Serenida"]["last"] == "2026-06-10"
    rep = phr.build_report(db_session, "cosmetics", s, e, group="category")
    rows = {r["name"]: r for r in rep["rows"]}
    assert rows["Şampuan"]["total"] == 124 and rows["Tonik"]["total"] == 50
    assert rep["totals"]["monthly"] == {"2026-05": 130, "2026-06": 44}


def test_months_start_at_first_record(db_session: Session):
    """'Son 12 ay' seçilse de sistemdeki ilk kayıttan önceki aylar sütun olmaz."""
    _seed(db_session)                                         # ilk kayıt 20.05.2026
    rep = phr.build_report(db_session, "cosmetics", dt.date(2025, 10, 6), dt.date(2026, 7, 15))
    assert rep["months"] == ["2026-05", "2026-06", "2026-07"]
    assert rep["first_record"] == "2026-05-20"


def test_history_quantities(db_session: Session):
    sh, tonik = _seed(db_session)
    _ph(db_session, None, 9, dt.datetime(2026, 5, 25, 9, 0), name="Hedefsiz Eski Kayıt")
    db_session.commit()
    got = phr.history_quantities(db_session, "cosmetics", dt.date(2026, 5, 1), dt.date(2026, 5, 31))
    assert got == {sh.id: 100, tonik.id: 30}


def test_domain_isolation(authed_client: TestClient, db_session: Session):
    _seed(db_session)
    sup = _item(db_session, "Minerva 108 Bor+Geven Food Supplement", domain="supplement")
    _ph(db_session, sup, 500, dt.datetime(2026, 5, 20, 9, 0), domain="supplement")
    db_session.commit()
    d = authed_client.get(URL, params={"start": "2026-05-01", "end": "2026-06-30"}).json()
    names = {r["name"] for r in d["rows"]}
    assert sup.name not in names and len(names) == 2
    got = phr.history_quantities(db_session, "supplement", dt.date(2026, 5, 1), dt.date(2026, 6, 30))
    assert got == {sup.id: 500}


def test_endpoint_json_and_validation(authed_client: TestClient, db_session: Session):
    _seed(db_session)
    r = authed_client.get(URL, params={"start": "2026-05-01", "end": "2026-06-30",
                                       "group": "brand", "brand": "Serenida"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["group"] == "brand" and [x["name"] for x in d["rows"]] == ["Serenida"]
    assert d["first_record"] == "2026-05-20"
    assert d["note"] == "Sistemde üretim kaydı 20.05.2026 tarihinden beri var."
    # Varsayılan aralık (son 12 ay) parametresiz çalışır
    assert authed_client.get(URL).status_code == 200
    assert authed_client.get(URL, params={"start": "2026-07-01", "end": "2026-06-01"}).status_code == 400
    assert authed_client.get(URL, params={"group": "x"}).status_code == 422
    # Uç tarihler OverflowError → 500 yerine 400 (yalnız end verilince başlangıç end−364 gün)
    for params in ({"start": "0001-01-01", "end": "0001-01-05"}, {"end": "0001-01-05"},
                   {"end": "9999-12-31"}, {"start": "9999-12-30", "end": "9999-12-31"},
                   {"start": "1999-12-31"}):
        r = authed_client.get(URL, params=params)
        assert r.status_code == 400, (params, r.status_code, r.text)
    assert authed_client.get(URL, params={"start": "2000-01-01", "end": "2000-01-31"}).status_code == 200
    assert authed_client.get(URL, params={"end": "2100-12-31"}).status_code == 200


def test_endpoint_xlsx_export(authed_client: TestClient, db_session: Session):
    from openpyxl import load_workbook
    _seed(db_session)
    for group in ("product", "brand", "category"):
        r = authed_client.get(URL, params={"start": "2026-05-01", "end": "2026-06-30",
                                           "group": group, "format": "xlsx"})
        assert r.status_code == 200
        assert "attachment" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Özet", "Aylık"]
        ozet = [c.value for c in wb["Özet"]["B"]]
        assert "TOPLAM" in ozet
        aylik_head = [c.value for c in wb["Aylık"][3]]
        assert "May 2026" in aylik_head and "Haz 2026" in aylik_head and aylik_head[-1] == "Toplam"


def test_labtech_can_view(labtech_client: TestClient, db_session: Session):
    _seed(db_session)
    r = labtech_client.get(URL, params={"start": "2026-05-01", "end": "2026-06-30"})
    assert r.status_code == 200 and len(r.json()["rows"]) == 2
