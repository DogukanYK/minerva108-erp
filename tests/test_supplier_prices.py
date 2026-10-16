"""
Tedarikçi fiyat listesi — Işık Hanım'ın "Stok Son Durum" Excel'inden satın alma
raporunu otomatik dolduran modülün testleri.

  • parse_stok_son_durum  — onun sütun düzenini (Ted-1 paket ADından ÖNCE) doğru okur
  • import_prices         — malzeme/tedarikçi eşler, eşleşmeyeni raporlar, replace eder
  • prices_for_items      — en ucuzdan, None sona, limit
  • build_workbook(prices)— Satın Alma sayfası tedarikçi sütunlarıyla genişler
  • uçlar                 — içe aktar (finans), liste (reports.view), sil; yetki
"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api_main import app
from database import Item, Supplier, SupplierPrice

_H = {"Origin": "http://testserver"}
_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_xlsx(rows) -> bytes:
    """rows = [(material, kategori, birim, [(ad,paket,fiyat), …]), …] → Işık'ın düzeni."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Malzeme", "Kategori", "Birim", "Toplam Gereken (fireli)", "Mevcut Stok",
               "ALINACAK (eksik)", "ALINABİLECEK MİKTAR (KG)", "TEDARİKÇİ-1", "BİRİM FİYAT ",
               "TEDARİKÇİ-2", "ALINABİLECEK MİKTAR ", "BİRİM FİYAT ",
               "TEDARİKÇİ-3", "ALINABİLECEK MİKTAR ", "BİRİM FİYAT "])
    for material, kat, birim, sups in rows:
        r = [material, kat, birim, 100, 10, 90]
        s1 = sups[0] if len(sups) > 0 else (None, None, None)
        s2 = sups[1] if len(sups) > 1 else (None, None, None)
        s3 = sups[2] if len(sups) > 2 else (None, None, None)
        r += [s1[1], s1[0], s1[2]]   # Ted-1: paket, ad, fiyat
        r += [s2[0], s2[1], s2[2]]   # Ted-2: ad, paket, fiyat
        r += [s3[0], s3[1], s3[2]]   # Ted-3: ad, paket, fiyat
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _seed_items_suppliers(db):
    for nm in ["ARGAN YAĞI", "COCO BETAİNE%30"]:
        db.add(Item(name=nm, sku="SKU-" + nm[:6], category="Hammadde", unit="g",
                    current_stock=10, domain="cosmetics"))
    for sp in ["ULUDAĞ HERBAL", "NATURALYA", "UMAYCHEM"]:
        db.add(Supplier(name=sp, domain="cosmetics"))
    db.commit()


# ─── parse ──────────────────────────────────────────────────────────────────

def test_parse_layout():
    from core.supplier_prices import parse_stok_son_durum
    data = _make_xlsx([
        ("ARGAN YAĞI", "Hammadde", "g", [("ULUDAĞ HERBAL ", 25, 119.46),
                                         ("SURYA KİMYA ", 25, 22.0),
                                         ("NATURALYA ", 25, 25.8)]),
        ("COCO BETAİNE%30", "Hammadde", "g", [("UMAYCHEM ", 220, 1.37)]),
        ("TUZ", "Hammadde", "kg", []),
    ])
    rows = parse_stok_son_durum(data)
    assert len(rows) == 3
    argan = next(r for r in rows if r["material"] == "ARGAN YAĞI")
    assert [s["name"] for s in argan["suppliers"]] == ["ULUDAĞ HERBAL", "SURYA KİMYA", "NATURALYA"]
    assert argan["suppliers"][0]["package"] == 25 and argan["suppliers"][0]["price"] == 119.46
    coco = next(r for r in rows if r["material"].startswith("COCO"))
    assert len(coco["suppliers"]) == 1 and coco["suppliers"][0]["name"] == "UMAYCHEM"
    tuz = next(r for r in rows if r["material"] == "TUZ")
    assert tuz["suppliers"] == []


# ─── import ─────────────────────────────────────────────────────────────────

def test_import_matches_and_reports_unmatched(db_session):
    from core.supplier_prices import parse_stok_son_durum, import_prices
    _seed_items_suppliers(db_session)
    data = _make_xlsx([
        ("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8), ("ULUDAĞ HERBAL", 25, 119.46)]),
        ("COCO BETAİNE%30", "Hammadde", "g", [("UMAYCHEM", 220, 1.37)]),
        ("BİLİNMEYEN MALZEME", "Hammadde", "g", [("GİZEM TEDARİK", 1, 5.0)]),  # eşleşmeyen malzeme
    ])
    rows = parse_stok_son_durum(data)
    summ = import_prices(db_session, rows, "cosmetics")
    assert summ["items_updated"] == 2
    assert summ["prices_inserted"] == 3
    assert "BİLİNMEYEN MALZEME" in summ["unmatched_materials"]
    assert "GİZEM TEDARİK" not in summ["unmatched_suppliers"]  # eşleşmeyen malzeme → satır hiç eklenmedi
    argan = db_session.query(Item).filter(Item.name == "ARGAN YAĞI").first()
    prices = db_session.query(SupplierPrice).filter(SupplierPrice.item_id == argan.id).all()
    assert len(prices) == 2
    assert all(p.supplier_id is not None for p in prices)   # ikisi de eşleşti


def test_import_unmatched_supplier_kept_as_text(db_session):
    from core.supplier_prices import parse_stok_son_durum, import_prices
    _seed_items_suppliers(db_session)
    data = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("YENİ TEDARİKÇİ A.Ş.", 25, 10.0)])])
    summ = import_prices(db_session, parse_stok_son_durum(data), "cosmetics")
    assert "YENİ TEDARİKÇİ A.Ş." in summ["unmatched_suppliers"]
    row = db_session.query(SupplierPrice).first()
    assert row.supplier_id is None and row.supplier_name == "YENİ TEDARİKÇİ A.Ş."


def test_import_replaces_existing(db_session):
    from core.supplier_prices import parse_stok_son_durum, import_prices
    _seed_items_suppliers(db_session)
    import_prices(db_session, parse_stok_son_durum(
        _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8)])])), "cosmetics")
    # ikinci içe aktarma aynı malzemenin eski satırını DEĞİŞTİRİR
    import_prices(db_session, parse_stok_son_durum(
        _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("ULUDAĞ HERBAL", 25, 99.0)])])), "cosmetics")
    rows = db_session.query(SupplierPrice).all()
    assert len(rows) == 1 and rows[0].supplier_name == "ULUDAĞ HERBAL" and rows[0].unit_price == 99.0


# ─── lookup ─────────────────────────────────────────────────────────────────

def test_prices_for_items_cheapest_first(db_session):
    from core.supplier_prices import prices_for_items
    _seed_items_suppliers(db_session)
    argan = db_session.query(Item).filter(Item.name == "ARGAN YAĞI").first()
    db_session.add_all([
        SupplierPrice(item_id=argan.id, supplier_name="PAHALI", unit_price=119.0, package_size=25, domain="cosmetics"),
        SupplierPrice(item_id=argan.id, supplier_name="UCUZ", unit_price=22.0, package_size=25, domain="cosmetics"),
        SupplierPrice(item_id=argan.id, supplier_name="FİYATSIZ", unit_price=None, package_size=1, domain="cosmetics"),
    ])
    db_session.commit()
    out = prices_for_items(db_session, [argan.id], "cosmetics")
    names = [s["supplier_name"] for s in out[argan.id]]
    assert names == ["UCUZ", "PAHALI", "FİYATSIZ"]   # ucuzdan; fiyatsız sona


# ─── rapor (saf fonksiyon) ──────────────────────────────────────────────────

def test_build_workbook_adds_supplier_columns():
    import openpyxl
    from core.production_sim import build_workbook
    rep = {"summary": {"products": 1, "quantity": 250, "total_units": 250, "language": "TR",
                       "materials": 1, "short_count": 1, "ok_count": 0, "total_raw": 1.0},
           "materials": [], "producible": [],
           "purchase": [{"item_id": 7, "name": "ARGAN YAĞI", "category": "Hammadde",
                         "unit": "g", "used": 100, "current": 10, "shortfall": 90}],
           "skipped_labels": []}
    prices = {7: [{"supplier_name": "UCUZ", "package_size": 25, "unit_price": 22.0},
                  {"supplier_name": "PAHALI", "package_size": 25, "unit_price": 119.0}]}
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(rep, "x", prices=prices)))
    ws = wb["Satın Alma Listesi"]
    headers = [c.value for c in ws[1]]
    assert "TEDARİKÇİ-1" in headers and "TEDARİKÇİ-3" in headers
    row2 = [c.value for c in ws[2]]
    assert "UCUZ" in row2 and 22.0 in row2
    # prices boş → eski sade düzen
    ws2 = openpyxl.load_workbook(io.BytesIO(build_workbook(rep, "x", prices={})))["Satın Alma Listesi"]
    assert "TEDARİKÇİ-1" not in [c.value for c in ws2[1]]


# ─── uçlar ──────────────────────────────────────────────────────────────────

def test_endpoints_import_list_delete(authed_client: TestClient, db_session):
    _seed_items_suppliers(db_session)
    xls = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8), ("ULUDAĞ HERBAL", 25, 119.46)])])
    r = authed_client.post("/api/supplier-prices/import",
                           files={"file": ("STOK SON DURUM.xlsx", xls, _MIME)}, headers=_H)
    assert r.status_code == 200, r.text
    assert r.json()["items_updated"] == 1 and r.json()["prices_inserted"] == 2

    lst = authed_client.get("/api/supplier-prices")
    assert lst.status_code == 200
    body = lst.json()
    assert body["total_items"] == 1 and body["total_prices"] == 2
    item = body["items"][0]
    assert item["material"] == "ARGAN YAĞI"
    assert item["suppliers"][0]["supplier_name"] == "NATURALYA"   # en ucuz önce

    pid = item["suppliers"][0]["id"]
    d = authed_client.delete(f"/api/supplier-prices/{pid}", headers=_H)
    assert d.status_code == 200
    assert authed_client.get("/api/supplier-prices").json()["total_prices"] == 1


def test_import_requires_finance_role(labtech_client: TestClient, db_session):
    _seed_items_suppliers(db_session)
    xls = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8)])])
    r = labtech_client.post("/api/supplier-prices/import",
                            files={"file": ("x.xlsx", xls, _MIME)}, headers=_H)
    assert r.status_code == 403
    # ama listeyi görebilir (reports.view)
    assert labtech_client.get("/api/supplier-prices").status_code == 200


# ─── fiyat temeli (C1): para birimi + birim + kaynak ────────────────────────

def _seed_unit_items(db):
    """g / ml / adet birimli üç malzeme + bir tedarikçi."""
    db.add_all([
        Item(name="ARGAN YAĞI", sku="SKU-ARG", category="Hammadde", unit="g", current_stock=0, domain="cosmetics"),
        Item(name="GÜL SUYU", sku="SKU-GUL", category="Hammadde", unit="ml", current_stock=0, domain="cosmetics"),
        Item(name="POMPA 24/410", sku="SKU-POM", category="Ambalaj", unit="adet", current_stock=0, domain="cosmetics"),
        Supplier(name="NATURALYA", domain="cosmetics"),
    ])
    db.commit()


def _basis_xlsx():
    return _make_xlsx([
        ("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8)]),
        ("GÜL SUYU", "Hammadde", "ml", [("NATURALYA", 30, 4.2)]),
        ("POMPA 24/410", "Ambalaj", "adet", [("NATURALYA", 1000, 0.11)]),
    ])


def _rows_by_item(db):
    out = {}
    for sp in db.query(SupplierPrice).all():
        out[db.get(Item, sp.item_id).name] = sp
    return out


def test_import_endpoint_defaults_usd_kg_and_adet(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    r = authed_client.post("/api/supplier-prices/import",
                           files={"file": ("STOK SON DURUM.xlsx", _basis_xlsx(), _MIME)}, headers=_H)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["currency"] == "USD" and body["price_unit"] == "kg"
    assert body["source_label"] == "STOK SON DURUM.xlsx"       # etiket boş → dosya adı
    db_session.expire_all()
    rows = _rows_by_item(db_session)
    assert rows["ARGAN YAĞI"].currency == "USD" and rows["ARGAN YAĞI"].price_unit == "kg"
    assert rows["GÜL SUYU"].price_unit == "kg"                 # liste USD/kg; ml kalemi de kg fiyatı
    assert rows["POMPA 24/410"].price_unit == "adet"           # adet birimli → daima adet
    assert all(sp.source == "stok_son_durum" for sp in rows.values())
    assert all(sp.quoted_at is None for sp in rows.values())


def test_import_endpoint_form_sets_currency_unit_label_date(authed_client: TestClient, db_session):
    import datetime as dt
    _seed_unit_items(db_session)
    r = authed_client.post("/api/supplier-prices/import",
                           files={"file": ("x.xlsx", _basis_xlsx(), _MIME)},
                           data={"currency": "eur", "price_unit": "l",
                                 "label": "  Ekim   teklifleri ", "quoted_at": "2026-10-01"},
                           headers=_H)
    assert r.status_code == 200, r.text
    assert r.json()["currency"] == "EUR" and r.json()["quoted_at"] == "2026-10-01"
    db_session.expire_all()
    rows = _rows_by_item(db_session)
    assert rows["GÜL SUYU"].currency == "EUR" and rows["GÜL SUYU"].price_unit == "l"
    assert rows["ARGAN YAĞI"].price_unit == "l"                # kullanıcının seçtiği temel
    assert rows["POMPA 24/410"].price_unit == "adet"
    assert rows["ARGAN YAĞI"].source_label == "Ekim teklifleri"
    assert rows["ARGAN YAĞI"].quoted_at == dt.date(2026, 10, 1)

    lst = authed_client.get("/api/supplier-prices").json()
    by_name = {it["material"]: it["suppliers"][0] for it in lst["items"]}
    assert by_name["GÜL SUYU"]["currency"] == "EUR" and by_name["GÜL SUYU"]["price_unit"] == "l"
    assert by_name["POMPA 24/410"]["price_unit"] == "adet"
    assert by_name["ARGAN YAĞI"]["quoted_at"] == "2026-10-01"
    assert by_name["ARGAN YAĞI"]["source_label"] == "Ekim teklifleri"


def test_import_endpoint_rejects_bad_basis(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    for data, frag in (({"currency": "GBP"}, "Para birimi"),
                       ({"price_unit": "ton"}, "Fiyat birimi"),
                       # liste birimi 'adet' olamaz: her g/kg/ml/lt fiyatını
                       # satın alma planında kullanılamaz yapardı
                       ({"price_unit": "adet"}, "Fiyat birimi"),
                       ({"quoted_at": "01.10.2026"}, "tarih"),
                       ({"label": "x" * 121}, "Etiket")):
        r = authed_client.post("/api/supplier-prices/import",
                               files={"file": ("x.xlsx", _basis_xlsx(), _MIME)}, data=data, headers=_H)
        assert r.status_code == 400, (data, r.text)
        assert frag in r.json()["detail"]
    assert db_session.query(SupplierPrice).count() == 0       # hiçbiri yazılmadı


def test_import_endpoint_writes_audit_row(authed_client: TestClient, db_session):
    import json
    from database import AdminAuditLog
    _seed_unit_items(db_session)
    r = authed_client.post("/api/supplier-prices/import",
                           files={"file": ("liste.xlsx", _basis_xlsx(), _MIME)},
                           data={"currency": "USD", "price_unit": "kg", "label": "Eylül"}, headers=_H)
    assert r.status_code == 200, r.text
    logs = db_session.query(AdminAuditLog).filter(AdminAuditLog.action == "supplier_prices.import").all()
    assert len(logs) == 1
    det = json.loads(logs[0].details)
    assert logs[0].target_name == "Eylül"
    assert det["currency"] == "USD" and det["price_unit"] == "kg"
    assert det["items_updated"] == 3 and det["prices_inserted"] == 3
    assert det["filename"] == "liste.xlsx" and det["domain"] == "cosmetics"


def test_import_prices_core_kwargs_and_validation(db_session):
    import datetime as dt
    import pytest
    from core.supplier_prices import parse_stok_son_durum, import_prices, prices_for_items
    _seed_unit_items(db_session)
    rows = parse_stok_son_durum(_basis_xlsx())
    summ = import_prices(db_session, rows, "cosmetics", currency="TRY", price_unit="kg",
                         source_label="Elle", quoted_at=dt.date(2026, 9, 30))
    assert summ["prices_inserted"] == 3
    pump = db_session.query(Item).filter(Item.name == "POMPA 24/410").first()
    argan = db_session.query(Item).filter(Item.name == "ARGAN YAĞI").first()
    out = prices_for_items(db_session, [pump.id, argan.id], "cosmetics")
    assert out[pump.id][0]["price_unit"] == "adet" and out[pump.id][0]["currency"] == "TRY"
    assert out[argan.id][0]["price_unit"] == "kg"
    assert out[argan.id][0]["source_label"] == "Elle"
    assert out[argan.id][0]["quoted_at"] == dt.date(2026, 9, 30)
    with pytest.raises(ValueError):
        import_prices(db_session, rows, "cosmetics", currency="GBP")
    with pytest.raises(ValueError):
        import_prices(db_session, rows, "cosmetics", price_unit="ton")
    with pytest.raises(ValueError):
        import_prices(db_session, rows, "cosmetics", price_unit="adet")


# ─── tek seferlik geri doldurma (sentinel'li) ───────────────────────────────

def test_backfill_supplier_price_units_runs_once(db_session):
    from database import AppSetting, _backfill_supplier_price_units
    SENT = "backfill.supplier_price_units.v1"
    _seed_unit_items(db_session)
    # init_db (conftest) boş tabloda sentinel'i zaten yazdı → eski DB'yi taklit et
    db_session.query(AppSetting).filter(AppSetting.key == SENT).delete()
    items = {it.name: it for it in db_session.query(Item).all()}
    db_session.add_all([   # eski satırlar: model varsayılanı TRY, birimsiz, kaynaksız
        SupplierPrice(item_id=items["ARGAN YAĞI"].id, supplier_name="A", unit_price=25.8, domain="cosmetics"),
        SupplierPrice(item_id=items["GÜL SUYU"].id, supplier_name="B", unit_price=4.2, domain="cosmetics"),
        SupplierPrice(item_id=items["POMPA 24/410"].id, supplier_name="C", unit_price=0.11, domain="cosmetics"),
        # zaten temeli olan satıra DOKUNULMAZ
        SupplierPrice(item_id=items["ARGAN YAĞI"].id, supplier_name="D", unit_price=30.0, currency="EUR",
                      price_unit="l", source="manual", domain="cosmetics"),
    ])
    db_session.commit()
    assert db_session.query(SupplierPrice).filter(SupplierPrice.supplier_name == "A").one().currency == "TRY"

    _backfill_supplier_price_units()
    db_session.expire_all()
    got = {sp.supplier_name: sp for sp in db_session.query(SupplierPrice).all()}
    assert (got["A"].currency, got["A"].price_unit, got["A"].source) == ("USD", "kg", "stok_son_durum")
    assert (got["B"].currency, got["B"].price_unit) == ("USD", "kg")
    assert (got["C"].currency, got["C"].price_unit) == ("USD", "adet")
    assert (got["D"].currency, got["D"].price_unit, got["D"].source) == ("EUR", "l", "manual")
    sent = db_session.query(AppSetting).filter(AppSetting.key == SENT).one()
    assert sent.value == "3"

    # İkinci çağrı no-op: sentinel var → yeni birimsiz satır olduğu gibi kalır
    db_session.add(SupplierPrice(item_id=items["ARGAN YAĞI"].id, supplier_name="E",
                                 unit_price=1.0, domain="cosmetics"))
    db_session.commit()
    _backfill_supplier_price_units()
    db_session.expire_all()
    e = db_session.query(SupplierPrice).filter(SupplierPrice.supplier_name == "E").one()
    assert e.price_unit is None and e.currency == "TRY"


# ─── fiyatın alım birimine çevrilmesi ───────────────────────────────────────

def test_price_per_purchase_unit_cases():
    from core.supplier_prices import price_per_purchase_unit as ppu, LITRE_KG_NOTE
    kg = {"unit_price": 66.0, "price_unit": "kg"}
    lt = {"unit_price": 5.0, "price_unit": "l"}
    ad = {"unit_price": 0.11, "price_unit": "adet"}
    # g/kg kalem → kg temeli
    assert ppu(kg, "g") == (66.0, "kg", None)
    assert ppu(kg, "kg") == (66.0, "kg", None)
    assert ppu(kg, " KG ") == (66.0, "kg", None)
    # ml/l kalem → l temeli; kg fiyatı uygulanınca not düşülür
    assert ppu(kg, "ml") == (66.0, "l", LITRE_KG_NOTE)
    assert LITRE_KG_NOTE == "1 l ≈ 1 kg kabulüyle"
    assert ppu(lt, "ml") == (5.0, "l", None)
    assert ppu(lt, "lt") == (5.0, "l", None)
    assert ppu(lt, "g") == (5.0, "kg", LITRE_KG_NOTE)          # simetrik
    # adet
    assert ppu(ad, "adet") == (0.11, "adet", None)
    assert ppu(ad, "kutu") == (0.11, "adet", None)             # sayılan birim
    # çevrilemeyen temel → fiyat kullanılmaz, not açıklar
    price, unit, note = ppu(kg, "adet")
    assert price is None and unit == "adet" and "uyuşmuyor" in note
    price, unit, note = ppu(ad, "g")
    assert price is None and unit == "kg" and "uyuşmuyor" in note
    # fiyatsız teklif
    assert ppu({"unit_price": None, "price_unit": "kg"}, "g") == (None, "kg", None)
    # birimsiz eski satır → varsayılan kural (adet kalemde adet, diğerinde kg)
    assert ppu({"unit_price": 2.0, "price_unit": None}, "adet") == (2.0, "adet", None)
    assert ppu({"unit_price": 2.0}, "ml") == (2.0, "l", LITRE_KG_NOTE)
    # SupplierPrice nesnesi de kabul edilir
    assert ppu(SupplierPrice(unit_price=7.5, price_unit="kg"), "g") == (7.5, "kg", None)


# ─── Elle fiyat (lab — suppliers.prices, 07.10.2026) ────────────────────────

import json  # noqa: E402

from database import AdminAuditLog  # noqa: E402


def _login(c: TestClient, username: str) -> TestClient:
    c.post("/api/logout", headers=_H)
    r = c.post("/api/login", json={"username": username, "password": "minerva123"}, headers=_H)
    assert r.status_code == 200, r.text
    return c


def _ids(db):
    db.expire_all()
    items = {i.name: i.id for i in db.query(Item).all()}
    sups = {s.name: s.id for s in db.query(Supplier).all()}
    return items, sups


def _audit_actions(db):
    db.expire_all()
    return [a.action for a in db.query(AdminAuditLog).order_by(AdminAuditLog.id).all()]


def test_lablead_manual_price_crud(client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    _login(client, "songul")                                   # LabLead
    r = client.post("/api/supplier-prices", headers=_H, json={
        "item_id": items["ARGAN YAĞI"], "supplier_id": sups["NATURALYA"], "unit_price": 24.5,
        "currency": "eur", "price_unit": "kg", "package_size": 25, "quoted_at": "2026-10-07",
        "note": "telefonla teyit"})
    assert r.status_code == 201, r.text
    row = r.json()
    assert row["source"] == "manual" and row["created_by"] and row["currency"] == "EUR"
    assert row["material"] == "ARGAN YAĞI" and row["quoted_at"] == "2026-10-07"
    pid = row["id"]
    r = client.put(f"/api/supplier-prices/{pid}", headers=_H, json={"unit_price": 23.0})
    assert r.status_code == 200, r.text
    assert r.json()["unit_price"] == 23.0 and r.json()["package_size"] == 25   # kısmi
    assert r.json()["updated_by"]
    lst = client.get("/api/supplier-prices").json()
    s = lst["items"][0]["suppliers"][0]
    assert s["source"] == "manual" and s["note"] == "telefonla teyit" and s["supplier_status"] == "normal"
    assert client.delete(f"/api/supplier-prices/{pid}", headers=_H).status_code == 200
    acts = _audit_actions(db_session)
    for a in ("supplier_prices.create", "supplier_prices.update", "supplier_prices.delete"):
        assert a in acts
    db_session.expire_all()
    dl = db_session.query(AdminAuditLog).filter(AdminAuditLog.action == "supplier_prices.delete").one()
    d = json.loads(dl.details)
    assert d["unit_price"] == 23.0 and d["supplier_name"] == "NATURALYA" and d["source"] == "manual"


def test_labtech_cannot_write_prices(client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    sp = SupplierPrice(item_id=items["ARGAN YAĞI"], supplier_id=sups["NATURALYA"],
                       supplier_name="NATURALYA", unit_price=1, price_unit="kg", domain="cosmetics")
    db_session.add(sp); db_session.commit()
    _login(client, "meltem")                                   # LabTech
    assert client.post("/api/supplier-prices", headers=_H, json={
        "item_id": items["ARGAN YAĞI"], "supplier_id": sups["NATURALYA"], "unit_price": 2,
        "price_unit": "kg"}).status_code == 403
    assert client.put(f"/api/supplier-prices/{sp.id}", headers=_H,
                      json={"unit_price": 3}).status_code == 403
    assert client.delete(f"/api/supplier-prices/{sp.id}", headers=_H).status_code == 403


def test_manual_price_validation(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    base = {"supplier_id": sups["NATURALYA"], "unit_price": 1.5}
    # Birim ailesi: adetli kart → yalnız adet; g kart → kg|l
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "item_id": items["POMPA 24/410"], "price_unit": "kg"})
    assert r.status_code == 400
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "item_id": items["ARGAN YAĞI"], "price_unit": "adet"})
    assert r.status_code == 400
    for bad in ({"unit_price": 0}, {"unit_price": -1}, {"currency": "GBP"},
                {"quoted_at": "07.10.2026"}, {"package_size": 0}):
        r = authed_client.post("/api/supplier-prices", headers=_H,
                               json={**base, "item_id": items["ARGAN YAĞI"], "price_unit": "kg", **bad})
        assert r.status_code == 400, bad
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "item_id": items["POMPA 24/410"], "price_unit": "adet"})
    assert r.status_code == 201
    # Mükerrer (malzeme + tedarikçi + birim) → 409 + mevcut id
    r2 = authed_client.post("/api/supplier-prices", headers=_H,
                            json={**base, "item_id": items["POMPA 24/410"], "price_unit": "adet"})
    assert r2.status_code == 409
    assert r2.json()["code"] == "price_exists" and r2.json()["id"] == r.json()["id"]
    # Pasif tedarikçiye fiyat girilmez
    authed_client.delete(f"/api/suppliers/{sups['NATURALYA']}", headers=_H)
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "item_id": items["GÜL SUYU"], "price_unit": "l"})
    assert r.status_code == 400


def test_manual_price_other_domain_rejected(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    authed_client.post("/api/domain/switch", json={"domain": "supplement"}, headers=_H)
    r = authed_client.post("/api/supplier-prices", headers=_H, json={
        "item_id": items["ARGAN YAĞI"], "supplier_id": sups["NATURALYA"], "unit_price": 2,
        "price_unit": "kg"})
    assert r.status_code == 404


def test_import_keeps_manual_rows(authed_client: TestClient, db_session):
    _seed_items_suppliers(db_session)
    items, sups = _ids(db_session)
    r = authed_client.post("/api/supplier-prices", headers=_H, json={
        "item_id": items["ARGAN YAĞI"], "supplier_id": sups["NATURALYA"], "unit_price": 30,
        "price_unit": "kg"})
    assert r.status_code == 201
    manual_id = r.json()["id"]
    xls = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8),
                                                     ("ULUDAĞ HERBAL", 25, 119.46)])])
    r = authed_client.post("/api/supplier-prices/import",
                           files={"file": ("x.xlsx", xls, _MIME)}, headers=_H)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["kept_manual"] == 1 and body["skipped_manual"] == 1
    db_session.expire_all()
    rows = db_session.query(SupplierPrice).all()
    assert len(rows) == 2                                        # elle NATURALYA + Excel ULUDAĞ
    m = db_session.get(SupplierPrice, manual_id)
    assert m is not None and m.unit_price == 30 and m.source == "manual"
    # Tekrar içe aktarma da elle satıra dokunmaz
    authed_client.post("/api/supplier-prices/import",
                       files={"file": ("x.xlsx", xls, _MIME)}, headers=_H)
    db_session.expire_all()
    assert db_session.get(SupplierPrice, manual_id).unit_price == 30


def test_edited_import_row_becomes_manual_and_survives(authed_client: TestClient, db_session):
    _seed_items_suppliers(db_session)
    xls = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8)])])
    authed_client.post("/api/supplier-prices/import",
                       files={"file": ("x.xlsx", xls, _MIME)}, headers=_H)
    db_session.expire_all()
    pid = db_session.query(SupplierPrice).one().id
    r = authed_client.put(f"/api/supplier-prices/{pid}", headers=_H, json={"unit_price": 21.0})
    assert r.status_code == 200 and r.json()["source"] == "manual"
    authed_client.post("/api/supplier-prices/import",
                       files={"file": ("x.xlsx", xls, _MIME)}, headers=_H)
    db_session.expire_all()
    rows = db_session.query(SupplierPrice).all()
    assert len(rows) == 1 and rows[0].id == pid and rows[0].unit_price == 21.0


def test_import_matches_only_active_suppliers(db_session):
    from core.supplier_prices import parse_stok_son_durum, import_prices
    db_session.add(Item(name="ARGAN YAĞI", category="Hammadde", unit="g", current_stock=0,
                        domain="cosmetics"))
    old = Supplier(name="TATLIDİLİMLER", domain="cosmetics", is_active=False)
    db_session.add(old); db_session.commit()
    new = Supplier(name="TATLİDİLİMLER", domain="cosmetics")
    db_session.add(new); db_session.commit()
    import_prices(db_session, parse_stok_son_durum(
        _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("Tatlıdilimler", 1, 9.0)])])), "cosmetics")
    assert db_session.query(SupplierPrice).one().supplier_id == new.id


def test_prices_for_items_manual_hides_same_firm_import(db_session):
    from core.supplier_prices import prices_for_items
    _seed_items_suppliers(db_session)
    items, sups = _ids(db_session)
    iid = items["ARGAN YAĞI"]
    db_session.add_all([
        SupplierPrice(item_id=iid, supplier_id=sups["NATURALYA"], supplier_name="NATURALYA",
                      unit_price=25.8, source="stok_son_durum", domain="cosmetics"),
        SupplierPrice(item_id=iid, supplier_id=sups["NATURALYA"], supplier_name="NATURALYA",
                      unit_price=30.0, source="manual", domain="cosmetics"),
        SupplierPrice(item_id=iid, supplier_id=sups["UMAYCHEM"], supplier_name="UMAYCHEM",
                      unit_price=40.0, source="stok_son_durum", domain="cosmetics"),
    ])
    db_session.commit()
    out = prices_for_items(db_session, [iid], "cosmetics")
    assert [(s["supplier_name"], s["unit_price"]) for s in out[iid]] == \
        [("NATURALYA", 30.0), ("UMAYCHEM", 40.0)]


def test_price_unit_ok_rules():
    from core.supplier_prices import price_unit_ok
    assert price_unit_ok("adet", "adet") and not price_unit_ok("adet", "kg")
    assert price_unit_ok("kutu", "adet")
    assert price_unit_ok("g", "kg") and price_unit_ok("g", "l") and not price_unit_ok("g", "adet")
    assert price_unit_ok("ml", "l") and price_unit_ok("lt", "kg")


# ─── Pasif tedarikçi + düzenlenen Excel satırının etiketi (08.10.2026) ──────
# Yumuşak silme (P1a) sonrası pasif firmanın fiyatı en ucuz seçilmeye devam
# ediyordu: Excel'lerde Tedarikçi-1, panelde yeşil "en ucuz".

def _inactive_case(db):
    """Jojoba: Eski Firma (pasif) 5 $, Yeni Firma 9 $, TATLİDİLİMLER kopyası
    pasif ama TATLIDİLİMLER aktif (aynı firma) 7 $, kartsız serbest metin
    "Kapanan kimya" 4 $ (adı yalnız pasif KAPANAN KİMYA kartında)."""
    it = Item(name="JOJOBA", category="Hammadde", unit="g", domain="cosmetics")
    old = Supplier(name="ESKİ FİRMA", domain="cosmetics", is_active=False)
    new = Supplier(name="YENİ FİRMA", domain="cosmetics")
    dup = Supplier(name="TATLİDİLİMLER", domain="cosmetics", is_active=False)
    twin = Supplier(name="TATLIDİLİMLER", domain="cosmetics")
    gone = Supplier(name="KAPANAN KİMYA", domain="cosmetics", is_active=False)
    db.add_all([it, old, new, dup, twin, gone])
    db.flush()
    db.add_all([
        SupplierPrice(item_id=it.id, supplier_id=old.id, supplier_name="ESKİ FİRMA", unit_price=5.0,
                      currency="USD", price_unit="kg", source="manual", domain="cosmetics"),
        SupplierPrice(item_id=it.id, supplier_id=new.id, supplier_name="YENİ FİRMA", unit_price=9.0,
                      currency="USD", price_unit="kg", source="manual", domain="cosmetics"),
        SupplierPrice(item_id=it.id, supplier_id=dup.id, supplier_name="TATLİDİLİMLER", unit_price=7.0,
                      currency="USD", price_unit="kg", source="stok_son_durum", domain="cosmetics"),
        SupplierPrice(item_id=it.id, supplier_id=None, supplier_name="Kapanan kimya", unit_price=4.0,
                      currency="USD", price_unit="kg", source="stok_son_durum", domain="cosmetics"),
    ])
    db.commit()
    return it.id


def test_prices_for_items_inactive_firm_last_and_flagged(db_session):
    from core.supplier_prices import prices_for_items
    iid = _inactive_case(db_session)
    out = prices_for_items(db_session, [iid], "cosmetics", limit=5)
    assert [(s["supplier_name"], s["unit_price"], s["inactive"]) for s in out[iid]] == [
        ("TATLİDİLİMLER", 7.0, False),          # pasif kopya kart, aktif ikiz → geçerli
        ("YENİ FİRMA", 9.0, False),
        ("Kapanan kimya", 4.0, True),           # serbest metin, adı yalnız pasif kartta
        ("ESKİ FİRMA", 5.0, True)]
    assert [s["supplier_name"] for s in prices_for_items(db_session, [iid], "cosmetics")[iid]][:2] == \
        ["TATLİDİLİMLER", "YENİ FİRMA"]


def test_supplier_price_panel_marks_inactive_firm(authed_client: TestClient, db_session):
    iid = _inactive_case(db_session)
    r = authed_client.get("/api/supplier-prices")
    assert r.status_code == 200
    g = next(x for x in r.json()["items"] if x["item_id"] == iid)
    assert [(s["supplier_name"], s["supplier_inactive"]) for s in g["suppliers"]] == [
        ("TATLİDİLİMLER", False), ("YENİ FİRMA", False), ("Kapanan kimya", True), ("ESKİ FİRMA", True)]


def test_excel_inactive_supplier_grey_and_labelled():
    import openpyxl
    from core.production_sim import build_workbook
    rep = {"summary": {"products": 1, "quantity": 1, "total_units": 1, "language": "TR",
                       "materials": 1, "short_count": 1, "ok_count": 0, "total_raw": 1.0},
           "materials": [], "producible": [],
           "purchase": [{"item_id": 7, "name": "JOJOBA", "category": "Hammadde",
                         "unit": "g", "used": 100, "current": 10, "shortfall": 90}],
           "skipped_labels": []}
    prices = {7: [{"supplier_name": "ESKİ FİRMA", "package_size": 25, "unit_price": 5.0, "inactive": True}]}
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(rep, "x", prices=prices)))["Satın Alma Listesi"]
    c = ws.cell(2, 7)
    assert c.value == "ESKİ FİRMA (pasif tedarikçi)"
    assert not c.font.bold and c.font.color.rgb.endswith("9CA3AF")


def test_edited_import_row_drops_stale_list_label(authed_client: TestClient, db_session):
    """Excel satırı elle düzenlenince eski liste etiketi düşer — satın alma
    planı elle fiyatı "<dosya adı>" kaynağıyla yazmasın; etiket gönderilirse o."""
    _seed_items_suppliers(db_session)
    xls = _make_xlsx([("ARGAN YAĞI", "Hammadde", "g", [("NATURALYA", 25, 25.8)])])
    authed_client.post("/api/supplier-prices/import", files={"file": ("liste.xlsx", xls, _MIME)}, headers=_H)
    db_session.expire_all()
    sp = db_session.query(SupplierPrice).one()
    assert sp.source_label == "liste.xlsx"
    r = authed_client.put(f"/api/supplier-prices/{sp.id}", headers=_H, json={"unit_price": 21.0})
    assert r.status_code == 200
    assert (r.json()["source"], r.json()["source_label"]) == ("manual", None)
    r = authed_client.put(f"/api/supplier-prices/{sp.id}", headers=_H,
                          json={"unit_price": 22.0, "source_label": "Telefon teklifi"})
    assert r.json()["source_label"] == "Telefon teklifi"
    r = authed_client.put(f"/api/supplier-prices/{sp.id}", headers=_H, json={"unit_price": 23.0})
    assert r.json()["source_label"] == "Telefon teklifi"          # zaten elle: etiketine dokunulmaz


# ─── KDV (lab fiyat notları, 08.10.2026) ────────────────────────────────────

def test_net_unit_price_and_vat_label():
    from core.supplier_prices import net_unit_price, vat_label
    assert net_unit_price({"unit_price": 82.56, "vat_included": True, "vat_rate": 20}) == pytest.approx(68.8)
    assert vat_label({"unit_price": 82.56, "vat_included": True, "vat_rate": 20}) == "KDV %20 dahil"
    # Boş / hariç / oransız / geçersiz oran → fiyat AYNEN (tip bile değişmez)
    for o in ({"unit_price": 5}, {"unit_price": 5, "vat_included": False, "vat_rate": 20},
              {"unit_price": 5, "vat_included": True}, {"unit_price": 5, "vat_included": True, "vat_rate": 150},
              {"unit_price": 5, "vat_included": True, "vat_rate": float("nan")}):
        assert net_unit_price(o) == 5 and type(net_unit_price(o)) is int
    assert net_unit_price({"unit_price": None, "vat_included": True, "vat_rate": 20}) is None
    assert vat_label({"unit_price": 5, "vat_included": True}) == "KDV dahil"
    assert vat_label({"unit_price": 5, "vat_included": False, "vat_rate": 20}) is None
    sp = SupplierPrice(unit_price=12.0, vat_included=True, vat_rate=1.0)      # nesne de olur
    assert net_unit_price(sp) == pytest.approx(12.0 / 1.01) and vat_label(sp) == "KDV %1 dahil"


def test_prices_for_items_sorts_by_net_and_carries_vat(db_session):
    from core.supplier_prices import prices_for_items
    _seed_unit_items(db_session)
    argan = db_session.query(Item).filter(Item.name == "ARGAN YAĞI").first()
    db_session.add_all([
        SupplierPrice(item_id=argan.id, supplier_name="BRÜT", unit_price=12.0, vat_included=True, vat_rate=20,
                      domain="cosmetics"),                                         # net 10
        SupplierPrice(item_id=argan.id, supplier_name="NET", unit_price=11.0, domain="cosmetics"),
    ])
    db_session.commit()
    lst = prices_for_items(db_session, [argan.id], "cosmetics")[argan.id]
    assert [s["supplier_name"] for s in lst] == ["BRÜT", "NET"]
    assert lst[0]["vat_included"] is True and lst[0]["vat_rate"] == 20 and lst[1]["vat_included"] is None


def test_manual_price_vat_fields_api(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    base = {"item_id": items["ARGAN YAĞI"], "supplier_id": sups["NATURALYA"], "unit_price": 12.0,
            "currency": "EUR", "price_unit": "kg"}
    # KDV dahil + oran yok → 400; oran aralık dışı → 400
    assert authed_client.post("/api/supplier-prices", headers=_H,
                              json={**base, "vat_included": True}).status_code == 400
    for bad in (-1, 101, "abc"):
        r = authed_client.post("/api/supplier-prices", headers=_H, json={**base, "vat_rate": bad})
        assert r.status_code in (400, 422), bad
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "vat_included": True, "vat_rate": 20})
    assert r.status_code == 201, r.text
    row = r.json()
    assert row["vat_included"] is True and row["vat_rate"] == 20.0
    pid = row["id"]
    # Kısmi PUT: yalnız oranı silmek "dahil + oran yok" bırakır → 400
    assert authed_client.put(f"/api/supplier-prices/{pid}", headers=_H,
                             json={"vat_rate": None}).status_code == 400
    r = authed_client.put(f"/api/supplier-prices/{pid}", headers=_H, json={"vat_included": False})
    assert r.status_code == 200 and r.json()["vat_included"] is False and r.json()["vat_rate"] == 20.0
    r = authed_client.put(f"/api/supplier-prices/{pid}", headers=_H, json={"vat_included": None, "vat_rate": None})
    assert r.status_code == 200 and r.json()["vat_included"] is None and r.json()["vat_rate"] is None
    # Alanlar gönderilmezse boş kalır (eski davranış) — başka malzemede
    r = authed_client.post("/api/supplier-prices", headers=_H,
                           json={**base, "item_id": items["GÜL SUYU"], "price_unit": "l"})
    assert r.status_code == 201 and r.json()["vat_included"] is None and r.json()["vat_rate"] is None
    # Audit: KDV değişikliği eski → yeni
    db_session.expire_all()
    upd = (db_session.query(AdminAuditLog).filter(AdminAuditLog.action == "supplier_prices.update")
           .order_by(AdminAuditLog.id).all())
    assert json.loads(upd[0].details)["changes"]["vat_included"] == {"eski": True, "yeni": False}


def test_supplier_price_panel_sorts_by_net(authed_client: TestClient, db_session):
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    db_session.add_all([
        SupplierPrice(item_id=items["ARGAN YAĞI"], supplier_name="NET", unit_price=11.0, currency="EUR",
                      price_unit="kg", domain="cosmetics"),
        SupplierPrice(item_id=items["ARGAN YAĞI"], supplier_id=sups["NATURALYA"], supplier_name="NATURALYA",
                      unit_price=12.0, currency="EUR", price_unit="kg", vat_included=True, vat_rate=20,
                      source="lab_notu", source_label="Lab fiyat notları 07.10.2026 — s.1", domain="cosmetics"),
    ])
    db_session.commit()
    sups_out = authed_client.get("/api/supplier-prices").json()["items"][0]["suppliers"]
    assert [s["supplier_name"] for s in sups_out] == ["NATURALYA", "NET"]
    assert sups_out[0]["vat_included"] is True and sups_out[0]["vat_rate"] == 20


def test_excel_exports_label_vat_included_supplier():
    """Üretim stok analizi Excel'i: KDV dahil satırın adı etiketli, fiyat brüt."""
    import openpyxl
    from core.production_sim import build_workbook
    rep = {"summary": {"products": 1, "quantity": 250, "total_units": 250, "language": "TR",
                       "materials": 1, "short_count": 1, "ok_count": 0, "total_raw": 1.0},
           "materials": [], "producible": [],
           "purchase": [{"item_id": 7, "name": "VANİLYA", "category": "Hammadde",
                         "unit": "ml", "used": 100, "current": 10, "shortfall": 90}],
           "skipped_labels": []}
    prices = {7: [{"supplier_name": "PHARMATERM", "package_size": 25, "unit_price": 82.56,
                   "vat_included": True, "vat_rate": 20},
                  {"supplier_name": "ULUDAĞ", "package_size": 25, "unit_price": 70.0}]}
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(rep, "x", prices=prices)))["Satın Alma Listesi"]
    row2 = [c.value for c in ws[2]]
    assert "PHARMATERM (KDV %20 dahil)" in row2 and 82.56 in row2 and "ULUDAĞ" in row2


def test_import_keeps_lab_note_rows_and_still_writes_same_firm(db_session):
    """Stok Son Durum içe aktarma lab notu satırını silmez; aynı firmanın Excel
    satırını yalnız ELLE satır engeller (lab satırı engellemez)."""
    from core.supplier_prices import import_prices
    _seed_unit_items(db_session)
    items, sups = _ids(db_session)
    lab_row = SupplierPrice(item_id=items["ARGAN YAĞI"], supplier_id=sups["NATURALYA"], supplier_name="NATURALYA",
                            unit_price=25.8, currency="EUR", price_unit="kg", source="lab_notu",
                            source_label="Lab fiyat notları 07.10.2026 — s.1", domain="cosmetics")
    db_session.add(lab_row)
    db_session.commit()
    summary = import_prices(db_session, [{"material": "ARGAN YAĞI",
                                          "suppliers": [{"name": "NATURALYA", "package": 25, "price": 28.0}]}],
                            "cosmetics")
    assert summary["prices_inserted"] == 1 and summary["kept_manual"] == 0 and summary["skipped_manual"] == 0
    db_session.expire_all()
    got = sorted((p.source, p.unit_price) for p in db_session.query(SupplierPrice).all())
    assert got == [("lab_notu", 25.8), ("stok_son_durum", 28.0)]


# ─── Para birimi karışık malzeme (08.10.2026, lab notları EUR/TRY) ───────────
# Ham fiyatı kur çevirmeden sıralamak 250 ₺/kg'ı (≈5 €) en pahalı gösteriyor,
# ilk 3'te kalacak ucuz teklifi kesebiliyordu; Excel de "Tedarikçi-1 en uygun"
# diyordu.

def _mixed_case(db):
    """Lavanta: Uludağ 1,88 €, BEFCHEM 9 €, Pharmaterm 250 ₺/kg (≈5,3 €), Stok
    Son Durum 12 $ — dört teklif, limit 3."""
    _seed_unit_items(db)
    items, _sups = _ids(db)
    iid = items["ARGAN YAĞI"]
    db.add_all([
        SupplierPrice(item_id=iid, supplier_name="ULUDAĞ", unit_price=1.88, currency="EUR", price_unit="kg",
                      domain="cosmetics"),
        SupplierPrice(item_id=iid, supplier_name="BEFCHEM", unit_price=9.0, currency="EUR", price_unit="kg",
                      domain="cosmetics"),
        SupplierPrice(item_id=iid, supplier_name="PHARMATERM", unit_price=250.0, currency="TRY", price_unit="kg",
                      domain="cosmetics"),
        SupplierPrice(item_id=iid, supplier_name="ESKİ LİSTE", unit_price=12.0, currency="USD", price_unit="kg",
                      domain="cosmetics"),
    ])
    db.commit()
    return iid


def _rates():
    from core.fx import Rates
    return Rates(try_per={"USD": 40.0, "EUR": 47.0, "TRY": 1.0}, source="test", as_of=None,
                 stale=False, warning=None)


def test_prices_for_items_mixed_currency_sorted_with_rates_lazily(db_session):
    from core.supplier_prices import prices_for_items
    iid = _mixed_case(db_session)
    calls = []

    def lazy():
        calls.append(1)
        return _rates()

    lst = prices_for_items(db_session, [iid], "cosmetics", rates=lazy)[iid]
    # 1,88 € = 2,21 $ · 250 ₺ = 6,25 $ · 9 € = 10,58 $ · 12 $ → kesilen ESKİ LİSTE
    assert [s["supplier_name"] for s in lst] == ["ULUDAĞ", "PHARMATERM", "BEFCHEM"]
    assert all(s["comparable"] for s in lst) and calls == [1]
    # Tek para birimli malzemede kur hiç istenmez
    calls.clear()
    other = Item(name="TEK", sku="SKU-TEK", category="Hammadde", unit="g", domain="cosmetics")
    db_session.add(other)
    db_session.flush()
    db_session.add(SupplierPrice(item_id=other.id, supplier_name="A", unit_price=3.0, currency="EUR",
                                 price_unit="kg", domain="cosmetics"))
    db_session.commit()
    assert prices_for_items(db_session, [other.id], "cosmetics", rates=lazy)[other.id][0]["comparable"]
    assert calls == []


def test_prices_for_items_mixed_currency_without_rates_not_comparable(db_session):
    from core.supplier_prices import prices_for_items
    iid = _mixed_case(db_session)
    for rates in (None, lambda: None):
        lst = prices_for_items(db_session, [iid], "cosmetics", rates=rates)[iid]
        assert [s["supplier_name"] for s in lst] == ["ULUDAĞ", "BEFCHEM", "ESKİ LİSTE"]   # ham sıra (eski)
        assert not any(s["comparable"] for s in lst)


def test_excel_price_cells_show_currency_and_skip_cheapest_claim_when_mixed():
    """Fiyat hücresi sayı kalır ama biçimi para birimi/birimi gösterir;
    karşılaştırılamayan grupta Tedarikçi-1 kalın değil."""
    import openpyxl
    from core.production_sim import build_workbook
    from core.supplier_prices import price_cell_format
    rep = {"summary": {"products": 1, "quantity": 250, "total_units": 250, "language": "TR",
                       "materials": 2, "short_count": 2, "ok_count": 0, "total_raw": 1.0},
           "materials": [], "producible": [],
           "purchase": [{"item_id": 1, "name": "LAVANTA", "category": "Hammadde", "unit": "g", "used": 100,
                         "current": 10, "shortfall": 90},
                        {"item_id": 2, "name": "SUSAM", "category": "Hammadde", "unit": "g", "used": 100,
                         "current": 10, "shortfall": 90}],
           "skipped_labels": []}
    prices = {1: [{"supplier_name": "ULUDAĞ", "package_size": 25, "unit_price": 1.88, "currency": "EUR",
                   "price_unit": "kg", "comparable": False},
                  {"supplier_name": "PHARMATERM", "package_size": 5, "unit_price": 250.0, "currency": "TRY",
                   "price_unit": "kg", "comparable": False}],
              2: [{"supplier_name": "KRK", "package_size": 5, "unit_price": 7.0, "currency": "EUR",
                   "price_unit": "kg", "comparable": True}]}
    ws = openpyxl.load_workbook(io.BytesIO(build_workbook(rep, "x", prices=prices)))["Satın Alma Listesi"]
    assert ws.cell(2, 7).value == "ULUDAĞ" and not ws.cell(2, 7).font.bold
    assert ws.cell(3, 7).value == "KRK" and ws.cell(3, 7).font.bold
    assert ws.cell(2, 9).value == 1.88 and ws.cell(2, 9).number_format == '#,##0.00##" €/kg"'
    assert ws.cell(2, 12).value == 250.0 and "₺/kg" in ws.cell(2, 12).number_format
    assert price_cell_format({"unit_price": 3}) == "#,##0.00##"


# ─── Malzeme Fiyatları sayfası (/fiyatlar, 16.10.2026) ───────────────────────
# GET /api/supplier-prices/overview → core.price_overview.build_overview.
# Grup → kart → satır; planın durum bayrakları; elle fiyat aynı firmanın Excel
# satırını gizler; net fiyat seçilen para biriminde (kur sahte — ağa çıkılmaz).

from datetime import date as _date  # noqa: E402

from database import AppSetting, MaterialGroup, MaterialSupplierPref, User  # noqa: E402


def _ov(db, **kw):
    from core.price_overview import build_overview
    kw.setdefault("rates_fn", _rates)
    return build_overview(db, "cosmetics", **kw)


def _cards(ov):
    return [c for b in ov["blocks"] for c in b["cards"]]


def _card(ov, name):
    return next(c for c in _cards(ov) if c["name"] == name)


def _sp(item, sup=None, price=1.0, *, name=None, cur="USD", unit="kg", pkg=None, source="stok_son_durum",
        domain="cosmetics", **kw):
    return SupplierPrice(item_id=item.id, supplier_id=(sup.id if sup is not None else None),
                         supplier_name=name or (sup.name if sup is not None else None), unit_price=price,
                         currency=cur, price_unit=unit, package_size=pkg, source=source, domain=domain, **kw)


def _overview_case(db):
    """SETİL / CETYL aynı grupta (iki görünür kart → tek blok), ARGAN grupsuz,
    pasif kart ve bitmiş ürün fiyatlı ama görünmez; JOJOBA grubunun tek
    görünür kartı (diğeri pasif) blok açmaz."""
    g = MaterialGroup(name="Setil stearil alkol", domain="cosmetics")
    gj = MaterialGroup(name="Jojoba grubu", domain="cosmetics")
    db.add_all([g, gj]); db.flush()
    um = Supplier(name="UMAYCHEM", domain="cosmetics")
    nat = Supplier(name="NATURALYA", domain="cosmetics")
    db.add_all([um, nat]); db.flush()
    setil = Item(name="SETİL STEARİL ALKOL", category="Hammadde", unit="g", domain="cosmetics",
                 material_group_id=g.id, supplier_id=um.id)
    cetyl = Item(name="CETYL STEARYL ALCOHOL", category="Hammadde", unit="g", domain="cosmetics",
                 material_group_id=g.id)
    argan = Item(name="ARGAN YAĞI", category="Hammadde", unit="g", domain="cosmetics")
    old = Item(name="ESKİ KART", category="Hammadde", unit="g", domain="cosmetics", is_active=False)
    krem = Item(name="KREM", category="Bitmiş Ürün", unit="adet", domain="cosmetics")
    jo1 = Item(name="JOJOBA", category="Hammadde", unit="g", domain="cosmetics", material_group_id=gj.id)
    jo2 = Item(name="JOJOBA ESKİ", category="Hammadde", unit="g", domain="cosmetics",
               material_group_id=gj.id, is_active=False)
    db.add_all([setil, cetyl, argan, old, krem, jo1, jo2]); db.flush()
    db.add_all([
        _sp(setil, um, 5.0, pkg=25),                               # Excel — elle satır gizler
        _sp(setil, um, 4.8, pkg=25, source="manual", created_by="Songül"),
        _sp(cetyl, nat, 6.0, pkg=20),
        _sp(argan, nat, 10.0, pkg=5, unit="l"),
        _sp(old, nat, 1.0),
        _sp(krem, nat, 3.0, unit="adet"),
        _sp(jo1, um, 9.0), _sp(jo2, um, 8.0),
    ])
    db.commit()
    return {"setil": setil.id, "cetyl": cetyl.id, "argan": argan.id, "group": g.id}


def test_overview_groups_dedup_and_excludes_inactive_and_finished(db_session):
    ids = _overview_case(db_session)
    ov = _ov(db_session, has_price="priced")
    names = [[c["name"] for c in b["cards"]] for b in ov["blocks"]]
    assert names == [["ARGAN YAĞI"], ["JOJOBA"], ["CETYL STEARYL ALCOHOL", "SETİL STEARİL ALKOL"]]
    gb = ov["blocks"][2]
    assert gb["group"] == {"id": ids["group"], "name": "Setil stearil alkol"}
    assert ov["blocks"][1]["group"] is None and _card(ov, "JOJOBA")["group_name"] == "Jojoba grubu"
    setil = _card(ov, "SETİL STEARİL ALKOL")
    assert [(r["unit_price"], r["source"]) for r in setil["rows"]] == [(4.8, "manual")]
    assert setil["hidden_rows"] == 1 and "elle girilen fiyatla gizlendi" in setil["hidden_text"]
    assert setil["card_supplier"] == "UMAYCHEM" and setil["purchase_unit"] == "kg"
    r = setil["rows"][0]
    # Ambalaj FİYAT BİRİMİNDE (25 kg), kartın g birimiyle değil; satır SPF formunun şekli
    assert r["package_text"] == "25 kg" and r["material"] == "SETİL STEARİL ALKOL" and r["unit"] == "g"
    assert r["net_text"] == "4,80 $/kg" and r["orig_text"] == "4,80 $/kg" and r["matched"] is True
    assert _card(ov, "ARGAN YAĞI")["rows"][0]["package_text"] == "5 litre"
    s = ov["summary"]
    assert s == {"cards": 4, "priced_cards": 4, "unpriced_cards": 0, "rows": 4, "hidden_rows": 1, "groups": 1}
    assert ov["summary_text"] == "4 malzeme · 4 fiyat satırı · 1 aynı malzeme grubu"
    assert ov["fx"]["needed"] is False                     # hepsi USD → kur yok, ağ yok


def test_overview_search_expands_to_group_and_matches_supplier(db_session):
    _overview_case(db_session)
    # Grubun bir üyesine uyan arama grubun bütün kartlarını getirir (karşılaştırma için)
    ov = _ov(db_session, q="setil")
    assert [[c["name"] for c in b["cards"]] for b in ov["blocks"]] == [
        ["CETYL STEARYL ALCOHOL", "SETİL STEARİL ALKOL"]]
    assert [c["name"] for c in _cards(_ov(db_session, q="naturalya"))] == [
        "ARGAN YAĞI", "CETYL STEARYL ALCOHOL", "SETİL STEARİL ALKOL"]
    assert _cards(_ov(db_session, q="yok böyle")) == []


def test_overview_status_flags_follow_purchase_plan_rules(db_session):
    """Firma (tercih edilen / bitirilecek / pasif / atlanacaklar) ve malzeme
    (tercih sırası / bu malzemede alma) bayrakları planın `_offer_status`
    kuralıyla; en ucuz pasif olmayanlardan, planın seçtiği tercih kademesinden."""
    it = Item(name="LAVANTA YAĞI", category="Hammadde", unit="g", domain="cosmetics")
    pref = Supplier(name="TERCİHLİ KİMYA", domain="cosmetics", purchase_status="preferred")
    out = Supplier(name="BİTEN FİRMA", domain="cosmetics", purchase_status="phase_out",
                   status_reason="kalite sorunu")
    avo = Supplier(name="ALINMAZ LTD", domain="cosmetics")
    gone = Supplier(name="KAPANAN KİMYA", domain="cosmetics", is_active=False)
    mat = Supplier(name="MALZEME TERCİHİ", domain="cosmetics")
    skip = Supplier(name="MİNERVA", domain="cosmetics")
    db_session.add_all([it, pref, out, avo, gone, mat, skip]); db_session.flush()
    db_session.add_all([
        _sp(it, pref, 9.0), _sp(it, out, 3.0), _sp(it, avo, 4.0), _sp(it, gone, 1.0),
        _sp(it, mat, 12.0), _sp(it, skip, 2.0),
        MaterialSupplierPref(domain="cosmetics", item_id=it.id, supplier_id=mat.id, preference="preferred", rank=1),
        MaterialSupplierPref(domain="cosmetics", item_id=it.id, supplier_id=avo.id, preference="avoid", rank=1),
    ])
    db_session.commit()
    c = _card(_ov(db_session), "LAVANTA YAĞI")
    st = {r["supplier_name"]: (r["status"], r["eligible"], r["status_tags"], r["status_blocked"])
          for r in c["rows"]}
    assert st == {
        "TERCİHLİ KİMYA": ("preferred", True, ["tercih edilen firma"], None),
        "BİTEN FİRMA": ("phase_out", False, [], "bitirilecek, alma"),
        "ALINMAZ LTD": ("avoid", False, [], "bu malzemede alma"),
        "KAPANAN KİMYA": ("inactive", False, [], "pasif tedarikçi"),
        "MALZEME TERCİHİ": ("preferred", True, ["tercih"], None),
        "MİNERVA": ("skip", False, [], "atlanacaklar listesinde"),
    }
    by = {r["supplier_name"]: r for r in c["rows"]}
    assert by["BİTEN FİRMA"]["status_reason"] == "kalite sorunu"
    assert by["KAPANAN KİMYA"]["supplier_inactive"] is True and by["KAPANAN KİMYA"]["status_kind"] == "inactive"
    # En ucuz: pasif firma (1 $) sayılmaz → MİNERVA 2 $; plan malzeme tercihini seçer (12 $)
    assert [r["supplier_name"] for r in c["rows"] if r["cheapest"]] == ["MİNERVA"]
    assert [r["supplier_name"] for r in c["rows"] if r["plan_pick"]] == ["MALZEME TERCİHİ"]
    # Sıra: uygunlar ucuzdan → uygun olmayanlar → pasif sonda
    assert [r["supplier_name"] for r in c["rows"]] == [
        "TERCİHLİ KİMYA", "MALZEME TERCİHİ", "MİNERVA", "BİTEN FİRMA", "ALINMAZ LTD", "KAPANAN KİMYA"]
    assert c["pref_text"] == "Tercih: MALZEME TERCİHİ · Bu malzemede alma: ALINMAZ LTD"
    # Atlanacaklar listesi AppSetting'ten okunur (plan ile aynı ayar)
    db_session.add(AppSetting(key="purchase_plan.skip_suppliers", value="BAŞKA FİRMA"))
    db_session.commit()
    by = {r["supplier_name"]: r for r in _card(_ov(db_session), "LAVANTA YAĞI")["rows"]}
    assert by["MİNERVA"]["status"] == "normal"


def test_overview_group_prefs_apply_to_member_cards(db_session):
    g = MaterialGroup(name="Argan", domain="cosmetics")
    db_session.add(g); db_session.flush()
    a = Item(name="ARGAN A", category="Hammadde", unit="g", domain="cosmetics", material_group_id=g.id)
    b = Item(name="ARGAN B", category="Hammadde", unit="g", domain="cosmetics", material_group_id=g.id)
    s1 = Supplier(name="BİRİNCİ", domain="cosmetics")
    s2 = Supplier(name="İKİNCİ", domain="cosmetics")
    db_session.add_all([a, b, s1, s2]); db_session.flush()
    db_session.add_all([
        _sp(a, s1, 5.0), _sp(a, s2, 6.0), _sp(b, s1, 5.0),
        MaterialSupplierPref(domain="cosmetics", material_group_id=g.id, supplier_id=s2.id,
                             preference="preferred", rank=1),
        # Kartın kendi satırı grubunkini ezer
        MaterialSupplierPref(domain="cosmetics", item_id=b.id, supplier_id=s1.id, preference="avoid", rank=1),
    ])
    db_session.commit()
    ov = _ov(db_session)
    ca, cb = _card(ov, "ARGAN A"), _card(ov, "ARGAN B")
    assert [r["supplier_name"] for r in ca["rows"] if r["plan_pick"]] == ["İKİNCİ"]
    assert [(r["supplier_name"], r["status"]) for r in cb["rows"]] == [("BİRİNCİ", "avoid")]
    assert not any(r["plan_pick"] for r in cb["rows"])


def test_overview_fx_conversion_and_vat_net(db_session, monkeypatch):
    """Net = KDV hariç fiyat → alım birimi → seçilen para birimi; kur TCMB
    sahte fetcher'ından BİR KEZ (hepsi aynı para biriminde ise hiç)."""
    import core.fx as fx
    calls = []

    def fake_fetch():
        calls.append(1)
        return {"USD": 40.0, "EUR": 44.0, "as_of": _date(2026, 10, 16)}

    fx.reset_cache()
    monkeypatch.setattr(fx, "_fetch_tcmb_rates", fake_fetch)
    it = Item(name="BİBERİYE", category="Hammadde", unit="g", domain="cosmetics")
    db_session.add(it); db_session.flush()
    db_session.add_all([
        _sp(it, name="ULUDAĞ", price=10.0, cur="EUR"),
        _sp(it, name="PHARMATERM", price=400.0, cur="TRY"),
        _sp(it, name="KDVLİ", price=12.6, vat_included=True, vat_rate=20),     # net 10,50 $
    ])
    db_session.commit()
    from core.price_overview import build_overview
    try:
        ov = build_overview(db_session, "cosmetics")                 # gerçek today_rates → sahte fetch
        assert calls == [1]
        by = {r["supplier_name"]: r for r in _card(ov, "BİBERİYE")["rows"]}
        assert by["ULUDAĞ"]["net"] == pytest.approx(11.0) and by["ULUDAĞ"]["net_text"] == "11,00 $/kg"
        assert by["ULUDAĞ"]["orig_text"] == "10,00 €/kg"
        assert by["PHARMATERM"]["net"] == pytest.approx(10.0)
        assert by["KDVLİ"]["net"] == pytest.approx(10.5) and by["KDVLİ"]["vat_label"] == "KDV %20 dahil"
        assert by["KDVLİ"]["orig_text"] == "12,60 $/kg"
        assert [r["supplier_name"] for r in _card(ov, "BİBERİYE")["rows"]] == ["PHARMATERM", "KDVLİ", "ULUDAĞ"]
        assert by["PHARMATERM"]["cheapest"] and _card(ov, "BİBERİYE")["comparable"] is True
        assert ov["fx"]["text"] == "kur: TCMB 16.10" and ov["fx"]["needed"] is True
        assert ov["fx"]["rates_text"] == "1 $ = 40,0000 ₺ · 1 € = 44,0000 ₺"
        ov = build_overview(db_session, "cosmetics", currency="EUR")  # 1 saat önbellek → ağ yok
        assert calls == [1] and ov["symbol"] == "€"
        by = {r["supplier_name"]: r for r in _card(ov, "BİBERİYE")["rows"]}
        assert by["ULUDAĞ"]["net"] == pytest.approx(10.0) and by["ULUDAĞ"]["net_text"] == "10,00 €/kg"
        assert by["KDVLİ"]["net"] == pytest.approx(10.5 * 40 / 44)
    finally:
        fx.reset_cache()


def test_overview_without_rates_marks_rows_and_card(db_session):
    it = Item(name="BİBERİYE", category="Hammadde", unit="g", domain="cosmetics")
    db_session.add(it); db_session.flush()
    db_session.add_all([_sp(it, name="ULUDAĞ", price=10.0, cur="EUR"), _sp(it, name="ESKİ LİSTE", price=12.0)])
    db_session.commit()

    def boom():
        raise RuntimeError("TCMB yok")
    ov = _ov(db_session, rates_fn=boom)
    c = _card(ov, "BİBERİYE")
    by = {r["supplier_name"]: r for r in c["rows"]}
    assert by["ULUDAĞ"]["net"] is None and by["ULUDAĞ"]["net_text"] is None
    assert "kur yok" in by["ULUDAĞ"]["net_note"] and by["ESKİ LİSTE"]["net_text"] == "12,00 $/kg"
    assert c["comparable"] is False and "kur yok" in c["compare_note"]
    assert not any(r["cheapest"] for r in c["rows"])
    assert [r["supplier_name"] for r in c["rows"] if r["plan_pick"]] == ["ESKİ LİSTE"]   # planla aynı: çevrilemeyen fiyatsız
    assert ov["fx"]["text"] == "kur yok" and ov["fx"]["available"] is False
    # Tek para birimi (EUR) + USD görünüm + kur yok → ham fiyatla karşılaştırılır
    db_session.query(SupplierPrice).filter(SupplierPrice.supplier_name == "ESKİ LİSTE").update(
        {"currency": "EUR", "unit_price": 9.0})
    db_session.commit()
    c = _card(_ov(db_session, rates_fn=boom), "BİBERİYE")
    assert c["comparable"] is True and [r["supplier_name"] for r in c["rows"] if r["cheapest"]] == ["ESKİ LİSTE"]


def test_overview_mixed_units_not_comparable_and_litre_note(db_session):
    it = Item(name="GÜL SUYU", category="Hammadde", unit="ml", domain="cosmetics")
    lab = Item(name="ŞİŞE 100 ML", category="Ambalaj", pkg_type="şişe", unit="adet", domain="cosmetics")
    db_session.add_all([it, lab]); db_session.flush()
    db_session.add_all([_sp(it, name="A FİRMA", price=3.0, unit="kg", pkg=20),
                        _sp(it, name="B FİRMA", price=4.0, unit="l", pkg=5),
                        _sp(lab, name="C FİRMA", price=0.1193, unit="adet", pkg=1000)])
    db_session.commit()
    ov = _ov(db_session)
    c = _card(ov, "GÜL SUYU")
    assert c["comparable"] is False and c["compare_note"] == "Farklı fiyat birimi (kg / litre) — karşılaştırılamaz"
    assert not any(r["cheapest"] for r in c["rows"])
    by = {r["supplier_name"]: r for r in c["rows"]}
    assert by["A FİRMA"]["litre_kg"] is True and by["A FİRMA"]["net_text"] == "3,00 $/litre"
    assert by["B FİRMA"]["package_text"] == "5 litre" and by["A FİRMA"]["package_text"] == "20 kg"
    r = _card(ov, "ŞİŞE 100 ML")["rows"][0]
    assert r["net_text"] == "0,1193 $/adet" and r["package_text"] == "1.000 adet"


def test_overview_filters(db_session):
    _overview_case(db_session)
    db_session.add_all([
        Item(name="KAPAK", category="Ambalaj", pkg_type="kapak", unit="adet", domain="cosmetics"),
        Item(name="ETİKET X", category="Ambalaj", pkg_type="etiket", unit="adet", domain="cosmetics"),
        Item(name="YARDIMCI", category="Sarf", unit="adet", domain="cosmetics"),
    ])
    db_session.commit()
    items, _ = _ids(db_session)
    db_session.add(SupplierPrice(item_id=items["ARGAN YAĞI"], supplier_name="LAB FİRMA", unit_price=7.0,
                                 currency="USD", price_unit="l", source="lab_notu",
                                 source_label="Lab fiyat notları 07.10.2026 — s.3", domain="cosmetics"))
    db_session.commit()
    names = lambda ov: [c["name"] for c in _cards(ov)]                        # noqa: E731
    # Fiyatsız: yalnız hammadde / ambalaj / etiket kartları (Sarf ve bitmiş ürün yok)
    un = _ov(db_session, has_price="unpriced")
    assert names(un) == ["ETİKET X", "KAPAK"]
    assert all(not c["priced"] and c["rows"] == [] for c in _cards(un))
    assert un["summary"]["unpriced_cards"] == 2 and "2 fiyatsız" in un["summary_text"]
    allc = names(_ov(db_session))                                              # varsayılan: hepsi
    assert "KAPAK" in allc and "ARGAN YAĞI" in allc and "YARDIMCI" not in allc and "KREM" not in allc
    assert names(_ov(db_session, category="etiket")) == ["ETİKET X"]
    assert names(_ov(db_session, category="ambalaj")) == ["KAPAK"]
    assert "KAPAK" not in names(_ov(db_session, category="hammadde"))
    # Kaynak: kartta o kaynaktan satır olmalı; satırlar daralır, sayısı kartta
    lab = _ov(db_session, source="lab")
    assert names(lab) == ["ARGAN YAĞI"]
    c = _card(lab, "ARGAN YAĞI")
    assert [r["supplier_name"] for r in c["rows"]] == ["LAB FİRMA"] and c["filtered_rows"] == 1
    assert "1 fiyat satırı daha" in c["filtered_text"]
    assert names(_ov(db_session, source="manual")) == ["SETİL STEARİL ALKOL"]
    assert "SETİL STEARİL ALKOL" not in names(_ov(db_session, source="stok_son_durum"))   # Excel satırı gizli
    # "Diğer" kovası (Sarf / Kimyasal / Yardımcı — plan hammadde sayıp "Fiyat gir →"
    # bağlantısı verir) aramada ve "diğer" süzgecinde gelir; bitmiş ürün hiç gelmez
    assert names(_ov(db_session, q="yardımcı")) == ["YARDIMCI"]
    assert names(_ov(db_session, category="other")) == ["YARDIMCI"]
    assert names(_ov(db_session, q="krem")) == []
    assert "YARDIMCI" not in names(_ov(db_session, q="yardımcı", has_price="priced"))
    with pytest.raises(ValueError):
        _ov(db_session, currency="GBP")
    with pytest.raises(ValueError):
        _ov(db_session, has_price="belki")


def test_overview_plan_pick_spans_plan_merged_twin_cards(db_session):
    """Plan aynı katlanmış ad + birimli iki aktif hammadde kartını (kept kararı /
    grup yoksa) tek satırda birleştirir ve teklifi iki kartın fiyatlarından
    seçer — "planda seçilir" küme başına BİR kez; fiyatsız ikizi fiyatlı kart
    "fiyatı olmayanlar"a düşmez; kept kararı birleştirmeyi kapatır."""
    from database import DuplicateItemDecision
    a = Item(name="ALOE VERA EKSTRAKTI", category="Hammadde", unit="ml", domain="cosmetics")
    b = Item(name="ALOEVERA EKSTRAKTI", category="Hammadde", unit="ml", domain="cosmetics")
    k1 = Item(name="KAPAK 24 MM", category="Ambalaj", pkg_type="kapak", unit="adet", domain="cosmetics")
    k2 = Item(name="KAPAK 24MM", category="Ambalaj", pkg_type="kapak", unit="adet", domain="cosmetics")
    db_session.add_all([a, b, k1, k2]); db_session.flush()
    db_session.add_all([_sp(a, name="NATURALYA", price=12.0, unit="l"), _sp(a, name="AROMA", price=13.0, unit="l"),
                        _sp(b, name="KRK", price=10.0, unit="l"), _sp(k1, name="KAPAKÇI", price=0.05, unit="adet")])
    db_session.commit()
    ov = _ov(db_session)
    ca, cb = _card(ov, "ALOE VERA EKSTRAKTI"), _card(ov, "ALOEVERA EKSTRAKTI")
    assert not any(r["plan_pick"] for r in ca["rows"])                  # plan KRK'dan alır
    assert [r["supplier_name"] for r in cb["rows"] if r["plan_pick"]] == ["KRK"]
    assert ca["merged_with"] == [{"item_id": b.id, "name": "ALOEVERA EKSTRAKTI"}]
    assert "ALOEVERA EKSTRAKTI ile tek satırda birleştirir" in ca["merge_text"]
    assert [r["supplier_name"] for r in ca["rows"] if r["cheapest"]] == ["NATURALYA"]   # en ucuz kart başına
    # İkiz görünmese de (arama yalnız birini buluyor) seçim kümeden yapılır
    only = _ov(db_session, q="aloe vera")
    assert [c["name"] for c in _cards(only)] == ["ALOE VERA EKSTRAKTI"]
    assert not any(r["plan_pick"] for r in _card(only, "ALOE VERA EKSTRAKTI")["rows"])
    # Fiyatsız ikiz: "hepsi"de görünür (planda fiyatı var), "fiyatı olmayanlar"da yok
    k2c = _card(ov, "KAPAK 24MM")
    assert k2c["priced"] is False and k2c["twin_priced"] is True and k2c["rows"] == []
    un = _ov(db_session, has_price="unpriced")
    assert "KAPAK 24MM" not in [c["name"] for c in _cards(un)]
    assert ov["summary"]["unpriced_cards"] == 0 and "fiyatsız" not in ov["summary_text"]
    # Lab "ayrı kalsın" dediyse plan birleştirmez → kart kendi en iyisini seçer
    db_session.add(DuplicateItemDecision(cluster_key="aloe", title="Aloe", status="kept",
                                         item_ids=json.dumps([a.id, b.id])))
    db_session.commit()
    ov = _ov(db_session)
    ca = _card(ov, "ALOE VERA EKSTRAKTI")
    assert [r["supplier_name"] for r in ca["rows"] if r["plan_pick"]] == ["NATURALYA"]
    assert ca["merged_with"] == [] and ca["merge_text"] is None


def test_overview_endpoint_domain_scoped_and_validated(authed_client: TestClient, db_session):
    _overview_case(db_session)
    sup_item = Item(name="KAPSÜL TOZU", category="Hammadde", unit="g", domain="supplement")
    db_session.add(sup_item); db_session.flush()
    db_session.add(SupplierPrice(item_id=sup_item.id, supplier_name="TAKVİYE FİRMA", unit_price=2.0,
                                 currency="USD", price_unit="kg", domain="supplement"))
    db_session.commit()
    r = authed_client.get("/api/supplier-prices/overview?has_price=priced")
    assert r.status_code == 200, r.text
    names = [c["name"] for b in r.json()["blocks"] for c in b["cards"]]
    assert "KAPSÜL TOZU" not in names and "ARGAN YAĞI" in names
    authed_client.cookies.set("active_domain", "supplement")
    r = authed_client.get("/api/supplier-prices/overview")
    assert [c["name"] for b in r.json()["blocks"] for c in b["cards"]] == ["KAPSÜL TOZU"]
    for bad in ("currency=GBP", "category=bitmis", "has_price=x", "source=excel"):
        r = authed_client.get("/api/supplier-prices/overview?" + bad)
        assert r.status_code == 400 and isinstance(r.json()["detail"], str), bad
    assert authed_client.get("/api/supplier-prices/overview?q=" + "x" * 101).status_code == 422


def test_overview_endpoint_requires_reports_view(client: TestClient, db_session):
    import bcrypt
    u = User(username="raporsuz", full_name="raporsuz", role="Staff", is_active=True,
             password_hash=bcrypt.hashpw(b"minerva123", bcrypt.gensalt()).decode(),
             permissions=json.dumps({"items": {"view": True}}))
    db_session.add(u); db_session.commit()
    assert client.get("/api/supplier-prices/overview").status_code == 401
    r = client.get("/fiyatlar", follow_redirects=False)
    assert r.status_code == 302 and r.headers["location"] == "/login"
    _login(client, "raporsuz")
    assert client.get("/api/supplier-prices/overview").status_code == 403
    r = client.get("/fiyatlar", follow_redirects=False)
    assert r.status_code == 302 and r.headers["location"] == "/"
    _login(client, "meltem")                                                    # LabTech: reports.view var
    assert client.get("/api/supplier-prices/overview").status_code == 200
