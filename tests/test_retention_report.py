"""
Şahit numune dolabı raporu (PDF + Excel) testleri.

Kritik sözleşmeler:
  • rapor, liste ucuyla AYNI filtre yolundan geçer — ekranda görünen kayıt
    kümesi ile çıktıdaki küme birebir aynı (limit hariç)
  • imha edilenler ve yanlış girilip silinenler (is_active=False) varsayılan
    raporda YOK; `status=destroyed` ile istenirse gelir
  • supplement panelinin kaydı kozmetik raporuna sızmaz (ve tersi)
  • LabTech `retention.view` ile raporu alabilir
  • gruplama dolap → raf → göz, raf/göz DOĞAL sırayla ("2" < "10")
"""
from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from core.retention_report import build_report, filters_text, natural_key
from database import Item, RetentionSample, RetentionSampleCheck

_HDR = {"Origin": "http://testserver"}
_URL = "/api/retention/samples/export"


def _item(db: Session, name: str, domain="cosmetics", **kw) -> Item:
    it = Item(name=name, sku=f"sku-{name}", category="Bitmiş Ürün", unit="adet",
              current_stock=0, domain=domain, **kw)
    db.add(it); db.flush()
    return it


def _row(db: Session, item: Item, lot: str, *, brand: str, shelf=None, slot=None,
         qty=2.0, status="stored", until=None, domain=None, active=True, note=None):
    r = RetentionSample(item_id=item.id, item_name=item.name, lot_number=lot,
                        brand=brand, shelf=shelf, slot=slot, quantity=qty,
                        initial_quantity=max(qty, 3.0), unit="adet", status=status,
                        source="manual", placed_by="test", retention_until=until,
                        note=note, is_active=active,
                        domain=domain or item.domain or "cosmetics")
    db.add(r); db.flush()
    return r


@pytest.fixture
def seeded(db_session):
    """İki dolap (Minerva 108 + Serenida), raf/göz karışık, bir imha + bir
    silinmiş kayıt + supplement panelinde bir kayıt."""
    db = db_session
    m1 = _item(db, "Minerva 108 Şampuan Isırgan")
    m2 = _item(db, "Minerva 108 Saç Kremi")
    s1 = _item(db, "Serenida Duş Jeli")
    sup = _item(db, "Supplement Kapsül", domain="supplement")
    today = date.today()
    _row(db, m1, "MNR010-S", brand="Minerva 108", shelf="10", slot="A", qty=3)
    _row(db, m1, "MNR002-S", brand="Minerva 108", shelf="2", slot="C10", qty=1,
         until=today - timedelta(days=5), note="Kapak çatlak")          # süresi dolmuş
    _row(db, m2, "MNR003-S", brand="Minerva 108", shelf="2", slot="C2", qty=4,
         until=today + timedelta(days=30))                               # yaklaşıyor
    _row(db, m2, "MNR004-S", brand="Minerva 108", qty=2)                 # rafsız
    _row(db, s1, "SR001-S", brand="Serenida", shelf="1", slot="B", qty=5)
    _row(db, s1, "SR002-S", brand="Serenida", shelf="1", slot="A", qty=0,
         status="depleted")
    _row(db, s1, "SR003-S", brand="Serenida", shelf="1", qty=2, status="destroyed")
    _row(db, s1, "SR004-S", brand="Serenida", shelf="1", qty=2, active=False)
    _row(db, sup, "SUP001-S", brand="Supplement", shelf="1", qty=9)
    db.commit()
    return db


def _xlsx(resp):
    assert resp.status_code == 200, resp.text
    return load_workbook(BytesIO(resp.content))


def _lots(wb) -> list:
    ws = wb["Dolap"]
    head = [c.value for c in ws[1]]
    col = head.index("Lot")
    return [r[col] for r in ws.iter_rows(min_row=2, values_only=True)]


def _pdf_text(resp) -> str:
    assert resp.status_code == 200, resp.text
    assert resp.content.startswith(b"%PDF")
    return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(resp.content)).pages)


# ─── Saf birim ──────────────────────────────────────────────────────────────

def test_natural_key_orders_shelves_and_slots():
    assert sorted(["10", "2", "", "B", "A10", "A2", "1"], key=natural_key) == \
        ["1", "2", "10", "A2", "A10", "B", ""]


def test_natural_key_survives_non_decimal_digits():
    """"²"/"①" isdigit'e True ama int()'e girmez — rapor 500'e düşmemeli."""
    for s in ("²", "①", "1²", "A①3"):
        natural_key(s)
    assert sorted(["1²", "1"], key=natural_key) == ["1", "1²"]


def test_filters_text_describes_screen_filters():
    assert filters_text() == "Tüm dolaplar · İmha edilenler hariç"
    t = filters_text(brand="Serenida", q=" krem ", expired=1, unsized=1,
                     item_label="Serenida Duş Jeli 200ml")
    assert "Dolap: Serenida" in t and "Arama: “krem”" in t
    assert "Yalnız süresi dolanlar" in t and "Boyu belirsiz" in t
    assert "Ürün: Serenida Duş Jeli 200ml" in t
    assert "Durum: İmha edildi" in filters_text(status="destroyed")


def test_build_report_counts_expiry_only_for_stored():
    """Tükenmiş kaydın geçmiş tarihi "süresi dolan" sayılmaz — dolapta yok."""
    base = dict(brand="Serenida", shelf="1", slot="A", item_name="X", lot_number="L",
                quantity=1, initial_quantity=1, retention_until_label="—")
    rep = build_report([
        {**base, "status": "stored", "expiry_state": "expired"},
        {**base, "status": "depleted", "expiry_state": "expired", "quantity": 0},
        {**base, "status": "stored", "expiry_state": "due_soon", "is_parent": True},
    ])
    t = rep["totals"]
    assert (t["samples"], t["stored"], t["expired"], t["due_soon"], t["unsized"]) == (3, 2, 1, 1, 1)
    rows = rep["cabinets"][0]["shelves"][0]["rows"]
    assert {r["expiry"] for r in rows} == {"expired", "", "due_soon"}


# ─── Uç: biçim + içerik ─────────────────────────────────────────────────────

def test_pdf_export_basic(authed_client, seeded):
    r = authed_client.get(f"{_URL}?format=pdf")
    text = _pdf_text(r)
    assert r.headers["content-type"] == "application/pdf"
    cd = r.headers["content-disposition"]
    assert cd.startswith("inline;") and "sahit_numune_dolabi_" in cd and ".pdf" in cd
    for lot in ("MNR010-S", "MNR002-S", "MNR003-S", "MNR004-S", "SR001-S", "SR002-S"):
        assert lot in text
    assert "SR003-S" not in text            # imha → varsayılan raporda yok
    assert "SR004-S" not in text            # silinmiş kayıt
    assert "SUP001-S" not in text           # supplement paneli sızmaz
    assert "Dolap: Minerva 108" in text and "Dolap: Serenida" in text
    assert "Raf belirtilmemiş" in text      # rafsız kayıt kendi grubunda
    assert "Sayfa 1 /" in text
    assert "Şampuan Isırgan" in text        # Türkçe karakterler bozulmadan


def test_pdf_multi_page_repeats_header(authed_client, db_session):
    it = _item(db_session, "Minerva 108 Losyon")
    for i in range(70):
        _row(db_session, it, f"MNR{i:03d}-S", brand="Minerva 108", shelf="1",
             slot=str(i % 9 + 1), note="Uzun bir not metni satırı iki satıra kaydırsın diye")
    db_session.commit()
    r = authed_client.get(f"{_URL}?format=pdf")
    pages = PdfReader(BytesIO(r.content)).pages
    assert len(pages) >= 3
    for p in pages:
        txt = p.extract_text() or ""
        assert "Saklama bitişi" in txt      # başlık satırı her sayfada
        assert f"/ {len(pages)}" in txt     # "Sayfa X / N" altbilgisi


def test_xlsx_export_rows_order_and_summary(authed_client, seeded):
    r = authed_client.get(f"{_URL}?format=xlsx")
    wb = _xlsx(r)
    assert r.headers["content-disposition"].startswith("attachment;")
    assert ".xlsx" in r.headers["content-disposition"]
    assert wb.sheetnames == ["Dolap", "Özet"]
    ws = wb["Dolap"]
    head = [c.value for c in ws[1]]
    assert head[:6] == ["Dolap", "Raf", "Göz", "Ürün", "Boy", "Lot"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    # dolap (TR alfabe) → raf (doğal: 2 < 10, boş en sonda) → göz (doğal: C2 < C10)
    assert _lots(wb) == ["MNR003-S", "MNR002-S", "MNR010-S", "MNR004-S",
                         "SR002-S", "SR001-S"]
    by_lot = {row[head.index("Lot")]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert by_lot["SR001-S"][head.index("Kalan")] == 5
    assert by_lot["MNR002-S"][head.index("Not")] == "Kapak çatlak"
    assert by_lot["MNR002-S"][head.index("Saklama bitişi")] is not None   # gerçek tarih

    summary = list(wb["Özet"].iter_rows(values_only=True))
    tot = next(x for x in summary if x and x[0] == "TOPLAM")
    # 6 kayıt, 5'i dolapta, kalan 3+1+4+2+5+0 = 15, süresi dolan 1, yaklaşan 1
    assert tot[1:6] == (6, 5, 15, 1, 1)
    assert any(x and x[0] == "Filtre" for x in summary)


def test_xlsx_free_text_starting_with_equals_stays_text(authed_client, db_session):
    """Serbest metin "=" ile başlarsa openpyxl formül yazar: Excel "onarım"
    uyarısı + kayıp not, ya da canlı =HYPERLINK.  Düz metin kalmalı."""
    from zipfile import ZipFile
    it = _item(db_session, "=Formül Ürün")
    _row(db_session, it, '=HYPERLINK("http://x","t")', brand="Minerva 108",
         shelf="1", note="=2 adet kırık")
    db_session.commit()
    r = authed_client.get(f"{_URL}?format=xlsx")
    wb = _xlsx(r)
    ws = wb["Dolap"]
    head = [c.value for c in ws[1]]
    row = next(ws.iter_rows(min_row=2, values_only=True))
    assert row[head.index("Not")] == "=2 adet kırık"
    assert row[head.index("Lot")] == '=HYPERLINK("http://x","t")'
    assert row[head.index("Ürün")].startswith("=Formül Ürün")
    with ZipFile(BytesIO(r.content)) as z:
        sheets = [z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/")]
    assert sheets and not any(b"<f>" in x or b"<f " in x for x in sheets)


def test_brand_filter(authed_client, seeded):
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx&brand=Serenida"))
    assert sorted(_lots(wb)) == ["SR001-S", "SR002-S"]
    assert "Dolap: Serenida" in _pdf_text(authed_client.get(f"{_URL}?format=pdf&brand=Serenida"))
    filt = dict(x[:2] for x in wb["Özet"].iter_rows(values_only=True) if x and x[0])
    assert "Dolap: Serenida" in filt["Filtre"]


def test_status_destroyed_filter_includes_destroyed(authed_client, seeded):
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx&status=destroyed"))
    assert _lots(wb) == ["SR003-S"]


@pytest.mark.parametrize("qs", [
    "", "brand=Minerva%20108", "q=Isırgan", "q=sr00", "expired=1",
    "unchecked=1", "needs_review=1", "status=depleted",
])
def test_export_matches_list_endpoint(authed_client, seeded, qs):
    """Ekrandaki liste ile rapor AYNI kayıt kümesini döner (tek filtre yolu)."""
    seeded.add(RetentionSampleCheck(
        sample_id=seeded.query(RetentionSample).filter_by(lot_number="SR001-S").one().id,
        checked_on=date.today(), properties="[]", result="uygun"))
    seeded.query(RetentionSample).filter_by(lot_number="MNR004-S").update({"needs_review": True})
    seeded.commit()
    listed = {r["lot_number"] for r in
              authed_client.get(f"/api/retention/samples?{qs}").json()["rows"]}
    exported = set(_lots(_xlsx(authed_client.get(f"{_URL}?format=xlsx&{qs}"))))
    assert exported == listed


def test_unsized_and_item_filters(authed_client, db_session):
    parent = _item(db_session, "Evanira Losyon")
    small = _item(db_session, "Evanira Losyon 200ml", parent_id=parent.id,
                  variation_name="200ml")
    _row(db_session, small, "EV006", brand="Evanira", shelf="1")
    _row(db_session, parent, "EV002", brand="Evanira", shelf="1")      # boyu belirsiz
    db_session.commit()
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx&unsized=1"))
    assert _lots(wb) == ["EV002"]
    ws = wb["Dolap"]
    head = [c.value for c in ws[1]]
    row = next(ws.iter_rows(min_row=2, values_only=True))
    assert row[head.index("Boy")] == "boy seçilmedi"
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx&item_id={small.id}"))
    assert _lots(wb) == ["EV006"]
    row = next(wb["Dolap"].iter_rows(min_row=2, values_only=True))
    assert row[head.index("Ürün")] == "Evanira Losyon" and row[head.index("Boy")] == "200ml"
    filt = dict(x[:2] for x in wb["Özet"].iter_rows(values_only=True) if x and x[0])
    assert "Ürün: Evanira Losyon 200ml" in filt["Filtre"]


def test_domain_isolation(authed_client, seeded):
    authed_client.post("/api/domain/switch", json={"domain": "supplement"}, headers=_HDR)
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx"))
    assert _lots(wb) == ["SUP001-S"]
    text = _pdf_text(authed_client.get(f"{_URL}?format=pdf"))
    assert "SUP001-S" in text and "MNR010-S" not in text and "SR001-S" not in text


def test_empty_report_still_renders(authed_client, db_session):
    text = _pdf_text(authed_client.get(f"{_URL}?format=pdf"))
    assert "kayıt bulunamadı" in text
    wb = _xlsx(authed_client.get(f"{_URL}?format=xlsx"))
    assert _lots(wb) == []


def test_invalid_format_rejected(authed_client):
    assert authed_client.get(f"{_URL}?format=csv").status_code == 422


# ─── RBAC + sayfa ───────────────────────────────────────────────────────────

def test_labtech_can_export(labtech_client, seeded):
    """Dolabı fiilen laborant kullanıyor — `retention.view` yeterli."""
    assert "SR001-S" in _pdf_text(labtech_client.get(f"{_URL}?format=pdf"))
    assert "SR001-S" in _lots(_xlsx(labtech_client.get(f"{_URL}?format=xlsx")))


def test_anonymous_blocked(client):
    assert client.get(f"{_URL}?format=pdf").status_code in (401, 403)


def test_page_has_report_controls(authed_client):
    html = authed_client.get("/sahit-numune").text
    assert "Rapor al" in html
    assert "/api/retention/samples/export?" in html
    assert "exportReport('pdf')" in html and "exportReport('xlsx')" in html
