"""
Eksik Hammaddeler raporu (core/stock_gaps.py) testleri.

  • classify() saf fonksiyon tablosu
  • Uç nokta: durumlar, min tanımsız anahtarı, kategori negasyonu, pasif kart,
    reçete etkisi, son tedarikçi/hareket, domain izolasyonu
  • Sipariş bayrağı: oluştur/yenile, elle kapat, mal kabul + numune stoğa
    çevirmede otomatik kapanış, domain/kart doğrulama
  • RBAC: LabTech görebilir ama işaretleyemez
  • Dışa aktarma: Excel 3 sayfa, PDF A4 antetli
  • Dashboard zero_stock_raw_count
  • Eski /api/reports/low-stock-alert kaldırıldı
"""
from datetime import date, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from database import Inventory, Item, Recipe, RecipeIngredient, StockOrderFlag, Supplier, SupplierPrice, Transaction

_HDR = {"Origin": "http://testserver"}
_API = "/api/reports/stock-gaps"
_A4_W, _A4_H = 595.276, 841.89


def _switch(client, domain):
    r = client.post("/api/domain/switch", json={"domain": domain}, headers=_HDR)
    assert r.status_code == 200, r.text
    return r


_sku_seq = 0


def _next_sku(prefix):
    global _sku_seq
    _sku_seq += 1
    return f"{prefix}-{_sku_seq:05d}"


def _item(db, name="TEST HAMMADDE", stock=0.0, min_=0.0, unit="g", domain="cosmetics",
         category="Hammadde", supplier_id=None, active=True):
    it = Item(name=name, sku=_next_sku("sg"), category=category, unit=unit,
              current_stock=stock, min_stock_level=min_, domain=domain,
              supplier_id=supplier_id, is_active=active)
    db.add(it); db.flush()
    return it


def _supplier(db, name="TEST TEDARİKÇİ", domain="cosmetics"):
    s = Supplier(name=name, domain=domain)
    db.add(s); db.flush()
    return s


def _sample_lot(db, item_id, qty=5.0, lot="NUM-X", supplier_id=None, domain="cosmetics"):
    inv = Inventory(item_id=item_id, lot_number=lot, quantity=qty, status="APPROVED",
                    is_sample=True, domain=domain, supplier_id=supplier_id)
    db.add(inv); db.flush()
    return inv


def _recipe(db, name="TEST REÇETE", domain="cosmetics", active=True):
    tgt = Item(name=f"{name} ürün", sku=_next_sku("sg-r"), category="Bitmiş Ürün",
              unit="adet", domain=domain)
    db.add(tgt); db.flush()
    rec = Recipe(name=name, output_quantity=1, output_unit="adet", target_item_id=tgt.id,
                domain=domain, is_active=active)
    db.add(rec); db.flush()
    return rec


def _use_in_recipe(db, recipe_id, item_id, qty=1.0):
    db.add(RecipeIngredient(recipe_id=recipe_id, item_id=item_id, quantity=qty, unit="g"))
    db.flush()


def _all_pages_a4(pdf_bytes: bytes) -> bool:
    pages = PdfReader(BytesIO(pdf_bytes)).pages
    assert len(pages) >= 1
    for pg in pages:
        b = pg.mediabox
        if not (abs(float(b.width) - _A4_W) < 2 and abs(float(b.height) - _A4_H) < 2):
            return False
    return True


# ─── classify() saf fonksiyon ────────────────────────────────────────────────

def test_classify_table():
    from core.stock_gaps import classify
    assert classify(0, 0, 0) == "sifir"
    assert classify(0, 5, 0) == "sifir"          # sıfır + min>0 asla 'kritik' değil
    assert classify(0, 0, 3) == "numune"
    assert classify(2, 5, 0) == "kritik"
    assert classify(5, 5, 0) == "kritik"          # eşitlik de kritik
    assert classify(6, 5, 0) is None
    assert classify(3, 0, 0) is None              # min tanımsız + stok var → None (ipucu ayrı sayılır)
    assert classify(-1, 0, 0) == "sifir"


# ─── Uç nokta: durumlar + özet ───────────────────────────────────────────────

def test_statuses_and_summary(authed_client: TestClient, db_session: Session):
    a = _item(db_session, "SIFIR MADDE", stock=0)
    b = _item(db_session, "KRİTİK MADDE", stock=2, min_=5)
    c = _item(db_session, "NUMUNE MADDE", stock=0)
    _sample_lot(db_session, c.id, qty=4)
    d = _item(db_session, "SIFIR NUMUNE MADDE", stock=0)
    _sample_lot(db_session, d.id, qty=0)          # quantity=0 → numune SAYILMAZ
    e = _item(db_session, "YETERLİ MADDE", stock=10, min_=5)
    db_session.commit()

    r = authed_client.get(_API, headers=_HDR)
    assert r.status_code == 200
    d_ = r.json()
    by_name = {row["name"]: row for row in d_["rows"]}
    assert by_name["SIFIR MADDE"]["status"] == "sifir"
    assert by_name["KRİTİK MADDE"]["status"] == "kritik"
    assert by_name["NUMUNE MADDE"]["status"] == "numune"
    assert by_name["SIFIR NUMUNE MADDE"]["status"] == "sifir"     # quantity=0 lot yok sayılır
    assert "YETERLİ MADDE" not in by_name
    assert d_["summary"]["sifir"] == 2
    assert d_["summary"]["kritik"] == 1
    assert d_["summary"]["numune"] == 1


def test_min_undefined_toggle(authed_client: TestClient, db_session: Session):
    _item(db_session, "MİN TANIMSIZ", stock=10, min_=0)
    db_session.commit()

    r = authed_client.get(_API, headers=_HDR)
    d = r.json()
    assert d["summary"]["min_undefined"] == 1
    assert not any(row["name"] == "MİN TANIMSIZ" for row in d["rows"])

    r2 = authed_client.get(f"{_API}?include_undefined_min=1", headers=_HDR)
    d2 = r2.json()
    assert any(row["name"] == "MİN TANIMSIZ" and row["status"] == "min_tanimsiz" for row in d2["rows"])


def test_category_negation(authed_client: TestClient, db_session: Session):
    _item(db_session, "HAM", stock=0, category="Hammadde")
    _item(db_session, "KIM", stock=0, category="Kimyasal")
    _item(db_session, "BOS_KAT", stock=0, category=None)
    _item(db_session, "AMB", stock=0, category="Ambalaj")
    _item(db_session, "BIT", stock=0, category="Bitmiş Ürün")
    db_session.commit()

    ham = {r["name"] for r in authed_client.get(f"{_API}?category=hammadde", headers=_HDR).json()["rows"]}
    assert ham == {"HAM", "KIM", "BOS_KAT"}

    amb = {r["name"] for r in authed_client.get(f"{_API}?category=ambalaj", headers=_HDR).json()["rows"]}
    assert amb == {"AMB"}

    allc = {r["name"] for r in authed_client.get(f"{_API}?category=all", headers=_HDR).json()["rows"]}
    assert allc == {"HAM", "KIM", "BOS_KAT", "AMB", "BIT"}


def test_inactive_items_excluded(authed_client: TestClient, db_session: Session):
    _item(db_session, "PASİF", stock=0, active=False)
    db_session.commit()
    rows = authed_client.get(f"{_API}?category=all", headers=_HDR).json()["rows"]
    assert not any(r["name"] == "PASİF" for r in rows)


def test_recipe_impact(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "REÇETELİ MADDE", stock=0)
    r1 = _recipe(db_session, "B Reçete", active=True)
    r2 = _recipe(db_session, "A Reçete", active=True)
    r3 = _recipe(db_session, "Pasif Reçete", active=False)
    r4 = _recipe(db_session, "Supplement Reçete", domain="supplement", active=True)
    for rec in (r1, r2, r3, r4):
        _use_in_recipe(db_session, rec.id, it.id)
    db_session.commit()

    row = authed_client.get(_API, headers=_HDR).json()["rows"][0]
    assert row["recipe_count"] == 2
    assert [r["name"] for r in row["recipes"]] == ["A Reçete", "B Reçete"]   # tr_key sıralı


def test_last_supplier_and_movement(authed_client: TestClient, db_session: Session):
    sup1 = _supplier(db_session, "Eski Tedarikçi")
    sup2 = _supplier(db_session, "Yeni Tedarikçi")
    it = _item(db_session, "TEDARİKÇİLİ MADDE", stock=0)
    db_session.add(Inventory(item_id=it.id, lot_number="L1", quantity=5, status="APPROVED",
                             domain="cosmetics", supplier_id=sup1.id))
    db_session.flush()
    db_session.add(Inventory(item_id=it.id, lot_number="L2", quantity=3, status="APPROVED",
                             domain="cosmetics", supplier_id=sup2.id))
    db_session.add(Transaction(item_id=it.id, transaction_type="Output", quantity=1, performed_by="test"))
    db_session.commit()

    row = authed_client.get(_API, headers=_HDR).json()["rows"][0]
    assert row["last_supplier"] == "Yeni Tedarikçi"
    assert row["last_movement"] is not None


def test_domain_isolation(authed_client: TestClient, db_session: Session):
    _item(db_session, "SUPP MADDE", stock=0, domain="supplement")
    db_session.commit()
    rows = authed_client.get(f"{_API}?category=all", headers=_HDR).json()["rows"]
    assert not any(r["name"] == "SUPP MADDE" for r in rows)
    _switch(authed_client, "supplement")
    rows2 = authed_client.get(f"{_API}?category=all", headers=_HDR).json()["rows"]
    assert any(r["name"] == "SUPP MADDE" for r in rows2)
    _switch(authed_client, "cosmetics")


def test_search_q(authed_client: TestClient, db_session: Session):
    _item(db_session, "LAVANTA HİDROSOLÜ", stock=0)
    _item(db_session, "PORTAKAL YAĞI", stock=0)
    db_session.commit()
    rows = authed_client.get(f"{_API}?q=lavanta", headers=_HDR).json()["rows"]
    assert len(rows) == 1 and rows[0]["name"] == "LAVANTA HİDROSOLÜ"


# ─── Sipariş bayrağı ─────────────────────────────────────────────────────────

def test_order_flag_create_and_replace(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "SİPARİŞLİ MADDE", stock=0)
    sup = _supplier(db_session, "Sipariş Tedarikçisi")
    db_session.commit()

    r1 = authed_client.post(f"{_API}/{it.id}/order", headers=_HDR,
                            json={"supplier_id": sup.id, "quantity": 10, "note": "ilk"})
    assert r1.status_code == 201, r1.text

    r2 = authed_client.post(f"{_API}/{it.id}/order", headers=_HDR,
                            json={"supplier_id": sup.id, "quantity": 20, "note": "ikinci"})
    assert r2.status_code == 201

    db_session.expire_all()
    flags = db_session.query(StockOrderFlag).filter(StockOrderFlag.item_id == it.id).order_by(StockOrderFlag.id).all()
    assert len(flags) == 2
    assert flags[0].closed_reason == "manual" and flags[0].closed_at is not None
    assert flags[1].closed_at is None

    d = authed_client.get(_API, headers=_HDR).json()
    row = next(r for r in d["rows"] if r["name"] == "SİPARİŞLİ MADDE")
    assert row["order"]["quantity"] == 20 and row["order"]["note"] == "ikinci"
    assert d["summary"]["open_orders"] == 1


def test_order_flag_manual_close(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "KAPATILACAK MADDE", stock=0)
    db_session.commit()
    authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 5})
    r = authed_client.delete(f"{_API}/{it.id}/order", headers=_HDR)
    assert r.status_code == 200
    r2 = authed_client.delete(f"{_API}/{it.id}/order", headers=_HDR)
    assert r2.status_code == 404


def test_order_flag_auto_close_on_receive(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "MAL KABUL MADDE", stock=0, unit="g")
    db_session.commit()
    authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 5})

    r = authed_client.post("/api/inventory/receive", headers=_HDR,
                           json={"item_id": it.id, "lot_number": "L-RCV", "quantity": 5})
    assert r.status_code == 201, r.text

    db_session.expire_all()
    flag = db_session.query(StockOrderFlag).filter(StockOrderFlag.item_id == it.id).one()
    assert flag.closed_reason == "received"


def test_sample_receive_does_not_close_flag(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "NUMUNE KAPANMAZ MADDE", stock=0, unit="g")
    db_session.commit()
    authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 5})

    r = authed_client.post("/api/inventory/receive", headers=_HDR,
                           json={"item_id": it.id, "lot_number": "L-NUM", "quantity": 5, "is_sample": True})
    assert r.status_code == 201, r.text

    db_session.expire_all()
    flag = db_session.query(StockOrderFlag).filter(StockOrderFlag.item_id == it.id).one()
    assert flag.closed_at is None


def test_order_flag_auto_close_on_sample_convert(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "ÇEVRİM MADDE", stock=0, unit="g")
    db_session.commit()
    inv = _sample_lot(db_session, it.id, qty=5, lot="L-CONV")
    db_session.commit()
    authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 5})

    r = authed_client.post(f"/api/inventory/samples/{inv.id}/convert", headers=_HDR)
    assert r.status_code == 200, r.text

    db_session.expire_all()
    flag = db_session.query(StockOrderFlag).filter(StockOrderFlag.item_id == it.id).one()
    assert flag.closed_reason == "received"


def test_order_flag_validation(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "SUPP MADDE 2", stock=0, domain="supplement")
    db_session.commit()
    r = authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 1})
    assert r.status_code == 404       # başka domain kart

    it2 = _item(db_session, "GEÇERLİ MADDE", stock=0)
    db_session.commit()
    r2 = authed_client.post(f"{_API}/{it2.id}/order", headers=_HDR, json={"supplier_id": 999999, "quantity": 1})
    assert r2.status_code == 400      # olmayan tedarikçi


def test_rbac_labtech_view_only(authed_client: TestClient, labtech_client: TestClient, db_session: Session):
    it = _item(db_session, "RBAC MADDE", stock=0)
    db_session.commit()
    assert labtech_client.get(_API).status_code == 200
    assert labtech_client.get(f"{_API}/export?format=xlsx").status_code == 200
    assert labtech_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"quantity": 1}).status_code == 403
    assert labtech_client.delete(f"{_API}/{it.id}/order", headers=_HDR).status_code == 403


# ─── Dışa aktarma ─────────────────────────────────────────────────────────────

def test_export_xlsx_sheets(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "EXCEL MADDE", stock=0, unit="g")
    sup = _supplier(db_session, "Excel Tedarikçi")
    db_session.add(SupplierPrice(item_id=it.id, supplier_id=sup.id, supplier_name=sup.name,
                                 package_size=25, unit_price=120.5, domain="cosmetics"))
    db_session.commit()
    authed_client.post(f"{_API}/{it.id}/order", headers=_HDR, json={"supplier_id": sup.id, "quantity": 10})

    r = authed_client.get(f"{_API}/export?format=xlsx", headers=_HDR)
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Eksikler", "Sipariş Listesi", "Reçete Etkisi"]
    ws = wb["Sipariş Listesi"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "Excel Tedarikçi" in values
    assert 120.5 in values


def test_export_pdf_is_a4_letterhead(authed_client: TestClient, db_session: Session):
    _item(db_session, "PDF MADDE", stock=0)
    db_session.commit()
    r = authed_client.get(f"{_API}/export?format=pdf", headers=_HDR)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    r.headers["content-disposition"].encode("latin-1")
    assert _all_pages_a4(r.content)


# ─── Dashboard ────────────────────────────────────────────────────────────────

def test_dashboard_zero_stock_raw_count(authed_client: TestClient, db_session: Session):
    _item(db_session, "DASH SIFIR", stock=0, category="Hammadde")
    _item(db_session, "DASH AMBALAJ", stock=0, category="Ambalaj")   # hariç
    _item(db_session, "DASH KRİTİK", stock=2, min_=5)                # kritik ayrı sayaç, burada sayılmaz
    db_session.commit()
    d = authed_client.get("/api/dashboard/stats", headers=_HDR).json()
    assert d["zero_stock_raw_count"] == 1
    assert d["critical_stock_count"] == 1


# ─── Eski uç nokta kaldırıldı ─────────────────────────────────────────────────

def test_low_stock_alert_removed(authed_client: TestClient):
    r = authed_client.get("/api/reports/low-stock-alert", headers=_HDR)
    assert r.status_code == 404


def test_reports_page_has_eksik_section(authed_client: TestClient):
    r = authed_client.get("/reports")
    assert r.status_code == 200
    assert 'id="eksik"' in r.text


# ─── Seçim modu (satır bazlı dışa aktarma) ────────────────────────────────────

def test_assemble_item_ids_ignores_filters(authed_client: TestClient, db_session: Session):
    """Seçim modu: kategori/durum süzgeçleri yok sayılır — hammadde + ambalaj
    birlikte, hatta stoğu yeterli olan kalem bile seçildiyse satır olarak döner."""
    from core.stock_gaps import assemble
    ham = _item(db_session, "UÇUCU YAĞ A", stock=0, category="Hammadde")
    amb = _item(db_session, "ŞİŞE 200 ML", stock=0, category="Ambalaj")
    dolu = _item(db_session, "STOKLU MADDE", stock=50, min_=5, category="Hammadde")
    _item(db_session, "SEÇİLMEYEN", stock=0, category="Hammadde")
    db_session.commit()

    rep = assemble(db_session, "cosmetics", category="hammadde",
                   statuses=["sifir"], item_ids=[ham.id, amb.id, dolu.id])
    by_name = {r["name"]: r for r in rep["rows"]}
    assert set(by_name) == {"UÇUCU YAĞ A", "ŞİŞE 200 ML", "STOKLU MADDE"}
    assert by_name["ŞİŞE 200 ML"]["status"] == "sifir"          # kategori süzgeci yok sayıldı
    assert by_name["STOKLU MADDE"]["status"] == "yeterli"        # seçilen satır düşmez
    assert by_name["STOKLU MADDE"]["status_label"] == "Yeterli"


def test_export_with_ids(authed_client: TestClient, db_session: Session):
    ham = _item(db_session, "UÇUCU YAĞ B", stock=0, category="Hammadde")
    amb = _item(db_session, "KAVANOZ 50 ML", stock=0, category="Ambalaj")
    _item(db_session, "DIŞARIDA KALAN", stock=0, category="Hammadde")
    db_session.commit()

    r = authed_client.get(f"{_API}/export?format=xlsx&ids={ham.id},{amb.id}", headers=_HDR)
    assert r.status_code == 200
    assert "secim" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    names = [c.value for row in wb["Eksikler"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "UÇUCU YAĞ B" in names and "KAVANOZ 50 ML" in names
    assert "DIŞARIDA KALAN" not in names


def test_export_ids_pdf_and_validation(authed_client: TestClient, db_session: Session):
    it = _item(db_session, "PDF SEÇİM", stock=0)
    db_session.commit()
    r = authed_client.get(f"{_API}/export?format=pdf&ids={it.id}", headers=_HDR)
    assert r.status_code == 200 and _all_pages_a4(r.content)
    assert authed_client.get(f"{_API}/export?format=xlsx&ids=abc", headers=_HDR).status_code == 400
    assert authed_client.get(f"{_API}/export?format=xlsx&ids=,", headers=_HDR).status_code == 400


def test_export_ids_domain_isolation(authed_client: TestClient, db_session: Session):
    supp = _item(db_session, "SUPP SEÇİM", stock=0, domain="supplement")
    ok = _item(db_session, "KOZMETİK SEÇİM", stock=0)
    db_session.commit()
    r = authed_client.get(f"{_API}/export?format=xlsx&ids={supp.id},{ok.id}", headers=_HDR)
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    names = [c.value for row in wb["Eksikler"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "KOZMETİK SEÇİM" in names and "SUPP SEÇİM" not in names
