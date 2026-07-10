"""
CRM Faz C1 — Excel/CSV içe/dışa aktarma testleri.

  • export: firma/kişi/fırsat → xlsx (başlık + veri)
  • import: xlsx'ten firma oluştur (source='import'); kişiyi firmaya eşle
  • şablon indir; RBAC (import crm.create ister)
"""
import io

from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

_H = {"Origin": "http://testserver"}
_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(headers, rows) -> bytes:
    wb = Workbook(); ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b)
    return b.getvalue()


def test_export_companies_xlsx(authed_client: TestClient):
    authed_client.post("/api/crm/companies", json={"name": "Export Co", "city": "İzmir"}, headers=_H)
    r = authed_client.get("/api/crm/export?entity=companies")
    assert r.status_code == 200
    assert "spreadsheet" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert "Firma" in headers and "Kaynak" in headers
    names = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert "Export Co" in names


def test_import_companies_creates_rows(authed_client: TestClient):
    data = _xlsx(["Firma", "Telefon", "Şehir"],
                 [["İçe Firma A", "0555 111 22 33", "Bursa"], ["İçe Firma B", "", "Ankara"]])
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("c.xlsx", data, _MIME)}, headers=_H)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 2
    # 'import' artık kendi filtresi — Elle (manual) kovasına düşmez
    rows = authed_client.get("/api/crm/companies?source=import").json()
    imported = [c for c in rows if c["name"].startswith("İçe Firma")]
    assert len(imported) == 2
    assert all(c["source"] == "import" for c in imported)
    manual = authed_client.get("/api/crm/companies?source=manual").json()
    assert not [c for c in manual if c["name"].startswith("İçe Firma")]


def test_import_contacts_links_company(authed_client: TestClient):
    co = authed_client.post("/api/crm/companies", json={"name": "Bağlı Firma"}, headers=_H).json()
    data = _xlsx(["Ad Soyad", "Firma", "Telefon"], [["Veli İçe", "Bağlı Firma", "0500 000 00 00"]])
    r = authed_client.post("/api/crm/import", data={"entity": "contacts"},
                           files={"file": ("k.xlsx", data, _MIME)}, headers=_H)
    assert r.status_code == 200 and r.json()["created"] == 1
    ct = authed_client.get("/api/crm/contacts?q=Veli").json()
    assert ct and ct[0]["company_id"] == co["id"]


def test_import_csv_companies(authed_client: TestClient):
    csv_bytes = "Firma;Şehir\nCSV Firma;Konya\n".encode("utf-8")
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("c.csv", csv_bytes, "text/csv")}, headers=_H)
    assert r.status_code == 200 and r.json()["created"] == 1


def test_import_unknown_headers_rejected(authed_client: TestClient):
    data = _xlsx(["Kolon1", "Kolon2"], [["a", "b"]])
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("x.xlsx", data, _MIME)}, headers=_H)
    assert r.status_code == 400


def test_import_requires_create(labtech_client: TestClient):
    data = _xlsx(["Firma"], [["X"]])
    r = labtech_client.post("/api/crm/import", data={"entity": "companies"},
                            files={"file": ("c.xlsx", data, _MIME)}, headers=_H)
    assert r.status_code == 403


def test_template_download(authed_client: TestClient):
    r = authed_client.get("/api/crm/import/template?entity=contacts")
    assert r.status_code == 200 and "spreadsheet" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert "Ad Soyad" in headers


# ─── İçe aktarmada eşleştir-güncelle (upsert) ────────────────────────────────

def test_import_companies_upserts_by_folded_name(authed_client: TestClient):
    co = authed_client.post("/api/crm/companies",
                            json={"name": "Işık Kozmetik", "city": "İstanbul"}, headers=_H).json()
    # Türkçe-katlanmış ad eşleşir (küçük harf, diakritiksiz) → güncelleme, yeni kayıt yok
    data = _xlsx(["Firma", "Telefon"], [["isik kozmetik", "0212 999 88 77"]])
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("c.xlsx", data, _MIME)}, headers=_H)
    assert r.status_code == 200
    assert r.json()["created"] == 0 and r.json()["updated"] == 1
    rows = [c for c in authed_client.get("/api/crm/companies?q=Kozmetik").json()
            if "Kozmetik" in c["name"]]
    assert len(rows) == 1 and rows[0]["id"] == co["id"]
    assert rows[0]["phone"] == "0212 999 88 77"       # gelen alan güncellendi
    assert rows[0]["city"] == "İstanbul"              # gelmeyen alan korundu
    assert rows[0]["source"] == "manual"              # source'a dokunulmadı


def test_import_companies_matches_by_tax_no(authed_client: TestClient):
    authed_client.post("/api/crm/companies",
                       json={"name": "Eski Unvan A.Ş.", "tax_no": "1234567890"}, headers=_H)
    data = _xlsx(["Firma", "Vergi No", "Şehir"], [["Bambaşka Ad Ltd.", "1234567890", "İzmir"]])
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("c.xlsx", data, _MIME)}, headers=_H)
    assert r.json()["updated"] == 1 and r.json()["created"] == 0
    rows = authed_client.get("/api/crm/companies?q=Eski Unvan").json()
    assert rows and rows[0]["city"] == "İzmir"        # vergi no eşleşti, ad değişmedi


def test_import_contacts_upserts_by_email_then_phone(authed_client: TestClient):
    authed_client.post("/api/crm/contacts",
                       json={"full_name": "Posta Kişisi", "email": "posta@x.com"}, headers=_H)
    authed_client.post("/api/crm/contacts",
                       json={"full_name": "Telefon Kişisi", "mobile": "+90 555 111 22 33"}, headers=_H)
    data = _xlsx(["Ad Soyad", "E-posta", "Unvan"], [["Posta Kişisi", "posta@x.com", "Müdür"]])
    r1 = authed_client.post("/api/crm/import", data={"entity": "contacts"},
                            files={"file": ("k.xlsx", data, _MIME)}, headers=_H)
    assert r1.json()["updated"] == 1 and r1.json()["created"] == 0
    ct = authed_client.get("/api/crm/contacts?q=Posta").json()[0]
    assert ct["title"] == "Müdür"
    # telefon format farkı eşleşmeyi bozmaz (son 10 hane)
    data2 = _xlsx(["Ad Soyad", "Telefon", "Unvan"], [["Telefon Kişisi", "05551112233", "Uzman"]])
    r2 = authed_client.post("/api/crm/import", data={"entity": "contacts"},
                            files={"file": ("k2.xlsx", data2, _MIME)}, headers=_H)
    assert r2.json()["updated"] == 1 and r2.json()["created"] == 0
    ct2 = [c for c in authed_client.get("/api/crm/contacts?q=Telefon Kişisi").json()]
    assert len(ct2) == 1 and ct2[0]["title"] == "Uzman"


def test_import_new_rows_get_source_import_and_counts(authed_client: TestClient):
    authed_client.post("/api/crm/companies", json={"name": "Var Olan"}, headers=_H)
    data = _xlsx(["Firma", "Şehir"], [["Var Olan", "Adana"], ["Yepyeni Firma", "Mersin"]])
    r = authed_client.post("/api/crm/import", data={"entity": "companies"},
                           files={"file": ("c.xlsx", data, _MIME)}, headers=_H)
    body = r.json()
    assert body["created"] == 1 and body["updated"] == 1 and body["total"] == 2
    yeni = authed_client.get("/api/crm/companies?q=Yepyeni").json()[0]
    assert yeni["source"] == "import"
