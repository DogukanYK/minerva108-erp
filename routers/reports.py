# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Reports router — dashboard stats + reporting data feeds.
"""
import calendar
import datetime as _dt
import io
import re
from typing import Optional, List

from fastapi import APIRouter, Depends, Query, Request, UploadFile, File, Form
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import case, func
from sqlalchemy.orm import Session

from database import (
    to_tr,
    get_db, Item, Supplier, Recipe, Transaction, ProductionHistory, Inventory,
    StockSnapshot, SupplierPrice, StockOrderFlag,
)
from core.permissions import require_permission
from core.auth import require_role
from core.snapshots import compute_stock_at, snapshot_exists
from core.domain import active_domain
from core.audit import log_admin_event

router = APIRouter(prefix="/api", tags=["reports"])


# ─── Aylık Stok Raporu — kategori bazlı, ay sonu mevcut stok ───────────────

# UI'da görünen 4 kategori → DB'de Item.category + pkg_type kombinasyonu
def _category_filter(category_key: str):
    """Reports UI'daki kategori anahtarını Item filtresine çevir."""
    k = (category_key or "all").lower()
    if k == "hammadde":
        return (Item.category == "Hammadde",)
    if k == "ambalaj":   # etiket hariç
        return (Item.category == "Ambalaj",
                (Item.pkg_type.is_(None)) | (Item.pkg_type != "etiket"))
    if k == "etiket":
        return (Item.category == "Ambalaj", Item.pkg_type == "etiket")
    if k == "finished":
        return (Item.category == "Bitmiş Ürün",)
    return ()


def _month_window(year: int, month: int) -> tuple[_dt.datetime, _dt.datetime]:
    """Verilen yıl/ay için (ayın 1'i 00:00, ayın son günü 23:59:59.999999)."""
    if not (1 <= month <= 12) or year < 2000 or year > 2100:
        raise ValueError("Geçersiz yıl/ay")
    first = _dt.datetime(year, month, 1, 0, 0, 0)
    last_day = calendar.monthrange(year, month)[1]
    last  = _dt.datetime(year, month, last_day, 23, 59, 59, 999999)
    return first, last


def _snapshot_category_filter(category_key: str):
    """`_category_filter`'ın StockSnapshot satırı için karşılığı."""
    k = (category_key or "all").lower()
    if k == "hammadde":
        return (StockSnapshot.category == "Hammadde",)
    if k == "ambalaj":   # etiket hariç
        return (StockSnapshot.category == "Ambalaj",
                (StockSnapshot.pkg_type.is_(None)) | (StockSnapshot.pkg_type != "etiket"))
    if k == "etiket":
        return (StockSnapshot.category == "Ambalaj", StockSnapshot.pkg_type == "etiket")
    if k == "finished":
        return (StockSnapshot.category == "Bitmiş Ürün",)
    return ()


def _rows_from_snapshot(db: Session, year: int, month: int, category_key: str,
                        domain: str = "cosmetics") -> list[dict]:
    """Dondurulmuş `stock_snapshot` satırlarından kategori bazlı rapor satırları."""
    rows = (
        db.query(StockSnapshot)
        .filter(StockSnapshot.year == year, StockSnapshot.month == month,
                StockSnapshot.domain == domain,
                *_snapshot_category_filter(category_key))
        .all()
    )
    out = []
    for s in rows:
        stock = float(s.stock or 0.0)
        if abs(stock) < 1e-9:
            continue
        out.append({
            "item_id":  s.item_id,
            "name":     s.item_name or "—",
            "unit":     s.unit or "",
            "pkg_type": s.pkg_type or "",
            "language": "",
            "stock":    round(stock, 4),
        })
    out.sort(key=lambda r: (r["name"] or "").lower())
    return out


def _rows_live(db: Session, eom: _dt.datetime, category_key: str,
               domain: str = "cosmetics") -> list[dict]:
    """Snapshot yoksa: `core.snapshots.compute_stock_at` ile canlı rekonstrüksiyon."""
    stocks, items = compute_stock_at(db, eom, Item.domain == domain, *_category_filter(category_key))
    out = []
    for it in items:
        stock = float(stocks.get(it.id, 0.0))
        if abs(stock) < 1e-9:
            continue
        out.append({
            "item_id":  it.id,
            "name":     it.name,
            "unit":     it.unit or "",
            "pkg_type": it.pkg_type or "",
            "language": it.language or "",
            "stock":    round(stock, 4),
        })
    out.sort(key=lambda r: r["name"].lower())
    return out


def _attach_groups(db: Session, sections: list, domain: str) -> bool:
    """Aylık rapor satırlarına kartın BUGÜNKÜ "aynı malzeme" grubunu ekle
    (13.10.2026).  Yalnız görünüm — bölüm toplamları DEĞİŞMEZ.

    Bir bölümde aynı aktif gruptan ≥2 satır varsa `sec["groups"]`'a grup
    girdisi (taban birimde toplam + kart dökümü, sunucu biçimli) eklenir,
    satırlar `group_id` alır ve üyeler ilk üyenin yerinde ARDIŞIK dizilir
    (ekran ve Excel aynı sırayı okur).  Toplam o ayın satır stoklarından
    `core.group_stock` kuralıyla (g↔kg, ml↔l; başka aile `aside_text`).
    Dondurulmuş eski aylarda grup bilgisi yoktur — bugünkü gruplama
    uygulanır (ekranda dipnot).  Döner: herhangi bir bölümde grup var mı."""
    from core import group_stock as GS
    from database import MaterialGroup

    ids = {r["item_id"] for sec in sections for r in sec["rows"] if r.get("item_id")}
    info = {}
    if ids:
        q = (db.query(Item.id, Item.material_group_id, MaterialGroup.name, Supplier.name)
             .join(MaterialGroup, MaterialGroup.id == Item.material_group_id)
             .outerjoin(Supplier, Supplier.id == Item.supplier_id)
             .filter(Item.id.in_(ids), Item.domain == domain,
                     MaterialGroup.domain == domain, MaterialGroup.is_active == True))  # noqa: E712
        info = {iid: (gid, gname or f"Grup #{gid}", sname) for iid, gid, gname, sname in q.all()}
    any_group = False
    for sec in sections:
        by_gid: dict = {}
        for r in sec["rows"]:
            r["group_id"] = None
            inf = info.get(r.get("item_id"))
            if inf:
                by_gid.setdefault(inf[0], []).append(r)
        groups = {}
        for gid, rws in by_gid.items():
            if len(rws) < 2:
                continue
            members = [GS.GroupMember(
                id=r["item_id"], name=r["name"], unit=r["unit"], stock=max(0.0, float(r["stock"])),
                raw_stock=float(r["stock"]), min_stock=0.0, category="", pkg_type=r.get("pkg_type") or None,
                supplier_id=None, supplier_name=info[r["item_id"]][2]) for r in rws]
            g = GS.build_group(gid, info[rws[0]["item_id"]][1], sec["category"], members)
            for r in rws:
                r["group_id"] = gid
            groups[gid] = {
                "group_id": gid, "name": g.name, "base_unit": g.base_unit,
                "total": round(g.total, 4), "total_text": g.total_text,
                "breakdown_text": g.breakdown_text, "aside_text": g.aside_text,
                "member_count": len(rws),
            }
        # Üyeler ilk üyenin yerinde ardışık — ad sırası bölüm içinde korunur
        ordered, done = [], set()
        for r in sec["rows"]:
            gid = r["group_id"]
            if gid is None:
                ordered.append(r)
            elif gid not in done:
                done.add(gid)
                ordered.extend(x for x in sec["rows"] if x["group_id"] == gid)
        sec["rows"] = ordered
        sec["groups"] = [groups[gid] for gid in
                         dict.fromkeys(r["group_id"] for r in ordered if r["group_id"] is not None)]
        any_group = any_group or bool(groups)
    return any_group


@router.get("/reports/monthly-stock")
def report_monthly_stock(
    year:     int = Query(..., ge=2000, le=2100),
    month:    int = Query(..., ge=1,    le=12),
    category: str = Query("all", regex="^(all|hammadde|ambalaj|etiket|finished)$"),
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """
    Verilen ay sonunda mevcut stok dökümü.  Tarihçe rekonstrüksiyonu —
    Transaction'lardaki signed delta'ların kümülatif toplamı.  Aktif panele
    (domain) göre süzülür.  Kartın bugünkü "aynı malzeme" grubu her zaman
    eklenir (`_attach_groups`: `groups[]`, satırda `group_id`, `group_note`);
    bölüm toplamları değişmez.

    category: all | hammadde | ambalaj | etiket | finished
    """
    try:
        _, eom_raw = _month_window(year, month)
    except ValueError as e:
        return JSONResponse(status_code=400, content={"detail": str(e)})

    # Veri kesim tarihi = ayın sonu VEYA şimdi (hangisi önce).  Mevcut ay
    # seçilince ayın sonu henüz GELMEMİŞTİR; gelecekteki bir tarihe kadar
    # rekonstrüksiyon anlamsız — bugüne kadar olan stok gösterilir.
    now = _dt.datetime.utcnow()
    eom = min(eom_raw, now)

    sections = (
        [("hammadde", "Hammadde"),
         ("ambalaj",  "Ambalaj"),
         ("etiket",   "Etiket"),
         ("finished", "Bitmiş Ürün")]
        if category == "all" else
        [(category, {"hammadde": "Hammadde", "ambalaj": "Ambalaj",
                     "etiket": "Etiket", "finished": "Bitmiş Ürün"}[category])]
    )

    # Geçmiş bir ay için dondurulmuş snapshot varsa onu kullan — kesin ve
    # transaction değişikliklerinden etkilenmez.  Yoksa canlı rekonstrüksiyon.
    use_snapshot = snapshot_exists(db, year, month, domain)

    result = []
    for key, label in sections:
        rows = (_rows_from_snapshot(db, year, month, key, domain)
                if use_snapshot else
                _rows_live(db, eom, key, domain))
        result.append({
            "category": key,
            "label":    label,
            "rows":     rows,
            "row_count": len(rows),
            "total":    round(sum(r["stock"] for r in rows), 4),
        })

    # Kartın BUGÜNKÜ "aynı malzeme" grubu (her zaman — yalnız görünüm)
    has_groups = _attach_groups(db, result, domain)
    group_note = ""
    if has_groups:
        group_note = ("Aynı malzemenin tedarikçi kartları bugünkü «aynı malzeme» gruplarına göre "
                      "toplandı (g↔kg, ml↔l çevrilir; başka birim ayrı yazılır). Bölüm toplamları değişmez.")
        if use_snapshot:
            group_note += (" Dondurulmuş kayıtta grup bilgisi yoktur — eski ay bugünkü gruplamayla "
                           "gösterilir.")

    return {
        "year":         year,
        "month":        month,
        "group_note":   group_note,
        # Verinin kaynağı: dondurulmuş snapshot mu, canlı hesap mı?
        "source":       "snapshot" if use_snapshot else "canlı",
        # Raporun fiilen oluşturulduğu an — "Rapor tarihi" olarak gösterilir.
        "generated_at": to_tr(now).strftime("%d.%m.%Y %H:%M"),
        # Stok verisinin hangi ana kadar olduğu (ay sonu ya da bugün).
        "as_of":        to_tr(eom).strftime("%d.%m.%Y %H:%M"),
        "sections":     result,
        "category":     category,
    }


@router.get("/reports/monthly-stock/export")
def report_monthly_stock_export(
    year:     int = Query(..., ge=2000, le=2100),
    month:    int = Query(..., ge=1,    le=12),
    category: str = Query("all", regex="^(all|hammadde|ambalaj|etiket|finished)$"),
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Aylık stok raporu .xlsx — her kategori ayrı sheet."""
    data = report_monthly_stock(year, month, category, db, _, domain)
    if isinstance(data, JSONResponse):
        return data

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    bold  = Font(bold=True, color="FFFFFF")
    fill  = PatternFill("solid", fgColor="232E6E")
    box   = Border(*([Side(style="thin", color="999999")] * 4))
    title_font = Font(bold=True, size=12)

    month_label = ["", "Ocak","Şubat","Mart","Nisan","Mayıs","Haziran",
                   "Temmuz","Ağustos","Eylül","Ekim","Kasım","Aralık"][month]

    for sec in data["sections"]:
        ws = wb.create_sheet(title=sec["label"][:31])
        ws["A1"] = f"{sec['label']} — {month_label} {year} stok durumu"
        ws["A1"].font = title_font
        ws["A2"] = (f"Rapor tarihi: {data['generated_at']}  ·  "
                    f"Stok kesim: {data['as_of']} itibarıyla  ·  "
                    f"{sec['row_count']} kalem  ·  toplam {sec['total']}")
        ws["A2"].font = Font(italic=True, color="6b7280")

        if data.get("group_note"):
            ws["A3"] = data["group_note"]
            ws["A3"].font = Font(italic=True, color="6b7280", size=9)

        headers = ["#", "Ürün Adı", "Birim", "Alt-Tip", "Dil", "Stok", "Grup dökümü"]
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=4, column=ci, value=h)
            c.font = bold; c.fill = fill; c.border = box
            c.alignment = Alignment(horizontal="center")

        # Grup satırı (kalın) + girintili üyeler.  Grup toplamı F'de METİN —
        # SUM(F) kartları iki kez saymasın; bölüm TOPLAM'ı değişmez.
        groups = {g["group_id"]: g for g in sec.get("groups") or []}
        ri, n, done = 5, 0, set()
        bold_plain = Font(bold=True)
        for r in sec["rows"]:
            gid = r.get("group_id")
            if gid in groups and gid not in done:
                done.add(gid)
                g = groups[gid]
                dok = g["breakdown_text"] + (f" · ayrı birim: {g['aside_text']}" if g["aside_text"] else "")
                vals = ["", f"{g['name']} (aynı malzeme, {g['member_count']} kart)", g["base_unit"],
                        "", "", g["total_text"], dok]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=ri, column=ci, value=v)
                    c.border = box; c.font = bold_plain
                ws.cell(row=ri, column=6).alignment = Alignment(horizontal="right")
                ri += 1
            n += 1
            ws.cell(row=ri, column=1, value=n).border = box
            name_c = ws.cell(row=ri, column=2, value=r["name"])
            name_c.border = box
            if gid in groups:
                name_c.alignment = Alignment(indent=2)
            ws.cell(row=ri, column=3, value=r["unit"]).border = box
            ws.cell(row=ri, column=4, value=r["pkg_type"]).border = box
            ws.cell(row=ri, column=5, value=r["language"]).border = box
            ws.cell(row=ri, column=6, value=r["stock"]).border = box
            ws.cell(row=ri, column=7, value="").border = box
            ri += 1

        # TOPLAM satırı
        total_row = ri
        ws.cell(row=total_row, column=2, value="TOPLAM").font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=sec["total"]).font = Font(bold=True)

        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 14
        ws.column_dimensions["G"].width = 60

    if not wb.worksheets:
        wb.create_sheet(title="Boş")["A1"] = "Bu ay için kayıt yok."

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"stok_raporu_{year}-{month:02d}_{category}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ─── Reports Endpoints ───────────────────────────────────────────────────────

@router.get("/reports/top-usage")
def report_top_usage(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    import datetime as _dt
    since = _dt.datetime.utcnow() - _dt.timedelta(days=30)
    return usage_between(db, domain, since)[:5]


def usage_between(db: Session, domain: str, since, until=None) -> list:
    """Dönem tüketimi (UTC `since` ≤ t < `until`), çoktan aza — top-usage ucu ve
    Minerva Asistan `consumption` aracı AYNI kuralı kullanır: Output toplamı;
    iptal edilen üretimin tüketimi ve fason SEVK Output'ları (transfer) hariç,
    fason fiilî tüketim + fire (`movement_totals`) dahil."""
    from sqlalchemy import func
    from core.production_cancel import cancelled_tx_ids_subq
    from core.outsourcing_reporting import transfer_transaction_ids, movement_totals
    filters = [
        Transaction.transaction_type == "Output",
        Transaction.timestamp >= since,
        Item.domain == domain,
        # İptal edilen üretimin tüketimi gerçek kullanım değil
        ~Transaction.id.in_(cancelled_tx_ids_subq(db)),
        ~Transaction.id.in_(transfer_transaction_ids(db)),
    ]
    if until is not None:
        filters.append(Transaction.timestamp < until)
    rows = (
        db.query(
            Item.id,
            Item.name,
            Item.unit,
            func.sum(Transaction.quantity).label("total_used"),
        )
        .join(Transaction, Transaction.item_id == Item.id)
        .filter(*filters)
        .group_by(Item.id)
        .order_by(func.sum(Transaction.quantity).desc())
        .all()
    )
    usage = {
        r.id: {"item_id": r.id, "name": r.name, "unit": r.unit, "total_used": float(r.total_used)}
        for r in rows
    }
    outside = movement_totals(db, since, until, domain=domain)
    if outside:
        for item in db.query(Item).filter(Item.id.in_(outside), Item.domain == domain).all():
            row = usage.setdefault(item.id, {"item_id": item.id, "name": item.name, "unit": item.unit, "total_used": 0.0})
            row["total_used"] += outside[item.id]["consumption"] + outside[item.id]["waste"]
    for row in usage.values():
        row["total_used"] = round(row["total_used"], 4)
    return sorted(usage.values(), key=lambda row: row["total_used"], reverse=True)


@router.get("/reports/production-trends")
def report_production_trends(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    import datetime as _dt
    from sqlalchemy import func
    today = _dt.date.today()
    since = _dt.datetime.combine(today - _dt.timedelta(days=6), _dt.time.min)
    rows = (db.query(ProductionHistory)
            .filter(ProductionHistory.produced_at >= since,
                    ProductionHistory.domain == domain,
                    ProductionHistory.cancelled_at.is_(None)).all())
    # Gün gün topla
    day_map = {}
    for i in range(7):
        d = (today - _dt.timedelta(days=6 - i)).isoformat()
        day_map[d] = {"date": d, "total_quantity": 0.0, "count": 0}
    for r in rows:
        d = r.produced_at.date().isoformat()
        if d in day_map:
            day_map[d]["total_quantity"] = round(day_map[d]["total_quantity"] + r.produced_quantity, 4)
            day_map[d]["count"] += 1
    return list(day_map.values())


class _OrderFlagIn(BaseModel):
    supplier_id: Optional[int] = None
    quantity: Optional[float] = Field(None, gt=0)
    expected_date: Optional[_dt.date] = None
    note: Optional[str] = Field(None, max_length=300)


@router.get("/reports/stock-gaps")
def stock_gaps_report(
    category: str = Query("hammadde", pattern="^(hammadde|ambalaj|all)$"),
    status: Optional[str] = None,
    include_undefined_min: int = Query(0, ge=0, le=1),
    q: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Eksik Hammaddeler — sıfır / kritik altı / sadece numune hammaddeler.

    core/stock_gaps.py TEK KAYNAK. `current_stock<=0 and min_stock_level=0`
    (min. tanımsız) hiçbir alarma girmiyordu — `include_undefined_min=1` ile
    görünür kılınır.  Üretimde kaynak seçimi açıkken (`group_stock_enabled`)
    "aynı malzeme" grubu toplamla değerlendirilir: eksik grup tek satır
    (`kind="group"`, `members[]`), grubu yeterli kart düşer."""
    from core.stock_gaps import assemble
    statuses = [s.strip() for s in status.split(",") if s.strip()] if status else None
    return assemble(db, domain, category=category, statuses=statuses,
                    include_undefined_min=bool(include_undefined_min), q=q)


@router.post("/reports/stock-gaps/{item_id}/order", status_code=201)
def stock_gaps_order(
    item_id: int,
    data: _OrderFlagIn,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "adjust")),
    domain: str = Depends(active_domain),
):
    """'Sipariş verildi' işaretle — ürün başına en fazla bir açık bayrak
    (yeniden işaretleme eskisini 'manual' kapatır). Gerçek mal kabulünde
    (routers/inventory.py receive_stock / samples/convert) otomatik kapanır."""
    from core.stock_gaps import open_flag
    item = db.query(Item).filter(Item.id == item_id).with_for_update().first()
    if not item or not item.is_active or item.domain != domain:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})
    if data.supplier_id is not None:
        sup = db.query(Supplier).filter(Supplier.id == data.supplier_id).first()
        if not sup or sup.domain != domain:
            return JSONResponse(status_code=400, content={"detail": "Tedarikçi bu panelde değil."})
    actor = current_user.get("full_name") or current_user.get("username") or "—"
    flag = open_flag(db, item, supplier_id=data.supplier_id, quantity=data.quantity,
                     expected_date=data.expected_date, note=data.note, actor=actor, domain=domain)
    log_admin_event(db, request, actor=current_user, action="stock_gap.order",
                    target_type="item", target_id=item.id, target_name=item.name,
                    details={"supplier_id": data.supplier_id, "quantity": data.quantity,
                             "expected_date": str(data.expected_date) if data.expected_date else None})
    db.commit()
    return {"id": flag.id, "message": "Sipariş işaretlendi."}


@router.delete("/reports/stock-gaps/{item_id}/order")
def stock_gaps_order_cancel(
    item_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "adjust")),
    domain: str = Depends(active_domain),
):
    from core.stock_gaps import close_open_flags
    item = db.query(Item).filter(Item.id == item_id, Item.domain == domain).first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})
    actor = current_user.get("full_name") or current_user.get("username") or "—"
    closed = close_open_flags(db, item_id, "manual", actor)
    if not closed:
        return JSONResponse(status_code=404, content={"detail": "Açık sipariş bulunamadı."})
    log_admin_event(db, request, actor=current_user, action="stock_gap.order_cancel",
                    target_type="item", target_id=item.id, target_name=item.name)
    db.commit()
    return {"message": "Sipariş işareti kaldırıldı."}


@router.get("/reports/stock-gaps/export")
def stock_gaps_export(
    format: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    category: str = Query("hammadde", pattern="^(hammadde|ambalaj|all)$"),
    status: Optional[str] = None,
    include_undefined_min: int = Query(0, ge=0, le=1),
    q: Optional[str] = None,
    ids: Optional[str] = None,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Ekrandaki süzgeçlerle ya da SEÇİLEN kartlarla (`ids`) dışa aktar.

    `ids` verilirse kategori/durum/arama yok sayılır — lab yalnız istediği
    uçucu yağları, yalnız istediği ambalajları ya da ikisini birlikte tek
    Excel'de alabilsin diye (10.09.2026 talebi)."""
    from core.stock_gaps import assemble, build_workbook, flat_rows, render_pdf, export_filename
    from core.supplier_prices import prices_for_items, today_rates_or_none
    from core.delivery_note import content_disposition
    statuses = [s.strip() for s in status.split(",") if s.strip()] if status else None
    picked = None
    if ids:
        try:
            picked = [int(x) for x in ids.split(",") if x.strip()][:2000]
        except ValueError:
            return JSONResponse(status_code=400, content={"detail": "Geçersiz ürün seçimi."})
        if not picked:
            return JSONResponse(status_code=400, content={"detail": "Seçili ürün yok."})
    report = assemble(db, domain, category=category, statuses=statuses,
                      include_undefined_min=bool(include_undefined_min), q=q,
                      item_ids=picked)
    try:
        if format == "pdf":
            content = render_pdf(report)
            media = "application/pdf"
        else:
            # bayrak KART başına — grup satırının siparişli üyeleri dahil
            order_ids = [r["item_id"] for r in flat_rows(report) if r.get("order")]
            # para birimi karışık malzemede bugünün kuruyla sıralı (kur tembel)
            prices = prices_for_items(db, order_ids, domain, rates=today_rates_or_none)
            content = build_workbook(report, prices=prices)
            media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    except Exception:
        return JSONResponse(status_code=500, content={"detail": "Rapor üretilemedi."})
    fname = export_filename("secim" if picked else category, format)
    return StreamingResponse(
        io.BytesIO(content), media_type=media,
        headers={"Content-Disposition": content_disposition(fname, inline=(format == "pdf"))},
    )


# ─── Dashboard Stats ─────────────────────────────────────────────────────────

@router.get("/dashboard/stats")
def dashboard_stats(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    total_items     = db.query(Item).filter(Item.is_active == True, Item.domain == domain).count()
    total_suppliers = db.query(Supplier).filter(Supplier.is_active == True, Supplier.domain == domain).count()
    total_recipes   = db.query(Recipe).filter(Recipe.is_active == True, Recipe.domain == domain).count()
    # "Aynı malzeme" grup kuralı — yalnız üretimdeki kaynak seçimi açıkken
    # (core.group_stock TEK ANAHTAR; kapalı → eski kart başı sayım birebir).
    # ≥2 sayılan kartlı grubun kartları grup TOPLAMIYLA değerlendirilir:
    # grup toplamı ≤ 0 → tek "sıfır", toplam ≤ grup min → tek kritik satır.
    from core.stock_gaps import group_index
    agg = group_index(db, domain)[2]

    zero_rows = (
        db.query(Item.id)
        .filter(Item.is_active == True, Item.domain == domain,
                Item.current_stock <= 0, Item.category.notin_(("Ambalaj", "Bitmiş Ürün")))
        .all()
    )
    zero_cards, zero_groups = 0, set()
    for (iid,) in zero_rows:
        g = agg.get(iid)
        if g is None:
            zero_cards += 1
        elif g.total <= 1e-9:
            zero_groups.add((g.group_id, g.bucket))     # grubun tamamı sıfır → bir kez
    zero_stock_raw_count = zero_cards + len(zero_groups)

    # ── Critical stock: items at or below min level ────────────────────────
    critical_items_raw = (
        db.query(Item)
        .filter(
            Item.is_active == True,
            Item.min_stock_level > 0,
            Item.current_stock <= Item.min_stock_level,
            Item.domain == domain,
        )
        .order_by(Item.current_stock.asc())
        .all()
    )
    # Kritik grup en az bir kritik üye içerir (toplam ≥ en büyük minimumlu
    # üyenin stoğu) → yalnız bu listeyi taramak tüm kritik grupları bulur.
    from core.group_stock import amount_exact
    critical_stock_list, seen_groups = [], set()
    for i in critical_items_raw:
        g = agg.get(i.id)
        if g is None:
            critical_stock_list.append({
                "id":              i.id,
                "name":            i.name,
                "sku":             i.sku or "—",
                "category":        i.category or "—",
                "unit":            i.unit,
                "current_stock":   round(i.current_stock, 4),
                "min_stock_level": round(i.min_stock_level, 4),
                "deficit":         round(i.min_stock_level - i.current_stock, 4),
                "group":           None,
            })
            continue
        gk = (g.group_id, g.bucket)
        if gk in seen_groups:
            continue
        seen_groups.add(gk)
        if not (g.min_stock > 1e-9 and g.total <= g.min_stock + 1e-9):
            continue                                  # grubu yeterli → kart uyarısı düşer
        critical_stock_list.append({
            "id":              None,
            "name":            g.name,
            "sku":             "",
            "category":        i.category or "—",
            "unit":            g.base_unit,
            "current_stock":   round(g.total, 4),
            "min_stock_level": round(g.min_stock, 4),
            "deficit":         round(g.min_stock - g.total, 4),
            "group": {
                "id": g.group_id, "name": g.name, "member_count": len(g.counted),
                # kesin metin: 999,6 g / 1.000 g "1 kg / 1 kg · 0 g" görünmesin
                "total_text": amount_exact(g.total, g.base_unit),
                "min_text": amount_exact(g.min_stock, g.base_unit),
                "deficit_text": amount_exact(max(0.0, g.min_stock - g.total), g.base_unit),
                "breakdown_text": g.breakdown_text,
            },
        })
    critical_stock_list.sort(key=lambda r: r["current_stock"])

    # ── Recent transactions (last 5) ───────────────────────────────────────
    recent_txs = (
        db.query(Transaction)
        .join(Item, Item.id == Transaction.item_id)
        .filter(Item.domain == domain)
        .order_by(Transaction.id.desc())
        .limit(5)
        .all()
    )
    recent_transactions = [
        {
            "id":               t.id,
            "item_name":        t.item.name if t.item else "—",
            "transaction_type": t.transaction_type,
            "quantity":         t.quantity,
            "notes":            (t.notes or "")[:90],   # truncate for display
            "timestamp":        to_tr(t.timestamp).strftime("%d.%m.%Y %H:%M") if t.timestamp else "",
        }
        for t in recent_txs
    ]

    # ── Recent production (last 5) ─────────────────────────────────────────
    recent_prod = (
        db.query(ProductionHistory)
        .filter(ProductionHistory.domain == domain,
                ProductionHistory.cancelled_at.is_(None))       # iptaller panoda yok
        .order_by(ProductionHistory.id.desc())
        .limit(5)
        .all()
    )

    return {
        "total_items":          total_items,
        "total_suppliers":      total_suppliers,
        "total_recipes":        total_recipes,
        "critical_stock_count": len(critical_stock_list),
        "critical_stock_list":  critical_stock_list,
        "zero_stock_raw_count": zero_stock_raw_count,
        "recent_transactions":  recent_transactions,
        "recent_productions": [
            {
                "recipe_name":       r.recipe_name,
                "target_item_name":  r.target_item_name,
                "produced_quantity": r.produced_quantity,
                "produced_at":       to_tr(r.produced_at).strftime("%d.%m.%Y %H:%M") if r.produced_at else "",
            }
            for r in recent_prod
        ],
    }


# ─── Üretim Stok Analizi (what-if simülasyonu) ──────────────────────────────
# Lab ekibi: seçilen ürünlerden N'er adet üretsek malzeme yeter mi? Excel iner.

class _PlanRequest(BaseModel):
    recipe_ids: List[int] = Field(..., min_length=1, max_length=500)
    quantity:   float     = Field(..., gt=0, le=1_000_000)
    language:   Optional[str] = "TR"


@router.get("/reports/production-plan/products")
def production_plan_products(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Aktif panelde reçetesi olan ürünler + reçetesi olmayan bitmiş ürünler (uyarı için)."""
    from core.production_sim import list_plan_products, list_recipeless_products
    return {"products": list_plan_products(db, domain),
            "recipeless": list_recipeless_products(db, domain)}


@router.post("/reports/production-plan")
def production_plan(
    data: _PlanRequest,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Seçilen ürünlerden quantity'şer adet üretim senaryosu — JSON rapor."""
    from core.production_sim import simulate
    return simulate(db, data.recipe_ids, data.quantity, data.language or "TR", domain)


@router.post("/reports/production-plan/export")
def production_plan_export(
    data: _PlanRequest,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Aynı senaryoyu Excel olarak indir (Satın Alma tedarikçi/fiyatla + Reçetesiz Ürünler sayfası)."""
    from core.production_sim import simulate, build_workbook, list_recipeless_products
    from core.supplier_prices import prices_for_items, today_rates_or_none
    rep = simulate(db, data.recipe_ids, data.quantity, data.language or "TR", domain)
    if not rep["materials"]:
        return JSONResponse(status_code=400, content={"detail": "Seçilen ürünlerde malzeme bulunamadı."})
    purchase_ids = [m["item_id"] for m in rep["purchase"] if m.get("item_id")]
    prices = prices_for_items(db, purchase_ids, domain, rates=today_rates_or_none)   # kur tembel
    recipeless = list_recipeless_products(db, domain)
    try:
        content = build_workbook(rep, title_suffix=f"{rep['summary']['products']} ürün × {int(data.quantity)}",
                                 prices=prices, recipeless=recipeless)
    except Exception:
        return JSONResponse(status_code=500, content={"detail": "Excel üretilemedi."})
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="uretim_stok_analizi.xlsx"'},
    )


# ─── Üretim geçmişi (ürün / marka / ürün grubu bazında) ─────────────────────

@router.get("/reports/production-history")
def production_history_report(
    start:  Optional[_dt.date] = Query(None),
    end:    Optional[_dt.date] = Query(None),
    brand:  Optional[str] = Query(None, max_length=300),
    q:      Optional[str] = Query(None, max_length=100),
    group:  str = Query("product", regex="^(product|brand|category)$"),
    format: str = Query("json", regex="^(json|xlsx)$"),
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """ProductionHistory toplamları — tarihler TR-yerel gün (varsayılan son 12 ay).

    `brand` virgüllü olabilir; `format=xlsx` → 'Özet' + 'Aylık' sayfalı Excel.
    """
    from core import production_history_report as phr
    # Varsayılan başlangıç hesabından (end − 364 gün) ÖNCE: uç tarih OverflowError → 500
    if not (phr.date_in_range(start) and phr.date_in_range(end)):
        return JSONResponse(status_code=400, content={
            "detail": f"Tarihler {phr.MIN_DATE:%d.%m.%Y} – {phr.MAX_DATE:%d.%m.%Y} aralığında olmalı."})
    d_start, d_end = phr.default_range()
    start = start or (d_start if end is None else end - (d_end - d_start))
    end = end or d_end
    if start > end:
        return JSONResponse(status_code=400, content={"detail": "Başlangıç tarihi bitiş tarihinden sonra olamaz."})
    if (end - start).days > 3660:
        return JSONResponse(status_code=400, content={"detail": "Tarih aralığı en fazla 10 yıl olabilir."})
    rep = phr.build_report(db, domain, start, end, brand=brand, q=q, group=group)
    rep["note"] = phr.since_note(rep["first_record"])
    if format == "json":
        return rep
    try:
        content = phr.build_workbook(rep)
    except Exception:
        return JSONResponse(status_code=500, content={"detail": "Excel üretilemedi."})
    fname = f"uretim_gecmisi_{start.isoformat()}_{end.isoformat()}_{group}.xlsx"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ─── Tedarikçi fiyat listesi (satın alma raporunu besler) ───────────────────
# Excel içe aktarma finans rolünde (Işık Hanım'ın "Stok Son Durum" dosyası);
# elle ekle/düzenle/sil `suppliers.prices` (Manager + LabLead — lab fiyatı
# kendisi girer, 07.10.2026).  Hepsi audit'li: supplier_prices.import |
# create | update | delete.
_FINANCE = ["SuperAdmin", "Manager"]


def _sp_actor(current_user: dict) -> str:
    return ((current_user or {}).get("full_name") or (current_user or {}).get("username") or "—")[:100]


def _sp_status_map(db: Session, domain: str) -> dict:
    """{supplier_id: purchase_status} — fiyat satırında firma durumu rozeti."""
    return dict(db.query(Supplier.id, Supplier.purchase_status)
                .filter(Supplier.domain == domain).all())


@router.get("/supplier-prices")
def list_supplier_prices(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """Malzemeye göre gruplu tedarikçi fiyat listesi (görüntüleme paneli).
    Satır şekli `core.supplier_prices.serialize_price` (+ kaynak, kim, not,
    firma durumu).  Pasif firmanın satırı (`inactive_firm_check`) grubun
    sonunda, `supplier_inactive: true` — panel "en ucuz" saymaz, gri rozet.
    KDV dahil satır NET fiyatıyla sıralanır (`net_unit_price`; KDV boşken
    net = fiyat → eski sıra), panel fiyatın yanına "KDV dahil" yazar."""
    from core.supplier_prices import inactive_firm_check, net_unit_price, serialize_price
    rows = db.query(SupplierPrice).filter(SupplierPrice.domain == domain).all()
    is_inactive = inactive_firm_check(db, domain, rows)
    item_ids = {r.item_id for r in rows}
    meta = {}
    if item_ids:
        for it in db.query(Item).filter(Item.id.in_(item_ids)).all():
            meta[it.id] = {"name": it.name, "unit": it.unit or "", "category": it.category or ""}
    status = _sp_status_map(db, domain)
    groups: dict = {}
    for r in rows:
        g = groups.setdefault(r.item_id, {
            "item_id": r.item_id,
            "material": meta.get(r.item_id, {}).get("name", "—"),
            "unit": meta.get(r.item_id, {}).get("unit", ""),
            "category": meta.get(r.item_id, {}).get("category", ""),
            "suppliers": [],
        })
        g["suppliers"].append(serialize_price(r, supplier_status=status.get(r.supplier_id),
                                              supplier_inactive=is_inactive(r)))
    out = sorted(groups.values(), key=lambda g: (g["material"] or "").lower())
    for g in out:
        g["suppliers"].sort(key=lambda s: (s["supplier_inactive"], s["unit_price"] is None,
                                           net_unit_price(s) or 0.0))
    return {"items": out, "total_items": len(out), "total_prices": len(rows)}


@router.get("/supplier-prices/overview")
def supplier_prices_overview(
    currency: str = Query("USD", max_length=8),
    q: Optional[str] = Query(None, max_length=100),
    category: str = Query("all", max_length=20),
    has_price: str = Query("all", max_length=20),
    source: str = Query("all", max_length=20),
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("reports", "view")),
    domain: str = Depends(active_domain),
):
    """"Malzeme Fiyatları" sayfası (`/fiyatlar`) — `core.price_overview.build_overview`.
    Aktif kartlar (bitmiş ürün hariç) → "aynı malzeme" grubu blokları →
    fiyat satırları: planın tercih / bitirilecek bayrakları, elle-fiyat
    çakışmasında gizlenen satır sayısı, seçilen para biriminde net fiyat
    (`currency` USD | EUR | TRY; kur yalnız gerekirse, sunucuda), en ucuz ve
    planın seçeceği teklif.  Süzgeçler: `q` (malzeme / tedarikçi / grup),
    `category` all | hammadde | ambalaj | etiket | other, `has_price`
    all | priced | unpriced, `source` all | manual | stok_son_durum | lab | parasut.
    Geçersiz değer → 400.  Görüntüleme `reports.view` (fiyatlar raporu gören
    herkese açık, 05.10.2026); yazma uçları `suppliers.prices`."""
    from core.price_overview import build_overview
    try:
        return build_overview(db, domain, currency=currency, q=q, category=category,
                              has_price=has_price, source=source)
    except ValueError as e:
        return JSONResponse(status_code=400, content={"detail": str(e)})


class SupplierPriceIn(BaseModel):
    item_id: int
    supplier_id: int
    unit_price: float
    currency: str = Field("USD", max_length=8)
    price_unit: str = Field(..., max_length=8)
    package_size: Optional[float] = None
    quoted_at: Optional[str] = Field(None, max_length=20)
    note: Optional[str] = Field(None, max_length=2000)
    source_label: Optional[str] = Field(None, max_length=120)
    # KDV — boş = bilinmiyor (eski davranış); dahilse oran zorunlu
    vat_included: Optional[bool] = None
    vat_rate: Optional[float] = None


class SupplierPriceUpdate(BaseModel):
    supplier_id: Optional[int] = None
    unit_price: Optional[float] = None
    currency: Optional[str] = Field(None, max_length=8)
    price_unit: Optional[str] = Field(None, max_length=8)
    package_size: Optional[float] = None
    quoted_at: Optional[str] = Field(None, max_length=20)
    note: Optional[str] = Field(None, max_length=2000)
    source_label: Optional[str] = Field(None, max_length=120)
    vat_included: Optional[bool] = None
    vat_rate: Optional[float] = None


def _sp_bad(msg: str, code: int = 400) -> JSONResponse:
    return JSONResponse(status_code=code, content={"detail": msg})


def _sp_finite(v) -> bool:
    import math
    try:
        return math.isfinite(float(v))
    except (TypeError, ValueError):
        return False


def _sp_clean(fields: dict, item: Item):
    """Elle fiyat alanlarını doğrula/normalize et → (temiz dict, hata|None).
    Yalnız `fields`'ta OLAN anahtarlar işlenir (PUT kısmi)."""
    from core import supplier_prices as SP
    out = {}
    if "unit_price" in fields:
        v = fields["unit_price"]
        # NaN/Infinity her karşılaştırmada False — `<= 0` kontrolünü atlardı.
        if v is None or not _sp_finite(v) or float(v) <= 0:
            return None, _sp_bad("Birim fiyat sıfırdan büyük bir sayı olmalı.")
        out["unit_price"] = float(v)
    if "currency" in fields:
        cur = (fields["currency"] or "").strip().upper()
        if cur not in SP.CURRENCIES:
            return None, _sp_bad("Para birimi USD, EUR ya da TRY olmalı.")
        out["currency"] = cur
    if "price_unit" in fields:
        pu = (fields["price_unit"] or "").strip().lower()
        if pu not in SP.PRICE_UNITS or not SP.price_unit_ok(item.unit, pu):
            fam = "adet" if SP.default_price_unit(item.unit) == "adet" else "kg ya da l"
            return None, _sp_bad(f"«{item.name}» ({item.unit or 'birimsiz'}) için fiyat birimi "
                                 f"{fam} olmalı.")
        out["price_unit"] = pu
    if "package_size" in fields:
        v = fields["package_size"]
        if v is not None and (not _sp_finite(v) or float(v) <= 0):
            return None, _sp_bad("Alınabilecek miktar sıfırdan büyük olmalı (ya da boş).")
        out["package_size"] = float(v) if v is not None else None
    if "quoted_at" in fields:
        raw = (fields["quoted_at"] or "").strip()
        try:
            out["quoted_at"] = _dt.date.fromisoformat(raw) if raw else None
        except ValueError:
            return None, _sp_bad("Geçersiz teklif tarihi.")
    if "note" in fields:
        out["note"] = (fields["note"] or "").strip() or None
    if "source_label" in fields:
        out["source_label"] = SP.clean_name(fields["source_label"])
    if "vat_included" in fields:
        out["vat_included"] = fields["vat_included"]          # None | True | False (pydantic doğruladı)
    if "vat_rate" in fields:
        v = fields["vat_rate"]
        if v is not None and (not _sp_finite(v) or not 0 <= float(v) <= SP.VAT_RATE_MAX):
            return None, _sp_bad("KDV oranı 0 ile 100 arasında olmalı (ya da boş).")
        out["vat_rate"] = float(v) if v is not None else None
    return out, None


def _sp_vat_ok(vat_included, vat_rate) -> Optional[JSONResponse]:
    """KDV dahil fiyatta oran şart — yoksa net hesaplanamaz, karşılaştırma
    brüt fiyatla yapılırdı.  Satırın SON hâline uygulanır (PUT kısmi)."""
    if vat_included is True and vat_rate is None:
        return _sp_bad("KDV dahil fiyat için KDV oranını yazın (ör. 20).")
    return None


def _sp_supplier(db: Session, domain: str, supplier_id: int):
    # FOR SHARE: eşzamanlı tedarikçi birleştirmesini bekle (pasif kaybedene fiyat yazılmasın)
    sup = (db.query(Supplier)
           .filter(Supplier.id == supplier_id, Supplier.domain == domain)
           .with_for_update(read=True).populate_existing().first())
    if sup is None:
        return None, _sp_bad("Tedarikçi bu panelde bulunamadı.", 404)
    if sup.is_active is False:
        return None, _sp_bad(f"«{sup.name}» pasif tedarikçi — fiyat girilemez.")
    return sup, None


def _sp_duplicate(db: Session, domain: str, item_id: int, supplier_id: int, price_unit: str,
                  exclude_id: Optional[int] = None):
    q = db.query(SupplierPrice).filter(
        SupplierPrice.domain == domain, SupplierPrice.item_id == item_id,
        SupplierPrice.supplier_id == supplier_id, SupplierPrice.price_unit == price_unit)
    if exclude_id is not None:
        q = q.filter(SupplierPrice.id != exclude_id)
    return q.order_by(SupplierPrice.id).first()


def _sp_audit_view(sp: SupplierPrice) -> dict:
    return {"item_id": sp.item_id, "supplier_id": sp.supplier_id, "supplier_name": sp.supplier_name,
            "unit_price": sp.unit_price, "currency": sp.currency, "price_unit": sp.price_unit,
            "package_size": sp.package_size, "source": sp.source, "source_label": sp.source_label,
            "quoted_at": sp.quoted_at.isoformat() if sp.quoted_at else None, "note": sp.note,
            "vat_included": sp.vat_included, "vat_rate": sp.vat_rate}


@router.post("/supplier-prices", status_code=201)
def create_supplier_price(
    data: SupplierPriceIn,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("suppliers", "prices")),
    domain: str = Depends(active_domain),
):
    """Elle fiyat ekle (lab) → `source='manual'` + `created_by`.  Kart ve
    tedarikçi aktif panelde ve aktif olmalı; fiyat birimi malzemenin birim
    ailesine uymalı (adetli → adet, kütle/hacim → kg|l).  Aynı (malzeme,
    tedarikçi, fiyat birimi) satırı varsa 409 `{code: 'price_exists', id}` —
    arayüz o satırı düzenlemeyi önerir.  KDV isteğe bağlı: `vat_included`
    (boş = bilinmiyor) + `vat_rate` (0–100); KDV dahilse oran zorunlu (400)."""
    from core.supplier_prices import SOURCE_MANUAL, serialize_price
    item = (db.query(Item)
            .filter(Item.id == data.item_id, Item.domain == domain,
                    Item.is_active == True).first())                              # noqa: E712
    if item is None:
        return _sp_bad("Malzeme bu panelde bulunamadı.", 404)
    sup, err = _sp_supplier(db, domain, data.supplier_id)
    if err:
        return err
    vals, err = _sp_clean(data.model_dump(), item)
    if err:
        return err
    err = _sp_vat_ok(vals.get("vat_included"), vals.get("vat_rate"))
    if err:
        return err
    dup = _sp_duplicate(db, domain, item.id, sup.id, vals["price_unit"])
    if dup is not None:
        return JSONResponse(status_code=409, content={
            "detail": f"«{item.name}» için «{sup.name}» fiyatı ({vals['price_unit']} başına) zaten "
                      f"kayıtlı — o satırı düzenleyin.",
            "code": "price_exists", "id": dup.id})
    actor = _sp_actor(current_user)
    sp = SupplierPrice(item_id=item.id, supplier_id=sup.id, supplier_name=sup.name,
                       source=SOURCE_MANUAL, created_by=actor, domain=domain, **vals)
    db.add(sp)
    db.commit()
    db.refresh(sp)
    log_admin_event(db, request, actor=current_user, action="supplier_prices.create",
                    target_type="supplier_price", target_id=sp.id,
                    target_name=f"{item.name} — {sup.name}"[:150],
                    details={"domain": domain, **_sp_audit_view(sp)})
    return serialize_price(sp, item, supplier_status=sup.purchase_status)


@router.put("/supplier-prices/{price_id}")
def update_supplier_price(
    price_id: int,
    data: SupplierPriceUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("suppliers", "prices")),
    domain: str = Depends(active_domain),
):
    """Fiyat satırını kısmen düzenle (yalnız gönderilen alanlar).  Excel'den
    gelen satır düzenlenince `source='manual'` olur — sonraki içe aktarma onu
    silmesin — ve eski liste etiketi (`source_label`) düşer; `updated_by`
    yazılır.  KDV alanları da kısmi; satırın son hâli "KDV dahil + oran yok"
    olamaz (400).  Audit: değişen alanlar eski → yeni."""
    from core.supplier_prices import SOURCE_MANUAL, serialize_price
    sp = (db.query(SupplierPrice)
          .filter(SupplierPrice.id == price_id, SupplierPrice.domain == domain).first())
    if sp is None:
        return _sp_bad("Kayıt bulunamadı.", 404)
    item = db.query(Item).filter(Item.id == sp.item_id).first()
    if item is None:
        return _sp_bad("Fiyat satırının malzemesi bulunamadı.", 404)
    fields = {k: getattr(data, k) for k in data.model_fields_set}
    sup = None
    if "supplier_id" in fields:
        if fields["supplier_id"] is None:
            return _sp_bad("Tedarikçi seçilmeli.")
        if fields["supplier_id"] != sp.supplier_id:
            sup, err = _sp_supplier(db, domain, fields["supplier_id"])
            if err:
                return err
        fields.pop("supplier_id")
    vals, err = _sp_clean(fields, item)
    if err:
        return err
    err = _sp_vat_ok(vals.get("vat_included", sp.vat_included), vals.get("vat_rate", sp.vat_rate))
    if err:
        return err
    if sup is not None:
        vals["supplier_id"] = sup.id
        vals["supplier_name"] = sup.name
    new_sid = vals.get("supplier_id", sp.supplier_id)
    new_pu = vals.get("price_unit", sp.price_unit)
    # Çift kontrolü yalnız tedarikçi / fiyat birimi DEĞİŞİYORSA — tedarikçi
    # birleştirmesinden sonra aynı (kart, firma, birim) iki satır kalabilir
    # (önizleme "ikisi de kalır" der); yalnız fiyatı düzeltmek engellenmesin.
    if new_sid is not None and new_pu and (new_sid != sp.supplier_id or new_pu != sp.price_unit):
        dup = _sp_duplicate(db, domain, sp.item_id, new_sid, new_pu, exclude_id=sp.id)
        if dup is not None:
            return JSONResponse(status_code=409, content={
                "detail": "Bu malzeme + tedarikçi + fiyat birimi için başka bir satır var.",
                "code": "price_exists", "id": dup.id})
    before = _sp_audit_view(sp)
    for k, v in vals.items():
        setattr(sp, k, v)
    # Excel satırı elle düzenlenince artık o listenin fiyatı değil: eski dosya
    # etiketi kalırsa satın alma planı elle fiyatı "<liste adı>" kaynağıyla
    # yazardı.  Etiket gönderildiyse o yazılır; eski değer audit'te kalır.
    if sp.source != SOURCE_MANUAL and "source_label" not in vals:
        sp.source_label = None
    sp.source = SOURCE_MANUAL
    sp.updated_by = _sp_actor(current_user)
    db.commit()
    db.refresh(sp)
    after = _sp_audit_view(sp)
    changes = {k: {"eski": before[k], "yeni": after[k]} for k in after if before[k] != after[k]}
    log_admin_event(db, request, actor=current_user, action="supplier_prices.update",
                    target_type="supplier_price", target_id=sp.id,
                    target_name=f"{item.name} — {sp.supplier_name or ''}"[:150],
                    details={"domain": domain, "changes": changes})
    status = None
    if sp.supplier_id is not None:
        status = db.query(Supplier.purchase_status).filter(Supplier.id == sp.supplier_id).scalar()
    return serialize_price(sp, item, supplier_status=status)


@router.post("/supplier-prices/import")
async def import_supplier_prices(
    request: Request,
    file: UploadFile = File(...),
    currency: str = Form("USD"),
    price_unit: str = Form("kg"),
    label: Optional[str] = Form(None),
    quoted_at: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_role(_FINANCE)),
    domain: str = Depends(active_domain),
):
    """Işık Hanım'ın 'Stok Son Durum' Excel'ini içe aktarır (malzeme başına eski
    Excel satırlarını değiştirir; lab'ın elle girdiği satırlar korunur).

    Form alanları fiyat temelini belirler: `currency` (USD|EUR|TRY, vars. USD),
    `price_unit` (kg|l, vars. kg — listenin ağırlık/hacim temeli; adet birimli
    malzemeler her zaman 'adet', 'adet' seçeneği bu yüzden yok), `label` (kaynak etiketi, boşsa dosya adı), `quoted_at`
    (YYYY-MM-DD liste/teklif tarihi).  Audit: `supplier_prices.import`.
    """
    from core import supplier_prices as SP
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        return JSONResponse(status_code=400, content={"detail": "Lütfen .xlsx dosyası yükleyin."})
    currency = (currency or "USD").strip().upper()
    if currency not in SP.CURRENCIES:
        return JSONResponse(status_code=400, content={"detail": "Para birimi USD, EUR ya da TRY olmalı."})
    price_unit = (price_unit or "kg").strip().lower()
    if price_unit not in SP.LIST_PRICE_UNITS:
        return JSONResponse(status_code=400, content={
            "detail": "Fiyat birimi kg ya da l olmalı (adet birimli malzemeler zaten adet başına yazılır)."})
    label = SP.clean_name(label)
    if label and len(label) > SP.SOURCE_LABEL_MAX:
        return JSONResponse(status_code=400, content={
            "detail": f"Etiket en fazla {SP.SOURCE_LABEL_MAX} karakter olabilir."})
    quoted = None
    if quoted_at and quoted_at.strip():
        try:
            quoted = _dt.date.fromisoformat(quoted_at.strip())
        except ValueError:
            return JSONResponse(status_code=400, content={"detail": "Geçersiz liste tarihi."})
    source_label = label or SP.clean_name(file.filename)
    data = await file.read()
    try:
        rows = SP.parse_stok_son_durum(data)
    except Exception:
        return JSONResponse(status_code=400, content={"detail": "Excel okunamadı — beklenen 'Stok Son Durum' düzeni mi?"})
    if not rows:
        return JSONResponse(status_code=400, content={"detail": "Dosyada veri satırı bulunamadı."})
    summary = SP.import_prices(db, rows, domain, currency=currency, price_unit=price_unit,
                               source_label=source_label, quoted_at=quoted)
    log_admin_event(db, request, actor=current_user, action="supplier_prices.import",
                    target_type="supplier_prices", target_name=summary["source_label"],
                    details={"domain": domain, "filename": file.filename,
                             "currency": summary["currency"], "price_unit": summary["price_unit"],
                             "quoted_at": summary["quoted_at"],
                             "items_updated": summary["items_updated"],
                             "prices_inserted": summary["prices_inserted"],
                             "kept_manual": summary["kept_manual"],
                             "skipped_manual": summary["skipped_manual"],
                             "unmatched_materials": len(summary["unmatched_materials"]),
                             "unmatched_suppliers": len(summary["unmatched_suppliers"])})
    return summary


@router.delete("/supplier-prices/{price_id}")
def delete_supplier_price(
    price_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("suppliers", "prices")),
    domain: str = Depends(active_domain),
):
    """Tek bir tedarikçi fiyat satırını siler.  Yetki `suppliers.prices`
    (izin tablosundan — JWT'deki rol değil; rolü düşürülen kullanıcının eski
    oturumu silemez).  Audit `supplier_prices.delete` silinen değerlerle."""
    row = db.query(SupplierPrice).filter(
        SupplierPrice.id == price_id, SupplierPrice.domain == domain
    ).first()
    if not row:
        return JSONResponse(status_code=404, content={"detail": "Kayıt bulunamadı."})
    item = db.query(Item).filter(Item.id == row.item_id).first()
    details = {"domain": domain, **_sp_audit_view(row)}
    name = f"{item.name if item else row.item_id} — {row.supplier_name or ''}"[:150]
    db.delete(row)
    db.commit()
    log_admin_event(db, request, actor=current_user, action="supplier_prices.delete",
                    target_type="supplier_price", target_id=price_id, target_name=name,
                    details=details)
    return {"ok": True}
