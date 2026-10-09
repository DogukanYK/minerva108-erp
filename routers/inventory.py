# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Inventory router — items, suppliers, receiving, transactions, traceability,
and Excel imports (both classic templated import + smart auto-detect import).
"""
import math
import re

from fastapi import APIRouter, Depends, UploadFile, File, BackgroundTasks, Request
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload
from pydantic import BaseModel, Field
from datetime import datetime
from typing import Literal, Optional, List


# ── Etiket dil/grup yardımcıları ────────────────────────────────────────────
# Migration (d579c1629598) ile aynı mantık — isimden dil belirteçlerini at,
# normalize et → label_group anahtarı.  TR ve EN kardeşleri aynı anahtarda
# buluşur, üretimde dil seçilince doğru kardeşe inilir.
_LANG_TOKEN_RE = re.compile(r'\(\s*(?:eng?|ing|tur|tr|t[üu]rk(?:[çc]e)?|english)\s*\)', re.IGNORECASE)


def _label_group_key(name: str) -> str:
    s = _LANG_TOKEN_RE.sub(' ', name or '')
    s = re.sub(r'\s+', ' ', s.lower()).strip()
    return s


def _norm_language(val: Optional[str]) -> Optional[str]:
    """Kullanıcı girdisini 'TR' / 'EN' / None'a normalize et."""
    v = (val or '').strip().upper()
    if v in ('TR', 'TUR', 'TÜRKÇE', 'TURKCE'):    return 'TR'
    if v in ('EN', 'ENG', 'İNG', 'ING', 'ENGLISH'): return 'EN'
    return None

from database import (
    to_tr,
    get_db, DuplicateItemDecision, Item, MaterialGroup, Supplier, Inventory, Transaction,
    Recipe, RecipeIngredient, RetentionSample, User,
)
from core.auth import get_current_user
from core.permissions import (_can_see_finance, require_any_permission, require_internal_user,
                              require_permission)
from core.notifications import notify_low_stock
from core.undo import record as record_undoable
from core.domain import active_domain
from core.audit import log_admin_event
from core.supplier_prices import normalize as _tr_fold
from core import material_groups, stock_lots
from core.stock_lots import LotMoveError

router = APIRouter(prefix="/api", tags=["inventory"])


# ─── Schemas ────────────────────────────────────────────────────────────────

class ItemCreateRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=150)
    name_tr: Optional[str] = Field(None, max_length=150)   # Türkçe ad (belge dili + çift-dilli arama)
    category: Optional[str] = Field(None, max_length=50)
    # Varsayılan YOK — 24.08.2026 olayında modal 'adet'i ön-seçili tutuyordu,
    # kullanıcı dokunmadan hammadde 'adet' birimiyle açılıyordu (bkz. incident
    # notu CLAUDE.md). Birim artık AÇIKÇA seçilmeli.
    unit: str = Field(..., min_length=1, max_length=20)
    min_stock:      Optional[float] = 0.0
    cost_price:     Optional[float] = 0.0
    parent_id:      Optional[int]   = None
    variation_name: Optional[str]   = Field(None, max_length=100)
    barcode:        Optional[str]   = Field(None, max_length=64)
    pkg_type:       Optional[str]   = Field(None, max_length=20)
    # Etiket dili — 'TR' / 'EN' / None (dilsiz).  Sadece etiketlerde anlamlı.
    language:       Optional[str]   = Field(None, max_length=8)
    # Varsayılan tedarikçi — mal kabulde bu ürün seçilince oto-doldurulur
    supplier_id:    Optional[int]   = None
    # Ad-çakışması uyarısı görüldükten sonra kullanıcı yine de yeni kart
    # açmak isterse True gönderilir (bkz. _find_name_conflict).
    force:          bool = False
    # Numune giriş penceresinin "bu tedarikçi için yeni kart aç" akışı: yeni
    # kart bu kartın "aynı malzeme" grubuna katılır (grubu yoksa ikisi için
    # açılır — core.material_groups.join).  Aynı panel, aktif, aynı tür.
    join_group_of_item_id: Optional[int] = None


class BulkDeleteRequest(BaseModel):
    # max_items 1000: tek istekte 1000 ürün silmek operasyonel kapsamımızdan büyük
    item_ids: List[int] = Field(..., min_length=1, max_length=1000)


class SupplierCreateRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=150)
    contact_person: Optional[str] = Field(None, max_length=100)
    email: Optional[str] = Field(None, max_length=150)
    phone: Optional[str] = Field(None, max_length=30)
    # PUT'ta YALNIZ gönderilirse yazılır (model_fields_set) — adres alanını
    # bilmeyen eski istemci kayıtlı adresi silmesin.
    address: Optional[str] = Field(None, max_length=1000)
    notes: Optional[str] = Field(None, max_length=2000)


class SupplierBulkDeleteRequest(BaseModel):
    supplier_ids: List[int] = Field(..., min_length=1, max_length=500)


class StockReceiveRequest(BaseModel):
    item_id: int
    supplier_id: Optional[int] = None
    lot_number: str = Field(..., min_length=1, max_length=100)
    expiry_date: Optional[str] = Field(None, max_length=20)
    quantity: float
    location: Optional[str] = Field(None, max_length=100)
    # Numune kabulü — alternatif tedarikçiden gelen numune partisi.  Ayrı lot
    # olarak işaretlenir, normal lotla birleşmez, ÜRETİMDE KULLANILAMAZ ve
    # current_stock'a GİRMEZ (bkz. receive_stock).  Stoğa geçmek için
    # POST /inventory/samples/{id}/convert kullanılır.
    is_sample: Optional[bool] = False


def _finite(v) -> bool:
    """NaN/Infinity değil mi.  JSON gövdesinde `NaN` geçer (Starlette
    json.loads), pydantic float'ı da kabul eder; NaN her karşılaştırmada False
    döndüğü için `<= 0` kontrollerini atlayıp deftere NaN yazardı (defter
    değişmez → compute_stock_at kalıcı NaN).  `allow_inf_nan=False` YETMEZ:
    422 gövdesi girdiyi (nan) geri yazar, JSON'a dökülemez → 500."""
    try:
        return math.isfinite(float(v))
    except (TypeError, ValueError):
        return False


class StockAdjustRequest(BaseModel):
    """
    Manual stock correction. The user supplies the *target* quantity that
    Item.current_stock should hold; backend computes the delta and writes
    an immutable Transaction(type='Adjustment') so the change is auditable.
    Reason is mandatory — no silent corrections.
    """
    item_id: int
    new_quantity: float
    reason: str = Field(..., min_length=3, max_length=200)
    lot_number: Optional[str] = Field(None, max_length=100)


# ─── Items Endpoints ─────────────────────────────────────────────────────────

@router.get("/items")
def list_items(
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("items", "view")),
    domain: str = Depends(active_domain),
):
    items = (db.query(Item)
             .filter(Item.is_active == True, Item.domain == domain)
             .order_by(Item.id.desc()).all())

    # Resolve parent names + child counts in O(n) — avoids N+1 queries
    name_by_id = {i.id: i.name for i in items}
    child_count = {}
    for i in items:
        if i.parent_id:
            child_count[i.parent_id] = child_count.get(i.parent_id, 0) + 1

    # Supplier lookup — listing'de supplier_name göstermek için.  Pasif
    # (yumuşak silinmiş) tedarikçinin adı da çözülür; `supplier_active` kart
    # penceresinin onu "(pasif)" seçeneği olarak tutmasını sağlar.
    supplier_name_by_id, inactive_sup = {}, set()
    for sid, sname, sact in (db.query(Supplier.id, Supplier.name, Supplier.is_active)
                             .filter(Supplier.domain == domain).all()):
        supplier_name_by_id[sid] = sname
        if sact is False:
            inactive_sup.add(sid)
    # Kartta bekleyen numune miktarı — stoğa DAHİL DEĞİL, ama "Mevcut Stok"
    # yanında görünmezse lab numuneyi kayıp sanıyor (Songül Hanım, 05.10).
    # Tek GROUP BY sorgusu.
    sample_qty = dict(
        db.query(Inventory.item_id, func.sum(Inventory.quantity))
        .filter(Inventory.is_sample == True, Inventory.quantity > 0,     # noqa: E712
                Inventory.domain == domain)
        .group_by(Inventory.item_id).all()
    )
    group_name_by_id = dict(
        db.query(MaterialGroup.id, MaterialGroup.name)
        .filter(MaterialGroup.domain == domain, MaterialGroup.is_active == True).all()  # noqa: E712
    )

    finance_ok = _can_see_finance(current_user)
    return [
        {
            "id":              i.id,
            "name":            i.name,
            "name_tr":         i.name_tr or "",
            "category":        i.category,
            "pkg_type":        i.pkg_type or "",                # Phase 10 — Ambalaj alt-tipi
            "unit":            i.unit,
            "min_stock_level": i.min_stock_level,
            "current_stock":   i.current_stock,
            # ── Finance-gated: zeroed for lab roles (defense in depth vs. DevTools snooping) ──
            "cost_price":      round(i.cost_price or 0.0, 4) if finance_ok else 0.0,
            "parent_id":       i.parent_id,
            "parent_name":     name_by_id.get(i.parent_id) if i.parent_id else None,
            "variation_name":  i.variation_name,
            "barcode":         i.barcode or "",                 # Phase 9
            "language":        i.language or "",                 # Phase 15 — etiket dili
            "label_group":     i.label_group or "",
            "child_count":     child_count.get(i.id, 0),
            "is_parent":       child_count.get(i.id, 0) > 0,
            "is_variation":    i.parent_id is not None,
            "supplier_id":     i.supplier_id,
            "supplier_name":   supplier_name_by_id.get(i.supplier_id) if i.supplier_id else None,
            "supplier_active": (i.supplier_id not in inactive_sup) if i.supplier_id else None,
            "created_at":      to_tr(i.created_at).strftime("%d.%m.%Y") if i.created_at else "",
            "sample_qty":      round(float(sample_qty.get(i.id) or 0.0), 4),
            "material_group_id":   i.material_group_id if i.material_group_id in group_name_by_id else None,
            "material_group_name": group_name_by_id.get(i.material_group_id),
        }
        for i in items
    ]


_PRINT_PKG_LABELS = {"şişe": "Şişe", "kavanoz": "Kavanoz", "pompa": "Pompa",
                     "kapak": "Kapak", "etiket": "Etiket"}


@router.get("/items/print")
def print_items(
    tab: str = "hammadde",
    pkg: Optional[str] = None,       # ambalaj alt-tipi; '__none__' = pkg_type boş
    q: Optional[str] = None,         # arama (name/category/pkg_type/barcode)
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("items", "view")),
    domain: str = Depends(active_domain),
):
    """Aktif sekme + filtreyle antetli A4 ürün listesi PDF'i (Ürünler → Yazdır).

    Sekme filtreleri items.html TABS ile birebir; Maliyet kolonu yalnız finans
    yetkisi olan rollerde PDF'e girer (sunucu tarafı gate — sızma yolu yok).
    """
    import io as _io
    from datetime import datetime as _dt
    from core.items_report import (render_items_pdf, items_filename, tr_key, TAB_LABELS)
    from core.delivery_note import content_disposition
    from database import tr_now

    tab = (tab or "hammadde").lower()
    if tab not in ("hammadde", "ambalaj", "bitmis_urun", "numune"):
        return JSONResponse(status_code=400, content={"detail": "Geçersiz sekme."})
    qn = (q or "").strip().casefold() or None

    finance_ok = _can_see_finance(current_user)
    rows = []
    pkg_label = None

    if tab == "numune":
        recs = (db.query(Inventory)
                .options(joinedload(Inventory.item), joinedload(Inventory.supplier))
                .filter(Inventory.is_sample == True, Inventory.domain == domain)  # noqa: E712
                .order_by(Inventory.id.desc()).all())
        for r in recs:
            item_name = r.item.name if r.item else "—"
            supplier = r.supplier.name if r.supplier else "—"
            if qn and not any(qn in str(v or "").casefold()
                              for v in (item_name, supplier, r.lot_number)):
                continue
            rows.append({
                "item_name": item_name, "supplier": supplier, "lot": r.lot_number,
                "qty_unit": f"{round(float(r.quantity or 0), 4):g} {(r.item.unit if r.item else '') or ''}".strip(),
                "expiry": r.expiry_date or "—",
                "received": to_tr(r.created_at).strftime("%d.%m.%Y") if r.created_at else "—",
            })
        rows.sort(key=lambda x: tr_key(x["item_name"]))
    else:
        items = (db.query(Item)
                 .filter(Item.is_active == True, Item.domain == domain).all())  # noqa: E712
        # Sekme filtresi — items.html TABS ile birebir
        if tab == "hammadde":
            items = [i for i in items if i.category not in ("Ambalaj", "Bitmiş Ürün")]
        elif tab == "ambalaj":
            items = [i for i in items if i.category == "Ambalaj"]
            if pkg == "__none__":
                items = [i for i in items if not (i.pkg_type or "").strip()]
                pkg_label = "Tipi yok"
            elif pkg:
                items = [i for i in items if (i.pkg_type or "").lower() == pkg.lower()]
                pkg_label = _PRINT_PKG_LABELS.get(pkg.lower(), pkg)
        else:
            items = [i for i in items if i.category == "Bitmiş Ürün"]
        # Arama — items.html JS aramasıyla birebir (name/category/pkg_type/barcode + name_tr)
        if qn:
            items = [i for i in items if any(
                qn in str(v or "").casefold()
                for v in (i.name, i.name_tr, i.category, i.pkg_type, i.barcode))]

        suppliers = _supplier_names(db, domain)       # pasif firmanın adı da

        def flat_row(i):
            return {"name": i.name, "unit": i.unit or "", "stock": i.current_stock,
                    "min": i.min_stock_level,
                    "supplier": suppliers.get(i.supplier_id) if i.supplier_id else None,
                    "cost": (i.cost_price if finance_ok else None),
                    "pkg_label": _PRINT_PKG_LABELS.get((i.pkg_type or "").lower(),
                                                       (i.pkg_type or "").strip() or None),
                    "language": i.language or None}

        if tab == "bitmis_urun":
            # Hiyerarşi — items.html renderTable ile birebir: parents → varyasyonları → standalone
            def fin_row(i, kind):
                return {"kind": kind,
                        "name": (i.variation_name or i.name) if kind == "variation" else i.name,
                        "name_tr": i.name_tr or None, "unit": i.unit or "",
                        "stock": i.current_stock, "min": i.min_stock_level,
                        "cost": (i.cost_price if finance_ok else None)}
            by_parent = {}
            for i in items:
                if i.parent_id:
                    by_parent.setdefault(i.parent_id, []).append(i)
            parents = [i for i in items if not i.parent_id and i.id in by_parent]
            standalones = [i for i in items if not i.parent_id and i.id not in by_parent]
            parents.sort(key=lambda i: tr_key(i.name))
            standalones.sort(key=lambda i: tr_key(i.name))
            for p in parents:
                rows.append(fin_row(p, "parent"))
                kids = sorted(by_parent[p.id], key=lambda i: tr_key(i.variation_name or i.name))
                rows.extend(fin_row(k, "variation") for k in kids)
            # Ana ürünü listede olmayan (orphan) varyasyonlar — düz satır
            parent_ids = {p.id for p in parents}
            orphans = [i for i in items if i.parent_id and i.parent_id not in parent_ids]
            for o in sorted(orphans, key=lambda i: tr_key(i.name)):
                rows.append(fin_row(o, "standalone"))
            rows.extend(fin_row(i, "standalone") for i in standalones)
        else:
            items.sort(key=lambda i: tr_key(i.name))
            rows = [flat_row(i) for i in items]

    view = {"tab": tab, "tab_label": TAB_LABELS[tab], "pkg_label": pkg_label,
            "q": (q or "").strip() or None, "finance": finance_ok,
            "date": tr_now().strftime("%d.%m.%Y %H:%M"), "rows": rows}
    try:
        content = render_items_pdf(view)
    except Exception:
        return JSONResponse(status_code=500, content={"detail": "Liste üretilemedi."})
    fname = items_filename(tab, pkg_label, (q or "").strip() or None)
    return StreamingResponse(
        _io.BytesIO(content),
        media_type="application/pdf",
        headers={"Content-Disposition": content_disposition(fname, inline=True)},
    )


def _validate_variation(
    db: Session,
    parent_id: Optional[int],
    variation_name: Optional[str],
    self_id: Optional[int] = None,
) -> Optional[JSONResponse]:
    """
    Shared parent-child validation for both create_item and update_item.
    Returns a JSONResponse on validation failure, or None if OK.
    """
    if parent_id is None:
        return None  # Standalone or future-parent — nothing to validate

    if self_id is not None and parent_id == self_id:
        return JSONResponse(status_code=400, content={"detail": "Bir ürün kendisinin varyasyonu olamaz."})

    parent = db.query(Item).filter(Item.id == parent_id).first()
    if not parent:
        return JSONResponse(status_code=400, content={"detail": "Seçilen ana ürün bulunamadı."})

    if parent.parent_id is not None:
        return JSONResponse(status_code=400, content={
            "detail": "Varyasyonlar bir başka varyasyonun altına eklenemez (depth=1 sınırı)."
        })

    if not variation_name or not variation_name.strip():
        return JSONResponse(status_code=400, content={"detail": "Varyasyon adı zorunludur (örn: '200ml')."})

    return None


def _find_name_conflict(db: Session, domain: str, name: str,
                        name_tr: Optional[str], exclude_id: Optional[int] = None):
    """Türkçe-katlanmış ad çakışması ara — 24.08.2026 olayının kapısı.

    Stajyer mevcut "BADEM YAĞI" kartını bulamayıp "BADEM YAGI" adıyla yeni
    kart açtı (Item.name'de unique kısıt yok, eski arama da Türkçe harfleri
    katlamıyordu).  Burada `core.supplier_prices.normalize` (İ/ı/ğ/ş/ö/ü/ç
    katlama + boşluk normalize) ile YENİ adın hem mevcut `name` hem
    `name_tr`'ye karşı çapraz kontrolü yapılır.  Bulunca engellemez —
    çağıran `force=True` ile geçmeyi seçebilir.
    """
    new_key = _tr_fold(name)
    new_tr_key = _tr_fold(name_tr) if name_tr else ""
    if not new_key and not new_tr_key:
        return None
    q = db.query(Item).filter(Item.domain == domain, Item.is_active == True)  # noqa: E712
    if exclude_id:
        q = q.filter(Item.id != exclude_id)
    for it in q.all():
        existing_key = _tr_fold(it.name)
        existing_tr_key = _tr_fold(it.name_tr) if it.name_tr else ""
        if new_key and new_key in (existing_key, existing_tr_key):
            return it
        if new_tr_key and new_tr_key in (existing_key, existing_tr_key):
            return it
    return None


def _conflict_payload(it: Item) -> dict:
    return {"id": it.id, "name": it.name, "name_tr": it.name_tr or "",
            "category": it.category or "", "unit": it.unit or "",
            "current_stock": round(float(it.current_stock or 0), 4)}


@router.post("/items", status_code=201)
def create_item(data: ItemCreateRequest, request: Request, db: Session = Depends(get_db),
                current_user: dict = Depends(require_permission("items", "create")),
                domain: str = Depends(active_domain)):
    err = _validate_variation(db, data.parent_id, data.variation_name)
    if err: return err
    err = _check_supplier_ref(db, domain, data.supplier_id)
    if err: return err

    # `join_group_of_item_id` BİLİNÇLİ İSTİSNA: grup yazma kuralı items.edit'tir
    # (/api/material-groups) ama bu yol items.create ile açılır — numune giriş
    # penceresinin "bu tedarikçi için yeni kart aç" akışı.  Numuneyi giren
    # LabTech'tir (items.create var, items.edit yok); P1'in çevirme penceresindeki
    # new_item yolu da (inventory.receive) yeni kartı aynı şekilde gruba katar.
    # Yalnız YENİ açılan kart gruba girer; var olan kartlar taşınamaz,
    # çıkarılamaz, grup adı değişmez — onlar items.edit ister.
    group_source = None
    if data.join_group_of_item_id is not None:
        group_source = (db.query(Item)
                        .filter(Item.id == data.join_group_of_item_id, Item.domain == domain,
                                Item.is_active == True).first())               # noqa: E712
        if group_source is None:
            return JSONResponse(status_code=400, content={
                "detail": "Grubuna katılınacak kart bulunamadı ya da pasif."})
        kind = stock_lots.lot_kind(data.category)
        if kind == "finished" or kind != stock_lots.lot_kind(group_source.category):
            return JSONResponse(status_code=400, content={
                "detail": f"Yeni kart «{group_source.name}» ile aynı malzeme grubuna giremez — "
                          f"tür farklı (hammadde / ambalaj) ya da bitmiş ürün."})

    if not data.force:
        conflict = _find_name_conflict(db, domain, data.name, data.name_tr)
        if conflict:
            return JSONResponse(status_code=409, content={
                "detail": f"Aynı isimde ürün zaten kayıtlı: {conflict.name}",
                "existing": _conflict_payload(conflict),
            })

    pkg = (data.pkg_type.strip() if data.pkg_type and data.pkg_type.strip() else None)
    is_label = (data.category == "Ambalaj" and pkg == "etiket")
    lang     = _norm_language(data.language) if is_label else None
    # label_group sadece etiketlerde — isimden türetilir (TR/EN kardeşleri buluşsun)
    grp      = _label_group_key(data.name) if is_label else None

    item = Item(
        name=data.name,
        name_tr=((data.name_tr or "").strip() or None),
        category=data.category,
        unit=data.unit,
        min_stock_level=data.min_stock  or 0.0,
        # Maliyet yalnız finans rolünden — LabLead/LabTech'in gönderdiği
        # cost_price (API ile bile) yok sayılır.
        cost_price=(data.cost_price or 0.0) if _can_see_finance(current_user) else 0.0,
        parent_id=data.parent_id,
        variation_name=(data.variation_name.strip() if data.parent_id and data.variation_name else None),
        barcode=(data.barcode.strip() if data.barcode and data.barcode.strip() else None),
        pkg_type=pkg,
        language=lang,
        label_group=grp,
        supplier_id=data.supplier_id,
        domain=domain,                     # Faz 3 — aktif panele damgala
    )
    db.add(item)
    grp = None
    if group_source is not None:
        db.flush()
        actor = (current_user or {}).get("full_name") or (current_user or {}).get("username") or ""
        grp = material_groups.join(db, group_source, item, actor=actor)
    db.commit()
    db.refresh(item)
    out = {"id": item.id, "message": "Ürün başarıyla eklendi.", "material_group": None}
    if grp is not None:
        out["material_group"] = {"id": grp.id, "name": grp.name}
        out["warning"] = material_groups.unit_warning([group_source.unit, item.unit])
        log_admin_event(db, request, actor=current_user, action="material_group.add",
                        target_type="material_group", target_id=grp.id, target_name=grp.name,
                        details={"item_id": item.id, "item": item.name,
                                 "kaynak_kart": group_source.id, "yol": "yeni kart"})
    return out


# ─── Kopya kartı kararları — lab popup'ı ────────────────────────────────────
# 24.08.2026 numune olayı taramasında bulunan kopya hammadde kümeleri.  Hangi
# kartların aynı ürün olduğuna LAB karar verir; "birleştir" kararı sunucuda
# core/item_merge.merge_items ile anında uygulanır (Adjustment çifti + FK
# taşıma + pasifleştirme — defter kuralları o modülün docstring'inde).

class DupDecisionBody(BaseModel):
    action: str = Field(..., max_length=10)              # merge | keep
    target_item_id: Optional[int] = None                 # merge için zorunlu


@router.get("/items/dup-decisions")
def list_dup_decisions(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("items", "edit")),
):
    from core.item_merge import pending_clusters
    clusters = pending_clusters(db)
    db.commit()          # pending_clusters tek-kartlı kümeleri kapatmış olabilir
    return {"clusters": clusters, "count": len(clusters)}


@router.post("/items/dup-decisions/{decision_id}")
def decide_duplicate(
    decision_id: int,
    data: DupDecisionBody,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("items", "edit")),
):
    from core.item_merge import MergeError, merge_items
    import json as _json

    row = (db.query(DuplicateItemDecision)
           .filter(DuplicateItemDecision.id == decision_id,
                   DuplicateItemDecision.status == "pending")
           .with_for_update().first())
    if not row:
        return JSONResponse(status_code=404,
                            content={"detail": "Karar kaydı bulunamadı ya da kapanmış."})
    actor = (current_user or {}).get("full_name") or (current_user or {}).get("username") or "sistem"

    if data.action == "keep":
        row.status = "kept"
        row.decided_by = actor
        row.decided_at = datetime.utcnow()
        row.result_note = "Lab kararı: farklı ürünler, kartlar ayrı kalacak."
        # "Ayrı kalsın" = aynı malzemenin farklı tedarikçi kartları → "aynı
        # malzeme" grubu (satın alma alternatifleri + numune çevirme hedefi).
        groups = material_groups.group_decision(db, row, actor=actor)
        db.commit()
        grp_info = [{"id": g.id, "name": g.name} for g in groups]
        log_admin_event(db, request, actor=current_user, action="items.dup_keep",
                        target_type="dup_decision", target_id=row.id,
                        target_name=row.title,
                        details={"karar": "ayrı kalsın",
                                 "ayni_malzeme_grubu": [g["id"] for g in grp_info]})
        msg = f"«{row.title}» — kartlar ayrı bırakıldı."
        if grp_info:
            msg += f" «{grp_info[0]['name']}» aynı malzeme grubunda bağlandı."
        return {"message": msg,
                "material_group": grp_info[0] if grp_info else None}

    if data.action != "merge":
        return JSONResponse(status_code=400, content={"detail": "Geçersiz karar."})
    if not data.target_item_id:
        return JSONResponse(status_code=400,
                            content={"detail": "Birleştirme için kalacak kartı seçin."})
    try:
        ids = [int(x) for x in _json.loads(row.item_ids)]
    except (TypeError, ValueError):
        return JSONResponse(status_code=500, content={"detail": "Küme verisi bozuk."})
    if data.target_item_id not in ids:
        return JSONResponse(status_code=400,
                            content={"detail": "Hedef kart bu kümeye ait değil."})

    losers = [i for i in ids if i != data.target_item_id
              and db.query(Item.id).filter(Item.id == i,
                                           Item.is_active == True).first()]  # noqa: E712
    summaries = []
    try:
        for loser_id in losers:
            summaries.append(merge_items(db, loser_id, data.target_item_id, actor))
        row.status = "merged"
        row.target_item_id = data.target_item_id
        row.decided_by = actor
        row.decided_at = datetime.utcnow()
        row.result_note = _json.dumps(summaries, ensure_ascii=False)[:2000]
        db.commit()
    except MergeError as exc:
        db.rollback()
        return JSONResponse(status_code=400, content={"detail": str(exc)})
    except Exception:
        db.rollback()
        return JSONResponse(status_code=500,
                            content={"detail": "Birleştirme sırasında hata — hiçbir şey değişmedi."})

    dissolved = [dict(g, loser_id=s["loser_id"]) for s in summaries
                 for g in s.get("dissolved_groups") or []]
    log_admin_event(db, request, actor=current_user, action="items.dup_merge",
                    target_type="dup_decision", target_id=row.id,
                    target_name=row.title,
                    details={"hedef": data.target_item_id,
                             "birlesen": [s["loser_id"] for s in summaries],
                             "tasinan_stok": [f'{s["moved_stock"]:g} {s["unit"]}'
                                              for s in summaries],
                             "dagilan_gruplar": [g["id"] for g in dissolved]})
    # Kaybeden kartın "aynı malzeme" grubu tek aktif karta inip dağıldıysa —
    # /api/material-groups'un remove ucuyla aynı iz.
    for g in dissolved:
        log_admin_event(db, request, actor=current_user, action="material_group.dissolve",
                        target_type="material_group", target_id=g["id"], target_name=g["name"],
                        details={"sebep": "kart birleştirme", "karar": row.id,
                                 "kaybeden": g["loser_id"], "kazanan": data.target_item_id})
    moved = sum(s["moved_stock"] for s in summaries)
    unit = summaries[0]["unit"] if summaries else ""
    return {"message": f"«{row.title}» birleştirildi — {len(summaries)} kart kapandı, "
                       f"{moved:g} {unit} stok tek karta taşındı."}


@router.post("/items/import-names")
async def import_item_names(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("items", "import")),
):
    """Excel'den ürünlere Türkçe ad (name_tr) toplu yükle.

    Esnek kolon: bir kolon SKU **veya** mevcut (İngilizce) ad ile eşleştirir,
    diğer kolon Türkçe ad'dır.  Başlıkta 'türkçe'/'turkish' geçen → TR ad;
    'sku' geçen → eşleştirici.  Başlık yoksa: 1. kolon eşleştirici, 2. kolon TR.
    """
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        return JSONResponse(status_code=400, content={"detail": "Lütfen .xlsx dosyası yükleyin."})
    data = await file.read()
    try:
        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
        rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    except Exception:
        return JSONResponse(status_code=400, content={"detail": "Excel okunamadı."})
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return JSONResponse(status_code=400, content={"detail": "Dosyada veri satırı yok."})

    hdr = [str(c or "").strip().lower() for c in rows[0]]

    def _find(keys):
        return next((i for i, h in enumerate(hdr) if any(k in h for k in keys)), None)

    bc_col   = _find(("barcode", "barkod", "ean"))
    tr_col   = _find(("türkçe", "turkce", "turkish", "tr ad", "tr_ad"))
    sku_col  = _find(("sku", "stok kod", "kod", "code"))
    # Ürün-adı kolonu: önce açık 'product name/ürün adı', yoksa 'brand'/'barcode'
    # OLMAYAN herhangi bir '...name' kolonu (ör. 'BRAND NAME' yanlışlıkla seçilmesin).
    name_col = _find(("product name", "ürün ad", "urun ad", "product"))
    if name_col is None:
        name_col = next((i for i, h in enumerate(hdr)
                         if "name" in h and "brand" not in h and "barcode" not in h), None)
    qty_col  = _find(("nominal", "quantity", "hacim", "ağırlık", "agirlik", "miktar", "(ml)"))
    has_header = any(c is not None for c in (bc_col, tr_col, sku_col, name_col))

    # TR-ad kaynağı: açık 'türkçe' kolonu > (barkod varsa) ürün-adı kolonu > 2. kolon.
    # Barkod-formatında (ör. fuar listeleri) "PRODUCT NAME" zaten Türkçe addır.
    if tr_col is None:
        if bc_col is not None and name_col is not None and name_col != bc_col:
            tr_col = name_col
        else:
            tr_col = 1
    # Eşleştirici: barkod > sku > İngilizce ad.
    if bc_col is not None:
        match_mode, id_col = "barcode", bc_col
    elif sku_col is not None:
        match_mode, id_col = "ident", sku_col
    elif name_col is not None and name_col != tr_col:
        match_mode, id_col = "ident", name_col
    else:
        match_mode, id_col = "ident", 0
    if match_mode == "ident" and id_col == tr_col:
        id_col, tr_col = 0, 1
    body = rows[1:] if has_header else rows

    from core.supplier_prices import normalize

    def _bc(v):
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    def _digits(s):
        return "".join(ch for ch in str(s) if ch.isdigit())

    by_sku, by_name, by_bc = {}, {}, {}
    for it in db.query(Item).filter(Item.is_active == True).all():
        if it.sku:
            by_sku[it.sku.strip().lower()] = it
        by_name.setdefault(normalize(it.name), it)
        if it.barcode:
            b = str(it.barcode).strip()
            by_bc[b] = it
            d = _digits(b)
            if d:
                by_bc.setdefault(d, it)

    updated, unmatched = 0, []
    for r in body:
        tr = str(r[tr_col]).strip() if (tr_col < len(r) and r[tr_col] is not None) else ""
        if not tr:
            continue
        # Barkod formatında ölçü/hacmi ada ekle (aynı ürün adının ölçü varyantlarını ayırır)
        if (match_mode == "barcode" and qty_col is not None
                and qty_col < len(r) and r[qty_col] not in (None, "")):
            size = " ".join(str(r[qty_col]).split())
            if size and size.lower() not in tr.lower():
                tr = f"{tr} {size}"
        if match_mode == "barcode":
            raw = _bc(r[id_col]) if id_col < len(r) else ""
            it = (by_bc.get(raw) or by_bc.get(_digits(raw))) if raw else None
            label = raw
        else:
            ident = str(r[id_col]).strip() if (id_col < len(r) and r[id_col] is not None) else ""
            it = (by_sku.get(ident.lower()) or by_name.get(normalize(ident))) if ident else None
            label = ident
        if not it:
            unmatched.append(label or tr)
            continue
        it.name_tr = tr[:150]
        updated += 1
    db.commit()
    return {"updated": updated, "matched": updated, "match_mode": match_mode,
            "unmatched": unmatched[:30], "unmatched_count": len(unmatched)}


def _item_delete_blockers(db: Session, item_id: int) -> list[str]:
    """
    HARD-delete'in mantıksal/şemasal olarak imkansız olduğu gerekçeleri döner.

    Audit kayıtları (transaction/inventory) ARTIK BURADA YOK — onlar için
    soft-delete devreye girer (bkz. delete_item).  Sadece şu durumlar HARD
    blok: silinirse veri orphan/corrupt olur.

    Kapsam:
      1) RecipeIngredient — bu ürünü malzeme olarak kullanan reçete varsa,
         silinirse reçete bozulur (NOT NULL constraint).
      2) Recipe.target_item_id — bu ürünün çıktı olduğu reçete var (silinirse
         reçete bir hedefe işaret edemez; nullable ama mantıksal kayıp).
      3) Item.parent_id — varyasyonlar parent'a bağlı, silinirse orphan.
    """
    blockers: list[str] = []

    # 1) Reçete malzemesi
    using = (
        db.query(Recipe.name)
        .join(RecipeIngredient, RecipeIngredient.recipe_id == Recipe.id)
        .filter(RecipeIngredient.item_id == item_id)
        .distinct()
        .limit(6)
        .all()
    )
    if using:
        names = ", ".join(r[0] for r in using[:5])
        more  = " …" if len(using) > 5 else ""
        blockers.append(f"şu reçetelerde malzeme: {names}{more}")

    # 2) Reçete hedefi
    targets = db.query(Recipe.name).filter(Recipe.target_item_id == item_id).limit(6).all()
    if targets:
        names = ", ".join(r[0] for r in targets[:5])
        more  = " …" if len(targets) > 5 else ""
        blockers.append(f"şu reçetenin hedef ürünü: {names}{more}")

    # 3) Varyasyonlar
    kids = db.query(Item.name).filter(Item.parent_id == item_id).limit(6).all()
    if kids:
        names = ", ".join(r[0] for r in kids[:5])
        more  = " …" if len(kids) > 5 else ""
        blockers.append(f"varyasyonu mevcut: {names}{more}")

    return blockers


def _item_has_audit(db: Session, item_id: int) -> bool:
    """Transaction veya inventory satırı varsa True — hard-delete yerine
    soft-delete kullanılır."""
    if db.query(Transaction.id).filter(Transaction.item_id == item_id).first():
        return True
    if db.query(Inventory.id).filter(Inventory.item_id == item_id).first():
        return True
    # Lotu başka karta taşınmış kart ("Karta taşı" / hedefli çevirme): kartta
    # Transaction da lot da kalmayabilir ama lot izi (`moved_from_item_id`, FK)
    # ona bakar — hard-delete FK ihlaliyle 500 verirdi; soft-delete, iz korunur.
    if db.query(Inventory.id).filter(Inventory.moved_from_item_id == item_id).first():
        return True
    # Numune analiz formunda bileşen olarak geçen kart (pending satırda ne
    # Transaction ne Inventory olabilir) — FK kırılmasın, soft-delete.
    from database import SampleAnalysisIngredient, ProductionConsumption
    from sqlalchemy import or_
    if db.query(SampleAnalysisIngredient.id).filter(SampleAnalysisIngredient.item_id == item_id).first():
        return True
    # Üretim tüketim dökümünde geçen kart — etiket kardeşi çözümünde reçetedeki
    # orijinal kart (ör. TR etiket) YALNIZ `recipe_item_id`'de durur, kartın
    # Transaction'ı/lotu olmayabilir; hard-delete FK ihlaliyle 500 verirdi.
    if db.query(ProductionConsumption.id).filter(
            or_(ProductionConsumption.item_id == item_id,
                ProductionConsumption.recipe_item_id == item_id)).first():
        return True
    # Fiyat satırı ya da kart kapsamlı tedarikçi tercihi olan kart — ikisi de
    # FK ile karta bakar (ondelete yok); hard-delete 500 verirdi, lab'ın fiyat
    # ve tercih kaydı da kaybolurdu.
    from database import MaterialSupplierPref, SupplierPrice
    if db.query(SupplierPrice.id).filter(SupplierPrice.item_id == item_id).first():
        return True
    if db.query(MaterialSupplierPref.id).filter(MaterialSupplierPref.item_id == item_id).first():
        return True
    return False


def _prune_material_groups(db: Session, group_ids) -> list:
    """Kart arşivlendi / silindi / gruptan düştü → etkilenen "aynı malzeme"
    grupları budanır: aktif üyesi 2'nin altına düşen grup dağılır
    (core.material_groups.prune — /api/material-groups uçlarıyla aynı kural;
    yoksa tek kartlık grup çipte durur, adı rezerve kalırdı).  Commit
    ÇAĞIRANA aittir.  Dönüş: dağılan grupların [{id, name}]'i."""
    out = []
    for gid in sorted({g for g in group_ids if g}):
        if material_groups.prune(db, gid):
            grp = db.query(MaterialGroup).filter(MaterialGroup.id == gid).first()
            out.append({"id": gid, "name": grp.name if grp else ""})
    return out


def _log_dissolved(db: Session, request: Request, actor: dict, dissolved: list,
                   reason: str, item_ids: list) -> None:
    """`_prune_material_groups`'un dağıttığı gruplar için audit (commit SONRASI)."""
    for g in dissolved:
        log_admin_event(db, request, actor=actor, action="material_group.dissolve",
                        target_type="material_group", target_id=g["id"], target_name=g["name"],
                        details={"sebep": reason, "item_ids": item_ids})


@router.delete("/items/{item_id}")
def delete_item(item_id: int, request: Request, db: Session = Depends(get_db),
                current_user: dict = Depends(require_permission("items", "delete"))):
    """
    Ürün silme — iki davranış birden:

    1) **Hard delete** — hiç bağ yoksa (audit / recipe / variation) satır
       DB'den kalkar.  Yeni oluşturulmuş hatalı ürünler için ideal.

    2) **Soft delete** — transaction veya inventory kaydı varsa
       `is_active=False` yapılır.  Ürün /api/items listesinden kaybolur
       (zaten is_active=True filtresi var), ama audit kayıtları + lot
       geçmişi DB'de okunur kalır (rapor / raporlama / restore için).

    3) **Block** — sadece reçete bağı veya varyasyon orphan riski olursa.
       Bu gerçekten silinmemeli; aksi halde reçeteler bozulur.

    Soft-delete kullanıcıya "Ürün arşivlendi" mesajı ile bildirilir; toast
    aynı yeşil tonda, lab fark etmez ama biz audit'i koruruz.

    Kart bir "aynı malzeme" grubundaysa ve grupta tek aktif kart kalırsa grup
    dağılır (`_prune_material_groups`).
    """
    item = db.query(Item).filter(Item.id == item_id).first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    blockers = _item_delete_blockers(db, item_id)
    if blockers:
        msg = "Bu ürün silinemez — " + "; ".join(blockers) + "."
        return JSONResponse(status_code=400, content={"detail": msg})

    gid = item.material_group_id
    if _item_has_audit(db, item_id):
        # Soft-delete: kayıtlar korunsun
        item.is_active = False
        dissolved = _prune_material_groups(db, [gid])
        db.commit()
        _log_dissolved(db, request, current_user, dissolved, "kart arşivlendi", [item_id])
        return {
            "message": f"'{item.name}' arşivlendi (geçmiş kayıtlar korundu).",
            "soft_deleted": True,
        }

    # Hard-delete: hiç bağ yok
    db.delete(item)
    dissolved = _prune_material_groups(db, [gid])
    db.commit()
    _log_dissolved(db, request, current_user, dissolved, "kart silindi", [item_id])
    return {"message": "Ürün silindi.", "soft_deleted": False}


@router.put("/items/{item_id}")
def update_item(
    item_id: int,
    data: ItemCreateRequest,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("items", "edit")),
    domain: str = Depends(active_domain),
):
    item = db.query(Item).filter(Item.id == item_id, Item.domain == domain).first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    if not data.force and _tr_fold(data.name) != _tr_fold(item.name):
        # Yalnız ad fiilen değişiyorsa kontrol et — aksi hâlde ürünün kendi
        # adı kendine "çakışma" olarak dönerdi.
        conflict = _find_name_conflict(db, domain, data.name, data.name_tr,
                                       exclude_id=item_id)
        if conflict:
            return JSONResponse(status_code=409, content={
                "detail": f"Aynı isimde ürün zaten kayıtlı: {conflict.name}",
                "existing": _conflict_payload(conflict),
            })

    # If converting to a variation, ensure this item itself has no children (would orphan them)
    if data.parent_id is not None and item.parent_id is None:
        has_children = db.query(Item).filter(Item.parent_id == item_id).count() > 0
        if has_children:
            return JSONResponse(status_code=400, content={
                "detail": "Bu ürünün varyasyonları mevcut. Önce varyasyonları silip sonra dönüştürebilirsiniz."
            })

    err = _validate_variation(db, data.parent_id, data.variation_name, self_id=item_id)
    if err: return err
    err = _check_supplier_ref(db, domain, data.supplier_id, current_id=item.supplier_id)
    if err: return err

    # Maliyet YALNIZ finans rolü alanı fiilen gönderdiyse yazılır.  Eskiden
    # `data.cost_price or 0.0` koşulsuzdu: items.html finans dışı kullanıcıda
    # alanı hiç göndermiyor (undefined), Pydantic varsayılanı 0.0 → LabLead
    # min stoğu düzeltip kaydedince Manager'ın girdiği maliyet 0'lanıyordu.
    write_cost = ("cost_price" in data.model_fields_set and data.cost_price is not None
                  and _can_see_finance(current_user))

    # "Aynı malzeme" grubu tek türdür (hammadde | ambalaj; bitmiş ürün hiç) —
    # gruplu kartın türü değişirse grup karışır ve /members her eklemede 400
    # verirdi.  Önce gruptan çıkarılmalı (düzenleme penceresindeki bölüm).
    if item.material_group_id and \
            stock_lots.lot_kind(data.category) != stock_lots.lot_kind(item.category):
        grp = (db.query(MaterialGroup)
               .filter(MaterialGroup.id == item.material_group_id,
                       MaterialGroup.is_active == True).first())             # noqa: E712
        if grp is not None:
            return JSONResponse(status_code=400, content={
                "detail": f"«{item.name}» «{grp.name}» aynı malzeme grubunda — kart türü "
                          f"(hammadde / ambalaj / bitmiş ürün) grupta değiştirilemez. Önce "
                          f"kartı gruptan çıkarın."})

    # ── Undo için BEFORE snapshot (mutation öncesi mevcut değerleri yakala) ──
    before_snapshot = {
        "name":            item.name,
        "category":        item.category,
        "unit":            item.unit,
        "min_stock_level": float(item.min_stock_level or 0.0),
        "parent_id":       item.parent_id,
        "variation_name":  item.variation_name,
        "barcode":         item.barcode,
        "pkg_type":        item.pkg_type,
        "language":        item.language,
        "label_group":     item.label_group,
        "supplier_id":     item.supplier_id,
    }
    # cost_price YALNIZ bu düzenleme ona dokunduysa geri alınır — yoksa
    # LabLead'in "Geri al"ı arada Manager'ın girdiği maliyeti sessizce eski
    # değere çekerdi (core/undo._undo_item_edit `if col in before`).
    if write_cost:
        before_snapshot["cost_price"] = float(item.cost_price or 0.0)

    item.name            = data.name
    item.name_tr         = ((data.name_tr or "").strip() or None)
    item.category        = data.category
    item.unit            = data.unit
    item.min_stock_level = data.min_stock  or 0.0
    if write_cost:
        item.cost_price  = data.cost_price
    item.parent_id       = data.parent_id
    item.variation_name  = (data.variation_name.strip() if data.parent_id and data.variation_name else None)
    item.barcode         = (data.barcode.strip() if data.barcode and data.barcode.strip() else None)
    item.pkg_type        = (data.pkg_type.strip().lower() if data.pkg_type and data.pkg_type.strip() else None)
    item.supplier_id     = data.supplier_id
    # Etiket dil/grup — sadece category=Ambalaj + pkg_type=etiket kombinasyonunda
    _is_label = (item.category == "Ambalaj" and item.pkg_type == "etiket")
    item.language    = _norm_language(data.language) if _is_label else None
    item.label_group = _label_group_key(data.name)   if _is_label else None

    # ── Undo entry ───────────────────────────────────────────────────────
    try:
        uid = int(current_user.get("sub", 0))
        if uid:
            record_undoable(
                db,
                user_id=uid,
                action_type="item_edit",
                target_table="items",
                target_id=item.id,
                # after_supplier_id: geri alırken bağ arada (tedarikçi birleştirmesi) değiştiyse ezilmesin
                payload={"item_id": item.id, "before": before_snapshot,
                         "after_supplier_id": item.supplier_id},
                description=f"Ürün düzenlendi: {item.name}",
            )
    except Exception:
        pass

    db.commit()

    # ── Low-stock alert: if the edit (typically a min_stock_level bump) leaves
    #     the item at or below threshold, queue a notification on the response.
    if item.min_stock_level > 0 and item.current_stock <= item.min_stock_level:
        background_tasks.add_task(
            notify_low_stock,
            item.name, item.current_stock, item.min_stock_level, item.unit or "",
        )

    return {"id": item.id, "message": "Ürün güncellendi."}


@router.get("/items/by-barcode/{barcode}")
def get_item_by_barcode(
    barcode: str,
    db: Session = Depends(get_db),
    _: dict = Depends(require_internal_user(("items", "view"), ("inventory", "view"))),
    domain: str = Depends(active_domain),
):
    """
    Server-side resolver — useful when the client doesn't have the full item
    list cached (e.g. mobile receiving flow on a slow connection). Frontend
    pages that already loaded /api/items can match locally without this call.
    Returns 404 when no active item carries the given barcode.
    """
    code = (barcode or "").strip()
    if not code:
        return JSONResponse(status_code=400, content={"detail": "Barkod boş."})
    item = (
        db.query(Item)
        .filter(Item.barcode == code, Item.is_active == True, Item.domain == domain)
        .first()
    )
    if not item:
        return JSONResponse(status_code=404, content={"detail": f"Bu barkoda sahip ürün yok: {code}"})
    return {
        "id":              item.id,
        "name":            item.name,
        "category":        item.category,
        "unit":            item.unit,
        "current_stock":   item.current_stock,
        "barcode":         item.barcode,
        "parent_id":       item.parent_id,
        "variation_name":  item.variation_name,
    }


@router.post("/items/bulk-delete")
def bulk_delete_items(data: BulkDeleteRequest, request: Request, db: Session = Depends(get_db),
                      current_user: dict = Depends(require_permission("items", "delete"))):
    """
    Toplu silme — single-delete'le aynı 3 davranış:
      • Reçete/varyasyon bağı → block, batch iptal
      • Audit kaydı → soft-delete (is_active=False)
      • Bağsız → hard-delete

    Hard-blok'lar mevcutsa hiçbir şey silinmez (yarım iş kalmasın).
    Aksi halde her ürün uygun yola gönderilir; kullanıcıya kaç hard +
    kaç soft yapıldığı raporlanır.  Etkilenen "aynı malzeme" grupları
    budanır (tek aktif kart kalan grup dağılır).
    """
    blocked:  list[str] = []   # "ürün adı (gerekçe)"
    soft_ids: list[int] = []
    hard_ids: list[int] = []
    group_ids: set = set()

    for iid in data.item_ids:
        it = db.query(Item).filter(Item.id == iid).first()
        if not it:
            continue
        if it.material_group_id:
            group_ids.add(it.material_group_id)
        reasons = _item_delete_blockers(db, iid)
        if reasons:
            blocked.append(f"{it.name} — {'; '.join(reasons)}")
            continue
        if _item_has_audit(db, iid):
            soft_ids.append(iid)
        else:
            hard_ids.append(iid)

    if blocked:
        msg = (
            f"{len(blocked)} ürün silinemediği için toplu silme iptal edildi:\n• "
            + "\n• ".join(blocked[:10])
            + ("\n…" if len(blocked) > 10 else "")
        )
        return JSONResponse(status_code=400, content={"detail": msg})

    hard_count = 0
    soft_count = 0
    if hard_ids:
        hard_count = db.query(Item).filter(Item.id.in_(hard_ids)).delete(synchronize_session=False)
    if soft_ids:
        soft_count = (
            db.query(Item).filter(Item.id.in_(soft_ids))
            .update({Item.is_active: False}, synchronize_session=False)
        )
    # Toplu UPDATE/DELETE oturumdaki nesneleri tazelemez (synchronize_session=False)
    # — budama sorguları DB'den okur ama oturumdaki eski nesneler karışmasın.
    db.expire_all()
    dissolved = _prune_material_groups(db, group_ids)
    db.commit()
    _log_dissolved(db, request, current_user, dissolved, "toplu silme",
                   sorted(set(soft_ids) | set(hard_ids)))

    parts = []
    if hard_count: parts.append(f"{hard_count} ürün silindi")
    if soft_count: parts.append(f"{soft_count} ürün arşivlendi (geçmiş korundu)")
    return {"message": ", ".join(parts) + ".", "hard": hard_count, "soft": soft_count}


# ─── Suppliers Endpoints ─────────────────────────────────────────────────────
# Silme YUMUŞAKTIR (07.10.2026): tedarikçiye lot, kart, fiyat, sipariş ve
# üretim tüketimi bağlıdır — kalıcı silme ya IntegrityError (500) veriyordu ya
# da izi koparırdı.  Pasif kart listede gizlenir (`?include_inactive=1` ile
# görünür), adı kart/lot listelerinde çözülmeye devam eder, yeni atamada
# kabul edilmez (`_check_supplier_ref`).  Geri dönüş: POST .../activate.

def _clean(v) -> Optional[str]:
    return (v or "").strip() or None


def _check_supplier_ref(db: Session, domain: str, supplier_id: Optional[int],
                        current_id: Optional[int] = None) -> Optional[JSONResponse]:
    """Karta/lota yazılacak `supplier_id`'yi doğrula; sorun yoksa None.

    Aktif panelin tedarikçisi olmalı (başka panelin firması bağlanırsa P1b
    birleştirmesi lotu taşıyamaz, raporlarda panel karışır) ve AKTİF olmalı —
    pasife alınmış firma yeniden atanamaz.  İstisna: kartın DEĞİŞMEYEN mevcut
    tedarikçisi (`current_id`) olduğu gibi kabul edilir — pasif ya da eski
    (panel öncesi) bir bağ olsa da; yoksa o kartı düzenleyip kaydetmek
    imkânsız olurdu.

    Satır FOR SHARE ile okunur: eşzamanlı bir tedarikçi birleştirmesi
    (FOR UPDATE) commit edene kadar bekler ve pasifleşmiş kaybedeni görür —
    yoksa kontrolden sonra yazılan lot/kart pasif kaybedene bağlı kalırdı."""
    if supplier_id is None or (current_id is not None and supplier_id == current_id):
        return None
    sup = (db.query(Supplier)
           .filter(Supplier.id == supplier_id, Supplier.domain == domain)
           .with_for_update(read=True).populate_existing().first())
    if sup is None:
        return JSONResponse(status_code=400, content={
            "detail": "Seçilen tedarikçi bu panelde bulunamadı."})
    if sup.is_active is False:
        return JSONResponse(status_code=400, content={
            "detail": f"«{sup.name}» pasif tedarikçi — yeni kayda bağlanamaz."})
    return None


def _supplier_names(db: Session, domain: str) -> dict:
    """Panelin BÜTÜN tedarikçileri (pasifler dahil) → {id: ad}.  Kart/lot
    listelerindeki ad çözümü için — pasife alınan firmanın adı kartlarda
    sessizce kaybolmasın."""
    return dict(db.query(Supplier.id, Supplier.name).filter(Supplier.domain == domain).all())


@router.get("/suppliers")
def list_suppliers(include_inactive: bool = False,
                   include_ids: Optional[str] = None,
                   db: Session = Depends(get_db), domain: str = Depends(active_domain),
                   _: dict = Depends(require_any_permission(("items", "view"), ("inventory", "view"),
                                                            ("reports", "view")))):
    """Aktif panelin tedarikçileri (varsayılan yalnız aktifler).

    `include_inactive=1` pasifleri de döndürür (Tedarikçiler sayfasının
    "Pasifleri göster" anahtarı).  `include_ids=3,7` yalnız bu id'lerdeki
    pasifleri ekler — kart penceresi, kartın mevcut (pasif) tedarikçisini
    seçenek olarak tutabilsin.  Yanıt düz dizi; satır şekli
    `core.suppliers.serialize_supplier`."""
    from core.suppliers import serialize_supplier
    extra: set = set()
    for x in (include_ids or "").split(","):
        x = x.strip()
        if x.isdigit():
            extra.add(int(x))
    q = db.query(Supplier).filter(Supplier.domain == domain)
    if not include_inactive:
        if extra:
            q = q.filter((Supplier.is_active == True) | Supplier.id.in_(sorted(extra)))  # noqa: E712
        else:
            q = q.filter(Supplier.is_active == True)                                    # noqa: E712
    return [serialize_supplier(s) for s in q.order_by(Supplier.id.desc()).all()]


@router.post("/suppliers", status_code=201)
def create_supplier(data: SupplierCreateRequest, request: Request, db: Session = Depends(get_db),
                    current_user: dict = Depends(require_permission("items", "create")),
                    domain: str = Depends(active_domain)):
    name = data.name.strip()
    if not name:
        return JSONResponse(status_code=400, content={"detail": "Firma adı zorunludur."})
    supplier = Supplier(
        name=name,
        contact_person=_clean(data.contact_person),
        email=_clean(data.email),
        phone=_clean(data.phone),
        address=_clean(data.address),
        notes=_clean(data.notes),
        domain=domain,                     # Faz 3 — aktif panele damgala
    )
    db.add(supplier)
    db.commit()
    db.refresh(supplier)
    log_admin_event(db, request, actor=current_user, action="supplier.create",
                    target_type="supplier", target_id=supplier.id, target_name=supplier.name,
                    details={"domain": domain})
    return {"id": supplier.id, "message": "Tedarikçi başarıyla eklendi."}


@router.put("/suppliers/{supplier_id}")
def update_supplier(supplier_id: int, data: SupplierCreateRequest, request: Request,
                    db: Session = Depends(get_db),
                    current_user: dict = Depends(require_permission("items", "edit")),
                    domain: str = Depends(active_domain)):
    """Tedarikçi bilgilerini düzenle.  Sayfada düzenleme hiç yoktu (yalnız ekle/sil) —
    yanlış girilen firma adı/telefonu silip yeniden eklemek gerekiyordu, bu da kayda
    bağlı mal kabul lotlarını koparırdı.  Aktif panelin (domain) dışındaki tedarikçi
    değiştirilemez.  `address` yalnız gövdede varsa yazılır.  Audit
    `supplier.update` (değişen alanlar eski → yeni)."""
    supplier = (db.query(Supplier)
                .filter(Supplier.id == supplier_id, Supplier.is_active == True,   # noqa: E712
                        Supplier.domain == domain)
                .first())
    if not supplier:
        return JSONResponse(status_code=404, content={"detail": "Tedarikçi bulunamadı."})
    name = data.name.strip()
    if not name:
        return JSONResponse(status_code=400, content={"detail": "Firma adı zorunludur."})
    new = {"name": name, "contact_person": _clean(data.contact_person),
           "email": _clean(data.email), "phone": _clean(data.phone), "notes": _clean(data.notes)}
    if "address" in data.model_fields_set:
        new["address"] = _clean(data.address)
    changes = {}
    for col, val in new.items():
        old = getattr(supplier, col)
        if old != val:
            changes[col] = {"eski": old, "yeni": val}
            setattr(supplier, col, val)
    db.commit()
    if changes:
        log_admin_event(db, request, actor=current_user, action="supplier.update",
                        target_type="supplier", target_id=supplier.id, target_name=supplier.name,
                        details={"changes": changes})
    return {"id": supplier.id, "message": "Tedarikçi güncellendi."}


def _log_deactivated(db: Session, request: Request, actor: dict, sup: Supplier, via: str) -> None:
    log_admin_event(db, request, actor=actor, action="supplier.deactivate",
                    target_type="supplier", target_id=sup.id, target_name=sup.name,
                    details={"via": via, "domain": sup.domain,
                             "purchase_status": sup.purchase_status})


@router.delete("/suppliers/{supplier_id}")
def delete_supplier(supplier_id: int, request: Request, db: Session = Depends(get_db),
                    current_user: dict = Depends(require_permission("items", "delete")),
                    domain: str = Depends(active_domain)):
    """Tedarikçiyi PASİFE al (yumuşak silme) — lot/kart/fiyat bağları korunur."""
    supplier = (db.query(Supplier)
                .filter(Supplier.id == supplier_id, Supplier.domain == domain).first())
    if not supplier:
        return JSONResponse(status_code=404, content={"detail": "Tedarikçi bulunamadı."})
    if supplier.is_active is False:
        return {"message": "Tedarikçi zaten pasif."}
    supplier.is_active = False
    db.commit()
    _log_deactivated(db, request, current_user, supplier, "tekil")
    return {"message": "Tedarikçi pasife alındı."}


@router.post("/suppliers/bulk-delete")
def bulk_delete_suppliers(data: SupplierBulkDeleteRequest, request: Request,
                          db: Session = Depends(get_db),
                          current_user: dict = Depends(require_permission("items", "delete")),
                          domain: str = Depends(active_domain)):
    """Seçilenleri pasife al — yalnız aktif paneldeki AKTİF kayıtlar; diğer
    id'ler sessizce atlanır (sayıya girmez)."""
    rows = (db.query(Supplier)
            .filter(Supplier.id.in_(data.supplier_ids), Supplier.domain == domain,
                    Supplier.is_active == True)                                  # noqa: E712
            .order_by(Supplier.id).all())
    for s in rows:
        s.is_active = False
    db.commit()
    for s in rows:
        _log_deactivated(db, request, current_user, s, "toplu")
    return {"message": f"{len(rows)} tedarikçi pasife alındı.", "deactivated": len(rows)}


@router.post("/suppliers/{supplier_id}/activate")
def activate_supplier(supplier_id: int, request: Request, db: Session = Depends(get_db),
                      current_user: dict = Depends(require_permission("items", "delete")),
                      domain: str = Depends(active_domain)):
    """Yanlışlıkla pasife alınan tedarikçiyi geri aç (pasife almanın tersi,
    aynı yetki).  Birleştirilmiş kart açılırsa yeniden kendi firmasıdır —
    `merged_into_id` (takma ad bağı) temizlenir; taşınmış kayıtlar kazananda
    kalır.  Audit `supplier.activate`."""
    supplier = (db.query(Supplier)
                .filter(Supplier.id == supplier_id, Supplier.domain == domain).first())
    if not supplier:
        return JSONResponse(status_code=404, content={"detail": "Tedarikçi bulunamadı."})
    if supplier.is_active is not False:
        return {"message": "Tedarikçi zaten aktif."}
    supplier.is_active = True
    supplier.merged_into_id = None
    db.commit()
    log_admin_event(db, request, actor=current_user, action="supplier.activate",
                    target_type="supplier", target_id=supplier.id, target_name=supplier.name,
                    details={"domain": domain})
    return {"message": "Tedarikçi yeniden etkinleştirildi."}


# ─── Inventory / Receiving Endpoints ────────────────────────────────────────

@router.get("/inventory")
def list_inventory(db: Session = Depends(get_db), domain: str = Depends(active_domain),
                   _: dict = Depends(require_permission("inventory", "view"))):
    # joinedload — item + supplier ilişkileri tek query'de gelir (N+1 önler)
    rows = (
        db.query(Inventory)
        .options(joinedload(Inventory.item), joinedload(Inventory.supplier))
        .filter(Inventory.domain == domain)
        .order_by(Inventory.id.desc())
        .all()
    )
    return [
        {
            "id": r.id,
            "item_name": r.item.name if r.item else "—",
            "item_id": r.item_id,
            "supplier_name": r.supplier.name if r.supplier else "—",
            "lot_number": r.lot_number,
            "expiry_date": r.expiry_date or "—",
            "quantity": r.quantity,
            "location": r.location or "—",
            "status": r.status,
            "created_at": to_tr(r.created_at).strftime("%d.%m.%Y") if r.created_at else "",
        }
        for r in rows
    ]


@router.post("/inventory/receive", status_code=201)
def receive_stock(
    data: StockReceiveRequest,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "receive")),
    domain: str = Depends(active_domain),
):
    if not _finite(data.quantity) or data.quantity <= 0:
        return JSONResponse(status_code=400, content={"detail": "Miktar sıfırdan büyük olmalıdır."})

    item = db.query(Item).filter(Item.id == data.item_id).with_for_update().first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})
    if (item.domain or "cosmetics") != domain:
        return JSONResponse(status_code=400, content={"detail": "Bu ürün aktif panelde değil."})
    err = _check_supplier_ref(db, domain, data.supplier_id)
    if err:
        return err

    actor = current_user.get("full_name") or current_user.get("username") or "—"

    is_sample = bool(data.is_sample)
    # Numune için varsayılan konum "Numune" (kullanıcı vermezse) — stok
    # sayfasında normal stoktan görsel olarak ayrışsın.
    location = data.location or ("Numune" if is_sample else None)

    try:
        # ── Upsert Logic (4.7) ──────────────────────────────────────────────
        # Numune lotları normal lotlarla BİRLEŞMEZ — is_sample da eşleşme
        # anahtarına dahil.  DİKKAT: supplier_id anahtara DAHİL DEĞİL — aynı
        # lot no farklı tedarikçiden ikinci kez gelirse bu satıra birleşir
        # (ilk tedarikçi COALESCE ile korunur, ikincisi sessizce atlanır).
        existing = (db.query(Inventory)
                   .filter(Inventory.item_id == data.item_id,
                           Inventory.lot_number == data.lot_number,
                           Inventory.outsourcing_receipt_id.is_(None),
                           Inventory.is_sample == is_sample)
                   .with_for_update().first())

        new_inventory: Optional[Inventory] = None
        stock_before = float(item.current_stock or 0.0)

        if existing:
            # Miktarı topla; location her zaman güncellenir, supplier_id/expiry
            # yalnız BOŞSA doldurulur (COALESCE — mevcut tedarikçi ezilmez).
            existing.quantity   += data.quantity
            existing.updated_at  = __import__("datetime").datetime.utcnow()
            if location:
                existing.location = location
            if data.supplier_id is not None and existing.supplier_id is None:
                existing.supplier_id = data.supplier_id
            if data.expiry_date and not existing.expiry_date:
                existing.expiry_date = data.expiry_date
            # received_by sadece ilk kabul edende kalır (audit immutability)
        else:
            # Yeni satır — statü kesinlikle APPROVED
            new_inventory = Inventory(
                item_id=data.item_id,
                supplier_id=data.supplier_id,
                lot_number=data.lot_number,
                expiry_date=data.expiry_date,
                quantity=data.quantity,
                location=location,
                status="APPROVED",
                received_by=actor,                      # Audit trail
                is_sample=is_sample,
                domain=(item.domain or "cosmetics"),    # Faz 3 — lot, item ile aynı panelde
            )
            db.add(new_inventory)

        if is_sample:
            # ── Numune STOK DEĞİLDİR (2026-08-24 olayı) ─────────────────────
            # Eskiden numune de current_stock'a eklenip Input Transaction'ı
            # yazıyordu — kritik-stok uyarısını maskeledi, üretim FIFO'suna
            # karışabildi, Shopify'a satılabilir stok diye gitti.  Artık
            # Inventory satırı (numune envanteri) yazılır ama STOK ve DEFTER'e
            # dokunulmaz; denetim izi audit log'dadır.  Gerçek stoğa geçmek
            # için POST /inventory/samples/{id}/convert kullanılır.
            db.flush()
            log_admin_event(db, request, actor=current_user, action="inventory.sample_receive",
                            target_type="inventory", target_id=(existing or new_inventory).id,
                            target_name=item.name,
                            details={"lot": data.lot_number, "quantity": data.quantity,
                                     "unit": item.unit, "supplier_id": data.supplier_id})
            db.commit()
            return {"message": f"Numune kabul başarılı. {data.quantity} {item.unit} "
                               f"numune kaydedildi (stok toplamına dahil edilmez)."}

        # ── Transaction kaydı (2.4) — yalnız GERÇEK stok girişi ─────────────
        tx = Transaction(
            item_id=data.item_id,
            lot_number=data.lot_number,
            transaction_type="Input",
            quantity=data.quantity,
            notes=f"Mal kabul — Lot: {data.lot_number}" + (f", Konum: {location}" if location else ""),
            performed_by=actor,                          # Audit trail
        )
        db.add(tx)

        # ── items.current_stock güncelle (üretim modülü ile uyum) ───────────
        item.current_stock = round(item.current_stock + data.quantity, 6)
        stock_after = float(item.current_stock)

        # ── Undo log — sadece yeni lot (upsert değil) için ─────────────────
        # Upsert durumunda undo karmaşık (mevcut lot'tan subtract); şimdilik skip.
        if new_inventory is not None:
            try:
                db.flush()
                uid = int(current_user.get("sub", 0))
                if uid:
                    record_undoable(
                        db,
                        user_id=uid,
                        action_type="inventory_receive",
                        target_table="inventory",
                        target_id=new_inventory.id,
                        payload={
                            "inventory_id":      new_inventory.id,
                            "transaction_id":    tx.id,
                            "item_id":           item.id,
                            "received_quantity": float(data.quantity),
                            "item_stock_before": stock_before,
                            "item_stock_after":  stock_after,
                            "lot_number":        data.lot_number,
                        },
                        description=(
                            f"Lot kabul: {data.quantity} {item.unit or ''} "
                            f"{item.name} (Lot {data.lot_number})"
                        ),
                    )
            except Exception:
                pass

        # Eksik Hammaddeler raporunda açık "sipariş verildi" bayrağı varsa
        # gerçek mal kabulü kapatır (core/stock_gaps.close_open_flags) — elle
        # kaldırmaya gerek kalmaz.
        from core.stock_gaps import close_open_flags
        close_open_flags(db, item.id, "received", actor)

        db.commit()
        return {"message": f"Mal kabul başarılı. {data.quantity} {item.unit} stoka eklendi."}

    except Exception:
        db.rollback()
        return JSONResponse(status_code=500, content={"detail": "Mal kabul sırasında hata oluştu."})


# ─── Stock Adjustment (manual correction) ──────────────────────────────────
# Black-box invariant: stock changes outside of receiving/production/QC must
# still leave a fingerprint. This endpoint never deletes or rewrites history;
# it appends an immutable Transaction(type='Adjustment') and updates the
# canonical Item.current_stock to the requested value. The reason string is
# required so audit reviewers can see *why* the correction happened.

@router.post("/inventory/adjust", status_code=201)
def adjust_stock(
    data: StockAdjustRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "adjust")),
):
    if not _finite(data.new_quantity):
        return JSONResponse(status_code=400, content={"detail": "Stok miktarı geçerli bir sayı olmalıdır."})
    if data.new_quantity < 0:
        return JSONResponse(status_code=400, content={"detail": "Stok miktarı negatif olamaz."})

    reason = (data.reason or "").strip()
    if len(reason) < 3:
        return JSONResponse(status_code=422, content={"detail": "Sebep en az 3 karakter olmalıdır."})

    item = db.query(Item).filter(Item.id == data.item_id).with_for_update().first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    actor    = current_user.get("full_name") or current_user.get("username") or "—"
    old_qty  = float(item.current_stock or 0)
    new_qty  = float(data.new_quantity)
    delta    = round(new_qty - old_qty, 6)

    if abs(delta) < 1e-9:
        return JSONResponse(status_code=400, content={
            "detail": f"Yeni miktar mevcut stokla aynı ({old_qty} {item.unit or ''}). Düzeltme gerekmez."
        })

    try:
        # ── Item.current_stock güncelle ────────────────────────────────────
        item.current_stock = round(new_qty, 6)

        # ── Lot-level adjustment opsiyonel: belirli bir lot'un quantity'sini
        #     hedef değere düşür/yükselt. Yoksa sadece item-level düzeltme.
        lot_note = ""
        if data.lot_number:
            inv = (
                db.query(Inventory)
                .filter(Inventory.item_id == data.item_id, Inventory.lot_number == data.lot_number)
                .with_for_update()
                .first()
            )
            if inv:
                if inv.outsourcing_receipt_id and (inv.status == "QUARANTINE" or inv.qc_required):
                    db.rollback()
                    return JSONResponse(status_code=400, content={
                        "detail": "Fason kabul lotu kalite kontrolü tamamlanmadan stok düzeltmesine konu olamaz."
                    })
                # Bu lot için yeni miktar mantıklı mı kontrol etmiyoruz — kullanıcı
                # zaten gerekçeyi girdi. Sadece lot satırını da güncelleyelim.
                inv.quantity = round(new_qty, 6)
                inv.updated_at = __import__("datetime").datetime.utcnow()
                lot_note = f" | Lot: {data.lot_number}"
        elif delta < 0:
            # Lot SEÇİLMEDİ ve stok DÜŞÜYOR → lotlardan da FIFO düş.  Aksi hâlde
            # `current_stock` iner ama lot satırları olduğu yerde kalır ve iki
            # sayaç kalıcı olarak ayrışır (09.09.2026 tespiti, 25 üründe 274
            # adet sapma).  Defter kaydı yine TEK imzalı Adjustment'tır —
            # birleştirme script'leri bu konvansiyona dayanıyor.
            from core.stock_lots import draw_down, lot_summary
            touched, uncovered = draw_down(db, item, -delta)
            if touched:
                lot_note = f" | Lot düşümü: {lot_summary(touched)}"
            if uncovered > 1e-9:
                lot_note += f" | lot kaydı dışı: {round(uncovered, 6):g}"

        # ── Immutable audit kaydı (delta hem +/- olabilir) ─────────────────
        sign = "+" if delta > 0 else ""
        tx = Transaction(
            item_id=item.id,
            lot_number=data.lot_number,
            transaction_type="Adjustment",
            quantity=delta,            # signed delta (-8.0 veya +3.0)
            notes=(
                f"Stok düzeltme — Eski: {old_qty} {item.unit or ''} → "
                f"Yeni: {new_qty} {item.unit or ''} "
                f"(Δ {sign}{delta}){lot_note} | Sebep: {reason[:200]}"
            ),
            performed_by=actor,
        )
        db.add(tx)
        db.flush()    # tx.id'i undo payload'a koyabilmek için

        # ── Undo log — Ctrl+Z için ─────────────────────────────────────────
        try:
            uid = int(current_user.get("sub", 0))
            if uid:
                record_undoable(
                    db,
                    user_id=uid,
                    action_type="stock_adjust",
                    target_table="items",
                    target_id=item.id,
                    payload={
                        "item_id":        item.id,
                        "before_stock":   old_qty,
                        "after_stock":    new_qty,
                        "transaction_id": tx.id,
                    },
                    description=f"Stok düzeltildi: {old_qty} → {new_qty} {item.unit or ''} ({item.name})",
                )
        except Exception:
            pass  # Undo başarısız olsa ana mutation hâlâ commit edilir.

        db.commit()
        return {
            "message": f"Stok düzeltildi: {old_qty} → {new_qty} {item.unit or ''} (Δ {sign}{delta}).",
            "item_id":    item.id,
            "old_qty":    old_qty,
            "new_qty":    new_qty,
            "delta":      delta,
        }

    except Exception:
        db.rollback()
        return JSONResponse(status_code=500, content={"detail": "Stok düzeltme sırasında hata oluştu."})


# ─── Inventory summary + transactions feed ──────────────────────────────────

@router.get("/inventory/summary")
def inventory_summary(db: Session = Depends(get_db), domain: str = Depends(active_domain),
                      _: dict = Depends(require_any_permission(("inventory", "view"), ("reports", "view")))):
    """
    Mevcut stok özeti — `Item.current_stock` source-of-truth olarak
    kullanılır. Önceden APPROVED Inventory satırlarının quantity'lerini
    topluyorduk; ama production output'u eskiden QUARANTINE'de kalıyordu
    ve current_stock'a yansımıyordu. Phase 8 / Bug 4'le production direkt
    current_stock'u arttırıyor — bu endpoint de aynı kanonik değeri okur,
    böylece /stocks ve canlı önizleme birbirine uyumlu.
    """
    items = (
        db.query(Item)
        .filter(Item.is_active == True, Item.domain == domain)
        .order_by(Item.category, Item.name)
        .all()
    )
    # Supplier lookup — listing'de supplier_name göstermek için tek query
    # (pasife alınan tedarikçinin adı da — kart tedarikçisiz görünmesin)
    supplier_name_by_id = _supplier_names(db, domain)
    return [
        {
            "item_id":       i.id,
            "name":          i.name,
            "category":      i.category or "Diğer",
            "unit":          i.unit,
            "pkg_type":      i.pkg_type or "",        # /stocks Etiket vs Ambalaj ayırımı için
            "total_stock":   round(float(i.current_stock or 0), 4),
            "supplier_id":   i.supplier_id,
            "supplier_name": supplier_name_by_id.get(i.supplier_id) if i.supplier_id else None,
        }
        for i in items
    ]


@router.get("/inventory/by-item/{item_id}")
def inventory_by_item(
    item_id: int,
    db: Session = Depends(get_db),
    _: dict = Depends(require_internal_user(("inventory", "view"), ("reports", "view"))),
    domain: str = Depends(active_domain),
):
    """
    Tek bir ürünün lot bazlı envanter dökümü — Stoklar sayfasında satıra
    tıklayınca açılan detay.  Hangi lot nerede (Showroom / Şahit Numune
    Dolabı / diğer konum), ne kadarı, hangi durumda.

    Lokasyon kırılımı `by_location` alanında özetlenir — üretimde şahit
    numune ayrımı yapılan ürünlerde "kaçı showroom, kaçı şahit" tek bakışta
    görülür.  Her lot `inventory_id` + `supplier_id` taşır ("Taşı" düğmesi
    POST /inventory/lots/{id}/move'a gider); başka karttan taşındıysa
    `moved_from` ilk kartı gösterir.  `item_group` kartın "aynı malzeme"
    grubu ve üyeleridir (yoksa null).
    """
    item = db.query(Item).filter(Item.id == item_id, Item.domain == domain).first()
    if not item:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    rows = (
        db.query(Inventory)
        .options(joinedload(Inventory.supplier))
        .filter(Inventory.item_id == item_id, Inventory.domain == domain)
        .order_by(Inventory.id.desc())
        .all()
    )
    moved_ids = {r.moved_from_item_id for r in rows if r.moved_from_item_id}
    moved_names = (dict(db.query(Item.id, Item.name).filter(Item.id.in_(moved_ids)).all())
                   if moved_ids else {})

    lots = []
    by_location: dict[str, float] = {}
    for r in rows:
        loc = (r.location or "").strip() or "(Konum belirtilmemiş)"
        qty = float(r.quantity or 0)
        by_location[loc] = round(by_location.get(loc, 0.0) + qty, 4)
        lots.append({
            "inventory_id":  r.id,
            "lot_number":    r.lot_number,
            "location":      loc,
            "quantity":      round(qty, 4),
            "status":        r.status or "",
            "qc_required":   bool(r.qc_required),
            "is_sample":     bool(r.is_sample),
            "supplier_id":   r.supplier_id,
            "supplier_name": (r.supplier.name if r.supplier else "—"),
            "expiry_date":   r.expiry_date or "",
            "received_by":   r.received_by or "",
            "created_at":    to_tr(r.created_at).strftime("%d.%m.%Y %H:%M") if r.created_at else "",
            "moved_from":    ({"id": r.moved_from_item_id,
                               "name": moved_names.get(r.moved_from_item_id, "—")}
                              if r.moved_from_item_id else None),
            "sample_converted_at": (to_tr(r.sample_converted_at).strftime("%d.%m.%Y %H:%M")
                                    if r.sample_converted_at else ""),
        })

    # lot_total STOK karşılığı — numune satırları hariç (numune current_stock'a
    # hiç girmiyor, dahil edilirse her numunesi olan üründe kalıcı "stok ile
    # lot toplamı uyuşmuyor" uyarısı basardı).  sample_total ayrı gösterilir.
    non_sample_total = round(sum(l["quantity"] for l in lots if not l["is_sample"]), 4)
    sample_total = round(sum(l["quantity"] for l in lots if l["is_sample"]), 4)

    return {
        "item_id":       item.id,
        "item_name":     item.name,
        "category":      item.category or "",
        "unit":          item.unit or "",
        # Item.current_stock source-of-truth; lot toplamı bundan sapabilir
        # (ör. üretim çıktısı doğrudan current_stock'a yazılır).
        "current_stock": round(float(item.current_stock or 0), 4),
        "lot_total":     non_sample_total,
        "sample_total":  sample_total,
        "by_location":   [
            {"location": loc, "total": tot}
            for loc, tot in sorted(by_location.items(), key=lambda x: -x[1])
        ],
        "lots":          lots,
        "item_group":    material_groups.group_payload(db, item.material_group_id),
    }


@router.get("/inventory/samples")
def list_samples(include_empty: bool = False,
                 db: Session = Depends(get_db),
                 _: dict = Depends(require_internal_user(("items", "view"), ("qc", "view"))),
                 domain: str = Depends(active_domain)):
    """
    Numune lotları — Ürünler sayfası "Numune" sekmesini ve Numune Analizi
    lot seçicisini besler.  Var olan hammaddelere bağlı, alternatif
    tedarikçilerden gelen numune partileri.

    Varsayılan yalnız miktarı kalan (> 0) satırlar — analizde tükenmiş
    numuneler listeyi doldurmasın; `?include_empty=1` hepsini verir.  Pasif
    karta bağlı numune GİZLENMEZ (`item_active=false` ile işaretli) — yoksa
    lot görünmez olur, kimse taşıyamaz.  `card_supplier_*` kartın kendi
    tedarikçisidir: numunenin tedarikçisinden farklıysa (`supplier_mismatch`)
    numune büyük ihtimalle yanlış tedarikçinin kartına girilmiştir (Naturalya
    jojobası KRK GIDA kartında, 05.10) → arayüz "Karta taşı" önerir.
    """
    q = (
        db.query(Inventory)
        .options(joinedload(Inventory.item).joinedload(Item.supplier),
                 joinedload(Inventory.supplier))
        .filter(Inventory.is_sample == True, Inventory.domain == domain)   # noqa: E712
    )
    if not include_empty:
        q = q.filter(Inventory.quantity > 0)
    rows = q.order_by(Inventory.id.desc()).all()
    group_name_by_id = dict(
        db.query(MaterialGroup.id, MaterialGroup.name)
        .filter(MaterialGroup.domain == domain, MaterialGroup.is_active == True).all()  # noqa: E712
    )
    out = []
    for r in rows:
        it = r.item
        card_sup_id = it.supplier_id if it else None
        gid = it.material_group_id if it else None
        out.append({
            "inventory_id":  r.id,
            "item_id":       r.item_id,
            "item_name":     it.name if it else "—",
            "item_active":   bool(it.is_active) if it else False,
            "unit":          (it.unit if it else "") or "",
            "supplier_id":   r.supplier_id,
            "supplier_name": (r.supplier.name if r.supplier else "—"),
            "card_supplier_id":   card_sup_id,
            "card_supplier_name": (it.supplier.name if it and it.supplier else None),
            "supplier_mismatch":  bool(r.supplier_id and card_sup_id and r.supplier_id != card_sup_id),
            "material_group_id":   gid if gid in group_name_by_id else None,
            "material_group_name": group_name_by_id.get(gid),
            "moved_from_item_id":  r.moved_from_item_id,
            "lot_number":    r.lot_number,
            "quantity":      round(float(r.quantity or 0), 4),
            "expiry_date":   r.expiry_date or "—",
            "created_at":    to_tr(r.created_at).strftime("%d.%m.%Y") if r.created_at else "—",
        })
    return out


# ─── Numune → stok: hedef seçimi + çift sayım koruması (06.10.2026) ─────────
# 05.10'da Naturalya'nın jojoba/portakal/lavanta numuneleri KRK GIDA ve İPEDA
# kartlarına girilmiş olduğu için "Stoğa çevir" stoğu yanlış tedarikçinin
# kartına yazdı; Songül Hanım 5 dk sonra Naturalya kartını açtı, o kart 0
# gösterdi.  Ayrıca 758 JOJOBA UÇUCU YAĞI kartında numune zaten 10.09'da elle
# +20 sayılmıştı, çevirme +20 daha ekledi (stok 40, lot 20).  Çevirme artık
# hedef kart (aynı malzeme grubundan / yeni kart) seçtirir ve hedefte numune
# geldikten sonra yapılmış elle "Stok düzeltme"leri gösterip onay ister.

#: Tek bir elle düzeltme numune miktarının bu oranı içindeyse "aynı numune
#: zaten sayılmış olabilir" denir (+20 düzeltme ↔ 20 ml numune).
COUNT_GUARD_TOLERANCE = 0.05


class SampleConvertNewItem(BaseModel):
    name: str = Field(..., min_length=1, max_length=150)
    name_tr: Optional[str] = Field(None, max_length=150)
    force: bool = False                  # ad çakışması (409) görüldükten sonra


class SampleConvertBody(BaseModel):
    """Numune çevirme seçenekleri — HEPSİ isteğe bağlı.  Gövdesiz eski çağrı
    "bu numuneyi bağlı olduğu karta stok olarak ekle" demektir."""
    target_item_id: Optional[int] = None             # başka (aynı birimli) kart
    new_item: Optional[SampleConvertNewItem] = None  # numunenin tedarikçisiyle yeni kart
    # add: Input + current_stock (bugünkü davranış).  link_only: numune zaten
    # elle sayılmış — yalnız lot normal lota döner, defter/stok DEĞİŞMEZ.
    mode: Literal["add", "link_only"] = "add"
    acknowledge_counted: bool = False                # çift sayım uyarısını gördüm, yine de ekle
    # Yalnız bu kartların uyarısı görüldü (arayüz).  Hedef değişince yeni bir
    # kartın bulgusu çıkarsa yine 409 — görülmemiş uyarı onaylanmış sayılmaz.
    acknowledged_item_ids: List[int] = Field(default_factory=list, max_length=20)
    lot_number: Optional[str] = Field(None, max_length=100)   # lot_collision çözümü


def _count_guard(db: Session, item_id: int, since: Optional[datetime], qty: float) -> dict:
    """Numune zaten elle sayılmış olabilir mi?

    Kartta numune geldikten SONRA (`since` = numune satırının created_at'i)
    yazılmış "Stok düzeltme…" Adjustment'larına bakar (adjust_stock'un not
    öneki; lot taşıma ve çevirme bu öneki KULLANMAZ).  Tek bir pozitif kayıt
    numune miktarının ±%5'i içindeyse ya da bu kayıtların net toplamı numune
    miktarını karşılıyorsa bayrak kalkar.  Numuneden önceki düzeltmeler
    sayılmaz — numune o gün yoktu.
    """
    q = (db.query(Transaction)
         .filter(Transaction.item_id == item_id,
                 Transaction.transaction_type == "Adjustment",
                 Transaction.notes.like("Stok düzeltme%")))
    if since is not None:
        q = q.filter(Transaction.timestamp >= since)
    rows = q.order_by(Transaction.timestamp.asc(), Transaction.id.asc()).all()
    qty = float(qty or 0.0)
    net = sum(float(t.quantity or 0.0) for t in rows)
    single = any(float(t.quantity or 0.0) > 0
                 and abs(float(t.quantity) - qty) <= qty * COUNT_GUARD_TOLERANCE
                 for t in rows)
    flagged = qty > 0 and bool(rows) and (single or net >= qty - 1e-9)
    return {
        "maybe_already_counted": flagged,
        "adjustments": [{
            "id": t.id,
            "date": to_tr(t.timestamp).strftime("%d.%m.%Y %H:%M") if t.timestamp else "",
            "qty": round(float(t.quantity or 0.0), 4),
            "by": t.performed_by or "",
            "note": (t.notes or "")[:300],
        } for t in rows[-30:]],
    }


def _convert_guards(db: Session, row: Inventory, source: Item, target: Optional[Item],
                    *, moving: bool) -> List[dict]:
    """Çevirmede çift sayım bakılacak kartlar — yalnız bayrak kalkanlar döner.

    Sıra: hedef kart (`target`; yeni kartta None — hareketi yok), numune başka
    karta gidiyorsa numunenin DURDUĞU kart (`on='source'`) ve numune daha önce
    "Karta taşı" ile taşınmışsa İLK girildiği kart (`moved_from_item_id`,
    `on='origin'`).  Kaynakta elle sayılmış numune başka karta eklenirse aynı
    fiziksel miktar İKİ kartta sayılır: kaynak +20 (lotsuz), hedef +20 —
    758'in iki kart arasındaki tekrarı (06.10 incelemesi, prob ile doğrulandı).
    """
    qty = float(row.quantity or 0.0)
    checks = [("target", target)] if target is not None else []
    if moving:
        checks.append(("source", source))
    if row.moved_from_item_id:
        origin = db.query(Item).filter(Item.id == row.moved_from_item_id).first()
        if origin is not None:
            checks.append(("origin", origin))
    out, seen = [], set()
    for on, it in checks:
        if it.id in seen:
            continue
        seen.add(it.id)
        g = _count_guard(db, it.id, row.created_at, qty)
        if g["maybe_already_counted"]:
            out.append({"on": on, "item_id": it.id, "item_name": it.name,
                        "adjustments": g["adjustments"]})
    return out


def _guard_detail(g: dict, target_name: str) -> str:
    """409 `maybe_already_counted` mesajı — bulgunun hangi kartta olduğuna göre
    doğru yolu söyler."""
    if g["on"] == "source":
        return (f"«{g['item_name']}» kartında (numunenin bağlı olduğu kart) numune geldikten "
                f"sonra elle yapılmış stok düzeltmesi var — numune orada zaten sayılmış "
                f"olabilir; «{target_name}» kartına da eklenirse iki kartta sayılır.  Doğru "
                f"yol: önce numuneyi kendi kartında \"Zaten sayıldı — yalnız lotu bağla\" ile "
                f"çevirin, sonra lotu \"Taşı\" ile hedef karta taşıyın.")
    if g["on"] == "origin":
        return (f"«{g['item_name']}» kartında (numunenin ilk girildiği kart) numune geldikten "
                f"sonra elle yapılmış stok düzeltmesi var — numune orada zaten sayılmış "
                f"olabilir.  Doğru yol: numuneyi «{g['item_name']}» kartına geri taşıyıp orada "
                f"\"yalnız lotu bağla\" deyin, sonra lotu \"Taşı\" ile hedef karta taşıyın.")
    return (f"«{g['item_name']}» kartında numune geldikten sonra elle yapılmış stok "
            f"düzeltmesi var — bu numune zaten sayılmış olabilir.")


def _sample_analysis_use(db: Session, inventory_id: int) -> tuple:
    """Numune Analizi'nde bu numune lotundan şu an düşülmüş miktar + belge
    numaraları (`consumed_qty > 0`, source='sample')."""
    from database import SampleAnalysis, SampleAnalysisIngredient
    rows = (db.query(SampleAnalysisIngredient.consumed_qty, SampleAnalysis.document_no)
            .join(SampleAnalysis, SampleAnalysis.id == SampleAnalysisIngredient.analysis_id)
            .filter(SampleAnalysisIngredient.inventory_id == inventory_id,
                    SampleAnalysisIngredient.source == "sample",
                    SampleAnalysisIngredient.consumed_qty > 0).all())
    used = round(sum(float(q or 0.0) for q, _ in rows), 6)
    docs = sorted({d for _, d in rows if d})
    return used, docs


def _name_tokens(name) -> set:
    """Benzer ad karşılaştırması için TR-katlanmış sözcükler — parantez içi
    ("(NUMUNE)") ve noktalama atılır."""
    s = re.sub(r"\([^)]*\)", " ", _tr_fold(name or ""))
    return set(re.sub(r"[^0-9a-z]+", " ", s).split())


def _names_similar(a: Item, b: Item) -> bool:
    """Basit benzer ad: birinin sözcükleri diğerinde tamamen geçiyor (name ya
    da name_tr).  Tek sözcükte en az 5 harf aranır — "YAĞ" her yağa uymasın.
    Ayrıca malzeme anahtarı (`material_groups.material_key`, TR/EN eşanlamlı)
    aynıysa benzer: SETİL STEARİL ALKOL ↔ CETYL STEARYL ALCOHOL."""
    for x in (_name_tokens(a.name), _name_tokens(a.name_tr)):
        for y in (_name_tokens(b.name), _name_tokens(b.name_tr)):
            small, big = (x, y) if len(x) <= len(y) else (y, x)
            if not small or not small <= big:
                continue
            if len(small) >= 2 or len(next(iter(small))) >= 5:
                return True
    ka = {material_groups.material_key(n) for n in (a.name, a.name_tr) if n} - {()}
    kb = {material_groups.material_key(n) for n in (b.name, b.name_tr) if n} - {()}
    return bool(ka & kb)


def _active_groups(db: Session, domain: str) -> dict:
    """{grup id: ad} — bu paneldeki AKTİF "aynı malzeme" grupları."""
    return dict(db.query(MaterialGroup.id, MaterialGroup.name)
                .filter(MaterialGroup.domain == domain,
                        MaterialGroup.is_active == True).all())          # noqa: E712


def _target_candidates(db: Session, domain: str, source: Item, supplier_id: Optional[int], *,
                       who: str = "Seçilen tedarikçi", sup_names: Optional[dict] = None,
                       groups: Optional[dict] = None) -> List[dict]:
    """`source` kartı yerine stoğun yazılabileceği kartlar — numune "Stoğa
    çevir" hedefi (convert-options) ve mal kabul ipucu (receive-options) ortak
    sıralaması.

    Sıra: aynı grup + `supplier_id` kartı > aynı grup > kart tedarikçisi =
    `supplier_id` ve benzer ad > benzer ad (aynı sırada TR-katlanmış ad, id).
    Yalnız aktif, aynı panel, aynı tür (hammadde/ambalaj; Bitmiş Ürün asla) ve
    aynı birim ailesi (`core.purchase_plan.unit_norm`) kartlar; `source`
    hariç.  `who`: gerekçe metnindeki tedarikçi ("Numunenin tedarikçisi")."""
    from core.purchase_plan import unit_norm
    if sup_names is None:
        sup_names = dict(db.query(Supplier.id, Supplier.name).all())
    if groups is None:
        groups = _active_groups(db, domain)
    src_gid = source.material_group_id if source.material_group_id in groups else None
    kind = stock_lots.lot_kind(source.category)
    unit_key = unit_norm(source.unit)
    ranked = []
    if kind == "finished":
        return []
    cards = (db.query(Item)
             .filter(Item.domain == domain, Item.is_active == True,       # noqa: E712
                     Item.id != source.id).all())
    for it in cards:
        if stock_lots.lot_kind(it.category) != kind or unit_norm(it.unit) != unit_key:
            continue
        same_group = src_gid is not None and it.material_group_id == src_gid
        same_sup = supplier_id is not None and it.supplier_id == supplier_id
        if same_group and same_sup:
            tier, reason = 0, f"Aynı malzeme grubu · {who.lower()}"
        elif same_group:
            tier, reason = 1, "Aynı malzeme grubu"
        elif _names_similar(source, it):
            tier, reason = ((2, f"{who} · benzer ad") if same_sup else (3, "Benzer ad"))
        else:
            continue
        ranked.append((tier, _tr_fold(it.name), it.id, {
            "item_id": it.id, "name": it.name, "name_tr": it.name_tr or "",
            "unit": it.unit or "",
            "supplier_id": it.supplier_id,
            "supplier_name": sup_names.get(it.supplier_id) if it.supplier_id else None,
            "stock": round(float(it.current_stock or 0.0), 4),
            "material_group_id": it.material_group_id if it.material_group_id in groups else None,
            "reason": reason,
        }))
    ranked.sort(key=lambda x: x[:3])
    return [c for *_, c in ranked]


@router.get("/inventory/receive-options")
def receive_options(
    item_id: int,
    supplier_id: int,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("inventory", "receive")),
    domain: str = Depends(active_domain),
):
    """Mal kabul ipucu — seçilen tedarikçi kartın tedarikçisinden farklıysa
    "bu tedarikçinin kartı «…» — oraya kabul et" / "yeni kart aç (aynı
    malzeme grubuna)" önerisi; tedarikçi "bitirilecek" ise uyarı.  Lab'ın
    düzeni her tedarikçiye ayrı kart (core/material_groups) — yanlış karta
    kabul jojoba vakasını tekrarlar.  ASLA engellemez; yalnız bilgi.

    Yanıt: `mismatch` (kartın tedarikçisi var VE seçilenle aynı firma değil —
    mükerrer firma kartları `supplier_key` ile aynı sayılır), `card_supplier`,
    `supplier` {id, name, status, status_label, status_reason} (durum firma
    anahtarı düzeyinde: mükerrer kartlardan biri bitirilecekse bitirilecek),
    `supplier_card` (bu tedarikçinin aynı malzemedeki ilk adayı ya da null),
    `candidates` (convert-options sırası: aynı grup + bu tedarikçi önce, ≤30),
    `new_card` {suggested_name "KART — FİRMA", join_group_of_item_id, unit,
    category, supplier_id, supplier_name, conflict}."""
    from core.purchase_pricing import SupplierIndex
    from core.suppliers import STATUS_LABELS, normalize_status, status_by_key
    item = db.query(Item).filter(Item.id == item_id, Item.domain == domain).first()
    if item is None:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})
    sup = db.query(Supplier).filter(Supplier.id == supplier_id, Supplier.domain == domain).first()
    if sup is None:
        return JSONResponse(status_code=404, content={"detail": "Tedarikçi bulunamadı."})
    all_sups = db.query(Supplier).filter(Supplier.domain == domain).all()
    ix = SupplierIndex([s.name for s in all_sups])
    by_key, _ = status_by_key(all_sups, ix)
    k = ix.key(sup.name)
    status = by_key.get(k) or normalize_status(sup.purchase_status) or "normal"
    card_sup = next((s for s in all_sups if s.id == item.supplier_id), None) if item.supplier_id else None
    mismatch = card_sup is not None and card_sup.id != sup.id and ix.key(card_sup.name) != k
    sup_names = {s.id: s.name for s in all_sups}
    groups = _active_groups(db, domain)
    cands = _target_candidates(db, domain, item, sup.id, sup_names=sup_names, groups=groups)
    same_firm = {s.id for s in all_sups if ix.key(s.name) == k}
    supplier_card = next((c for c in cands if c["supplier_id"] in same_firm), None)
    suggested = f"{item.name} — {sup.name}"
    conflict = _find_name_conflict(db, domain, suggested, None)
    reason = None
    if status == "phase_out":
        reason = sup.status_reason or next(
            (s.status_reason for s in all_sups if ix.key(s.name) == k and s.status_reason
             and normalize_status(s.purchase_status) == "phase_out"), None)
    return {
        "item": {"id": item.id, "name": item.name, "unit": item.unit or "",
                 "group": ({"id": item.material_group_id, "name": groups[item.material_group_id]}
                           if item.material_group_id in groups else None)},
        "mismatch": bool(mismatch),
        "card_supplier": ({"id": card_sup.id, "name": card_sup.name} if card_sup is not None else None),
        "supplier": {"id": sup.id, "name": sup.name, "status": status,
                     "status_label": STATUS_LABELS[status], "status_reason": reason},
        "supplier_status": status,
        "supplier_card": supplier_card,
        "candidates": cands[:30],
        "new_card": {
            "suggested_name": suggested[:150], "join_group_of_item_id": item.id,
            "unit": item.unit or "", "category": item.category or "",
            "supplier_id": sup.id, "supplier_name": sup.name,
            "conflict": _conflict_payload(conflict) if conflict else None,
        },
    }


@router.get("/inventory/samples/{inventory_id}/convert-options")
def sample_convert_options(
    inventory_id: int,
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("inventory", "receive")),
    domain: str = Depends(active_domain),
):
    """"Stoğa çevir" penceresinin verisi — numune, bağlı olduğu kart, aday
    hedef kartlar, çift sayım bulguları ve önerilen yeni kart adı.

    Aday sırası: aynı grup + numunenin tedarikçisi > aynı grup > kart
    tedarikçisi = numune tedarikçisi ve benzer ad > benzer ad.  Yalnız aktif,
    aynı panel, aynı tür (hammadde/ambalaj; Bitmiş Ürün asla) ve aynı birim
    ailesi (`core.purchase_plan.unit_norm`) kartlar.  `guard` bağlı olduğu
    kart içindir (eski alan); `guards` HER hedefte geçerli bulgulardır —
    numunenin durduğu kart ve ilk girildiği kart (`_convert_guards`).  Seçilen
    başka hedefin kendi bulgusunu POST 409 `maybe_already_counted` söyler.
    `sample.analysis_used` Numune Analizi'nde bu lottan düşülmüş miktardır
    ("yalnız lotu bağla" uyarısı için).
    """
    row = (db.query(Inventory).options(joinedload(Inventory.supplier))
           .filter(Inventory.id == inventory_id).first())
    if not row or not row.is_sample or row.domain != domain:
        return JSONResponse(status_code=404, content={"detail": "Numune kaydı bulunamadı."})
    source = db.query(Item).filter(Item.id == row.item_id).first()
    if not source:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    sup_name = row.supplier.name if row.supplier else None
    sup_names = dict(db.query(Supplier.id, Supplier.name).all())
    groups = _active_groups(db, domain)
    src_gid = source.material_group_id if source.material_group_id in groups else None
    candidates = _target_candidates(db, domain, source, row.supplier_id, who="Numunenin tedarikçisi",
                                    sup_names=sup_names, groups=groups)

    suggested = f"{source.name} — {sup_name}" if sup_name else source.name
    conflict = _find_name_conflict(db, domain, suggested, None)
    qty = round(float(row.quantity or 0.0), 4)
    used, used_docs = _sample_analysis_use(db, row.id)
    return {
        "sample": {
            "id": row.id, "item_id": row.item_id, "lot": row.lot_number, "qty": qty,
            "unit": source.unit or "", "supplier_id": row.supplier_id,
            "supplier_name": sup_name,
            "created_at": to_tr(row.created_at).strftime("%d.%m.%Y") if row.created_at else "",
            "analysis_used": round(used, 4), "analysis_docs": used_docs,
        },
        "source": {
            "id": source.id, "name": source.name, "unit": source.unit or "",
            "category": source.category or "",
            "supplier_id": source.supplier_id,
            "supplier_name": sup_names.get(source.supplier_id) if source.supplier_id else None,
            "stock": round(float(source.current_stock or 0.0), 4),
            "is_active": bool(source.is_active),
            "group": {"id": src_gid, "name": groups[src_gid]} if src_gid else None,
        },
        "candidates": candidates[:30],
        "guard": _count_guard(db, source.id, row.created_at, qty),
        "guards": _convert_guards(db, row, source, None, moving=True),
        "new_card": {
            "suggested_name": suggested[:150], "unit": source.unit or "",
            "category": source.category or "",
            "supplier_id": row.supplier_id, "supplier_name": sup_name,
            "conflict": _conflict_payload(conflict) if conflict else None,
        },
    }


@router.post("/inventory/samples/{inventory_id}/convert", status_code=200)
def convert_sample_to_stock(
    inventory_id: int,
    request: Request,
    data: Optional[SampleConvertBody] = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "receive")),
    domain: str = Depends(active_domain),
):
    """Numuneyi GERÇEK stoğa çevirir — lab değerlendirmesi olumlu geldiğinde.

    Numune artık current_stock'a hiç girmiyor (2026-08-24 olayı sonrası).
    Bir numune üretimde kullanılabilir hâle gelmesi gerekiyorsa (yeterli
    miktar geldi, tedarikçi onaylandı) bu uç onu deftere işler: tam satır
    (kısmi bölme kapsam dışı) `is_sample=False`'a döner ve TAM BURADA
    `Transaction(Input)` + `current_stock` artışı yazılır — böylece snapshot/
    trace/aylık rapor hep tutarlı kalır (numune hiçbir aşamada iki kere
    sayılmaz).

    Gövde isteğe bağlı (`SampleConvertBody`; gövdesiz çağrı eskisi gibi):
      • `target_item_id` — numune başka karta (aynı panel, aynı tür, aynı
        birim ailesi, aktif, Bitmiş Ürün değil) taşınıp orada çevrilir;
        `new_item` — numunenin tedarikçisiyle kaynak kartın kategori/birimini
        alan yeni kart açılır (ad çakışması 409, `force`) ve kaynakla aynı
        malzeme grubuna katılır.
      • Hedefte aynı lot no'lu normal satır: tedarikçi aynı ya da biri boşsa
        birleşir (analiz satırları silinmeden önce yönlenir); farklıysa 409
        `lot_collision` — istemci `lot_number` ile tekrar dener.
      • Çift sayım koruması (`_convert_guards`): hedef kart, numune başka
        karta gidiyorsa numunenin durduğu kart ve ilk girildiği kart.  "add"
        modunda bayrak kalkarsa 409 `maybe_already_counted` (`guards` hepsini
        listeler) — `acknowledge_counted` hepsini, `acknowledged_item_ids`
        yalnız görülen kartları onaylar.
      • `mode="link_only"`: numune zaten sayılmış — lot normal lota döner,
        `sample_converted_at` damgalanır; Transaction YOK, stok DEĞİŞMEZ,
        sipariş bayrağı kapanmaz.  Yalnız numunenin kendi kartında.  Numune
        Analizi bu lottan düşmüşse yanıt uyarır (bkz. aşağıdaki not).
    Not öneki "Numune stoğa çevrildi" KORUNUR — satın alma motoru çevrilmiş
    lotu bu işaretle tanır (core/purchase_plan.SAMPLE_CONVERTED_MARK).
    """
    from core.purchase_plan import SAMPLE_CONVERTED_MARK
    from core.stock_gaps import close_open_flags

    data = data or SampleConvertBody()
    row = (db.query(Inventory).filter(Inventory.id == inventory_id)
          .with_for_update().first())
    if not row or not row.is_sample:
        return JSONResponse(status_code=404, content={"detail": "Numune kaydı bulunamadı."})
    if row.domain != domain:
        return JSONResponse(status_code=404, content={"detail": "Numune kaydı bulunamadı."})

    source = db.query(Item).filter(Item.id == row.item_id).with_for_update().first()
    if not source:
        return JSONResponse(status_code=404, content={"detail": "Ürün bulunamadı."})

    qty = round(float(row.quantity or 0), 6)
    if qty <= 1e-9:
        return JSONResponse(status_code=400, content={
            "detail": "Numune lotunda miktar kalmamış — stoğa çevrilecek bir şey yok."})

    # Şahit numune dolabı bu lotu izlemiyor olmalı — production'daki -S
    # lotları ayrı bir akış (routers/production.py); burada sadece emin ol.
    if db.query(RetentionSample.id).filter(RetentionSample.inventory_id == row.id).first():
        return JSONResponse(status_code=400, content={
            "detail": "Bu kayıt şahit numune dolabına bağlı, buradan çevrilemez."})

    actor = current_user.get("full_name") or current_user.get("username") or "—"

    # ── Hedef kartı çöz (mutasyon yok — önce tüm kontroller) ────────────────
    if data.target_item_id and data.new_item:
        return JSONResponse(status_code=400, content={
            "detail": "Ya mevcut bir kart ya da yeni kart seçin — ikisi birden olmaz."})
    target = source
    new_name = new_name_tr = None
    if data.target_item_id and data.target_item_id != source.id:
        target = db.query(Item).filter(Item.id == data.target_item_id).with_for_update().first()
        if not target or (target.domain or "cosmetics") != domain:
            return JSONResponse(status_code=404, content={"detail": "Hedef kart bulunamadı."})
        problem = stock_lots.target_problem(source, target)
        if problem:
            return JSONResponse(status_code=400, content={"detail": problem})
    elif data.new_item:
        user = db.query(User).filter(User.id == int(current_user.get("sub", 0))).first()
        from core.permissions import _has_permission
        if not (user and _has_permission(user, "items", "create")):
            return JSONResponse(status_code=403, content={
                "detail": "Yeni kart açma yetkiniz yok (items.create)."})
        if stock_lots.lot_kind(source.category) == "finished":
            return JSONResponse(status_code=400, content={
                "detail": "Bitmiş ürün numunesinden yeni kart açılamaz."})
        new_name = data.new_item.name.strip()
        new_name_tr = (data.new_item.name_tr or "").strip() or None
        if not new_name:
            return JSONResponse(status_code=400, content={"detail": "Yeni kartın adı boş olamaz."})
        if not data.new_item.force:
            conflict = _find_name_conflict(db, domain, new_name, new_name_tr)
            if conflict:
                return JSONResponse(status_code=409, content={
                    "code": "name_conflict",
                    "detail": f"Aynı isimde ürün zaten kayıtlı: {conflict.name}",
                    "existing": _conflict_payload(conflict),
                })
    if target is source and not data.new_item and not source.is_active:
        return JSONResponse(status_code=400, content={
            "detail": f"«{source.name}» kartı pasif — numuneyi aktif bir karta çevirmek için "
                      f"hedef kart seçin."})
    moving = bool(data.new_item) or target.id != source.id
    if data.mode == "link_only" and moving:
        return JSONResponse(status_code=400, content={
            "detail": "\"Yalnız lotu bağla\" yalnız numunenin kendi kartında kullanılabilir — "
                      "önce lotu karta taşıyın."})

    new_lot = (data.lot_number or "").strip() or row.lot_number
    twin = None
    if not data.new_item:
        twin = stock_lots.find_twin(db, target.id, new_lot, is_sample=False, exclude_id=row.id)
        if twin is not None and not stock_lots.suppliers_compatible(twin.supplier_id, row.supplier_id):
            return JSONResponse(status_code=409,
                                content=stock_lots.collision_error(db, twin, new_lot).payload())
    guards = _convert_guards(db, row, source, None if data.new_item else target, moving=moving)
    acked = set(data.acknowledged_item_ids or ())
    pending = [g for g in guards
               if not (data.acknowledge_counted or g["item_id"] in acked)]
    if data.mode == "add" and pending:
        first = pending[0]
        return JSONResponse(status_code=409, content={
            "code": "maybe_already_counted",
            "detail": _guard_detail(first, data.new_item.name if data.new_item else target.name),
            # Eski alanlar ilk bulgudan; `guards` hepsini taşır.
            "item_id": first["item_id"], "item_name": first["item_name"], "on": first["on"],
            "sample_qty": qty, "unit": source.unit or "",
            "adjustments": first["adjustments"],
            "guards": guards,
        })
    used, used_docs = _sample_analysis_use(db, row.id) if data.mode == "link_only" else (0.0, [])

    # ── Uygula ──────────────────────────────────────────────────────────────
    sup_name = row.supplier.name if row.supplier else "—"
    supplier_id = row.supplier_id
    old_lot = row.lot_number
    created_item_id = group_id = None
    warning = None
    now = datetime.utcnow()
    try:
        if data.new_item:
            target = Item(name=new_name, name_tr=new_name_tr, category=source.category,
                          unit=source.unit, supplier_id=row.supplier_id, domain=domain,
                          min_stock_level=0.0, cost_price=0.0, current_stock=0.0)
            db.add(target)
            db.flush()
            created_item_id = target.id
            grp = material_groups.join(db, source, target, actor=actor)
            group_id = grp.id if grp else None
        moved = target.id != source.id
        if moved:
            stock_lots.move_sample_row(db, row, target)
        row.lot_number = new_lot

        if twin is not None:
            stock_lots.absorb_row(db, row, twin, item_id=target.id)
            if twin.sample_converted_at is None:
                twin.sample_converted_at = now
            lot_row = twin
            lot_row_msg = f"lot {new_lot} — mevcut normal lotla birleşti"
        else:
            row.is_sample = False
            if row.location == "Numune":
                row.location = None
            row.sample_converted_at = now
            row.updated_at = now
            lot_row = row
            lot_row_msg = f"lot {new_lot}"

        if data.mode == "add":
            note = f"{SAMPLE_CONVERTED_MARK} — Lot: {new_lot} | Tedarikçi: {sup_name}"
            if moved:
                note += f" | Kart: {source.name} → {target.name}"
            db.add(Transaction(
                item_id=target.id, lot_number=new_lot, transaction_type="Input",
                quantity=qty, notes=note[:500], performed_by=actor,
            ))
            target.current_stock = round((target.current_stock or 0.0) + qty, 6)
            db.flush()
            # Numune gerçek Input'a döndü — Eksik Hammaddeler'de açık sipariş
            # bayrağı varsa (core/stock_gaps.py) kapanır.
            close_open_flags(db, target.id, "received", actor)
        else:
            db.flush()
            lot_total = float(
                db.query(func.coalesce(func.sum(Inventory.quantity), 0.0))
                .filter(Inventory.item_id == target.id, Inventory.is_sample == False,  # noqa: E712
                        Inventory.quantity > 0).scalar() or 0.0)
            stock = float(target.current_stock or 0.0)
            warns = []
            if lot_total > stock + 1e-6:
                warns.append(f"Kartın normal lot toplamı ({lot_total:g} {target.unit or ''}) stoktan "
                             f"({stock:g} {target.unit or ''}) fazla — numune elle sayılmamış "
                             f"olabilir, stoğu kontrol edin.")
            # Analizde numuneden düşülen kısım stoğa hiç yazılmadı (numune stok
            # değildi); elle sayım numunenin TAMAMINI saydıysa bu kısım stokta
            # fazladan durur.  Analiz silinir/azaltılırsa core.sample_trial_stock
            # ._release lot artık normal olduğu için +Adjustment yazar → stok bir
            # kez daha artar.  Hangisi olduğunu sistem bilemez: sayım ister.
            if used > 1e-9:
                warns.append(f"Bu numuneden Numune Analizi'nde {used:g} {target.unit or ''} "
                             f"kullanılmış ({', '.join(used_docs) or 'analiz'}). Elle sayım "
                             f"numunenin tamamını kapsadıysa bu kısım stokta fazladan duruyor "
                             f"olabilir — sayımla kontrol edip Stok düzeltme yapın. Analiz "
                             f"silinir ya da miktarı azaltılırsa iade stoğa da eklenir.")
            warning = " ".join(warns) or None
        lot_row_id = lot_row.id
        db.commit()
    except Exception:
        db.rollback()
        return JSONResponse(status_code=500, content={"detail": "Numune stoğa çevrilemedi."})

    log_admin_event(
        db, request, actor=current_user,
        action="inventory.sample_convert" if data.mode == "add" else "inventory.sample_link",
        target_type="inventory", target_id=inventory_id, target_name=target.name,
        details={"lot": new_lot, "old_lot": old_lot if old_lot != new_lot else None,
                 "quantity": qty, "unit": target.unit, "supplier_id": supplier_id,
                 "source_item_id": source.id, "target_item_id": target.id,
                 "mode": data.mode, "created_item_id": created_item_id,
                 "group_id": group_id, "merged_into": twin.id if twin is not None else None,
                 "guard_item_ids": [g["item_id"] for g in guards],
                 "guard_adjustment_ids": [a["id"] for g in guards for a in g["adjustments"]],
                 "analysis_used": used or None})

    unit = target.unit or ""
    if data.mode == "add":
        msg = f"Numune stoğa çevrildi — {qty:g} {unit} {target.name} ({lot_row_msg})."
    else:
        msg = (f"Numune lotu karta bağlandı — {qty:g} {unit} {target.name} ({lot_row_msg}); "
               f"stok değişmedi.")
    return {"message": msg, "mode": data.mode, "item_id": target.id, "item_name": target.name,
            "inventory_id": lot_row_id, "lot_number": new_lot, "quantity": qty,
            "moved": target.id != source.id, "source_item_id": source.id,
            "created_item_id": created_item_id, "group_id": group_id, "warning": warning}


# ─── Lotu başka karta taşı (06.10.2026) ─────────────────────────────────────
# Naturalya lotlarını lab KENDİSİ doğru karta taşısın diye (kullanıcı kararı).
# Defter kuralı core/stock_lots.move_lot docstring'inde: normal lot Adjustment
# çiftiyle, numune lotu deftersiz.

class LotMoveBody(BaseModel):
    target_item_id: int
    # boş = lotun tamamı (numunede zorunlu tam).  NaN/Infinity'yi
    # core.stock_lots.move_lot 400 ile keser (bkz. `_finite`).
    quantity: Optional[float] = None
    supplier_id: Optional[int] = None                # lot tedarikçisini düzelt (boş = aynı)
    lot_number: Optional[str] = Field(None, max_length=100)   # lot_collision çözümü
    reason: str = Field(..., min_length=3, max_length=200)


@router.post("/inventory/lots/{inventory_id}/move")
def move_inventory_lot(
    inventory_id: int,
    data: LotMoveBody,
    request: Request,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_permission("inventory", "receive")),
    domain: str = Depends(active_domain),
):
    """Lotu başka karta taşı — numune lotu `inventory.receive` ile, normal
    (stoktaki) lot ayrıca `inventory.adjust` ister (iki kartın stoğunu
    değiştirir; LabTech varsayılanında yok).  Kurallar ve defter kaydı
    `core.stock_lots.move_lot`'ta; engeller 400/409 (`lot_collision`)."""
    reason = (data.reason or "").strip()
    if len(reason) < 3:
        return JSONResponse(status_code=422, content={"detail": "Sebep en az 3 karakter olmalıdır."})

    lot = db.query(Inventory).filter(Inventory.id == inventory_id).with_for_update().first()
    if not lot or lot.domain != domain:
        return JSONResponse(status_code=404, content={"detail": "Lot bulunamadı."})
    if not lot.is_sample:
        from core.permissions import _has_permission
        user = db.query(User).filter(User.id == int(current_user.get("sub", 0))).first()
        if not (user and _has_permission(user, "inventory", "adjust")):
            return JSONResponse(status_code=403, content={
                "detail": "Stoktaki lotu taşımak stok düzeltme yetkisi ister (inventory.adjust)."})

    source = db.query(Item).filter(Item.id == lot.item_id).with_for_update().first()
    if not source or (source.domain or "cosmetics") != domain:
        return JSONResponse(status_code=404, content={"detail": "Lotun kartı bulunamadı."})
    target = db.query(Item).filter(Item.id == data.target_item_id).with_for_update().first()
    if not target or (target.domain or "cosmetics") != domain:
        return JSONResponse(status_code=404, content={"detail": "Hedef kart bulunamadı."})
    if data.supplier_id is not None:
        sup = db.query(Supplier).filter(Supplier.id == data.supplier_id).first()
        if not sup or (sup.domain or "cosmetics") != domain:
            return JSONResponse(status_code=400, content={"detail": "Tedarikçi bulunamadı."})

    actor = current_user.get("full_name") or current_user.get("username") or "—"
    try:
        res = stock_lots.move_lot(db, lot, source, target, data.quantity, actor=actor,
                                  reason=reason, supplier_id=data.supplier_id,
                                  lot_number=data.lot_number)
        db.commit()
    except LotMoveError as exc:
        db.rollback()
        return JSONResponse(status_code=exc.status, content=exc.payload())
    except Exception:
        db.rollback()
        return JSONResponse(status_code=500, content={"detail": "Lot taşınamadı — hiçbir şey değişmedi."})

    log_admin_event(db, request, actor=current_user, action="inventory.lot_move",
                    target_type="inventory", target_id=inventory_id, target_name=target.name,
                    details={**res, "reason": reason, "source_name": source.name})
    tail = " (hedefteki aynı lotla birleşti)" if res["merged"] else ""
    return {
        "message": (f"Lot {res['lot_number']} taşındı: {res['quantity']:g} {res['unit']} "
                    f"«{source.name}» → «{target.name}»{tail}."),
        **res,
    }


class _AvailableLotsRequest(BaseModel):
    item_ids: List[int] = Field(..., max_length=300)


@router.post("/inventory/available-lots")
def available_lots(
    data: _AvailableLotsRequest,
    db: Session = Depends(get_db),
    _: dict = Depends(require_internal_user(("production", "view"))),
    domain: str = Depends(active_domain),
):
    """
    Üretim ekranı için: verilen hammaddelerin TÜKETİLEBİLİR lotları
    (APPROVED + miktar>0 + numune DEĞİL), tedarikçi bilgisiyle, FIFO sırasında
    (en eski önce).  `{item_id: [ {inventory_id, lot_number, supplier_name,
    quantity, is_sample, location, expiry_date}, … ]}` döner.

    Numune lotları hariç (2026-08-24 olayı) — numune stok değildir, üretimde
    kullanılamaz.  `is_sample` alanı geriye dönük uyum için hâlâ dönüyor,
    artık her zaman `false`.
    """
    if not data.item_ids:
        return {}
    rows = (
        db.query(Inventory)
        .options(joinedload(Inventory.supplier))
        .filter(
            Inventory.item_id.in_(data.item_ids),
            Inventory.status == "APPROVED",
            Inventory.quantity > 0,
            Inventory.is_sample == False,   # noqa: E712
            Inventory.domain == domain,
        )
        .order_by(Inventory.created_at.asc(), Inventory.id.asc())
        .all()
    )
    out: dict = {}
    for r in rows:
        out.setdefault(r.item_id, []).append({
            "inventory_id":  r.id,
            "lot_number":    r.lot_number,
            "supplier_name": (r.supplier.name if r.supplier else "—"),
            "quantity":      round(float(r.quantity or 0), 4),
            "is_sample":     bool(r.is_sample),
            "location":      r.location or "",
            "expiry_date":   r.expiry_date or "",
        })
    return out


@router.get("/transactions")
def list_transactions(db: Session = Depends(get_db), domain: str = Depends(active_domain),
                      _: dict = Depends(require_any_permission(("inventory", "view"), ("reports", "view")))):
    # 500 satır × N+1 ürün lookup'ı yerine joinedload ile tek query.
    # Transaction'da domain kolonu yok → Item join'iyle aktif panele süzülür.
    rows = (
        db.query(Transaction)
        .options(joinedload(Transaction.item))
        .join(Item, Item.id == Transaction.item_id)
        .filter(Item.domain == domain)
        .order_by(Transaction.id.desc())
        .limit(500)
        .all()
    )
    return [
        {
            "id": r.id,
            "item_name": r.item.name if r.item else "—",
            "lot_number": r.lot_number or "—",
            "transaction_type": r.transaction_type,
            "quantity": r.quantity,
            "notes": r.notes or "",
            "timestamp": to_tr(r.timestamp).strftime("%d.%m.%Y %H:%M") if r.timestamp else "",
        }
        for r in rows
    ]


# ─── Traceability / Audit Trail (Phase 2 / Task 9) ──────────────────────────

def _import_dt():
    """Tiny helper — import datetime lazily without polluting module top level."""
    import datetime as _d
    return _d.datetime


@router.get("/traceability/lot/{lot_number}")
def trace_lot(lot_number: str, item_id: Optional[int] = None, db: Session = Depends(get_db),
              _: dict = Depends(require_internal_user(("inventory", "view")))):
    """
    Full genealogy tree for a lot. Resolves:
      • Lot identity (Inventory record + supplier)
      • Production record (if internally produced)
      • Ingredients consumed during that production (with their own supplier/expiry/received_by)
      • All transactions for this lot — chronological audit trail.

    Lot no ÜRÜN BAZLIDIR (SR005 prod'da 7 üründe var): lot satırı ile üretim
    kaydı AYNI ürüne bağlanır — önce üretim (iptal edilmemiş, en yeni), lot
    satırı onun hedef kartından.  `item_id` verilirse o ürün seçilir; aynı lot
    no'lu öbür ürünler `other_items`'ta döner (arayüz ürün seçtirir).
    """
    # Imported here to avoid cross-router top-level dependency on production model
    from database import ProductionHistory, ProductionConsumption
    from sqlalchemy import case
    from core.production_cancel import cancelled_view

    # İptal edilmiş üretim / iptalle kapanmış lot satırı yalnız başka aday
    # yoksa gösterilir (lot no serbest bırakılıp yeniden kullanılmış olabilir).
    prod_q = db.query(ProductionHistory).filter(ProductionHistory.lot_number == lot_number)
    if item_id is not None:
        prod_q = prod_q.filter(ProductionHistory.target_item_id == item_id)
    prod = (prod_q.order_by(ProductionHistory.cancelled_at.isnot(None),
                            ProductionHistory.id.desc()).first())
    inv_q = db.query(Inventory).filter(Inventory.lot_number == lot_number)
    if item_id is not None:
        inv_q = inv_q.filter(Inventory.item_id == item_id)
    elif prod is not None and prod.target_item_id:
        inv_q = inv_q.filter(Inventory.item_id == prod.target_item_id)
    inv = (inv_q.order_by(case((Inventory.status == "CANCELLED", 1), else_=0), Inventory.id.asc())
           .first())

    if not inv and not prod:
        return JSONResponse(status_code=404, content={"detail": f"Lot bulunamadı: {lot_number}"})

    out = {"lot_number": lot_number}
    chosen = inv.item_id if inv is not None else prod.target_item_id
    other_ids = ({i for (i,) in db.query(Inventory.item_id)
                  .filter(Inventory.lot_number == lot_number).distinct().all()}
                 | {i for (i,) in db.query(ProductionHistory.target_item_id)
                    .filter(ProductionHistory.lot_number == lot_number,
                            ProductionHistory.target_item_id.isnot(None)).distinct().all()})
    other_ids.discard(chosen)
    out["item_id"] = chosen
    out["other_items"] = [{"item_id": it.id, "item_name": it.name}
                          for it in (db.query(Item).filter(Item.id.in_(other_ids))
                                     .order_by(Item.name.asc()).all() if other_ids else [])]

    # ── Lot bilgisi (Inventory) ──────────────────────────────────────────────
    if inv:
        item     = db.query(Item).filter(Item.id == inv.item_id).first()
        supplier = db.query(Supplier).filter(Supplier.id == inv.supplier_id).first() if inv.supplier_id else None
        out["lot_info"] = {
            "id":             inv.id,
            "item_id":        inv.item_id,
            "item_name":      item.name if item else "—",
            "item_category":  item.category if item else "—",
            "item_unit":      item.unit if item else "",
            "supplier_name":  supplier.name if supplier else "—",
            "supplier_phone": (supplier.phone if supplier else "") or "",
            "expiry_date":    inv.expiry_date or "—",
            "quantity":       inv.quantity,
            "location":       inv.location or "—",
            "status":         inv.status,
            "received_by":    inv.received_by or "—",
            "qc_approved_by": inv.qc_approved_by or "—",
            "qc_notes":       inv.qc_notes or "",
            "created_at":     to_tr(inv.created_at).strftime("%d.%m.%Y %H:%M") if inv.created_at else "",
            "updated_at":     to_tr(inv.updated_at).strftime("%d.%m.%Y %H:%M") if inv.updated_at else "",
        }
        # ── QC formu (varsa) — soru metinleriyle etiketlenmiş okunur görünüm ──
        from core.qc_report import parse_qc_form
        out["qc_form"] = parse_qc_form(inv, item)
    else:
        out["lot_info"] = None
        out["qc_form"]  = None

    # ── Üretim kaydı + tüketilen hammaddeler ────────────────────────────────
    if prod:
        # Tüketim dökümü (production_consumptions) varsa kesin kaynak odur —
        # aynı lot no'lu başka ürünün üretimi karışmaz.  Yoksa (eski kayıt)
        # üretimin defter imzasıyla yeniden kurulur (core/production_cancel —
        # reçete adı + lot + kişi + zaman; salt okur); o da boşsa Output
        # notundaki "Üretim Lot: {lot}" damgasına düşülür.
        snap = (db.query(ProductionConsumption)
                .filter(ProductionConsumption.production_id == prod.id,
                        ProductionConsumption.kind != "output",
                        ProductionConsumption.transaction_id.isnot(None))
                .order_by(ProductionConsumption.id.asc()).all())
        if not snap and prod.produced_at:
            from core.production_cancel import reconstruct_from_ledger
            snap = [pc for pc in reconstruct_from_ledger(db, prod)
                    if pc.kind != "output" and pc.transaction_id]
        if snap:
            ing_outputs = (db.query(Transaction)
                           .filter(Transaction.id.in_([pc.transaction_id for pc in snap]))
                           .order_by(Transaction.id.asc()).all())
        else:
            marker = f"Üretim Lot: {lot_number}"
            ing_outputs = (
                db.query(Transaction)
                .filter(
                    Transaction.transaction_type == "Output",
                    Transaction.notes.like(f"%{marker}%"),
                )
                .order_by(Transaction.id.asc())
                .all()
            )
        snap_inv = {pc.transaction_id: pc.inventory_id for pc in snap}
        # P2 — bölünmüş üretimde Output reçetedeki karttan değil gruptaki
        # başka karttan (tedarikçiden) düşmüş olabilir; döküm reçete kartını
        # ayrıca taşır (`recipe_item_id`).  Etikette fark dil kardeşidir.
        snap_by_tx = {pc.transaction_id: pc for pc in snap}
        rec_ids = {pc.recipe_item_id for pc in snap if pc.recipe_item_id}
        rec_names = ({i: n for i, n in db.query(Item.id, Item.name).filter(Item.id.in_(rec_ids)).all()}
                     if rec_ids else {})

        ingredients_consumed = []
        for tx in ing_outputs:
            tx_item = db.query(Item).filter(Item.id == tx.item_id).first()
            # Faz 2 — tüketim Output'u artık KAYNAK lotu Transaction.lot_number'a
            # yazıyor.  Varsa o kesin lottan tedarikçi/SKT/numune bilgisi gelir;
            # yoksa (eski kayıtlar) eski "en yakın APPROVED lot" tahminine düşülür.
            tx_inv = None
            if snap_inv.get(tx.id):
                tx_inv = db.query(Inventory).filter(Inventory.id == snap_inv[tx.id]).first()
            if tx_inv is None and tx.lot_number:
                tx_inv = (
                    db.query(Inventory)
                    .filter(Inventory.item_id == tx.item_id,
                            Inventory.lot_number == tx.lot_number)
                    .order_by(Inventory.id.desc())
                    .first()
                )
            if tx_inv is None:
                tx_inv = (
                    db.query(Inventory)
                    .filter(
                        Inventory.item_id == tx.item_id,
                        Inventory.status == "APPROVED",
                        Inventory.created_at <= (prod.produced_at or _import_dt().utcnow()),
                    )
                    .order_by(Inventory.created_at.desc())
                    .first()
                )
            tx_supplier = db.query(Supplier).filter(Supplier.id == tx_inv.supplier_id).first() if tx_inv and tx_inv.supplier_id else None
            pc = snap_by_tx.get(tx.id)
            rec_id = pc.recipe_item_id if pc is not None else None
            substituted = bool(rec_id and rec_id != tx.item_id)
            ingredients_consumed.append({
                "item_id":       tx.item_id,
                "item_name":     tx_item.name if tx_item else "—",
                "item_category": tx_item.category if tx_item else "—",
                "quantity":      tx.quantity,
                "unit":          tx_item.unit if tx_item else "",
                "source_lot":    (tx.lot_number or (tx_inv.lot_number if tx_inv else None)) or "—",
                # Reçetedeki kart (döküm varsa) — fiilen düşülenden farklıysa
                # `substituted` (kaynak kart seçimi ya da etiket dil kardeşi).
                "recipe_item_id":   rec_id,
                "recipe_item_name": rec_names.get(rec_id) if substituted else None,
                "substituted":      substituted,
                "kind":             pc.kind if pc is not None else None,
                "supplier_name": (tx_supplier.name if tx_supplier
                                  else ((pc.supplier_name if pc is not None else None) or "—")),
                "is_sample":     bool(tx_inv.is_sample) if tx_inv else False,
                "expiry_date":   (tx_inv.expiry_date if tx_inv else None) or "—",
                "received_by":   (tx_inv.received_by if tx_inv else None) or "—",
                "performed_by":  tx.performed_by or "—",
            })

        out["production"] = {
            "id":                prod.id,
            "recipe_id":         prod.recipe_id,
            "recipe_name":       prod.recipe_name or "—",
            "target_item_id":    prod.target_item_id,
            "target_item_name":  prod.target_item_name or "—",
            "produced_quantity": prod.produced_quantity,
            "produced_at":       to_tr(prod.produced_at).strftime("%d.%m.%Y %H:%M") if prod.produced_at else "—",
            "produced_by":       prod.produced_by or "—",
            "ingredients_consumed": ingredients_consumed,
            "cancelled":         cancelled_view(prod),
        }
    else:
        out["production"] = None

    # ── İşlem geçmişi ───────────────────────────────────────────────────────
    txs = (
        db.query(Transaction)
        .filter(Transaction.lot_number == lot_number)
        .order_by(Transaction.id.asc())
        .all()
    )
    out["transactions"] = [
        {
            "id":               t.id,
            "transaction_type": t.transaction_type,
            "quantity":         t.quantity,
            "timestamp":        to_tr(t.timestamp).strftime("%d.%m.%Y %H:%M") if t.timestamp else "—",
            "performed_by":     t.performed_by or "—",
            "notes":            (t.notes or "")[:200],
        }
        for t in txs
    ]

    return out


@router.get("/traceability/expiring")
def list_expiring(db: Session = Depends(get_db), _: dict = Depends(require_internal_user(("inventory", "view"))),
                  domain: str = Depends(active_domain)):
    """All APPROVED inventory lots expiring within the next 60 days, sorted most-urgent first."""
    from datetime import datetime as _dt, timedelta

    today  = _dt.utcnow().date()
    cutoff = today + timedelta(days=60)

    rows = (
        db.query(Inventory)
        .filter(
            Inventory.expiry_date.isnot(None),
            Inventory.expiry_date != "",
            Inventory.status == "APPROVED",
            Inventory.domain == domain,
        )
        .all()
    )

    result = []
    for r in rows:
        try:
            exp_date = _dt.strptime(r.expiry_date, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue
        days_left = (exp_date - today).days
        if days_left < 0 or days_left > 60:
            continue
        item = db.query(Item).filter(Item.id == r.item_id).first()
        result.append({
            "id":            r.id,
            "lot_number":    r.lot_number,
            "item_name":     item.name if item else "—",
            "item_category": item.category if item else "—",
            "item_unit":     item.unit if item else "",
            "expiry_date":   r.expiry_date,
            "days_left":     days_left,
            "quantity":      r.quantity,
            "location":      r.location or "—",
            "received_by":   r.received_by or "—",
        })

    result.sort(key=lambda x: x["days_left"])
    return result


# ─── User-specific audit feed (Phase 8 / Bug 5) ─────────────────────────────
# Lets management drill into what a specific user has done over a time window.
# Gated on admin.view_audit so Manager (Işık Hanım) and SuperAdmin can see it
# while lab roles cannot.

@router.get("/traceability/audit-users")
def list_audit_users(
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("admin", "view_audit")),
):
    """
    Lightweight user list — drives the dropdown on the traceability page.
    Each entry includes `tx_count` so the UI can show "(N işlem)" hints,
    avoiding the "Songül seçildi → boş" surprise.

    Match is case-insensitive and tries both full_name and username because
    the legacy code wrote actor names in mixed casings (e.g. "Doğukan
    YALÇINKAYA" vs "Doğukan Yalçınkaya").
    """
    from sqlalchemy import func
    from database import User

    users = (
        db.query(User)
        .filter(User.is_active == True)
        .order_by(User.full_name.asc())
        .all()
    )

    # Tek seferde tüm performed_by sayımlarını al — N+1 yok
    counts_raw = (
        db.query(func.lower(Transaction.performed_by), func.count(Transaction.id))
        .filter(Transaction.performed_by.isnot(None))
        .group_by(func.lower(Transaction.performed_by))
        .all()
    )
    counts = {k: v for k, v in counts_raw}   # {lowercase_name: count}

    out = []
    for u in users:
        candidates = [c.lower() for c in (u.full_name, u.username) if c]
        tx_count = sum(counts.get(c, 0) for c in candidates)
        out.append({
            "id":        u.id,
            "username":  u.username,
            "full_name": u.full_name,
            "role":      u.role,
            "tx_count":  tx_count,
        })
    return out


@router.get("/traceability/user-activity")
def user_activity(
    user_id: int,
    days:       Optional[int] = None,             # legacy preset support (last N days)
    date_from:  Optional[str] = None,             # YYYY-MM-DD inclusive (custom range)
    date_to:    Optional[str] = None,             # YYYY-MM-DD inclusive
    db: Session = Depends(get_db),
    _: dict = Depends(require_permission("admin", "view_audit")),
):
    """
    Audit trail for a specific user. Two modes:

      • Preset window:  days=30  → "last 30 days from now"
      • Custom range:   date_from=2025-01-01 & date_to=2025-12-31

    If both are sent, custom range wins. If neither, defaults to 30 days.
    Transaction.performed_by stores `actor` strings (full_name preferred,
    username fallback) so we match against both candidates of the target user.
    """
    from datetime import datetime as _dt, timedelta
    from database import User

    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        return JSONResponse(status_code=404, content={"detail": "Kullanıcı bulunamadı."})

    # ── Resolve the window ──────────────────────────────────────────────
    range_label = None
    if date_from or date_to:
        try:
            since = _dt.strptime(date_from, "%Y-%m-%d") if date_from else _dt(1970, 1, 1)
            until = _dt.strptime(date_to,   "%Y-%m-%d") + timedelta(days=1) if date_to else _dt.utcnow()
        except ValueError:
            return JSONResponse(status_code=422, content={
                "detail": "Geçersiz tarih formatı. YYYY-MM-DD bekleniyor."
            })
        range_label = f"{date_from or '∞'} → {date_to or 'bugün'}"
    else:
        n = max(1, min(int(days or 30), 1095))    # clamp 1 day .. 3 years
        since = _dt.utcnow() - timedelta(days=n)
        until = _dt.utcnow() + timedelta(days=1)
        range_label = f"son {n} gün"

    # Case-insensitive eşleşme — eski veriler "Doğukan YALÇINKAYA" gibi
    # büyük harf, yeni veriler "Doğukan Yalçınkaya" olabilir; ikisini de yakala.
    from sqlalchemy import func
    candidates_lower = [c.lower() for c in (user.full_name, user.username) if c]

    rows = (
        db.query(Transaction)
        .filter(
            func.lower(Transaction.performed_by).in_(candidates_lower),
            Transaction.timestamp >= since,
            Transaction.timestamp <  until,
        )
        .order_by(Transaction.id.desc())
        .limit(500)
        .all()
    )
    return {
        "user": {
            "id":        user.id,
            "username":  user.username,
            "full_name": user.full_name,
            "role":      user.role,
        },
        "range":       range_label,
        "since":       since.isoformat() + "Z",
        "until":       until.isoformat() + "Z",
        "count":       len(rows),
        "transactions": [
            {
                "id":               t.id,
                "item_name":        t.item.name if t.item else "—",
                "lot_number":       t.lot_number or "—",
                "transaction_type": t.transaction_type,
                "quantity":         t.quantity,
                "notes":            (t.notes or "")[:200],
                "timestamp":        to_tr(t.timestamp).strftime("%d.%m.%Y %H:%M") if t.timestamp else "—",
            }
            for t in rows
        ],
    }


# ─── Excel Import Endpoints (classic templated upload) ──────────────────────

# ─── Excel upload guards ────────────────────────────────────────────────────
# Lab dosyaları büyük olabiliyor (Sayfa10 + sayfa-sayfa hammadde sayımları),
# 100 MB tavanı pratik; üstüne çıkarsa OOM riskine girer.
MAX_EXCEL_BYTES = 100 * 1024 * 1024     # 100 MB

# XLSX = ZIP container; ZIP header magic bytes "PK\x03\x04" (veya bazen
# "PK\x05\x06" boş arşiv).  Saldırgan .xlsx uzantılı text/JS/HTML gönderirse
# openpyxl açmaya çalışmadan reddederiz.
_XLSX_MAGIC = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


async def _read_xlsx_safely(file) -> bytes:
    """
    Upload'ı validate edip bytes döner.  Hata olursa HTTPException raise eder.
    Sıra: uzantı → boyut (Content-Length spoof'a karşı stream-based) → magic bytes.
    """
    from fastapi import HTTPException

    # 1) Uzantı (hızlı reddetme — yine de zorunlu değil ama UX'i iyi)
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Yalnızca .xlsx dosyaları desteklenir.")

    # 2) Stream-based read — saldırgan Content-Length'i yalan söylese bile
    #     biz okurken anlık byte sayısını sayıyoruz, MAX_EXCEL_BYTES'i geçince keseriz
    contents = bytearray()
    while True:
        chunk = await file.read(1024 * 1024)   # 1 MB
        if not chunk:
            break
        contents.extend(chunk)
        if len(contents) > MAX_EXCEL_BYTES:
            raise HTTPException(
                status_code=413,
                detail=f"Dosya çok büyük (max {MAX_EXCEL_BYTES // (1024*1024)} MB)."
            )

    if len(contents) == 0:
        raise HTTPException(status_code=400, detail="Yüklenen dosya boş.")

    # 3) Magic bytes — uzantısı .xlsx ama içeriği farklıysa reddet
    if not any(bytes(contents[:4]).startswith(m) for m in _XLSX_MAGIC):
        raise HTTPException(
            status_code=400,
            detail="Dosya gerçekten bir Excel (.xlsx) değil. Header doğrulaması başarısız."
        )

    return bytes(contents)


_REQUIRED_COLS = {"Item_Name", "SKU", "Category", "Unit", "Stock", "Cost_Price", "Min_Stock_Level"}


@router.get("/import-items/template")
def download_import_template(_: dict = Depends(require_permission("items", "import"))):
    """Doldurulabilir örnek Excel şablonunu indir."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ürün Şablonu"

    headers = list(_REQUIRED_COLS)
    # Sabit sıralama
    headers = ["Item_Name", "SKU", "Category", "Unit", "Stock", "Cost_Price", "Min_Stock_Level"]
    ws.append(headers)

    # Örnek satırlar
    ws.append(["Argan Yağı",      "ARG-001",     "Hammadde",    "kg",   50,  25.50, 10])
    ws.append(["Cam Şişe 100ml",  "CAM-100",     "Ambalaj",     "adet", 200,  3.75, 50])
    ws.append(["Vitamin C Serum", "VIT-SRM-001", "Bitmiş Ürün", "adet",  0,  45.00, 20])

    # Header stili — lacivert/altın
    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill  = PatternFill("solid", fgColor="2C2C73")
    hdr_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="B8965A"),
        right=Side(style="thin",  color="E5E7EB"),
    )
    for cell in ws[1]:
        cell.font   = hdr_font
        cell.fill   = hdr_fill
        cell.alignment = hdr_align
        cell.border = thin_border

    # Zebra satırları
    alt_fill = PatternFill("solid", fgColor="F5F0E8")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.row % 2 == 0:
                cell.fill = alt_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Otomatik sütun genişliği
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col) + 4
        ws.column_dimensions[col[0].column_letter].width = min(max_len, 30)

    ws.row_dimensions[1].height = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=minerva108_urun_sablonu.xlsx"},
    )


@router.post("/import-items")
async def import_items_from_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    _: dict = Depends(require_permission("items", "import")),
    domain: str = Depends(active_domain),
):
    """Excel (.xlsx) dosyasından toplu ürün içe aktarma — SKU bazlı upsert."""
    import io, datetime as _dt
    actor = current_user.get("full_name") or current_user.get("username") or "Excel Import"

    # Uzantı + boyut + magic-bytes doğrulamayla güvenli okuma
    contents = await _read_xlsx_safely(file)

    # ── Pandas ile oku ──────────────────────────────────────────────────────
    try:
        import pandas as pd
        df = pd.read_excel(io.BytesIO(contents), dtype=str)   # hepsini str oku, sonra cast
        df.columns = [str(c).strip() for c in df.columns]     # boşluk temizle
    except Exception as exc:
        return JSONResponse(status_code=400, content={
            "detail": f"Excel dosyası okunamadı. Dosya bozuk olabilir. ({exc})"
        })

    # ── 3. Zorunlu sütun kontrolü ───────────────────────────────────────────
    missing_cols = _REQUIRED_COLS - set(df.columns)
    if missing_cols:
        return JSONResponse(status_code=422, content={
            "detail": f"Eksik sütunlar: {', '.join(sorted(missing_cols))}. "
                      "Lütfen örnek şablonu indirip kullanın."
        })

    # ── 4. Satır satır upsert ───────────────────────────────────────────────
    created, updated = 0, 0
    errors = []

    # Maliyet YALNIZ finans rolünden yazılır (update_item/create_item ile aynı
    # kural, "API ile bile"): items.import LabLead'de de açık; eskiden boş
    # Cost_Price hücresi NaN olarak karta yazılıp finans kullanıcısının
    # /api/items yanıtını 500'e düşürüyor, 0 yazılırsa maliyeti sessizce 0'lıyordu.
    finance_ok = _can_see_finance(current_user)

    def _num(val):
        """Hücreyi sonlu sayıya çevirir; boş/NaN/metin → None (alan yazılmaz)."""
        try:
            f = float(str(val).replace(",", "."))
        except (TypeError, ValueError):
            return None
        return f if math.isfinite(f) else None

    def _text(val) -> str:
        """Hücre metni; pandas boş hücreyi NaN (float) okur — "nan" yazılmasın."""
        if val is None or (isinstance(val, float) and math.isnan(val)):
            return ""
        return str(val).strip()

    for idx, row in df.iterrows():
        row_num = int(idx) + 2  # Excel satır no (header = 1)
        try:
            item_name = _text(row.get("Item_Name"))
            sku       = _text(row.get("SKU"))

            if not item_name:
                errors.append(f"Satır {row_num}: 'Item_Name' boş bırakılamaz — satır atlandı.")
                continue
            if not sku:
                errors.append(f"Satır {row_num}: 'SKU' boş bırakılamaz — satır atlandı.")
                continue

            category   = _text(row.get("Category")) or None
            unit       = _text(row.get("Unit")) or None
            stock      = _num(row.get("Stock"))
            min_stock  = _num(row.get("Min_Stock_Level"))
            cost_in    = _num(row.get("Cost_Price")) if finance_ok else None

            # SKU global tekil (DB kısıtı) → arama paneli ayırmaz; ama başka
            # panelin kartı bu panelden GÜNCELLENMEZ (domain sızıntısı).
            existing = db.query(Item).filter(Item.sku == sku).first()
            if existing and (existing.domain or "cosmetics") != domain:
                errors.append(f"Satır {row_num}: SKU '{sku}' başka panelin kartında — satır atlandı.")
                continue

            if existing:
                # ── GÜNCELLE ────────────────────────────────────────────
                # Boş/geçersiz hücre mevcut değeri KORUR — metinde de: boş
                # Unit kartı sessizce "adet" yapıp birim ailesini (g/kg → adet)
                # değiştirirdi (fiyatlar "uyuşmuyor", plan çevrimleri bozulur);
                # boş Category kartı Mal Kabul Hammadde/Ambalaj listesinden ve
                # plandan düşürürdü.
                existing.name           = item_name
                if category:
                    existing.category = category
                if unit:
                    existing.unit = unit
                if cost_in is not None:
                    existing.cost_price = cost_in
                if min_stock is not None:
                    existing.min_stock_level = min_stock
                if stock is not None and stock >= 0:
                    existing.current_stock = round(stock, 6)
                updated += 1

            else:
                # ── OLUŞTUR ─────────────────────────────────────────────
                stock = stock or 0.0
                new_item = Item(
                    name=item_name, sku=sku, category=category,
                    unit=unit or "adet", current_stock=round(stock, 6),
                    cost_price=cost_in if cost_in is not None else 0.0,
                    min_stock_level=min_stock or 0.0,
                    domain=domain,                      # Faz 3 — aktif panel
                )
                db.add(new_item)
                db.flush()   # ID'yi al

                if stock > 0:
                    lot = f"IMP-{_dt.datetime.utcnow().strftime('%Y%m%d-%H%M%S')}-{new_item.id}"
                    db.add(Inventory(
                        item_id=new_item.id, lot_number=lot,
                        quantity=stock, status="APPROVED",
                        received_by=actor, domain=domain,
                    ))
                    db.add(Transaction(
                        item_id=new_item.id, lot_number=lot,
                        transaction_type="Input", quantity=stock,
                        notes=f"Excel içe aktarım — SKU: {sku}",
                        performed_by=actor,
                    ))
                created += 1

        except Exception as exc:
            db.rollback()
            errors.append(f"Satır {row_num}: İşlem hatası — {str(exc)[:100]}")

    # ── 5. Commit ────────────────────────────────────────────────────────────
    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        return JSONResponse(status_code=500, content={
            "detail": f"Veritabanına kayıt sırasında hata oluştu: {str(exc)[:120]}"
        })

    suffix = f" {len(errors)} satırda hata oluştu." if errors else ""
    return {
        "created": created,
        "updated": updated,
        "error_count": len(errors),
        "errors": errors,
        "message": f"{created} yeni ürün eklendi, {updated} ürün güncellendi.{suffix}",
    }


# ─── Smart Excel Importer (Phase 4 / Task 1) ────────────────────────────────

def _parse_qty_unit(raw):
    """
    Parse messy quantity strings from the Hammadde sheets.
      '4.276GR'  → (4276, 'g')      (Turkish dot = thousands separator)
      '1 KG'     → (1000, 'g')
      '90GR'     → (90, 'g')
      '162ML'    → (162, 'ml')
      '21KG'     → (21000, 'g')
      ''/None    → (0, 'adet')
    """
    import re
    if raw is None:
        return (0.0, "adet")
    s = str(raw).strip().upper()
    if not s or s in ("—", "-", "NAN"):
        return (0.0, "adet")

    # Detect unit suffix
    m = re.search(r'(KG|ML|GR|G)\b\s*$', s)
    if m:
        unit_str = m.group(1)
        s = s[:m.start()].strip()
    else:
        unit_str = "GR"

    # Comma → dot for decimals (Turkish comma)
    s = s.replace(",", ".")
    # Treat XX.XXX (3 trailing digits after dot) as a thousands separator
    if re.match(r'^\d+\.\d{3}$', s):
        s = s.replace(".", "")

    try:
        qty = float(s)
    except ValueError:
        qty = 0.0

    if unit_str == "KG":
        return (qty * 1000, "g")
    if unit_str == "ML":
        return (qty, "ml")
    return (qty, "g")


def _detect_excel_schema(workbook) -> str:
    """
    Sniff the workbook to identify which schema it follows. Returns one of:
      'hammadde' — multi-sheet raw-material catalogue with HAMMADDE/GR/MARKASI columns
      'etiket'   — packaging/label inventory (ÜRÜN İSMİ + ML + ETİKET SAYISI)
      'standard' — the existing minimal Item_Name/SKU/Category template
      'unknown'
    """
    keywords_hammadde = {"HAMMADDE", "MARKASI"}
    keywords_etiket   = {"ETİKET", "ÜRÜN İSİM", "ÜRÜN İSMİ"}
    keywords_standard = {"ITEM_NAME", "SKU"}

    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        # Scan first 5 rows
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            cells = [str(c).strip().upper() for c in row if c is not None]
            joined = " | ".join(cells)
            if any(k in joined for k in keywords_hammadde):
                return "hammadde"
            if any(k in joined for k in keywords_etiket):
                return "etiket"
            if any(k in joined for k in keywords_standard):
                return "standard"
    return "unknown"


def _parse_hammadde_workbook(workbook) -> list:
    """
    Walk every sheet in a Hammadde workbook. Returns list of dicts:
      { sheet, name, quantity, unit, supplier, raw_qty }
    Skips header rows ('KONTROL TARİHİ' or starts with 'DOLAP') and blank rows.
    """
    out = []
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        for row in ws.iter_rows(min_row=1, values_only=True):
            row_norm = [c if c is not None else "" for c in row]
            if len(row_norm) < 4:
                continue

            col_a = str(row_norm[0]).strip()
            col_b = str(row_norm[1]).strip()
            col_c = row_norm[2]
            col_d = str(row_norm[3]).strip() if len(row_norm) > 3 else ""

            # Skip blanks and section headers
            if not col_b:
                continue
            joined_left = f"{col_a} {col_b}".upper()
            if "KONTROL TARİHİ" in joined_left:        # Header row (multiple sheets repeat it)
                continue
            if "DOLAP" in joined_left and "RAF" in joined_left:   # Cabinet/shelf section header — anywhere in left two cols
                continue
            if "DOLAP" in joined_left and "ÜST RAF" in joined_left:
                continue
            # Skip rows where col_b is itself the column header
            if col_b.upper().strip() in ("HAMMADDE İSİM", "HAMMADDE İSIM"):
                continue
            # Skip if col_b looks like a section header by itself
            if col_b.upper().strip().startswith("DOLAP"):
                continue

            qty, unit = _parse_qty_unit(col_c)
            out.append({
                "sheet":    sheet_name,
                "name":     col_b,
                "quantity": qty,
                "unit":     unit,
                "supplier": col_d if col_d and col_d not in ("—", "-", "NAN") else None,
                "raw_qty":  str(col_c) if col_c is not None else "",
            })
    return out


@router.post("/admin/smart-import")
async def smart_excel_import(
    file: UploadFile = File(...),
    commit: bool = False,
    _: dict = Depends(require_permission("admin", "import_excel")),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
    domain: str = Depends(active_domain),
):
    """
    Smart Excel importer.
      - Auto-detects schema (Hammadde / Etiket / Standard)
      - Always returns a dry-run preview (commit=false default)
      - When commit=true: creates Items, Suppliers, Inventory + Transaction logs
      - Existing items are upserted (stock added, supplier filled if missing)
    """
    import io
    import openpyxl
    from datetime import datetime as _dt

    actor = current_user.get("full_name") or current_user.get("username") or "Excel Bulk Import"

    # Uzantı + boyut + magic-bytes doğrulama (HTTPException döner — global handler yakalar)
    contents = await _read_xlsx_safely(file)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=False)
    except Exception as e:
        return JSONResponse(status_code=400, content={"detail": f"Excel dosyası okunamadı: {e}"})

    schema = _detect_excel_schema(wb)
    if schema == "unknown":
        return JSONResponse(status_code=422, content={
            "detail": "Dosya yapısı tanınamadı. Hammadde, Etiket veya Standart formatlardan biri olmalı.",
            "detected_schema": schema,
        })

    # ── Currently Hammadde is the production-ready parser; others return preview only.
    if schema != "hammadde":
        return JSONResponse(status_code=422, content={
            "detail": f"Bu sürümde sadece Hammadde formatı işlenebiliyor. Algılanan: '{schema}'. "
                      "Etiket dosyaları için ayrı içe aktarım akışı yakında eklenecektir.",
            "detected_schema": schema,
        })

    parsed = _parse_hammadde_workbook(wb)
    if not parsed:
        return JSONResponse(status_code=422, content={"detail": "Sayfalardan veri okunamadı."})

    # ── Pre-flight analysis: classify each row as create vs update ──────────
    name_index = {i.name.strip().upper(): i for i in db.query(Item).filter(Item.is_active == True).all()}
    # Tedarikçi dizini aktif panelin BÜTÜN firmaları (pasifler dahil, aktif
    # olan aynı adı ezer): pasife alınmış firmanın adıyla gelen satır aynı
    # adlı YENİ bir tedarikçi açmasın.
    supp_index = {}
    for s in (db.query(Supplier).filter(Supplier.domain == domain)
              .order_by(Supplier.is_active.asc(), Supplier.id.asc()).all()):
        supp_index[s.name.strip().upper()] = s

    plan = {
        "schema":               schema,
        "sheets_processed":     len(wb.sheetnames),
        "rows_parsed":          len(parsed),
        "items_to_create":      0,
        "items_to_update":      0,
        "suppliers_to_create":  0,
        "warnings":             [],
        "preview":              [],
    }

    new_supplier_keys = set()
    for row in parsed:
        key = row["name"].strip().upper()
        existing = name_index.get(key)
        action = "update" if existing else "create"
        if action == "create":
            plan["items_to_create"] += 1
        else:
            plan["items_to_update"] += 1

        if row["supplier"]:
            sk = row["supplier"].strip().upper()
            if sk not in supp_index and sk not in new_supplier_keys:
                new_supplier_keys.add(sk)
                plan["suppliers_to_create"] += 1

        if row["quantity"] <= 0:
            plan["warnings"].append(f"Sıfır miktar: '{row['name']}' (sayfa: {row['sheet']}, ham değer: '{row['raw_qty']}')")

        if len(plan["preview"]) < 25:   # Cap preview for response size
            plan["preview"].append({
                "sheet":    row["sheet"],
                "name":     row["name"],
                "quantity": row["quantity"],
                "unit":     row["unit"],
                "supplier": row["supplier"] or "—",
                "action":   action,
            })

    # ── If dry run, stop here ──
    if not commit:
        plan["committed"] = False
        return plan

    # ── COMMIT: idempotent upserts ──────────────────────────────────────────
    items_created    = 0
    items_updated    = 0
    suppliers_added  = 0
    inv_rows_created = 0
    txs_logged       = 0

    try:
        for row in parsed:
            key = row["name"].strip().upper()
            supplier_obj = None
            if row["supplier"]:
                sk = row["supplier"].strip().upper()
                supplier_obj = supp_index.get(sk)
                if not supplier_obj:
                    supplier_obj = Supplier(name=row["supplier"], is_active=True, domain=domain)
                    db.add(supplier_obj)
                    db.flush()
                    supp_index[sk] = supplier_obj
                    suppliers_added += 1

            existing = name_index.get(key)
            if existing:
                # Add stock to existing item
                existing.current_stock = round((existing.current_stock or 0) + row["quantity"], 6)
                # Kartın varsayılan tedarikçisi yalnız AKTİF firmadan (pasif
                # firma lota bağlanır, karta yeni atanmaz).
                if supplier_obj and supplier_obj.is_active is not False and not existing.supplier_id:
                    existing.supplier_id = supplier_obj.id
                item = existing
                items_updated += 1
            else:
                item = Item(
                    name=row["name"],
                    category="Hammadde",
                    unit=row["unit"],
                    current_stock=row["quantity"],
                    cost_price=0.0,
                    supplier_id=(supplier_obj.id if supplier_obj and supplier_obj.is_active is not False
                                 else None),
                    is_active=True,
                    domain=domain,                      # Faz 3 — aktif panel
                )
                db.add(item)
                db.flush()
                name_index[key] = item
                items_created += 1

            # Inventory + Transaction (only when there's actual stock)
            if row["quantity"] > 0:
                lot_no = f"XLS-{_dt.utcnow().strftime('%Y%m%d')}-{item.id}"
                db.add(Inventory(
                    item_id=item.id,
                    supplier_id=supplier_obj.id if supplier_obj else None,
                    lot_number=lot_no,
                    quantity=row["quantity"],
                    status="APPROVED",
                    received_by=actor,
                    domain=domain,
                ))
                inv_rows_created += 1

                db.add(Transaction(
                    item_id=item.id,
                    lot_number=lot_no,
                    transaction_type="Input",
                    quantity=row["quantity"],
                    notes=f"Excel Bulk Import — Sheet: {row['sheet']} · Raw: '{row['raw_qty']}'",
                    performed_by=actor,
                ))
                txs_logged += 1

        db.commit()
    except Exception as e:
        db.rollback()
        return JSONResponse(status_code=500, content={"detail": f"İçe aktarım sırasında hata: {e}"})

    plan.update({
        "committed":       True,
        "items_created":   items_created,
        "items_updated":   items_updated,
        "suppliers_added": suppliers_added,
        "inventory_rows":  inv_rows_created,
        "transactions":    txs_logged,
    })
    return plan
