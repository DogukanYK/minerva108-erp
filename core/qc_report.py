# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Kalite Kontrol formu — okunur görünüm + PDF/Excel üretimi.

  parse_qc_form(inv, item) -> dict | None    Inventory.qc_form_data JSON'unu
                                             soru metinleriyle etiketlenmiş,
                                             tek bir okunur sözlüğe çevirir.
  render_qc_pdf(view)   -> bytes             reportlab ile resmi PDF
  render_qc_excel(view) -> bytes             openpyxl ile tek sayfalı xlsx

Soru metinleri core/qc_questions.py'den (tek kaynak) gelir; font yönetimi
core/monthly_report.py'deki Türkçe-uyumlu `_register_fonts` ile paylaşılır.
"""
import json
from datetime import datetime
from io import BytesIO
from typing import Optional
from xml.sax.saxutils import escape

from database import to_tr
from core.qc_questions import QC_QUESTIONS, QC_QUESTION_TEXT


# ─── Parse: qc_form_data JSON → okunur görünüm ──────────────────────────────

def parse_qc_form(inv, item=None) -> Optional[dict]:
    """Inventory satırının QC formunu etiketlenmiş bir dict'e çevirir.

    Form yoksa veya bozuksa None döner."""
    raw = getattr(inv, "qc_form_data", None)
    if not raw:
        return None
    try:
        data = json.loads(raw)
    except Exception:
        return None

    checklist_raw = data.get("checklist") or {}
    checklist = []
    seen = set()
    # Kanonik sırada soru metinleriyle eşle
    for q in QC_QUESTIONS:
        seen.add(q["id"])
        checklist.append({"text": q["text"], "answer": checklist_raw.get(q["id"], "—")})
    # core/qc_questions.py'de olmayan (gelecekte eklenmiş) anahtarları da göster
    for k, v in checklist_raw.items():
        if k not in seen:
            checklist.append({"text": QC_QUESTION_TEXT.get(k, k), "answer": v})

    status = (data.get("status") or getattr(inv, "status", "") or "").upper()
    submitted_tr = "—"
    submitted = data.get("submitted_at")
    if submitted:
        try:
            submitted_tr = to_tr(datetime.fromisoformat(submitted)).strftime("%d.%m.%Y %H:%M")
        except Exception:
            submitted_tr = "—"

    return {
        "lot_number":    inv.lot_number,
        "item_name":     item.name if item else "—",
        "item_category": item.category if item else "—",
        "quantity":      inv.quantity,
        "unit":          (item.unit if item else "") or "",
        "status":        status,
        "status_label":  "Onaylandı" if status == "APPROVED"
                         else ("Reddedildi" if status == "REJECTED" else (status or "—")),
        "is_approved":   status == "APPROVED",
        "approved_by":   getattr(inv, "qc_approved_by", None) or "—",
        "submitted_at":  submitted_tr,
        "checklist":     checklist,
        "lab_ml":        data.get("lab_ml"),
        "lab_density":   data.get("lab_density"),
        "lab_color":     data.get("lab_color"),
        "notes":         data.get("notes") or "",
    }


def qc_export_filename(view: dict, ext: str) -> str:
    """qc_formu_{lot}.{ext} — güvenli dosya adı (yalnızca alfasayısal + - _)."""
    import re
    lot = re.sub(r"[^A-Za-z0-9_-]+", "_", str(view.get("lot_number") or "lot"))
    return f"qc_formu_{lot}.{ext}"


# ─── PDF (reportlab) ────────────────────────────────────────────────────────

def render_qc_pdf(view: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)
    from core.monthly_report import _register_fonts

    font, font_b = _register_fonts()
    NAVY  = colors.HexColor("#232E6E")
    LIGHT = colors.HexColor("#f5f0e8")
    GREY  = colors.HexColor("#e5e7eb")
    GREEN = colors.HexColor("#15803d")
    RED   = colors.HexColor("#b91c1c")

    st_h1   = ParagraphStyle("h1", fontName=font_b, fontSize=16, textColor=NAVY,
                             spaceAfter=2, leading=19)
    st_meta = ParagraphStyle("meta", fontName=font, fontSize=8,
                             textColor=colors.HexColor("#6b7280"), leading=11)
    st_h2   = ParagraphStyle("h2", fontName=font_b, fontSize=11, textColor=NAVY,
                             spaceBefore=13, spaceAfter=5, leading=14)
    st_cell = ParagraphStyle("cell", fontName=font, fontSize=8.5,
                             textColor=colors.HexColor("#374151"), leading=11)
    st_hcell = ParagraphStyle("hcell", fontName=font_b, fontSize=8.5,
                              textColor=colors.white, leading=11)
    st_badge = ParagraphStyle("badge", fontName=font_b, fontSize=12,
                              textColor=(GREEN if view["is_approved"] else RED), leading=15)
    st_note = ParagraphStyle("note", fontName=font, fontSize=9,
                             textColor=colors.HexColor("#374151"), leading=13)

    W = 174.0  # kullanılabilir içerik genişliği (mm)
    story = []

    def _kv_table(rows, widths):
        data = [[Paragraph(escape(str(k)), st_cell), Paragraph(escape(str(v)), st_cell)]
                for k, v in rows]
        t = Table(data, colWidths=widths)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, GREY),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        return t

    def _grid_table(headers, rows, widths):
        head = [Paragraph(escape(str(h)), st_hcell) for h in headers]
        body = [[Paragraph(escape(str(c)), st_cell) for c in r] for r in rows]
        t = Table([head] + body, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("GRID", (0, 0), (-1, -1), 0.4, GREY),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        return t

    # Başlık + onay rozeti
    story.append(Paragraph("Minerva 108 — Kalite Kontrol Formu", st_h1))
    story.append(Paragraph(
        f"Ürün: {escape(str(view['item_name']))}  ·  Lot: {escape(str(view['lot_number']))}", st_meta))
    story.append(Spacer(1, 6))
    badge = "● UYGUNDUR (APPROVED)" if view["is_approved"] else "● UYGUN DEĞİLDİR (REJECTED)"
    story.append(Paragraph(badge, st_badge))

    # Künye
    story.append(Paragraph("Lot Bilgisi", st_h2))
    story.append(_kv_table([
        ("Ürün",        view["item_name"]),
        ("Kategori",    view["item_category"]),
        ("Lot / Parti", view["lot_number"]),
        ("Miktar",      f"{view['quantity']} {view['unit']}".strip()),
        ("QC Yapan",    view["approved_by"]),
        ("Form Tarihi", view["submitted_at"]),
        ("Sonuç",       view["status_label"]),
    ], [W * 0.32 * mm, W * 0.68 * mm]))

    # Kontrol listesi
    story.append(Paragraph("Kontrol Listesi", st_h2))
    story.append(_grid_table(
        ["#", "Soru", "Cevap"],
        [[i + 1, c["text"], c["answer"]] for i, c in enumerate(view["checklist"])],
        [W * 0.08 * mm, W * 0.72 * mm, W * 0.20 * mm]))

    # Lab kriterleri (varsa)
    lab_rows = []
    if view.get("lab_ml") is not None:
        lab_rows.append(("İçindeki ml", str(view["lab_ml"])))
    if view.get("lab_density") is not None:
        lab_rows.append(("Yoğunluk (g/cm³)", str(view["lab_density"])))
    if view.get("lab_color"):
        lab_rows.append(("Renk", str(view["lab_color"])))
    if lab_rows:
        story.append(Paragraph("Laboratuvar Kriterleri", st_h2))
        story.append(_kv_table(lab_rows, [W * 0.32 * mm, W * 0.68 * mm]))

    # Notlar
    story.append(Paragraph("Gözlem / Notlar", st_h2))
    story.append(Paragraph(escape(view["notes"] or "—") or "—", st_note))

    story.append(Spacer(1, 16))
    story.append(Paragraph(
        "Bu form Minerva 108 ERP kalite kontrol kaydından üretilmiştir.", st_meta))

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"QC Formu — {view['lot_number']}", author="Minerva 108 ERP")
    doc.build(story)
    return buf.getvalue()


# ─── Excel (openpyxl) ───────────────────────────────────────────────────────

def render_qc_excel(view: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "QC Formu"

    navy = PatternFill("solid", fgColor="232E6E")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=13, color="232E6E")
    label_font = Font(bold=True)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 22

    r = 1
    ws.cell(r, 1, "Minerva 108 — Kalite Kontrol Formu").font = title_font
    r += 2

    def kv(label, value):
        nonlocal r
        c1 = ws.cell(r, 1, label); c1.font = label_font
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 2, value)
        r += 1

    kv("Ürün", view["item_name"])
    kv("Kategori", view["item_category"])
    kv("Lot / Parti", view["lot_number"])
    kv("Miktar", f"{view['quantity']} {view['unit']}".strip())
    kv("Sonuç", view["status_label"])
    kv("QC Yapan", view["approved_by"])
    kv("Form Tarihi", view["submitted_at"])
    r += 1

    # Kontrol listesi başlığı
    for ci, h in enumerate(["#", "Soru", "Cevap"], 1):
        c = ws.cell(r, ci, h); c.fill = navy; c.font = head_font
        c.alignment = Alignment(horizontal="left")
    r += 1
    for i, item in enumerate(view["checklist"], 1):
        ws.cell(r, 1, i)
        ws.cell(r, 2, item["text"])
        ws.cell(r, 3, item["answer"])
        r += 1
    r += 1

    # Lab kriterleri
    if view.get("lab_ml") is not None or view.get("lab_density") is not None or view.get("lab_color"):
        ws.cell(r, 1, "Laboratuvar Kriterleri").font = label_font
        r += 1
        if view.get("lab_ml") is not None:        kv("İçindeki ml", view["lab_ml"])
        if view.get("lab_density") is not None:   kv("Yoğunluk (g/cm³)", view["lab_density"])
        if view.get("lab_color"):                 kv("Renk", view["lab_color"])
        r += 1

    ws.cell(r, 1, "Notlar").font = label_font
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, view["notes"] or "—")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
