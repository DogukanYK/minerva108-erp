# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""SKT kontrol raporu — SKT'si eksik / okunamayan / geçmiş hammadde lotları.

Neden: B2B sipariş partisi (core/b2b_orders.strict_lot_ok) ve fason sevk,
SKT'si okunur ve geçmemiş olmayan hammadde lotunu KULLANMAZ.  Mal kabulde SKT
serbest metin olduğu için eski lotların bir kısmı boş ya da okunamaz durumda;
lab bunları ilk siparişten önce görüp düzeltebilmeli (İzlenebilirlik sayfası,
`PATCH /api/inventory/lots/{id}/expiry`).

Kapsam: aktif panelin numune OLMAYAN, miktarı > 0, APPROVED ya da QC bekleyen
(QUARANTINE) lotları; aktif kartlar.  "Hammadde" = tüketim motorundaki tür
(`core.consumption._kind`: Ambalaj dışı her şey) — sıkı havuzun denetlediği.
  scope="recipes" (varsayılan): aktif reçetelerde geçen hammaddeler
  scope="all": ayrıca reçetede olmayan hammadde kartları (Bitmiş Ürün hariç)
Ayrıştırıcı ve "bugün" sıkı havuzla AYNI (core/lots.parse_expiry / expiry_today)
→ raporda olmayan lot sıkı havuzda SKT yüzünden elenmez.
"""
from collections import defaultdict
from io import BytesIO

from sqlalchemy import or_

from core.lots import expiry_today, parse_expiry
from database import Inventory, Item, Recipe, RecipeIngredient, Supplier, to_tr

ISSUES = {"expired": "SKT geçmiş", "missing": "SKT girilmemiş", "unreadable": "SKT okunamıyor"}
_ORDER = {"expired": 0, "missing": 1, "unreadable": 2}
SCOPES = ("recipes", "all")


def _fold(text):
    from core.supplier_prices import normalize
    return normalize(text or "")


def lot_issue(lot, today):
    """(sorun, ayrıştırılan tarih) — sorun yoksa (None, tarih)."""
    raw = (lot.expiry_date or "").strip()
    if not raw:
        return "missing", None
    exp = parse_expiry(raw)
    if exp is None:
        return "unreadable", None
    if exp < today:
        return "expired", exp
    return None, exp


def editable(lot) -> bool:
    """SKT bu raporda düzeltilebilir mi — fason kabul lotu kendi kaydından gelir."""
    return (not lot.is_sample and lot.outsourcing_receipt_id is None
            and (lot.status or "") in ("APPROVED", "QUARANTINE"))


def expiry_issues(db, domain, scope="recipes", today=None):
    scope = scope if scope in SCOPES else "recipes"
    today = today or expiry_today()
    recipes = defaultdict(set)                           # item_id → aktif reçete adları
    for item_id, name in (db.query(RecipeIngredient.item_id, Recipe.name)
                          .join(Recipe, Recipe.id == RecipeIngredient.recipe_id)
                          .filter(Recipe.domain == domain, Recipe.is_active == True)):   # noqa: E712
        recipes[item_id].add(name)

    q = (db.query(Inventory, Item, Supplier)
         .join(Item, Item.id == Inventory.item_id)
         .outerjoin(Supplier, Supplier.id == Inventory.supplier_id)
         .filter(Inventory.domain == domain, Inventory.is_sample == False,                 # noqa: E712
                 Inventory.status.in_(("APPROVED", "QUARANTINE")), Inventory.quantity > 1e-9,
                 Item.is_active == True,                                                   # noqa: E712
                 or_(Item.category.is_(None), Item.category != "Ambalaj")))
    rows, summary = [], {k: 0 for k in ISSUES}
    for lot, item, sup in q:
        used = recipes.get(item.id)
        if scope == "recipes" and not used:
            continue
        if scope == "all" and not used and item.category == "Bitmiş Ürün":
            continue
        issue, exp = lot_issue(lot, today)
        if issue is None:
            continue
        summary[issue] += 1
        names = sorted(used or (), key=_fold)
        rows.append({
            "id": lot.id, "item_id": item.id, "item_name": item.name, "sku": item.sku or "",
            "category": item.category or "", "unit": item.unit or "",
            "lot_number": lot.lot_number or "", "supplier": sup.name if sup else "",
            "quantity": round(float(lot.quantity or 0), 6), "status": lot.status,
            "qc_pending": bool(lot.qc_required) or lot.status == "QUARANTINE",
            "location": lot.location or "", "expiry_date": lot.expiry_date or "",
            "expiry_iso": exp.isoformat() if exp else "",
            "days_expired": (today - exp).days if (issue == "expired" and exp) else None,
            "issue": issue, "issue_label": ISSUES[issue],
            "recipes": names[:5], "recipe_count": len(names),
            "received_by": lot.received_by or "",
            "received_at": to_tr(lot.created_at).strftime("%d.%m.%Y") if lot.created_at else "",
            "editable": editable(lot),
        })
    rows.sort(key=lambda r: (_ORDER[r["issue"]], _fold(r["item_name"]), r["lot_number"]))
    summary["total"] = len(rows)
    return {"scope": scope, "today": today.isoformat(), "summary": summary, "rows": rows}


def build_workbook(data, domain_label="") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "SKT Sorunları"
    scope = "Aktif reçetelerde geçen hammaddeler" if data["scope"] == "recipes" else "Tüm hammaddeler"
    s = data["summary"]
    ws.append([f"SKT kontrol raporu — {domain_label}".strip(" —")])
    ws.append([f"Kapsam: {scope} · Tarih: {data['today']} · Geçmiş {s['expired']} · "
               f"Girilmemiş {s['missing']} · Okunamayan {s['unreadable']}"])
    ws.append([])
    head = ["Sorun", "Malzeme", "SKU", "Lot no", "Tedarikçi", "Miktar", "Birim", "Durum",
            "Kayıtlı SKT", "Geçen gün", "Konum", "Kabul tarihi", "Kabul eden", "Reçeteler"]
    ws.append(head)
    hrow = ws.max_row
    fill = PatternFill("solid", fgColor="D9D9D9")
    for c in ws[hrow]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    colors = {"expired": "FDE2E2", "missing": "FEF3C7", "unreadable": "E0E7FF"}
    for r in data["rows"]:
        recipes = ", ".join(r["recipes"]) + (f" (+{r['recipe_count'] - len(r['recipes'])})"
                                             if r["recipe_count"] > len(r["recipes"]) else "")
        ws.append([r["issue_label"], r["item_name"], r["sku"], r["lot_number"], r["supplier"],
                   r["quantity"], r["unit"], "QC bekliyor" if r["qc_pending"] else "Onaylı",
                   r["expiry_date"], r["days_expired"], r["location"], r["received_at"],
                   r["received_by"], recipes])
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=colors[r["issue"]])
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    widths = [16, 34, 14, 16, 22, 10, 7, 12, 14, 10, 14, 13, 18, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(head))}{max(ws.max_row, hrow)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
