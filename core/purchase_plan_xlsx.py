# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Satın Alma Planı — Excel çıktısı.

    build_workbook(report) -> bytes

`report` = `compute()` + `attach()` çıktısı (PDF ile aynı sözlük).  Sayfalar:
  • "Yönetici Özeti" (16.10.2026 yeniden tasarım; İLK sayfa) — PDF Yönetici
    Özeti'yle aynı veri (`purchase_pricing.summary_data`): göstergeler, tür
    bazında maliyet (sütunlar PDF'teki gibi: Stoktan yalnız net modda, Yoldan
    yalnız açık siparişte), tedarikçi bazında alınacaklar (ilk 8 + Diğer), hammadde
    ihtiyacı (kg / l ayrı, ilk 12 + Diğer; alınacak hammadde toplamı), fiyatı
    eksik kalemler, uyarılar ve bekletilenler.  Değer hücreleri sayı (formül
    yok); öteki sayfaların sırası değişmez.
  • "Alım listesi"  — bölüm sırası `purchase_pricing.sections()` ile AYNI
    (1–5: hammadde/ambalaj/etiket, fiyatlı → fiyatsız), sonra yeni etiket /
    koli / palet satırları.  Formüller:
        G (Alınacak) = MAX(0, ROUND(D−E−F, 3))      (brüt modda = D)
        K (Tutar)    = IF(J="", "", ROUND(G*J, 0))   (Excel ROUND = round_half_up)
    + bölüm ara toplamları ve GENEL TOPLAM (SUM).  Lab Excel'de fiyatı ya da
    elimizdekini değiştirince tutar kendiliğinden güncellenir.  "Not" (X)
    sonrasında Y "Aynı malzeme (diğer kartlar)", Z "Sipariş geçmişi" ve AA
    "Grup stoğu (kart dökümü)" (`purchase_plan.stock_note`, önizleme/PDF ile
    aynı metin).
  • "Stoktan karşılananlar" (14.10.2026; YALNIZ net modda ve satır varsa,
    "Alım listesi"nden hemen sonra) — önce "Yeterli" satırlar, sonra
    kısmen stoktan / yoldan karşılanan alınacaklar.  Formüller:
        G (Kullanılan)  = MIN(D, E)                      (planın stoktan kullandığı)
        K (Değer)       = IF(J="", "", ROUND(G*J, 0))
      açık sipariş seçeneği açıkken ayrıca
        L (Yoldan)      = MIN(F, ROUND(D−G, 6))
        M (Yoldan değer)= IF(J="", "", ROUND(L*J, 0))
    G / L satırın `stock_qty_used` / `open_qty_used`'ıdır; formülle aynı
    sayıyı vermiyorsa (düz yuvarlamada alınacağı tama yuvarlanan adet
    satırı: stoktan = gereken − alınacak) SAYI yazılır — Alım listesi G'si
    gibi.  K / M formülü her satırda kalır.
    + TOPLAM (SUM) — `pricing.cost_summary` stoktan / yoldan değeriyle aynı.
    "Alım listesi"nin sütunları, formülleri ve GENEL TOPLAM'ı DEĞİŞMEZ.
  • "Tedarikçiler" — seçilen + aday firmalar, iletişim, kart adları, K
    "İlişki türleri" (stok kartı / alım / numune / sipariş / eşdeğer kart …).
  • "Ürünler"      — hedef adet, bitmiş stok, kapasite, kısıtlayan malzeme.
  • "Ürün maliyeti" (16.10.2026; fiyatlı raporda, "Ürünler"den sonra) —
    üstte ürün özeti (1 adet maliyeti + toplam, fiyatı eksik kalemler) ve
    TOPLAM; altında ürün × malzeme dökümü (H = IF(G="","",F*G), J = F*I,
    K = IF(G="","",J*G), ara toplam H = SUM / K = ROUND(SUM,0)); tür
    bazında maliyet (Σ = planın toplam malzeme maliyeti), fiyatı eksik
    kalemler, notlar.  Alım listesi'nde ambalaj uyarılı satırın K'si
    turuncu + uyarı hücre yorumu (X "Not"ta da yazar).
  • "Ürün detayı"  — ürün × malzeme ihtiyaç matrisi (kalem biriminde).
  • "Özet"         — kapsam, toplamlar (net modda GENEL TOPLAM'ın altında
    stoktan / yoldan değer + PLANIN TOPLAM MALZEME MALİYETİ + fiyatsız kalem
    sayıları ve uyarılar — `cost_summary_lines` / `_caveats`), bölüm
    listesi, notlar.

Formül hücrelerinin HESAPLANMIŞ değeri `core.xlsx_cache.inject_cached_values`
ile gömülür: formül hesaplamayan önizlemelerde (Quick Look, WhatsApp, Drive)
"Tutar" boş görünmesin.  Gömülen değer Excel'in formülden bulacağı değerin
AYNISIDIR (G için Excel ROUND taklidi; tutmuyorsa — `safe_rounding` kapalı
gibi — G formül yerine raporun alınacak sayısı olarak yazılır ki PDF ile
Excel toplamı ayrışmasın).

Biçim (16.10.2026): başlık satırı lacivert dolgu / beyaz kalın yazı, veri
satırlarında kıl alt çizgi + zebra (ara toplam / başlık satırları hariç),
tek bloklu sayfalarda (Tedarikçiler, Ürünler, Ürün detayı) otomatik süzgeç —
ara toplamlı Alım listesi / Stoktan karşılananlar'da YOK (sıralama SUM
aralığını bozardı) —, sabit başlık (freeze), sayı biçimleri;
her sayfada A4 yazdırma ayarı (`ingredients_report._setup_a4`, üst bilgi
"Minerva 108 — Satın Alma Planı · <başlık>", 1. satır her baskı sayfasında).

Formül kalkanı: openpyxl "=" ile başlayan HER metni formül sayar.  Malzeme /
tedarikçi adı, not, iletişim gibi serbest metinler kullanıcıdan gelir
("=HYPERLINK(…)" ya da "=2 adet kırık" notu) → hepsi `_put_text`/
`_append_as_text` üzerinden düz metin + quotePrefix yazılır; formül yalnız
bu modülün kendi kurduğu hücrelerde bulunur (core/retention_report.py ile
aynı kalıp).  Aynı iki kapı XML'e yazılamayan kontrol karakterlerini
(\x00-\x08, \x0b, \x0c, \x0e-\x1f) de siler: openpyxl bunlarda
IllegalCharacterError atar ve adında tek bir \x01 olan kart bütün Excel
indirmesini 500'e düşürürdü (`clean_text` yalnız boşluk sınıfını temizler).
"""
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Dict, List, Optional

from core.purchase_plan import _amount_near, stock_note
from core.purchase_pricing import (CURRENCY_SYMBOL, KIND_TITLE, REL_TYPE_TEXT, UNIT_TEXT, alt_texts, amount_warn,
                                   contact_line, cost_overview_view, cost_summary_caveats, cost_summary_lines,
                                   money, notes_lines, offer_tags, order_texts, pkg_round_text,
                                   product_cost_lines, round6, round_half_up, sections, summary_data,
                                   unpriced_offer_text)
from core.xlsx_cache import inject_cached_values

SHEET_LIST = "Alım listesi"
SHEET_STOCK = "Stoktan karşılananlar"
SHEET_SUP = "Tedarikçiler"
SHEET_PROD = "Ürünler"
SHEET_COST = "Ürün maliyeti"
SHEET_DETAIL = "Ürün detayı"
SHEET_SUMMARY = "Özet"
SHEET_EXEC = "Yönetici Özeti"

HEADER_FILL = "232E6E"       # lacivert başlık dolgusu (report_theme PALETTE navy), beyaz kalın yazı
HEADER_TEXT = "FFFFFF"
ZEBRA_FILL = "F5F6F9"
LINE_COLOR = "D9DDE6"
TITLE_COLOR = "232E6E"
MUTED_COLOR = "5B6170"
MONEY_FMT = "#,##0"
WARN_COLOR = "C2410C"         # ambalaj uyarılı Tutar (K) hücresi — turuncu yazı + hücre yorumu
KIND_TEXT = {"raw": "Hammadde", "packaging": "Ambalaj", "label": "Etiket"}
GROUP_TEXT = {"list": "Fiyat listesinde var", "none": "Fiyat yok — teklif alınacak"}
MANUAL_SECTION_TEXT = {"label": "Etiket", "shipping": "Koli / palet", "other": "Diğer"}
MATERIAL_SECTIONS = ("raw_priced", "raw_unpriced", "pkg_priced", "pkg_unpriced", "labels")

# Alım listesi sütunları (spec C4 tablosu).  "Not" X'te (24) sabit; "Aynı
# malzeme (diğer kartlar)" Y, "Sipariş geçmişi" Z ve "Grup stoğu (kart
# dökümü)" AA (27, 13.10.2026) onun SONRASINA eklendi — formüller D–K, eski
# sütunlar kaymaz.
COLS = list("ABCDEFGHIJKLMNOPQRSTUVWXYZ") + ["AA"]
WIDTHS = (10, 22, 34, 12, 12, 10, 12, 7, 20, 11, 11, 24, 18, 9, 10, 18, 9, 18, 9, 22, 24, 24, 40, 60, 50, 40,
          50)


def _excel_round(x: float, digits: int) -> float:
    """Excel ROUND(x, n) — .5 sıfırdan uzağa (Decimal(repr) ile, bkz. round_half_up)."""
    q = Decimal(1).scaleb(-digits)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def _clean(v):
    """XML'e yazılamayan kontrol karakterlerini atar (metin değilse aynen)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _put_text(ws, row: int, col: int, value):
    """Serbest metni hücreye DAİMA metin olarak yazar ("=…" formüle dönmez)."""
    value = _clean(value)
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and c.data_type == "f":
        c.data_type = "s"
        c.quotePrefix = True
    return c


def _append_as_text(ws, values) -> int:
    """`ws.append` + formül kalkanı + kontrol karakteri temizliği; eklenen
    satır numarasını döner."""
    values = [_clean(v) for v in values]
    ws.append(values)
    r = ws.max_row
    for col, v in enumerate(values, start=1):
        if isinstance(v, str):
            c = ws.cell(row=r, column=col)
            if c.data_type == "f":
                c.data_type = "s"
                c.quotePrefix = True
    return r


def _style_header(ws, row: int = 1, wrap: bool = True):
    """Başlık satırı: lacivert dolgu, beyaz kalın yazı, kalın alt çizgi."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fill = PatternFill("solid", fgColor=HEADER_FILL)
    border = Border(bottom=Side(style="medium", color=HEADER_FILL))
    for c in ws[row]:
        c.font = Font(bold=True, color=HEADER_TEXT)
        c.fill = fill
        c.border = border
        c.alignment = Alignment(wrap_text=wrap, vertical="center")


def _polish(ws, *, zebra: bool = True, filt: bool = False, header_row: int = 1) -> None:
    """Veri satırlarına kıl alt çizgi + zebra; `filt` ile otomatik süzgeç
    (tek bloklu sayfalar).  Veri satırı = dolu ve kalın yazısız satır; başlık,
    ara toplam ve boş satırlar atlanır, zebra her blokta yeniden başlar."""
    from openpyxl.styles import Border, PatternFill, Side
    from openpyxl.utils import get_column_letter

    last = ws.max_column
    if ws.max_row <= header_row or not last:
        return
    hair = Border(bottom=Side(style="hair", color=LINE_COLOR))
    fill = PatternFill("solid", fgColor=ZEBRA_FILL)
    n = 0
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, max_col=last):
        filled = [c for c in row if c.value not in (None, "")]
        # İlk üç dolu hücreden biri kalınsa başlık / ara toplam satırı (ambalaj
        # uyarılı Tutar hücresi kalın olsa da satır veri satırıdır)
        if not filled or any(c.font is not None and c.font.bold for c in filled[:3]):
            n = 0
            continue
        n += 1
        for c in row:
            c.border = hair
            if zebra and n % 2 == 0:
                c.fill = fill
    if filt:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last)}{ws.max_row}"


def _print_setup(ws, *, landscape: bool, title: str) -> None:
    """A4 yazdırma ayarı — `ingredients_report._setup_a4` (tembel içe aktarma:
    o modül `database`'i yükler); üst bilgi Satın Alma Planı başlığı, 1.
    satır her baskı sayfasında."""
    from core.ingredients_report import _setup_a4

    _setup_a4(ws, landscape=landscape, header_text=f"Minerva 108 — Satın Alma Planı · {title}", title_rows="1:1")


# ─── Yönetici Özeti (ilk sayfa) ─────────────────────────────────────────────

def _sheet_exec(wb, report: dict) -> None:
    """İlk sayfa: `summary_data` (PDF Yönetici Özeti ile aynı sayılar).
    Bloklar arası boş satır; her bloğun başlığı `_style_header` satırı.
    Serbest metin (ad, iletişim, uyarı) formül kalkanından geçer."""
    from openpyxl.styles import Alignment, Font

    from database import to_tr

    sd = summary_data(report)
    sym = sd["symbol"]
    ws = wb.active
    ws.title = SHEET_EXEC
    bold = Font(bold=True)
    r = _append_as_text(ws, [f"{sd['title']} — Yönetici Özeti"])
    ws.cell(r, 1).font = Font(bold=True, size=14, color=TITLE_COLOR)
    gen = sd.get("generated_at")
    try:
        gen_tr = to_tr(datetime.fromisoformat(gen)).strftime("%d.%m.%Y %H:%M") if gen else ""
    except ValueError:
        gen_tr = gen or ""
    sub = f"Stoklar {sd['stock_as_of_tr']} · fiyatlar {sd['currency']} ({sym})" + (f" · oluşturma {gen_tr}" if gen_tr else "")
    if sd["brands"]:
        sub += " · markalar: " + ", ".join(sd["brands"])
    r = _append_as_text(ws, [sub])
    ws.cell(r, 1).font = Font(color=MUTED_COLOR)
    if sd["scope_text"]:
        r = _append_as_text(ws, [sd["scope_text"]])
        ws.cell(r, 1).font = Font(color=MUTED_COLOR)

    def head(cells):
        ws.append([])
        rr = _append_as_text(ws, cells)
        _style_header(ws, rr, wrap=False)
        return rr

    # Göstergeler
    head(["Göstergeler", "Değer", "Gösterim", "", "", "", "", "Not"])
    for k in sd["kpis"]:
        rr = _append_as_text(ws, [k["label"], k["value"], k["text"], "", "", "", "", k.get("note") or ""])
        ws.cell(rr, 2).number_format = MONEY_FMT if k["key"] in ("plan_total", "buy_total", "stock_total") else "#,##0"

    # Maliyet — tür bazında
    if sd["categories"]:
        # Sütunlar PDF'teki tabloyla aynı: Stoktan yalnız net modda, Yoldan
        # yalnız açık sipariş seçeneğiyle (brüt modda stok değeri dili yok)
        cols = [("buy", f"Alınacak ({sym})")] + ([("stock", f"Stoktan ({sym})")] if sd["show_stock"] else []) \
            + ([("open", f"Yoldan ({sym})")] if sd["show_open"] else []) \
            + [("total", f"Toplam ({sym})"), ("priced", "Fiyatlı"), ("unpriced", "Fiyatsız")]
        head(["Maliyet — tür bazında"] + [h for _, h in cols])
        rows = list(sd["categories"]) + [dict(sd["categories_total"], label="Toplam")]
        for c in rows:
            rr = _append_as_text(ws, [c["label"]] + [c.get(k) or 0 for k, _ in cols])
            for i, (k, _) in enumerate(cols, start=2):
                if k not in ("priced", "unpriced"):
                    ws.cell(rr, i).number_format = MONEY_FMT
        for col in range(1, len(cols) + 2):
            ws.cell(rr, col).font = bold

    # Alınacaklar — tedarikçi bazında
    if sd["top_suppliers"]:
        head([f"Alınacaklar — tedarikçi bazında ({sd['suppliers_count']} firma)", "Kalem", f"Tutar ({sym})",
              "", "", "", "", "İletişim"])
        for b in sd["top_suppliers"]:
            contact = "" if b["other"] else (b.get("contact_line") or "İletişim bilgisi sistemde yok")
            rr = _append_as_text(ws, [b["name"], b["count"], b["total"], "", "", "", "", contact])
            ws.cell(rr, 3).number_format = MONEY_FMT
            if b["other"]:
                ws.cell(rr, 1).font = Font(color=MUTED_COLOR)

    # Hammadde ihtiyacı
    if sd["raw_need"] or any(sd["raw_to_buy"].values()):
        head(["Hammadde ihtiyacı", "Birim", "Toplam gereken", "Alınacak"])
        for u in ("kg", "l", "adet"):
            if sd["raw_to_buy"].get(u) or sd["raw_need_totals"].get(u):
                ut = UNIT_TEXT.get(u, u)          # "litre" — PDF ve Alım listesiyle aynı söz
                rr = _append_as_text(ws, [f"Alınacak hammadde — {ut} (alım birimine göre toplam)", ut,
                                          sd["raw_need_totals"].get(u) or 0, sd["raw_to_buy"].get(u) or 0])
                ws.cell(rr, 1).font = bold
                for col in (3, 4):
                    ws.cell(rr, col).number_format = _exec_qty_fmt(u)
                    ws.cell(rr, col).font = bold
        for blk in sd["raw_need"]:
            u = blk["unit"]
            for row in blk["rows"]:
                rr = _append_as_text(ws, [row["name"], UNIT_TEXT.get(u, u), row["need"],
                                          row["buy"] if row["buy"] else None])
                for col in (3, 4):
                    # Malzeme satırı Alım listesinin biçimiyle (1,25 → "1,25", orada
                    # da öyle); "Diğer" toplamı 1 ondalık
                    ws.cell(rr, col).number_format = _exec_qty_fmt(u) if row["other"] else _qty_fmt(u)
                if row["other"]:
                    ws.cell(rr, 1).font = Font(color=MUTED_COLOR)

    # Fiyatı eksik kalemler
    if sd["unpriced_groups"]:
        head([f"Fiyatı eksik kalemler ({sd['unpriced_count']})", "Kalem", "Alınacak", "", "", "", "", "Örnekler"])
        for g in sd["unpriced_groups"]:
            ex = ", ".join(g["names"]) + (f" (+{g['more']} kalem daha)" if g["more"] else "")
            _append_as_text(ws, [g["label"], g["count"], g["to_buy"], "", "", "", "", ex])
        rr = _append_as_text(ws, [sd["howto"]])
        ws.cell(rr, 1).font = Font(color=MUTED_COLOR)

    # Uyarılar ve bekletilenler
    if sd["warnings"] or sd["held"]:
        head(["Uyarılar ve bekletilenler"])
        for w in sd["warnings"]:
            _append_as_text(ws, [w])
        for h in sd["held"]:
            t = f"{h['name']} ({h['buy_text']}) " + (
                f"bekletiliyor: {(h.get('reason') or '').rstrip('. ')}" if h.get("reason") else "teyit bekliyor")
            if h.get("amount") is not None and h.get("supplier"):
                t += f" — alınırsa yaklaşık {money(h['amount'], sd['currency'])} ({h['supplier']})"
            _append_as_text(ws, [t])
        if sd["held"] and sd["held_what_if"]:
            _append_as_text(ws, [f"Bekletilenlerin tamamı alınırsa yaklaşık "
                                 f"{money(sd['held_what_if'], sd['currency'])} eklenir."])

    for col, w in zip("ABCDEFGH", (46, 15, 15, 15, 15, 10, 10, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8):
        for c in row[1:7]:
            if isinstance(c.value, (int, float)):
                c.alignment = Alignment(horizontal="right")
        row[7].alignment = Alignment(wrap_text=True, vertical="top")


def _qty_fmt(unit: Optional[str]) -> str:
    return "#,##0" if unit == "adet" else "#,##0.0##"


# Yönetici Özeti TOPLAM miktarı ("Alınacak hammadde", "Diğer"): PDF `_amt` ile
# aynı yazım — kg/litre 1 ondalık ("442.720,6"), 1'in altı 3 ondalığa kadar
# (0,026 sıfır görünmesin).  Malzeme satırı Alım listesinin `_qty_fmt`'i (1,25)
EXEC_QTY_FMT = "[<1]0.0##;#,##0.0"


def _exec_qty_fmt(unit: Optional[str]) -> str:
    return "#,##0" if unit == "adet" else EXEC_QTY_FMT


def _rel_texts(m: dict):
    rel = m.get("relations") or {}
    card = " / ".join(x["name"] for x in rel.get("card") or [])
    last = rel.get("last")
    last_t = (f"{last['name']} ({last['date']})" if last.get("date") else last["name"]) if last else ""
    samples = "; ".join(f"{s['name']} ({s['date']})" if s.get("date") else s["name"]
                        for s in rel.get("samples") or [])
    return card, last_t, samples


def _offer_name(o: Optional[dict]) -> str:
    """T1–T3 sütunu: firma adı + durum etiketi ("tercih", "eşdeğer kart «X»",
    " — bitirilecek, alma") — PDF/önizlemedeki etiketlerin aynısı, fiyat
    karşılaştırması ("daha pahalı") hariç.  Etiketsiz teklifte yalnız ad."""
    if not o:
        return ""
    tags, blocked = offer_tags(o)
    t = o["name"] + (" (" + ", ".join(tags) + ")" if tags else "")
    return t + (f" — {blocked}" if blocked else "")


def _products_text(m: dict) -> str:
    return "; ".join(f"{p['no']}. {p['product']}" for p in m.get("per_product") or [])


def _note_text(m: dict, cur: str) -> str:
    """"Not" sütunu — PDF tedarikçi hücresindeki açıklamalarla aynı bilgi:
    kullanılamayan teklifin nedeni (birim uyuşmazlığı / kur yok), ambalaj
    katına yuvarlanmış miktar + tutar (seçenek açıkken), uyarılar."""
    parts = [unpriced_offer_text(o) + "." for o in m.get("offers") or []
             if o.get("price") is None and o.get("note")]
    pkg = pkg_round_text(m, cur)          # önizleme / PDF ile aynı metin
    if pkg:
        parts.append(pkg + ".")
    parts += [x["text"] for x in m.get("cautions") or []]
    return " ".join(parts)


# ─── Alım listesi ───────────────────────────────────────────────────────────

def _sheet_list(wb, report: dict, secs: List[dict], cache: Dict[str, object]) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    meta = report.get("meta") or {}
    opts = meta.get("options") or {}
    pr = report.get("pricing") or {}
    cur = pr.get("currency") or opts.get("currency") or "USD"
    sym = CURRENCY_SYMBOL.get(cur, cur)
    gross = opts.get("stock_mode") == "gross"
    open_on = bool(opts.get("subtract_open_orders"))
    bold = Font(bold=True)

    ws = wb.create_sheet(SHEET_LIST)
    ws.append(["Bölüm", "Fiyat durumu", "Malzeme", "Gereken",
               "Elimizde (bilgi)" if gross else "Elimizde", "Yolda (bilgi)" if gross else "Yolda",
               "Alınacak", "Birim", "Seçilen tedarikçi", f"Birim fiyat ({sym})", f"Tutar ({sym})",
               "Fiyat kaynağı", "T1", f"Fiyat 1 ({sym})", "Ambalaj 1", "T2", f"Fiyat 2 ({sym})",
               "T3", f"Fiyat 3 ({sym})", "Stok kartında yazan", "Son alım", "Numune",
               "Kullanıldığı ürünler", "Not", "Aynı malzeme (diğer kartlar)", "Sipariş geçmişi",
               "Grup stoğu (kart dökümü)"])
    _style_header(ws)

    subtotal_cells: List[str] = []
    for sec in secs:
        if sec["key"] not in MATERIAL_SECTIONS:
            continue
        first = None
        amounts: List[int] = []
        for m in sec["rows"]:
            d = m.get("display") or {}
            unit = d.get("buy_unit") or "adet"
            offers = [o for o in m.get("offers") or []][:3]
            o = offers + [None] * (3 - len(offers))
            card, last_t, samples = _rel_texts(m)
            D = float(d.get("need_num") or 0.0)
            E = float(d.get("stock_num") or 0.0)
            F = float(d.get("open_num") or 0.0) if open_on else None
            buy = float(d.get("buy_num") or 0.0)
            price = m.get("price")
            vals = [KIND_TEXT.get(m["kind"], m["kind"]), GROUP_TEXT.get(m.get("group") or "none"),
                    m["name"], D, E, F, None, unit, m.get("supplier") or "", price, None,
                    m.get("price_source") or "",
                    _offer_name(o[0]), o[0]["price"] if o[0] else None,
                    o[0]["package"] if o[0] else None,
                    _offer_name(o[1]), o[1]["price"] if o[1] else None,
                    _offer_name(o[2]), o[2]["price"] if o[2] else None,
                    card, last_t, samples, _products_text(m), _note_text(m, cur),
                    "; ".join(alt_texts(m)), "; ".join(order_texts(m)), stock_note(m) or ""]
            r = _append_as_text(ws, vals)
            first = first or r
            # G — Alınacak
            if gross:
                g_formula, g_val = f"=D{r}", D
            else:
                g_formula, g_val = f"=MAX(0,ROUND(D{r}-E{r}-F{r},3))", max(0.0, _excel_round(D - E - (F or 0.0), 3))
            if abs(g_val - buy) <= 1e-9:
                ws.cell(r, 7).value = g_formula
                cache[f"G{r}"] = g_val
            else:
                ws.cell(r, 7).value = buy     # düz yuvarlama: formül raporla ayrışırdı
            # K — Tutar
            ws.cell(r, 11).value = f'=IF(J{r}="","",ROUND(G{r}*J{r},0))'
            amount = round_half_up(buy * price) if price is not None else None
            cache[f"K{r}"] = amount if amount is not None else ""
            if amount is not None:
                amounts.append(amount)
            fmt = _qty_fmt(unit)
            for col in (4, 5, 6, 7):
                ws.cell(r, col).number_format = fmt
            for col in (10, 14, 17, 19):
                v = ws.cell(r, col).value
                if isinstance(v, (int, float)):
                    ws.cell(r, col).number_format = "0.0000" if v < 1 else "#,##0.00"
            ws.cell(r, 11).number_format = "#,##0"
            warn = amount_warn(m)
            if warn:
                # Ambalaj uyarısı: tutar turuncu, uyarı metni hücre yorumu (X "Not"ta da yazar)
                ws.cell(r, 11).font = Font(color=WARN_COLOR, bold=True)
                cm = Comment(_clean(warn), "Minerva 108")
                cm.width, cm.height = 320, 110
                ws.cell(r, 11).comment = cm
        if first and any(m.get("price") is not None for m in sec["rows"]):
            last = ws.max_row
            r = _append_as_text(ws, ["", "", f"{sec['no']}. {sec['title']} — toplam"])
            ws.cell(r, 11).value = f"=SUM(K{first}:K{last})"
            cache[f"K{r}"] = sum(amounts)
            ws.cell(r, 3).font = bold
            ws.cell(r, 11).font = bold
            ws.cell(r, 11).number_format = "#,##0"
            subtotal_cells.append(f"K{r}")

    ws.append([])
    r = _append_as_text(ws, ["", "", "GENEL TOPLAM (fiyatı belli olanlar)"])
    if subtotal_cells:
        ws.cell(r, 11).value = "=" + "+".join(subtotal_cells)
        cache[f"K{r}"] = sum(cache[c] for c in subtotal_cells)
    else:
        ws.cell(r, 11).value = 0
    ws.cell(r, 3).font = bold
    ws.cell(r, 11).font = bold
    ws.cell(r, 11).number_format = "#,##0"

    # Yeni etiket / koli / palet
    new_sec = next((s for s in secs if s["key"] == "new_items"), None)
    if new_sec:
        ws.append([])
        r = _append_as_text(ws, ["", GROUP_TEXT["none"],
                                 f"{new_sec['no']}. {new_sec['title']} (sistemde fiyat ve stok kaydı yok)"])
        ws.cell(r, 3).font = bold
        for row in new_sec["rows"]:
            if row.get("type") == "label":
                vals = ["Etiket", GROUP_TEXT["none"], row["name"], None, None, None, row["total"], "adet"]
                vals += [""] * 15 + [row.get("face_text") or ""]
            else:
                # Sayı yoksa serbest miktar metni ("22 palet için") — birim o zaman metnin içinde
                num = row.get("qty") is not None
                vals = [MANUAL_SECTION_TEXT.get(row.get("section"), "Diğer"), GROUP_TEXT["none"], row["name"],
                        None, None, None, row["qty"] if num else (row.get("qty_text") or ""),
                        (row.get("unit") or "") if num else ""]
                vals += [""] * 15 + ["Sistemde stok kaydı yok" + (f" · {row['note']}" if row.get("note") else "")]
            r = _append_as_text(ws, vals)
            q = ws.cell(r, 7).value
            # Serbest satır miktarı ondalıklı olabilir (2,5 rulo) — PDF gibi göster
            ws.cell(r, 7).number_format = ("#,##0.0#" if isinstance(q, float) and not q.is_integer()
                                           else "#,##0")

    for col, w in zip(COLS, WIDTHS):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"


# ─── Stoktan karşılananlar ──────────────────────────────────────────────────

def _covered_rows(report: dict, secs: List[dict]) -> List[dict]:
    """Önce "Yeterli" bölümünün satırları (sections() sırası), sonra Alım
    listesi sırasıyla stoktan / yoldan kısmen karşılanan alınacaklar."""
    def used(m):
        return (m.get("stock_qty_used") or 0) > 0 or (m.get("open_qty_used") or 0) > 0

    suff = next((s["rows"] for s in secs if s["key"] == "sufficient"), [])
    part = [m for s in secs if s["key"] in MATERIAL_SECTIONS for m in s["rows"] if used(m)]
    return [m for m in suff if used(m)] + part


def _sheet_stock(wb, report: dict, secs: List[dict], cache: Dict[str, object]) -> None:
    """"Stoktan karşılananlar" — planın stoktan (ve açık siparişten) karşıladığı
    malzemenin değeri.  Yalnız net modda ve satır varsa kurulur."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    meta = report.get("meta") or {}
    opts = meta.get("options") or {}
    pr = report.get("pricing") or {}
    if not (pr.get("cost_summary") or {}).get("applies"):
        return
    rows = _covered_rows(report, secs)
    if not rows:
        return
    cur = pr.get("currency") or opts.get("currency") or "USD"
    sym = CURRENCY_SYMBOL.get(cur, cur)
    open_on = bool(opts.get("subtract_open_orders"))
    bold = Font(bold=True)
    ws = wb.create_sheet(SHEET_STOCK, index=wb.sheetnames.index(SHEET_LIST) + 1)
    head = ["Durum", "Tür", "Malzeme", "Gereken", "Elimizde", "Yolda", "Stoktan kullanılan", "Birim",
            "Seçilen tedarikçi", f"Birim fiyat ({sym})", f"Değer ({sym})"]
    if open_on:
        head += ["Yoldan kullanılan", f"Yoldan değer ({sym})"]
    head += ["Fiyat kaynağı", "Not"]
    ws.append(head)
    _style_header(ws)
    first = None
    k_vals: List[int] = []
    m_vals: List[int] = []
    for m in rows:
        d = m.get("display") or {}
        unit = d.get("buy_unit") or "adet"
        D = float(d.get("need_num") or 0.0)
        E = float(d.get("stock_num") or 0.0)
        F = float(d.get("open_num") or 0.0) if open_on else None
        price = m.get("price")
        if m["status"] == "yeterli":
            status = "Yeterli"
        else:
            parts = (["stoktan"] if (m.get("stock_qty_used") or 0) > 0 else []) + \
                    (["yoldan"] if (m.get("open_qty_used") or 0) > 0 else [])
            status = "Kısmen " + " + ".join(parts)
        note = []
        if price is None:
            note.append("Fiyat yok; değer toplama girmedi.")
        sn = stock_note(m)
        if sn:
            note.append(f"Grup stoğu: {sn}.")
        vals = [status, KIND_TEXT.get(m["kind"], m["kind"]), m["name"], D, E, F, None, unit,
                m.get("supplier") or "", price, None]
        if open_on:
            vals += [None, None]
        vals += [m.get("price_source") or "", " ".join(note)]
        r = _append_as_text(ws, vals)
        first = first or r
        # G = satırın `stock_qty_used`'ı; MIN(D,E)'ye eşitse formül (yeterli
        # satır, güvenli yuvarlama), değilse sayı — düz yuvarlamada adet
        # satırının alınacağı tama yuvarlanır, stoktan = gereken − alınacak
        # olur (Alım listesi G'sinin formül/sayı korumasıyla aynı kalıp)
        G = float(m.get("stock_qty_used") or 0.0)
        if abs(G - min(D, E)) <= 1e-9:
            ws.cell(r, 7).value = f"=MIN(D{r},E{r})"
            cache[f"G{r}"] = G
        else:
            ws.cell(r, 7).value = G
        ws.cell(r, 11).value = f'=IF(J{r}="","",ROUND(G{r}*J{r},0))'
        kv = round_half_up(G * price) if price is not None else None
        cache[f"K{r}"] = kv if kv is not None else ""
        if kv is not None:
            k_vals.append(kv)
        fmt = _qty_fmt(unit)
        for col in (4, 5, 6, 7):
            ws.cell(r, col).number_format = fmt
        if isinstance(price, (int, float)):
            ws.cell(r, 10).number_format = "0.0000" if price < 1 else "#,##0.00"
        ws.cell(r, 11).number_format = "#,##0"
        if open_on:
            L = float(m.get("open_qty_used") or 0.0)
            if abs(L - min(F, round6(D - G))) <= 1e-9:
                ws.cell(r, 12).value = f"=MIN(F{r},ROUND(D{r}-G{r},6))"
                cache[f"L{r}"] = L
            else:
                ws.cell(r, 12).value = L
            ws.cell(r, 13).value = f'=IF(J{r}="","",ROUND(L{r}*J{r},0))'
            mv = round_half_up(L * price) if price is not None else None
            cache[f"M{r}"] = mv if mv is not None else ""
            if mv is not None:
                m_vals.append(mv)
            ws.cell(r, 12).number_format = fmt
            ws.cell(r, 13).number_format = "#,##0"
    last = ws.max_row
    r = _append_as_text(ws, ["", "", "TOPLAM (fiyatı belli olanlar)"])
    ws.cell(r, 11).value = f"=SUM(K{first}:K{last})"
    cache[f"K{r}"] = sum(k_vals)
    cells = [11]
    if open_on:
        ws.cell(r, 13).value = f"=SUM(M{first}:M{last})"
        cache[f"M{r}"] = sum(m_vals)
        cells.append(13)
    ws.cell(r, 3).font = bold
    for col in cells:
        ws.cell(r, col).font = bold
        ws.cell(r, col).number_format = "#,##0"
    ws.append([])
    for t in cost_summary_caveats(report):          # fiyatsız kalem sayıları + değerin dayanağı
        _append_as_text(ws, ["", "", t])
    widths = [16, 11, 34, 12, 12, 10, 14, 7, 22, 12, 12] + ([14, 14] if open_on else []) + [24, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"


# ─── Diğer sayfalar ─────────────────────────────────────────────────────────

def _sheet_suppliers(wb, report: dict, cur: str) -> None:
    ws = wb.create_sheet(SHEET_SUP)
    sym = CURRENCY_SYMBOL.get(cur, cur)
    ws.append(["Firma", "Kalem (seçildiği)" if (report.get("pricing") or {}).get("status_used")
               else "Kalem (en ucuz olduğu)", f"Tutar ({sym})", "Yetkili", "Telefon", "E-posta",
               "Adres / şehir", "Bilgi kaynağı", "Sistemdeki kart adı", "Fiyatı olmayan kalemlerde aday",
               "İlişki türleri"])
    _style_header(ws)
    sup = report.get("suppliers") or {}
    chosen = {b["key"]: b for b in sup.get("chosen") or []}
    cands = {b["key"]: b for b in sup.get("candidates") or []}
    keys = list(chosen) + [k for k in cands if k not in chosen]
    for k in keys:
        b = chosen.get(k) or cands[k]
        c = b.get("contact") or {}
        if c and contact_line(c):
            src = "IMS tedarikçi kartı"
        elif c:
            src = "IMS kartında bilgi yok — girilecek"
        else:
            src = "Sistemde kart yok — girilecek"
        cand_items = cands.get(k, {}).get("items") or []
        r = _append_as_text(ws, [
            b["name"], chosen[k]["count"] if k in chosen else 0,
            chosen[k]["total"] if k in chosen else None,
            c.get("person") or "", c.get("phone") or "", c.get("email") or "", c.get("address") or "",
            src, " / ".join(c.get("cards") or []),
            ", ".join(f"{it['name']} ({', '.join(it.get('reasons') or [])})" if it.get("reasons") else it["name"]
                      for it in cand_items),
            ", ".join(REL_TYPE_TEXT.get(t, t) for t in (b.get("types") or []))])
        ws.cell(r, 3).number_format = "#,##0"
    unrelated = sup.get("unrelated") or []
    if unrelated:
        _append_as_text(ws, ["Sistemde hiçbir firmayla ilişkisi yok", 0, None, "", "", "", "",
                             "Yeni tedarikçi bulunmalı", "", ", ".join(u["name"] for u in unrelated), ""])
    for col, w in zip("ABCDEFGHIJK", (26, 12, 12, 20, 16, 28, 28, 26, 34, 90, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"


def _sheet_products(wb, report: dict) -> None:
    ws = wb.create_sheet(SHEET_PROD)
    ws.append(["No", "Ürün", "Marka", "Reçete", "Ölçek", "Hedef adet", "Bitmiş stok", "Üretilecek",
               "Reçetede ambalaj", "Tahmini", "Elimizdekiyle üretilebilir", "Kısıtlayan malzeme",
               "Uyarılar", "Not"])
    _style_header(ws)
    for p in report.get("products") or []:
        cap = p.get("capacity") or {}
        r = _append_as_text(ws, [
            p["no"], p["name"], p.get("brand") or "", p.get("recipe_name") or "", p.get("scale"),
            p["qty"], p.get("finished_stock"), p.get("produce_qty"),
            "Var" if p.get("has_packaging") else "Yok", "Evet" if p.get("estimated") else "",
            cap.get("producible"),
            "; ".join(f"{x['name']} ({_amount_near(x['stock'], x['unit'])})" for x in cap.get("limiting") or []),
            " ".join(x["text"] for x in p.get("cautions") or []), p.get("note") or ""])
        for col in (6, 7, 8, 11):
            ws.cell(r, col).number_format = "#,##0"
    for col, w in zip("ABCDEFGHIJKLMN", (6, 40, 14, 40, 8, 11, 11, 11, 10, 9, 14, 40, 60, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"


def _sheet_cost(wb, report: dict, cache: Dict[str, object]) -> None:
    """"Ürün maliyeti" (16.10.2026; "Ürünler"den sonra, fiyatlı raporda).

    Üstte ürün özeti (No · Ürün · Adet · 1 adet maliyeti · Toplam · fiyatı
    eksik kalem sayısı + adları) ve TOPLAM; altında ürün × malzeme dökümü
    (`product_cost_lines` — aynı sayılar):
        F (1 adet için, fiyat biriminde) · G birim fiyat · H = IF(G="","",F*G)
        I adet · J = F*I (toplam tüketim) · K = IF(G="","",J*G)
    + ürün başına ara toplam (H = SUM, K = ROUND(SUM, 0)); özetin D/E'si ara
    toplama bağlanır.  Sonra tür bazında maliyet (Σ = planın toplam
    malzeme maliyeti), fiyatı eksik kalemler ve notlar
    (`cost_overview_view`, önizleme / PDF ile aynı metin).  Bütün
    formüllerin önbellek değeri yazılır."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    pr = report.get("pricing") or {}
    co = pr.get("cost_overview")
    v = cost_overview_view(report)
    if not co or not v:
        return
    cur = pr.get("currency") or "USD"
    sym = CURRENCY_SYMBOL.get(cur, cur)
    bold = Font(bold=True)
    ws = wb.create_sheet(SHEET_COST, index=wb.sheetnames.index(SHEET_PROD) + 1)
    prods = co["products"]
    lines = product_cost_lines(report)
    by_no: Dict[int, List[dict]] = {}
    for ln in lines:
        by_no.setdefault(ln["no"], []).append(ln)

    # Satır yerleşimi önceden: özet (başlık + ürünler + TOPLAM), boş satır,
    # döküm başlığı, sonra ürün başına satırlar + ara toplam
    top_first = 2
    top_total = top_first + len(prods)
    det_head = top_total + 2
    sub_row: Dict[int, int] = {}
    r = det_head
    for p in prods:
        ls = by_no.get(p["no"]) or []
        if ls:
            r += len(ls) + 1
            sub_row[p["no"]] = r

    ws.append(["No", "Ürün", "Adet", f"1 adet maliyeti ({sym})", f"Toplam ({sym})", "Fiyatı eksik kalem",
               "Eksik kalemler"])
    _style_header(ws)
    units = 0.0
    for p in prods:
        row = _append_as_text(ws, [p["no"], p["name"], p["qty"], None, None, p["missing"],
                                   ", ".join(p["missing_names"]) or ("Reçetesi yok" if p["recipeless"] else "")])
        units += float(p["qty"] or 0)
        sr = sub_row.get(p["no"])
        if sr and p["priced"]:
            ws.cell(row, 4).value = f"=H{sr}"
            cache[f"D{row}"] = p["unit_cost"]
            ws.cell(row, 5).value = f"=K{sr}"
            cache[f"E{row}"] = p["total"]
        ws.cell(row, 3).number_format = "#,##0"
        ws.cell(row, 4).number_format = "#,##0.00"
        ws.cell(row, 5).number_format = "#,##0"
    row = _append_as_text(ws, ["", "TOPLAM", None, None, None])
    assert row == top_total
    ws.cell(row, 3).value = f"=SUM(C{top_first}:C{top_total - 1})" if prods else 0
    ws.cell(row, 5).value = f"=SUM(E{top_first}:E{top_total - 1})" if prods else 0
    if prods:
        cache[f"C{row}"] = units
        cache[f"E{row}"] = co["products_total"]
    for col in (2, 3, 5):
        ws.cell(row, col).font = bold
    ws.cell(row, 3).number_format = ws.cell(row, 5).number_format = "#,##0"

    ws.append([])
    _append_as_text(ws, ["No", "Ürün", "Malzeme", "Tür", "Birim", "1 adet için", f"Birim fiyat ({sym})",
                         f"1 adet maliyeti ({sym})", "Adet", "Toplam tüketim", f"Toplam ({sym})", "Tedarikçi",
                         "Not"])
    _style_header(ws, det_head)
    for p in prods:
        ls = by_no.get(p["no"]) or []
        if not ls:
            continue
        first = ws.max_row + 1
        h_sum, k_sum = 0.0, 0.0
        for ln in ls:
            unit = UNIT_TEXT.get(ln["unit"], ln["unit"])
            row = _append_as_text(ws, [p["no"], p["name"], ln["name"], KIND_TITLE.get(ln["kind"], ln["kind"]), unit,
                                       ln["per_unit"], ln["price"], None, ln["qty"], None, None,
                                       ln.get("supplier") or "",
                                       ("Fiyat yok — toplama girmedi" if ln["price"] is None else "")
                                       + (" · ek ambalaj" if ln["extra"] else "")])
            ws.cell(row, 8).value = f'=IF(G{row}="","",F{row}*G{row})'
            ws.cell(row, 10).value = f"=F{row}*I{row}"
            ws.cell(row, 11).value = f'=IF(G{row}="","",J{row}*G{row})'
            cache[f"H{row}"] = ln["unit_cost"] if ln["unit_cost"] is not None else ""
            cache[f"J{row}"] = ln["total_qty"]
            cache[f"K{row}"] = ln["total"] if ln["total"] is not None else ""
            if ln["unit_cost"] is not None:
                h_sum += ln["unit_cost"]
                k_sum += ln["total"]
            qfmt = "#,##0" if ln["unit"] == "adet" else "#,##0.000###"
            ws.cell(row, 6).number_format = qfmt
            ws.cell(row, 10).number_format = qfmt
            if isinstance(ln["price"], (int, float)):
                ws.cell(row, 7).number_format = "0.0000" if ln["price"] < 1 else "#,##0.00"
            ws.cell(row, 8).number_format = "#,##0.0000"
            ws.cell(row, 9).number_format = "#,##0"
            ws.cell(row, 11).number_format = "#,##0.00"
        last = ws.max_row
        row = _append_as_text(ws, ["", "", f"{p['no']}. {p['name']} — ara toplam"])
        assert row == sub_row[p["no"]]
        ws.cell(row, 8).value = f"=SUM(H{first}:H{last})"
        ws.cell(row, 11).value = f"=ROUND(SUM(K{first}:K{last}),0)"
        cache[f"H{row}"] = h_sum
        cache[f"K{row}"] = round_half_up(k_sum) if p["priced"] else 0
        for col in (3, 8, 11):
            ws.cell(row, col).font = bold
        ws.cell(row, 8).number_format = "#,##0.00"
        ws.cell(row, 11).number_format = "#,##0"

    # Tür bazında maliyet (Σ = planın toplam malzeme maliyeti) — değer, formül yok
    ws.append([])
    row = _append_as_text(ws, ["", "Tür bazında maliyet", "", f"Alınacak ({sym})", f"Stoktan ({sym})",
                               f"Yoldan ({sym})", f"Toplam ({sym})", "Fiyatlı", "Fiyatsız"])
    _style_header(ws, row)
    tot_label = "TOPLAM (planın toplam malzeme maliyeti)" if co["net"] else "TOPLAM (alınacakların tutarı)"
    for c in co["categories"] + [dict(co["total"], label=tot_label)]:
        row = _append_as_text(ws, ["", c["label"], "", c["buy"], c["stock"], c["open"], c["total"],
                                   c["priced"], c["unpriced"]])
        for col in (4, 5, 6, 7):
            ws.cell(row, col).number_format = "#,##0"
    ws.cell(row, 2).font = ws.cell(row, 7).font = bold
    for t in v["caveats"]:
        _append_as_text(ws, ["", t])
    # Gereken miktarlar
    ws.append([])
    row = _append_as_text(ws, ["", "Gereken miktarlar"])
    ws.cell(row, 2).font = bold
    for q in v["quantities"]:
        _append_as_text(ws, ["", q["label"], q["text"]])
    # Fiyatı eksik kalemler (tam liste)
    if v["unpriced"]:
        ws.append([])
        row = _append_as_text(ws, ["", f"Fiyatı eksik kalemler ({v['unpriced_count']})", "Tür", "İhtiyaç",
                                   "Not"])
        _style_header(ws, row)
        for u in v["unpriced"]:
            _append_as_text(ws, ["", u["name"], u["kind"], u["need"], u["note"]])
    ws.append([])
    row = _append_as_text(ws, ["", "Notlar"])
    ws.cell(row, 2).font = bold
    for t in v["notes"]:
        _append_as_text(ws, ["", t])
    for i, w in enumerate((8, 40, 34, 14, 14, 14, 14, 16, 10, 16, 14, 24, 30), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Yalnız A:B (No, Ürün) iki blokta da aynı anlamda — üst blok başlığı
    # dökümün üstünde donarsa sütun adları yanlış olurdu
    ws.freeze_panes = "C1"


def _sheet_detail(wb, report: dict, cache: Dict[str, object]) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(SHEET_DETAIL)
    prods = report.get("products") or []
    col_of = {p["no"]: 4 + i for i, p in enumerate(prods)}
    tot_col = 4 + len(prods)
    _append_as_text(ws, ["Malzeme", "Tür", "Birim"] + [f"{p['no']}. {p['name']}" for p in prods]
                    + ["Toplam (üretim)", "Gereken (alım firesiyle)", "Durum"])
    _style_header(ws)
    ws.row_dimensions[1].height = 60
    for c in ws[1][3:]:
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    rows = [(m, "Alınacak" if m["status"] == "to_buy" else "Yeterli") for m in report.get("materials") or []]
    rows += [(h, "Bekletiliyor") for h in report.get("held") or []]
    for m, status in rows:
        vals = [m["name"], KIND_TEXT.get(m["kind"], m["kind"]), m.get("unit") or ""] + [None] * len(prods)
        for pp in m.get("per_product") or []:
            if pp["no"] in col_of:
                vals[col_of[pp["no"]] - 1] = round(float(pp["need"]), 3)
        written = [v for v in vals[3:3 + len(prods)] if v is not None]     # sütun sırası
        vals += [None, round(float(m.get("need") or 0.0), 3), status]
        r = _append_as_text(ws, vals)
        if prods:
            a, b = get_column_letter(4), get_column_letter(tot_col - 1)
            ws.cell(r, tot_col).value = f"=SUM({a}{r}:{b}{r})"
            total = 0.0
            for v in written:          # Excel'in soldan sağa toplamıyla aynı sıra
                total += v
            cache[f"{get_column_letter(tot_col)}{r}"] = total
        fmt = "#,##0" if m.get("unit") == "adet" else "#,##0.###"
        for col in range(4, tot_col + 2):
            ws.cell(r, col).number_format = fmt
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["C"].width = 7
    for i in range(len(prods)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 14
    for col, w in ((tot_col, 14), (tot_col + 1, 14), (tot_col + 2, 12)):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "D2"


def _sheet_summary(wb, report: dict, secs: List[dict], cur: str) -> None:
    from openpyxl.styles import Alignment, Font

    from database import to_tr

    ws = wb.create_sheet(SHEET_SUMMARY)
    meta = report.get("meta") or {}
    pr = report.get("pricing") or {}
    tot = pr.get("totals") or {}
    cnt = pr.get("counts") or {}
    sym = CURRENCY_SYMBOL.get(cur, cur)
    bold = Font(bold=True)
    ws.append(["Alan", "Değer"])
    _style_header(ws)
    gen = meta.get("generated_at")
    try:
        gen_tr = to_tr(datetime.fromisoformat(gen)).strftime("%d.%m.%Y %H:%M") if gen else ""
    except ValueError:
        gen_tr = gen or ""
    c = meta.get("counts") or {}
    rows = [
        ("Başlık", meta.get("title") or ""),
        ("Kapsam", meta.get("scope_text") or ""),
        ("Stoklar", meta.get("stock_as_of_tr") or ""),
        ("Oluşturulma", gen_tr),
        ("Para birimi", f"{cur} ({sym})"),
        ("Ürün sayısı", c.get("products")),
        ("Toplam adet", c.get("units")),
        (f"Hammadde — fiyatı belli tutar ({sym})", tot.get("raw")),
        (f"Ambalaj — fiyatı belli tutar ({sym})", tot.get("packaging")),
        (f"Etiket — fiyatı belli tutar ({sym})", tot.get("label")),
        (f"GENEL TOPLAM — fiyatı belli ({sym})", tot.get("all")),
    ]
    # Net modda: stoktan (+ yoldan) karşılanan değer + planın toplam malzeme
    # maliyeti (alınacak + stoktan + yoldan) ve fiyatsız kalem sayıları.
    cs = pr.get("cost_summary") or {}
    plan_label = None
    for ln in cost_summary_lines(report):
        if ln["key"] == "buy":
            continue
        label = (f"PLANIN TOPLAM MALZEME MALİYETİ ({sym})" if ln["key"] == "plan"
                 else f"{ln['label']} — fiyatı belli ({sym})")
        plan_label = label if ln["key"] == "plan" else plan_label
        rows.append((label, ln["value"]))
    if cs.get("applies"):
        rows.append(("Fiyatlı / fiyatsız stoktan karşılanan",
                     f"{cs.get('stock_rows', 0) - cs.get('stock_unpriced', 0)} / {cs.get('stock_unpriced', 0)}"))
        if cs.get("open_rows"):
            rows.append(("Fiyatlı / fiyatsız yoldan karşılanan",
                         f"{cs['open_rows'] - cs.get('open_unpriced', 0)} / {cs.get('open_unpriced', 0)}"))
        rows += [("Maliyet notu", t) for t in cost_summary_caveats(report)]
    rows += [
        ("Fiyatlı / fiyatsız hammadde", f"{cnt.get('raw_priced', 0)} / {cnt.get('raw_unpriced', 0)}"),
        ("Fiyatlı / fiyatsız ambalaj", f"{cnt.get('packaging_priced', 0)} / {cnt.get('packaging_unpriced', 0)}"),
        ("Fiyatlı / fiyatsız etiket", f"{cnt.get('label_priced', 0)} / {cnt.get('label_unpriced', 0)}"),
        ("Yeni etiket (adet)", c.get("labels_new_total")),
        ("Yeterli kalem", c.get("sufficient")),
        ("Bekletilen kalem", c.get("held")),
        ("Hariç kalem", c.get("excluded")),
    ]
    # Plan maliyet özeti (16.10.2026) — yalnız YENİ satırlar, eskilerin sırası aynı
    co = pr.get("cost_overview")
    cv = cost_overview_view(report) if co else None
    if co and cv:
        rows += [(f"Ürün başına maliyet toplamı ({sym})", co["products_total"]),
                 ("Fiyatı eksik kalem (maliyet özeti)", co["total"]["unpriced"])]
        rows += [(f"Toplam gereken — {q['label'].lower()}", q["text"]) for q in cv["quantities"]
                 if q["label"] != "Ürün"]            # ürün / adet yukarıda zaten var
        if pr.get("package_warned"):
            rows.append(("Ambalaj uyarılı satır", pr["package_warned"]))
    for k, v in rows:
        r = _append_as_text(ws, [k, v])
        if isinstance(v, (int, float)):
            ws.cell(r, 2).number_format = "#,##0"
            ws.cell(r, 2).alignment = Alignment(horizontal="left")
        if k == plan_label:
            ws.cell(r, 1).font = bold
            ws.cell(r, 2).font = bold
    ws.append([])
    r = _append_as_text(ws, ["Bölümler", ""])
    ws.cell(r, 1).font = bold
    for s in secs:
        _append_as_text(ws, [f"{s['no']}. {s['title']}", s.get("summary") or ""])
    ws.append([])
    r = _append_as_text(ws, ["Notlar ve kaynaklar", ""])
    ws.cell(r, 1).font = bold
    for t in notes_lines(report):
        _append_as_text(ws, ["", t])
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 120


# ─── Giriş ──────────────────────────────────────────────────────────────────

def build_workbook(report: dict) -> bytes:
    """Rapor dict'inden formüllü + önbellek değerli xlsx baytları."""
    from openpyxl import Workbook

    meta = report.get("meta") or {}
    cur = (report.get("pricing") or {}).get("currency") or (meta.get("options") or {}).get("currency") or "USD"
    secs = sections(report)
    wb = Workbook()
    list_cache: Dict[str, object] = {}
    stock_cache: Dict[str, object] = {}
    detail_cache: Dict[str, object] = {}
    _sheet_exec(wb, report)                 # ilk sayfa (wb.active)
    _sheet_list(wb, report, secs, list_cache)
    _sheet_stock(wb, report, secs, stock_cache)
    _sheet_suppliers(wb, report, cur)
    _sheet_products(wb, report)
    cost_cache: Dict[str, object] = {}
    _sheet_cost(wb, report, cost_cache)
    _sheet_detail(wb, report, detail_cache)
    _sheet_summary(wb, report, secs, cur)
    title = _clean(meta.get("title")) or "Satın Alma Planı"
    for name in wb.sheetnames:
        text_sheet = name in (SHEET_EXEC, SHEET_SUMMARY)
        # Süzgeç yalnız TEK bloklu sayfalarda: Alım listesi / Stoktan karşılananlar
        # bölüm başlığı + sabit aralıklı ara toplam (=SUM(K2:K42)) satırları taşır;
        # süzgeçten sıralamak ara toplamı veri arasına atar ve aralığı bozar.
        _polish(wb[name], zebra=not text_sheet, filt=name in (SHEET_SUP, SHEET_PROD, SHEET_DETAIL))
        _print_setup(wb[name], landscape=not text_sheet, title=title)
    wb.properties.title = _clean(meta.get("title")) or "Satın Alma Planı"
    wb.properties.creator = "Minerva 108"
    buf = BytesIO()
    wb.save(buf)
    return inject_cached_values(buf.getvalue(), {SHEET_LIST: list_cache, SHEET_STOCK: stock_cache,
                                                 SHEET_COST: cost_cache, SHEET_DETAIL: detail_cache})
