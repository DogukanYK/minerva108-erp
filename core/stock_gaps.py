# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Eksik Hammaddeler raporu — TEK KAYNAK (core/production_sim.py kalıbı).

Sistemde sıfır stoğu listeleyen hiçbir ekran yoktu: tüm uyarılar (dashboard,
eski /api/reports/low-stock-alert, push bildirimleri) tek koşula bağlıydı —
`min_stock_level > 0 and current_stock <= min_stock_level`. Min. stoğu
tanımsız (0, DB varsayılanı) bir hammadde stoğu bitse bile hiçbir yerde
görünmüyordu (09.09.2026 tespiti — bkz. CLAUDE.md "Eksik Hammaddeler").

  classify(current, min, sample_qty) -> 'sifir'|'numune'|'kritik'|None
  assemble(db, domain, ...)          -> rapor dict'i (JSON + Excel + PDF ortak)
  close_open_flags / open_flag       -> "sipariş verildi" bayrağı yaşam döngüsü
  build_workbook / render_pdf        -> dışa aktarma

DEĞİŞMEZLER
───────────
`Item.current_stock` TEK OTORİTE — `Inventory` toplanmaz (stok_lots.py'deki
aynı kural). Numune miktarı yalnız `Inventory.is_sample=True, quantity>0`
lotlarından; numune STOK DEĞİLDİR (24.08.2026 kuralı), bu yüzden ayrı bir
durum (`numune`) — `current_stock<=0` olsa da bir şey elde var demektir.
Sıfır stok + min>0 HER ZAMAN 'sifir' kalır, asla 'kritik' değildir (rank
sırası bunu garanti eder) — "elimizde hiç yok" ile "azaldı" karıştırılmaz.

"AYNI MALZEME" GRUP KURALI (13.10.2026, `core/group_stock`)
───────────────────────────────────────────────────────────
Yalnız `group_stock.group_stock_enabled(db)` açıkken (= üretimdeki
`production.source_choice.enabled`; 0 → eski kart başı davranış birebir):
≥2 sayılan kartlı grubun kartları grup TOPLAMIYLA sınıflanır —
`classify(toplam, grup_min, çevrilmiş numune)`.  Eksikteki grup tek satır
(`kind="group"`, kart satırları `members[]`), grubu yeterli kart listeden
düşer.  Sipariş bayrağı KART başına kalır; grup satırı üyelerden birinde
bayrak varsa "sipariş verildi" görünür (`order` + `ordered_count`).  Seçim
modu (`item_ids`) kart başıdır, gruptaki karta `group` notu eklenir.
Excel/PDF önce grup satırı, sonra üyeleri yazar; kart düzeyindeki işler
(sipariş listesi, fiyat sorgusu) `flat_rows` ile üyelere iner.
"""
from datetime import date, datetime
from typing import Dict, List, Optional

from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from database import Inventory, Item, Recipe, RecipeIngredient, StockOrderFlag, Supplier, Transaction, to_tr
from core import group_stock
from core.items_report import tr_key
from core.supplier_prices import normalize

STATUS_RANK = {"sifir": 0, "numune": 1, "kritik": 2, "min_tanimsiz": 3, "yeterli": 4}
STATUS_LABELS = {
    "sifir": "Sıfır stok",
    "numune": "Sadece numune",
    "kritik": "Kritik altı",
    "min_tanimsiz": "Min. tanımsız",
    "yeterli": "Yeterli",
}
CATEGORIES = ("hammadde", "ambalaj", "all")
_NOT_HAMMADDE = ("Ambalaj", "Bitmiş Ürün")
EPS = 1e-9


def category_filter(category: str):
    """UI kategori anahtarını Item filtre koşullarına çevir (Ürünler sekmesiyle
    aynı negasyon — templates/items.html TABS['hammadde'])."""
    k = (category or "hammadde").lower()
    if k == "ambalaj":
        return (Item.category == "Ambalaj",)
    if k == "all":
        return ()
    return (Item.category.notin_(_NOT_HAMMADDE) | Item.category.is_(None),)


def classify(current: float, min_level: float, sample_qty: float) -> Optional[str]:
    current = float(current or 0.0)
    min_level = float(min_level or 0.0)
    sample_qty = float(sample_qty or 0.0)
    if current <= EPS:
        return "numune" if sample_qty > EPS else "sifir"
    if min_level > EPS and current <= min_level + EPS:
        return "kritik"
    return None


def status_label(status: str) -> str:
    return STATUS_LABELS.get(status, status or "—")


# ─── "Aynı malzeme" grup kuralı (13.10.2026) ──────────────────────────────────

def group_index(db: Session, domain: str):
    """Grup kuralı açık mı + kart → GroupStock eşlemeleri.

    Döner `(açık, tümü, toplayan)`: `tümü` her grup üyesi (sayılan + ayrı
    birimli) → grubu; `toplayan` yalnız ≥2 sayılan kartlı grupların SAYILAN
    kartları.  Kural kapalıysa (`group_stock_enabled` = üretimdeki
    `production.source_choice.enabled`) iki eşleme de boştur → eski kart başı
    davranış birebir."""
    if not group_stock.group_stock_enabled(db):
        return False, {}, {}
    idx = group_stock.index_by_item(group_stock.load_groups(db, domain))
    agg = {iid: g for iid, g in idx.items() if g.aggregates and iid in g.counted_ids}
    return True, idx, agg


def group_sample_qty(g, sample_qty_by_item: Dict[int, float]) -> float:
    """Sayılan üyelerin numune toplamı, grubun taban biriminde."""
    tot = 0.0
    for m in g.counted:
        c = group_stock.convert_qty(sample_qty_by_item.get(m.id, 0.0), m.unit, g.base_unit)
        tot += c or 0.0
    return round(tot, 6)


def group_status(g, sample_qty: float) -> str:
    """Grubun durumu — `classify(toplam, grup_min, numune)`; eksik değilse
    `min_tanimsiz` (min yok + stok var) ya da `yeterli`."""
    st = classify(g.total, g.min_stock, sample_qty)
    if st is not None:
        return st
    return "min_tanimsiz" if (g.min_stock <= EPS and g.total > EPS) else "yeterli"


def group_key(g) -> str:
    return f"g{g.group_id}-{g.bucket}"


def flat_rows(report: dict):
    """Kart düzeyinde satırlar — grup satırı yerine üyeleri (sipariş listesi,
    fiyat sorgusu, seçim)."""
    for r in report.get("rows") or []:
        if r.get("kind") == "group":
            for m in r.get("members") or []:
                yield m
        else:
            yield r


# ─── Rapor derleme ────────────────────────────────────────────────────────────

def assemble(db: Session, domain: str, *, category: str = "hammadde",
             statuses: Optional[List[str]] = None, include_undefined_min: bool = False,
             q: Optional[str] = None, item_ids: Optional[List[int]] = None) -> dict:
    """Rapor derle.

    `item_ids` verilirse SEÇİM MODU: kategori/durum/arama süzgeçleri yok sayılır,
    tam olarak seçilen kartlar döner (kullanıcı hammadde + ambalajı birlikte tek
    Excel'e alabilsin diye).  Seçilen bir kalem bu arada stoklanmışsa satır
    düşmez, `yeterli` durumuyla görünür — "12 seçtim, 12 satır aldım".

    GRUP KURALI (yalnız `group_stock_enabled` açıkken): ≥2 sayılan kartlı bir
    "aynı malzeme" grubunun kartları tek tek değil grup toplamıyla
    sınıflanır — `classify(toplam, grup_min, numune)`.  Eksikteki grup TEK
    satırdır (`kind="group"`, `members[]` altında kart satırları, sipariş
    bayrağı kart başına); grubu yeterli kart listeden düşer.  Ayrı birimli
    üye ve tek kartlı grup eski kart kuralıyla kalır.  Seçim modunda satırlar
    yine KART başınadır, gruptaki karta `group` notu eklenir.
    """
    category = (category or "hammadde").lower()
    if category not in CATEGORIES:
        category = "hammadde"
    picked = [int(i) for i in item_ids] if item_ids else None

    # (1) aday ürünler — aktif + domain + kategori (seçim modunda: yalnız seçilenler)
    q_items = (
        db.query(Item)
        .options(joinedload(Item.supplier))
        .filter(Item.is_active == True, Item.domain == domain)  # noqa: E712
    )
    q_items = (q_items.filter(Item.id.in_(picked)) if picked
               else q_items.filter(*category_filter(category)))
    items = q_items.all()
    if category == "all" and not picked:
        # Konteyner (parent) kartlar doğası gereği stoksuzdur — "sıfır" gibi
        # görünüp 25 kap satırıyla raporu boğmasın.
        parent_ids = {i.parent_id for i in items if i.parent_id}
        items = [i for i in items if i.id not in parent_ids]

    # (2) numune toplamı — item_id → sample_qty
    sample_rows = (
        db.query(Inventory.item_id, func.sum(Inventory.quantity))
        .filter(Inventory.is_sample == True, Inventory.quantity > 0,  # noqa: E712
                Inventory.domain == domain)
        .group_by(Inventory.item_id)
        .all()
    )
    sample_qty_by_item = {iid: float(qty or 0) for iid, qty in sample_rows}

    use_groups, groups_by_item, agg = group_index(db, domain)
    agg_eval = {} if picked else agg            # seçim modu kart başı kalır

    # Python'da sınıflandır — min tanımsız olanları da say (durum filtresinden ÖNCE)
    min_undefined = 0
    classified = []   # (Item | GroupStock, status)
    seen_groups = set()
    for it in items:
        g = agg_eval.get(it.id)
        if g is not None:
            gk = (g.group_id, g.bucket)
            if gk in seen_groups:
                continue
            seen_groups.add(gk)
            gst = group_status(g, group_sample_qty(g, sample_qty_by_item))
            if gst == "yeterli":
                continue                          # grubu yeterli kart listeden düşer
            if gst == "min_tanimsiz":
                min_undefined += 1
                if include_undefined_min:
                    classified.append((g, gst))
                continue
            classified.append((g, gst))
            continue
        sq = sample_qty_by_item.get(it.id, 0.0)
        st = classify(it.current_stock, it.min_stock_level, sq)
        if st is None:
            undefined_min = ((it.min_stock_level or 0.0) <= EPS and (it.current_stock or 0.0) > EPS)
            if undefined_min:
                min_undefined += 1
            if picked:
                classified.append((it, "min_tanimsiz" if undefined_min else "yeterli"))
            elif undefined_min and include_undefined_min:
                classified.append((it, "min_tanimsiz"))
            continue
        classified.append((it, st))

    if statuses and not picked:
        wanted = {s for s in statuses if s in STATUS_RANK}
        classified = [(e, st) for e, st in classified if st in wanted]

    item_by_id = {it.id: it for it in items}
    card_ids = [e.id for e, _ in classified if isinstance(e, Item)]
    member_ids = [m.id for e, _ in classified if not isinstance(e, Item) for m in e.counted]
    missing = [i for i in member_ids if i not in item_by_id]
    if missing:      # konteyner süzgecine takılan üye — yine de grubun altında görünsün
        for it in (db.query(Item).options(joinedload(Item.supplier))
                   .filter(Item.id.in_(missing)).all()):
            item_by_id[it.id] = it
    ids = card_ids + member_ids

    # (3) numune lot detayı — yalnız 'numune' durumundaki kartlar / grup üyeleri
    sample_lots_by_item: Dict[int, list] = {}
    numune_ids = ([e.id for e, st in classified if st == "numune" and isinstance(e, Item)]
                  + [m.id for e, st in classified if st == "numune" and not isinstance(e, Item)
                     for m in e.counted])
    if numune_ids:
        lots = (
            db.query(Inventory)
            .options(joinedload(Inventory.supplier))
            .filter(Inventory.is_sample == True, Inventory.quantity > 0,  # noqa: E712
                    Inventory.item_id.in_(numune_ids), Inventory.domain == domain)
            .order_by(Inventory.created_at.asc())
            .all()
        )
        for l in lots:
            sample_lots_by_item.setdefault(l.item_id, []).append({
                "lot_number": l.lot_number,
                "quantity": round(float(l.quantity or 0), 4),
                "supplier_name": l.supplier.name if l.supplier else None,
                "expiry_date": l.expiry_date or None,
            })

    # (4) son hareket — item_id → en son Transaction.timestamp
    last_movement_by_item: Dict[int, datetime] = {}
    if ids:
        rows = (
            db.query(Transaction.item_id, func.max(Transaction.timestamp))
            .filter(Transaction.item_id.in_(ids))
            .group_by(Transaction.item_id)
            .all()
        )
        last_movement_by_item = {iid: ts for iid, ts in rows}

    # (5) son tedarikçi — item_id → en son supplier'lı Inventory satırı (N+1 yerine tek sorgu)
    last_supplier_by_item: Dict[int, str] = {}
    if ids:
        sub = (
            db.query(Inventory.item_id, func.max(Inventory.id).label("max_id"))
            .filter(Inventory.item_id.in_(ids), Inventory.supplier_id.isnot(None))
            .group_by(Inventory.item_id)
            .subquery()
        )
        rows = (
            db.query(Inventory.item_id, Supplier.name)
            .join(sub, Inventory.id == sub.c.max_id)
            .join(Supplier, Supplier.id == Inventory.supplier_id)
            .all()
        )
        last_supplier_by_item = {iid: name for iid, name in rows}

    # (6) reçete etkisi — item_id → [{id, name}], yalnız aktif + domain reçeteler
    recipes_by_item: Dict[int, list] = {}
    if ids:
        rows = (
            db.query(RecipeIngredient.item_id, Recipe.id, Recipe.name)
            .join(Recipe, Recipe.id == RecipeIngredient.recipe_id)
            .filter(Recipe.is_active == True, Recipe.domain == domain,  # noqa: E712
                    RecipeIngredient.item_id.in_(ids))
            .distinct()
            .all()
        )
        for iid, rid, rname in rows:
            recipes_by_item.setdefault(iid, []).append({"id": rid, "name": rname})
        for iid in recipes_by_item:
            recipes_by_item[iid].sort(key=lambda r: tr_key(r["name"]))

    # (7) açık sipariş bayrakları — KART başına (grup üyeleri dahil)
    open_flags_by_item: Dict[int, StockOrderFlag] = {}
    open_orders_total = 0
    if ids:
        flags = (
            db.query(StockOrderFlag)
            .options(joinedload(StockOrderFlag.supplier))
            .filter(StockOrderFlag.closed_at.is_(None), StockOrderFlag.domain == domain,
                    StockOrderFlag.item_id.in_(ids))
            .all()
        )
        for f in flags:
            open_flags_by_item[f.item_id] = f
        open_orders_total = len(flags)

    def own_status(it) -> str:
        st = classify(it.current_stock, it.min_stock_level, sample_qty_by_item.get(it.id, 0.0))
        if st is not None:
            return st
        undefined = ((it.min_stock_level or 0.0) <= EPS and (it.current_stock or 0.0) > EPS)
        return "min_tanimsiz" if undefined else "yeterli"

    def card_row(it, st, kind="item") -> dict:
        sup_name = it.supplier.name if it.supplier else None
        flag = open_flags_by_item.get(it.id)
        min_level = float(it.min_stock_level or 0.0)
        current = float(it.current_stock or 0.0)
        lm = last_movement_by_item.get(it.id)
        return {
            "kind": kind,
            "key": str(it.id),
            "item_id": it.id,
            "name": it.name,
            "name_tr": it.name_tr or "",
            "category": it.category or "—",
            "unit": it.unit or "",
            "current_stock": round(current, 4),
            "min_stock_level": round(min_level, 4),
            "deficit": round(min_level - current, 4) if min_level > EPS else None,
            "status": st,
            "status_label": status_label(st),
            "sample_qty": round(sample_qty_by_item.get(it.id, 0.0), 4),
            "sample_lots": sample_lots_by_item.get(it.id, []),
            "supplier_id": it.supplier_id,
            "supplier_name": sup_name,
            "last_supplier": last_supplier_by_item.get(it.id) or sup_name,
            "last_movement": to_tr(lm).strftime("%d.%m.%Y %H:%M") if lm else None,
            "recipes": recipes_by_item.get(it.id, []),
            "recipe_count": len(recipes_by_item.get(it.id, [])),
            "order": _flag_view(flag) if flag else None,
            "group": None,
        }

    def group_note(it) -> Optional[dict]:
        """Seçim modunda gruptaki karta grup bilgisi (satır kart başı kalır)."""
        g = groups_by_item.get(it.id)
        if g is None:
            return None
        counted = g.aggregates and it.id in g.counted_ids
        gst = group_status(g, group_sample_qty(g, sample_qty_by_item)) if counted else None
        return {"id": g.group_id, "name": g.name, "counted": counted,
                "total": round(g.total, 4), "unit": g.base_unit,
                "total_text": group_stock.amount_exact(g.total, g.base_unit),
                "min_text": group_stock.amount_exact(g.min_stock, g.base_unit) if g.min_stock > EPS else "",
                "breakdown_text": g.breakdown_text, "aside_text": g.aside_text,
                "status": gst, "status_label": status_label(gst) if gst else ""}

    def group_row(g, st) -> dict:
        base = g.base_unit
        members = [card_row(item_by_id[m.id], own_status(item_by_id[m.id]), kind="member")
                   for m in g.counted if m.id in item_by_id]
        for mr in members:
            mr["stock_text"] = group_stock.amount_short(max(0.0, mr["current_stock"]), mr["unit"])
        flagged = [mr for mr in members if mr["order"]]
        recipes, seen_r = [], set()
        for mr in members:
            for rec in mr["recipes"]:
                if rec["id"] not in seen_r:
                    seen_r.add(rec["id"])
                    recipes.append(rec)
        recipes.sort(key=lambda r: tr_key(r["name"]))
        moves = [last_movement_by_item[m.id] for m in g.counted if last_movement_by_item.get(m.id)]
        lm = max(moves) if moves else None
        gsq = group_sample_qty(g, sample_qty_by_item)
        lots = [dict(l, item_name=mr["name"], unit=mr["unit"])
                for mr in members for l in mr["sample_lots"]]
        total, gmin = float(g.total), float(g.min_stock)
        first = item_by_id.get(g.counted[0].id) if g.counted else None
        return {
            "kind": "group",
            "key": group_key(g),
            "item_id": None,
            "group_id": g.group_id,
            "bucket": g.bucket,
            "name": g.name,
            "name_tr": "",
            "category": (first.category if first is not None else None) or "—",
            "unit": base,
            "current_stock": round(total, 4),
            "min_stock_level": round(gmin, 4),
            "deficit": round(gmin - total, 4) if gmin > EPS else None,
            # Kesin metin (taban birim korunur, pozitif eksik asla "0"):
            # amount_short 999,6 g / 1.000 g'ı "1 kg / 1 kg" gösteriyordu.
            "current_text": group_stock.amount_exact(total, base),
            "min_text": group_stock.amount_exact(gmin, base) if gmin > EPS else "",
            "deficit_text": (group_stock.amount_exact(gmin - total, base)
                             if gmin > EPS and gmin - total > EPS else ""),
            "status": st,
            "status_label": status_label(st),
            "sample_qty": round(gsq, 4),
            "sample_text": group_stock.amount_exact(gsq, base) if gsq > EPS else "",
            "sample_lots": lots,
            "supplier_id": None,
            "supplier_name": None,
            "last_supplier": None,
            "last_movement": to_tr(lm).strftime("%d.%m.%Y %H:%M") if lm else None,
            "recipes": recipes,
            "recipe_count": len(recipes),
            "order": flagged[0]["order"] if flagged else None,
            "ordered_count": len(flagged),
            "member_count": len(members),
            "member_ids": [mr["item_id"] for mr in members],
            "members": members,
            "breakdown_text": g.breakdown_text,
            "aside_text": g.aside_text,
            "group": {"id": g.group_id, "name": g.name},
        }

    qn = normalize(q) if (q and not picked) else None
    rows = []
    for e, st in classified:
        if isinstance(e, Item):
            row = card_row(e, st)
            if qn:
                hay = normalize(f"{e.name} {e.name_tr or ''} {row['supplier_name'] or ''} "
                                f"{row['last_supplier'] or ''}")
                if qn not in hay:
                    continue
            if use_groups and picked:
                row["group"] = group_note(e)
            rows.append(row)
            continue
        row = group_row(e, st)
        if qn:
            hay = normalize(" ".join([row["name"]] + [
                f"{m['name']} {m['name_tr']} {m['supplier_name'] or ''} {m['last_supplier'] or ''}"
                for m in row["members"]]))
            if qn not in hay:
                continue
        rows.append(row)

    rows.sort(key=lambda r: (STATUS_RANK.get(r["status"], 9), tr_key(r["name"])))

    summary = {
        "sifir": sum(1 for _, st in classified if st == "sifir"),
        "kritik": sum(1 for _, st in classified if st == "kritik"),
        "numune": sum(1 for _, st in classified if st == "numune"),
        "min_undefined": min_undefined,
        "total": len(rows),
        "open_orders": open_orders_total,
        "groups": sum(1 for r in rows if r["kind"] == "group"),
    }

    return {
        "generated_at": to_tr(datetime.utcnow()).strftime("%d.%m.%Y %H:%M"),
        "category": category,
        "group_rule": use_groups,
        "summary": summary,
        "rows": rows,
    }


def _flag_view(f: StockOrderFlag) -> dict:
    return {
        "id": f.id,
        "supplier_id": f.supplier_id,
        "supplier_name": f.supplier.name if f.supplier else None,
        "quantity": round(float(f.quantity), 4) if f.quantity is not None else None,
        "unit": f.unit or "",
        "expected_date": f.expected_date.strftime("%Y-%m-%d") if f.expected_date else None,
        "note": f.note or "",
        "ordered_by": f.ordered_by or "",
        "ordered_at": to_tr(f.ordered_at).strftime("%d.%m.%Y %H:%M") if f.ordered_at else None,
    }


# ─── Sipariş bayrağı yaşam döngüsü ────────────────────────────────────────────

def close_open_flags(db: Session, item_id: int, reason: str, actor: Optional[str] = None) -> int:
    """Ürünün açık bayraklarını kapat — COMMIT YAPMAZ, çağıran kendi işleminde eder.

    `reason='received'`: gerçek mal kabul (routers/inventory.py receive_stock
    numune-olmayan dal + samples/convert). `reason='manual'`: kullanıcı elle
    kapattı veya yeniden işaretledi."""
    open_flags = (
        db.query(StockOrderFlag)
        .filter(StockOrderFlag.item_id == item_id, StockOrderFlag.closed_at.is_(None))
        .with_for_update()
        .all()
    )
    now = datetime.utcnow()
    for f in open_flags:
        f.closed_at = now
        f.closed_reason = reason
    return len(open_flags)


def open_flag(db: Session, item: Item, *, supplier_id: Optional[int], quantity: Optional[float],
             expected_date: Optional[date], note: Optional[str], actor: str, domain: str) -> StockOrderFlag:
    """Yeni sipariş bayrağı aç — varsa eskisini 'manual' ile kapatır. flush eder, COMMIT ÇAĞIRANDA."""
    close_open_flags(db, item.id, "manual", actor)
    f = StockOrderFlag(
        item_id=item.id, supplier_id=supplier_id, quantity=quantity, unit=item.unit,
        expected_date=expected_date, note=(note or "").strip()[:300] or None,
        ordered_by=actor, domain=domain,
    )
    db.add(f)
    db.flush()
    return f


# ─── Dışa aktarma ─────────────────────────────────────────────────────────────

def export_filename(category: str, ext: str) -> str:
    return f"eksik_hammaddeler_{category}_{datetime.utcnow():%Y%m%d}.{ext}"


def group_note_text(g: Optional[dict]) -> str:
    """Seçim modunda gruptaki kartın grup notu (`assemble` → `row["group"]`)
    tek satır metin: Excel "Grup dökümü" sütunu."""
    if not g or not g.get("total_text"):
        return ""
    t = f"«{g['name']}» grubu: {g['total_text']}"
    if g.get("min_text"):
        t += f" / min {g['min_text']}"
    if g.get("status_label"):
        t += f" · {g['status_label']}"
    if g.get("breakdown_text"):
        t += f" — {g['breakdown_text']}"
    if g.get("aside_text"):
        t += f" · ayrı birim: {g['aside_text']}"
    return t


def build_workbook(report: dict, prices: Optional[dict] = None) -> bytes:
    """Rapordan 3-sayfalı .xlsx üretir (Eksikler / Sipariş Listesi / Reçete Etkisi).

    `prices` verilirse ({item_id: [{supplier_name, package_size, unit_price}, …]},
    core.supplier_prices.prices_for_items), Sipariş Listesi'nde her satıra 3
    tedarikçiye kadar [Tedarikçi · Alınabilecek Miktar · Birim Fiyat] eklenir
    (core.production_sim.build_workbook ile aynı kalıp)."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    s = report["summary"]
    prices = prices or {}
    wb = Workbook()
    NAVY = PatternFill("solid", fgColor="232E6E"); HF = Font(bold=True, color="FFFFFF")
    RED = Font(color="B91C1C", bold=True); GRAY = Font(color="9CA3AF")
    TITLE = Font(bold=True, size=13, color="232E6E"); BLD = Font(bold=True)
    thin = Border(*([Side(style="thin", color="DDDDDD")] * 4))

    def hrow(ws, headers, row=1):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row, ci, h); c.fill = NAVY; c.font = HF
            c.alignment = Alignment(horizontal="left", vertical="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Eksikler
    ws = wb.active; ws.title = "Eksikler"
    ws["A1"] = f"Minerva 108 — Eksik Hammaddeler ({report.get('category', 'hammadde')})"
    ws["A1"].font = TITLE
    meta = [
        ("Tarih", report["generated_at"]),
        ("Sıfır stok", s["sifir"]), ("Kritik altı", s["kritik"]),
        ("Sadece numune", s["numune"]), ("Min. tanımsız", s["min_undefined"]),
        ("Açık sipariş", s["open_orders"]),
    ]
    if s.get("groups"):
        meta.append(("Aynı malzeme grubu", s["groups"]))
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(i, 1, k).font = BLD; ws.cell(i, 2, v)
    hstart = 3 + len(meta) + 1
    # "Grup dökümü" sütunu yalnız grup kuralı açıkken (ya da grup satırı /
    # grup notu varsa) — kural kapalıyken sayfa eski 13 sütun, birebir.
    show_grp = bool(report.get("group_rule")) or any(
        r.get("kind") == "group" or r.get("group") for r in report["rows"])
    hrow2 = ["Hammadde", "Türkçe Ad", "Kategori", "Durum", "Mevcut", "Min", "Eksik", "Birim",
             "Numune", "Reçete Sayısı", "Son Hareket", "Tedarikçi", "Sipariş"]
    if show_grp:
        hrow2.append("Grup dökümü")
    for ci, h in enumerate(hrow2, 1):
        c = ws.cell(hstart, ci, h); c.fill = NAVY; c.font = HF

    def _order_txt(r):
        return (f"{r['order']['supplier_name'] or '—'} · {r['order']['quantity'] or ''} {r['order']['unit']}"
                if r.get("order") else "")

    ri = hstart + 1
    for r in report["rows"]:
        is_group = r.get("kind") == "group"
        if is_group:
            # Önce grup satırı (kalın; miktarlar grubun taban biriminde),
            # sonra üyeler girintili — kart başı sipariş bayraklarıyla.
            dok = r["breakdown_text"] + (f" · ayrı birim: {r['aside_text']}" if r.get("aside_text") else "")
            vals = [f"{r['name']} (aynı malzeme, {r['member_count']} kart)", "", r["category"],
                    r["status_label"], r["current_stock"], r["min_stock_level"],
                    r["deficit"] if r["deficit"] is not None else "", r["unit"], r["sample_qty"],
                    r["recipe_count"], r["last_movement"] or "—", "—",
                    f"{r['ordered_count']} kartta sipariş" if r.get("ordered_count") else "", dok]
            subs = r.get("members") or []
        else:
            vals = [r["name"], r["name_tr"], r["category"], r["status_label"], r["current_stock"],
                    r["min_stock_level"], r["deficit"] if r["deficit"] is not None else "", r["unit"],
                    r["sample_qty"], r["recipe_count"], r["last_movement"] or "—",
                    r["last_supplier"] or "—", _order_txt(r), group_note_text(r.get("group"))]
            subs = []
        if not show_grp:
            vals = vals[:13]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v); c.border = thin
            if is_group:
                c.font = BLD
            if r["status"] in ("sifir",) and ci == 4:
                c.font = RED
        ri += 1
        for m in subs:
            mv = [m["name"], m["name_tr"], m["category"], m["status_label"], m["current_stock"],
                  m["min_stock_level"], m["deficit"] if m["deficit"] is not None else "", m["unit"],
                  m["sample_qty"], m["recipe_count"], m["last_movement"] or "—",
                  m["last_supplier"] or "—", _order_txt(m), f"«{r['name']}» grubunda"]
            for ci, v in enumerate(mv, 1):
                c = ws.cell(ri, ci, v); c.border = thin
                if ci == 1:
                    c.alignment = Alignment(indent=2)
                if ci == 14:
                    c.font = GRAY
            ri += 1
    ws.freeze_panes = f"A{hstart + 1}"
    widths(ws, [30, 24, 12, 14, 10, 10, 10, 8, 10, 12, 16, 20, 26] + ([40] if show_grp else []))

    # Sipariş Listesi — yalnız açık siparişli satırlar, tedarikçiye göre gruplu
    ws = wb.create_sheet("Sipariş Listesi")
    ordered = [r for r in flat_rows(report) if r.get("order")]   # bayrak KART başına
    show_sup = bool(prices)
    SUP_SLOTS = 3
    base = ["Hammadde", "Miktar", "Birim", "Beklenen", "Not", "Sipariş Veren"]
    headers = list(base)
    if show_sup:
        for n in range(1, SUP_SLOTS + 1):
            headers += [f"TEDARİKÇİ-{n}", "Alınabilecek Miktar", "Birim Fiyat"]
    if not ordered:
        ws["A1"] = "Açık sipariş yok."; ws["A1"].font = BLD
    else:
        by_supplier: Dict[str, list] = {}
        for r in ordered:
            by_supplier.setdefault(r["order"]["supplier_name"] or "Tedarikçi belirtilmedi", []).append(r)
        ri = 1
        for sup_name in sorted(by_supplier, key=tr_key):
            c = ws.cell(ri, 1, sup_name); c.font = HF; c.fill = NAVY
            ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=len(headers))
            ri += 1
            hrow(ws, headers, row=ri)
            ri += 1
            for r in by_supplier[sup_name]:
                o = r["order"]
                vals = [r["name"], o["quantity"] or "", o["unit"], o["expected_date"] or "",
                        o["note"] or "", o["ordered_by"] or ""]
                gray_cols = set()
                fmt_cols = {}                            # fiyat hücresi → "#,##0.00## ₺/kg"
                if show_sup:
                    from core.supplier_prices import price_cell_format, vat_label
                    plist = prices.get(r["item_id"], [])
                    for n in range(SUP_SLOTS):
                        if n < len(plist):
                            p = plist[n]
                            name = p["supplier_name"]
                            if p.get("vat_included"):    # brüt fiyat — sıralama netle yapıldı
                                name = f"{name} ({vat_label(p)})"
                            if p.get("inactive"):        # pasif firma: sonda, gri, etiketli
                                name = f"{name} (pasif tedarikçi)"
                                gray_cols.update(range(len(vals) + 1, len(vals) + 4))
                            if p.get("unit_price"):
                                fmt_cols[len(vals) + 3] = price_cell_format(p)
                            vals += [name, p["package_size"] or "", p["unit_price"] or ""]
                        else:
                            vals += ["", "", ""]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(ri, ci, v); c.border = thin
                    if ci in gray_cols:
                        c.font = GRAY
                    if ci in fmt_cols:
                        c.number_format = fmt_cols[ci]
                ri += 1
            ri += 1
    widths(ws, [30, 10, 8, 12, 30, 16] + ([16, 16, 12] * SUP_SLOTS if show_sup else []))

    # Reçete Etkisi
    ws = wb.create_sheet("Reçete Etkisi")
    hrow(ws, ["Hammadde", "Durum", "Reçete"])
    ri = 2
    for r in report["rows"]:
        if not r["recipes"]:
            continue
        for rec in r["recipes"]:
            vals = [r["name"], r["status_label"], rec["name"]]
            for ci, v in enumerate(vals, 1):
                ws.cell(ri, ci, v).border = thin
            ri += 1
    ws.freeze_panes = "A2"; widths(ws, [30, 14, 46])

    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()


def _story(report: dict, s: float = 1.0):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer

    from core.monthly_report import _register_fonts
    from core.shipment_note import _grid

    font, font_b = _register_fonts()
    NAVY = colors.HexColor("#232E6E")
    MUTED = colors.HexColor("#6b7280")
    st = {
        "font": font, "font_b": font_b, "NAVY": NAVY,
        "LIGHT": colors.HexColor("#f5f0e8"), "GREY": colors.HexColor("#e5e7eb"),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=7.6 * s, textColor=colors.HexColor("#374151"), leading=10 * s),
        "hcell": ParagraphStyle("hcell", fontName=font_b, fontSize=7.6 * s, textColor=colors.white, leading=10 * s),
    }
    h1 = ParagraphStyle("h1", fontName=font_b, fontSize=15 * s, textColor=NAVY, spaceAfter=2 * s, leading=18 * s)
    meta = ParagraphStyle("meta", fontName=font, fontSize=8 * s, textColor=MUTED, leading=11 * s)

    sm = report["summary"]
    cat_label = {"hammadde": "Hammadde", "ambalaj": "Ambalaj", "all": "Tümü"}.get(report.get("category"), "Hammadde")
    story = [
        Paragraph("EKSİK HAMMADDELER", h1),
        Paragraph(f"{report['generated_at']}  ·  Kapsam: {cat_label}  ·  "
                 f"Sıfır: {sm['sifir']}  ·  Kritik altı: {sm['kritik']}  ·  "
                 f"Sadece numune: {sm['numune']}  ·  Min. tanımsız: {sm['min_undefined']}  ·  "
                 f"Açık sipariş: {sm['open_orders']}", meta),
        Spacer(1, 8 * s),
    ]
    headers = ["HAMMADDE", "DURUM", "MEVCUT / MİN", "EKSİK", "NUMUNE", "REÇETE", "SİPARİŞ"]
    W = 174.0
    widths = [W * f * mm for f in (0.27, 0.13, 0.16, 0.11, 0.10, 0.09, 0.14)]
    def _card_cells(r, name):
        cur_min = f"{r['current_stock']:g} / {r['min_stock_level']:g} {r['unit']}"
        deficit = f"{r['deficit']:g} {r['unit']}" if r["deficit"] is not None else "—"
        numune = f"{r['sample_qty']:g} {r['unit']}" if r["sample_qty"] else "—"
        order = (f"{r['order']['supplier_name'] or '—'} · {r['order']['quantity'] or ''} {r['order']['unit']}"
                 if r.get("order") else "—")
        return [name, r["status_label"], cur_min, deficit, numune,
                str(r["recipe_count"]) if r["recipe_count"] else "—", order]

    rows = []
    for r in report["rows"]:
        if r.get("kind") != "group":
            g = r.get("group")
            # seçim modunda gruptaki karta grup notu — ekrandaki grup satırıyla çelişmesin
            gn = (f" («{g['name']}» grubu: {g['total_text']}"
                  + (f" · {g['status_label']}" if g.get("status_label") else "") + ")"
                  if g and g.get("total_text") else "")
            rows.append(_card_cells(r, r["name"] + gn))
            continue
        # Grup satırı (taban birimde, sunucu biçimli metinler) + üyeleri
        rows.append([f"{r['name']} — aynı malzeme, {r['member_count']} kart: {r['breakdown_text']}",
                     r["status_label"], f"{r['current_text']} / {r['min_text'] or '—'}",
                     r["deficit_text"] or "—", r["sample_text"] or "—",
                     str(r["recipe_count"]) if r["recipe_count"] else "—",
                     f"{r['ordered_count']} kartta" if r.get("ordered_count") else "—"])
        for m in r.get("members") or []:
            label = f"– {m['name']}" + (f" ({m['supplier_name']})" if m.get("supplier_name") else "")
            rows.append(_card_cells(m, label))
    story.append(_grid(st, headers, rows, widths, s))
    return story


def render_pdf(report: dict) -> bytes:
    from reportlab.lib.units import mm
    from core.delivery_note import render_autofit, merge_letterhead

    content = render_autofit(
        lambda s: _story(report, s),
        margins=(20 * mm, 20 * mm, 48 * mm, 40 * mm),
        doc_kwargs={"title": "Eksik Hammaddeler", "author": "Minerva 108"})
    return merge_letterhead(content)
