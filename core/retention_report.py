# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Şahit numune dolabı raporu — yazdırılabilir PDF + Excel.

  build_report(views, *, filters, generated_by, domain_label) -> dict
  filters_text(...)          -> str     ekrandaki filtrelerin Türkçe özeti
  render_pdf(report)         -> bytes   A4 yatay, çok sayfa, başlık satırı her sayfada
  render_xlsx(report)        -> bytes   'Dolap' (tek sayfa + Dolap sütunu) + 'Özet'
  report_filename(ext)       -> str

Satırlar `routers/retention._view()` çıktısından gelir (liste ucuyla AYNI
filtre + görünüm yolu) — rapor ekranda görünenden başka bir şey basamaz.
DB session/sorgu YOK; burası yalnız gruplar, sıralar ve çizer.

Gruplama: dolap (= marka) → raf → göz.  Raf/göz serbest metin ("2", "10",
"B") olduğu için DOĞAL sıralama kullanılır: "Raf 10", "Raf 2"'nin önüne
düşmez.  Rafı boş kayıtlar her dolabın sonunda "Raf belirtilmemiş" grubunda.

Renk kuralı yalnız DOLAPTAKİ (status='stored') kayıtlara uygulanır: tükenmiş
bir kaydın geçmiş tarihi kırmızı basılsa "imha edilmesi gereken numune var"
diye yanlış alarm verirdi.  Özet sayaçları da aynı kuralı kullanır
(`list_cabinets` ve `expired=1` filtresiyle tutarlı).

PDF `render_autofit` KULLANMAZ — tek sayfaya sıkıştırma değil, okunaklı çok
sayfa isteniyor (220+ kayıt).  Yazı tipleri `monthly_report._register_fonts`.
"""
import re
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from core.delivery_note import _fmt
from core.retention import DUE_SOON_DAYS, status_label
from database import tr_now

TITLE = "Şahit Numune Dolabı — Envanter Listesi"
NO_SHELF = "Raf belirtilmemiş"

PDF_COLUMNS = ("Göz", "Ürün", "Boy", "Lot", "Kalan / İlk", "Üretim",
               "Saklama bitişi", "Son kontrol", "Durum", "Not")
# A4 yatay, 14 mm kenar boşluğu → 269 mm kullanılabilir genişlik
_PDF_WIDTHS_MM = (12, 62, 18, 24, 17, 18, 21, 33, 26, 38)

XLSX_COLUMNS = ("Dolap", "Raf", "Göz", "Ürün", "Boy", "Lot", "Kalan", "İlk",
                "Birim", "Üretim", "Saklama bitişi", "Son kontrol",
                "Kontrol sonucu", "Durum", "Not")
_XLSX_WIDTHS = (16, 8, 7, 40, 14, 14, 9, 9, 8, 12, 14, 12, 15, 22, 44)

SUMMARY_COLUMNS = ("Dolap", "Kayıt", "Dolapta (kayıt)", "Kalan adet",
                   "Süresi dolan", "Süresi yaklaşan", "Boyu belirsiz")


# ─── Sıralama ────────────────────────────────────────────────────────────────

_TR_ALPHA = "abcçdefgğhıijklmnoöprsştuüvyz"
_TR_ORDER = {c: i for i, c in enumerate(_TR_ALPHA)}
_NUM = re.compile(r"(\d+)")


def _tr_key(s) -> tuple:
    """Türkçe alfabe sırası (Ç C'den sonra, İ/ı ayrı).  Rakam/boşluk/noktalama
    harflerden önce, alfabe dışı harfler (q, w, x, é) sonra gelir."""
    s = str(s or "").replace("I", "ı").replace("İ", "i").lower()
    return tuple(_TR_ORDER[ch] if ch in _TR_ORDER
                 else (ord(ch) - 100_000 if not ch.isalpha() else 1000 + ord(ch))
                 for ch in s)


def natural_key(s) -> tuple:
    """Doğal sıra: "2" < "10", "A2" < "A10".  Boş değer EN SONA düşer."""
    s = str(s or "").strip()
    if not s:
        return (1, ())
    parts = []
    for tok in _NUM.split(s):
        if not tok:
            continue
        # isdecimal = `\d`'nin yakaladığı küme.  isdigit "²"/"①" için de True
        # döner, int() ise patlar → tek bir "1²" rafı tüm raporu 500'e düşürürdü.
        parts.append((0, int(tok), ()) if tok.isdecimal() else (1, 0, _tr_key(tok)))
    return (0, tuple(parts))


# ─── Veri ────────────────────────────────────────────────────────────────────

def filters_text(*, brand=None, q=None, status=None, expired=None,
                 needs_review=None, unchecked=None, item_label=None,
                 unsized=None) -> str:
    """Raporun başına basılan filtre satırı — kağıt elden ele geçtiğinde
    "bu liste eksik mi, süzülmüş mü?" sorusunun cevabı."""
    parts = [f"Dolap: {brand}" if brand else "Tüm dolaplar"]
    if item_label:
        parts.append(f"Ürün: {item_label}")
    if q and q.strip():
        parts.append(f"Arama: “{q.strip()}”")
    parts.append(f"Durum: {status_label(status)}" if status else "İmha edilenler hariç")
    if expired:
        parts.append("Yalnız süresi dolanlar")
    if needs_review:
        parts.append("Teyit bekleyenler")
    if unchecked:
        parts.append("Hiç kontrol edilmemiş")
    if unsized:
        parts.append("Boyu belirsiz")
    return " · ".join(parts)


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _row(v: dict) -> dict:
    """`_view()` sözlüğünden rapor satırı."""
    stored = v.get("status") == "stored"
    is_parent = bool(v.get("is_parent"))
    qty, init = _num(v.get("quantity")), _num(v.get("initial_quantity"))
    unit = v.get("unit") or "adet"
    qty_label = f"{_fmt(qty)} / {_fmt(init)}" + ("" if unit == "adet" else f" {unit}")

    status = v.get("status_label") or status_label(v.get("status"))
    flags = []
    if v.get("qc_status") == "rejected":
        flags.append("QC RED")
    if v.get("needs_review"):
        flags.append("teyit bekliyor")
    if flags:
        status = f"{status} · {' · '.join(flags)}"

    last = ""
    if v.get("last_check_on"):
        last = v["last_check_on"]
        if v.get("last_check_result_label"):
            last += f" · {v['last_check_result_label']}"
        if (v.get("check_count") or 0) > 1:
            last += f" ({v['check_count']} kontrol)"

    return {
        "brand": v.get("brand") or "—",
        "shelf": (v.get("shelf") or "").strip(),
        "slot": (v.get("slot") or "").strip(),
        # Ekrandaki gibi aile adı + ayrı Boy sütunu (varyasyon adı zaten
        # boyu içeriyor; ikisini birden basmak "Losyon 200ml · 200ml" olurdu).
        "product": v.get("parent_name") or v.get("item_name") or "",
        "size": v.get("variation_name") or ("boy seçilmedi" if is_parent else ""),
        "unsized": is_parent,
        "lot": v.get("lot_number") or "",
        "quantity": qty,
        "initial": init,
        "unit": unit,
        "qty_label": qty_label,
        "produced": v.get("produced_at") or "",
        "until": v.get("retention_until") or "",
        "until_label": v.get("retention_until_label") or "—",
        "expiry": (v.get("expiry_state") or "") if stored else "",
        "last_check": last,
        "last_check_on": v.get("last_check_on") or "",
        "last_check_result": v.get("last_check_result_label") or "",
        "last_check_bad": bool(v.get("last_check_on"))
                          and v.get("last_check_result") == "uygun_degil",
        "status": status,
        "stored": stored,
        "note": v.get("note") or "",
    }


def _stats() -> dict:
    return {"samples": 0, "stored": 0, "quantity": 0.0, "initial": 0.0,
            "expired": 0, "due_soon": 0, "unsized": 0}


def _add(st: dict, r: dict) -> None:
    st["samples"] += 1
    st["quantity"] += r["quantity"]
    st["initial"] += r["initial"]
    if r["stored"]:
        st["stored"] += 1
    if r["expiry"] == "expired":
        st["expired"] += 1
    elif r["expiry"] == "due_soon":
        st["due_soon"] += 1
    if r["unsized"]:
        st["unsized"] += 1


def build_report(views, *, filters: str = "", generated_by: str = "",
                 domain_label: str = "", now=None) -> dict:
    """Dolap → raf → göz gruplu rapor sözlüğü (PDF ve Excel aynı yapıyı okur)."""
    now = now or tr_now()
    cabs: dict = {}
    totals = _stats()
    for v in views:
        r = _row(v)
        cab = cabs.setdefault(r["brand"], {"brand": r["brand"], "_shelves": {},
                                           **_stats()})
        sh = cab["_shelves"].setdefault(r["shelf"], {
            "shelf": r["shelf"],
            "label": f"Raf {r['shelf']}" if r["shelf"] else NO_SHELF,
            "rows": [], **_stats()})
        sh["rows"].append(r)
        for st in (sh, cab, totals):
            _add(st, r)

    cabinets = []
    for cab in sorted(cabs.values(), key=lambda c: _tr_key(c["brand"])):
        shelves = sorted(cab.pop("_shelves").values(),
                         key=lambda s: natural_key(s["shelf"]))
        for sh in shelves:
            sh["rows"].sort(key=lambda r: (natural_key(r["slot"]), _tr_key(r["product"]),
                                           natural_key(r["size"]), natural_key(r["lot"])))
        cab["shelves"] = shelves
        cabinets.append(cab)
    return {
        "title": TITLE,
        "generated_at": now.strftime("%d.%m.%Y %H:%M"),
        "generated_by": generated_by or "",
        "domain_label": domain_label or "",
        "filters": filters or "",
        "cabinets": cabinets,
        "totals": totals,
        "count": totals["samples"],
    }


def report_filename(ext: str) -> str:
    return f"sahit_numune_dolabi_{tr_now():%Y%m%d}.{ext}"


def _summary_rows(report) -> list:
    rows = [[c["brand"], c["samples"], c["stored"], c["quantity"], c["expired"],
             c["due_soon"], c["unsized"]] for c in report["cabinets"]]
    t = report["totals"]
    rows.append(["TOPLAM", t["samples"], t["stored"], t["quantity"], t["expired"],
                 t["due_soon"], t["unsized"]])
    return rows


# ─── PDF ─────────────────────────────────────────────────────────────────────

def _numbered_canvas(font: str, footer_left: str):
    """"Sayfa X / Y" altbilgili canvas — toplam sayfa ancak build bitince
    bilindiği için sayfalar biriktirilip `save()`'de numaralanır (reportlab'ın
    bilinen iki geçişsiz tarifi)."""
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas

    class _NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._footer(total)
                super().showPage()
            super().save()

        def _footer(self, total):
            w, _h = self._pagesize
            self.saveState()
            self.setFont(font, 7.5)
            self.setFillColor(colors.HexColor("#6b7280"))
            self.setStrokeColor(colors.HexColor("#e5e7eb"))
            self.line(14 * mm, 11 * mm, w - 14 * mm, 11 * mm)
            self.drawString(14 * mm, 7 * mm, footer_left)
            self.drawRightString(w - 14 * mm, 7 * mm,
                                 f"Sayfa {self._pageNumber} / {total}")
            self.restoreState()

    return _NumberedCanvas


def render_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (CondPageBreak, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from core.monthly_report import _register_fonts

    font, font_b = _register_fonts()
    NAVY = colors.HexColor("#232E6E")
    LIGHT = colors.HexColor("#f5f0e8")
    GREY = colors.HexColor("#e5e7eb")
    SUB = colors.HexColor("#ece7dc")
    RED, RED_BG = colors.HexColor("#b91c1c"), colors.HexColor("#fee2e2")
    AMBER, AMBER_BG = colors.HexColor("#92400e"), colors.HexColor("#fef3c7")
    INK = colors.HexColor("#374151")

    def ps(name, **kw):
        base = {"fontName": font, "fontSize": 7.6, "leading": 9.5, "textColor": INK}
        base.update(kw)
        return ParagraphStyle(name, **base)

    st = {
        "h1": ps("h1", fontName=font_b, fontSize=16, leading=19, textColor=NAVY, spaceAfter=2),
        "meta": ps("meta", fontSize=8, leading=11, textColor=colors.HexColor("#6b7280")),
        "h2": ps("h2", fontName=font_b, fontSize=12.5, leading=15, textColor=NAVY,
                 spaceBefore=6, spaceAfter=2),
        "h3": ps("h3", fontName=font_b, fontSize=9.5, leading=12, textColor=NAVY,
                 spaceBefore=8, spaceAfter=3),
        "cell": ps("cell"),
        "cell_b": ps("cell_b", fontName=font_b),
        "hcell": ps("hcell", fontName=font_b, textColor=colors.white),
        "bad": ps("bad", fontName=font_b, textColor=RED),
        "warn": ps("warn", fontName=font_b, textColor=AMBER),
        "empty": ps("empty", fontSize=9, leading=12, textColor=colors.HexColor("#9ca3af")),
    }
    P = lambda text, style="cell": Paragraph(escape(str(text)), st[style])   # noqa: E731

    def base_style(n_rows):
        return [
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("ROWBACKGROUNDS", (0, 1), (-1, n_rows), [colors.white, LIGHT]),
            ("GRID", (0, 0), (-1, -1), 0.4, GREY),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]

    story = [
        P(report["title"], "h1"),
        P(" · ".join(x for x in (
            f"{report['domain_label']} paneli" if report["domain_label"] else "",
            f"Oluşturulma: {report['generated_at']}",
            f"Oluşturan: {report['generated_by']}" if report["generated_by"] else "",
        ) if x), "meta"),
        P(f"Filtre: {report['filters'] or 'Tüm dolaplar'}", "meta"),
        Spacer(1, 5),
    ]

    if not report["count"]:
        story.append(P("Bu filtrelerle kayıt bulunamadı.", "empty"))
    else:
        # ── Özet ──────────────────────────────────────────────────────────
        story.append(P("Özet", "h3"))
        srows = _summary_rows(report)
        data = [[P(h, "hcell") for h in SUMMARY_COLUMNS]]
        for i, r in enumerate(srows):
            last = i == len(srows) - 1
            sty = "cell_b" if last else "cell"
            data.append([P(r[0], sty), P(r[1], sty), P(r[2], sty), P(_fmt(r[3]), sty),
                         P(r[4], "bad" if r[4] else sty),
                         P(r[5], "warn" if r[5] else sty),
                         P(r[6], "warn" if r[6] else sty)])
        t = Table(data, colWidths=[w * mm for w in (62, 22, 28, 26, 26, 28, 26)],
                  repeatRows=1, hAlign="LEFT")
        sstyle = base_style(len(srows) - 1)
        sstyle.append(("BACKGROUND", (0, len(srows)), (-1, len(srows)), SUB))
        t.setStyle(TableStyle(sstyle))
        story.append(t)
        story.append(Spacer(1, 3))
        story.append(P(
            f"Saklama bitişi: kırmızı = süresi dolmuş, turuncu = {DUE_SOON_DAYS} gün "
            f"içinde doluyor (yalnız dolaptaki kayıtlar).  Kalan / İlk = dolapta "
            f"kalan adet / ilk konulan adet.", "meta"))

        # ── Dolaplar ──────────────────────────────────────────────────────
        widths = [w * mm for w in _PDF_WIDTHS_MM]
        for ci, cab in enumerate(report["cabinets"]):
            # Her dolap ayrı sayfadan başlar (dolabın kapağına asılabilsin).
            # `keepWithNext` KULLANILMIYOR: başlığı uzun tabloyla KeepTogether'a
            # sarıp tabloyu komple sonraki sayfaya itiyor, başlık öksüz kalıyordu.
            # CondPageBreak yalnız başlığın altında birkaç satırlık yer arar.
            if ci:
                story.append(PageBreak())
            else:
                story.append(Spacer(1, 8))
                story.append(CondPageBreak(55 * mm))
            story.append(P(f"Dolap: {cab['brand']}", "h2"))
            story.append(P(
                f"{cab['samples']} kayıt · {_fmt(cab['quantity'])} adet kalan · "
                f"süresi dolan {cab['expired']} · süresi yaklaşan {cab['due_soon']} · "
                f"boyu belirsiz {cab['unsized']}", "meta"))
            for sh in cab["shelves"]:
                story.append(CondPageBreak(35 * mm))
                story.append(P(f"{sh['label']}  —  {sh['samples']} kayıt · "
                               f"{_fmt(sh['quantity'])} adet", "h3"))
                data = [[P(h, "hcell") for h in PDF_COLUMNS]]
                extra = []
                for ri, r in enumerate(sh["rows"], start=1):
                    until_style = {"expired": "bad", "due_soon": "warn"}.get(r["expiry"], "cell")
                    data.append([
                        P(r["slot"] or "—"), P(r["product"], "cell_b"),
                        P(r["size"] or "—", "warn" if r["unsized"] else "cell"),
                        P(r["lot"]), P(r["qty_label"]), P(r["produced"] or "—"),
                        P(r["until_label"], until_style),
                        P(r["last_check"] or "—", "bad" if r["last_check_bad"] else "cell"),
                        P(r["status"]), P(r["note"]),
                    ])
                    if r["expiry"] == "expired":
                        extra.append(("BACKGROUND", (6, ri), (6, ri), RED_BG))
                    elif r["expiry"] == "due_soon":
                        extra.append(("BACKGROUND", (6, ri), (6, ri), AMBER_BG))
                n = len(sh["rows"])
                data.append([P(f"{sh['label']} toplamı — {n} kayıt", "cell_b"), "", "", "",
                             P(f"{_fmt(sh['quantity'])} / {_fmt(sh['initial'])}", "cell_b"),
                             "", "", "", "", ""])
                t = Table(data, colWidths=widths, repeatRows=1)
                t.setStyle(TableStyle(base_style(n) + extra + [
                    ("SPAN", (0, n + 1), (3, n + 1)),
                    ("BACKGROUND", (0, n + 1), (-1, n + 1), SUB),
                ]))
                story.append(t)
            story.append(Spacer(1, 6))
            story.append(P(
                f"{cab['brand']} dolabı toplamı: {cab['samples']} kayıt · "
                f"{_fmt(cab['quantity'])} / {_fmt(cab['initial'])} adet (kalan / ilk)",
                "cell_b"))

    footer = f"Minerva 108 ERP · {TITLE} · {report['generated_at']}"
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=16 * mm,
        title="Şahit Numune Dolabı Raporu", author="Minerva 108 ERP")
    doc.build(story, canvasmaker=_numbered_canvas(font, footer))
    return buf.getvalue()


# ─── Excel ───────────────────────────────────────────────────────────────────

def _as_date(s):
    """'dd.mm.yyyy' / ISO metni gerçek tarihe — Excel'de süzme/sıralama doğru
    çalışsın (metin tarih "10.01" < "09.12" diye yanlış sıralanır)."""
    if not s:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(s)[:10], fmt).date()
        except ValueError:
            continue
    return None


def _append_as_text(ws, values) -> None:
    """`ws.append` + formül kalkanı.  openpyxl "=" ile başlayan HER metni formül
    sayar (data_type "f"): "=2 adet kırık" notu Excel'de "onarım" uyarısı açıp
    hücreyi siler, "=HYPERLINK(...)" ise canlı formül olarak yazılırdı.  Bu
    raporda hiçbir hücre formül değil; öyle işaretlenen düz metne çevrilir.
    quotePrefix = Excel'in kendi gizli "'" işareti — hücre sonradan düzenlense
    de formüle dönmez.  "+", "-", "@" dokunulmaz: xlsx'te metin hücresi olarak
    yazılır, Excel çalıştırmaz (risk yalnız CSV içe aktarımında)."""
    ws.append(values)
    row = ws.max_row
    for col, v in enumerate(values, start=1):
        if isinstance(v, str):
            c = ws.cell(row=row, column=col)
            if c.data_type == "f":
                c.data_type = "s"
                c.quotePrefix = True


def render_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy = PatternFill("solid", fgColor="232E6E")
    head_font = Font(bold=True, color="FFFFFF")
    red_fill, red_font = PatternFill("solid", fgColor="FEE2E2"), Font(bold=True, color="B91C1C")
    amb_fill, amb_font = PatternFill("solid", fgColor="FEF3C7"), Font(bold=True, color="92400E")
    bold = Font(bold=True)
    DATE_FMT = "DD.MM.YYYY"

    wb = Workbook()
    ws = wb.active
    ws.title = "Dolap"
    ws.append(list(XLSX_COLUMNS))
    for c in ws[1]:
        c.fill, c.font = navy, head_font
        c.alignment = Alignment(vertical="center")
    for cab in report["cabinets"]:
        for sh in cab["shelves"]:
            for r in sh["rows"]:
                _append_as_text(ws, [
                    r["brand"], r["shelf"] or None, r["slot"] or None, r["product"],
                    r["size"] or None, r["lot"], r["quantity"], r["initial"], r["unit"],
                    _as_date(r["produced"]), _as_date(r["until"]),
                    _as_date(r["last_check_on"]), r["last_check_result"] or None,
                    r["status"], r["note"] or None,
                ])
                row = ws.max_row
                for col in (10, 11, 12):
                    ws.cell(row=row, column=col).number_format = DATE_FMT
                until = ws.cell(row=row, column=11)
                if r["expiry"] == "expired":
                    until.fill, until.font = red_fill, red_font
                elif r["expiry"] == "due_soon":
                    until.fill, until.font = amb_fill, amb_font
                if r["unsized"]:
                    ws.cell(row=row, column=5).font = amb_font
                if r["last_check_bad"]:
                    ws.cell(row=row, column=13).font = red_font
    last_col = get_column_letter(len(XLSX_COLUMNS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{max(ws.max_row, 1)}"
    for i, w in enumerate(_XLSX_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.oddFooter.right.text = "Sayfa &P / &N"

    so = wb.create_sheet("Özet")
    for k, v in (("Rapor", report["title"]),
                 ("Panel", report["domain_label"]),
                 ("Oluşturulma", report["generated_at"]),
                 ("Oluşturan", report["generated_by"]),
                 ("Filtre", report["filters"] or "Tüm dolaplar")):
        _append_as_text(so, [k, v])
        so.cell(row=so.max_row, column=1).font = bold
    so.append([])
    so.append(list(SUMMARY_COLUMNS))
    for c in so[so.max_row]:
        c.fill, c.font = navy, head_font
    srows = _summary_rows(report)
    for i, r in enumerate(srows):
        _append_as_text(so, r)
        if i == len(srows) - 1:
            for c in so[so.max_row]:
                c.font = bold
    so.column_dimensions["A"].width = 18
    so.column_dimensions["B"].width = 60
    for col in "CDEFG":
        so.column_dimensions[col].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
