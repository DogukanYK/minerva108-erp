# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Üretim geçmişi raporu — "hangi üründen, hangi markadan, hangi ürün grubundan
ne kadar ürettik?" (Işık Hanım, 05.10.2026).

Kaynak `ProductionHistory` (domain-kapsamlı).  Sistemde üretim kaydı ancak
13.05.2026'dan beri var; daha eskisi yok — en eski kayıt tarihi rapora
dinamik olarak eklenir (UI notu sabit tarih yazmaz).

Tarihler TR-yerel GÜN olarak alınır (kullanıcı "13 Mayıs" der, UTC değil):
DB naive UTC tuttuğu için sınırlar `TR_OFFSET` kadar geri kaydırılır, aylık
kırılım ve ilk/son üretim tarihleri `to_tr()` ile TR gününe çevrilir.

Gruplama:
  • product  — `target_item_id` ile (aynı ürünün '(100ml)' / '100 ml' gibi
    farklı ad yazımları TEK satır); ad = ürün kartının BUGÜNKÜ adı, kart
    silinmişse en son kayıttaki ad.  Hedefsiz eski kayıtlar katlanmış ada göre.
  • brand    — `core.brands.product_brand` (MİNERVA-108 / Minerva108 → Minerva 108)
  • category — ad içindeki anahtar kelimelerden (CATEGORY_RULES, sıra önemli)

`history_quantities()` Satın Alma Planı'nın "adetleri geçmiş üretimden doldur"
düğmesi içindir: {target_item_id: toplam adet}.
"""
import datetime as _dt
import re
from typing import Dict, List, Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from core.brands import product_brand
from core.supplier_prices import normalize as fold
from database import Item, ProductionHistory, TR_OFFSET, to_tr

GROUPS = ("product", "brand", "category")
OTHER_CATEGORY = "Diğer"

# Ürün grubu anahtar kelime tablosu — TR-katlanmış küçük harf (ş→s, ı→i …).
# (etiket, anahtar kelimeler, hariç tutulan kelimeler).  İLK EŞLEŞEN KAZANIR,
# sıra bilinçli: güneş ürünleri vücut losyonundan önce (SPF'li losyon güneş
# ürünüdür), saç kremi/maskesi ve yüz/el/ayak/göz kremleri genel "Krem"den önce
# (genel kural EN SONDA — "Leke Karşıtı Krem", "Body Cream" oraya düşer).
# Anahtar kelime ancak bir kelime başında eşleşir ("jel kremi" içindeki
# "el krem" EL KREMİ sayılmaz).  Çok kelimeli anahtar kelimede kelimeler
# arasına isteğe bağlı TEK bir "bakım"/"care" girebilir: "Saç Bakım Kremi" =
# saç kremi, "Hand Care Cream" = el kremi (ara kelime serbest DEĞİL — "El
# yapımı krem" el kremi sayılmasın).
# Yeni grup = yeni satır; eşleşmeyen ürün "Diğer"e düşer.
CATEGORY_RULES = (
    ("Güneş Ürünleri", ("gunes", "spf", "sunscreen", "sun protection", "sun cream",
                        "after sun"), ()),
    ("Şampuan",        ("sampuan", "shampoo"), ()),
    ("Saç Kremi",      ("sac krem", "conditioner"), ()),
    ("Saç Maskesi",    ("sac maske", "hair mask"), ()),
    ("Duş Jeli",       ("dus jel", "shower gel", "shower jel", "body wash"), ()),
    ("Vücut Losyonu",  ("vucut losyon", "body lotion"),
                       ("spf", "gunes", "sunscreen", "sun protection")),
    ("Yüz Kremi",      ("yuz krem", "face cream", "day cream", "night cream",
                        "gece krem", "gunduz krem"), ()),
    ("El Kremi",       ("el krem", "hand cream"), ()),
    ("Ayak Kremi",     ("ayak krem", "ayakbakim", "ayak bakim", "foot cream",
                        "foot care"), ()),
    ("Göz Kremi",      ("goz krem", "goz cevresi", "goz alti", "eye cream",
                        "eye contour", "eye contuur"), ()),
    ("Serum",          ("serum",), ()),
    ("Peeling",        ("peeling", "scrub"), ()),
    ("Tonik",          ("tonik", "toner", "tonic"), ()),
    ("Krem",           ("krem", "cream"), ()),
)

_KW_CACHE: Dict[str, "re.Pattern"] = {}
_KW_FILLER = r" (?:bakim |care )?"          # "saç [bakım] kremi", "hand [care] cream"


def _kw(word: str):
    p = _KW_CACHE.get(word)
    if p is None:
        body = _KW_FILLER.join(re.escape(w) for w in word.split())
        p = _KW_CACHE[word] = re.compile(r"(?<![a-z0-9])" + body)
    return p


def product_category(name: str) -> str:
    """Ürün adı → ürün grubu etiketi (CATEGORY_RULES; eşleşmezse 'Diğer')."""
    s = fold(name)
    if not s:
        return OTHER_CATEGORY
    for label, words, excludes in CATEGORY_RULES:
        if any(_kw(w).search(s) for w in words) and not any(_kw(x).search(s) for x in excludes):
            return label
    return OTHER_CATEGORY


# ─── Tarih yardımcıları ─────────────────────────────────────────────────────

# Kabul edilen tarih penceresi.  Uç değerler (0001-01-01, 9999-12-31) TR_OFFSET
# kaydırması / +1 gün / varsayılan başlangıç hesabında OverflowError atıp 500'e
# düşüyordu; router bu pencerenin dışını 400 ile reddeder.
MIN_DATE = _dt.date(2000, 1, 1)
MAX_DATE = _dt.date(2100, 12, 31)


def date_in_range(d: Optional[_dt.date]) -> bool:
    return d is None or MIN_DATE <= d <= MAX_DATE


def utc_bounds(start: _dt.date, end: _dt.date):
    """TR-yerel [start, end] gün aralığı → naive UTC [lo, hi) sınırları."""
    lo = _dt.datetime.combine(start, _dt.time.min) - TR_OFFSET
    hi = _dt.datetime.combine(end + _dt.timedelta(days=1), _dt.time.min) - TR_OFFSET
    return lo, hi


def default_range(today: Optional[_dt.date] = None):
    """Son 12 ay: bir yıl önceki günün ertesinden bugüne (TR günü)."""
    from database import tr_now
    end = today or tr_now().date()
    try:
        start = end.replace(year=end.year - 1)
    except ValueError:                       # 29 Şubat
        start = end.replace(year=end.year - 1, day=28)
    return start + _dt.timedelta(days=1), end


def month_keys(start: _dt.date, end: _dt.date) -> List[str]:
    out, y, m = [], start.year, start.month
    while (y, m) <= (end.year, end.month):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def record_span(db: Session, domain: str):
    """Paneldeki EN ESKİ ve EN YENİ üretim kaydının TR günü (filtreden bağımsız)."""
    lo, hi = (db.query(func.min(ProductionHistory.produced_at),
                       func.max(ProductionHistory.produced_at))
              .filter(ProductionHistory.domain == domain).one())
    return (to_tr(lo).date() if lo else None), (to_tr(hi).date() if hi else None)


# ─── Sorgu ───────────────────────────────────────────────────────────────────

def _rows_in_range(db: Session, domain: str, start: _dt.date, end: _dt.date):
    lo, hi = utc_bounds(start, end)
    return (db.query(ProductionHistory)
            .filter(ProductionHistory.domain == domain,
                    ProductionHistory.cancelled_at.is_(None),     # iptal edilen üretim sayılmaz
                    ProductionHistory.produced_at >= lo,
                    ProductionHistory.produced_at < hi)
            .order_by(ProductionHistory.produced_at.asc(), ProductionHistory.id.asc())
            .all())


def history_quantities(db: Session, domain: str, start: _dt.date, end: _dt.date) -> Dict[int, float]:
    """{target_item_id: toplam üretilen} — hedefi olmayan eski kayıtlar ve
    iptal edilen üretimler dahil edilmez (Satın Alma Planı history-fill de bunu
    okur)."""
    lo, hi = utc_bounds(start, end)
    rows = (db.query(ProductionHistory.target_item_id,
                     func.sum(ProductionHistory.produced_quantity))
            .filter(ProductionHistory.domain == domain,
                    ProductionHistory.target_item_id.isnot(None),
                    ProductionHistory.cancelled_at.is_(None),
                    ProductionHistory.produced_at >= lo,
                    ProductionHistory.produced_at < hi)
            .group_by(ProductionHistory.target_item_id)
            .all())
    return {int(tid): round(float(q or 0.0), 4) for tid, q in rows}


def _parse_brands(brand) -> set:
    if not brand:
        return set()
    if isinstance(brand, str):
        brand = brand.split(",")
    return {fold(b) for b in brand if fold(b)}


def _product_groups(db: Session, rows) -> List[dict]:
    """Kayıtları ürün bazında birleştirir (target_item_id öncelikli)."""
    tids = {r.target_item_id for r in rows if r.target_item_id}
    current = {}
    if tids:
        for it in db.query(Item.id, Item.name).filter(Item.id.in_(tids)).all():
            current[it.id] = it.name
    prods: Dict[str, dict] = {}
    for r in rows:
        stored = (r.target_item_name or r.recipe_name or "").strip() or "—"
        key = f"i{r.target_item_id}" if r.target_item_id else f"n:{fold(stored)}"
        p = prods.get(key)
        if p is None:
            p = prods[key] = {"key": key, "target_item_id": r.target_item_id,
                              "stored_name": stored, "names": [], "recipe_name": r.recipe_name,
                              "total": 0.0, "count": 0, "first": None, "last": None,
                              "monthly": {}}
        if stored not in p["names"]:
            p["names"].append(stored)
        p["stored_name"] = stored                       # sıralı sorgu → en son kayıt
        if r.recipe_name:
            p["recipe_name"] = r.recipe_name
        q = float(r.produced_quantity or 0.0)
        d = to_tr(r.produced_at).date() if r.produced_at else None
        p["total"] += q
        p["count"] += 1
        if d:
            p["first"] = d if p["first"] is None or d < p["first"] else p["first"]
            p["last"] = d if p["last"] is None or d > p["last"] else p["last"]
            mk = f"{d.year:04d}-{d.month:02d}"
            p["monthly"][mk] = p["monthly"].get(mk, 0.0) + q
    out = []
    for p in prods.values():
        name = current.get(p["target_item_id"]) or p["stored_name"]
        p["name"] = name
        p["brand"] = product_brand(name, p["recipe_name"])
        p["category"] = product_category(name)
        out.append(p)
    return out


def _fmt_row(r: dict, months: List[str]) -> dict:
    return {
        **r,
        "total": round(r["total"], 4),
        "first": r["first"].isoformat() if r["first"] else None,
        "last": r["last"].isoformat() if r["last"] else None,
        "monthly": {m: round(r["monthly"].get(m, 0.0), 4) for m in months},
    }


def _sort_key(r: dict):
    return (-r["total"], fold(r["name"]))


def build_report(db: Session, domain: str, start: _dt.date, end: _dt.date, *,
                 brand=None, q: Optional[str] = None, group: str = "product") -> dict:
    """Üretim geçmişi raporu (JSON sözleşmesi — UI ve Excel aynı veriyi kullanır).

    `brand` virgüllü liste olabilir ('Minerva 108,Serenida'); karşılaştırma
    TR-katlanmış.  `q` ürün adı + eski ad yazımlarında katlanmış alt-dizgi.
    """
    if group not in GROUPS:
        raise ValueError("group product|brand|category olmalı")
    first_rec, last_rec = record_span(db, domain)
    # Aylık kırılım, sistemdeki ilk kayıttan önceki aylarla başlamaz (o aylarda
    # veri olamaz — "son 12 ay"ın ilk yarısı boş sütun olurdu).
    months = month_keys(max(start, first_rec) if first_rec else start, end)
    prods = _product_groups(db, _rows_in_range(db, domain, start, end))
    brands_available = sorted({p["brand"] for p in prods}, key=fold)

    want = _parse_brands(brand)
    if want:
        prods = [p for p in prods if fold(p["brand"]) in want]
    qf = fold(q)
    if qf:
        prods = [p for p in prods
                 if any(qf in fold(n) for n in [p["name"], *p["names"]])]

    if group == "product":
        rows = [{
            "key": p["key"], "name": p["name"], "brand": p["brand"], "category": p["category"],
            "target_item_id": p["target_item_id"],
            "names": [n for n in p["names"] if n != p["name"]],
            "products": 1, "total": p["total"], "count": p["count"],
            "first": p["first"], "last": p["last"], "monthly": p["monthly"],
        } for p in prods]
    else:
        field = "brand" if group == "brand" else "category"
        agg: Dict[str, dict] = {}
        for p in prods:
            k = p[field]
            a = agg.get(k)
            if a is None:
                a = agg[k] = {"key": k, "name": k,
                              "brand": k if group == "brand" else None,
                              "category": k if group == "category" else None,
                              "products": 0, "total": 0.0, "count": 0,
                              "first": None, "last": None, "monthly": {}}
            a["products"] += 1
            a["total"] += p["total"]
            a["count"] += p["count"]
            for b in ("first", "last"):
                v = p[b]
                if v is None:
                    continue
                if a[b] is None or (v < a[b] if b == "first" else v > a[b]):
                    a[b] = v
            for mk, v in p["monthly"].items():
                a["monthly"][mk] = a["monthly"].get(mk, 0.0) + v
        rows = list(agg.values())

    rows.sort(key=_sort_key)
    month_totals = {m: round(sum(r["monthly"].get(m, 0.0) for r in rows), 4) for m in months}
    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "group": group,
        "months": months,
        "rows": [_fmt_row(r, months) for r in rows],
        "brands": brands_available,
        "totals": {
            "quantity": round(sum(r["total"] for r in rows), 4),
            "count": sum(r["count"] for r in rows),
            "products": sum(r["products"] for r in rows),
            "monthly": month_totals,
        },
        "first_record": first_rec.isoformat() if first_rec else None,
        "last_record": last_rec.isoformat() if last_rec else None,
    }


# ─── Excel ───────────────────────────────────────────────────────────────────

_MONTHS_TR = ["", "Oca", "Şub", "Mar", "Nis", "May", "Haz",
              "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]
_GROUP_LABEL = {"product": "Ürün", "brand": "Marka", "category": "Ürün Grubu"}


def _tr_date(iso: Optional[str]) -> str:
    if not iso:
        return ""
    y, m, d = iso.split("-")
    return f"{d}.{m}.{y}"


def _month_label(mk: str) -> str:
    y, m = mk.split("-")
    return f"{_MONTHS_TR[int(m)]} {y}"


def since_note(first_record_iso: Optional[str]) -> str:
    """'Sistemde üretim kaydı 13.05.2026 tarihinden beri var.' — ek ünlü uyumu
    yıla göre değiştiği için ('2026'dan' / '2025'ten') ekli biçim kullanılmaz."""
    if not first_record_iso:
        return "Sistemde henüz üretim kaydı yok."
    return f"Sistemde üretim kaydı {_tr_date(first_record_iso)} tarihinden beri var."


def build_workbook(rep: dict) -> bytes:
    """'Özet' + 'Aylık' sayfalı .xlsx."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY = PatternFill("solid", fgColor="232E6E")
    HF = Font(bold=True, color="FFFFFF")
    TITLE = Font(bold=True, size=13, color="232E6E")
    NOTE = Font(italic=True, color="6B7280")
    BLD = Font(bold=True)
    thin = Border(*([Side(style="thin", color="DDDDDD")] * 4))
    group = rep["group"]
    glabel = _GROUP_LABEL[group]
    period = f"{_tr_date(rep['start'])} – {_tr_date(rep['end'])}"
    since = since_note(rep.get("first_record"))

    def header(ws, row, headers):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row, ci, h)
            c.fill = NAVY; c.font = HF
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # ── Özet ────────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = f"Minerva 108 — Üretim Geçmişi ({glabel} bazında)"
    ws["A1"].font = TITLE
    ws["A2"] = f"Dönem: {period}  ·  {since}"
    ws["A2"].font = NOTE
    if group == "product":
        headers = ["#", "Ürün", "Marka", "Ürün Grubu", "Toplam Üretim", "Üretim Sayısı",
                   "İlk Üretim", "Son Üretim"]
        w = [5, 48, 14, 16, 15, 13, 12, 12]
    else:
        headers = ["#", glabel, "Ürün Sayısı", "Toplam Üretim", "Üretim Sayısı",
                   "İlk Üretim", "Son Üretim"]
        w = [5, 30, 12, 15, 13, 12, 12]
    header(ws, 4, headers)
    r = 5
    for i, row in enumerate(rep["rows"], 1):
        if group == "product":
            vals = [i, row["name"], row["brand"], row["category"], row["total"], row["count"],
                    _tr_date(row["first"]), _tr_date(row["last"])]
        else:
            vals = [i, row["name"], row["products"], row["total"], row["count"],
                    _tr_date(row["first"]), _tr_date(row["last"])]
        for ci, v in enumerate(vals, 1):
            ws.cell(r, ci, v).border = thin
        r += 1
    t = rep["totals"]
    tot_col = 5 if group == "product" else 4
    ws.cell(r, 2, "TOPLAM").font = BLD
    ws.cell(r, tot_col, t["quantity"]).font = BLD
    ws.cell(r, tot_col + 1, t["count"]).font = BLD
    ws.cell(r + 2, 1, "Not: miktarlar reçete çıktı biriminde (çoğunlukla adet). Aynı ürünün "
                      "farklı ad yazımları ürün kartı üzerinden tek satırda birleştirilir.").font = NOTE
    ws.freeze_panes = "A5"
    widths(ws, w)

    # ── Aylık ───────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Aylık")
    ws["A1"] = f"Aylık üretim — {glabel} bazında · {period}"
    ws["A1"].font = TITLE
    lead = [glabel] + (["Marka"] if group == "product" else [])
    months = rep["months"]
    header(ws, 3, lead + [_month_label(m) for m in months] + ["Toplam"])
    r = 4
    for row in rep["rows"]:
        vals = [row["name"]] + ([row["brand"]] if group == "product" else [])
        vals += [row["monthly"].get(m) or None for m in months] + [row["total"]]
        for ci, v in enumerate(vals, 1):
            ws.cell(r, ci, v).border = thin
        r += 1
    ws.cell(r, 1, "TOPLAM").font = BLD
    off = len(lead)
    for j, m in enumerate(months, 1):
        ws.cell(r, off + j, t["monthly"].get(m) or None).font = BLD
    ws.cell(r, off + len(months) + 1, t["quantity"]).font = BLD
    ws.freeze_panes = ws.cell(4, off + 1).coordinate
    widths(ws, [44] + ([14] if group == "product" else []) + [10] * len(months) + [12])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
