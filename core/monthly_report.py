# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Aylık Detaylı Sistem Raporu — veri toplama + PDF/Excel üretimi.

Akış:
  gather_report_data(db, year, month) -> dict   bir ayın tüm sistem verisi
  render_pdf(data)   -> bytes                   reportlab ile resmi PDF
  render_excel(data) -> bytes                   openpyxl ile çok-sayfalı xlsx

Rapor SuperAdmin'e özeldir (bkz. routers/system.py) — üretim hattı, harcanan
malzeme, kullanıcı aktivitesi, yönetim olayları ve sistem/teknik bölümlerini
içerir.
"""
import calendar
import os
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from sqlalchemy import text

from database import (
    to_tr, tr_now,
    ProductionHistory, Transaction, Item, AdminAuditLog, StockSnapshot, SystemEvent,
)


# Otomatik üretilen aylık raporların saklandığı klasör.  Sunucuda
# /var/www/minerva/system_reports, lokalde proje kökü/system_reports.
SYSTEM_REPORT_DIR = Path(os.environ.get(
    "MINERVA_SYSTEM_REPORT_DIR",
    str(Path(__file__).resolve().parent.parent / "system_reports"),
))
SYSTEM_REPORT_DIR.mkdir(parents=True, exist_ok=True)
try:
    os.chmod(SYSTEM_REPORT_DIR, 0o700)   # sadece servis kullanıcısı
except OSError:
    pass


# ─── Küçük yardımcılar ──────────────────────────────────────────────────────

def _month_bounds(year: int, month: int):
    """Ayın [başlangıç, bitiş) yarı-açık datetime aralığı (naive UTC)."""
    start = datetime(year, month, 1)
    end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
    return start, end


def _humanize_bytes(n) -> str:
    n = float(n or 0)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024.0 or unit == "TB":
            return f"{int(n)} B" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0
    return f"{n:.1f} TB"


def _fmt_seconds(s: int) -> str:
    s = int(s)
    if s < 60:
        return f"~{s} sn"
    m, rem = divmod(s, 60)
    return f"~{m} dk {rem} sn" if rem else f"~{m} dk"


def report_filename(year: int, month: int, ext: str) -> str:
    """Standart rapor dosya adı — minerva_rapor_2026-04.pdf gibi."""
    return f"minerva_rapor_{year:04d}-{month:02d}.{ext}"


# ─── Veri toplama ───────────────────────────────────────────────────────────

def gather_report_data(db, year: int, month: int) -> dict:
    """Bir ayın tüm sistem verisini tek bir dict'te toplar."""
    start, end = _month_bounds(year, month)
    period = f"{month:02d}.{year}"

    # ── Üretim hattı ──────────────────────────────────────────────────────
    prods = (
        db.query(ProductionHistory)
        .filter(ProductionHistory.produced_at >= start,
                ProductionHistory.produced_at < end)
        .order_by(ProductionHistory.produced_at)
        .all()
    )
    production = [{
        "date":     to_tr(p.produced_at).strftime("%d.%m.%Y %H:%M") if p.produced_at else "—",
        "product":  p.target_item_name or p.recipe_name or "—",
        "recipe":   p.recipe_name or "—",
        "quantity": round(p.produced_quantity or 0.0, 2),
        "lot":      p.lot_number or "—",
        "by":       p.produced_by or "—",
    } for p in prods]

    # ── Stok hareketleri (transactions) ───────────────────────────────────
    txs = (
        db.query(Transaction)
        .filter(Transaction.timestamp >= start, Transaction.timestamp < end)
        .order_by(Transaction.timestamp)
        .all()
    )
    item_ids = {t.item_id for t in txs if t.item_id}
    items = {}
    if item_ids:
        for it in db.query(Item).filter(Item.id.in_(item_ids)).all():
            items[it.id] = it

    materials, receiving = {}, {}
    adjustments = []
    output_count = input_count = 0
    for t in txs:
        it = items.get(t.item_id)
        nm  = it.name if it else (f"#{t.item_id}" if t.item_id else "—")
        cat = (it.category if it else "") or ""
        un  = (it.unit if it else "") or ""
        qty = abs(t.quantity or 0.0)
        if t.transaction_type == "Output":
            output_count += 1
            d = materials.setdefault(t.item_id, {"name": nm, "category": cat,
                                                 "unit": un, "qty": 0.0, "count": 0})
            d["qty"] += qty
            d["count"] += 1
        elif t.transaction_type == "Input":
            input_count += 1
            d = receiving.setdefault(t.item_id, {"name": nm, "category": cat,
                                                 "unit": un, "qty": 0.0, "count": 0})
            d["qty"] += qty
            d["count"] += 1
        elif t.transaction_type == "Adjustment":
            adjustments.append({
                "date": to_tr(t.timestamp).strftime("%d.%m.%Y %H:%M") if t.timestamp else "—",
                "item": nm,
                "quantity": round(t.quantity or 0.0, 2),
                "by": t.performed_by or "—",
                "notes": (t.notes or "")[:90],
            })

    materials_list = sorted(
        [{**v, "qty": round(v["qty"], 2)} for v in materials.values()],
        key=lambda x: x["qty"], reverse=True)
    receiving_list = sorted(
        [{**v, "qty": round(v["qty"], 2)} for v in receiving.values()],
        key=lambda x: x["qty"], reverse=True)

    # ── Yönetim olayları (audit) ──────────────────────────────────────────
    audits = (
        db.query(AdminAuditLog)
        .filter(AdminAuditLog.timestamp >= start, AdminAuditLog.timestamp < end)
        .order_by(AdminAuditLog.timestamp)
        .all()
    )
    audit_events = [{
        "date":   to_tr(a.timestamp).strftime("%d.%m.%Y %H:%M") if a.timestamp else "—",
        "actor":  a.actor_name or "—",
        "action": a.action or "",
        "target": a.target_name or (a.target_type or "—"),
    } for a in audits]

    # ── Kullanıcı aktivitesi (kim ne yapmış) ──────────────────────────────
    ua = {}

    def _u(name):
        key = name or "—"
        return ua.setdefault(key, {"name": key, "tx": 0, "prod": 0, "audit": 0})

    for t in txs:
        _u(t.performed_by)["tx"] += 1
    for p in prods:
        _u(p.produced_by)["prod"] += 1
    for a in audits:
        _u(a.actor_name)["audit"] += 1
    users_list = sorted(
        [{**v, "total": v["tx"] + v["prod"] + v["audit"]} for v in ua.values()],
        key=lambda x: x["total"], reverse=True)

    # ── Sistem / teknik ───────────────────────────────────────────────────
    restarts = (
        db.query(SystemEvent)
        .filter(SystemEvent.event_type == "app_start",
                SystemEvent.created_at >= start, SystemEvent.created_at < end)
        .order_by(SystemEvent.created_at)
        .all()
    )
    restart_times = [
        to_tr(e.created_at).strftime("%d.%m.%Y %H:%M") for e in restarts if e.created_at
    ]

    try:
        db_size_b = db.execute(
            text("SELECT pg_database_size(current_database())")
        ).scalar() or 0
    except Exception:
        db_size_b = 0

    snap_count = (
        db.query(StockSnapshot)
        .filter(StockSnapshot.year == year, StockSnapshot.month == month)
        .count()
    )

    backup_count, last_backup, month_backup_count = 0, "—", 0
    try:
        from routers.backup import _list_backups
        bks = _list_backups()
        backup_count = len(bks)
        if bks:
            last_backup = bks[0]["created_at"]
        s_ts = calendar.timegm(start.timetuple())
        e_ts = calendar.timegm(end.timetuple())
        month_backup_count = sum(1 for b in bks if s_ts <= b["timestamp"] < e_ts)
    except Exception:
        pass

    system = {
        "restart_count":      len(restarts),
        "restart_times":      restart_times,
        "est_downtime":       _fmt_seconds(len(restarts) * 4),
        "db_size":            _humanize_bytes(db_size_b),
        "snapshot_count":     snap_count,
        "backup_count":       backup_count,
        "month_backup_count": month_backup_count,
        "last_backup":        last_backup,
    }

    summary = {
        "period":               period,
        "production_count":      len(production),
        "production_total_qty":  round(sum(p["quantity"] for p in production), 2),
        "transaction_count":     len(txs),
        "output_count":          output_count,
        "input_count":           input_count,
        "material_lines":        len(materials_list),
        "active_users":          len([u for u in users_list if u["name"] != "—"]),
        "restart_count":         len(restarts),
        "audit_count":           len(audits),
    }

    return {
        "year": year, "month": month, "period": period,
        "generated_at": tr_now().strftime("%d.%m.%Y %H:%M"),
        "summary":      summary,
        "production":   production,
        "materials":    materials_list,
        "receiving":    receiving_list,
        "adjustments":  adjustments,
        "users":        users_list,
        "audit_events": audit_events,
        "system":       system,
    }


# ─── PDF üretimi (reportlab) ────────────────────────────────────────────────

# Türkçe karakter (ş/ğ/ı/ç/ö/ü) içeren bir TTF gerekir — yerleşik Helvetica
# bunları gösteremez.  Aday yollar: prod (Linux/DejaVu) ve Mac (Arial).
_FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
]


def _register_fonts():
    """Türkçe-uyumlu bir font çiftini reportlab'a kaydeder; (normal, bold)
    font adlarını döner.  Hiçbir aday bulunamazsa Helvetica'ya düşer."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for reg, bold in _FONT_CANDIDATES:
        if os.path.isfile(reg):
            try:
                pdfmetrics.registerFont(TTFont("RPT", reg))
                pdfmetrics.registerFont(
                    TTFont("RPT-B", bold if os.path.isfile(bold) else reg))
                return "RPT", "RPT-B"
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


def render_pdf(data: dict) -> bytes:
    """Rapor dict'inden resmi görünümlü, çok-sayfalı bir PDF üretir."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)

    font, font_b = _register_fonts()
    NAVY  = colors.HexColor("#232E6E")
    LIGHT = colors.HexColor("#f5f0e8")
    GREY  = colors.HexColor("#e5e7eb")

    st_h1   = ParagraphStyle("h1", fontName=font_b, fontSize=17, textColor=NAVY,
                             spaceAfter=2, leading=20)
    st_meta = ParagraphStyle("meta", fontName=font, fontSize=8,
                             textColor=colors.HexColor("#6b7280"), leading=11)
    st_h2   = ParagraphStyle("h2", fontName=font_b, fontSize=11.5, textColor=NAVY,
                             spaceBefore=15, spaceAfter=5, leading=14)
    st_cell = ParagraphStyle("cell", fontName=font, fontSize=7.6,
                             textColor=colors.HexColor("#374151"), leading=9.5)
    st_hcell = ParagraphStyle("hcell", fontName=font_b, fontSize=7.6,
                              textColor=colors.white, leading=9.5)
    st_empty = ParagraphStyle("empty", fontName=font, fontSize=8.5,
                              textColor=colors.HexColor("#9ca3af"), leading=12)

    story = []

    def _table(headers, rows, col_widths):
        head = [Paragraph(escape(str(h)), st_hcell) for h in headers]
        data_rows = [[Paragraph(escape(str(c)), st_cell) for c in r] for r in rows]
        t = Table([head] + data_rows, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",     (0, 0), (-1, 0), NAVY),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("GRID",           (0, 0), (-1, -1), 0.4, GREY),
            ("VALIGN",         (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING",     (0, 0), (-1, -1), 3.5),
            ("BOTTOMPADDING",  (0, 0), (-1, -1), 3.5),
            ("LEFTPADDING",    (0, 0), (-1, -1), 5),
            ("RIGHTPADDING",   (0, 0), (-1, -1), 5),
        ]))
        return t

    def _section(title, headers, rows, col_widths, empty_msg):
        story.append(Paragraph(escape(title), st_h2))
        if rows:
            story.append(_table(headers, rows, col_widths))
        else:
            story.append(Paragraph(escape(empty_msg), st_empty))

    W = 174.0  # kullanılabilir içerik genişliği (mm)

    # ── Başlık ────────────────────────────────────────────────────────────
    s = data["summary"]
    story.append(Paragraph("Minerva 108 — Aylık Detaylı Sistem Raporu", st_h1))
    story.append(Paragraph(
        f"Dönem: {data['period']}  ·  Üretildi: {data['generated_at']}  "
        f"·  GİZLİ — yalnızca yönetim erişimi", st_meta))
    story.append(Spacer(1, 4))

    # ── 1) Özet ───────────────────────────────────────────────────────────
    summary_rows = [
        ["Üretim sayısı",         str(s["production_count"])],
        ["Üretilen toplam miktar", str(s["production_total_qty"])],
        ["Toplam stok hareketi",  f"{s['transaction_count']} "
                                  f"(çıkış {s['output_count']} · giriş {s['input_count']})"],
        ["Harcanan malzeme çeşidi", str(s["material_lines"])],
        ["Aktif kullanıcı",       str(s["active_users"])],
        ["Yönetim olayı",         str(s["audit_count"])],
        ["Uygulama restart sayısı", str(s["restart_count"])],
    ]
    _section("1 · Özet", ["Gösterge", "Değer"], summary_rows,
             [W * 0.55 * mm, W * 0.45 * mm], "—")

    # ── 2) Üretim hattı ───────────────────────────────────────────────────
    _section(
        "2 · Üretim Hattı",
        ["Tarih", "Ürün", "Reçete", "Miktar", "Lot", "Üreten"],
        [[p["date"], p["product"], p["recipe"], p["quantity"], p["lot"], p["by"]]
         for p in data["production"]],
        [W*0.15*mm, W*0.25*mm, W*0.22*mm, W*0.10*mm, W*0.15*mm, W*0.13*mm],
        "Bu dönemde üretim kaydı yok.")

    # ── 3) Harcanan malzemeler ────────────────────────────────────────────
    _section(
        "3 · Harcanan Malzemeler (stoktan çıkan)",
        ["Malzeme", "Kategori", "Toplam Miktar", "Birim", "İşlem Sayısı"],
        [[m["name"], m["category"], m["qty"], m["unit"], m["count"]]
         for m in data["materials"]],
        [W*0.38*mm, W*0.18*mm, W*0.18*mm, W*0.12*mm, W*0.14*mm],
        "Bu dönemde stoktan çıkış (tüketim) yok.")

    # ── 4) Mal kabul / stok girişi ────────────────────────────────────────
    _section(
        "4 · Mal Kabul / Stok Girişi",
        ["Malzeme", "Kategori", "Toplam Miktar", "Birim", "İşlem Sayısı"],
        [[r["name"], r["category"], r["qty"], r["unit"], r["count"]]
         for r in data["receiving"]],
        [W*0.38*mm, W*0.18*mm, W*0.18*mm, W*0.12*mm, W*0.14*mm],
        "Bu dönemde stok girişi yok.")

    # ── 5) Stok düzeltmeleri ──────────────────────────────────────────────
    _section(
        "5 · Stok Düzeltmeleri",
        ["Tarih", "Ürün", "Miktar", "Kullanıcı", "Not"],
        [[a["date"], a["item"], a["quantity"], a["by"], a["notes"]]
         for a in data["adjustments"]],
        [W*0.16*mm, W*0.27*mm, W*0.12*mm, W*0.17*mm, W*0.28*mm],
        "Bu dönemde stok düzeltmesi yok.")

    # ── 6) Kullanıcı aktivitesi ───────────────────────────────────────────
    _section(
        "6 · Kullanıcı Aktivitesi (kim ne yapmış)",
        ["Kullanıcı", "Stok Hareketi", "Üretim", "Yönetim Olayı", "Toplam"],
        [[u["name"], u["tx"], u["prod"], u["audit"], u["total"]]
         for u in data["users"]],
        [W*0.32*mm, W*0.20*mm, W*0.16*mm, W*0.18*mm, W*0.14*mm],
        "Bu dönemde kullanıcı aktivitesi yok.")

    # ── 7) Yönetim olayları ───────────────────────────────────────────────
    _section(
        "7 · Yönetim Olayları (audit log)",
        ["Tarih", "Kullanıcı", "İşlem", "Hedef"],
        [[a["date"], a["actor"], a["action"], a["target"]]
         for a in data["audit_events"]],
        [W*0.17*mm, W*0.22*mm, W*0.31*mm, W*0.30*mm],
        "Bu dönemde yönetim olayı yok.")

    # ── 8) Sistem / teknik ────────────────────────────────────────────────
    sysd = data["system"]
    restart_detail = (", ".join(sysd["restart_times"])
                      if sysd["restart_times"] else "Restart kaydı yok")
    sys_rows = [
        ["Uygulama restart sayısı", str(sysd["restart_count"])],
        ["Restart zamanları",       restart_detail],
        ["Tahmini toplam kesinti",  f"{sysd['est_downtime']} "
                                    f"(her restart ~4 sn varsayımı)"],
        ["Bu ay alınan yedek",      str(sysd["month_backup_count"])],
        ["Toplam yedek (mevcut)",   str(sysd["backup_count"])],
        ["Son yedek",               sysd["last_backup"]],
        ["Aylık stok snapshot",     f"{sysd['snapshot_count']} kalem donduruldu"],
        ["Veritabanı boyutu",       sysd["db_size"]],
    ]
    _section("8 · Sistem / Teknik", ["Gösterge", "Değer"], sys_rows,
             [W*0.32*mm, W*0.68*mm], "—")

    story.append(Spacer(1, 14))
    story.append(Paragraph(
        "Bu rapor Minerva 108 ERP tarafından otomatik üretilmiştir.  "
        "Downtime değerleri restart kayıtlarından tahmindir; saniye hassasiyetli "
        "erişilemezlik için /health endpoint'ini izleyen bir dış servis gerekir.",
        st_meta))

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18*mm, rightMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm,
        title=f"Minerva 108 — Aylik Sistem Raporu {data['period']}",
        author="Minerva 108 ERP")
    doc.build(story)
    return buf.getvalue()


# ─── Excel üretimi (openpyxl) ───────────────────────────────────────────────

def render_excel(data: dict) -> bytes:
    """Rapor dict'inden her bölümü ayrı sayfa olan bir .xlsx üretir."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)  # varsayılan boş sayfayı kaldır

    navy_fill = PatternFill("solid", fgColor="232E6E")
    head_font = Font(bold=True, color="FFFFFF")

    def _sheet(title, headers, rows):
        ws = wb.create_sheet(title=title[:31])  # Excel sayfa adı ≤ 31 karakter
        ws.append(headers)
        for c in ws[1]:
            c.fill = navy_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="left", vertical="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            cells = [str(h)] + [str(r[i - 1]) for r in rows]
            width = min(max(len(x) for x in cells) + 3, 55)
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    s = data["summary"]
    sysd = data["system"]

    _sheet("Özet", ["Gösterge", "Değer"], [
        ["Dönem",                    data["period"]],
        ["Rapor üretim tarihi",      data["generated_at"]],
        ["Üretim sayısı",            s["production_count"]],
        ["Üretilen toplam miktar",   s["production_total_qty"]],
        ["Toplam stok hareketi",     s["transaction_count"]],
        ["  — çıkış (Output)",       s["output_count"]],
        ["  — giriş (Input)",        s["input_count"]],
        ["Harcanan malzeme çeşidi",  s["material_lines"]],
        ["Aktif kullanıcı",          s["active_users"]],
        ["Yönetim olayı",            s["audit_count"]],
        ["Uygulama restart sayısı",  s["restart_count"]],
    ])

    _sheet("Üretim", ["Tarih", "Ürün", "Reçete", "Miktar", "Lot", "Üreten"],
           [[p["date"], p["product"], p["recipe"], p["quantity"], p["lot"], p["by"]]
            for p in data["production"]])

    _sheet("Harcanan Malzeme",
           ["Malzeme", "Kategori", "Toplam Miktar", "Birim", "İşlem Sayısı"],
           [[m["name"], m["category"], m["qty"], m["unit"], m["count"]]
            for m in data["materials"]])

    _sheet("Mal Kabul",
           ["Malzeme", "Kategori", "Toplam Miktar", "Birim", "İşlem Sayısı"],
           [[r["name"], r["category"], r["qty"], r["unit"], r["count"]]
            for r in data["receiving"]])

    _sheet("Stok Düzeltme",
           ["Tarih", "Ürün", "Miktar", "Kullanıcı", "Not"],
           [[a["date"], a["item"], a["quantity"], a["by"], a["notes"]]
            for a in data["adjustments"]])

    _sheet("Kullanıcı Aktivitesi",
           ["Kullanıcı", "Stok Hareketi", "Üretim", "Yönetim Olayı", "Toplam"],
           [[u["name"], u["tx"], u["prod"], u["audit"], u["total"]]
            for u in data["users"]])

    _sheet("Yönetim Olayları",
           ["Tarih", "Kullanıcı", "İşlem", "Hedef"],
           [[a["date"], a["actor"], a["action"], a["target"]]
            for a in data["audit_events"]])

    _sheet("Sistem", ["Gösterge", "Değer"], [
        ["Uygulama restart sayısı", sysd["restart_count"]],
        ["Restart zamanları",       ", ".join(sysd["restart_times"]) or "—"],
        ["Tahmini toplam kesinti",  sysd["est_downtime"]],
        ["Bu ay alınan yedek",      sysd["month_backup_count"]],
        ["Toplam yedek (mevcut)",   sysd["backup_count"]],
        ["Son yedek",               sysd["last_backup"]],
        ["Aylık stok snapshot",     sysd["snapshot_count"]],
        ["Veritabanı boyutu",       sysd["db_size"]],
    ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Diske kaydetme (scheduler için) ────────────────────────────────────────

def save_monthly_report(db, year: int, month: int) -> list:
    """Belirtilen ay için PDF + Excel üretip SYSTEM_REPORT_DIR'a yazar.

    Otomatik aylık scheduler job'u bunu çağırır.  Yazılan dosya adlarını döner.
    """
    data = gather_report_data(db, year, month)
    written = []
    for ext, renderer in (("pdf", render_pdf), ("xlsx", render_excel)):
        path = SYSTEM_REPORT_DIR / report_filename(year, month, ext)
        path.write_bytes(renderer(data))
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass
        written.append(path.name)
    return written
