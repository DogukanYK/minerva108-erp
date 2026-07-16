# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
İçindekiler Raporu — bitmiş ürünlerin reçete bileşimini dosyalanabilir rapora çevirir.

  list_report_products(db, domain) -> [dict]   seçim paneli verisi (marka + reçete durumu)
  assemble(db, item_ids, domain)   -> dict     rapor verisi (Hammadde/Ambalaj ayrımıyla)
  build_workbook(data)             -> bytes    4 sayfalık, A4-yazdırılabilir .xlsx
  render_pdf(data)                 -> bytes    antetli A4 PDF (ürün başına blok)

Excel sayfaları HER ZAMAN 4 tanedir (boşsa yalnız başlık satırı) ve veri sayfaları
%100 düzdür (birleşik hücre yok) — makinece işlenebilirlik için sayfa adları ve
sütun konumları deterministiktir. Rapor metadatası baskı üstbilgi/altbilgisindedir.

Maliyet BİLİNÇLİ olarak yoktur — bu rapor içerik dokümantasyonudur; finance
gating gerektirmez, `recipes:view` yeter (bileşim yine de ticari sırdır).
"""
from datetime import datetime

from sqlalchemy.orm import Session

from database import Item, Recipe, to_tr
from core.production_sim import brand_of
from core.items_report import tr_key
from core.supplier_prices import normalize as _tr_fold


def _canonical_brands(names: list) -> dict:
    """Marka yazım varyantlarını birleştirir: TR-katlanmış anahtar → en yaygın yazım.

    Ürün adları elle girildiğinden aynı marka 'Minerva' / 'MİNERVA' gibi
    varyantlarla yaşıyor (prod'da gerçek durum). Chip/gruplama tek markada
    toplanmalı; görüntü olarak en çok kullanılan yazım kazanır.
    """
    counts: dict = {}
    for nm in names:
        b = brand_of(nm)
        counts.setdefault(_tr_fold(b), {}).setdefault(b, 0)
        counts[_tr_fold(b)][b] += 1
    return {k: max(v, key=v.get) for k, v in counts.items()}


def _brand(name: str, canon: dict) -> str:
    return canon.get(_tr_fold(brand_of(name)), brand_of(name))


def report_filename(ext: str = "xlsx") -> str:
    """Saf ASCII dosya adı — Content-Disposition header'ına doğrudan girebilir."""
    return f"icindekiler_raporu_{datetime.utcnow():%Y%m%d}.{ext}"


def _abstract_parent_ids(db: Session) -> set:
    """Varyasyon ebeveyni (soyut) ürün id'leri — reçete taşıyamazlar."""
    return {pid for (pid,) in
            db.query(Item.parent_id).filter(Item.parent_id.isnot(None)).distinct()
            if pid}


def _recipe_map(db: Session, item_ids, domain: str) -> dict:
    """target_item_id -> aktif Recipe (aynı ürüne birden çoksa en yenisi kazanır)."""
    q = (db.query(Recipe)
         .filter(Recipe.is_active == True,          # noqa: E712
                 Recipe.domain == domain,
                 Recipe.target_item_id.in_(item_ids))
         .order_by(Recipe.id.asc()))
    return {r.target_item_id: r for r in q}


def _finished_items(db: Session, domain: str, item_ids=None):
    """Aktif paneldeki somut (soyut-ebeveyn olmayan) bitmiş ürünler."""
    parents = _abstract_parent_ids(db)
    q = (db.query(Item)
         .filter(Item.is_active == True,            # noqa: E712
                 Item.domain == domain,
                 Item.category == "Bitmiş Ürün"))
    if item_ids is not None:
        q = q.filter(Item.id.in_(item_ids))
    return [it for it in q.all() if it.id not in parents]


def list_report_products(db: Session, domain: str) -> list:
    """Seçim paneli: marka + reçete durumu ile tüm somut bitmiş ürünler."""
    items = _finished_items(db, domain)
    rmap = _recipe_map(db, [it.id for it in items], domain) if items else {}
    canon = _canonical_brands([it.name for it in items])
    parent_names = {}
    for it in items:
        if it.parent_id and it.parent_id not in parent_names:
            p = db.query(Item).filter(Item.id == it.parent_id).first()
            parent_names[it.parent_id] = p.name if p else ""
    out = []
    for it in items:
        rec = rmap.get(it.id)
        out.append({
            "item_id":        it.id,
            "name":           it.name,
            "name_tr":        it.name_tr or "",
            "sku":            it.sku or "",
            "brand":          _brand(it.name, canon),
            "unit":           it.unit or "adet",
            "barcode":        it.barcode or "",
            "variation_name": it.variation_name or "",
            "parent_name":    parent_names.get(it.parent_id, "") if it.parent_id else "",
            "has_recipe":     rec is not None,
            "recipe_id":      rec.id if rec else None,
        })
    out.sort(key=lambda x: (tr_key(x["brand"]), tr_key(x["name"])))
    return out


def assemble(db: Session, item_ids: list, domain: str) -> dict:
    """Seçilen ürünlerin içindekiler verisini toplar.

    Domain dışı / bitmiş-ürün olmayan / pasif id'ler sessizce düşer — domain
    izolasyonunun garantisi budur. Reçetesi olmayan ürünler `recipeless`e gider.
    """
    items = _finished_items(db, domain, item_ids=item_ids)
    rmap = _recipe_map(db, [it.id for it in items], domain) if items else {}
    canon = _canonical_brands([it.name for it in items])

    products, raw_rows, pkg_rows, recipeless = [], [], [], []
    for it in items:
        rec = rmap.get(it.id)
        base = {
            "brand":          _brand(it.name, canon),
            "product_name":   it.name,
            "product_name_tr": it.name_tr or "",
            "product_sku":    it.sku or "",
            "variation":      it.variation_name or "",
        }
        if not rec:
            recipeless.append({**base,
                               "barcode": it.barcode or "",
                               "unit":    it.unit or "adet"})
            continue

        waste_pct = float(rec.waste_percentage or 0.0)
        waste_factor = 1.0 + waste_pct / 100.0
        raws, pkgs = [], []
        for ing in rec.ingredients:
            item = ing.item
            if not item:
                continue
            entry = {
                **base,
                "recipe_id":   rec.id,
                "ing_name":    item.name,
                "ing_name_tr": item.name_tr or "",
                "ing_sku":     item.sku or "",
                "net":         float(ing.quantity or 0.0),
                "unit":        ing.unit or item.unit or "",
            }
            if item.category == "Ambalaj":   # fire muaf — brüt == net
                pkgs.append({**entry,
                             "pkg_type":    item.pkg_type or "",
                             "language":    item.language or "",
                             "label_group": item.label_group or ""})
            else:
                entry["phase"] = ing.phase or ""
                entry["waste_pct"] = waste_pct
                entry["gross"] = round(entry["net"] * waste_factor, 6)
                raws.append(entry)

        raw_total = sum(r["net"] for r in raws)
        for r in raws:
            r["pct"] = round(r["net"] / raw_total * 100.0, 4) if raw_total else 0.0

        products.append({
            **base,
            "barcode":        it.barcode or "",
            "recipe_id":      rec.id,
            "output_quantity": float(rec.output_quantity or 1.0),
            "output_unit":    rec.output_unit or "adet",
            "waste_pct":      waste_pct,
            "raw_count":      len(raws),
            "pkg_count":      len(pkgs),
            "raw_total":      round(raw_total, 6),
        })
        raw_rows.extend(raws)
        pkg_rows.extend(pkgs)

    products.sort(key=lambda x: (tr_key(x["brand"]), tr_key(x["product_name"])))
    recipeless.sort(key=lambda x: (tr_key(x["brand"]), tr_key(x["product_name"])))
    raw_rows.sort(key=lambda x: (tr_key(x["brand"]), tr_key(x["product_name"]),
                                 x["phase"] or "~", tr_key(x["ing_name"])))
    pkg_rows.sort(key=lambda x: (tr_key(x["brand"]), tr_key(x["product_name"]),
                                 x["pkg_type"] or "~", tr_key(x["ing_name"])))

    return {
        "generated_at": to_tr(datetime.utcnow()).strftime("%d.%m.%Y %H:%M"),
        "domain":       domain,
        "products":     products,
        "raw_rows":     raw_rows,
        "pkg_rows":     pkg_rows,
        "recipeless":   recipeless,
        "summary": {
            "product_count":    len(products),
            "recipeless_count": len(recipeless),
            "raw_line_count":   len(raw_rows),
            "pkg_line_count":   len(pkg_rows),
        },
    }


# ─── Excel ───────────────────────────────────────────────────────────────────

def _setup_a4(ws, *, landscape: bool = True, title: str = "") -> None:
    """Sayfayı A4 baskıya hazırlar — repo'daki ilk Excel print-setup kullanımı.

    fitToWidth tek başına yetmez: sheet_properties.pageSetUpPr.fitToPage
    işaretlenmezse Excel fitTo* değerlerini yok sayar.
    """
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.properties import PageSetupProperties

    ws.page_setup.paperSize = ws.PAPERSIZE_A4              # == 9
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0                          # boyuna serbest sayfa
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"                            # başlık her baskı sayfasında
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.7, bottom=0.6,
                                  header=0.3, footer=0.3)
    ws.oddHeader.center.text = f"Minerva 108 — İçindekiler Raporu · {title}"
    ws.oddHeader.center.size = 8
    ws.oddFooter.left.text = "&D &T"
    ws.oddFooter.right.text = "Sayfa &P / &N"


_SHEET_RAW = "Hammaddeler"
_SHEET_PKG = "Ambalaj-Paketleme"
_SHEET_SUMMARY = "Özet"
_SHEET_RECIPELESS = "Reçetesiz Ürünler"

_HDR_SUMMARY = ["Marka", "Ürün (EN)", "Ürün (TR)", "SKU", "Varyasyon", "Barkod",
                "Reçete ID", "Reçete", "Batch Çıktısı", "Çıktı Birimi", "Fire %",
                "Hammadde Kalemi", "Ambalaj Kalemi", "Toplam Hammadde (Net)"]
_HDR_RAW = ["Marka", "Ürün (EN)", "Ürün (TR)", "Ürün SKU", "Varyasyon", "Reçete ID",
            "Faz", "Hammadde (EN)", "Hammadde (TR)", "Hammadde SKU", "Net Miktar",
            "Fire %", "Brüt Miktar", "Birim", "Bileşim %"]
_HDR_PKG = ["Marka", "Ürün (EN)", "Ürün (TR)", "Ürün SKU", "Varyasyon", "Reçete ID",
            "Bileşen (EN)", "Bileşen (TR)", "Bileşen SKU", "Ambalaj Tipi",
            "Etiket Dili", "Etiket Grubu", "Miktar", "Birim"]
_HDR_RECIPELESS = ["Marka", "Ürün (EN)", "Ürün (TR)", "SKU", "Varyasyon", "Barkod", "Birim"]


def build_workbook(data: dict) -> bytes:
    """4 sayfalık, A4-yazdırılabilir içindekiler Excel'i (bkz. modül docstring'i)."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "Minerva 108 — İçindekiler Raporu"
    wb.properties.creator = "Minerva 108 ERP"

    navy_fill = PatternFill("solid", fgColor="232E6E")
    head_font = Font(bold=True, color="FFFFFF")

    def _sheet(title, headers, rows, *, landscape=True, num_formats=None):
        ws = wb.create_sheet(title=title[:31])  # Excel sayfa adı ≤ 31 karakter
        ws.append(headers)
        for c in ws[1]:
            c.fill = navy_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="left", vertical="center")
        for r in rows:
            ws.append(r)
        if num_formats:
            for col, fmt in num_formats.items():
                for row_cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    row_cells[0].number_format = fmt
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            cells = [str(h)] + [str(r[i - 1]) for r in rows]
            width = min(max(len(x) for x in cells) + 3, 55)
            ws.column_dimensions[get_column_letter(i)].width = width
        _setup_a4(ws, landscape=landscape, title=title)
        return ws

    _sheet(_SHEET_SUMMARY, _HDR_SUMMARY,
           [[p["brand"], p["product_name"], p["product_name_tr"], p["product_sku"],
             p["variation"], p["barcode"], p["recipe_id"], "Var",
             p["output_quantity"], p["output_unit"], p["waste_pct"],
             p["raw_count"], p["pkg_count"], p["raw_total"]]
            for p in data["products"]],
           num_formats={9: "0.####", 11: "0.00", 14: "0.####"})

    _sheet(_SHEET_RAW, _HDR_RAW,
           [[r["brand"], r["product_name"], r["product_name_tr"], r["product_sku"],
             r["variation"], r["recipe_id"], r["phase"], r["ing_name"],
             r["ing_name_tr"], r["ing_sku"], r["net"], r["waste_pct"], r["gross"],
             r["unit"], r["pct"]]
            for r in data["raw_rows"]],
           num_formats={11: "0.####", 12: "0.00", 13: "0.####", 15: "0.00"})

    _sheet(_SHEET_PKG, _HDR_PKG,
           [[r["brand"], r["product_name"], r["product_name_tr"], r["product_sku"],
             r["variation"], r["recipe_id"], r["ing_name"], r["ing_name_tr"],
             r["ing_sku"], r["pkg_type"], r["language"], r["label_group"],
             r["net"], r["unit"]]
            for r in data["pkg_rows"]],
           num_formats={13: "0.####"})

    _sheet(_SHEET_RECIPELESS, _HDR_RECIPELESS,
           [[r["brand"], r["product_name"], r["product_name_tr"], r["product_sku"],
             r["variation"], r["barcode"], r["unit"]]
            for r in data["recipeless"]],
           landscape=False)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── PDF (antetli A4) ────────────────────────────────────────────────────────

def _report_story(data: dict, s: float = 1.0):
    from html import escape

    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, Spacer

    from core.delivery_note import _fmt
    from core.shipment_note import _grid, _styles

    st = _styles(s)
    W = 174.0

    def pct(*vals):
        return [W * v / 100.0 * mm for v in vals]

    sm = data["summary"]
    meta = (f"Tarih: {escape(str(data.get('generated_at') or ''))}"
            f" · {sm['product_count']} ürün"
            + (f" · {sm['recipeless_count']} reçetesiz" if sm["recipeless_count"] else ""))
    story = [Paragraph("İÇİNDEKİLER RAPORU", st["h1"]),
             Paragraph(meta, st["meta"]),
             Spacer(1, 6 * s)]

    if not data["products"] and not data["recipeless"]:
        story.append(Paragraph("Seçimle eşleşen ürün yok.", st["cell"]))
        return story

    raw_by, pkg_by = {}, {}
    for r in data["raw_rows"]:
        raw_by.setdefault(r["recipe_id"], []).append(r)
    for r in data["pkg_rows"]:
        pkg_by.setdefault(r["recipe_id"], []).append(r)

    for p in data["products"]:
        title = f"{p['brand']} — {p['product_name']}"
        if p["product_name_tr"]:
            title += f" / {p['product_name_tr']}"
        head_meta = (f"Barkod: {p['barcode'] or '—'} · Varyasyon: {p['variation'] or '—'}"
                     f" · Batch: {_fmt(p['output_quantity'])} {p['output_unit']}"
                     f" · Fire: %{_fmt(p['waste_pct'])}")
        # Başlık + meta sayfa dibinde öksüz kalmasın; tablolar serbest bölünür.
        story.append(KeepTogether([Paragraph(escape(title), st["h2"]),
                                   Paragraph(escape(head_meta), st["meta"])]))

        raws = raw_by.get(p["recipe_id"], [])
        if raws:
            story.append(Spacer(1, 3 * s))
            story.append(_grid(
                st, ["Faz", "Hammadde", "Türkçe Ad", "Net", "Brüt", "Birim", "Bileşim %"],
                [[r["phase"] or "—", r["ing_name"], r["ing_name_tr"] or "—",
                  _fmt(r["net"]), _fmt(r["gross"]), r["unit"], f"{r['pct']:.2f}"]
                 for r in raws],
                pct(8, 30, 24, 10, 10, 8, 10), s))

        pkgs = pkg_by.get(p["recipe_id"], [])
        if pkgs:
            story.append(Spacer(1, 4 * s))
            story.append(_grid(
                st, ["Ambalaj / Bileşen", "Tip", "Dil", "Miktar", "Birim"],
                [[r["ing_name"], r["pkg_type"] or "—", r["language"] or "—",
                  _fmt(r["net"]), r["unit"]]
                 for r in pkgs],
                pct(46, 16, 10, 14, 14), s))
        story.append(Spacer(1, 8 * s))

    if data["recipeless"]:
        story.append(Paragraph("Reçetesiz Ürünler", st["h2"]))
        story.append(Paragraph("Bu ürünlerin sistemde aktif reçetesi yok — içerik dökümü verilemedi.",
                               st["meta"]))
        story.append(Spacer(1, 3 * s))
        story.append(_grid(
            st, ["Marka", "Ürün", "Türkçe Ad", "Varyasyon", "Barkod"],
            [[r["brand"], r["product_name"], r["product_name_tr"] or "—",
              r["variation"] or "—", r["barcode"] or "—"]
             for r in data["recipeless"]],
            pct(14, 30, 26, 14, 16), s))
    return story


def render_pdf(data: dict) -> bytes:
    from reportlab.lib.units import mm

    from core.delivery_note import merge_letterhead, render_autofit

    content = render_autofit(
        lambda s: _report_story(data, s),
        margins=(20 * mm, 20 * mm, 48 * mm, 40 * mm),   # HANDOFF #2 — değiştirme
        doc_kwargs={"title": "İçindekiler Raporu", "author": "Minerva 108"})
    return merge_letterhead(content)
