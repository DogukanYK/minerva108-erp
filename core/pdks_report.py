# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
PDKS aylık puantaj Excel üreticisi (openpyxl).

Girdi: routers/pdks._report_data() çıktısı (serileştirilmiş — tüm süreler hem
dakika hem '8 sa 15 dk' etiketiyle gelir).  Sayfa 1 "Özet" = personel başına
ay toplamları; sonra her personele bir sayfa (gün satırları + toplam satırı).
Stil kalıbı core/monthly_report.render_excel ile aynı (lacivert başlık).
"""


import re

from core.pdks import fmt_minutes

# Excel sayfa adında yasak karakterler (openpyxl ValueError fırlatır)
_SHEET_FORBIDDEN = re.compile(r"[\\/?*\[\]:]")


def _sheet_title(name: str, used: set) -> str:
    """Sayfa adını sanitize et: yasak karakterleri temizle, 31 karaktere kırp,
    çakışırsa numara ekle.  used kümesine ekleyip döndürür."""
    t = (_SHEET_FORBIDDEN.sub(" ", name or "").strip() or "Personel")[:31]
    base, n = t, 2
    while t in used:
        t = f"{base[:27]} ({n})"
        n += 1
    used.add(t)
    return t


def puantaj_filename(year: int, month: int) -> str:
    return f"puantaj_{year}-{month:02d}.xlsx"


def render_puantaj_excel(data: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    navy_fill = PatternFill("solid", fgColor="232E6E")
    head_font = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    warn_fill = PatternFill("solid", fgColor="FFE8CC")   # eksik çıkış (turuncu)
    miss_fill = PatternFill("solid", fgColor="FFD6D6")   # devamsız (kırmızı)
    off_fill = PatternFill("solid", fgColor="EFEFEF")    # tatil/hafta tatili (gri)
    leave_fill = PatternFill("solid", fgColor="D6F0EF")  # izinli (turkuaz)

    def _style_header(ws):
        for c in ws[1]:
            c.fill = navy_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="left", vertical="center")
        ws.freeze_panes = "A2"

    def _autofit(ws, headers, rows):
        for i, h in enumerate(headers, 1):
            cells = [str(h)] + [str(r[i - 1]) for r in rows]
            ws.column_dimensions[get_column_letter(i)].width = \
                min(max(len(x) for x in cells) + 3, 45)

    # ── Sayfa 1: Özet ────────────────────────────────────────────────────────
    headers = ["Personel", "Ünvan", "Çalışma", "Fazla Mesai", "Eksik",
               "Devamsızlık", "Eksik Çıkış", "İzin Günleri"]
    rows = []
    for item in data["employees"]:
        emp, tot = item["employee"], item["month"]["totals"]
        name = emp["full_name"] + ("" if emp.get("is_active", True) else " (ayrıldı)")
        rows.append([
            name, emp.get("title") or "—",
            tot["toplam_calisma_label"], tot["toplam_fazla_mesai_label"],
            tot["toplam_eksik_label"],
            tot["devamsizlik_gun"], tot["eksik_cikis_sayisi"],
            tot["izin_gunleri_label"],
        ])
    ws = wb.create_sheet(title="Özet")
    ws.append([f"Puantaj — {data['month_label']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(headers)
    for c in ws[3]:
        c.fill = navy_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="left", vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A4"
    _autofit(ws, headers, rows)

    # ── Personel sayfaları ───────────────────────────────────────────────────
    day_headers = ["Tarih", "Gün", "Durum", "Giriş", "Çıkış", "Çalışma",
                   "Fazla Mesai", "Eksik", "Not"]
    used_titles = {"Özet"}
    for item in data["employees"]:
        emp, month = item["employee"], item["month"]
        ws = wb.create_sheet(title=_sheet_title(emp["full_name"], used_titles))
        ws.append(day_headers)
        _style_header(ws)
        day_rows = []
        for d in month["days"]:
            notes = []
            if d["holiday_name"]:
                notes.append(d["holiday_name"])
            if d["missing_checkout"]:
                notes.append("Çıkış eksik — toplam dışı")
            if d["break_deducted"]:
                notes.append("Mola düşüldü")
            row = [
                d["date_label"], d["weekday_label"], d["status_label"],
                d["first_in"] or "—", d["last_out"] or "—",
                d["worked_label"], d["overtime_label"],
                fmt_minutes(d["missing_minutes"]),
                "; ".join(notes),
            ]
            day_rows.append(row)
            ws.append(row)
            fill = None
            if d["missing_checkout"]:
                fill = warn_fill
            elif d["status"] == "devamsiz":
                fill = miss_fill
            elif d["status"] == "izinli":
                fill = leave_fill
            elif d["status"] in ("hafta_tatili", "resmi_tatil"):
                fill = off_fill
            if fill:
                for c in ws[ws.max_row]:
                    c.fill = fill
        tot = month["totals"]
        ws.append([])
        total_row = ["TOPLAM", "", "", "", "",
                     tot["toplam_calisma_label"], tot["toplam_fazla_mesai_label"],
                     tot["toplam_eksik_label"],
                     f"Devamsız: {tot['devamsizlik_gun']} gün · "
                     f"{tot['izin_gunleri_label']}"]
        ws.append(total_row)
        for c in ws[ws.max_row]:
            c.font = bold
        _autofit(ws, day_headers, day_rows + [total_row])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
