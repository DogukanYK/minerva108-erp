"""
Tedarikçi fiyat listesi — Işık Hanım'ın "Stok Son Durum" Excel'ini sisteme alıp
satın alma raporunu otomatik dolduran çekirdek mantık.

Akış:
  • parse_stok_son_durum(data)  → Excel'i (onun sütun düzeniyle) satırlara çevirir
  • import_prices(db, rows, …)  → malzeme/tedarikçi eşler, SupplierPrice'a upsert eder
                                  (yalnız Excel kaynaklı satırları değiştirir; elle
                                  girilen `manual` satır korunur)
  • price_unit_ok / serialize_price → elle fiyat uçları (routers/reports.py)
  • prices_for_items(db, …)     → rapor için malzeme başına (ucuzdan) tedarikçi listesi
                                  (pasif firmanın teklifi sonda, `inactive` işaretli;
                                  para birimi karışıksa bugünün kuruyla sıralı,
                                  `comparable` = "Tedarikçi-1 en uygun" denebilir mi)
  • price_cell_format(offer)    → Excel fiyat hücresinin sayı biçimi ("#,##0.00## €/kg")
  • firm_keys / inactive_firm_check → "aynı firma" ve "pasif firma" kuralları
                                  (satın alma planı `core.purchase_pricing` de kullanır)
  • price_per_purchase_unit(…)  → fiyatı malzemenin alım birimine (kg / l / adet) çevirir
  • net_unit_price / vat_label  → KDV dahil fiyatın neti + "KDV dahil" etiketi

Fiyat temeli (2026-10): Stok Son Durum listesi **USD / kg**. Eskiden para birimi
kolonu model varsayılanıyla 'TRY' yazılıyor, birim hiç tutulmuyordu — panel her
fiyatı "₺" gösteriyordu. Artık her satır `currency` + `price_unit` (kg|l|adet) +
`source` / `source_label` / `quoted_at` taşır; adet birimli malzemenin fiyatı
DAİMA adet başınadır (`default_price_unit`). Eski satırlar `database.
_backfill_supplier_price_units()` ile bir kez USD/kg(adet) olarak işaretlendi.

KDV (08.10.2026): `vat_included` (NULL = bilinmiyor, False = hariç, True =
dahil) + `vat_rate` (yüzde).  Karşılaştırma ve tutar NET fiyatla yapılır
(`net_unit_price`); KDV dahil satır her gösterimde "KDV dahil" etiketi taşır.
İki alan da boşken her hesap eskisiyle birebir aynıdır.

Beklenen Excel düzeni (Işık Hanım'ın dosyası):
  A Malzeme · B Kategori · C Birim · D Toplam Gereken · E Mevcut Stok · F ALINACAK
  G Alınabilecek Miktar(1) · H TEDARİKÇİ-1 · I Birim Fiyat(1)
  J TEDARİKÇİ-2 · K Alınabilecek Miktar(2) · L Birim Fiyat(2)
  M TEDARİKÇİ-3 · N Alınabilecek Miktar(3) · O Birim Fiyat(3)
(Tedarikçi-1'in paketi adından ÖNCE, 2 ve 3'ün paketi adından SONRA — onun düzeni.)
"""
from datetime import date
from io import BytesIO
from typing import List, Optional, Tuple

from sqlalchemy import or_
from sqlalchemy.orm import Session

from database import Item, Supplier, SupplierPrice

# 0-indeksli sütun konumları — Işık Hanım'ın STOK SON DURUM düzeni
COL_MATERIAL = 0
_SUPPLIER_BLOCKS = [
    {"name": 7,  "package": 6,  "price": 8},   # Tedarikçi-1: G paket, H ad, I fiyat
    {"name": 9,  "package": 10, "price": 11},  # Tedarikçi-2: J ad, K paket, L fiyat
    {"name": 12, "package": 13, "price": 14},  # Tedarikçi-3: M ad, N paket, O fiyat
]

# Fiyat temeli sözlüğü — form/uç doğrulaması ve rapor aynı listeyi kullanır
CURRENCIES = ("USD", "EUR", "TRY")
PRICE_UNITS = ("kg", "l", "adet")
# İçe aktarmada seçilebilen LİSTE birimi yalnız ağırlık/hacim temelidir: adet
# birimli malzeme zaten otomatik 'adet' alır, 'adet' seçeneği ise listedeki her
# g/kg/ml/lt kalemine 'adet' yazıp `price_per_purchase_unit`'te hepsini
# "uyuşmuyor"a düşürürdü (satın alma planında hammadde fiyatı kalmazdı).
LIST_PRICE_UNITS = ("kg", "l")
SOURCE_STOK_SON_DURUM = "stok_son_durum"
# Lab'ın elle girdiği (ya da düzenlediği) satır — POST/PUT /api/supplier-prices.
# İçe aktarma bu satırları SİLMEZ ve aynı firmanın teklifini üstüne yazmaz.
SOURCE_MANUAL = "manual"
SOURCE_LABEL_MAX = 120
VAT_RATE_MAX = 100.0

# Malzeme birimi aileleri (Item.unit: g/kg/ml/lt/adet/kutu/rulo …).  Ağırlık ve
# hacim dışındaki her birim SAYILAN birimdir → fiyatı adet başına.
_MASS_UNITS = {"g", "gr", "kg"}
_VOLUME_UNITS = {"ml", "l", "lt"}
LITRE_KG_NOTE = "1 l ≈ 1 kg kabulüyle"

# Türkçe harf katlama (eşleştirme için) — ç/ğ/ı/İ/ö/ş/ü → ascii
_TR_FOLD = str.maketrans({
    "ç": "c", "Ç": "c", "ğ": "g", "Ğ": "g", "ı": "i", "İ": "i",
    "ö": "o", "Ö": "o", "ş": "s", "Ş": "s", "ü": "u", "Ü": "u", "I": "i",
})


def clean_name(v) -> Optional[str]:
    """Görüntü adı: trim + iç boşlukları teke indir.  Boşsa None."""
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s or None


def normalize(v) -> str:
    """Eşleştirme anahtarı: Türkçe-katlanmış, küçük harf, boşluk-normalize."""
    s = clean_name(v)
    if not s:
        return ""
    return " ".join(s.translate(_TR_FOLD).lower().split())


def _num(v) -> Optional[float]:
    """Hücreyi sayıya çevir; Türkçe ondalık virgülünü destekler.  Yoksa None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_stok_son_durum(data: bytes) -> List[dict]:
    """Yüklenen .xlsx baytlarını satırlara çevirir.

    Döner: [{ "material": str, "suppliers": [{name, package, price}, …] }, …]
    Başlık satırı (1. satır) atlanır; malzeme adı boş satırlar atlanır.
    """
    wb = openpyxl_load(data)
    ws = wb[wb.sheetnames[0]]
    rows: List[dict] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r:
            continue
        material = clean_name(r[COL_MATERIAL] if len(r) > COL_MATERIAL else None)
        if not material:
            continue
        suppliers = []
        for blk in _SUPPLIER_BLOCKS:
            name = clean_name(r[blk["name"]] if len(r) > blk["name"] else None)
            if not name:
                continue
            suppliers.append({
                "name": name,
                "package": _num(r[blk["package"]] if len(r) > blk["package"] else None),
                "price": _num(r[blk["price"]] if len(r) > blk["price"] else None),
            })
        rows.append({"material": material, "suppliers": suppliers})
    return rows


def openpyxl_load(data: bytes):
    import openpyxl
    return openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)


def _unit_family(item_unit) -> str:
    """'mass' | 'volume' | 'count' — boş/bilinmeyen birim sayılan birimdir."""
    u = (item_unit or "").strip().lower()
    if u in _MASS_UNITS:
        return "mass"
    if u in _VOLUME_UNITS:
        return "volume"
    return "count"


def price_unit_ok(item_unit, price_unit) -> bool:
    """Elle girilen fiyat biriminin malzemeye uygunluğu (birim ailesi).

    Sayılan (adet/kutu/rulo…) malzeme → yalnız 'adet'; kütle/hacim malzemesi
    → 'kg' ya da 'l' (çapraz kg↔l `price_per_purchase_unit`'te "1 l ≈ 1 kg"
    notuyla kullanılır).  Adet ↔ kg/l hiçbir zaman çevrilemez — öyle bir satır
    satın alma planında "fiyat yok" olurdu."""
    pu = (price_unit or "").strip().lower()
    if _unit_family(item_unit) == "count":
        return pu == "adet"
    return pu in LIST_PRICE_UNITS


def default_price_unit(item_unit, list_unit: str = "kg") -> str:
    """Bir malzemenin fiyat satırına yazılacak `price_unit`.

    Adet (ve kutu/rulo gibi sayılan) birimli malzemede fiyat DAİMA adet
    başınadır; ağırlık/hacim malzemesinde listenin birimi (`list_unit`,
    Stok Son Durum için 'kg') geçerlidir.  İçe aktarma ve tek seferlik geri
    doldurma bu tek kuralı paylaşır."""
    return "adet" if _unit_family(item_unit) == "count" else list_unit


def firm_keys(supplier_id, supplier_name) -> set:
    """Fiyat satırının firma anahtarları: {("id", supplier_id), ("name", normalize(ad))}.

    "Elle girilen fiyat, aynı firmanın Excel teklifinin yerine geçer" kuralının
    TEK tanımı — `import_prices` (Excel satırını yazmaz), `prices_for_items`
    (listeye almaz) ve satın alma planı `_price_row` (teklif saymaz) aynı
    anahtarı kullanır; iki rapor aynı malzeme için farklı fiyat göstermesin."""
    keys = set()
    if supplier_id is not None:
        keys.add(("id", supplier_id))
    if normalize(supplier_name):
        keys.add(("name", normalize(supplier_name)))
    return keys


def inactive_firm_check(db: Session, domain: str, rows=()):
    """Fiyat satırı → firması PASİF mi (yumuşak silinmiş) — `check(sp) -> bool`.

    Kural satın alma planıyla TEK: `core.purchase_pricing.FirmActivity`
    (ölçüt kart değil firma; mükerrer kartın pasif ikizi ve birleştirilmiş
    kartın adı firmayı pasif yapmaz; aktif karta bağlı satır asla pasif
    değildir).  `rows`: kontrol edilecek SupplierPrice satırları (ad anahtar
    uzayına girer)."""
    from core.purchase_pricing import FirmActivity, SupplierIndex   # döngüsel import: geç yükle
    sups = (db.query(Supplier.id, Supplier.name, Supplier.is_active, Supplier.merged_into_id)
            .filter(Supplier.domain == domain).all())
    ix = SupplierIndex([n for _, n, _, _ in sups] + [r.supplier_name for r in rows if r.supplier_name])
    firms = FirmActivity({sid: (n or "", a is not False, m) for sid, n, a, m in sups}, ix)
    return lambda sp: firms.is_inactive(sp.supplier_id, sp.supplier_name)


def import_prices(db: Session, rows: List[dict], domain: str, *,
                  currency: str = "USD", price_unit: str = "kg",
                  source_label: Optional[str] = None,
                  quoted_at: Optional[date] = None) -> dict:
    """Satırları SupplierPrice'a yazar (malzeme başına ESKİ Excel satırlarını değiştirir).

    Malzeme `Item.name` ile, tedarikçi AKTİF `Supplier.name` ile (Türkçe-
    katlanmış; aynı anahtarda en eski kart) eşlenir — pasife alınmış mükerrer
    kart (TATLİDİLİMLER / TATLIDİLİMLER) fiyatı üstüne çekmesin.  İstisna:
    BİRLEŞTİRİLMİŞ pasif kartın adı (`merged_into_id`) kazananın takma adıdır
    → satır kazanana bağlanır (aktif kart adı önce gelir).  Eşleşmeyen
    tedarikçi serbest metin olarak saklanır (supplier_id NULL).
    Silme YALNIZ `source` stok_son_durum ya da boş (eski) satırlarda; lab'ın
    elle girdiği `manual` satır korunur (`kept_manual`) ve aynı firmanın Excel
    teklifi o malzemeye YAZILMAZ (`skipped_manual`) — aynı firma iki kez
    görünmesin, elle girilen değer ezilmesin.  Lab fiyat notlarının satırları
    (lab_notu | proforma | fatura | siparis) silinmez ve engellemez.
    Her satır `currency` + `price_unit` + kaynak bilgisiyle yazılır; adet
    birimli malzemeler `price_unit`'ten bağımsız olarak 'adet' alır.
    `price_unit` listenin ağırlık/hacim temelidir → yalnız `LIST_PRICE_UNITS`
    (kg | l); 'adet' ValueError.
    Döner: özet (kaç malzeme güncellendi, kaç fiyat satırı, eşleşmeyenler).
    """
    currency = (currency or "").strip().upper()
    price_unit = (price_unit or "").strip().lower()
    if currency not in CURRENCIES:
        raise ValueError(f"Geçersiz para birimi: {currency!r}")
    if price_unit not in LIST_PRICE_UNITS:
        raise ValueError(f"Geçersiz liste fiyat birimi: {price_unit!r} (kg ya da l)")
    source_label = clean_name(source_label)
    if source_label:
        source_label = source_label[:SOURCE_LABEL_MAX]

    item_by_norm = {}
    for it in db.query(Item).filter(Item.domain == domain, Item.is_active == True).all():
        item_by_norm.setdefault(normalize(it.name), it)
    from core.suppliers import live_card_id
    sups = db.query(Supplier).filter(Supplier.domain == domain).order_by(Supplier.id).all()
    sup_by_norm = {}
    for sp in sups:
        if sp.is_active is not False:
            sup_by_norm.setdefault(normalize(sp.name), sp)
    # Birleştirilmiş (pasif) kartın adı kazananın takma adı: listede eski
    # yazım kalsa da satır kazanana bağlanır (yoksa serbest metin kalıp
    # yalnız pasif kaybedenle eşleşir → "pasif tedarikçi" sayılırdı).
    # Aktif kartın adı her zaman önce gelir.
    by_id = {sp.id: sp for sp in sups}
    cards = {sp.id: (sp.is_active is not False, sp.merged_into_id) for sp in sups}
    for sp in sups:
        if sp.is_active is False and sp.merged_into_id is not None:
            live = live_card_id(cards, sp.id)
            if live is not None:
                sup_by_norm.setdefault(normalize(sp.name), by_id[live])

    items_updated = 0
    prices_inserted = 0
    kept_manual = 0
    skipped_manual = 0
    unmatched_materials: List[str] = []
    unmatched_suppliers = set()

    for row in rows:
        it = item_by_norm.get(normalize(row["material"]))
        if not it:
            unmatched_materials.append(row["material"])
            continue
        # Replace semantics: bu malzemenin eski EXCEL satırlarını sil (elle
        # girilenler kalır).  Kaynağı boş satır içe aktarma öncesi döneme ait.
        db.query(SupplierPrice).filter(
            SupplierPrice.item_id == it.id, SupplierPrice.domain == domain,
            or_(SupplierPrice.source == SOURCE_STOK_SON_DURUM, SupplierPrice.source.is_(None)),
        ).delete(synchronize_session=False)
        # Yalnız ELLE satır aynı firmanın Excel teklifini engeller; lab fiyat
        # notlarından gelen satırlar (lab_notu / proforma / fatura / siparis)
        # da korunur ama engellemez — ikisi de listede kalır (08.10.2026 öncesi
        # bu kaynaklar yoktu: eski veride davranış aynı).
        manual_keys = set()
        for m in (db.query(SupplierPrice)
                  .filter(SupplierPrice.item_id == it.id, SupplierPrice.domain == domain,
                          SupplierPrice.source == SOURCE_MANUAL).all()):
            kept_manual += 1
            manual_keys |= firm_keys(m.supplier_id, m.supplier_name)
        for s in row["suppliers"]:
            sup = sup_by_norm.get(normalize(s["name"]))
            if firm_keys(sup.id if sup is not None else None, s["name"]) & manual_keys:
                skipped_manual += 1
                continue
            if not sup:
                unmatched_suppliers.add(s["name"])
            db.add(SupplierPrice(
                item_id=it.id,
                supplier_id=(sup.id if sup else None),
                supplier_name=s["name"],
                package_size=s["package"],
                unit_price=s["price"],
                currency=currency,
                price_unit=default_price_unit(it.unit, price_unit),
                source=SOURCE_STOK_SON_DURUM,
                source_label=source_label,
                quoted_at=quoted_at,
                domain=domain,
            ))
            prices_inserted += 1
        items_updated += 1

    db.commit()
    return {
        "items_updated": items_updated,
        "prices_inserted": prices_inserted,
        "unmatched_materials": unmatched_materials,
        "unmatched_suppliers": sorted(unmatched_suppliers),
        "matched_suppliers": prices_inserted - len(unmatched_suppliers),
        "kept_manual": kept_manual,
        "skipped_manual": skipped_manual,
        "currency": currency,
        "price_unit": price_unit,
        "source_label": source_label,
        "quoted_at": quoted_at.isoformat() if quoted_at else None,
    }


def prices_for_items(db: Session, item_ids: List[int], domain: str, limit: int = 3,
                     rates=None) -> dict:
    """Rapor için: { item_id: [ {supplier_name, package_size, unit_price,
    currency, price_unit, source_label, quoted_at, inactive, vat_included,
    vat_rate, comparable}, … ] }.

    Her malzeme için fiyatı OLAN tedarikçiler ucuzdan pahalıya (KDV dahil
    satır net fiyatıyla — `net_unit_price`); fiyatı olmayanlar (None) sona.
    En çok `limit` tedarikçi.  Bir firmanın elle girilmiş
    (`manual`) satırı varsa aynı firmanın Excel satırı listeye girmez —
    yoksa aynı firma iki sütun kaplar ve 3. tedarikçi kesilirdi (`firm_keys`).
    PASİF firmanın teklifi (`inactive_firm_check`) en uygun sayılmaz: aktif
    tekliflerden SONRA gelir ve `inactive: True` taşır (Excel'de gri
    "pasif tedarikçi") — lab'ın sildiği firma Tedarikçi-1 olarak önerilmesin.

    **Para birimi karışık malzeme** (ör. Uludağ 1,88 €, Pharmaterm 250 ₺/kg):
    ham sayılar karşılaştırılamaz — 250 ₺ en pahalı görünür, ilk 3'te kalacak
    asıl ucuz teklif kesilebilirdi.  `rates` (`core.fx.Rates` ya da onu döndüren
    argümansız fonksiyon — yalnız karışık malzeme varsa BİR KEZ çağrılır, ağ
    isteği gereksiz yere yapılmaz) verilirse sıralama net fiyatın USD
    karşılığıyla yapılır (satın alma planı gibi).  Kur yoksa eski (ham) sıra
    korunur.  Her satırın `comparable` alanı: aktif fiyatlı teklifler aynı
    fiyat biriminde (kg / l / adet) VE (aynı para biriminde ya da kurla
    çevrilmiş) ise True — Excel "Tedarikçi-1 en uygun" vurgusunu yalnız o
    zaman yapar.
    """
    if not item_ids:
        return {}
    out: dict = {}
    q = (db.query(SupplierPrice)
         .filter(SupplierPrice.item_id.in_(item_ids), SupplierPrice.domain == domain)
         .all())
    is_inactive = inactive_firm_check(db, domain, q)

    manual_firms: dict = {}
    for sp in q:
        if sp.source == SOURCE_MANUAL:
            manual_firms.setdefault(sp.item_id, set()).update(firm_keys(sp.supplier_id, sp.supplier_name))
    for sp in q:
        if (sp.source != SOURCE_MANUAL
                and firm_keys(sp.supplier_id, sp.supplier_name) & manual_firms.get(sp.item_id, set())):
            continue
        out.setdefault(sp.item_id, []).append({
            "supplier_name": sp.supplier_name or (sp.supplier.name if sp.supplier else "—"),
            "package_size": sp.package_size,
            "unit_price": sp.unit_price,
            "currency": sp.currency,
            "price_unit": sp.price_unit,
            "source_label": sp.source_label,
            "quoted_at": sp.quoted_at,
            "inactive": is_inactive(sp),
            "vat_included": sp.vat_included,
            "vat_rate": sp.vat_rate,
        })
    fx = {"done": False, "rates": None}

    def get_rates():
        if not fx["done"]:
            fx["done"] = True
            try:
                fx["rates"] = rates() if callable(rates) else rates
            except Exception:                         # kur alınamadı → ham sıra
                fx["rates"] = None
        return fx["rates"]

    for iid, lst in out.items():
        live = [x for x in lst if x["unit_price"] is not None and not x["inactive"]]
        curs = {_price_currency(x) for x in live}
        units = {u for u in ((x["price_unit"] or "").strip().lower() for x in live) if u}
        keys = {id(x): net_unit_price(x) for x in lst if x["unit_price"] is not None}
        converted = False
        if len(curs) > 1:
            r = get_rates()
            if r is not None:
                from core.fx import convert
                try:
                    keys = {id(x): convert(net_unit_price(x), _price_currency(x), "USD", r)
                            for x in lst if x["unit_price"] is not None}
                    converted = True
                except (ValueError, TypeError):           # bilinmeyen para birimi → ham sıra
                    pass
        comparable = len(units) <= 1 and (len(curs) <= 1 or converted)
        # KDV dahil satır NET fiyatıyla yarışır (boş KDV'de net = fiyat → eski sıra)
        lst.sort(key=lambda x: (x["inactive"], x["unit_price"] is None,
                                keys[id(x)] if x["unit_price"] is not None else 0.0))
        for x in lst:
            x["comparable"] = comparable
        out[iid] = lst[:limit]
    return out


def _price_currency(offer) -> str:
    """Teklifin para birimi kodu (boş eski satır → 'USD', Stok Son Durum temeli)."""
    return (str(_offer_get(offer, "currency") or "USD")).strip().upper()


def today_rates_or_none():
    """Bugünün kuru (`core.fx.today_rates`) ya da alınamazsa None — rapor
    uçları `prices_for_items(..., rates=today_rates_or_none)` ile TEMBEL verir
    (yalnız para birimi karışık malzemede çağrılır)."""
    try:
        from core.fx import today_rates
        return today_rates()
    except Exception:
        return None


_CELL_SYM = {"USD": "$", "EUR": "€", "TRY": "₺"}


def price_cell_format(offer) -> str:
    """Excel fiyat hücresinin sayı biçimi: değer SAYI kalır, para birimi ve
    fiyat birimi görünür ("#,##0.00## €/kg").  Para birimi yoksa düz sayı."""
    cur = (str(_offer_get(offer, "currency") or "")).strip().upper()
    if not cur:
        return "#,##0.00##"
    unit = (str(_offer_get(offer, "price_unit") or "")).strip().lower()
    txt = _CELL_SYM.get(cur, cur) + (f"/{unit}" if unit else "")
    return '#,##0.00##" ' + txt.replace('"', "") + '"'


def _offer_get(offer, key):
    """Teklif dict'i (prices_for_items) ya da SupplierPrice nesnesi."""
    if isinstance(offer, dict):
        return offer.get(key)
    return getattr(offer, key, None)


def price_per_purchase_unit(offer, item_unit) -> Tuple[Optional[float], str, Optional[str]]:
    """Teklif fiyatını malzemenin ALIM birimine çevirir → (fiyat, alım birimi, not|None).

    Alım birimi: g/kg malzeme → 'kg', ml/l malzeme → 'l', diğerleri → 'adet'.
      • kg fiyatı ml/l malzemeye (ya da l fiyatı g/kg malzemeye) uygulanırsa
        fiyat aynen kullanılır, not "1 l ≈ 1 kg kabulüyle" eklenir.
      • adet ↔ kg/l çevrilemez: fiyat None döner + not (yanlış tutar
        üretmektense "fiyat yok" görünmesi güvenli).
      • `price_unit` boş eski satır → `default_price_unit` kuralı.
    Para birimine dokunmaz (çeviri `core.fx.convert`).
    """
    family = _unit_family(item_unit)
    purchase_unit = {"mass": "kg", "volume": "l"}.get(family, "adet")
    price = _offer_get(offer, "unit_price")
    basis = (_offer_get(offer, "price_unit") or "").strip().lower() or default_price_unit(item_unit)
    if price is None:
        return None, purchase_unit, None
    price = float(price)
    if basis == purchase_unit:
        return price, purchase_unit, None
    if {basis, purchase_unit} == {"kg", "l"}:
        return price, purchase_unit, LITRE_KG_NOTE
    return None, purchase_unit, (
        f"Fiyat birimi ({basis}) malzemenin alım birimiyle ({purchase_unit}) "
        f"uyuşmuyor — fiyat kullanılmadı")


def _vat_rate(offer) -> Optional[float]:
    """Geçerli KDV oranı (0–100, sonlu) ya da None."""
    import math
    r = _offer_get(offer, "vat_rate")
    try:
        r = float(r)
    except (TypeError, ValueError):
        return None
    return r if math.isfinite(r) and 0.0 <= r <= VAT_RATE_MAX else None


def net_unit_price(offer):
    """Teklifin KDV HARİÇ birim fiyatı — karşılaştırma ve tutar bununla yapılır.

    `vat_included` True VE geçerli oran varsa `fiyat / (1 + oran/100)`;
    aksi hâlde fiyat OLDUĞU GİBİ döner (KDV hariç, bilinmiyor ya da dahil ama
    oran yok — oran yoksa net hesaplanamaz, brüt kullanılır).  Teklif dict'i
    ya da SupplierPrice/OfferRec nesnesi."""
    price = _offer_get(offer, "unit_price")
    if price is None or _offer_get(offer, "vat_included") is not True:
        return price
    rate = _vat_rate(offer)
    if rate is None:
        return price
    return float(price) / (1.0 + rate / 100.0)


def vat_label(offer) -> Optional[str]:
    """KDV dahil satırın kısa etiketi ("KDV %20 dahil" / "KDV dahil") — diğer
    satırlarda None.  Raporlar paneli, Excel çıktıları ve satın alma planı
    aynı metni kullanır."""
    if _offer_get(offer, "vat_included") is not True:
        return None
    rate = _vat_rate(offer)
    return f"KDV %{rate:g} dahil" if rate is not None else "KDV dahil"


def serialize_price(sp, item=None, supplier_status: Optional[str] = None,
                    supplier_inactive: bool = False) -> dict:
    """Fiyat satırının uç yanıtı — GET /api/supplier-prices (grup içi satır),
    GET /api/suppliers/{id}/prices ve POST/PUT yanıtı aynı şekli paylaşır.
    `item` verilirse malzeme adı/birimi eklenir; tarih alanları ISO (quoted_at)
    ve TR saati (updated_at).  `supplier_inactive`: firma pasif
    (`inactive_firm_check`) — panelde gri "pasif" rozeti, "en ucuz" sayılmaz."""
    from database import to_tr
    out = {
        "id": sp.id,
        "item_id": sp.item_id,
        "supplier_id": sp.supplier_id,
        "supplier_name": sp.supplier_name or (sp.supplier.name if sp.supplier else "—"),
        "package_size": sp.package_size,
        "unit_price": sp.unit_price,
        "currency": sp.currency,
        "price_unit": sp.price_unit,
        "source": sp.source,
        "source_label": sp.source_label,
        "quoted_at": sp.quoted_at.isoformat() if sp.quoted_at else None,
        "note": sp.note,
        "vat_included": sp.vat_included,
        "vat_rate": sp.vat_rate,
        "created_by": sp.created_by,
        "updated_by": sp.updated_by,
        "updated_at": to_tr(sp.updated_at).strftime("%d.%m.%Y %H:%M") if sp.updated_at else "",
        "matched": sp.supplier_id is not None,
        "supplier_status": supplier_status,
        "supplier_inactive": bool(supplier_inactive),
    }
    if item is not None:
        out["material"] = item.name
        out["unit"] = item.unit or ""
        out["category"] = item.category or ""
    return out
