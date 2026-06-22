# ─────────────────────────────────────────────────────────────────────────────
# Copyright (c) 2026 Doğukan Yalçınkaya.  All rights reserved.
# Bu dosya FSEK kapsamında bir bilgisayar programı eserinin parçasıdır.
# Mali haklar yazılı devir olmadıkça eser sahibinde kalır (FSEK m.48).
# Bkz. LICENSE ve AUTHORS.md.
# ─────────────────────────────────────────────────────────────────────────────
"""
Üretim stok analizi (what-if simülasyonu).

"Seçilen ürünlerden N'er adet üretsek ne kadar malzeme harcanır, stok yeter
mi, yetmezse ne kadar eksik, ve tek başına kaç adet üretilebilir?" sorusunu
yanıtlar.  Lab ekibi Raporlar sayfasından çalıştırır; Excel olarak iner.

Hesap, gerçek üretim tüketim mantığıyla BİREBİR aynıdır (routers/production.py
`start_production`):
  • gross = net × (qty / output) × (1 + fire%/100)
  • Ambalaj/etiket fire MUAF (faktör 1.0)
  • Etiket dil çözümü: seçilen dile (TR/EN) göre kardeş etikete inilir; o
    dilde etiket yoksa kalem atlanır (üretimde de atlanır)
  • Aynı label_group iki kez sayılmaz

Domain (Faz 3): yalnızca aktif panele ait reçeteler simüle edilir.
"""
import math
from typing import Optional

from sqlalchemy.orm import Session

from database import Recipe, Item


def brand_of(name: str) -> str:
    """Ürün adından markayı türet — ilk kelime (Minerva / Serenida / Evanira…)."""
    return (name or "").strip().split(" ")[0] or "—"


def list_plan_products(db: Session, domain: str) -> list:
    """Aktif panelde reçetesi olan (üretilebilir) ürünler — seçim listesi için."""
    recs = (db.query(Recipe)
            .filter(Recipe.is_active == True, Recipe.domain == domain)
            .all())
    out = []
    for r in recs:
        tgt = db.query(Item).filter(Item.id == r.target_item_id).first() if r.target_item_id else None
        nm = (tgt.name if tgt else r.name) or "—"
        out.append({
            "recipe_id":  r.id,
            "name":       nm,
            "brand":      brand_of(nm),
            "output":     r.output_quantity,
        })
    out.sort(key=lambda x: (x["brand"].lower(), x["name"].lower()))
    return out


def _resolve_label(db: Session, item, lang: str):
    """Etiket dil çözümü — seçilen dilin kardeşine in; o dilde yoksa None."""
    if item.language and item.label_group and item.language != lang:
        return (db.query(Item)
                .filter(Item.label_group == item.label_group,
                        Item.language == lang, Item.is_active == True)
                .first())   # None → üretimde atlanır
    return item


def simulate(db: Session, recipe_ids: list, qty: float, lang: str, domain: str) -> dict:
    """
    Verilen reçetelerden `qty`'şer adet üretim senaryosunu hesaplar.

    Döner: { summary, materials[], producible[], purchase[], skipped_labels[] }
    """
    lang = "EN" if (lang or "TR").upper().startswith("EN") else "TR"
    recs = (db.query(Recipe)
            .filter(Recipe.id.in_(recipe_ids), Recipe.is_active == True,
                    Recipe.domain == domain)
            .all())

    used: dict = {}        # item_id -> {name,category,pkg,unit,current,used}
    producible: list = []
    skipped_labels: list = []

    for r in recs:
        tgt = db.query(Item).filter(Item.id == r.target_item_id).first() if r.target_item_id else None
        prod_name = (tgt.name if tgt else r.name) or "—"
        mult = qty / (r.output_quantity or 1.0)
        wf = 1.0 + (r.waste_percentage or 0.0) / 100.0
        processed = set()
        min_units = None
        lim = lim_stock = lim_per = lim_unit = None

        for ing in r.ingredients:
            it = db.query(Item).filter(Item.id == ing.item_id).first()
            if not it:
                continue
            if it.language and it.label_group:
                if it.label_group in processed:
                    continue
                processed.add(it.label_group)
                sib = _resolve_label(db, it, lang)
                if sib is None:
                    skipped_labels.append({"product": prod_name, "label": it.name})
                    continue
                it = sib
            is_amb = (it.category == "Ambalaj")
            factor = 1.0 if is_amb else wf
            per_unit = ing.quantity * (1.0 / (r.output_quantity or 1.0)) * factor
            gross = ing.quantity * mult * factor

            d = used.setdefault(it.id, {
                "item_id": it.id,
                "name": it.name, "category": it.category or "",
                "pkg": it.pkg_type or "", "unit": it.unit or "",
                "current": float(it.current_stock or 0.0), "used": 0.0,
            })
            d["used"] += gross

            if per_unit > 0:
                cap = float(it.current_stock or 0.0) / per_unit
                if min_units is None or cap < min_units:
                    min_units = cap
                    lim, lim_stock, lim_per, lim_unit = it.name, float(it.current_stock or 0.0), per_unit, (it.unit or "")

        producible.append({
            "product":     prod_name,
            "target":      qty,
            "producible":  0 if min_units is None else int(math.floor(min_units)),
            "limiting":    lim or "—",
            "limit_stock": round(lim_stock, 2) if lim_stock is not None else None,
            "per_unit":    round(lim_per, 4) if lim_per else None,
            "unit":        lim_unit or "",
        })

    # ── Kategori sırası + satır kurma ────────────────────────────────────
    def catlbl(d):
        return "Etiket" if (d["category"] == "Ambalaj" and d["pkg"] == "etiket") else (d["category"] or "—")

    def catord(d):
        return {"Hammadde": 0, "Ambalaj": 1, "Etiket": 2}.get(catlbl(d), 3)

    materials = []
    purchase = []
    short_count = 0
    total_raw = 0.0
    for d in sorted(used.values(), key=lambda d: (catord(d), d["name"].lower())):
        rem = round(d["current"] - d["used"], 2)
        is_short = rem < 0
        if is_short:
            short_count += 1
        if d["category"] == "Hammadde":
            total_raw += d["used"]
        row = {
            "item_id": d.get("item_id"),
            "name": d["name"], "category": catlbl(d), "unit": d["unit"],
            "used": round(d["used"], 2), "current": round(d["current"], 2),
            "remaining": rem, "status": "YETERSİZ" if is_short else "Yeterli",
            "shortfall": round(-rem, 2) if is_short else 0,
        }
        materials.append(row)
        if is_short:
            purchase.append(row)
    purchase.sort(key=lambda x: (-x["shortfall"]))

    producible.sort(key=lambda x: x["producible"])

    summary = {
        "products":     len(recs),
        "quantity":     qty,
        "total_units":  len(recs) * qty,
        "language":     lang,
        "materials":    len(materials),
        "short_count":  short_count,
        "ok_count":     len(materials) - short_count,
        "total_raw":    round(total_raw, 1),
    }
    return {
        "summary":        summary,
        "materials":      materials,
        "producible":     producible,
        "purchase":       purchase,
        "skipped_labels": skipped_labels,
    }


# ─── Excel üretimi ──────────────────────────────────────────────────────────

def build_workbook(report: dict, title_suffix: str = "", prices: dict = None) -> bytes:
    """Simülasyon raporundan 4-sayfalı .xlsx üretir (Özet / Tüketim / Üretilebilir / Satın Alma).

    `prices` verilirse ({item_id: [{supplier_name, package_size, unit_price}, …]}),
    'Satın Alma Listesi' sayfası her malzeme için 3 tedarikçiye kadar
    [Tedarikçi · Alınabilecek Miktar · Birim Fiyat] sütunlarıyla genişler
    (Işık Hanım'ın "Stok Son Durum" tablosu). Boşsa sayfa eski sade halinde kalır.
    """
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    s = report["summary"]
    wb = Workbook()
    NAVY = PatternFill("solid", fgColor="232E6E"); HF = Font(bold=True, color="FFFFFF")
    RED = Font(color="B91C1C", bold=True); GRN = Font(color="15803D")
    TITLE = Font(bold=True, size=13, color="232E6E"); BLD = Font(bold=True)
    thin = Border(*([Side(style="thin", color="DDDDDD")] * 4))

    def hrow(ws, headers):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(1, ci, h); c.fill = NAVY; c.font = HF
            c.alignment = Alignment(horizontal="left", vertical="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Özet
    ws = wb.active; ws.title = "Özet"
    ws["A1"] = f"Minerva 108 — Üretim Stok Analizi{(' · ' + title_suffix) if title_suffix else ''}"
    ws["A1"].font = TITLE
    meta = [
        ("Senaryo", f"{s['products']} üründen {int(s['quantity'])}'er adet = {int(s['total_units'])} adet"),
        ("Etiket dili", s["language"]),
        ("Varsayım", "Fire reçeteye göre (brüt) · ambalaj/etiket fire muaf"),
        ("", ""),
        ("Kullanılan farklı malzeme", s["materials"]),
        ("Stok YETERSİZ malzeme", s["short_count"]),
        ("Stok yeterli malzeme", s["ok_count"]),
        ("Toplam hammadde tüketimi (fireli, g+ml)", s["total_raw"]),
        ("", ""),
        ("Not", "Sheet 2: tüketim + eksik · Sheet 3: tek başına üretilebilir adet · Sheet 4: satın alma listesi"),
    ]
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(i, 1, k).font = BLD; ws.cell(i, 2, v)
    widths(ws, [42, 72])

    # Malzeme Tüketimi
    ws = wb.create_sheet("Malzeme Tüketimi")
    hrow(ws, ["Malzeme", "Kategori", "Birim", "Kullanılan (fireli)", "Mevcut Stok", "Kalan", "Durum", "Eksik Miktar"])
    for ri, m in enumerate(report["materials"], 2):
        short = m["remaining"] < 0
        vals = [m["name"], m["category"], m["unit"], m["used"], m["current"], m["remaining"],
                m["status"], m["shortfall"] if short else ""]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v); c.border = thin
            if short and ci in (6, 7, 8): c.font = RED
            if not short and ci == 6: c.font = GRN
    ws.freeze_panes = "A2"; widths(ws, [40, 11, 8, 18, 14, 14, 11, 14])

    # Üretilebilir Adet
    ws = wb.create_sheet("Üretilebilir Adet")
    hrow(ws, ["Ürün", "Hedef", "Üretilebilir (tek başına)", "Kısıtlayan Malzeme", "Kısıt Stok", "1 adet için ihtiyaç", "Birim"])
    for ri, p in enumerate(report["producible"], 2):
        vals = [p["product"], int(p["target"]), p["producible"], p["limiting"],
                p["limit_stock"] if p["limit_stock"] is not None else "",
                p["per_unit"] if p["per_unit"] else "", p["unit"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v); c.border = thin
            if ci == 3: c.font = RED if p["producible"] < p["target"] else GRN
    ws.freeze_panes = "A2"; widths(ws, [46, 8, 22, 40, 12, 18, 8])

    # Satın Alma Listesi  (+ varsa tedarikçi/paket/fiyat sütunları)
    ws = wb.create_sheet("Satın Alma Listesi")
    show_sup = bool(prices)
    SUP_SLOTS = 3
    base_headers = ["Malzeme", "Kategori", "Birim", "Toplam Gereken (fireli)", "Mevcut Stok", "ALINACAK (eksik)"]
    headers = list(base_headers)
    if show_sup:
        for n in range(1, SUP_SLOTS + 1):
            headers += [f"TEDARİKÇİ-{n}", "Alınabilecek Miktar", "Birim Fiyat"]
    hrow(ws, headers)
    for ri, m in enumerate(report["purchase"], 2):
        vals = [m["name"], m["category"], m["unit"], m["used"], m["current"], m["shortfall"]]
        if show_sup:
            plist = prices.get(m.get("item_id"), [])
            for i in range(SUP_SLOTS):
                p = plist[i] if i < len(plist) else None
                vals += [
                    (p["supplier_name"] if p else ""),
                    (p["package_size"] if p and p["package_size"] is not None else ""),
                    (p["unit_price"] if p and p["unit_price"] is not None else ""),
                ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v); c.border = thin
            if ci == 6: c.font = RED
            if show_sup and ci == 7 and v: c.font = BLD   # en ucuz tedarikçi vurgulansın
    last = len(report["purchase"]) + 3
    note = "Not: 'Alınacak' = brüt gereken − mevcut. Fire zaten gerekene dahildir."
    if show_sup:
        note += "  Tedarikçiler en ucuzdan pahalıya sıralı (Tedarikçi-1 en uygun)."
    ws.cell(last, 1, note).font = Font(italic=True, color="6b7280")
    ws.freeze_panes = "A2"
    w = [40, 11, 8, 22, 14, 18]
    if show_sup:
        for _n in range(SUP_SLOTS):
            w += [22, 16, 12]
    widths(ws, w)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
