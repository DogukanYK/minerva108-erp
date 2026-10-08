"""Lab fiyat notlarının (07.10.2026 tarama) tedarikçi fiyatlarına yüklenmesi
(idempotent, kuru çalıştırma varsayılan).

**Kaynak:** Songül Hanım'ın 30 sayfalık el yazılı fiyat notları
(`~/Desktop/Claude/Lab-Fiyat-Notlari/`), yazıya dökülmüş hâli
`scripts/data/lab_fiyat_notlari_202606.json`.  Kapsam (kullanıcı kararı):
  • A — lab karşılaştırma tablosu (s.1–2, 23.06.2026): malzeme × en çok 3
    tedarikçi, €/kg, "alınabilecek miktar" = ambalaj (kg), SARI hücre = seçim.
  • C — alım belgeleri (proforma / e-fatura / satış siparişi, s.8–10, s.12–16,
    s.19).
  Tedarikçi fiyat listeleri (D) YÜKLENMEZ: s.11 Doalin basılı listesi veri
  dosyasında `tur: fiyat_listesi` — fiyatı yazılmaz (durum `kapsam_disi`),
  yalnız tercih önerisinde "nottan alternatif" olur.  Sarı seçimler OTOMATİK
  tercih OLMAZ — kontrol Excel'iyle Songül Hanım'ın onayına gider
  (`--apply-prefs`); küçük satıcılar da yalnız öneridir.

**Eşleştirme (yazmadan önce, `plan`):**
  • Malzeme → aktif, bitmiş ürün OLMAYAN kart (domain cosmetics): TR-katlanmış
    birebir ad (`name` ya da `name_tr`; tablo IMS çıktısı) → (yalnız alım
    belgesi) notun "lab tablosu karşılığı: … (A1-08)" ipucu → `core.
    material_groups.material_key` eşitliği.  Birden çok aday = BELİRSİZ,
    yazılmaz.  Belgenin adı TEK karta uyuyor ama ipucu başka kartı
    gösteriyorsa (TR/EN mükerrer kart: BERGAMOT YAĞI ↔ BERGAMOT UÇUCU YAĞI)
    da BELİRSİZ ("ad ↔ lab tablosu ipucu çelişiyor") — yoksa (kart, firma)
    başına tek kayıt kuralı iki karta bölünerek delinirdi.  Toz/sıvı gibi
    ayrı form olabilecek kayıtlar (`SEPARATE_FORM`) otomatik eşlenmez.  Veri
    dosyasındaki isteğe bağlı `elle_eslesme` (`{"kart": {"A2-01": 512},
    "tedarikci": {"TİMAY": 45}}`) Excel incelemesinden sonra elle eşleme içindir.
  • Tedarikçi → aktif kart, `core.purchase_pricing.SupplierIndex` anahtarıyla
    (bitişik/ayrı yazım, "KİMYA" gibi jenerik kelimeler); bilinen takma adlar
    `SUPPLIER_ALIASES` (PHARMATEM → PHARMATERM, DOLAIN/DOALIN → DOALİNN …).
    Birleştirilmiş pasif kartın adı kazanana gider.  Bulunamayan firma
    "yeni tedarikçi" listesine girer; kartı YALNIZ onaylı Excel'in "Yeni
    tedarikçiler" sayfasında "Kartı açılsın mı"=E ise `--create-suppliers` ile
    açılır; açılmazsa o firmanın satırları ATLANIR (serbest metin fiyat
    yazılmaz).

**Kurallar:**
  • Belgede elle üstü çizili teklif (`cizili: true`) yazılmaz, öncelik
    yarışına girmez (durum `cizili`) — reddedilmiş satır başka kaydı ezemez.
  • Her (kart, firma) için TEK kayıt: alım belgesi (en yeni tarih) > lab
    tablosu (sarı seçim önce).  Ezilen değerler kazanan satırın notuna yazılır.
  • Fiyat birimi kartın birim ailesine uymuyorsa (`supplier_prices.
    price_unit_ok`) satır "uyumsuz" olur, öncelik yarışına girmez.
  • Kartta aynı firmanın ELLE girilmiş (`manual`) fiyatı varsa lab satırı
    YAZILMAZ (elle girilen kazanır — raporda).  `manual` ve `stok_son_durum`
    satırlarına HİÇ dokunulmaz (aynı firmanın Stok Son Durum fiyatı varsa
    ikisi de kalır — Excel'de not).
  • KDV: veri dosyasındaki `kdv_dahil` / `kdv_orani` → `vat_included` /
    `vat_rate` (Pharmaterm uçucu yağları KDV DAHİL brüt €; satın alma planı
    net fiyatla karşılaştırır).

**Yazım (`--commit --onayli ONAYLI.xlsx`):** önce `source_label` "Lab fiyat
notları 07.10.2026" önekli (kaynağı lab_notu | proforma | fatura | siparis)
satırlar silinir, sonra yazılacaklar eklenir → ikinci çalıştırma aynı sonucu
verir.  Satır: `source` (lab_notu | proforma | fatura | siparis),
`source_label` "Lab fiyat notları 07.10.2026 — s.N", `quoted_at`, ambalaj,
para birimi, birim, KDV, not, `created_by` "sistem (lab notları)".  Audit
`supplier_prices.lab_notes` (betik kullanıcısı yok → actor None, aktör adı
details'te).  `--commit` onaylı Excel'siz ÇALIŞMAZ: Songül Hanım'ın
"Belirsiz okumalar → Doğru değer" ve "Eşleşmeyenler → Doğru IMS kartı"
düzeltmeleri betikçe OTOMATİK işlenmez — veri dosyasına (fiyat /
`elle_eslesme.kart`) aktarılmamış dolu hücre varsa HİÇBİR ŞEY yazılmaz
(`fix_errors`; sessizce okunan değeri yazmak, ör. 2,15 ↔ 7,15, planı
yanıltırdı).  Boş bırakılan belirsiz okuma, notunda "okuma belirsiz"
etiketiyle yazılır.

**Kontrol Excel'i (`--xlsx YOL`, yazmaz):** Nasıl doldurulur · Fiyatlar ·
Eşleşmeyenler (Doğru IMS kartı) · Belirsiz okumalar (Doğru değer) · Tercih
önerisi (Onay E/H) · Küçük satıcı önerisi (Bitirilecek E/H) · Yeni
tedarikçiler (Açılsın mı E/H).  Onay hücreleri YALNIZ E / H (Evet / Hayır)
kabul eder — Excel listesi başka girişi reddeder, betik de "X", "✓", "1"
gibi değerleri hata sayar (çarpı Türkçede "hayır" da demektir).

**Onay uygulaması (`--apply-prefs ONAYLI.xlsx`, yalnız `--commit` ile yazar,
fiyatlara dokunmaz):** "Tercih önerisi"nde Onay=E satırları
`material_supplier_prefs`'e preferred yazılır (kartın aktif grubu varsa
gruba, yoksa karta; nottan alternatif satırı altında listelendiği lab
satırının kartına; aynı kapsamda birden çok E → Excel sırasıyla rank 1, 2…;
kapsamda zaten tercih varsa onlar korunur ve yeniler arkasına eklenir; aynı
tedarikçinin tercihi/"alma"sı varsa dokunulmaz).  "Küçük satıcı önerisi"nde
E olan firmanın aktif kartları `purchase_status='phase_out'` (sebep "Lab:
küçük miktarlı satıcı — 07.10.2026 onayı") olur.  Aynı firma hem tercih (E)
hem bitirilecek (E) ise ya da tercih edilen firma zaten bitirilecekse HATA
(bitirilecek tercihi geçersiz kılar — onay sessizce etkisiz kalırdı).
İşlenmemiş düzeltme hücresi burada da hatadır.  Audit `material_pref.create`
/ `supplier.status` (+ özet).

Kullanım (prod — ÖNCE pg_dump; kuru → Excel → Songül onayı → düzeltmeler
veri dosyasına → --commit → --apply-prefs):
    cd /var/www/minerva && set -a && source .env && set +a
    venv/bin/python scripts/import_lab_price_notes_20261007.py                       # kuru özet
    venv/bin/python scripts/import_lab_price_notes_20261007.py --xlsx /tmp/lab_fiyat_kontrol.xlsx
    venv/bin/python scripts/import_lab_price_notes_20261007.py --onayli /tmp/onayli.xlsx           # kuru (onaylı)
    venv/bin/python scripts/import_lab_price_notes_20261007.py --onayli /tmp/onayli.xlsx --commit --create-suppliers
    venv/bin/python scripts/import_lab_price_notes_20261007.py --apply-prefs /tmp/onayli.xlsx          # kuru
    venv/bin/python scripts/import_lab_price_notes_20261007.py --apply-prefs /tmp/onayli.xlsx --commit

Çıkış kodu: 0 = tamam; 2 = veri/Excel hatası (hiçbir şey yazılmadı).
"""
import argparse
import difflib
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy.orm import Session                                         # noqa: E402

from core.audit import log_admin_event                                     # noqa: E402
from core.material_groups import material_key                              # noqa: E402
from core.purchase_plan import alnum_fold                                  # noqa: E402
from core.purchase_pricing import SupplierIndex, supplier_key              # noqa: E402
from core.stock_lots import lot_kind                                       # noqa: E402
from core.supplier_prices import (SOURCE_MANUAL, SOURCE_STOK_SON_DURUM,   # noqa: E402
                                  _num, _unit_family, net_unit_price, normalize,
                                  price_unit_ok, vat_label)
from core.suppliers import live_card_id, normalize_status                  # noqa: E402
from database import (Item, MaterialGroup, MaterialSupplierPref,          # noqa: E402
                      SessionLocal, Supplier, SupplierPrice)

DATA_FILE = Path(__file__).resolve().parent / "data" / "lab_fiyat_notlari_202606.json"
DOMAIN = "cosmetics"
LABEL_PREFIX = "Lab fiyat notları 07.10.2026"
ACTOR = "sistem (lab notları)"
SOURCE_BY_TYPE = {"lab_tablo": "lab_notu", "proforma": "proforma", "fatura": "fatura",
                  "siparis": "siparis"}
LAB_SOURCES = tuple(SOURCE_BY_TYPE.values())
DOC_TYPES = ("proforma", "fatura", "siparis")
# Yalnız bilgi olan türler — fiyatı YAZILMAZ (plan P1c: tedarikçi fiyat
# listeleri D bu pakette yok); eşleşir, tercih önerisinde alternatif olabilir.
REFERENCE_TYPES = ("fiyat_listesi",)
TYPE_TEXT = {"lab_tablo": "lab tablosu", "proforma": "proforma", "fatura": "e-fatura",
             "siparis": "satış siparişi", "fiyat_listesi": "fiyat listesi"}
AUDIT_IMPORT = "supplier_prices.lab_notes"
AUDIT_PREFS = "material_pref.lab_notes"
PHASE_OUT_REASON = "Lab: küçük miktarlı satıcı — 07.10.2026 onayı"
PREF_NOTE = "Lab fiyat notları 07.10.2026 — Songül Hanım onayı"
SUPPLIER_NOTE = "Lab fiyat notlarından (07.10.2026) açıldı"

# Bilinen takma adlar — anahtar `alnum_fold(ad)` (parantez içi atılmış);
# değer prod'daki kartın adı (eşleme `SupplierIndex` anahtarıyla yapılır,
# kart "PHARMATERM İLAÇ" adlı olsa da bulunur).  ULUDAĞ AGRO / VESER
# KİMYA / DOĞASA anahtar kuralıyla zaten eşleşir; okunsun diye buradalar.
SUPPLIER_ALIASES = {
    "PHARMATEM": "PHARMATERM",
    "PHARMATERM": "PHARMATERM",
    "PHARMATERMILAC": "PHARMATERM",
    "DOLAIN": "DOALİNN",
    "DOALIN": "DOALİNN",
    "DOGASA": "DOGASA",
    "VESERKIMYA": "VESER KİMYEVİ",
    "ULUDGHERBAL": "ULUDAĞ HERBAL",
    "ULUDAGAGRO": "ULUDAĞ HERBAL",
    "KRKGIDA": "KRK GIDA",
    "TATLIDILIMLER": "TATLIDİLİMLER",
    "HAMMADDESEPETI": "HAMMADDESEPETİ",
}

# Ayrı form (toz ↔ sıvı kart) — otomatik eşlenmez, lab kartı söyler.
SEPARATE_FORM = {
    "A1-08-S3": "toz hali — sıvı ALEOVERA kartıyla eşleştirilmedi; toz kartı varsa lab belirtsin",
    "C-s8-1": "ALOEVERA EXTRACT PE (toz) — sıvı ALEOVERA kartıyla eşleştirilmedi; toz kartı varsa lab belirtsin",
    "C-s16-3": "ALOEVERA EXTRACT PE (toz) — sıvı ALEOVERA kartıyla eşleştirilmedi; toz kartı varsa lab belirtsin",
}

# Tercih önerisinde notların sarı seçimle çelişmesi (satır öneki → uyarı)
PREF_CONFLICTS = {
    "A2-13": "ÇELİŞKİ: s.5 el notu 'Aktif Kömür tatlı dilimlerden alınacak!' — tabloda sarı BEFCHEM.",
    "A2-17": "ÇELİŞKİ: s.10 el notu 'Itır → Doalin'den alınacak!' — tabloda sarı PHARMATERM.",
    "A2-28": "ÇELİŞKİ: s.9 bergamot 'Doalin'den Satın Alındı' (C-s11-1) — tabloda sarı NATURALYA.",
    "A2-34": "ÇELİŞKİ: s.9 misk adaçayı 'Doalin'den Satın Alındı' (C-s11-2) — tabloda sarı NATURALYA.",
    "A2-09": "Dikkat: BEFCHEM önünde elle çarpı (✗) — seçim belirsiz; S1 ATAMAN turuncu.",
    "A1-08": "s.17 kopyasında yalnız NATURALYA (10,32 €) işaretli ve daire içinde; SURYA işaretsiz.",
}
# s.17 el notu "Uludağ ucuz olanları alacağız!" — Uludağ daha ucuzken sarı
# başka firmada (Pharmaterm) olan satırlar → ek açıklama.  Seçim OTOMATİK
# DEĞİŞMEZ; sarı satıra uyarı yazılır (Songül Hanım karar verir).
ULUDAG_NOTE = "s.17 el notu 'Uludağ ucuz olanları alacağız!'"
_ORDERED = "seçim değiştirilmedi, yalnız uyarı"
_PAGE2 = "satır tablonun 2. sayfasında (s.18'de not yok) — genel not olarak"
ULUDAG_CHEAPER = {
    "A1-04": f"25.06 Pharmaterm siparişinde de var (C-s10-2, 'Okeylendi! Alınacak!') — {_ORDERED}",
    "A1-10": "", "A1-11": "", "A1-14": "",
    "A1-27": "s.17'de ULUDAĞ hücresi yeşil, PHARMATERM hücresi daire içinde ✗",
    "A1-28": (f"s.17'de ULUDAĞ hücresi yeşil, PHARMATERM daire içinde ✗; ancak 25.06'da "
              f"Pharmaterm'den satın alındı (C-s10-1) — {_ORDERED}"),
    "A2-01": _PAGE2, "A2-08": _PAGE2,
}
# Notlardan çıkan alternatif tercih satırları: (kayıt, malzeme satırı, açıklama)
# — tercih kapsamı altında listelendiği LAB SATIRININ kartıdır.
PREF_ALTERNATIVES = (
    ("A2-13-S1", "A2-13", "s.5 el notu: 'Aktif Kömür tatlı dilimlerden alınacak!'"),
    ("C-s11-1", "A2-28", "s.9: bergamot 'Doalin'den Satın Alındı' (s.11 fiyat listesi — fiyatı yüklenmez)"),
    ("C-s11-2", "A2-34", "s.9: misk adaçayı 'Doalin'den Satın Alındı' (s.11 fiyat listesi — fiyatı yüklenmez)"),
    ("A1-27-S1", "A1-27", "s.17: ULUDAĞ hücresi yeşil, PHARMATERM ✗; 'Uludağ ucuz olanları alacağız!'"),
    ("A1-28-S1", "A1-28", "s.17: ULUDAĞ hücresi yeşil, PHARMATERM ✗ (ama 25.06 Pharmaterm'den satın alındı)"),
)
# Belirsiz okumalar sayfasına "kesin" işaretli ama notu çelişen kayıtlar
EXTRA_CHECK = {
    "A2-19-S2": "149,64 ↔ 149,58 (s.3 notu) — hesapla 149,64 tutuyor",
    "A2-38-S2": "15,48 ↔ 15,84 (s.4 notu) — 18 $ × 0,86 = 15,48",
    "A2-17-S2": "70,17 (20,17 değil — s.10: 3.719 TL = 70,17 × 53)",
}
# Küçük satıcı önerisi — planın adlandırdığı firmalar (+ veriden: yalnız
# ≤ 1 kg ambalajla satanlar).  Doalin ayrıca not taşır.
SMALL_SELLER_NAMES = ("HAMMADDE SEPETİ", "TATLI DİLİMLER", "KİMYASAL EVİ", "ROSECE", "MARKET")
SMALL_PKG_MAX = 1.0
SMALL_SELLER_NOTES = {
    "TATLIDILIMLER": "Dikkat: s.5 el notu 'Aktif Kömür tatlı dilimlerden alınacak!'.",
    "MARKET": "Tabloda tuz 25 kg'lık; perakende market alımı.",
    "DOALINN": ("Doalin notu: 1 kg'lık şişeyle satıyor AMA bergamot ve misk adaçayı Doalin'den "
                "satın alındı (s.9/s.11) ve s.10 notu 'Itır → Doalin'den alınacak!' — "
                "bitirilecek yapılması önerilmez."),
}

SHEETS = ("Nasıl doldurulur", "Fiyatlar", "Eşleşmeyenler", "Belirsiz okumalar", "Tercih önerisi",
          "Küçük satıcı önerisi", "Yeni tedarikçiler")
APPROVE_COL = "Onay (E/H)"
PHASE_COL = "Bitirilecek yapılsın mı (E/H)"
NEW_COL = "Kartı açılsın mı (E/H)"
FIX_CARD_COL = "Doğru IMS kartı (no / ad)"
FIX_VALUE_COL = "Doğru değer (biliyorsanız)"
NOT_APPLICABLE = "—"                       # onay/düzeltme hücresi bu satırda yok

STATUS_TEXT = {
    "yazilacak": "Evet",
    "ezildi": "Hayır — aynı kart + firmada daha yeni/öncelikli kayıt var",
    "elle_var": "Hayır — elle girilmiş fiyat var (elle girilen kazanır)",
    "yeni_tedarikci": ("Hayır (şimdilik) — tedarikçi kartı yok; 'Yeni tedarikçiler' sayfasında E "
                       "verilirse kart açıldıktan sonra yazılır"),
    "cizili": "Hayır — belgede üstü çizili (alınmadı / reddedildi)",
    "kapsam_disi": "Hayır — tedarikçi fiyat listesi; bu yüklemede fiyatı yazılmaz (yalnız bilgi)",
    "pasif_tedarikci": "Hayır — tedarikçi kartı pasif (silinmiş); kart yeniden etkinleştirilmeli",
    "uyumsuz": "Hayır — fiyat birimi kartın birimiyle uyumsuz",
    "eslesmedi": "Hayır — IMS kartı bulunamadı",
    "belirsiz": "Hayır — birden çok aday kart",
    "ayri_form": "Hayır — ayrı form (toz/sıvı), kart lab tarafından belirlenmeli",
    "fiyatsiz": "Hayır — fiyat yok",
}

_PAREN = re.compile(r"\([^()]*\)")
_DOC_PKG = re.compile(r"\(\s*\d*[.,]?\d*\s*(?:KG|LT|L|ML|GR|G)\.?\s*\)"
                      r"|(?<![\w])\d+[.,]?\d*\s*(?:LT|KG|ML)\b\.?"
                      r"|(?<![\w])LT\b\.?", re.IGNORECASE)
_HINT = re.compile(r"lab tablosu karşılığı:\s*(.+)$", re.IGNORECASE)
_HINT_REF = re.compile(r"\((A\d-\d{2})")


# ─── Yardımcılar ────────────────────────────────────────────────────────────

def load_data(path=None) -> dict:
    with open(path or DATA_FILE, encoding="utf-8") as f:
        return json.load(f)


def _num_tr(v) -> str:
    """12.5 → '12,5'; 18999.67 → '18.999,67' (en çok 4 ondalık)."""
    if v is None:
        return "—"
    s = f"{float(v):,.4f}".rstrip("0").rstrip(".")
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _date(v) -> Optional[date]:
    return date.fromisoformat(v) if v else None


def _dmy(d: Optional[date]) -> str:
    return d.strftime("%d.%m.%Y") if d else "tarihsiz"


def _row_prefix(rec_id: str) -> str:
    """'A1-08-S2' → 'A1-08' (lab tablosu satırı); diğerleri kendisi."""
    parts = rec_id.split("-")
    return "-".join(parts[:2]) if rec_id.startswith("A") and len(parts) >= 3 else rec_id


def clean_supplier_name(name) -> str:
    """'GÜLER KİMYA (YEŞİL KİL)' → 'GÜLER KİMYA' (parantez içi atılır)."""
    s = _PAREN.sub(" ", name or "")
    return " ".join(s.split()) or " ".join((name or "").split())


def clean_doc_material(name) -> str:
    """Alım belgesindeki ambalaj eklerini at: 'BERGAMOT YAĞI 1LT' →
    'BERGAMOT YAĞI', 'SUSAM YAĞI (KG)' → 'SUSAM YAĞI', 'XANTHAN GUM (25 KG)
    (MEIUHA)' → 'XANTHAN GUM (MEIUHA)'.  Ad hiç boşalmaz."""
    s = " ".join(_DOC_PKG.sub(" ", name or "").split()).rstrip(" .")
    return s or " ".join((name or "").split())


def price_label(page) -> str:
    return f"{LABEL_PREFIX} — s.{page}"[:120]


def _yes_no(v) -> Optional[bool]:
    """Excel onay hücresi → True (E/Evet) | False (H/Hayır) | None (boş).

    YALNIZ E / EVET / H / HAYIR (büyük-küçük harf, nokta önemsiz).  Başka her
    dolu değer ValueError: "X" Türkçede çoğu zaman "hayır/iptal" demek (lab'ın
    notlarındaki ✗ gibi) — onay sayılırsa reddedilen öneri kalıcı tercih /
    bitirilecek olurdu; "✓", "+", "1" gibi işaretler de boş sayılıp onay
    sessizce düşmesin diye hatadır."""
    if v is None:
        return None
    raw = str(v).strip()
    if not raw:
        return None
    s = alnum_fold(raw)
    if s in ("E", "EVET"):
        return True
    if s in ("H", "HAYIR"):
        return False
    raise ValueError(raw)


# ─── Eşleştiriciler ─────────────────────────────────────────────────────────

class _Suppliers:
    """Paneldeki tedarikçi kartları + belge adları tek anahtar uzayında."""

    def __init__(self, db: Session, domain: str, doc_names, extra_names=(), overrides=None):
        sups = db.query(Supplier).filter(Supplier.domain == domain).order_by(Supplier.id).all()
        self.by_id = {s.id: s for s in sups}
        names = [s.name for s in sups if s.name]
        names += [clean_supplier_name(n) for n in doc_names if n]
        names += list(SUPPLIER_ALIASES.values()) + [n for n in extra_names if n]
        self.ix = SupplierIndex(names)
        self.active: Dict[str, List[Supplier]] = {}
        self.merged: Dict[str, int] = {}
        self.inactive: Dict[str, List[Supplier]] = {}
        cards = {s.id: (s.is_active is not False, s.merged_into_id) for s in sups}
        for s in sups:
            k = self.ix.key(s.name)
            if not k:
                continue
            if s.is_active is not False:
                self.active.setdefault(k, []).append(s)
            elif s.merged_into_id is not None:
                live = live_card_id(cards, s.id)
                if live is not None:
                    self.merged.setdefault(k, live)
            else:
                self.inactive.setdefault(k, []).append(s)
        self.overrides = {normalize(k): int(v) for k, v in (overrides or {}).items() if v}

    def key(self, name) -> str:
        return self.ix.key(name) if name else ""

    def resolve(self, doc_name) -> dict:
        """→ {kind: 'id'|'inactive'|'new', id, name, key, how, dupes: [aynı firmanın
        diğer aktif kartları]}.  'inactive': firmanın yalnız PASİF (silinmiş,
        birleştirilmemiş) kartı var — yeni kart AÇILMAZ, satır yazılmaz (lab
        firmayı bilerek silmiş olabilir)."""
        ov = self.overrides.get(normalize(doc_name))
        if ov is not None:
            s = self.by_id.get(ov)
            if s is None or s.is_active is False:
                raise ValueError(f"elle_eslesme.tedarikci «{doc_name}» → {ov}: aktif kart yok")
            return {"kind": "id", "id": s.id, "name": s.name, "key": self.key(s.name),
                    "how": "elle eşleme", "dupes": []}
        cleaned = clean_supplier_name(doc_name)
        target = SUPPLIER_ALIASES.get(alnum_fold(cleaned))
        for cand, how in ((target, "takma ad"), (cleaned, "ad")):
            if not cand:
                continue
            k = self.key(cand)
            act = self.active.get(k)
            if act:
                return {"kind": "id", "id": act[0].id, "name": act[0].name, "key": k,
                        "how": "ad" if alnum_fold(cand) == alnum_fold(cleaned) else how,
                        "dupes": [s.name for s in act[1:]]}
            live = self.merged.get(k)
            if live is not None and live in self.by_id:
                s = self.by_id[live]
                return {"kind": "id", "id": s.id, "name": s.name, "key": self.key(s.name),
                        "how": "birleştirilmiş kartın kazananı", "dupes": []}
        for cand in (target, cleaned):
            dead = self.inactive.get(self.key(cand)) if cand else None
            if dead:
                return {"kind": "inactive", "id": dead[0].id, "name": dead[0].name,
                        "key": self.key(cand), "how": "pasif kart", "dupes": []}
        name = target or cleaned
        return {"kind": "new", "id": None, "name": name, "key": self.key(name), "how": "yok",
                "dupes": []}

    def active_cards(self, k: str) -> List[Supplier]:
        return list(self.active.get(k, []))


class _Materials:
    """Aktif, bitmiş ürün olmayan kartlar — ad / ad_tr / material_key dizinleri."""

    def __init__(self, db: Session, domain: str, sup_keys):
        rows = (db.query(Item).filter(Item.domain == domain, Item.is_active == True)   # noqa: E712
                .order_by(Item.id).all())
        self.items = {it.id: it for it in rows if lot_kind(it.category) != "finished"}
        self.sup_keys = set(sup_keys)
        self.by_norm: Dict[str, set] = {}
        self.by_key: Dict[tuple, set] = {}
        for it in self.items.values():
            for nm in (it.name, it.name_tr):
                if not nm:
                    continue
                self.by_norm.setdefault(normalize(nm), set()).add(it.id)
                k = material_key(nm, self.sup_keys)
                if k:
                    self.by_key.setdefault(k, set()).add(it.id)
        gids = {it.material_group_id for it in self.items.values() if it.material_group_id}
        self.groups = {g.id: g for g in (db.query(MaterialGroup)
                                         .filter(MaterialGroup.id.in_(gids),
                                                 MaterialGroup.is_active == True).all())   # noqa: E712
                       } if gids else {}
        self._names = {}
        for it in self.items.values():
            for nm in (it.name, it.name_tr):
                if nm:
                    self._names.setdefault(normalize(nm), it.id)

    def exact(self, name) -> List[int]:
        return sorted(self.by_norm.get(normalize(name), ()))

    def keyed(self, name) -> List[int]:
        k = material_key(name, self.sup_keys)
        return sorted(self.by_key.get(k, ())) if k else []

    def suggest(self, name, n: int = 3) -> List[int]:
        got = difflib.get_close_matches(normalize(name), list(self._names), n=n, cutoff=0.75)
        out: List[int] = []
        for g in got:
            iid = self._names[g]
            if iid not in out:
                out.append(iid)
        return out

    def view(self, iid: Optional[int]) -> Optional[dict]:
        it = self.items.get(iid) if iid else None
        if it is None:
            return None
        grp = self.groups.get(it.material_group_id) if it.material_group_id else None
        return {"id": it.id, "name": it.name, "unit": it.unit or "",
                "group_id": grp.id if grp else None, "group_name": grp.name if grp else None}

    def label(self, iid) -> str:
        v = self.view(iid)
        return f"{v['id']} {v['name']} ({v['unit'] or 'birimsiz'})" if v else "—"


def _hint_rows(note) -> List[str]:
    m = _HINT.search(note or "")
    return _HINT_REF.findall(m.group(1)) if m else []


# ─── Plan (yazmaz) ──────────────────────────────────────────────────────────

def _match_material(r: dict, mats: _Materials, row_card: Dict[str, dict],
                    item_overrides: Dict[str, int]) -> dict:
    """→ {status: ok|belirsiz|eslesmedi|ayri_form, item_id, how, candidates, suggestions}."""
    rid = r["id"]
    ov = item_overrides.get(rid, item_overrides.get(_row_prefix(rid)))
    if ov:
        if ov not in mats.items:
            raise ValueError(f"elle_eslesme.kart {rid} → {ov}: aktif hammadde kartı yok")
        return {"status": "ok", "item_id": ov, "how": "elle eşleme", "candidates": [], "suggestions": []}
    if rid in SEPARATE_FORM:
        return {"status": "ayri_form", "item_id": None, "how": None, "candidates": [],
                "suggestions": [], "reason": SEPARATE_FORM[rid]}
    doc = r["tur"] in DOC_TYPES or r["tur"] in REFERENCE_TYPES
    name = clean_doc_material(r["malzeme"]) if doc else r["malzeme"]
    ex = mats.exact(name)
    hint: List[int] = []
    if doc:
        for ref in _hint_rows(r.get("not")):
            rc = row_card.get(ref)
            if rc and rc.get("item_id") and rc["item_id"] not in hint:
                hint.append(rc["item_id"])
    if len(ex) == 1:
        # Ad tek karta uyuyor ama notun ipucu BAŞKA kartı gösteriyor (TR/EN
        # mükerrer kart: BERGAMOT YAĞI ↔ BERGAMOT UÇUCU YAĞI) → belirsiz;
        # yoksa aynı firmanın fiyatı iki karta bölünür, öncelik kuralı delinirdi.
        if hint and ex[0] not in hint:
            return {"status": "belirsiz", "item_id": None, "how": None,
                    "candidates": ex + [h for h in hint if h not in ex], "suggestions": [],
                    "reason": "ad ↔ lab tablosu ipucu çelişiyor"}
        return {"status": "ok", "item_id": ex[0], "how": "ad", "candidates": [], "suggestions": []}
    if len(ex) > 1:                                   # aynı adlı 2+ kart: ipucu ayırabilir
        both = [i for i in ex if i in hint]
        if len(both) == 1:
            return {"status": "ok", "item_id": both[0], "how": "ad + lab tablosu ipucu",
                    "candidates": [], "suggestions": []}
        return {"status": "belirsiz", "item_id": None, "how": None, "candidates": ex, "suggestions": []}
    if doc:
        if len(hint) == 1:
            return {"status": "ok", "item_id": hint[0], "how": "lab tablosu ipucu",
                    "candidates": [], "suggestions": []}
    keyed = mats.keyed(name)
    if hint:                                          # ipucu 2+ kart: anahtar ayırsın
        both = [i for i in keyed if i in hint]
        if len(both) == 1:
            return {"status": "ok", "item_id": both[0], "how": "ipucu + material_key",
                    "candidates": [], "suggestions": []}
        return {"status": "belirsiz", "item_id": None, "how": None, "candidates": hint, "suggestions": []}
    if len(keyed) == 1:
        return {"status": "ok", "item_id": keyed[0], "how": "material_key", "candidates": [],
                "suggestions": []}
    if len(keyed) > 1:
        return {"status": "belirsiz", "item_id": None, "how": None, "candidates": keyed, "suggestions": []}
    return {"status": "eslesmedi", "item_id": None, "how": None, "candidates": [],
            "suggestions": mats.suggest(name)}


def _prio(r: dict) -> tuple:
    """Öncelik: alım belgesi (en yeni önce, tarihsiz sonda) > lab tablosu (sarı önce)."""
    if r["type"] in DOC_TYPES:
        return (0, -(r["date"].toordinal() if r["date"] else 0), 0, r["idx"])
    return (1, 0, 0 if r["selected"] else 1, r["idx"])


def _short(r: dict) -> str:
    return (f"{r['id']} ({TYPE_TEXT[r['type']]} {_dmy(r['date'])}) {_num_tr(r['price'])} "
            f"{r['currency']}/{r['price_unit']}")


def plan(db: Session, data: Optional[dict] = None, *, domain: str = DOMAIN) -> dict:
    """Ne yazılacağını hesaplar — hiçbir şey YAZMAZ.

    Döner: {records: [kayıt görünümü], new_suppliers, prefs (tercih önerisi
    satırları), small (küçük satıcı satırları), existing_lab_rows, counts}.
    Kayıt `status`: yazilacak | ezildi | elle_var | yeni_tedarikci |
    pasif_tedarikci | uyumsuz | eslesmedi | belirsiz | ayri_form | fiyatsiz |
    cizili (belgede üstü çizili) | kapsam_disi (fiyat listesi, D)."""
    data = data or load_data()
    recs_in = data.get("kayitlar") or []
    manual_ov = data.get("elle_eslesme") or {}
    item_ov = {k: int(v) for k, v in (manual_ov.get("kart") or {}).items() if v}

    existing = (db.query(SupplierPrice)
                .filter(SupplierPrice.domain == domain,
                        SupplierPrice.source.in_((SOURCE_MANUAL, SOURCE_STOK_SON_DURUM))
                        | SupplierPrice.source.is_(None))
                .all())
    sups = _Suppliers(db, domain, [r.get("tedarikci") for r in recs_in],
                      extra_names=[e.supplier_name for e in existing],
                      overrides=manual_ov.get("tedarikci"))
    sup_keys = {supplier_key(s.name) for s in sups.by_id.values() if s.name}
    sup_keys.discard("")
    mats = _Materials(db, domain, sup_keys)

    # ── 1) kayıt görünümleri + malzeme/tedarikçi eşleşmesi (A önce: ipucu) ──
    recs: List[dict] = []
    row_card: Dict[str, dict] = {}
    order = sorted(range(len(recs_in)), key=lambda i: (recs_in[i]["tur"] != "lab_tablo", i))
    views: Dict[int, dict] = {}
    for i in order:
        r = recs_in[i]
        if r["tur"] not in SOURCE_BY_TYPE and r["tur"] not in REFERENCE_TYPES:
            raise ValueError(f"{r.get('id')}: bilinmeyen tür {r.get('tur')!r}")
        m = _match_material(r, mats, row_card, item_ov)
        if r["tur"] == "lab_tablo":
            row_card.setdefault(_row_prefix(r["id"]), m if m["status"] == "ok" else {})
        s = sups.resolve(r.get("tedarikci"))
        views[i] = {
            "idx": i, "id": r["id"], "page": r.get("kaynak_sayfa"), "type": r["tur"],
            "source": SOURCE_BY_TYPE.get(r["tur"]), "date": _date(r.get("tarih")),
            "material": r.get("malzeme") or "", "supplier": r.get("tedarikci") or "",
            "price": r.get("fiyat"), "currency": (r.get("para_birimi") or "").upper(),
            "price_unit": (r.get("fiyat_birimi") or "").lower(), "package": r.get("ambalaj_kg"),
            "vat_included": r.get("kdv_dahil"), "vat_rate": r.get("kdv_orani"),
            "selected": bool(r.get("secili")), "struck": bool(r.get("cizili")),
            "confidence": r.get("guven") or "",
            "note": r.get("not") or "", "table_unit": r.get("birim_kart"),
            "mat": m, "item": mats.view(m.get("item_id")), "sup": s,
            "flags": [], "overridden": [], "status": None,
        }
    recs = [views[i] for i in range(len(recs_in))]

    # ── 2) tek tek uygunluk (kapsam, üstü çizili, fiyat, birim) ──
    for r in recs:
        it = r["item"]
        if r["type"] in REFERENCE_TYPES:              # fiyat listesi (D) — yalnız bilgi
            r["status"] = "kapsam_disi"
            continue
        if r["struck"]:                               # reddedilmiş teklif yarışa girmez
            r["status"] = "cizili"
            continue
        if r["mat"]["status"] != "ok":
            r["status"] = r["mat"]["status"]
            if r["mat"]["status"] == "belirsiz" and r["mat"].get("reason"):
                r["flags"].append(r["mat"]["reason"])
            continue
        if r["price"] is None or float(r["price"]) <= 0:
            r["status"] = "fiyatsiz"
            continue
        if not price_unit_ok(it["unit"], r["price_unit"]):
            r["status"] = "uyumsuz"
            r["flags"].append(f"fiyat birimi {r['price_unit']} ↔ kart birimi {it['unit'] or 'birimsiz'}")
            continue
        if r["table_unit"] and _unit_family(r["table_unit"]) != _unit_family(it["unit"]):
            r["flags"].append(f"tablodaki birim {r['table_unit']}, kartın birimi {it['unit'] or 'birimsiz'}")
        if r["sup"]["dupes"]:
            r["flags"].append("firmanın başka aktif kartı da var: " + ", ".join(r["sup"]["dupes"]))
        if r["sup"]["kind"] == "inactive":
            r["status"] = "pasif_tedarikci"
            r["flags"].append(f"tedarikçi kartı pasif: {r['sup']['id']} {r['sup']['name']}")

    # ── 3) öncelik: (kart, firma) başına tek kayıt ──
    groups: Dict[tuple, List[dict]] = {}
    for r in recs:
        if r["status"] is not None:
            continue
        firm = ("id", r["sup"]["id"]) if r["sup"]["kind"] == "id" else ("new", r["sup"]["key"])
        groups.setdefault((r["item"]["id"], firm), []).append(r)
    for lst in groups.values():
        lst.sort(key=_prio)
        win = lst[0]
        for lo in lst[1:]:
            lo["status"] = "ezildi"
            lo["flags"].append(f"kazanan: {win['id']}")
            win["overridden"].append(_short(lo))

    # ── 4) elle fiyat / Stok Son Durum / yeni tedarikçi ──
    by_item: Dict[int, List[SupplierPrice]] = {}
    for e in existing:
        by_item.setdefault(e.item_id, []).append(e)

    def firm_hit(e: SupplierPrice, r: dict) -> bool:
        if r["sup"]["kind"] == "id" and e.supplier_id == r["sup"]["id"]:
            return True
        nm = e.supplier_name or (sups.by_id[e.supplier_id].name if e.supplier_id in sups.by_id else None)
        return bool(nm) and sups.key(nm) == r["sup"]["key"]

    for r in recs:
        if r["status"] is not None:
            continue
        same = [e for e in by_item.get(r["item"]["id"], []) if firm_hit(e, r)]
        man = [e for e in same if e.source == SOURCE_MANUAL]
        ssd = [e for e in same if e.source != SOURCE_MANUAL]
        if man:
            e = man[0]
            r["status"] = "elle_var"
            r["flags"].append(f"elle fiyat: {_num_tr(e.unit_price)} {e.currency or ''}/{e.price_unit or ''}"
                              f" ({e.updated_by or e.created_by or '—'})")
            continue
        if ssd:
            e = ssd[0]
            r["flags"].append(f"aynı firmanın Stok Son Durum fiyatı da var: {_num_tr(e.unit_price)} "
                              f"{e.currency or ''}/{e.price_unit or ''} — ikisi de kalır")
        r["status"] = "yeni_tedarikci" if r["sup"]["kind"] == "new" else "yazilacak"

    for r in recs:
        r["write_note"] = _write_note(r)
    cand_names = {iid: mats.label(iid) for r in recs
                  for iid in (r["mat"].get("candidates") or []) + (r["mat"].get("suggestions") or [])}

    new_sups = _new_suppliers(recs)
    prefs = _pref_rows(db, recs, sups, domain)
    small = _small_rows(recs, sups)
    existing_lab = (db.query(SupplierPrice)
                    .filter(SupplierPrice.domain == domain, SupplierPrice.source.in_(LAB_SOURCES),
                            SupplierPrice.source_label.like(f"{LABEL_PREFIX}%")).count())
    counts = {
        "kayit": len(recs),
        "malzeme_eslesen": sum(1 for r in recs if r["mat"]["status"] == "ok"),
        "malzeme_eslesmeyen": sum(1 for r in recs if r["mat"]["status"] == "eslesmedi"),
        "malzeme_belirsiz": sum(1 for r in recs if r["mat"]["status"] == "belirsiz"),
        "ayri_form": sum(1 for r in recs if r["mat"]["status"] == "ayri_form"),
        "tedarikci_eslesen": sum(1 for r in recs if r["sup"]["kind"] == "id"),
        "yeni_firma": len(new_sups),
        "okuma_belirsiz": sum(1 for r in recs if r["confidence"] == "belirsiz"),
    }
    for st in STATUS_TEXT:
        counts[st] = sum(1 for r in recs if r["status"] == st)
    return {"domain": domain, "records": recs, "new_suppliers": new_sups, "prefs": prefs,
            "small": small, "existing_lab_rows": existing_lab, "counts": counts,
            "cand_names": cand_names}


def _write_note(r: dict) -> str:
    bits = [f"[{r['id']}] {TYPE_TEXT[r['type']]} s.{r['page']}"
            + (" — SARI seçim (tercih önerisi, onaya bağlı)" if r["selected"] else "")]
    if r["note"]:
        bits.append(r["note"])
    if r["confidence"] == "belirsiz":
        bits.append("OKUMA BELİRSİZ — Excel 'Belirsiz okumalar'")
    if r["overridden"]:
        bits.append("Ezilen: " + "; ".join(r["overridden"]))
    return " · ".join(bits)


def _new_suppliers(recs: List[dict]) -> List[dict]:
    out: Dict[str, dict] = {}
    for r in recs:
        s = r["sup"]
        if s["kind"] != "new" or not s["key"]:
            continue
        d = out.setdefault(s["key"], {"key": s["key"], "names": {}, "spellings": [], "records": [],
                                      "materials": []})
        d["names"][s["name"]] = d["names"].get(s["name"], 0) + 1
        if r["supplier"] not in d["spellings"]:
            d["spellings"].append(r["supplier"])
        d["records"].append(r["id"])
        if r["material"] not in d["materials"]:
            d["materials"].append(r["material"])
    lst = []
    for d in out.values():
        nm = sorted(d.pop("names").items(), key=lambda kv: (-kv[1], kv[0] != kv[0].upper(), kv[0]))
        d["name"] = nm[0][0]                       # en sık (büyük harf öncelikli) temiz ad
        d["writable"] = sum(1 for r in recs if r["sup"]["key"] == d["key"] and r["status"] == "yeni_tedarikci")
        lst.append(d)
    lst.sort(key=lambda d: alnum_fold(d["name"]))
    return lst


def _scope_of(item: Optional[dict]) -> Optional[tuple]:
    if not item:
        return None
    return ("group", item["group_id"]) if item.get("group_id") else ("item", item["id"])


def _existing_prefs(db: Session, domain: str) -> Dict[tuple, List[MaterialSupplierPref]]:
    out: Dict[tuple, List[MaterialSupplierPref]] = {}
    for p in db.query(MaterialSupplierPref).filter(MaterialSupplierPref.domain == domain).all():
        k = ("group", p.material_group_id) if p.material_group_id else ("item", p.item_id)
        out.setdefault(k, []).append(p)
    return out


def _pref_text(p: MaterialSupplierPref, by_id: dict) -> str:
    s = by_id.get(p.supplier_id)
    nm = s.name if s else f"#{p.supplier_id}"
    return f"{nm} ({'tercih ' + str(p.rank) if p.preference == 'preferred' else 'alma'})"


def _net(r: dict) -> Optional[float]:
    """Kayıt görünümünün KDV hariç fiyatı (karşılaştırma için)."""
    if r.get("price") is None:
        return None
    return net_unit_price({"unit_price": float(r["price"]), "vat_included": r.get("vat_included"),
                           "vat_rate": r.get("vat_rate")})


def _price_text(r: dict) -> str:
    """'82,56 EUR/kg (KDV %20 dahil; net 68,8)' — KDV dahil brüt fiyat etiketsiz
    görünürse onay veren firmayı ~%20 pahalı sanır."""
    txt = f"{_num_tr(r['price'])} {r['currency']}/{r['price_unit']}"
    lab_ = vat_label({"unit_price": r["price"], "vat_included": r.get("vat_included"),
                      "vat_rate": r.get("vat_rate")})
    if lab_:
        net = _net(r)
        txt += f" ({lab_}" + (f"; net {_num_tr(round(net, 2))}" if net is not None and net != r["price"] else "") + ")"
    return txt


def _is_uludag(r: dict) -> bool:
    return alnum_fold(clean_supplier_name(r.get("supplier"))).startswith("ULUD")


def _pref_rows(db: Session, recs: List[dict], sups: _Suppliers, domain: str) -> List[dict]:
    """Tercih önerisi: lab tablosunun SARI seçimleri (+ notlardan alternatifler).

    • Bir satırda 2+ sarı varsa (aloe) KDV hariç UCUZ olan üstte — ikisi de E
      ise üstteki 1. tercih olur.
    • Nottan alternatif satırın kartı/kapsamı, altında listelendiği LAB
      SATIRININ kartıdır (belge kaydı başka karta eşlenmiş ya da hiç
      eşlenmemiş olsa da tercih o malzemeye yazılır).
    • `ULUDAG_CHEAPER` satırlarında sarıya "Uludağ daha ucuz" uyarısı."""
    by_id = {r["id"]: r for r in recs}
    yellow_by_row: Dict[str, List[dict]] = {}
    row_recs: Dict[str, List[dict]] = {}
    for r in recs:
        if r["type"] == "lab_tablo":
            row_recs.setdefault(_row_prefix(r["id"]), []).append(r)
            if r["selected"]:
                yellow_by_row.setdefault(_row_prefix(r["id"]), []).append(r)
    for lst in yellow_by_row.values():
        if len(lst) > 1 and len({y["currency"] for y in lst}) == 1:
            lst.sort(key=lambda y: (_net(y) is None, _net(y) or 0.0, y["idx"]))
    alts: Dict[str, List[tuple]] = {}
    for rid, row, why in PREF_ALTERNATIVES:
        if rid in by_id:
            alts.setdefault(row, []).append((by_id[rid], why))
    ex = _existing_prefs(db, domain)
    rows: List[dict] = []
    seen_rows = []
    for r in recs:
        row = _row_prefix(r["id"])
        if r["type"] != "lab_tablo" or row in seen_rows or row not in yellow_by_row:
            continue
        seen_rows.append(row)
        yel = yellow_by_row[row]
        row_item = next((y["item"] for y in yel if y["item"]), None)
        entries = [(y, "sarı seçim", None) for y in yel] + [(a, "nottan alternatif", why) for a, why in alts.get(row, [])]
        for rec, kind, why in entries:
            item = rec["item"] if kind == "sarı seçim" else row_item
            warn = []
            if kind == "sarı seçim" and len(yel) > 1:
                warn.append(f"Aynı malzemede {len(yel)} sarı seçim — ikisi de E ise üstteki (ucuz olan) "
                            "1. tercih olur.")
            if row in PREF_CONFLICTS:
                warn.append(PREF_CONFLICTS[row])
            if kind == "sarı seçim" and row in ULUDAG_CHEAPER and not _is_uludag(rec):
                ul = next((u for u in row_recs.get(row, []) if _is_uludag(u) and u["price"] is not None), None)
                if ul is not None and rec["price"] is not None and _net(ul) < _net(rec):
                    extra = ULUDAG_CHEAPER[row]
                    warn.append(f"{ULUDAG_NOTE} — bu satırda Uludağ daha ucuz ({_num_tr(ul['price'])} "
                                f"{ul['currency']} ↔ sarı {_num_tr(rec['price'])} {rec['currency']})."
                                + (f" {extra}." if extra else ""))
            if why:
                warn.append(f"Alternatif: {why}.")
            if item is None:
                warn.append("IMS kartı eşleşmedi — onaylansa da uygulanamaz (önce kart eşlemesi).")
            elif kind == "nottan alternatif" and rec["item"] and rec["item"]["id"] != item["id"]:
                warn.append(f"Belge kaydı başka karta eşlendi ({rec['item']['id']} {rec['item']['name']}); "
                            "tercih bu lab satırının kartına yazılır.")
            if rec["sup"]["kind"] == "new":
                warn.append("Tedarikçi kartı yok — 'Yeni tedarikçiler' sayfasında bu firmaya E verilirse "
                            "kart açıldıktan sonra uygulanır.")
            elif rec["sup"]["kind"] == "inactive":
                warn.append("Tedarikçi kartı PASİF — onaylansa da uygulanamaz.")
            scope = _scope_of(item)
            cur = [_pref_text(q, sups.by_id) for q in ex.get(scope, [])] if scope else []
            rows.append({"id": rec["id"], "kind": kind, "row": row, "page": rec["page"],
                         "material": rec["material"], "supplier": rec["supplier"], "item": item,
                         "sup": rec["sup"], "price": rec["price"], "currency": rec["currency"],
                         "price_unit": rec["price_unit"], "package": rec["package"],
                         "vat_included": rec["vat_included"], "vat_rate": rec["vat_rate"],
                         "scope": scope, "current": cur, "warn": " ".join(warn), "note": rec["note"]})
    return rows


def _small_rows(recs: List[dict], sups: _Suppliers) -> List[dict]:
    """Küçük satıcı önerisi — lab tablosunda yalnız ≤ 1 kg ambalajla satanlar
    + planın adlandırdıkları; karışık (bazı kalemi 1 kg) satıcılar 'bilgi'."""
    lab = [r for r in recs if r["type"] == "lab_tablo"]
    per: Dict[str, dict] = {}
    for r in lab:
        k = r["sup"]["key"]
        if not k:
            continue
        d = per.setdefault(k, {"key": k, "sup": r["sup"], "names": [], "items": []})
        if r["supplier"] not in d["names"]:
            d["names"].append(r["supplier"])
        d["items"].append((r["material"], r["package"]))
    named = {sups.key(clean_supplier_name(n)) for n in SMALL_SELLER_NAMES}
    named |= {sups.key(SUPPLIER_ALIASES.get(alnum_fold(n), n)) for n in SMALL_SELLER_NAMES}
    out = []
    for k, d in per.items():
        pk = [p for _, p in d["items"] if p is not None]
        small = [p for p in pk if p <= SMALL_PKG_MAX + 1e-9]
        only_small = bool(pk) and len(small) == len(pk)
        if not (only_small or k in named or small):
            continue
        kind = "küçük satıcı" if (only_small or k in named) else "bilgi"
        disp = d["sup"]["name"]
        note_key = alnum_fold(disp) if d["sup"]["kind"] == "id" else alnum_fold(d["sup"]["name"])
        note = SMALL_SELLER_NOTES.get(note_key) or SMALL_SELLER_NOTES.get(k)
        if note and note.startswith("Doalin"):
            kind = "Doalin notu"
        if kind == "bilgi":
            note = "Yalnız bilgi — diğer kalemlerde büyük ambalajla satıyor; bu satır işlenmez."
        out.append({"id": f"KS-{k}", "key": k, "kind": kind, "names": d["names"], "sup": d["sup"],
                    "items": d["items"], "note": note or "",
                    "card_ids": [c.id for c in sups.active_cards(d["sup"]["key"])] if d["sup"]["kind"] == "id" else [],
                    "status": (normalize_status(sups.by_id[d["sup"]["id"]].purchase_status)
                               if d["sup"]["kind"] == "id" else None)})
    rank = {"küçük satıcı": 0, "Doalin notu": 1, "bilgi": 2}
    out.sort(key=lambda x: (rank[x["kind"]], alnum_fold(x["sup"]["name"])))
    return out


# ─── Yazım ──────────────────────────────────────────────────────────────────

def _delete_lab_rows(db: Session, domain: str) -> int:
    return (db.query(SupplierPrice)
            .filter(SupplierPrice.domain == domain, SupplierPrice.source.in_(LAB_SOURCES),
                    SupplierPrice.source_label.like(f"{LABEL_PREFIX}%"))
            .delete(synchronize_session=False))


def fix_errors(p: dict, approvals: dict, data: dict) -> List[str]:
    """Onaylı Excel'in DÜZELTME hücreleri ("Belirsiz okumalar → Doğru değer",
    "Eşleşmeyenler → Doğru IMS kartı") betikçe otomatik işlenmez; dolu hücre
    veri dosyasına aktarılmamışsa hata döner (hiçbir şey yazılmasın — düzeltme
    sessizce kaybolup okunan değer yazılmasın).  Aktarılmış sayılır:
      • Doğru değer: kaydın `fiyat`ı artık o sayıya eşit;
      • Doğru kart: `elle_eslesme.kart`ta kayıt ya da lab satırı var."""
    by_id = {r["id"]: r for r in p["records"]}
    kart = (data.get("elle_eslesme") or {}).get("kart") or {}
    errs: List[str] = []
    for f in approvals.get("fixes") or []:
        where = f"'{f['sheet']}' satır {f['row']} ({f['id']})"
        r = by_id.get(f["id"])
        if r is None:
            errs.append(f"{where}: kayıt bu planda yok")
            continue
        if f["sheet"] == SHEETS[3]:
            v = _num(f["value"])
            if v is not None and r["price"] is not None and abs(float(r["price"]) - v) < 1e-9:
                continue
            errs.append(f"{where}: doğru değer «{f['value']}» veri dosyasına işlenmemiş (şu an "
                        f"{_num_tr(r['price'])}) — kaydın 'fiyat'ını düzeltin ('guven': 'kesin'); "
                        "değer geçersizse hücreyi temizleyin")
        else:
            if f["id"] in kart or _row_prefix(f["id"]) in kart:
                continue
            errs.append(f"{where}: doğru kart «{f['value']}» veri dosyasına işlenmemiş — "
                        "'elle_eslesme.kart'a kart no'sunu ekleyin; gerek yoksa hücreyi temizleyin")
    return errs


def _approved_new_suppliers(p: dict, approvals: Optional[dict]) -> Tuple[List[dict], List[str]]:
    """'Yeni tedarikçiler' sayfasında E olan firmalar → (açılacaklar, hatalar).
    Yazılabilir fiyat satırı olmayan firma açılmaz (kartı boş kalırdı)."""
    if not p["new_suppliers"]:
        return [], []
    if approvals is None:
        return [], ["yeni tedarikçi kartı açmak için onaylı Excel gerekir "
                    "('Yeni tedarikçiler' → Kartı açılsın mı = E)"]
    by_id = {f"YT-{ns['key']}": ns for ns in p["new_suppliers"]}
    out, errs = [], []
    for rid, a in sorted((approvals.get("new") or {}).items(), key=lambda kv: kv[1]["row"]):
        ns = by_id.get(rid)
        if ns is None:              # kartı artık var (önceki çalıştırmada açıldı) — yapılacak yok
            continue
        if a.get("name") not in (None, "") and normalize(a["name"]) != normalize(ns["name"]):
            errs.append(f"Yeni tedarikçiler {rid}: ad '{a['name']}' ≠ '{ns['name']}' (satır kaymış?)")
            continue
        if a["ok"] is True and ns["writable"] > 0:
            out.append(ns)
    return out, errs


def apply(db: Session, data: Optional[dict] = None, *, create_suppliers: bool = False,
          approvals: Optional[dict] = None, domain: str = DOMAIN) -> dict:
    """Planı kurar; `approvals` (onaylı Excel, `read_approvals`) verilirse
    Excel hataları + işlenmemiş düzeltme hücreleri (`fix_errors`) varsa
    HİÇBİR ŞEY yazmaz (`applied` False, `errors`).  `create_suppliers` ise
    yalnız "Yeni tedarikçiler"de E olan firmaların kartını açıp yeniden eşler;
    önceki lab satırlarını silip yazılacakları ekler, TEK commit; sonra özet
    audit.  `manual` / `stok_son_durum` satırlarına dokunmaz."""
    data = data or load_data()
    p = plan(db, data, domain=domain)
    errors: List[str] = []
    if approvals is not None:
        errors += list(approvals.get("errors") or []) + fix_errors(p, approvals, data)
    to_create: List[dict] = []
    if create_suppliers:
        to_create, errs = _approved_new_suppliers(p, approvals)
        errors += errs
    if errors:
        db.rollback()
        p.update(applied=False, errors=errors)
        return p
    created: List[dict] = []
    for ns in to_create:
        s = Supplier(name=ns["name"][:150], domain=domain, is_active=True,
                     notes=f"{SUPPLIER_NOTE} — notta: {', '.join(ns['spellings'])}")
        db.add(s)
        db.flush()
        created.append({"id": s.id, "name": s.name})
    if created:
        p = plan(db, data, domain=domain)
    deleted = _delete_lab_rows(db, domain)
    written = 0
    for r in p["records"]:
        if r["status"] != "yazilacak":
            continue
        sup = db.query(Supplier).filter(Supplier.id == r["sup"]["id"]).first()
        db.add(SupplierPrice(
            item_id=r["item"]["id"], supplier_id=sup.id, supplier_name=sup.name,
            package_size=r["package"], unit_price=float(r["price"]), currency=r["currency"],
            price_unit=r["price_unit"], source=r["source"],
            source_label=price_label(r["page"]), quoted_at=r["date"],
            note=r["write_note"], vat_included=r["vat_included"],
            vat_rate=(float(r["vat_rate"]) if r["vat_rate"] is not None else None),
            created_by=ACTOR, domain=domain))
        written += 1
    db.commit()
    c = p["counts"]
    log_admin_event(
        db, None, actor=None, action=AUDIT_IMPORT, target_type="supplier_prices",
        target_name=LABEL_PREFIX,
        details={"actor": ACTOR, "script": Path(__file__).name, "domain": domain,
                 "deleted_previous": deleted, "written": written,
                 "created_suppliers": created,
                 "skipped": {k: c[k] for k in STATUS_TEXT if k != "yazilacak" and c.get(k)},
                 "skipped_manual": [r["id"] for r in p["records"] if r["status"] == "elle_var"]})
    p.update(applied=True, deleted=deleted, written=written, created_suppliers=created, errors=[])
    return p


# ─── Kontrol Excel'i ────────────────────────────────────────────────────────

def _sup_text(s: dict) -> str:
    if s["kind"] == "id":
        return f"{s['id']} {s['name']}" + ("" if s["how"] == "ad" else f" ({s['how']})")
    if s["kind"] == "inactive":
        return f"PASİF: {s['id']} {s['name']}"
    return f"YENİ: {s['name']}"


def _vat_text(r: dict) -> str:
    if r["vat_included"] is True:
        return "dahil"
    if r["vat_included"] is False:
        return "hariç"
    return ""


def _mat_how(r: dict) -> str:
    m = r["mat"]
    return {"ok": m.get("how") or "", "belirsiz": "birden çok aday", "eslesmedi": "bulunamadı",
            "ayri_form": "ayrı form"}.get(m["status"], "")


def _fix_na(r: dict) -> bool:
    """Düzeltme hücresi anlamsız mı (kayıt zaten yazılmayacak: fiyat listesi /
    üstü çizili)."""
    return r["status"] in ("kapsam_disi", "cizili")


def build_xlsx(p: dict, path) -> Path:
    """Kontrol Excel'ini yazar (okunaklı: başlık dondurulmuş, filtreli, sütun
    genişlikleri ayarlı, onay hücrelerinde E/H listesi — liste dışı giriş
    Excel'de REDDEDİLİR, `showErrorMessage`).  Onay/düzeltme hücresi olmayan
    satırda "—" yazar (okuyucu atlar)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="232E6E")
    input_fill = PatternFill("solid", fgColor="FFF4C2")
    warn_font = Font(color="9A3412")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, cols, rows, *, inputs=(), yesno=None, warn_cols=()):
        ws = wb.create_sheet(title)
        for ci, (name, width) in enumerate(cols, 1):
            c = ws.cell(1, ci, name)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(ci)].width = width
        names = [n for n, _ in cols]
        for ri, row in enumerate(rows, 2):
            for ci, v in enumerate(row, 1):
                c = ws.cell(ri, ci, v)
                c.alignment = wrap if cols[ci - 1][1] >= 30 else top
                if names[ci - 1] in inputs and v != NOT_APPLICABLE:
                    c.fill = input_fill
                if names[ci - 1] in warn_cols and v:
                    c.font = warn_font
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32
        last = max(len(rows) + 1, 2)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
        if yesno and rows:
            # showErrorMessage varsayılanı False — açılmazsa Excel her girişi
            # (X, ✓ …) sessizce kabul eder; "stop" yanlış girişi reddeder.
            dv = DataValidation(type="list", formula1='"E,H"', allow_blank=True,
                                showErrorMessage=True, errorStyle="stop",
                                errorTitle="E ya da H", error="Lütfen yalnız E (evet) ya da H (hayır) yazın.")
            col = get_column_letter(names.index(yesno) + 1)
            dv.add(f"{col}2:{col}{len(rows) + 1}")
            ws.add_data_validation(dv)
        return ws

    c = p["counts"]
    # ── Nasıl doldurulur ──
    ws = wb.create_sheet(SHEETS[0])
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Lab fiyat notları (07.10.2026) — kontrol listesi", True),
        (f"Kayıt: {c['kayit']} · yazılacak fiyat satırı: {c['yazilacak']} · eşleşmeyen/belirsiz/ayrı form: "
         f"{c['malzeme_eslesmeyen']}/{c['malzeme_belirsiz']}/{c['ayri_form']} · yeni tedarikçi: "
         f"{c['yeni_firma']} · okuma belirsiz: {c['okuma_belirsiz']}", False),
        ("", False),
        ("Nasıl doldurulur", True),
        ("Onay sütunlarına YALNIZ E (evet) ya da H (hayır) yazın — X, ✓, 1 gibi işaretler kabul edilmez. "
         "Boş bırakılan satır işlenmez. '—' olan hücre o satırda doldurulmaz.", True),
        ("1) 'Tercih önerisi' sayfası: tablodaki SARI seçimler (+ notlardan alternatifler). Bu malzeme bu "
         "firmadan alınsın diyorsanız 'Onay (E/H)' sütununa E, istemiyorsanız H yazın.", False),
        ("   Aynı malzemede birden çok E varsa üstteki satır 1. tercih, alttaki 2. tercih olur. "
         "'Uyarı' sütunundaki çelişkilere (ör. Aktif Kömür, Itır, bergamot, misk adaçayı, 'Uludağ ucuz "
         "olanları alacağız!') özellikle bakın. Fiyat sütununda 'KDV dahil' yazan fiyat brüttür; "
         "karşılaştırma parantezdeki net fiyatla yapılır.", False),
        ("2) 'Küçük satıcı önerisi' sayfası: yalnız 1 kg'lık ambalajla satan firmalar. Bu firmadan artık "
         "alınmasın ('bitirilecek — alma') diyorsanız 'Bitirilecek yapılsın mı (E/H)' sütununa E yazın. "
         "Aynı firmaya hem tercih hem bitirilecek E'si verilemez.", False),
        ("3) 'Eşleşmeyenler' sayfası: sistemde kartı bulunamayan ya da birden çok kartla eşleşen malzemeler. "
         "Doğru kartı biliyorsanız 'Doğru IMS kartı' sütununa kart adını/numarasını yazın.", False),
        ("4) 'Belirsiz okumalar' sayfası: el yazısı net okunamayan fiyatlar. Doğru değeri biliyorsanız "
         "'Doğru değer' sütununa yazın (ör. 7,15). Boş bıraktığınız satırda okunan değer, notunda "
         "'okuma belirsiz' yazarak sisteme girer.", False),
        ("5) 'Yeni tedarikçiler' sayfası: sistemde kartı olmayan firmalar. Kartı açılsın mı sütununa E ya da "
         "H yazın; yalnız E olanların kartı açılır ve fiyatları yazılır.", False),
        ("Doldurduktan sonra dosyayı geri gönderin. E/H onayları bu dosyadan işlenir; 'Doğru değer' ve "
         "'Doğru IMS kartı' düzeltmeleri önce elle sisteme aktarılır — aktarılmadan hiçbir fiyat ya da "
         "tercih yazılmaz.", True),
        ("", False),
        ("Bilgi", True),
        ("• 'Fiyatlar' sayfası bütün kayıtları gösterir: nottaki malzeme → eşleşen sistem kartı, nottaki "
         "tedarikçi → sistemdeki firma, yazılacak mı / neden yazılmayacak.", False),
        ("• Aynı kart + firma için tek fiyat tutulur: fatura/sipariş/proforma (en yeni) lab tablosundan "
         "önce gelir; ezilen değer kazanan satırın notuna yazılır.", False),
        ("• Sistemde elle girilmiş fiyatı olan kart + firma ezilmez (elle girilen kazanır).", False),
        ("• Pharmaterm uçucu yağ fiyatları tabloda KDV DAHİL; satın alma planı bunları net fiyatla "
         "karşılaştırır.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        cell = ws.cell(i, 1, t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True, size=12 if i == 1 else 11)

    recs = p["records"]
    # ── Fiyatlar ──
    cols = [("Kayıt", 11), ("Sayfa", 7), ("Tür", 13), ("Tarih", 11), ("Nottaki malzeme", 30),
            ("IMS kartı (no ad birim)", 36), ("Eşleşme", 16), ("Nottaki tedarikçi", 22),
            ("IMS tedarikçisi", 28), ("Fiyat", 10), ("Para birimi", 9), ("Birim", 7),
            ("Ambalaj", 9), ("KDV", 8), ("KDV %", 7), ("Seçili (sarı)", 9), ("Güven", 9),
            ("Yazılacak mı / neden", 40), ("Uyarılar", 40), ("Not", 60)]
    rows = []
    for r in recs:
        rows.append([r["id"], r["page"], TYPE_TEXT[r["type"]], _dmy(r["date"]) if r["date"] else "",
                     r["material"], p_label(r), _mat_how(r), r["supplier"], _sup_text(r["sup"]),
                     r["price"], r["currency"], r["price_unit"], r["package"], _vat_text(r),
                     r["vat_rate"], "evet" if r["selected"] else "", r["confidence"],
                     STATUS_TEXT.get(r["status"], r["status"]), "; ".join(r["flags"]), r["write_note"]])
    sheet(SHEETS[1], cols, rows, warn_cols=("Uyarılar",))

    # ── Eşleşmeyenler ──
    mats_label = {}
    for r in recs:
        if r["item"]:
            mats_label[r["item"]["id"]] = p_label(r)
    cols = [("Kayıt", 11), ("Sayfa", 7), ("Tür", 13), ("Nottaki malzeme", 32), ("Nottaki tedarikçi", 22),
            ("Sorun", 34), ("Adaylar / benzer kartlar", 50), (FIX_CARD_COL, 26), ("Not", 50)]
    rows = []
    for r in recs:
        m = r["mat"]
        if m["status"] == "ok" and r["status"] != "uyumsuz":
            continue
        if m["status"] == "ayri_form":
            prob, cands = "Ayrı form — " + m.get("reason", ""), ""
        elif m["status"] == "belirsiz":
            prob = "Birden çok aday kart" + (f" — {m['reason']}" if m.get("reason") else "")
            cands = ", ".join(p["cand_names"].get(i, str(i)) for i in m["candidates"])
        elif m["status"] == "eslesmedi":
            prob = "Kart bulunamadı"
            cands = ("Benzer: " + ", ".join(p["cand_names"].get(i, str(i)) for i in m["suggestions"])
                     if m["suggestions"] else "")
        else:
            prob, cands = "Fiyat birimi kartla uyumsuz — " + "; ".join(r["flags"]), p_label(r)
        if _fix_na(r):
            prob = f"{prob} ({STATUS_TEXT[r['status']]})"
        rows.append([r["id"], r["page"], TYPE_TEXT[r["type"]], r["material"], r["supplier"], prob, cands,
                     NOT_APPLICABLE if _fix_na(r) else None, r["note"]])
    for ns in p["new_suppliers"]:
        rows.append(["", "", "", ", ".join(ns["materials"][:6]) + (" …" if len(ns["materials"]) > 6 else ""),
                     ", ".join(ns["spellings"]), "Tedarikçi sistemde yok (bkz. 'Yeni tedarikçiler')",
                     "", NOT_APPLICABLE, f"{len(ns['records'])} kayıt"])
    sheet(SHEETS[2], cols, rows, inputs=(FIX_CARD_COL,))

    # ── Belirsiz okumalar ──
    cols = [("Kayıt", 11), ("Sayfa", 7), ("Nottaki malzeme", 30), ("Nottaki tedarikçi", 22),
            ("Okunan değer", 12), ("Para birimi", 9), ("Güven", 9), ("Karşılaştırma", 36),
            ("Yazılacak mı / neden", 34), (FIX_VALUE_COL, 16), ("Açıklama", 70)]
    rows = []
    for r in recs:
        if r["confidence"] != "belirsiz" and r["id"] not in EXTRA_CHECK:
            continue
        rows.append([r["id"], r["page"], r["material"], r["supplier"], r["price"], r["currency"],
                     r["confidence"] if r["confidence"] == "belirsiz" else "kesin (notla çelişen)",
                     EXTRA_CHECK.get(r["id"], ""), STATUS_TEXT.get(r["status"], r["status"]),
                     NOT_APPLICABLE if _fix_na(r) else None, r["note"]])
    sheet(SHEETS[3], cols, rows, inputs=(FIX_VALUE_COL,))

    # ── Tercih önerisi ──
    cols = [("Kayıt", 11), ("Öneri türü", 15), ("Sayfa", 7), ("Malzeme (not)", 30), ("IMS kartı", 32),
            ("Tercih kapsamı", 26), ("Tedarikçi (not)", 22), ("IMS tedarikçisi", 24), ("Fiyat", 22),
            ("Mevcut tercih", 26), (APPROVE_COL, 10), ("Uyarı", 50), ("Not", 60)]
    rows = []
    for x in p["prefs"]:
        scope = ""
        if x["item"]:
            scope = (f"grup «{x['item']['group_name']}» (bütün kartları)" if x["item"]["group_id"]
                     else "yalnız bu kart")
        rows.append([x["id"], x["kind"], x["page"], x["material"],
                     f"{x['item']['id']} {x['item']['name']}" if x["item"] else "— eşleşmedi", scope,
                     x["supplier"], _sup_text(x["sup"]), _price_text(x),
                     "; ".join(x["current"]), None, x["warn"], x["note"]])
    sheet(SHEETS[4], cols, rows, inputs=(APPROVE_COL,), yesno=APPROVE_COL, warn_cols=("Uyarı",))

    # ── Küçük satıcı önerisi ──
    cols = [("Kayıt", 22), ("Öneri türü", 14), ("Tedarikçi (not)", 26), ("IMS tedarikçisi", 26),
            ("Şu anki durum", 14), ("Lab tablosundaki kalemler (ambalaj kg)", 60), (PHASE_COL, 14),
            ("Not", 60)]
    rows = []
    for s in p["small"]:
        items = "; ".join(f"{m} ({_num_tr(k) if k is not None else '—'})" for m, k in s["items"])
        rows.append([s["id"], s["kind"], ", ".join(s["names"]), _sup_text(s["sup"]),
                     {"phase_out": "bitirilecek", "preferred": "tercih edilen", "normal": "normal"}
                     .get(s["status"] or "", "pasif kart" if s["sup"]["kind"] == "inactive" else "kart yok"),
                     items, "—" if s["kind"] == "bilgi" else None, s["note"]])
    sheet(SHEETS[5], cols, rows, inputs=(PHASE_COL,), yesno=PHASE_COL)

    # ── Yeni tedarikçiler ──  (yazılabilir satırı olmayan firma açılmaz → "—")
    cols = [("Kayıt", 22), ("Açılacak ad", 24), ("Nottaki yazımlar", 32), ("Kayıt sayısı", 10),
            ("Yazılabilecek fiyat satırı", 12), (NEW_COL, 14), ("Malzemeler", 70), ("Kayıtlar", 50)]
    rows = [[f"YT-{ns['key']}", ns["name"], ", ".join(ns["spellings"]), len(ns["records"]), ns["writable"],
             None if ns["writable"] > 0 else NOT_APPLICABLE,
             ", ".join(ns["materials"]), ", ".join(ns["records"])] for ns in p["new_suppliers"]]
    sheet(SHEETS[6], cols, rows, inputs=(NEW_COL,), yesno=NEW_COL)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def p_label(r: dict) -> str:
    it = r["item"]
    return f"{it['id']} {it['name']} ({it['unit'] or 'birimsiz'})" if it else "—"


# ─── Onaylı Excel → tercih + bitirilecek ────────────────────────────────────

def read_approvals(path) -> dict:
    """Onaylı Excel → {prefs: {kayıt: {ok, material, supplier, row}},
    small: {kayıt: {ok, row}}, new: {YT-…: {ok, name, row}},
    fixes: [{sheet, id, value, row}], errors: [...]}.

    E/H dışı onay değeri hata listesine (`_yes_no`).  Düzeltme sayfalarının
    ("Eşleşmeyenler → Doğru IMS kartı", "Belirsiz okumalar → Doğru değer")
    DOLU hücreleri `fixes`e girer — otomatik işlenmez, `fix_errors` veri
    dosyasına aktarılıp aktarılmadığını denetler."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {"prefs": {}, "small": {}, "new": {}, "fixes": [], "errors": []}

    def table(title, col):
        if title not in wb.sheetnames:
            out["errors"].append(f"'{title}' sayfası yok")
            return None
        it = wb[title].iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, ())]
        if "Kayıt" not in head or col not in head:
            out["errors"].append(f"'{title}': 'Kayıt' / '{col}' sütunu yok")
            return None
        return head, it

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    for title, col, dest, extra in ((SHEETS[4], APPROVE_COL, "prefs", ("Malzeme (not)", "Tedarikçi (not)")),
                                    (SHEETS[5], PHASE_COL, "small", ()),
                                    (SHEETS[6], NEW_COL, "new", ("Açılacak ad",))):
        t = table(title, col)
        if t is None:
            continue
        head, it = t
        ik, ic = head.index("Kayıt"), head.index(col)
        ix = {k: (head.index(k) if k in head else None) for k in extra}
        for n, row in enumerate(it, 2):
            rid = cell(row, ik)
            if not rid:
                continue
            raw = cell(row, ic)
            if isinstance(raw, str) and raw.strip() == NOT_APPLICABLE:
                continue
            try:
                ok = _yes_no(raw)
            except ValueError:
                out["errors"].append(f"'{title}' satır {n} ({rid}): '{raw}' — yalnız E ya da H olmalı")
                continue
            rec = {"ok": ok, "row": n}
            if dest == "prefs":
                rec.update(material=cell(row, ix["Malzeme (not)"]), supplier=cell(row, ix["Tedarikçi (not)"]))
            elif dest == "new":
                rec.update(name=cell(row, ix["Açılacak ad"]))
            out[dest][str(rid)] = rec
    for title, col in ((SHEETS[2], FIX_CARD_COL), (SHEETS[3], FIX_VALUE_COL)):
        t = table(title, col)
        if t is None:
            continue
        head, it = t
        ik, ic = head.index("Kayıt"), head.index(col)
        for n, row in enumerate(it, 2):
            v = cell(row, ic)
            if v is None or (isinstance(v, str) and v.strip() in ("", NOT_APPLICABLE)):
                continue
            rid = cell(row, ik)
            if not rid:
                out["errors"].append(f"'{title}' satır {n}: '{col}' dolu ama kayıt no'su yok — elle işleyin")
                continue
            out["fixes"].append({"sheet": title, "id": str(rid), "value": v, "row": n})
    wb.close()
    return out


def plan_prefs(db: Session, approvals: dict, data: Optional[dict] = None, *, domain: str = DOMAIN) -> dict:
    """Onaylardan yapılacakları çıkarır — YAZMAZ.
    → {prefs: [{id, action: create|skip, reason, scope, supplier_id, rank, …}],
       phase_out: [{id, action, reason, cards: [...]}], counts, errors}.

    Hata (→ `apply_prefs` hiçbir şey yazmaz): Excel değer hataları, veri
    dosyasına aktarılmamış düzeltme hücreleri (`fix_errors`), satır kayması,
    ve ÇELİŞKİ — aynı firma hem tercih (E) hem bitirilecek (E) ya da tercih
    edilen firma zaten bitirilecek: satın alma planında bitirilecek tercihten
    önce gelir (`purchase_pricing._offer_status`), onay sessizce etkisiz kalırdı."""
    data = data or load_data()
    p = plan(db, data, domain=domain)
    errors = list(approvals.get("errors") or []) + fix_errors(p, approvals, data)
    pref_by_id = {x["id"]: x for x in p["prefs"]}
    small_by_id = {s["id"]: s for s in p["small"]}
    ex = _existing_prefs(db, domain)
    sups = {s.id: s for s in db.query(Supplier).filter(Supplier.domain == domain).all()}
    base_rank: Dict[tuple, int] = {}
    kept_text: Dict[tuple, List[str]] = {}
    taken = set()
    actions: List[dict] = []
    # Excel satır sırası = rank sırası
    for rid, a in sorted((approvals.get("prefs") or {}).items(), key=lambda kv: kv[1]["row"]):
        x = pref_by_id.get(rid)
        if x is None:
            errors.append(f"Tercih önerisi: '{rid}' bu planda yok")
            continue
        if a.get("material") not in (None, "") and normalize(a["material"]) != normalize(x["material"]):
            errors.append(f"Tercih önerisi {rid}: malzeme '{a['material']}' ≠ '{x['material']}' (satır kaymış?)")
            continue
        if a.get("supplier") not in (None, "") and normalize(a["supplier"]) != normalize(x["supplier"]):
            errors.append(f"Tercih önerisi {rid}: tedarikçi '{a['supplier']}' ≠ '{x['supplier']}' (satır kaymış?)")
            continue
        act = {"id": rid, "material": x["material"], "supplier": x["supplier"], "ok": a["ok"],
               "scope": x["scope"], "item": x["item"], "supplier_id": x["sup"]["id"],
               "supplier_name": x["sup"]["name"], "action": "skip", "reason": "", "rank": None}
        actions.append(act)
        if a["ok"] is not True:
            act["reason"] = "H — işlenmez" if a["ok"] is False else "boş — işlenmez"
            continue
        if x["item"] is None:
            act["reason"] = "IMS kartı eşleşmedi"
            continue
        if x["sup"]["kind"] != "id":
            act["reason"] = ("tedarikçi kartı pasif" if x["sup"]["kind"] == "inactive"
                             else "tedarikçi kartı yok (önce --commit --onayli … --create-suppliers)")
            continue
        sc = x["scope"]
        cur = ex.get(sc, [])
        same = next((q for q in cur if q.supplier_id == x["sup"]["id"]), None)
        if same is not None:
            act["reason"] = ("zaten var: " + ("tercih " + str(same.rank) if same.preference == "preferred"
                                               else "bu malzemede alma") + " — dokunulmadı")
            continue
        if (sc, x["sup"]["id"]) in taken:
            act["reason"] = "aynı kapsam + tedarikçi bu Excel'de zaten onaylandı"
            continue
        if sc not in base_rank:
            kept = [q for q in cur if q.preference == "preferred"]
            base_rank[sc] = max([int(q.rank or 1) for q in kept], default=0)
            kept_text[sc] = [_pref_text(q, sups) for q in kept]
        base_rank[sc] += 1
        taken.add((sc, x["sup"]["id"]))
        act.update(action="create", rank=base_rank[sc],
                   reason=("mevcut tercih korundu (" + ", ".join(kept_text[sc]) + "), arkasına eklenir"
                           if kept_text[sc] else ""))
    phase: List[dict] = []
    for rid, a in sorted((approvals.get("small") or {}).items(), key=lambda kv: kv[1]["row"]):
        s = small_by_id.get(rid)
        if s is None:
            errors.append(f"Küçük satıcı önerisi: '{rid}' bu planda yok")
            continue
        act = {"id": rid, "name": s["sup"]["name"], "ok": a["ok"], "action": "skip", "reason": "", "cards": []}
        phase.append(act)
        if a["ok"] is not True:
            act["reason"] = "H — işlenmez" if a["ok"] is False else "boş — işlenmez"
            continue
        if s["kind"] == "bilgi":
            act["reason"] = "bilgi satırı — işlenmez"
            continue
        if s["sup"]["kind"] != "id":
            act["reason"] = "tedarikçi kartı pasif" if s["sup"]["kind"] == "inactive" else "tedarikçi kartı yok"
            continue
        cards = [sups[i] for i in s["card_ids"] if i in sups and sups[i].is_active is not False]
        todo = [c for c in cards if normalize_status(c.purchase_status) != "phase_out"]
        if not todo:
            act["reason"] = "zaten bitirilecek"
            continue
        act.update(action="phase_out", cards=[{"id": c.id, "name": c.name,
                                               "old": normalize_status(c.purchase_status) or "normal"}
                                              for c in todo])
    # Çelişki: tercih edilen firma bu Excel'de bitirilecek ya da zaten bitirilecek
    po_cards = {c["id"]: a["id"] for a in phase if a["action"] == "phase_out" for c in a["cards"]}
    for a in actions:
        if a["action"] != "create":
            continue
        sid = a["supplier_id"]
        if sid in po_cards:
            errors.append(f"Tercih önerisi {a['id']}: {a['supplier_name']} aynı Excel'de 'bitirilecek' olarak da "
                          f"onaylandı ({po_cards[sid]}) — bitirilecek tercihi geçersiz kılar; birine H yazın")
        elif sid in sups and normalize_status(sups[sid].purchase_status) == "phase_out":
            errors.append(f"Tercih önerisi {a['id']}: {a['supplier_name']} zaten 'bitirilecek' — tercih "
                          "etkisiz kalır; tedarikçinin durumunu değiştirin ya da bu satıra H yazın")
    counts = {"tercih_yazilacak": sum(1 for a in actions if a["action"] == "create"),
              "tercih_atlanan": sum(1 for a in actions if a["action"] != "create" and a["ok"] is True),
              "bitirilecek": sum(1 for a in phase if a["action"] == "phase_out")}
    return {"prefs": actions, "phase_out": phase, "counts": counts, "errors": errors}


def apply_prefs(db: Session, approvals: dict, data: Optional[dict] = None, *, domain: str = DOMAIN) -> dict:
    """`plan_prefs` hatasızsa tercihleri ve bitirilecek durumlarını yazar
    (TEK commit), sonra satır başına audit.  Hata varsa hiçbir şey yazmaz."""
    pp = plan_prefs(db, approvals, data, domain=domain)
    if pp["errors"]:
        db.rollback()
        pp["applied"] = False
        return pp
    made: List[Tuple[MaterialSupplierPref, dict]] = []
    for a in pp["prefs"]:
        if a["action"] != "create":
            continue
        kind, sid = a["scope"]
        pref = MaterialSupplierPref(
            domain=domain, material_group_id=sid if kind == "group" else None,
            item_id=sid if kind == "item" else None, supplier_id=a["supplier_id"],
            preference="preferred", rank=a["rank"], note=f"{PREF_NOTE} ({a['id']})", created_by=ACTOR)
        db.add(pref)
        made.append((pref, a))
    now = datetime.utcnow()
    changed: List[Tuple[Supplier, str, dict]] = []
    for a in pp["phase_out"]:
        if a["action"] != "phase_out":
            continue
        for c in a["cards"]:
            s = db.query(Supplier).filter(Supplier.id == c["id"]).first()
            old_reason = s.status_reason
            s.purchase_status = "phase_out"
            s.status_reason = PHASE_OUT_REASON
            s.status_by = ACTOR
            s.status_at = now
            changed.append((s, c["old"], {"eski_sebep": old_reason}))
    db.commit()
    for pref, a in made:
        db.refresh(pref)
        log_admin_event(db, None, actor=None, action="material_pref.create",
                        target_type="material_pref", target_id=pref.id,
                        target_name=(a["item"]["group_name"] or a["item"]["name"])[:150],
                        details={"actor": ACTOR, "script": Path(__file__).name,
                                 "kapsam": "group" if pref.material_group_id else "item",
                                 "grup_id": pref.material_group_id, "kart_id": pref.item_id,
                                 "tedarikci_id": pref.supplier_id, "tedarikci": a["supplier_name"],
                                 "tercih": "preferred", "sira": pref.rank, "kayit": a["id"],
                                 "domain": domain})
    for s, old, extra in changed:
        log_admin_event(db, None, actor=None, action="supplier.status", target_type="supplier",
                        target_id=s.id, target_name=s.name,
                        details={"actor": ACTOR, "script": Path(__file__).name, "eski": old,
                                 "yeni": "phase_out", "sebep": PHASE_OUT_REASON, "domain": domain,
                                 **extra})
    log_admin_event(db, None, actor=None, action=AUDIT_PREFS, target_type="material_pref",
                    target_name=LABEL_PREFIX,
                    details={"actor": ACTOR, "script": Path(__file__).name, "domain": domain,
                             "tercih": [a["id"] for _, a in made],
                             "bitirilecek": [s.name for s, _, _ in changed]})
    pp["applied"] = True
    return pp


# ─── Konsol ─────────────────────────────────────────────────────────────────

def _print_plan(p: dict, *, commit: bool, create: bool) -> None:
    c = p["counts"]
    print("=== Lab fiyat notları 07.10.2026 — içe aktarma ===")
    print(f"  Kayıt                 : {c['kayit']}")
    print(f"  Malzeme eşleşen       : {c['malzeme_eslesen']}  · bulunamayan {c['malzeme_eslesmeyen']}"
          f" · belirsiz {c['malzeme_belirsiz']} · ayrı form {c['ayri_form']}")
    print(f"  Tedarikçi eşleşen     : {c['tedarikci_eslesen']} kayıt · yeni firma {c['yeni_firma']}")
    print(f"  Okuma belirsiz        : {c['okuma_belirsiz']}")
    print(f"  Ezilen (öncelik)      : {c['ezildi']}  · elle fiyat var {c['elle_var']}"
          f" · birim uyumsuz {c['uyumsuz']} · fiyatsız {c['fiyatsiz']} · pasif tedarikçi {c['pasif_tedarikci']}")
    print(f"  Yazılmayan (belge)    : üstü çizili {c['cizili']} · fiyat listesi (kapsam dışı) {c['kapsam_disi']}")
    print(f"  Yeni firma bekleyen   : {c['yeni_tedarikci']} satır ({c['yeni_firma']} firma)"
          + (" — onaylı Excel'de E + --create-suppliers ile yazılır" if not create else ""))
    print(f"  YAZILACAK satır       : {c['yazilacak']}  (silinecek önceki lab satırı: {p['existing_lab_rows']})")
    print(f"  Tercih önerisi        : {len(p['prefs'])} satır · küçük satıcı önerisi: "
          f"{sum(1 for s in p['small'] if s['kind'] != 'bilgi')}")
    if p["new_suppliers"]:
        print("\n  Yeni tedarikçiler (sistemde kartı yok; yalnız 'Kartı açılsın mı'=E olanlar açılır):")
        for ns in p["new_suppliers"]:
            print(f"    · YT-{ns['key']} {ns['name']}  ({', '.join(ns['spellings'])}) — {len(ns['records'])} "
                  f"kayıt, yazılabilir {ns['writable']}" + ("" if ns["writable"] else " → AÇILMAZ"))
    bad = [r for r in p["records"] if r["mat"]["status"] != "ok"]
    if bad:
        print("\n  Eşleşmeyen / belirsiz malzemeler:")
        for r in bad:
            print(f"    · {r['id']} {r['material']} — {STATUS_TEXT.get(r['status'], r['status'])}")
    man = [r for r in p["records"] if r["status"] == "elle_var"]
    if man:
        print("\n  Elle fiyatı olduğu için yazılmayacaklar:")
        for r in man:
            print(f"    · {r['id']} {r['material']} / {r['supplier']} — {'; '.join(r['flags'])}")
    print()
    for e in p.get("errors") or []:
        print(f"  ✖ {e}")
    if p.get("errors"):
        print("⛔ DURDURULDU — onaylı Excel hataları / işlenmemiş düzeltmeler giderilmeden hiçbir şey yazılmaz.")
    elif p.get("applied"):
        print(f"✓ YAZILDI — {p['written']} satır (önceki {p['deleted']} lab satırı silindi)"
              + (f"; açılan tedarikçi: {', '.join(s['name'] for s in p['created_suppliers'])}"
                 if p.get("created_suppliers") else ""))
    elif not commit:
        print("KURU ÇALIŞTIRMA — hiçbir şey yazılmadı.  Yazmak için: --onayli ONAYLI.xlsx --commit"
              + ("" if create else " (E onaylı yeni firmalar için + --create-suppliers)"))


def _print_prefs(pp: dict, *, commit: bool) -> None:
    print("=== Onaylı Excel — tercih + bitirilecek ===")
    for e in pp["errors"]:
        print(f"  ✖ {e}")
    for a in pp["prefs"]:
        if a["ok"] is not True:
            continue
        mark = "+" if a["action"] == "create" else "·"
        extra = f" → rank {a['rank']}" if a["rank"] else ""
        print(f"  {mark} {a['id']} {a['material']} → {a['supplier_name']}{extra}"
              + (f" — {a['reason']}" if a["reason"] else ""))
    for a in pp["phase_out"]:
        if a["ok"] is not True:
            continue
        mark = "+" if a["action"] == "phase_out" else "·"
        cards = ", ".join(f"{c['id']} {c['name']} ({c['old']} → bitirilecek)" for c in a["cards"])
        print(f"  {mark} {a['id']} {a['name']}" + (f": {cards}" if cards else f" — {a['reason']}"))
    c = pp["counts"]
    no = sum(1 for a in pp["prefs"] + pp["phase_out"] if a["ok"] is False)
    blank = sum(1 for a in pp["prefs"] + pp["phase_out"] if a["ok"] is None)
    print(f"\n  Tercih yazılacak: {c['tercih_yazilacak']} · atlanan (E ama uygulanamaz): {c['tercih_atlanan']}"
          f" · bitirilecek: {c['bitirilecek']} · H: {no} · boş: {blank}")
    if pp["errors"]:
        print("⛔ DURDURULDU — Excel hataları düzeltilmeden hiçbir şey yazılmaz.")
    elif pp.get("applied"):
        print("✓ YAZILDI.")
    elif not commit:
        print("KURU ÇALIŞTIRMA — hiçbir şey yazılmadı.  Yazmak için: --commit")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Lab fiyat notları (07.10.2026) → tedarikçi fiyatları")
    ap.add_argument("--commit", action="store_true", help="yaz (varsayılan: kuru çalıştırma)")
    ap.add_argument("--onayli", metavar="ONAYLI_XLSX",
                    help="Songül Hanım'ın onayladığı kontrol Excel'i (fiyat --commit için ŞART)")
    ap.add_argument("--create-suppliers", action="store_true",
                    help="onaylı Excel'de 'Kartı açılsın mı'=E olan yeni tedarikçilerin kartını aç")
    ap.add_argument("--xlsx", metavar="YOL", help="kontrol Excel'ini yaz")
    ap.add_argument("--apply-prefs", metavar="ONAYLI_XLSX", help="onaylı Excel'den tercih + bitirilecek")
    ap.add_argument("--data", metavar="JSON", help="veri dosyası (varsayılan scripts/data/…)")
    args = ap.parse_args(sys.argv[1:] if argv is None else argv)
    try:
        data = load_data(args.data)
    except (OSError, ValueError) as e:
        print(f"⛔ veri dosyası okunamadı: {e}")
        return 2
    if args.commit and not args.apply_prefs and not args.onayli:
        print("⛔ Fiyat yazımı (--commit) onaylı kontrol Excel'i ister: --onayli ONAYLI.xlsx "
              "(Songül Hanım'ın düzeltmeleri ve yeni firma onayları oradan denetlenir)")
        return 2
    if args.create_suppliers and not args.onayli:
        print("⛔ --create-suppliers yalnız --onayli ONAYLI.xlsx ile ('Kartı açılsın mı'=E olan firmalar)")
        return 2
    approvals = None
    xlsx_in = args.apply_prefs or args.onayli
    if xlsx_in:
        try:
            approvals = read_approvals(xlsx_in)
        except Exception as e:                           # bozuk / yanlış dosya
            print(f"⛔ Excel okunamadı: {e}")
            return 2
    db = SessionLocal()
    try:
        if args.apply_prefs:
            pp = apply_prefs(db, approvals, data) if args.commit else plan_prefs(db, approvals, data)
            _print_prefs(pp, commit=args.commit)
            return 2 if pp["errors"] else 0
        p = plan(db, data)
        if args.xlsx:
            out = build_xlsx(p, args.xlsx)
            print(f"Kontrol Excel'i yazıldı: {out}")
        if args.commit:
            if approvals is not None and not args.create_suppliers:
                waiting, _ = _approved_new_suppliers(p, approvals)
                if waiting:
                    print("  ⚠ E onaylı yeni tedarikçi(ler) AÇILMADI (--create-suppliers yok), satırları "
                          "yazılmayacak: " + ", ".join(ns["name"] for ns in waiting))
            p = apply(db, data, create_suppliers=args.create_suppliers, approvals=approvals)
        elif approvals is not None:                      # kuru, onaylı Excel denetimi
            errs = list(approvals["errors"]) + fix_errors(p, approvals, data)
            to_open, nerr = _approved_new_suppliers(p, approvals)
            p["errors"] = errs + nerr
            if to_open:
                print("  Onaylı yeni tedarikçiler (--create-suppliers ile açılır): "
                      + ", ".join(ns["name"] for ns in to_open))
        _print_plan(p, commit=args.commit, create=args.create_suppliers)
        if p.get("errors"):
            return 2
    except ValueError as e:                              # veri dosyası / elle eşleme hatası
        db.rollback()
        print(f"⛔ DURDURULDU — {e}; hiçbir şey yazılmadı")
        return 2
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
