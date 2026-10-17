#!/usr/bin/env python3
"""Fiyat tamamlama kademesi (16.10.2026) — kuru çalıştırma varsayılan.

**Hedef:** aktif reçetelerde geçen, hiçbir tedarikçi fiyatı olmayan HAMMADDE
kartları (`core.price_overview.build_overview(has_price="unpriced")` — fiyat
sayfasının "fiyatı olmayanlar" mantığı: elle fiyatla gizlenen satır, plan
ikizi, pasif kart kuralları aynı).  Ambalaj ve etiket bu kademeye GİRMEZ
(kullanıcı `/fiyatlar`'dan kendisi girer).  Yalnız grup üzerinden fiyatı olan
kart formda "grup üzerinden fiyatlı" notuyla kalır (sayı ayrı verilir).

**Kademe (sıra):**
  1. **Paraşüt alış faturası satırları** — `parasut_purchase_lines` tablosundaki
     BEKLEYEN satırlar (gece okuması, `core.parasut_purchases`); ya da
     `--parasut-json DOSYA` ile döküm: JSON:API yükü (`{"data": [...],
     "included": [...]}`, `parse_bill` beyaz listesi) veya `parasut_read.py`
     düz dökümü (`{"bills": [...]}` — yalnız izin listesindeki alanlar alınır,
     `supplier_contact` ATILIR, satır indirimi net toplamdan türetilir;
     kimliği olmadığı için YAZILAMAZ, yalnız kuru/Excel).  Kart önerisi
     `suggest_item` (tek ve ≥60 puan → "kesin"; hedef kart yalnız aday
     listesinde → "eşleşme belirsiz").  Hedef kartla ilgisi olmayan satır
     listelenmez (sayısı Özet'te; /fiyatlar kuyruğunda kalır).
  2. **Songül Hanım'ın notları** — `scripts/import_lab_price_notes_20261007.py`
     modül olarak (`plan()`; kural ikinci kez YAZILMAZ): hedef karta giden
     kayıtlar `yazilacak` (yazılabilir, E/H alır) · `ezildi` (aynı kart +
     firmada daha yeni/öncelikli kayıt; bilgi) · `belirsiz` / `eslesmedi`
     (aday / benzer kart hedefte: "eşleşme belirsiz"; bilgi — veri
     dosyasında `elle_eslesme.kart` ile çözülür) · diğer durumlar bilgi.
  3. **Fiyat formu** — kalan hedef kartlar Songül Hanım'a: kart, birim, reçete
     sayısı, son alım tedarikçisi, (doldurulacak) fiyat / para birimi / fiyat
     birimi / ambalaj / tedarikçi / tarih / not.  Form satırları betikçe
     OKUNMAZ — veri dosyasına kayıt olarak ya da /fiyatlar elle fiyat formuna
     aktarılır (dolu form hücresi commit'te uyarı olarak listelenir).

**Çalışma kitabı:** Nasıl doldurulur · Paraşüt adayları · Songül adayları ·
Fiyat formu · Özet.  Onay hücreleri YALNIZ E / H (içe aktarma betiğinin
`_yes_no` kuralı: "X", "✓", "1" hata; Excel listesi `showErrorMessage`).
Paraşüt satırında "Kart no", "Ambalaj miktarı", "Fiyat birimi (kg/l/adet)"
ve isteğe bağlı "IMS tedarikçi no" hücreleri yazımda okunur (satırlar "Adet"
geldiği için AMBALAJ ŞART); Songül satırında "Kart no" + "IMS tedarikçi no"
kısmi aktarım manifestine girer (içe aktarma betiği bugünkü planla yeniden
doğrular — eşleme değiştiyse hiçbir şey yazılmaz).

**Yazım (`--onayli ONAYLI.xlsx --commit`):** TEK işlem — Paraşüt E satırları
`parasut_purchases.approve_line` (commit etmez), Songül E satırları içe aktarma
betiğinin `apply(record_selection=…)`'ı (seçili (kart, firma) çiftlerinin lab
fiyatını yeniler, diğer lab / elle / Stok Son Durum satırlarına dokunmaz);
herhangi bir doğrulama hatası (E/H dışı değer, ambalaj yok, kart/firma
bulunamadı, satır artık bekleyen değil, elle fiyat var, kayıt kazanan değil…)
→ HİÇBİR ŞEY yazılmaz.  İkinci çalıştırma aynı sonucu verir (onaylanmış
Paraşüt satırı aynı kartla "zaten yazılmış" sayılır; lab satırı yeniden
yazılır).  Yazımdan önce ilgili tüm kartlar, sonra tüm satırlar id sırasıyla
kilitlenir (`lock_for_approval` — /fiyatlar'daki eşzamanlı tekil onayla
kilitlenme döngüsü kurulmaz; kalan nadir kilit çakışmasında `busy` hatası,
hiçbir şey yazılmaz, yeniden çalıştırılır).  Audit
`supplier_prices.fiyat_tamamlama`.  `--commit` onaylı Excel'siz ÇALIŞMAZ.
Paraşüt'e hiçbir şey yazılmaz; gerçek API ÇAĞRILMAZ.

Kullanım (prod — ÖNCE pg_dump; kuru → Excel → Songül Hanım → --commit):
    cd /var/www/minerva && set -a && source .env && set +a
    venv/bin/python scripts/fiyat_tamamlama_20261016.py                               # kuru özet
    venv/bin/python scripts/fiyat_tamamlama_20261016.py --xlsx /tmp/fiyat_tamamlama.xlsx
    venv/bin/python scripts/fiyat_tamamlama_20261016.py --parasut-json /tmp/parasut.json --xlsx /tmp/ft.xlsx
    venv/bin/python scripts/fiyat_tamamlama_20261016.py --onayli /tmp/onayli.xlsx            # kuru denetim
    venv/bin/python scripts/fiyat_tamamlama_20261016.py --onayli /tmp/onayli.xlsx --commit

Çıkış kodu: 0 = tamam; 2 = veri / Excel / doğrulama hatası (hiçbir şey yazılmadı).
"""
import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

_HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE.parent))
sys.path.insert(0, str(_HERE))

from sqlalchemy import func                                                # noqa: E402
from sqlalchemy.exc import OperationalError                                # noqa: E402
from sqlalchemy.orm import Session                                         # noqa: E402

import import_lab_price_notes_20261007 as lab                              # noqa: E402
from core import parasut_purchases as PP                                   # noqa: E402
from core.audit import log_admin_event                                     # noqa: E402
from core.price_overview import build_overview                             # noqa: E402
from core.supplier_prices import (CURRENCIES, PRICE_UNITS, _num,           # noqa: E402
                                  normalize, price_unit_ok)
from database import (Inventory, Item, ParasutPurchaseLine, ParasutProductPack,   # noqa: E402
                      Recipe, RecipeIngredient, SessionLocal, Supplier, SupplierPrice)

DOMAIN = "cosmetics"
ACTOR = "sistem (fiyat tamamlama)"
TITLE = "Fiyat tamamlama 16.10.2026"
AUDIT_ACTION = "supplier_prices.fiyat_tamamlama"
SHEETS = ("Nasıl doldurulur", "Paraşüt adayları", "Songül adayları", "Fiyat formu", "Özet")
KEY_COL = "Kayıt"
APPROVE_COL = lab.APPROVE_COL                 # "Onay (E/H)"
NOT_APPLICABLE = lab.NOT_APPLICABLE           # "—"
CARD_COL = "Kart no"
PACK_COL = "Ambalaj miktarı"
UNIT_COL = "Fiyat birimi (kg/l/adet)"
SUP_COL = "IMS tedarikçi no"
FORM_PRICE_COL = "Fiyat"
EXACT = "kesin"
AMBIG = "eşleşme belirsiz"
AMBIG_MIN_SCORE = 40.0                        # hedef kart aday listesinde en az bu puanla → "eşleşme belirsiz"
PARASUT_LIMIT = 5000
KEY_DB = "P-"                                 # parasut_purchase_lines.id
KEY_JSON = "PJ-"                              # JSON:API dökümü: parasut_detail_id
KEY_JSON_NOID = "PJx-"                        # düz döküm (kimliksiz) — yazılamaz
SONGUL_STATUS_TEXT = {
    "yazilacak": "yazılabilir",
    "ezildi": "ezildi — aynı kart + firmada daha yeni / öncelikli kayıt var",
    "belirsiz": "eşleşme belirsiz — birden çok aday kart (veri dosyasında elle_eslesme.kart ile çözülür)",
    "eslesmedi": "kart bulunamadı — benzer kart hedefte (veri dosyasında takma ad / elle_eslesme ile çözülür)",
    "yeni_tedarikci": "firma kartı yok — içe aktarma betiği 'Yeni tedarikçiler' E + --create-suppliers ile açar",
}


# ─── Yardımcılar ────────────────────────────────────────────────────────────

def _no_rates():
    """Kur çekilmez (ağ yok): fiyatsız kartlarda gerekmiyor; gerekirse 'kur yok' notu düşer."""
    return None


def _int(v) -> Optional[int]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(str(v).strip().replace(",", "."))
    except ValueError:
        return None
    return int(f) if f.is_integer() and f > 0 else None


def _dmy(d) -> str:
    if isinstance(d, datetime):
        d = d.date()
    return d.strftime("%d.%m.%Y") if isinstance(d, date) else ""


def _num_tr(v) -> str:
    if v is None:
        return ""
    s = f"{float(v):,.4f}".rstrip("0").rstrip(".")
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _card_text(c: Optional[dict]) -> str:
    return f"{c['id']} {c['name']} ({c.get('unit') or 'birimsiz'})" if c else ""


# ─── 0) Hedef kartlar ──────────────────────────────────────────────────────

def load_targets(db: Session, domain: str = DOMAIN) -> Dict[int, dict]:
    """Aktif reçetelerde geçen fiyatsız hammadde kartları (fiyat sayfasının
    'fiyatı olmayanlar' kuralı: hammadde + 'diğer' kovası; ambalaj / etiket
    HARİÇ).  {item_id: {…, recipes, recipe_names, last_supplier, last_date,
    card_supplier, group_name, group_priced_via}}."""
    cards: Dict[int, dict] = {}
    for cat in ("hammadde", "other"):
        ov = build_overview(db, domain, currency="USD", category=cat, has_price="unpriced", rates_fn=_no_rates)
        for blk in ov["blocks"]:
            for c in blk["cards"]:
                cards[c["item_id"]] = c
    if not cards:
        return {}
    ids = list(cards)
    rows = (db.query(RecipeIngredient.item_id, Recipe.id, Recipe.name)
            .join(Recipe, Recipe.id == RecipeIngredient.recipe_id)
            .filter(Recipe.is_active == True, Recipe.domain == domain,            # noqa: E712
                    RecipeIngredient.item_id.in_(ids)).all())
    recipes: Dict[int, Dict[int, str]] = {}
    for iid, rid, rname in rows:
        recipes.setdefault(iid, {})[rid] = rname
    sub = (db.query(Inventory.item_id, func.max(Inventory.id).label("max_id"))
           .filter(Inventory.item_id.in_(ids), Inventory.supplier_id.isnot(None),
                   Inventory.is_sample == False)                                   # noqa: E712
           .group_by(Inventory.item_id).subquery())
    last = {iid: (name, at) for iid, name, at in
            (db.query(Inventory.item_id, Supplier.name, Inventory.created_at)
             .join(sub, Inventory.id == sub.c.max_id)
             .join(Supplier, Supplier.id == Inventory.supplier_id).all())}
    gids = {c["group_id"] for c in cards.values() if c.get("group_id")}
    by_gid: Dict[int, List[str]] = {}
    if gids:
        mates = (db.query(Item.id, Item.name, Item.material_group_id)
                 .filter(Item.domain == domain, Item.is_active == True,            # noqa: E712
                         Item.material_group_id.in_(gids)).all())
        mate_ids = [m.id for m in mates if m.id not in cards]
        priced = set()
        if mate_ids:
            priced = {r[0] for r in db.query(SupplierPrice.item_id)
                      .filter(SupplierPrice.domain == domain, SupplierPrice.item_id.in_(mate_ids),
                              SupplierPrice.unit_price > 0).distinct().all()}
        for m in mates:
            if m.id in priced:
                by_gid.setdefault(m.material_group_id, []).append(f"{m.id} {m.name}")
    out: Dict[int, dict] = {}
    for iid, c in cards.items():
        if iid not in recipes:
            continue
        rec = recipes[iid]
        l = last.get(iid)
        out[iid] = {
            "item_id": iid, "name": c["name"], "name_tr": c["name_tr"], "unit": c["unit"],
            "category": c["category"], "bucket": c["bucket"], "bucket_label": c["bucket_label"],
            "purchase_unit": c["purchase_unit"], "recipes": len(rec),
            "recipe_names": [rec[k] for k in sorted(rec)],
            "last_supplier": l[0] if l else None, "last_date": l[1].date() if l and l[1] else None,
            "card_supplier": c.get("card_supplier"), "group_name": c.get("group_name"),
            "group_priced_via": ", ".join(by_gid.get(c["group_id"], [])) or None if c.get("group_id") else None,
            "merged_with": [m["name"] for m in c.get("merged_with") or []],
        }
    return out


# ─── 1) Paraşüt satırları ──────────────────────────────────────────────────

def _derive_discount(qty, unit_price, vat_rate, net_total) -> Optional[float]:
    """Düz dökümde indirim alanı yok: net = adet × fiyat × (1 − indirim) × (1 + KDV) → indirim."""
    base = (qty or 0.0) * (unit_price or 0.0) * (1.0 + (vat_rate or 0.0) / 100.0)
    if base <= 0 or net_total is None:
        return None
    rate = 1.0 - float(net_total) / base
    return round(rate, 4) if rate > 0.005 else None


def bill_from_dump(d: dict) -> Optional[dict]:
    """`parasut_read.py` düz dökümü → `parse_bill` ile AYNI beyaz listeli sözlük
    (kimlikler None; `supplier_contact` ve diğer alanlar ATILIR)."""
    item_type = (str(d.get("item_type") or "purchase_bill")).strip().lower()
    if item_type != "purchase_bill":
        return None
    lines = []
    for l in d.get("lines") or []:
        if not isinstance(l, dict):
            continue
        qty, up, vat = PP._num(l.get("qty")), PP._num(l.get("unit_price")), PP._num(l.get("vat_rate"))
        net = PP._num(l.get("net_total"))
        pname, ldesc = PP._txt(l.get("product"), 150), PP._txt(l.get("desc"), 250)
        lines.append({
            "parasut_detail_id": None, "description": ldesc, "product_parasut_id": None,
            "product_name": pname, "product_unit": PP._txt(l.get("unit"), 20),
            "quantity": qty, "unit_price": up, "discount_rate": _derive_discount(qty, up, vat, net),
            "vat_rate": vat, "net_total": net, "expense_hint": PP.expense_hint(ldesc, pname),
        })
    sup_name, desc = PP._txt(d.get("supplier"), 150), PP._txt(d.get("description"), 250)
    return {
        "parasut_id": None, "item_type": item_type, "invoice_no": PP._txt(d.get("no"), 60),
        "issue_date": PP._date(d.get("date")), "due_date": None,
        "currency": PP.currency_code(d.get("currency")), "exchange_rate": PP._num(d.get("exchange_rate")),
        "net_total": PP._num(d.get("net_total")), "gross_total": None, "description": desc,
        "supplier_parasut_id": None, "supplier_name": sup_name,
        "expense_hint": PP.expense_hint(sup_name, desc) or (bool(lines) and all(ln["expense_hint"] for ln in lines)),
        "lines": lines,
    }


def load_parasut_json(path) -> Tuple[List[dict], str]:
    """Döküm dosyası → (beyaz listeli faturalar, 'jsonapi' | 'dump')."""
    with Path(path).open(encoding="utf-8") as f:
        payload = json.load(f)
    if not isinstance(payload, dict):
        raise ValueError("--parasut-json: JSON nesnesi bekleniyor ({\"data\": …} ya da {\"bills\": …})")
    if "data" in payload:
        inc = PP._incmap(payload)
        bills = []
        for raw in payload.get("data") or []:
            b = PP.parse_bill(raw, inc) if isinstance(raw, dict) else None
            if b:
                bills.append(b)
        return bills, "jsonapi"
    if "bills" in payload:
        bills = []
        for d in payload.get("bills") or []:
            b = bill_from_dump(d) if isinstance(d, dict) else None
            if b:
                bills.append(b)
        return bills, "dump"
    raise ValueError("--parasut-json: ne JSON:API ('data') ne düz döküm ('bills') biçimi")


def _line_rec(key: str, b: dict, ld: dict, *, line_id: Optional[int] = None) -> dict:
    issue = b.get("issue_date")
    return {
        "key": key, "line_id": line_id, "detail_id": ld.get("parasut_detail_id"),
        "product_parasut_id": ld.get("product_parasut_id"),
        "invoice_no": b.get("invoice_no"), "issue_date": issue, "supplier_name": b.get("supplier_name"),
        "currency": b.get("currency"), "bill_description": b.get("description"),
        "description": ld.get("description"), "product_name": ld.get("product_name"),
        "product_unit": ld.get("product_unit"), "quantity": ld.get("quantity"),
        "unit_price": ld.get("unit_price"), "discount_rate": ld.get("discount_rate"),
        "vat_rate": ld.get("vat_rate"), "net_total": ld.get("net_total"),
        "expense_hint": bool(ld.get("expense_hint")) or bool(b.get("expense_hint")),
        "stale": bool(issue and (date.today() - issue).days > PP.STALE_DAYS),
    }


def parasut_lines(db: Session, bills: Optional[List[dict]] = None, kind: Optional[str] = None) -> dict:
    """Değerlendirilecek satırlar: veritabanındaki BEKLEYEN satırlar (+ varsa
    döküm — veritabanında kaydı olan döküm satırı atlanır, DB esastır)."""
    recs: List[dict] = []
    for ln in PP.pending_lines(db, PARASUT_LIMIT):
        b = ln.bill
        bd = {"invoice_no": b.invoice_no, "issue_date": b.issue_date, "supplier_name": b.supplier_name,
              "currency": b.currency, "description": b.description, "expense_hint": b.expense_hint}
        ld = {k: getattr(ln, k) for k in PP.LINE_KEYS}
        recs.append(_line_rec(f"{KEY_DB}{ln.id}", bd, ld, line_id=ln.id))
    pending_total = db.query(ParasutPurchaseLine).filter(ParasutPurchaseLine.status == "pending").count()
    n_json, n_json_in_db = 0, 0
    if bills:
        known = set()
        det_ids = [ld["parasut_detail_id"] for b in bills for ld in b["lines"] if ld.get("parasut_detail_id")]
        if det_ids:
            known = {r[0] for r in db.query(ParasutPurchaseLine.parasut_detail_id)
                     .filter(ParasutPurchaseLine.parasut_detail_id.in_(det_ids)).all()}
        n = 0
        for b in bills:
            for ld in b["lines"]:
                n += 1
                did = ld.get("parasut_detail_id")
                if did and did in known:
                    n_json_in_db += 1
                    continue
                key = f"{KEY_JSON}{did}" if did else f"{KEY_JSON_NOID}{n}"
                recs.append(_line_rec(key, b, ld))
                n_json += 1
    return {"lines": recs, "pending_total": pending_total, "json_lines": n_json, "json_in_db": n_json_in_db,
            "source": ("db+" + kind) if bills else "db"}


def parasut_candidates(db: Session, lines: List[dict], targets: Dict[int, dict], domain: str = DOMAIN) -> dict:
    """Her satır için `suggest_item`; hedef kartla ilgili satırlar → aday satırı
    ('kesin' = tek ve ≥60 puanlı öneri hedefte; 'eşleşme belirsiz' = hedef
    kart yalnız aday listesinde).  Hiçbir şey YAZMAZ."""
    index = PP.build_item_index(db, domain) if lines else []
    sup_cache: Dict[str, Optional[dict]] = {}
    pack_cache: Dict[int, Optional[ParasutProductPack]] = {}
    rows: List[dict] = []
    n_off = n_expense = 0
    for ln in lines:
        sk = normalize(ln["supplier_name"])
        if sk not in sup_cache:
            sup_cache[sk] = PP.match_supplier(db, domain, ln["supplier_name"])
        sup = sup_cache[sk]
        res = PP.suggest_item(db, {"product_name": ln["product_name"], "description": ln["description"]},
                              domain, index=index, supplier_id=(sup or {}).get("id"))
        sug = res["suggested"]
        tc = [c for c in res["candidates"] if c["id"] in targets and c["score"] >= AMBIG_MIN_SCORE]
        if not tc:
            n_off += 1
            if ln["expense_hint"]:
                n_expense += 1
            continue
        if sug and sug["id"] in targets:
            match, card = EXACT, sug
        else:
            match, card = AMBIG, tc[0]
        others = [c for c in tc if c["id"] != card["id"]]
        pack = None
        pid = ln.get("product_parasut_id")
        if pid:
            if pid not in pack_cache:
                pack_cache[pid] = db.query(ParasutProductPack).filter(
                    ParasutProductPack.product_parasut_id == pid).first()
            pk = pack_cache[pid]
            if pk is not None:
                pack = {"package_qty": pk.package_qty, "price_unit": pk.price_unit, "item_id": pk.item_id}
        pack_qty = pack["package_qty"] if pack else None
        # Hatırlanan ambalaj yoksa kartın ALIM birimi (g → kg, ml → l, adet → adet) —
        # "Fiyat formu" ve kuyruk arayüzüyle aynı; default_price_unit ml kartına 'kg' verirdi
        unit = pack["price_unit"] if pack else targets[card["id"]]["purchase_unit"]
        net = None
        if pack_qty and ln["unit_price"]:
            net = round(float(ln["unit_price"]) * (1.0 - float(ln["discount_rate"] or 0.0)) / float(pack_qty), 6)
        notes = []
        if match == AMBIG and sug:
            notes.append(f"betiğin önerisi hedef dışı kart: {_card_text(sug)}")
        if match == AMBIG and not sug:
            notes.append("tek ve ≥60 puanlı öneri yok")
        if ln["stale"]:
            notes.append("eski fatura (12 aydan eski) — yazılır ama satın alma planı seçmez; güncel fiyat istenir")
        if ln["expense_hint"]:
            notes.append("gider ipucu (kargo / hizmet?) — hammadde değilse H")
        if pack:
            notes.append("ambalaj Paraşüt ürününden hatırlandı")
        if sup is None:
            notes.append("firma IMS kartına eşleşmedi — serbest metin yazılır (IMS tedarikçi no girilebilir)")
        if not ln["line_id"]:
            notes.append("döküm satırı" + ("" if ln["detail_id"] else " (kimliksiz — YAZILAMAZ, yalnız bilgi)"))
        rows.append({
            **ln, "sup": sup, "match": match, "score": card["score"], "card": card,
            "others": "; ".join(f"{c['id']} {c['name']} ({c['score']})" for c in others),
            "pack_qty": pack_qty, "unit": unit, "net_preview": net, "note": " · ".join(notes),
            "input": bool(ln["detail_id"] or ln["line_id"]),
        })
    rows.sort(key=lambda r: (r["match"] != EXACT, r["stale"], normalize(r["card"]["name"]), r["key"]))
    return {"rows": rows, "off_target": n_off, "expense": n_expense}


# ─── 2) Songül kayıtları ───────────────────────────────────────────────────

def songul_candidates(db: Session, data: dict, targets: Dict[int, dict], domain: str = DOMAIN) -> dict:
    """İçe aktarma betiğinin tam planı; hedef karta giden (ya da aday / benzer
    kartı hedefte olan) kayıtlar.  Yalnız `yazilacak` satır E/H alır."""
    p = lab.plan(db, data, domain=domain)
    rows: List[dict] = []
    for r in p["records"]:
        m, it = r["mat"], r["item"]
        tid = it["id"] if it and it["id"] in targets else None
        cand_ids = [c for c in (m.get("candidates") or []) + (m.get("suggestions") or []) if c in targets]
        if tid is None and not cand_ids:
            continue
        status = r["status"]
        exact = bool(tid) and m.get("guven") != "belirsiz" and m.get("rule") != "c"
        if tid:
            card = {"id": tid, "name": targets[tid]["name"], "unit": targets[tid]["unit"]}
        else:
            card = {"id": cand_ids[0], "name": targets[cand_ids[0]]["name"], "unit": targets[cand_ids[0]]["unit"]}
        cands = [f"{c} {targets[c]['name']}" for c in cand_ids if c != card["id"]]
        text = SONGUL_STATUS_TEXT.get(status) or lab.STATUS_TEXT.get(status, status)
        if status == "ezildi":
            text += " (" + "; ".join(f for f in r["flags"] if f.startswith("kazanan")) + ")"
        notes = list(r["flags"])
        if m.get("rule"):
            notes.append(f"kart kuralı {lab.RULE_TEXT.get(m['rule'], m['rule'])}")
        if r["confidence"] == "belirsiz":
            notes.append("OKUMA BELİRSİZ — değer teyit edilmeli")
        if r["overridden"]:
            notes.append("ezilen: " + "; ".join(r["overridden"]))
        rows.append({
            "key": r["id"], "page": r["page"], "type": lab.TYPE_TEXT.get(r["type"], r["type"]),
            "date": r["date"], "material": r["material"], "supplier": r["supplier"], "sup": r["sup"],
            "price": r["price"], "currency": r["currency"], "price_unit": r["price_unit"],
            "package": r["package"], "vat": lab._vat_text(r), "vat_rate": r["vat_rate"],
            "confidence": r["confidence"], "card": card, "others": "; ".join(cands),
            "match": EXACT if exact else AMBIG, "status": status, "status_text": text,
            "writable": status == "yazilacak" and r["sup"]["kind"] == "id",
            "note": " · ".join(n for n in notes if n), "doc_note": r["note"],
        })
    order = {"yazilacak": 0, "ezildi": 1, "belirsiz": 2, "eslesmedi": 3}
    rows.sort(key=lambda r: (order.get(r["status"], 9), normalize(r["card"]["name"]), r["key"]))
    return {"rows": rows, "plan": p}


# ─── 3) Plan + form + özet ─────────────────────────────────────────────────

def build_plan(db: Session, *, data: Optional[dict] = None, domain: str = DOMAIN,
               parasut_bills: Optional[List[dict]] = None, parasut_kind: Optional[str] = None) -> dict:
    """Kademeyi hesaplar — hiçbir şey YAZMAZ.  Döner: {targets, parasut, songul, form, summary}."""
    data = data or lab.load_data()
    targets = load_targets(db, domain)
    pl = parasut_lines(db, parasut_bills, parasut_kind)
    pc = parasut_candidates(db, pl["lines"], targets, domain)
    sc = songul_candidates(db, data, targets, domain)
    covered = ({r["card"]["id"] for r in pc["rows"] if r["match"] == EXACT and not r["stale"]}
               | {r["card"]["id"] for r in sc["rows"] if r["writable"]})
    hints: Dict[int, List[str]] = {}
    for r in pc["rows"]:
        tag = "Paraşüt" + (" (eski fatura)" if r["stale"] else "") + ("" if r["match"] == EXACT else f" ({AMBIG})")
        hints.setdefault(r["card"]["id"], []).append(f"{tag}: {r['invoice_no'] or r['key']} {r['product_name'] or ''}")
    for r in sc["rows"]:
        if not r["writable"]:
            tag = "Songül notu " + (r["status_text"].split(" — ")[0] if r["status"] != "yazilacak" else "")
            hints.setdefault(r["card"]["id"], []).append(f"{tag}: {r['key']} {r['material']} / {r['supplier']}")
    form = []
    for iid, t in sorted(targets.items(), key=lambda kv: normalize(kv[1]["name"])):
        if iid in covered:
            continue
        notes = []
        if t["group_priced_via"]:
            notes.append(f"grup üzerinden fiyatlı ({t['group_priced_via']})")
        if t["merged_with"]:
            notes.append("planda birleşik: " + ", ".join(t["merged_with"]))
        notes += hints.get(iid, [])
        form.append({**t, "note": " · ".join(notes)})
    summary = {
        "targets": len(targets),
        "group_priced": sum(1 for t in targets.values() if t["group_priced_via"]),
        "parasut_source": pl["source"], "parasut_pending_total": pl["pending_total"],
        "parasut_json_lines": pl["json_lines"], "parasut_json_in_db": pl["json_in_db"],
        "parasut_rows": len(pc["rows"]),
        "parasut_exact_cards": len({r["card"]["id"] for r in pc["rows"] if r["match"] == EXACT}),
        "parasut_ambig_cards": len({r["card"]["id"] for r in pc["rows"] if r["match"] == AMBIG}),
        "parasut_stale_rows": sum(1 for r in pc["rows"] if r["stale"]),
        "parasut_off_target": pc["off_target"], "parasut_expense": pc["expense"],
        "songul_records": len(sc["plan"]["records"]), "songul_rows": len(sc["rows"]),
        "songul_writable_cards": len({r["card"]["id"] for r in sc["rows"] if r["writable"]}),
        "songul_writable_rows": sum(1 for r in sc["rows"] if r["writable"]),
        "songul_overridden": sum(1 for r in sc["rows"] if r["status"] == "ezildi"),
        "songul_ambig": sum(1 for r in sc["rows"] if r["status"] in ("belirsiz", "eslesmedi")),
        "songul_other": sum(1 for r in sc["rows"] if r["status"] not in ("yazilacak", "ezildi", "belirsiz",
                                                                            "eslesmedi")),
        "covered": len(covered & set(targets)), "form": len(form),
        "generated": datetime.now().strftime("%d.%m.%Y %H:%M"),
    }
    return {"domain": domain, "targets": targets, "parasut": pc, "songul": sc, "form": form, "summary": summary}


# ─── Çalışma kitabı ────────────────────────────────────────────────────────

def _yes_no_list():
    from openpyxl.worksheet.datavalidation import DataValidation
    return DataValidation(type="list", formula1='"E,H"', allow_blank=True, showErrorMessage=True,
                          errorStyle="stop", errorTitle="E ya da H",
                          error="Lütfen yalnız E (evet) ya da H (hayır) yazın.")


def _list_validation(values: str, title: str):
    from openpyxl.worksheet.datavalidation import DataValidation
    return DataValidation(type="list", formula1=f'"{values}"', allow_blank=True, showErrorMessage=True,
                          errorStyle="stop", errorTitle=title, error=f"Lütfen listeden seçin: {values}")


def build_xlsx(plan: dict, path) -> Path:
    """Çalışma kitabını yazar (onay hücreleri E/H listesi + hata iletisi; birim
    ve para birimi listeleri; giriş hücreleri sarı; '—' = bu satırda giriş yok)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="232E6E")
    input_fill = PatternFill("solid", fgColor="FFF4C2")
    warn_font = Font(color="9A3412")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, cols, rows, *, inputs=(), yesno=None, lists=(), warn_cols=()):
        ws = wb.create_sheet(title)
        for ci, (name, width) in enumerate(cols, 1):
            c = ws.cell(1, ci, name)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(ci)].width = width
        names = [n for n, _ in cols]
        for ri, row in enumerate(rows, 2):
            for ci, v in enumerate(row, 1):
                c = ws.cell(ri, ci, v)
                c.alignment = wrap if isinstance(v, str) else top
                if names[ci - 1] in inputs and v != NOT_APPLICABLE:
                    c.fill = input_fill
                if names[ci - 1] in warn_cols and v:
                    c.font = warn_font
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32
        last = max(len(rows) + 1, 2)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
        if rows:
            if yesno:
                dv = _yes_no_list()
                col = get_column_letter(names.index(yesno) + 1)
                dv.add(f"{col}2:{col}{len(rows) + 1}")
                ws.add_data_validation(dv)
            for cname, values, title in lists:
                dv = _list_validation(values, title)
                col = get_column_letter(names.index(cname) + 1)
                dv.add(f"{col}2:{col}{len(rows) + 1}")
                ws.add_data_validation(dv)
        return ws

    s = plan["summary"]
    # ── Nasıl doldurulur ──
    ws = wb.create_sheet(SHEETS[0])
    ws.column_dimensions["A"].width = 120
    lines = [
        (f"{TITLE} — fiyatsız hammaddeler için kontrol listesi ({s['generated']})", True),
        (f"Hedef: aktif reçetelerde geçen fiyatsız hammadde kartı {s['targets']} (yalnız grup üzerinden "
         f"fiyatlı: {s['group_priced']}) · Paraşüt adayı {s['parasut_rows']} satır · Songül adayı "
         f"{s['songul_rows']} kayıt (yazılabilir {s['songul_writable_rows']}) · forma kalan {s['form']}", False),
        ("", False),
        ("Onay hücreleri YALNIZ E (evet) ya da H (hayır) alır. Boş = karar yok. 'X', '✓', '1' gibi işaretler "
         "hata sayılır ve HİÇBİR ŞEY yazılmaz.", False),
        ("'—' yazan hücre bu satırda giriş olmadığını söyler (bilgi satırı).", False),
        ("", False),
        ("1) Paraşüt adayları: alış faturası satırı → IMS kartı. Satırlar faturada 'Adet' geldiği için "
         "AMBALAJ MİKTARI şart (ör. 25 kg'lık bidon → 25; 10 ml şişe, kg fiyatı → 0,01). Fiyat birimi "
         "kartın birim ailesine uymalı (g/kg → kg, ml/l → l, adet → adet). Net fiyat = birim fiyat × "
         "(1 − indirim) ÷ ambalaj, KDV hariç. 'Kart no' önerilen karttır, farklıysa değiştirin. "
         "'eşleşme belirsiz' satırını yalnız kartı doğruladıysanız E yapın. 'IMS tedarikçi no' boşsa "
         "faturanın firması (eşleşen kart ya da serbest metin) yazılır.", False),
        ("2) Songül adayları: 07.10 / 09.10 notlarındaki kayıtlar. Yalnız 'yazılabilir' satır E/H alır; "
         "ezilen / belirsiz satırlar bilgidir (belirsiz kart, veri dosyasında elle eşleme ile çözülür). "
         "'Kart no' ve 'IMS tedarikçi no' DEĞİŞTİRİLMEZ — farklıysa H yazıp notunu 'Not' hücresine yazın.", False),
        ("3) Fiyat formu: hiçbir kaynakta fiyatı olmayan kartlar. Fiyat, para birimi, fiyat birimi, ambalaj, "
         "tedarikçi, tarih ve notu doldurun. Bu sayfa betikçe OTOMATİK yazılmaz — dönünce veri dosyasına ya "
         "da /fiyatlar elle fiyat formuna aktarılır.", False),
        ("", False),
        ("Yazım: scripts/fiyat_tamamlama_20261016.py --onayli BU_DOSYA.xlsx --commit — tek işlem; herhangi bir "
         "hata varsa hiçbir fiyat yazılmaz.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(i, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=13)

    # ── Paraşüt adayları ──
    cols = [(KEY_COL, 11), ("Fatura no", 18), ("Fatura tarihi", 11), ("Firma (fatura)", 30),
            ("IMS tedarikçi", 24), ("Ürün (Paraşüt)", 34), ("Satır açıklaması", 28), ("Miktar", 8),
            ("Birim (Paraşüt)", 8), ("Birim fiyat", 11), ("İndirim %", 8), ("KDV %", 7), ("Para birimi", 7),
            ("Önerilen kart", 36), ("Eşleşme", 14), ("Puan", 6), ("Diğer hedef adaylar", 30),
            ("Net önizleme", 11), ("Eski fatura", 8), ("Gider ipucu", 8), ("Not", 50),
            (APPROVE_COL, 10), (CARD_COL, 9), (PACK_COL, 11), (UNIT_COL, 12), (SUP_COL, 11)]
    rows = []
    for r in plan["parasut"]["rows"]:
        na = not r["input"]
        rows.append([
            r["key"], r["invoice_no"] or "", _dmy(r["issue_date"]), r["supplier_name"] or "",
            (f"{r['sup']['id']} {r['sup']['name']}" if r["sup"] else ""), r["product_name"] or "",
            r["description"] or "", r["quantity"], r["product_unit"] or "", r["unit_price"],
            (round(float(r["discount_rate"]) * 100.0, 2) if r["discount_rate"] else None), r["vat_rate"],
            r["currency"] or "", _card_text(r["card"]), r["match"], r["score"], r["others"],
            r["net_preview"], "EVET" if r["stale"] else "", "evet" if r["expense_hint"] else "", r["note"],
            NOT_APPLICABLE if na else None, NOT_APPLICABLE if na else r["card"]["id"],
            NOT_APPLICABLE if na else r["pack_qty"], NOT_APPLICABLE if na else r["unit"],
            NOT_APPLICABLE if na else (r["sup"]["id"] if r["sup"] else None),
        ])
    sheet(SHEETS[1], cols, rows, inputs=(APPROVE_COL, CARD_COL, PACK_COL, UNIT_COL, SUP_COL), yesno=APPROVE_COL,
          lists=[(UNIT_COL, ",".join(PRICE_UNITS), "Fiyat birimi")], warn_cols=("Eşleşme", "Eski fatura", "Not"))

    # ── Songül adayları ──
    cols = [(KEY_COL, 11), ("Sayfa", 14), ("Tür", 12), ("Tarih", 11), ("Malzeme (belge)", 30),
            ("Tedarikçi (belge)", 22), ("IMS tedarikçi", 24), ("Fiyat", 10), ("Para birimi", 7),
            ("Fiyat birimi", 7), ("Ambalaj", 8), ("KDV", 8), ("Okuma", 9), ("Hedef kart", 36),
            ("Eşleşme", 14), ("Diğer hedef adaylar", 28), ("Durum", 40), ("Not", 50), ("Belge notu", 40),
            (APPROVE_COL, 10), (CARD_COL, 9), (SUP_COL, 11)]
    rows = []
    for r in plan["songul"]["rows"]:
        na = not r["writable"]
        vat = (r["vat"] + (f" %{_num_tr(r['vat_rate'])}" if r["vat_rate"] is not None else "")).strip()
        rows.append([
            r["key"], str(r["page"]) if r["page"] is not None else "", r["type"], _dmy(r["date"]),
            r["material"], r["supplier"], lab._sup_text(r["sup"]) if r["sup"].get("name") or r["sup"].get("id")
            else "", r["price"], r["currency"], r["price_unit"], r["package"], vat,
            "BELİRSİZ" if r["confidence"] == "belirsiz" else "kesin", _card_text(r["card"]), r["match"],
            r["others"], r["status_text"], r["note"], r["doc_note"],
            NOT_APPLICABLE if na else None, NOT_APPLICABLE if na else r["card"]["id"],
            NOT_APPLICABLE if na else r["sup"]["id"],
        ])
    sheet(SHEETS[2], cols, rows, inputs=(APPROVE_COL,), yesno=APPROVE_COL,
          warn_cols=("Eşleşme", "Okuma", "Durum", "Not"))

    # ── Fiyat formu ──
    cols = [("Kart no", 8), ("Kart", 36), ("Birim", 7), ("Kategori", 12), ("Reçete sayısı", 8),
            ("Reçeteler", 40), ("Son alım tedarikçisi", 24), ("Son alım tarihi", 11), ("Kart tedarikçisi", 22),
            ("Grup", 20), ("Not (adaylar / grup)", 50),
            (FORM_PRICE_COL, 10), ("Para birimi", 8), ("Fiyat birimi (kg/l/adet)", 10), ("Ambalaj", 9),
            ("Tedarikçi", 24), ("Tarih", 11), ("Not", 30)]
    rows = []
    for t in plan["form"]:
        names = t["recipe_names"]
        rows.append([
            t["item_id"], t["name"] + (f" / {t['name_tr']}" if t["name_tr"] else ""), t["unit"] or "birimsiz",
            t["bucket_label"] or t["category"], t["recipes"],
            ", ".join(names[:3]) + (f" (+{len(names) - 3})" if len(names) > 3 else ""),
            t["last_supplier"] or "", _dmy(t["last_date"]), t["card_supplier"] or "", t["group_name"] or "",
            t["note"], None, None, t["purchase_unit"], None, t["last_supplier"] or t["card_supplier"] or None,
            None, None,
        ])
    sheet(SHEETS[3], cols, rows,
          inputs=(FORM_PRICE_COL, "Para birimi", "Fiyat birimi (kg/l/adet)", "Ambalaj", "Tedarikçi", "Tarih", "Not"),
          lists=[("Para birimi", ",".join(CURRENCIES), "Para birimi"),
                 ("Fiyat birimi (kg/l/adet)", ",".join(PRICE_UNITS), "Fiyat birimi")],
          warn_cols=("Not (adaylar / grup)",))

    # ── Özet ──
    ozet = [
        ("Üretim", s["generated"]),
        ("Hedef kart (aktif reçetede, fiyatsız hammadde)", s["targets"]),
        ("  ↳ yalnız grup üzerinden fiyatlı (formda bilgi)", s["group_priced"]),
        ("Paraşüt kaynağı", s["parasut_source"]),
        ("Paraşüt bekleyen satır (kuyruk toplamı)", s["parasut_pending_total"]),
        ("Paraşüt döküm satırı (veritabanında olmayan)", s["parasut_json_lines"]),
        ("Paraşüt döküm satırı veritabanında zaten var (atlandı)", s["parasut_json_in_db"]),
        ("Paraşüt aday satırı", s["parasut_rows"]),
        ("  ↳ kesin eşleşen kart", s["parasut_exact_cards"]),
        ("  ↳ eşleşme belirsiz kart", s["parasut_ambig_cards"]),
        ("  ↳ eski fatura (12 ay+) satırı", s["parasut_stale_rows"]),
        ("Paraşüt hedef dışı satır (kuyrukta kalır)", s["parasut_off_target"]),
        ("  ↳ gider ipuçlu", s["parasut_expense"]),
        ("Songül kaydı (veri dosyası)", s["songul_records"]),
        ("Songül aday kaydı (hedef kartla ilgili)", s["songul_rows"]),
        ("  ↳ yazılabilir kayıt / kart", f"{s['songul_writable_rows']} / {s['songul_writable_cards']}"),
        ("  ↳ ezilen", s["songul_overridden"]),
        ("  ↳ eşleşme belirsiz / bulunamadı", s["songul_ambig"]),
        ("  ↳ diğer (firma yok, birim, çizili…)", s["songul_other"]),
        ("Kademeyle karşılanan kart", s["covered"]),
        ("Fiyat formuna kalan kart", s["form"]),
    ]
    sheet(SHEETS[4], [("Ölçü", 60), ("Değer", 24)], ozet)
    path = Path(path)
    wb.save(path)
    return path


# ─── Onaylı Excel → onaylar ────────────────────────────────────────────────

def read_workbook(path) -> dict:
    """Onaylı Excel → {parasut: {kayıt: {ok, row, item_id, package, unit, supplier_id}},
    songul: {kayıt: {ok, row, item_id, supplier_id}}, form_filled: [(row, kart, fiyat)],
    errors: [...]}.  Yalnız E/H yazılan satırlar girer (boş = karar yok); E/H dışı
    onay değeri hataya girer (`lab._yes_no`)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {"parasut": {}, "songul": {}, "form_filled": [], "errors": []}

    def table(title, *cols):
        if title not in wb.sheetnames:
            out["errors"].append(f"'{title}' sayfası yok")
            return None
        it = wb[title].iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, ())]
        missing = [c for c in cols if c not in head]
        if missing:
            out["errors"].append(f"'{title}': sütun yok: {', '.join(missing)}")
            return None
        return head, it

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    def na(v):
        return isinstance(v, str) and v.strip() == NOT_APPLICABLE

    t = table(SHEETS[1], KEY_COL, APPROVE_COL, CARD_COL, PACK_COL, UNIT_COL, SUP_COL)
    if t:
        head, it = t
        ix = {c: head.index(c) for c in (KEY_COL, APPROVE_COL, CARD_COL, PACK_COL, UNIT_COL, SUP_COL)}
        for n, row in enumerate(it, 2):
            key = cell(row, ix[KEY_COL])
            if key is None or not str(key).strip():
                continue
            raw = cell(row, ix[APPROVE_COL])
            if na(raw):
                continue
            try:
                ok = lab._yes_no(raw)
            except ValueError:
                out["errors"].append(f"'{SHEETS[1]}' satır {n} ({key}): '{raw}' — yalnız E ya da H olmalı")
                continue
            if ok is None:                                   # boş = karar yok
                continue
            out["parasut"][str(key).strip()] = {
                "ok": ok, "row": n, "item_id": cell(row, ix[CARD_COL]), "package": cell(row, ix[PACK_COL]),
                "unit": cell(row, ix[UNIT_COL]), "supplier_id": cell(row, ix[SUP_COL])}
    t = table(SHEETS[2], KEY_COL, APPROVE_COL, CARD_COL, SUP_COL)
    if t:
        head, it = t
        ix = {c: head.index(c) for c in (KEY_COL, APPROVE_COL, CARD_COL, SUP_COL)}
        for n, row in enumerate(it, 2):
            key = cell(row, ix[KEY_COL])
            if key is None or not str(key).strip():
                continue
            raw = cell(row, ix[APPROVE_COL])
            if na(raw):
                continue
            try:
                ok = lab._yes_no(raw)
            except ValueError:
                out["errors"].append(f"'{SHEETS[2]}' satır {n} ({key}): '{raw}' — yalnız E ya da H olmalı")
                continue
            if ok is None:
                continue
            out["songul"][str(key).strip()] = {"ok": ok, "row": n, "item_id": cell(row, ix[CARD_COL]),
                                               "supplier_id": cell(row, ix[SUP_COL])}
    t = table(SHEETS[3], "Kart no", FORM_PRICE_COL)
    if t:
        head, it = t
        ik, ip = head.index("Kart no"), head.index(FORM_PRICE_COL)
        for n, row in enumerate(it, 2):
            v = cell(row, ip)
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            out["form_filled"].append((n, cell(row, ik), v))
    wb.close()
    return out


# ─── Doğrulama + yazım ─────────────────────────────────────────────────────

def _find_line(db: Session, key: str) -> Tuple[Optional[ParasutPurchaseLine], Optional[int]]:
    """Kayıt anahtarı → (satır | None, detay kimliği | None)."""
    if key.startswith(KEY_DB):
        lid = _int(key[len(KEY_DB):])
        return (db.query(ParasutPurchaseLine).filter(ParasutPurchaseLine.id == lid).first() if lid else None), None
    if key.startswith(KEY_JSON):
        did = _int(key[len(KEY_JSON):])
        if not did:
            return None, None
        return db.query(ParasutPurchaseLine).filter(ParasutPurchaseLine.parasut_detail_id == did).first(), did
    return None, None


def resolve(db: Session, approvals: dict, *, data: Optional[dict] = None, domain: str = DOMAIN,
            parasut_bills: Optional[List[dict]] = None) -> dict:
    """E satırlarını BUGÜNKÜ veritabanıyla doğrular; hiçbir şey yazmaz.
    Döner: {parasut: [{key, line, item, package, unit, supplier_id, upsert_bill}], parasut_done: [...],
    songul_selection: [{id, item_id, supplier_id}], songul_plan, errors, warnings}."""
    data = data or lab.load_data()
    errors: List[str] = list(approvals.get("errors") or [])
    warnings: List[str] = []
    todo: List[dict] = []
    done: List[str] = []
    by_detail = {}
    for b in parasut_bills or []:
        for ld in b["lines"]:
            if ld.get("parasut_detail_id"):
                by_detail[ld["parasut_detail_id"]] = b
    for key, a in sorted(approvals.get("parasut", {}).items(), key=lambda kv: kv[1]["row"]):
        if a["ok"] is not True:
            continue
        where = f"'{SHEETS[1]}' satır {a['row']} ({key})"
        ln, did = _find_line(db, key)
        upsert = None
        if ln is None and did and did in by_detail:
            upsert = by_detail[did]
        elif ln is None:
            errors.append(f"{where}: Paraşüt satırı veritabanında yok"
                          + (" — kimliksiz döküm satırı yazılamaz" if key.startswith(KEY_JSON_NOID)
                             else " (gece okuması bekleniyor ya da --parasut-json JSON:API dökümüyle --commit)"))
            continue
        item_id = _int(a.get("item_id"))
        if not item_id:
            errors.append(f"{where}: '{CARD_COL}' boş ya da geçersiz")
            continue
        item = (db.query(Item).filter(Item.id == item_id, Item.domain == domain, Item.is_active == True)   # noqa: E712
                .first())
        if item is None or item.category == "Bitmiş Ürün":
            errors.append(f"{where}: kart {item_id} bu panelde aktif hammadde değil")
            continue
        if ln is not None and ln.status == "approved":
            if ln.chosen_item_id == item.id:
                done.append(key)
                continue
            errors.append(f"{where}: satır zaten {ln.chosen_item_id} no'lu karta onaylanmış (bu satırda {item.id})")
            continue
        if ln is not None and ln.status != "pending":
            errors.append(f"{where}: satır bekleyen değil ({ln.status}) — /fiyatlar kuyruğundan bakın")
            continue
        pq = _num(a.get("package"))
        if pq is None or pq <= 0:
            errors.append(f"{where}: '{PACK_COL}' şart (sıfırdan büyük; ör. 25 kg bidon → 25)")
            continue
        unit = (str(a.get("unit") or "")).strip().lower()
        if unit not in PRICE_UNITS or not price_unit_ok(item.unit, unit):
            errors.append(f"{where}: '{UNIT_COL}' «{a.get('unit') or ''}» «{item.name}» ({item.unit or 'birimsiz'}) "
                          f"kartına uymuyor")
            continue
        sid = _int(a.get("supplier_id"))
        if a.get("supplier_id") not in (None, "") and not sid:
            errors.append(f"{where}: '{SUP_COL}' geçersiz: {a.get('supplier_id')!r}")
            continue
        if sid:
            sup = db.query(Supplier).filter(Supplier.id == sid, Supplier.domain == domain).first()
            if sup is None:
                errors.append(f"{where}: tedarikçi {sid} bu panelde yok")
                continue
            if sup.is_active is False:
                errors.append(f"{where}: tedarikçi {sid} {sup.name} pasif")
                continue
        else:
            sname = ln.bill.supplier_name if ln is not None else upsert.get("supplier_name")
            m = PP.match_supplier(db, domain, sname)
            sid = m["id"] if m else None
        todo.append({"key": key, "line": ln, "detail_id": did, "item": item, "package": pq, "unit": unit,
                     "supplier_id": sid, "upsert_bill": upsert})
    selection: List[dict] = []
    for key, a in sorted(approvals.get("songul", {}).items(), key=lambda kv: kv[1]["row"]):
        if a["ok"] is not True:
            continue
        where = f"'{SHEETS[2]}' satır {a['row']} ({key})"
        iid, sid = _int(a.get("item_id")), _int(a.get("supplier_id"))
        if not iid or not sid:
            errors.append(f"{where}: '{CARD_COL}' / '{SUP_COL}' boş ya da geçersiz")
            continue
        selection.append({"id": key, "item_id": iid, "supplier_id": sid})
    songul_plan = None
    if selection:
        songul_plan = lab.plan(db, data, domain=domain)
        _, errs = lab._selected_price_records(songul_plan, selection)
        errors += errs
    for n, card, v in approvals.get("form_filled") or []:
        warnings.append(f"'{SHEETS[3]}' satır {n} (kart {card}): fiyat «{v}» dolu — bu sayfa betikçe yazılmaz; "
                        "veri dosyasına kayıt olarak ya da /fiyatlar elle fiyat formuyla girin")
    return {"parasut": todo, "parasut_done": done, "songul_selection": selection, "songul_plan": songul_plan,
            "errors": errors, "warnings": warnings, "applied": False}


def apply(db: Session, approvals: dict, *, data: Optional[dict] = None, domain: str = DOMAIN,
          parasut_bills: Optional[List[dict]] = None) -> dict:
    """TEK işlem: önce `lock_for_approval` (kartlar → satırlar, id sırası),
    Paraşüt E → `approve_line` (flush), Songül E → içe aktarma betiğinin
    `apply(record_selection)`'ı (commit eder); hata ya da kilit çakışması
    (`busy`) → rollback, hiçbir şey yazılmaz.  Audit
    `supplier_prices.fiyat_tamamlama`."""
    data = data or lab.load_data()
    res = resolve(db, approvals, data=data, domain=domain, parasut_bills=parasut_bills)
    if res["errors"]:
        db.rollback()
        return res
    written: List[dict] = []
    songul_written = songul_deleted = 0
    try:
        # Tüm kartlar, sonra tüm satırlar id sırasıyla — /fiyatlar'dan aynı anda
        # gelen tekil onayla kilitlenme döngüsü kurulmasın (bkz. lock_for_approval)
        PP.lock_for_approval(db, [t["item"].id for t in res["parasut"]]
                             + [s["item_id"] for s in res["songul_selection"]],
                             [t["line"].id for t in res["parasut"] if t["line"] is not None])
        for t in res["parasut"]:
            ln = t["line"]
            if ln is None and t["upsert_bill"] is not None:
                PP.upsert_bill(db, t["upsert_bill"])
                ln = (db.query(ParasutPurchaseLine)
                      .filter(ParasutPurchaseLine.parasut_detail_id == t["detail_id"]).first())
                if ln is None:
                    raise PP.ApproveError("not_found", f"{t['key']}: döküm satırı tablolara yazılamadı")
            out = PP.approve_line(db, ln.id, item_id=t["item"].id, supplier_id=t["supplier_id"], supplier_name=None,
                                  package_qty=t["package"], price_unit=t["unit"], domain=domain, actor=ACTOR)
            written.append({"key": t["key"], **out["audit"]})
        if res["songul_selection"]:
            p2 = lab.apply(db, data, domain=domain, record_selection=res["songul_selection"])
            if not p2.get("applied"):                                  # içe aktarma betiği rollback etti
                res["errors"] = list(p2.get("errors") or ["içe aktarma betiği yazmadı"])
                return res
            songul_written, songul_deleted = p2["written"], p2["deleted"]
        else:
            db.commit()
    except PP.ApproveError as e:
        db.rollback()
        res["errors"] = [f"{e.code}: {e}"]
        return res
    except OperationalError as e:
        db.rollback()
        if not PP.is_lock_busy(e):
            raise
        res["errors"] = [f"busy: {PP.BUSY_TEXT} (hiçbir şey yazılmadı; betiği yeniden çalıştırın)"]
        return res
    except Exception:
        db.rollback()
        raise
    log_admin_event(
        db, None, actor=None, action=AUDIT_ACTION, target_type="supplier_prices", target_name=TITLE,
        details={"actor": ACTOR, "script": Path(__file__).name, "domain": domain,
                 "parasut_written": written, "parasut_already": res["parasut_done"],
                 "songul_selection": res["songul_selection"], "songul_written": songul_written,
                 "songul_deleted_previous": songul_deleted, "warnings": res["warnings"]})
    res.update(applied=True, parasut_written=written, songul_written=songul_written,
               songul_deleted=songul_deleted)
    return res


# ─── Çıktı ─────────────────────────────────────────────────────────────────

def _print_plan(plan: dict) -> None:
    s = plan["summary"]
    print(f"=== {TITLE} — kuru çalıştırma ({plan['domain']}) ===")
    print(f"  Hedef kart (aktif reçetede fiyatsız hammadde): {s['targets']}"
          f"  · yalnız grup üzerinden fiyatlı: {s['group_priced']}")
    print(f"  Paraşüt [{s['parasut_source']}]: bekleyen {s['parasut_pending_total']}"
          + (f" · döküm {s['parasut_json_lines']} (DB'de olan {s['parasut_json_in_db']} atlandı)"
             if s["parasut_json_lines"] or s["parasut_json_in_db"] else "")
          + f" → aday satır {s['parasut_rows']} (kesin kart {s['parasut_exact_cards']} · belirsiz kart "
            f"{s['parasut_ambig_cards']} · eski fatura {s['parasut_stale_rows']}) · hedef dışı "
            f"{s['parasut_off_target']} (gider ipuçlu {s['parasut_expense']})")
    print(f"  Songül notları: {s['songul_records']} kayıt → hedefle ilgili {s['songul_rows']} (yazılabilir "
          f"{s['songul_writable_rows']} kayıt / {s['songul_writable_cards']} kart · ezilen {s['songul_overridden']}"
          f" · belirsiz {s['songul_ambig']} · diğer {s['songul_other']})")
    print(f"  Karşılanan kart: {s['covered']}  · fiyat formuna kalan: {s['form']}")
    for r in plan["parasut"]["rows"]:
        print(f"    P {r['key']:<10} {r['match']:<16} {_card_text(r['card'])[:40]:<40} ← {r['invoice_no'] or '—'} "
              f"{(r['product_name'] or r['description'] or '')[:38]}" + ("  [eski]" if r["stale"] else ""))
    for r in plan["songul"]["rows"]:
        print(f"    S {r['key']:<10} {r['status']:<14} {_card_text(r['card'])[:40]:<40} ← {r['material'][:24]} / "
              f"{r['supplier'][:18]} {_num_tr(r['price'])} {r['currency']}/{r['price_unit']}")
    for t in plan["form"]:
        print(f"    F {t['item_id']:<5} {t['name'][:40]:<40} {t['unit'] or '—':<6} reçete {t['recipes']:<3} "
              f"son alım: {t['last_supplier'] or '—'}" + (f"  · {t['note'][:60]}" if t["note"] else ""))


def _print_result(res: dict, *, commit: bool) -> None:
    print("\n=== Onaylı Excel " + ("— YAZIM" if commit else "— kuru denetim") + " ===")
    print(f"  Paraşüt E satırı : {len(res['parasut'])} yazılacak · zaten yazılmış {len(res['parasut_done'])}")
    print(f"  Songül E kaydı   : {len(res['songul_selection'])}")
    for w in res["warnings"]:
        print(f"  ⚠ {w}")
    if res["errors"]:
        print("  ⛔ HATA — hiçbir şey yazılmadı:")
        for e in res["errors"]:
            print(f"     · {e}")
        return
    if commit and res.get("applied"):
        print(f"  ✓ yazıldı: Paraşüt {len(res['parasut_written'])} fiyat · Songül {res['songul_written']} fiyat "
              f"(önceki lab satırı silinen {res['songul_deleted']})")
        for w in res["parasut_written"]:
            print(f"     · {w['key']} → {w['item_id']} {w['item']}: {_num_tr(w['unit_price'])} {w['currency']}/"
                  f"{w['price_unit']}" + ("  (daha yeni fatura fiyatı korundu)" if w["kept_newer"] else ""))
    elif not commit:
        print("  (kuru — --commit ile yazılır)")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=f"{TITLE} — Paraşüt → Songül notları → fiyat formu")
    ap.add_argument("--commit", action="store_true", help="yaz (varsayılan: kuru çalıştırma)")
    ap.add_argument("--onayli", metavar="ONAYLI_XLSX", help="onaylı çalışma kitabı (--commit için ŞART)")
    ap.add_argument("--xlsx", metavar="YOL", help="çalışma kitabını yaz")
    ap.add_argument("--parasut-json", metavar="DOSYA",
                    help="Paraşüt dökümü (JSON:API yükü ya da parasut_read.py düz dökümü) — gece okuması yoksa")
    ap.add_argument("--data", metavar="JSON", help="lab veri dosyası (varsayılan scripts/data/…)")
    args = ap.parse_args(sys.argv[1:] if argv is None else argv)
    if args.commit and not args.onayli:
        print("⛔ Yazım (--commit) onaylı çalışma kitabı ister: --onayli ONAYLI.xlsx")
        return 2
    try:
        data = lab.load_data(args.data)
        bills, kind = load_parasut_json(args.parasut_json) if args.parasut_json else (None, None)
    except (OSError, ValueError) as e:
        print(f"⛔ veri dosyası okunamadı: {e}")
        return 2
    approvals = None
    if args.onayli:
        try:
            approvals = read_workbook(args.onayli)
        except Exception as e:                           # bozuk / yanlış dosya
            print(f"⛔ Excel okunamadı: {e}")
            return 2
    db = SessionLocal()
    try:
        plan = build_plan(db, data=data, parasut_bills=bills, parasut_kind=kind)
        _print_plan(plan)
        if args.xlsx:
            print(f"\nÇalışma kitabı yazıldı: {build_xlsx(plan, args.xlsx)}")
        if approvals is not None:
            jb = bills if kind == "jsonapi" else None
            res = (apply(db, approvals, data=data, parasut_bills=jb) if args.commit
                   else resolve(db, approvals, data=data, parasut_bills=jb))
            _print_result(res, commit=args.commit)
            if res["errors"]:
                return 2
    except ValueError as e:                              # veri dosyası / eşleme hatası
        db.rollback()
        print(f"⛔ DURDURULDU — {e}; hiçbir şey yazılmadı")
        return 2
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
